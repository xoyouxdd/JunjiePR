"""V3 employee workbook preflight. No writes until the caller confirms the exact batch."""
import hashlib
import json
from io import BytesIO
from app.employee_numbers import employee_number_cell, employee_number_aliases

from fastapi import HTTPException
from openpyxl import load_workbook

from app.services.employee_import_dates import normalize_import_date
from app.services.employee_types import normalize, is_due, business_today
from app.services.identity import role_at
from app.v2_models import Employee, UserAccount, Role, Attraction

HEADERS = ['员工号', '姓名', '角色代码', '景点圈', 'Type', '实习结束日期']
ROLE_MAP = {'CM':'CM', 'TR':'TR', 'LEAD':'SUPERVISOR', 'GSM':'GSM', 'AM':'AM', 'OM':'OM'}


def selection_map(decisions):
    if decisions is None:
        return {}
    if not isinstance(decisions, list) or len(decisions) > 10000:
        raise HTTPException(400, '导入选择格式不正确')
    choices = {}
    fields = {'row', 'action'}
    for choice in decisions:
        if not isinstance(choice, dict) or set(choice) - fields or type(choice.get('row')) is not int:
            raise HTTPException(400, '导入选择格式不正确')
        row = choice['row']
        if row in choices or choice.get('action') not in ('skip', 'update'):
            raise HTTPException(400, '导入选择重复或操作不正确')
        choices[row] = choice
    return choices


def replacement(old, raw):
    # Overwrite the two supported fields with the workbook, including blank cells.
    kind, _ = normalize(raw['Type'], None)
    end = normalize_import_date(raw['实习结束日期']) if kind == 'Intern' else None
    kind, end = normalize(kind, end)
    return {'name': old.name, 'employee_type': kind, 'internship_ends_on': end}


def preflight(db, user, content, allowed_circles, allowed_roles, decisions=None):
    choices=selection_map(decisions)
    try:
        wb=load_workbook(BytesIO(content),read_only=True,data_only=True)
    except Exception as exc:
        raise HTTPException(400,'Excel文件无法读取') from exc
    try:
        if not wb.worksheets:
            raise HTTPException(400,'Excel文件没有员工工作表')
        sheet=wb.worksheets[0]
        header=[str(v or '').strip() for v in next(sheet.iter_rows(min_row=1,max_row=1,values_only=True))]
        while header and not header[-1]: header.pop()
        if len(header)!=len(HEADERS) or len(set(header))!=len(header) or set(header)!=set(HEADERS):
            raise HTTPException(400,'表头与V3模板不一致，请重新下载六列模板；不再导入密码、在职、账号启用或任职期限列')
        positions={name:header.index(name) for name in HEADERS}
        roles={r.code:r for r in db.query(Role).all()}
        circles={c.name:c for c in db.query(Attraction).filter(Attraction.active.is_(True),Attraction.employee_circle.is_(True)).all()}
        parsed=[];duplicates=[];expired=[];errors=[];seen=set();updates=[];occupied=[];duplicate_rows=set();unchanged=0
        for rownum,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
            if rownum>10001: raise HTTPException(400,'单次最多导入10000行，请拆分文件')
            values=list(values)+[None]*max(0,len(header)-len(values))
            raw={name:values[i] for name,i in positions.items()}
            if not any(v is not None and str(v).strip() for v in raw.values()): continue
            no=employee_number_cell(raw['员工号']);name=str(raw['姓名'] or '').strip()
            role=ROLE_MAP.get(str(raw['角色代码'] or '').strip().upper())
            circle=circles.get(str(raw['景点圈'] or '').strip())
            prefix=f'第{rownum}行：'
            if len(no)!=7 or not no.isascii() or not no.isdigit(): errors.append(prefix+'员工号必须为7位纯数字，可在前面加一个0');continue
            if no in seen: errors.append(prefix+f'文件内部员工号{no}重复，请先修正');continue
            seen.add(no)
            if not name or len(name)>100: errors.append(prefix+'姓名必填，且不能超过100个字符');continue
            if not role or role not in roles: errors.append(prefix+'角色仅支持CM、TR、Lead、GSM、AM、OM；TA职务请在页面设置');continue
            if not circle:
                errors.append(prefix+'请选择有效景点圈');continue
            aliases=employee_number_aliases(no)
            employees=db.query(Employee).filter(Employee.employee_no.in_(aliases)).limit(2).all()
            accounts=db.query(UserAccount).filter(UserAccount.login_account.in_(aliases)).limit(2).all()
            if len(employees)>1 or len(accounts)>1:
                errors.append(prefix+'员工号存在关联冲突，请先核对账号档案');continue
            old=employees[0] if employees else None
            account=accounts[0] if accounts else None
            if old or account:
                duplicate_rows.add(rownum)
                choice=choices.get(rownum, {'row':rownum,'action':'skip'})
                same_circle=old is not None and old.attraction_id==circle.id
                in_scope=old is not None and (allowed_circles is None or old.attraction_id in allowed_circles)
                old_role=role_at(db,old.id) if old else None
                visible=in_scope and (old_role is None or old_role.code in allowed_roles)
                cross=old is not None and (not same_circle or not in_scope)
                can_update=bool(same_circle and visible and not old.account_deleted_at
                                and (account is None or account.employee_id==old.id) and old.name==name)
                item={'row':rownum,'employee_no':no,'name':'信息受限','excel_name':name,
                      'can_update':can_update,'name_conflict':False,'selected_action':'skip'}
                if cross:
                    home=db.get(Attraction,old.attraction_id) if old.attraction_id else None
                    item.update(name=old.name,attraction_name=home.name if home else '未设置景点圈',
                                reason='员工已存在，所属景点圈与本次可导入范围或目标圈不同；仅可跳过，如需变更请通过跨圈调动更新景点圈')
                elif visible:
                    item.update(name=old.name,id=old.id,name_conflict=old.name!=name,
                                current={'name':old.name,'employee_type':old.employee_type,'internship_ends_on':old.internship_ends_on},
                                incoming={'name':name,'employee_type':None if raw['Type'] is None else str(raw['Type']).strip(),
                                          'internship_ends_on':None if raw['实习结束日期'] is None else str(raw['实习结束日期'])},
                                reason='可覆盖为文件中的Type和实习结束日期；角色和景点圈不变')
                    if old.account_deleted_at:
                        item['reason']='已删除留档，本次跳过；重新入职请使用新建员工'
                    elif account is not None and account.employee_id!=old.id:
                        item['reason']='员工号关联账号冲突，仅可跳过，请先核对账号档案'
                    elif old.name!=name:
                        item['reason']='员工号一致、姓名不一致，仅可跳过；请到系统员工管理中核对并修改姓名后重新导入'
                else:
                    item['reason']='员工号已存在，当前HR无权查看该档案；本次跳过，不修改原信息'
                # Bind confirmation to current ownership, role, values and account state,
                # without returning restricted data to the client.
                own_account=db.query(UserAccount).filter_by(employee_id=old.id).first() if old else None
                occupied.append({'row':rownum,'id':old.id if old else None,
                    'name':old.name if old else None,'circle':old.attraction_id if old else None,
                    'role':old_role.code if old_role else None,'type':old.employee_type if old else None,
                    'end':old.internship_ends_on if old else None,'active':old.is_active if old else None,
                    'deleted':str(old.account_deleted_at) if old else None,
                    'account':None if account is None else [account.id,account.employee_id,account.enabled],
                    'own_account':None if own_account is None else [own_account.id,own_account.enabled,str(own_account.disabled_at)]})
                if can_update:
                    try:
                        after=replacement(old,raw)
                    except (ValueError,HTTPException) as exc:
                        raise HTTPException(400,prefix+str(exc.detail if isinstance(exc,HTTPException) else exc)) from exc
                    before={'name':old.name,'employee_type':old.employee_type,'internship_ends_on':old.internship_ends_on}
                    changes={key:{'before':before[key],'after':value} for key,value in after.items() if before[key]!=value}
                    item.update(after=after,changes=changes,unchanged=not bool(changes))
                    if not changes:
                        item['selected_action']='unchanged'
                        unchanged+=1
                    elif choice['action']=='update':
                        item['selected_action']='update'
                        updates.append({'row':rownum,'employee_id':old.id,'employee_no':no,'before':before,'after':after,'changes':changes})
                        if is_due(after['employee_type'],after['internship_ends_on']):
                            expired.append({'row':rownum,'employee_no':no,'name':after['name'],
                                            'internship_ends_on':after['internship_ends_on'],'action':'update'})
                elif choice['action']=='update':
                    raise HTTPException(403,prefix+item['reason'])
                duplicates.append(item)
                continue
            if role not in allowed_roles:
                errors.append(prefix+'当前HR权限仅可导入Lead及以下员工');continue
            if allowed_circles is not None and circle.id not in allowed_circles:
                errors.append(prefix+'请选择权限范围内的有效景点圈');continue
            try:
                from app.rotation.temporary import pending
                pending(db,no,name,circle.id)
                kind,end=normalize(raw['Type'],normalize_import_date(raw['实习结束日期']))
            except (ValueError,HTTPException) as exc:
                errors.append(prefix+str(exc.detail if isinstance(exc,HTTPException) else exc));continue
            row={'row':rownum,'employee_no':no,'name':name,'role_code':role,'attraction_id':circle.id,
                 'employee_type':kind,'internship_ends_on':end}
            parsed.append(row)
            if is_due(kind,end): expired.append({'row':rownum,'employee_no':no,'name':name,'internship_ends_on':end})
        if set(choices)-duplicate_rows:raise HTTPException(400,'导入选择包含不存在或非已有员工的行，请重新预检')
        if errors: raise HTTPException(400,'；'.join(errors[:20]))
        if not seen: raise HTTPException(400,'文件中没有员工数据')
        digest=hashlib.sha256(content+json.dumps({'user':user.id,'date':business_today().isoformat(),
            'rows':parsed,'duplicates':duplicates,'occupied':occupied,'decisions':choices,
            'allowed_circles':None if allowed_circles is None else sorted(allowed_circles),'allowed_roles':sorted(allowed_roles)},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
        return {'rows':parsed,'duplicates':duplicates,'updates':updates,'expired':expired,'digest':digest,
                'would_create':len(parsed),'would_update':len(updates),'would_rename':sum('name' in r['changes'] for r in updates),
                'unchanged':unchanged,'would_skip':sum(r['selected_action']=='skip' for r in duplicates)}
    finally:
        wb.close()

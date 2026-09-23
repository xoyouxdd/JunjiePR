"""Monthly sick-leave import from the HR transaction workbook."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
import secrets
from threading import Lock

import xlrd
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.v2_auth import V2User, require_permissions
from app.v2_database import get_db
from app.v2_models import Employee, EmployeeLOAPeriod, SickLeaveRecord
from app.v2_services import recalculate_attendance, role_at, write_audit
from app.routers._shared import ensure_month_open, ensure_scoped_hr_employee, invalidate_data_caches, parse_iso_date


router = APIRouter()
_PREVIEWS: dict[str, dict] = {}
_PREVIEWS_LOCK = Lock()
_MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
_MAX_CACHED_PREVIEWS = 8
_MAX_CACHED_BYTES = 40 * 1024 * 1024
_PREVIEW_TTL = timedelta(minutes=20)
_SICK_TYPES = {"法定病假", "全薪病假", "无薪病假"}
_SOURCE_ID_LENGTH = 8
_TRANSACTION_MARKERS = ("事务:", "事务：")
_SUMMARY_MARKERS = ("总数:", "总数：")
_PREVIEW_EXPIRED_MESSAGE = "预检结果已失效（超时、服务重启或被较新的预检挤出），请重新上传文件预检"


def _prune_previews(now: datetime) -> None:
    """Caller holds _PREVIEWS_LOCK; retain only recent, bounded previews."""
    for token, preview in list(_PREVIEWS.items()):
        if now - preview["created_at"] > _PREVIEW_TTL:
            _PREVIEWS.pop(token, None)
    while _PREVIEWS and (
        len(_PREVIEWS) > _MAX_CACHED_PREVIEWS
        or sum(len(item["content"]) for item in _PREVIEWS.values()) > _MAX_CACHED_BYTES
    ):
        oldest = min(_PREVIEWS, key=lambda token: _PREVIEWS[token]["created_at"])
        _PREVIEWS.pop(oldest, None)


def _store_preview(token: str, preview: dict) -> None:
    with _PREVIEWS_LOCK:
        _prune_previews(datetime.now())
        _PREVIEWS[token] = preview
        _prune_previews(datetime.now())


def _get_preview(token: str, user_id: int) -> dict | None:
    with _PREVIEWS_LOCK:
        _prune_previews(datetime.now())
        preview = _PREVIEWS.get(token)
        return preview if preview and preview["user_id"] == user_id else None


def _drop_preview(token: str) -> None:
    with _PREVIEWS_LOCK:
        _PREVIEWS.pop(token, None)


def _cell(row: list, index: int) -> str:
    return str(row[index]).strip() if index < len(row) and row[index] not in (None, "") else ""


def _find_row(sheet, labels: tuple[str, ...]) -> int:
    """Locate a block marker; half- and full-width colons are both accepted."""
    for index in range(sheet.nrows):
        if any(_cell(sheet.row_values(index), column) in labels for column in range(sheet.ncols)):
            return index
    raise HTTPException(400, f"未找到“{labels[0]}”区块")


def _header_index(row: list, name: str) -> int:
    for index, value in enumerate(row):
        if str(value).strip().rstrip(":：") == name:
            return index
    raise HTTPException(400, f"导入文件缺少“{name}”列")


def _date_value(book, value) -> str:
    try:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return xlrd.xldate_as_datetime(value, book.datemode).date().isoformat()
        return date.fromisoformat(str(value).strip().replace("/", "-")).isoformat()
    except (ValueError, OverflowError) as exc:
        raise ValueError(f"事务日期无效：{value}") from exc


def _hours_value(value) -> Decimal:
    """Parse an hours cell; xlrd returns numeric cells as float (e.g. 8.0)."""
    raw = repr(value) if isinstance(value, float) else str(value).strip()
    try:
        return Decimal(raw).quantize(Decimal("0.01"))
    except ArithmeticError as exc:
        raise ValueError(f"时数无效：{value}") from exc


def _chinese_name(value: str) -> str:
    return value.split(",")[-1].strip().replace("，", "").strip()


def _source_id(value) -> str:
    """Normalize an ID cell to the 8-digit source ID.

    Numeric-formatted cells come back from xlrd as float (1727264.0) and have
    lost their leading zeros, so they are converted to int text and left-padded.
    Text cells are taken as-is and validated by _system_number.
    """
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"源文件 ID 必须为 {_SOURCE_ID_LENGTH} 位数字：{value}")
        value = int(value)
    if isinstance(value, int):
        return str(value).zfill(_SOURCE_ID_LENGTH) if value >= 0 else str(value)
    return str(value).strip() if value is not None else ""


def _system_number(source_id: str) -> str:
    value = str(source_id).strip()
    if len(value) != _SOURCE_ID_LENGTH or not value.isdigit():
        raise ValueError(f"源文件 ID 必须为 {_SOURCE_ID_LENGTH} 位数字：{value or '空'}")
    return value[1:]


def _read_workbook(content: bytes) -> tuple[list[dict], dict[tuple[str, str], Decimal], list[str]]:
    try:
        book = xlrd.open_workbook(file_contents=content)
    except xlrd.biffh.XLRDError as exc:
        raise HTTPException(400, "仅支持 Excel 97-2003 格式（.xls）文件") from exc
    sheet = book.sheet_by_index(0)
    transaction_marker = _find_row(sheet, _TRANSACTION_MARKERS)
    summary_marker = _find_row(sheet, _SUMMARY_MARKERS)
    transaction_header = sheet.row_values(transaction_marker + 2)
    summary_header = sheet.row_values(summary_marker + 3)
    t_name, t_id, t_date, t_type, t_hours = (_header_index(transaction_header, label) for label in ("员工", "ID", "日期", "工资代码", "时数"))
    s_name, s_id, s_type, s_hours = (_header_index(summary_header, label) for label in ("员工", "ID", "工资代码", "时数"))
    records: list[dict] = []
    errors: list[str] = []
    for row_index in range(transaction_marker + 3, summary_marker):
        row = sheet.row_values(row_index)
        leave_type = _cell(row, t_type)
        if leave_type not in _SICK_TYPES:
            continue
        source_name = _cell(row, t_name)
        try:
            source_id = _source_id(row[t_id])
            hours = _hours_value(row[t_hours])
            employee_no = _system_number(source_id)
            leave_date = _date_value(book, row[t_date])
        except (ValueError, IndexError) as exc:
            errors.append(f"事务区第 {row_index + 1} 行：{exc}")
            continue
        if hours <= 0 or hours % 4 != 0:
            errors.append(f"事务区第 {row_index + 1} 行：时数必须是大于 0 的 4 小时倍数")
            continue
        records.append({
            "row": row_index + 1, "source_id": source_id, "employee_no": employee_no,
            "source_name": source_name, "name": _chinese_name(source_name), "date": leave_date,
            "leave_type": leave_type, "hours": hours, "days": hours / Decimal("8"),
        })
    summaries: dict[tuple[str, str], Decimal] = {}
    for row_index in range(summary_marker + 4, sheet.nrows):
        row = sheet.row_values(row_index)
        summary_type = _cell(row, s_type)
        if summary_type not in {"病假时间总计", "无薪病假"}:
            continue
        try:
            # Same normalization as the transaction block so _reconcile keys match.
            source_id = _source_id(row[s_id])
            _system_number(source_id)
            hours = _hours_value(row[s_hours])
        except (ValueError, IndexError) as exc:
            errors.append(f"总数区第 {row_index + 1} 行：{exc}")
            continue
        summaries[(source_id, summary_type)] = hours
    return records, summaries, errors


def _reconcile(records: list[dict], summaries: dict[tuple[str, str], Decimal]) -> list[str]:
    totals: dict[tuple[str, str], Decimal] = {}
    for row in records:
        summary_type = "无薪病假" if row["leave_type"] == "无薪病假" else "病假时间总计"
        key = (row["source_id"], summary_type)
        totals[key] = totals.get(key, Decimal("0")) + row["hours"]
    errors = []
    for key in set(totals) | set(summaries):
        if totals.get(key, Decimal("0")) != summaries.get(key, Decimal("0")):
            errors.append(f"总数区对账不一致：ID {key[0]} 的 {key[1]}，事务区 {totals.get(key, 0)} 小时，总数区 {summaries.get(key, 0)} 小时")
    return errors


@router.post("/sick-leave-imports/preview")
async def preview_sick_leave_import(
    workbook: UploadFile = File(...), db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SICK_LEAVE_IMPORT")),
):
    if not (workbook.filename or "").lower().endswith(".xls"):
        raise HTTPException(400, "请上传 .xls 格式的员工事务文件")
    content = await workbook.read(_MAX_WORKBOOK_BYTES + 1)
    if len(content) > _MAX_WORKBOOK_BYTES:
        raise HTTPException(413, "员工事务文件不能超过10MB")
    records, summaries, errors = _read_workbook(content)
    if not records:
        errors.append("事务区未找到法定病假、全薪病假或无薪病假记录")
    errors.extend(_reconcile(records, summaries))
    months = {row["date"][:7] for row in records}
    if len(months) != 1:
        errors.append("一次导入只能包含同一个自然月的病假事务")
    unmatched, matched, loa_protected = [], [], []
    for row in records:
        employee = db.query(Employee).filter(Employee.employee_no == row["employee_no"]).first()
        reason = ""
        if not employee:
            reason = "系统中不存在该员工号"
        elif employee.name != row["name"]:
            reason = f"姓名不匹配（系统：{employee.name}）"
        elif not employee.is_active:
            reason = "员工当前不在职"
        else:
            try:
                ensure_scoped_hr_employee(db, user, employee)
            except HTTPException:
                reason = "不在当前账号可导入的景点圈范围"
        if reason:
            unmatched.append({**row, "reason": reason})
        else:
            period = db.query(EmployeeLOAPeriod).filter(
                EmployeeLOAPeriod.employee_id == employee.id,
                EmployeeLOAPeriod.status != "cancelled",
                EmployeeLOAPeriod.starts_on <= row["date"],
                (EmployeeLOAPeriod.ends_on.is_(None) | (EmployeeLOAPeriod.ends_on >= row["date"])),
            ).first()
            if period:
                loa_protected.append({**row, "employee_id": employee.id, "reason": f"LOA保护：{period.starts_on} 至 {period.ends_on or '至今'}"})
            else:
                matched.append({**row, "employee_id": employee.id})
    token = secrets.token_urlsafe(24)
    _store_preview(token, {"created_at": datetime.now(), "filename": workbook.filename, "content": content, "month": next(iter(months), ""), "matched": matched, "unmatched": unmatched, "loa_protected": loa_protected, "errors": errors, "user_id": user.id})
    return {
        "token": token, "month": next(iter(months), ""), "blocking_errors": errors,
        "matched_employee_count": len({row["employee_id"] for row in matched}), "matched_record_count": len(matched),
        "unmatched": unmatched, "loa_protected": loa_protected, "can_commit": not errors and bool(matched),
    }


@router.post("/sick-leave-imports/commit")
def commit_sick_leave_import(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SICK_LEAVE_IMPORT"))):
    token = str(payload.get("token") or "")
    preview = _get_preview(token, user.id)
    if not preview:
        raise HTTPException(400, _PREVIEW_EXPIRED_MESSAGE)
    if preview["errors"] or not preview["matched"]:
        raise HTTPException(400, "预检未通过，不能覆盖")
    month = preview["month"]
    employee_ids = sorted({row["employee_id"] for row in preview["matched"]})
    employees = {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(employee_ids)).all()}
    if len(employees) != len(employee_ids) or any(not employee.is_active for employee in employees.values()):
        raise HTTPException(409, "员工资料已变更，请重新预检文件")
    try:
        loa_periods = db.query(EmployeeLOAPeriod).filter(
            EmployeeLOAPeriod.employee_id.in_(employee_ids),
            EmployeeLOAPeriod.status != "cancelled",
        ).all()
        loa_by_employee: dict[int, list[EmployeeLOAPeriod]] = {}
        for period in loa_periods:
            loa_by_employee.setdefault(period.employee_id, []).append(period)
        for item in preview["matched"]:
            if any(
                period.starts_on <= item["date"]
                and (period.ends_on is None or period.ends_on >= item["date"])
                for period in loa_by_employee.get(item["employee_id"], ())
            ):
                raise HTTPException(409, "预检后LOA记录已变化，请重新上传文件预检")
        for employee in employees.values():
            ensure_scoped_hr_employee(db, user, employee)
            ensure_month_open(db, month, employee.attraction_id, "覆盖月度病假")
            db.query(SickLeaveRecord).filter(
                SickLeaveRecord.employee_id == employee.id,
                SickLeaveRecord.attendance_month == month,
                SickLeaveRecord.status == "active",
                SickLeaveRecord.is_violation.is_(False),
                SickLeaveRecord.import_source == "monthly_transaction_import",
            ).update({"status": "void", "voided_by": user.id, "voided_by_name": user.name, "voided_at": datetime.now(), "void_reason": "月度病假事务文件覆盖"}, synchronize_session=False)
        for item in preview["matched"]:
            employee = employees[item["employee_id"]]
            role = role_at(db, employee.id)
            db.add(SickLeaveRecord(
                employee_id=employee.id, employee_no_snapshot=employee.employee_no, employee_name_snapshot=employee.name,
                employee_role_snapshot=role.name if role else "", attraction_id_snapshot=employee.attraction_id,
                attendance_month=month, leave_start_date=item["date"], leave_end_date=item["date"], leave_days=item["days"], charged_days=item["days"],
                proof_file_id=None, leave_type=item["leave_type"], import_source="monthly_transaction_import", status="active",
                submitted_by=user.id, submitted_by_name=user.name, is_violation=False,
            ))
        db.flush()
        for employee in employees.values():
            recalculate_attendance(db, employee, month)
        write_audit(db, user.employee, "覆盖月度病假事务", "sick_leave_import", 0, after={"month": month, "file": preview["filename"], "employees": len(employees), "records": len(preview["matched"]), "unmatched": len(preview["unmatched"]), "loa_protected": len(preview["loa_protected"])})
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    _drop_preview(token)
    return {"ok": True, "month": month, "covered_employee_count": len(employees), "covered_record_count": len(preview["matched"]), "unmatched_count": len(preview["unmatched"]), "loa_protected_count": len(preview["loa_protected"])}


@router.get("/sick-leave-imports/{token}/unmatched-file")
def download_unmatched_sick_leave_file(token: str, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SICK_LEAVE_IMPORT"))):
    preview = _get_preview(token, user.id)
    if not preview:
        raise HTTPException(404, _PREVIEW_EXPIRED_MESSAGE)
    try:
        source = xlrd.open_workbook(file_contents=preview["content"])
    except (KeyError, xlrd.biffh.XLRDError) as exc:
        raise HTTPException(404, "原始预检文件已不可用，请重新上传") from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    red = PatternFill("solid", fgColor="FFC7CE")
    unmatched_by_row = {item["row"]: item["reason"] for item in preview["unmatched"]}
    for sheet_number, source_sheet in enumerate(source.sheets()):
        sheet = workbook.create_sheet(title=source_sheet.name[:31] or "Sheet")
        for row_index in range(source_sheet.nrows):
            values = source_sheet.row_values(row_index)
            sheet.append(values)
            if sheet_number == 0 and row_index + 1 in unmatched_by_row:
                for cell in sheet[sheet.max_row]:
                    cell.fill = red
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)

    sheet = workbook.create_sheet("未覆盖说明")
    sheet.append(["文件行号", "源文件姓名", "源文件ID", "转换后系统员工号", "日期", "病假类型", "时数", "换算天数", "未覆盖原因"])
    for item in preview["unmatched"]:
        sheet.append([item["row"], item["source_name"], item["source_id"], item["employee_no"], item["date"], item["leave_type"], float(item["hours"]), float(item["days"]), item["reason"]])
        for cell in sheet[sheet.max_row]:
            cell.fill = red
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=\"sick-leave-unmatched.xlsx\""})

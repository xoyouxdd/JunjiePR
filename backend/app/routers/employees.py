"""Employee master data, employee-number and import endpoints."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.v2_auth import V2User, current_user, require_permissions
from app.v2_crypto import default_initial_password, hash_password
from app.v2_database import get_db, synchronize_gsm_management_scope
from app.v2_models import Attraction, Employee, EmployeeNumberHistory, EmployeeLOAPeriod, EmployeeRoleAssignment, GroupLeaderAssignment, GroupMembership, RecognitionRecord, Role, SystemAlert, UserAccount, UserSession, WorkGroup
from app.v2_services import FRONTLINE_CODES, LEADER_CODES, active_group_leader, active_group_memberships, groups_led_by, recalculate_attendance, role_at, write_audit
from app.v2_watermark import watermark_workbook
from app.excel_export import build_employee_import_template
from app.routers._shared import (
    CIRCLE_HR_MANAGED_ROLE_CODES,
    EMPLOYEE_TARGET_PERMISSIONS,
    SCOPED_HR_ROLE_CODE,
    client_ip,
    employee_payloads,
    ensure_month_open,
    ensure_scoped_hr_attraction,
    ensure_scoped_hr_employee,
    invalidate_data_caches,
    like_escaped_pattern,
    months_between,
    parse_iso_date,
    scoped_hr_attraction_ids,
    search_employee_targets,
)

router = APIRouter()


def ensure_employee_number_change_target(db: Session, user: V2User, employee: Employee, role: Role | None = None) -> Role:
    """Apply the employee-number migration scope without changing org data."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以变更员工号")
    target_role = role or role_at(db, employee.id)
    if not target_role:
        raise HTTPException(400, "该员工当前没有有效角色")
    if target_role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR和最高管理员账号不能在此处变更员工号")
    ensure_scoped_hr_employee(db, user, employee)
    if user.role.code == SCOPED_HR_ROLE_CODE and target_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能变更本圈CM、TR、TA主管或主管的员工号")
    return target_role


def ensure_hr_role_allowed(user: V2User, role_code: str, label: str = "角色") -> None:
    if "SYSTEM_ADMIN" not in user.permissions and role_code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, f"景点圈HR只能设置{label}为CM、TR、TA主管或主管")


def employee_payload(db: Session, employee: Employee) -> dict:
    return employee_payloads(db, [employee])[employee.id]


def employee_circle_id(db: Session, value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        attraction_id = int(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "员工景点圈无效") from exc
    circle = db.get(Attraction, attraction_id)
    if not circle or not circle.active or not circle.employee_circle:
        raise HTTPException(400, "员工只能归属有效景点圈")
    return circle.id


@router.get("/frontline-employees")
def frontline_employees(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if not ({"EMPLOYEE_ADD", "SICK_REGISTER", "DEDUCTION_ALL", "DEDUCTION_DIRECT"} & user.permissions):
        raise HTTPException(403, "没有权限查询CM/TR")
    allowed = scoped_hr_attraction_ids(db, user)
    attraction_id = next(iter(allowed)) if allowed else None
    return search_employee_targets(db, attraction_id=attraction_id, limit=500)["items"]


@router.get("/employee-targets")
def employee_targets(
    usage: str,
    keyword: str = "",
    attraction_id: int | None = None,
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    required = EMPLOYEE_TARGET_PERMISSIONS.get(usage)
    if not required:
        raise HTTPException(400, "员工查询用途无效")
    if not (required & user.permissions):
        raise HTTPException(403, "没有对应的员工登记权限")
    if usage == "poc":
        if user.role.code not in {"TA_GSM", "GSM", "AM"}:
            raise HTTPException(403, "仅TA GSM、GSM、AM可以查询POC被认可员工")
        value = like_escaped_pattern(keyword)
        query = db.query(Employee).join(UserAccount, UserAccount.employee_id == Employee.id).filter(
            Employee.is_active.is_(True), UserAccount.enabled.is_(True),
            or_(Employee.name.like(value, escape="\\"), Employee.employee_no.like(value, escape="\\")),
        )
        if attraction_id is not None:
            query = query.filter(Employee.attraction_id == attraction_id)
        rows = []
        for item in query.order_by(Employee.name, Employee.employee_no).limit(max(1, min(limit, 50)) * 4).all():
            role = role_at(db, item.id)
            if role and role.code in (FRONTLINE_CODES | LEADER_CODES):
                circle = db.get(Attraction, item.attraction_id) if item.attraction_id else None
                rows.append({"id": item.id, "employee_no": item.employee_no, "name": item.name, "role_code": role.code, "role_name": role.name, "attraction_id": item.attraction_id, "attraction_name": circle.name if circle else "", "group_id": None, "group_name": ""})
        return {"items": rows[:max(1, min(limit, 50))], "total": len(rows), "limit": max(1, min(limit, 50))}
    if usage == "loa":
        # LOA registration deliberately permits its three authorized roles to
        # locate every active employee.  This expands lookup only, not any
        # other HR management scope.
        if user.role.code not in {"TA_GSM", "GSM", "HR_CIRCLE"}:
            raise HTTPException(403, "仅TA GSM、GSM、景点圈HR可以登记LOA")
        if not keyword.strip():
            return {"items": [], "total": 0, "limit": max(1, min(limit, 50)), "search_scope": "全部在职员工（请输入姓名或员工号）"}
        value = like_escaped_pattern(keyword)
        query = db.query(Employee).filter(
            Employee.is_active.is_(True),
            or_(Employee.name.like(value, escape="\\"), Employee.employee_no.like(value, escape="\\")),
        )
        if attraction_id is not None:
            query = query.filter(Employee.attraction_id == attraction_id)
        rows = []
        for item in query.order_by(Employee.name, Employee.employee_no).limit(max(1, min(limit, 50)) * 2).all():
            role = role_at(db, item.id)
            circle = db.get(Attraction, item.attraction_id) if item.attraction_id else None
            active_loa = (
                db.query(EmployeeLOAPeriod)
                .filter(
                    EmployeeLOAPeriod.employee_id == item.id,
                    EmployeeLOAPeriod.status == "active",
                    or_(EmployeeLOAPeriod.ends_on.is_(None), EmployeeLOAPeriod.ends_on >= date.today().isoformat()),
                )
                .order_by(EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
                .first()
            )
            rows.append({"id": item.id, "employee_no": item.employee_no, "name": item.name, "role_code": role.code if role else "", "role_name": role.name if role else "未分配角色", "attraction_id": item.attraction_id, "attraction_name": circle.name if circle else "", "group_id": None, "group_name": "", "loa_active": bool(active_loa), "loa_period_id": active_loa.id if active_loa else None, "loa_starts_on": active_loa.starts_on if active_loa else "", "loa_ends_on": active_loa.ends_on if active_loa else ""})
        return {"items": rows[:max(1, min(limit, 50))], "total": len(rows), "limit": max(1, min(limit, 50)), "search_scope": "全部在职员工"}
    # Deduction/absence target searches are intentionally keyword-only for
    # cross-circle managers.  TA主管 retains its home-circle boundary, while
    # 主管、TA GSM、GSM may locate active CM/TR across all circles.
    global_target_search = usage in {"deduction", "attendance"} and user.role.code in {"SUPERVISOR", "TA_GSM", "GSM"}
    if global_target_search and not keyword.strip():
        return {"items": [], "total": 0, "limit": max(1, min(limit, 50)), "search_scope": "全部景点圈在职CM/TR（请输入姓名或员工号）"}
    if user.role.code == "TA_SUPERVISOR" and usage in {"deduction", "attendance"}:
        attraction_id = user.employee.attraction_id
    if attraction_id is not None:
        circle = db.get(Attraction, attraction_id)
        if not circle or not circle.active or not circle.employee_circle:
            raise HTTPException(400, "请选择有效员工景点圈")
    elif not global_target_search:
        allowed = scoped_hr_attraction_ids(db, user)
        attraction_id = next(iter(allowed)) if allowed else None
    if not global_target_search:
        ensure_scoped_hr_attraction(db, user, attraction_id)
    result = search_employee_targets(
        db,
        keyword=keyword,
        attraction_id=attraction_id,
        limit=max(1, min(limit, 50)),
    )
    result["search_scope"] = "全部景点圈在职CM/TR（请输入姓名或员工号）" if global_target_search else "本景点圈在职CM/TR"
    return result


def loa_period_payload(row: EmployeeLOAPeriod, employee: Employee | None = None) -> dict:
    return {
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": employee.employee_no if employee else "",
        "employee_name": employee.name if employee else "",
        "starts_on": row.starts_on,
        "ends_on": row.ends_on or "",
        "note": row.note or "",
        "status": row.status,
        "created_by_name": row.created_by_name,
        "created_at": row.created_at.strftime("%Y-%m-%d %H:%M:%S") if row.created_at else "",
    }


@router.get("/loa-periods")
def list_loa_periods(month: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("LOA_REGISTER"))):
    query = db.query(EmployeeLOAPeriod).filter(EmployeeLOAPeriod.status != "cancelled")
    if month:
        start = parse_iso_date(f"{month}-01", "月份")
        finish = start.replace(day=28) + timedelta(days=4)
        finish = finish - timedelta(days=finish.day)
        query = query.filter(EmployeeLOAPeriod.starts_on <= finish.isoformat(), or_(EmployeeLOAPeriod.ends_on.is_(None), EmployeeLOAPeriod.ends_on >= start.isoformat()))
    rows = query.order_by(EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc()).limit(200).all()
    employees = {row.id: row for row in db.query(Employee).filter(Employee.id.in_({row.employee_id for row in rows})).all()} if rows else {}
    return {"items": [loa_period_payload(row, employees.get(row.employee_id)) for row in rows]}


@router.post("/loa-periods")
def create_loa_period(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("LOA_REGISTER"))):
    if user.role.code not in {"TA_GSM", "GSM", "HR_CIRCLE"}:
        raise HTTPException(403, "仅TA GSM、GSM、景点圈HR可以登记LOA")
    try:
        employee_id = int(payload.get("employee_id") or 0)
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "请选择员工") from exc
    employee = db.get(Employee, employee_id)
    if not employee or not employee.is_active:
        raise HTTPException(400, "请选择在职员工")
    starts_on = parse_iso_date(str(payload.get("starts_on") or ""), "LOA进入日期")
    ends_value = str(payload.get("ends_on") or "").strip()
    ends_on = parse_iso_date(ends_value, "LOA结束日期") if ends_value else None
    open_period = (
        db.query(EmployeeLOAPeriod)
        .filter(EmployeeLOAPeriod.employee_id == employee.id, EmployeeLOAPeriod.status == "active", EmployeeLOAPeriod.ends_on.is_(None))
        .order_by(EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
        .first()
    )
    if open_period:
        if not ends_on:
            raise HTTPException(409, "该员工已处于LOA，请填写结束日期完成登记")
        starts_on = parse_iso_date(open_period.starts_on, "LOA进入日期")
        if ends_on < starts_on:
            raise HTTPException(400, "LOA结束日期不能早于进入日期")
        for month in months_between(starts_on, ends_on):
            ensure_month_open(db, month, employee.attraction_id, "登记LOA结束")
        try:
            before = loa_period_payload(open_period, employee)
            open_period.ends_on = ends_on.isoformat()
            open_period.ended_by = user.id
            open_period.ended_by_name = user.name
            open_period.ended_at = datetime.now()
            if str(payload.get("note") or "").strip():
                open_period.note = str(payload.get("note")).strip()
            for month in months_between(starts_on, ends_on):
                recalculate_attendance(db, employee, month)
            write_audit(db, user.employee, "登记LOA结束", "employee_loa_period", open_period.id, before=before, after=loa_period_payload(open_period, employee), ip_address=client_ip(request))
            db.commit()
        except Exception:
            db.rollback()
            raise
        invalidate_data_caches()
        return {"ok": True, "record": loa_period_payload(open_period, employee), "excluded_months": months_between(starts_on, ends_on), "completed": True}
    if ends_on and ends_on < starts_on:
        raise HTTPException(400, "LOA结束日期不能早于进入日期")
    affected_months = months_between(starts_on, ends_on or starts_on)
    for month in affected_months:
        ensure_month_open(db, month, employee.attraction_id, "登记LOA")
    overlap = db.query(EmployeeLOAPeriod).filter(
        EmployeeLOAPeriod.employee_id == employee.id,
        EmployeeLOAPeriod.status != "cancelled",
        EmployeeLOAPeriod.starts_on <= (ends_on or starts_on).isoformat(),
        or_(EmployeeLOAPeriod.ends_on.is_(None), EmployeeLOAPeriod.ends_on >= starts_on.isoformat()),
    ).first()
    if overlap:
        raise HTTPException(409, "该员工在所选日期内已有LOA记录，请先修改或撤销原记录")
    row = EmployeeLOAPeriod(
        employee_id=employee.id,
        starts_on=starts_on.isoformat(),
        ends_on=ends_on.isoformat() if ends_on else None,
        status="active",
        note=str(payload.get("note") or "").strip() or None,
        created_by=user.id,
        created_by_name=user.name,
    )
    try:
        db.add(row)
        db.flush()
        for month in affected_months:
            recalculate_attendance(db, employee, month)
        write_audit(db, user.employee, "登记LOA", "employee_loa_period", row.id, after=loa_period_payload(row, employee), ip_address=client_ip(request))
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True, "record": loa_period_payload(row, employee), "excluded_months": affected_months, "completed": bool(ends_on)}


@router.delete("/loa-periods/{period_id}")
def cancel_loa_period(period_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("LOA_REGISTER"))):
    row = db.get(EmployeeLOAPeriod, period_id)
    if not row or row.status == "cancelled":
        raise HTTPException(404, "LOA记录不存在或已撤销")
    employee = db.get(Employee, row.employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    starts_on = parse_iso_date(row.starts_on, "LOA开始日期")
    ends_on = parse_iso_date(row.ends_on or date.today().isoformat(), "LOA结束日期")
    for month in months_between(starts_on, ends_on):
        ensure_month_open(db, month, employee.attraction_id, "撤销LOA")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "请填写撤销原因")
    try:
        before = loa_period_payload(row, employee)
        row.status = "cancelled"
        row.ended_by = user.id
        row.ended_by_name = user.name
        row.ended_at = datetime.now()
        row.note = f"{row.note or ''}\n撤销原因：{reason}".strip()
        for month in months_between(starts_on, ends_on):
            recalculate_attendance(db, employee, month)
        write_audit(db, user.employee, "撤销LOA", "employee_loa_period", row.id, before=before, after=loa_period_payload(row, employee), reason=reason, ip_address=client_ip(request))
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True}


@router.get("/hr/employees")
def hr_employees(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee)
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.employee_no).all()
    payloads = employee_payloads(db, employees)
    return [payloads[employee.id] for employee in employees]


@router.get("/hr/import-template")
def employee_import_template(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    wb = build_employee_import_template()
    watermark_workbook(wb, user.employee.employee_no)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    write_audit(db, user.employee, "下载员工导入模板", "employee_import_template", None)
    db.commit()
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="employee_import_template_v2.xlsx"'})


@router.post("/hr/import-employees")
async def import_employees(request: Request, workbook: UploadFile = File(...), db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    if Path(workbook.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(400, "请上传xlsx文件")
    content = await workbook.read(5 * 1024 * 1024 + 1)
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "导入文件不能为空且不能超过5MB")
    try:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:
        raise HTTPException(400, "Excel文件无法读取") from exc
    expected = ["员工号", "姓名", "角色代码", "景点圈", "初始密码", "在职", "账号启用", "任职开始日", "任职结束日", "到期恢复角色代码"]
    header = [str(cell or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if header[:len(expected)] != expected:
        raise HTTPException(400, "表头与V2导入模板不一致")
    roles = {row.code: row for row in db.query(Role).all()}
    circles = {
        row.name: row
        for row in db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .all()
    }
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    allowed_circle_names = {
        row.name for row in db.query(Attraction).filter(Attraction.id.in_(allowed_attractions)).all()
    } if allowed_attractions is not None else None
    role_codes_allowed = set(roles) if "SYSTEM_ADMIN" in user.permissions else CIRCLE_HR_MANAGED_ROLE_CODES
    parsed, seen, errors = [], set(), []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values[:len(expected)]) + [None] * max(0, len(expected) - len(values))
        employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code = [str(v).strip() if v is not None else "" for v in values]
        if not any((employee_no, name, role_code, attraction_name)):
            continue
        if not employee_no or not name or role_code not in roles:
            errors.append(f"第{row_number}行：员工号、姓名必填且角色代码必须有效")
            continue
        if role_code not in role_codes_allowed:
            errors.append(f"第{row_number}行：景点圈HR只能创建LEAD及以下账号")
            continue
        if len(employee_no) != 7 or not employee_no.isdigit():
            errors.append(f"第{row_number}行：员工号必须为7位纯数字")
            continue
        if attraction_name and attraction_name not in circles:
            errors.append(f"第{row_number}行：景点圈只能填写热力追踪、矮人迷宫或小熊罐子")
            continue
        if allowed_circle_names is not None and attraction_name not in allowed_circle_names:
            errors.append(f"第{row_number}行：景点圈HR只能导入所属景点圈员工")
            continue
        if (
            employee_no in seen
            or db.query(Employee).filter(Employee.employee_no == employee_no).first()
            or db.query(UserAccount).filter(UserAccount.login_account == employee_no).first()
        ):
            errors.append(f"第{row_number}行：员工号{employee_no}重复")
            continue
        if ends_on and return_code not in roles:
            errors.append(f"第{row_number}行：临时角色必须填写有效的到期恢复角色")
            continue
        if ends_on and return_code not in role_codes_allowed:
            errors.append(f"第{row_number}行：景点圈HR只能将到期恢复角色设置为LEAD及以下")
            continue
        seen.add(employee_no)
        parsed.append((employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code))
    if errors:
        raise HTTPException(400, "；".join(errors[:20]))
    created = 0
    try:
        for employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code in parsed:
            attraction = None
            if attraction_name:
                attraction = circles[attraction_name]
            employee = Employee(employee_no=employee_no, name=name, attraction_id=attraction.id if attraction else None, is_active=active not in {"否", "0", "false", "False"}, hired_on=starts_on or date.today().isoformat())
            db.add(employee)
            db.flush()
            db.add(EmployeeRoleAssignment(employee_id=employee.id, role_id=roles[role_code].id, starts_on=starts_on or date.today().isoformat(), ends_on=ends_on or None, assignment_type="temporary" if ends_on else "permanent", return_role_id=roles[return_code].id if ends_on else None, status="active", reason="HR Excel导入", created_by=user.id))
            # Keep Excel imports consistent with the documented employee-account
            # rule. An explicitly supplied initial password still takes priority.
            initial_password = password or default_initial_password(employee_no)
            db.add(
                UserAccount(
                    employee_id=employee.id,
                    login_account=employee_no,
                    password_hash=hash_password(initial_password),
                    enabled=enabled not in {"否", "0", "false", "False"},
                    must_change_password=True,
                    credential_initialized=True,
                )
            )
            created += 1
        write_audit(db, user.employee, "Excel导入员工", "employee_import", after={"created": created, "filename": workbook.filename}, ip_address=client_ip(request))
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True, "created": created}


@router.post("/hr/employees")
def create_employee(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee_no = str(payload.get("employee_no") or "").strip()
    name = str(payload.get("name") or "").strip()
    role = db.query(Role).filter(Role.code == str(payload.get("role_code") or "")).first()
    if not employee_no or not name or not role:
        raise HTTPException(400, "员工号、姓名和角色必填")
    if len(employee_no) != 7 or not employee_no.isdigit():
        raise HTTPException(400, "员工号必须为7位纯数字")
    ensure_hr_role_allowed(user, role.code)
    if db.query(Employee).filter(Employee.employee_no == employee_no).first() or db.query(UserAccount).filter(UserAccount.login_account == employee_no).first():
        raise HTTPException(400, "该员工号已存在，不能重复创建账号")
    attraction_id = employee_circle_id(db, payload.get("attraction_id"))
    ensure_scoped_hr_attraction(db, user, attraction_id)
    employee = Employee(employee_no=employee_no, name=name, attraction_id=attraction_id, is_active=True, hired_on=str(payload.get("hired_on") or date.today().isoformat()))
    db.add(employee)
    db.flush()
    ends_on = str(payload.get("ends_on") or "").strip() or None
    return_role = db.query(Role).filter(Role.code == str(payload.get("return_role_code") or "")).first() if ends_on else None
    if ends_on:
        if not return_role:
            raise HTTPException(400, "临时角色必须选择有效的到期恢复角色")
        ensure_hr_role_allowed(user, return_role.code, "到期恢复角色")
    db.add(
        EmployeeRoleAssignment(
            employee_id=employee.id,
            role_id=role.id,
            starts_on=str(payload.get("starts_on") or date.today().isoformat()),
            ends_on=ends_on,
            assignment_type="temporary" if ends_on else "permanent",
            return_role_id=return_role.id if return_role else None,
            status="active",
            reason="HR新建员工",
            created_by=user.id,
        )
    )
    db.flush()
    synchronize_gsm_management_scope(db, employee, role.code)
    requested_password = str(payload.get("password") or "").strip()
    # New accounts use the employee-number suffix by default. The user must
    # change it on the first successful login; HR may still explicitly provide
    # a different initial password when needed.
    temporary_password = requested_password or default_initial_password(employee_no)
    db.add(
        UserAccount(
            employee_id=employee.id,
            login_account=employee_no,
            password_hash=hash_password(temporary_password),
            enabled=True,
            must_change_password=True,
            credential_initialized=True,
        )
    )
    write_audit(db, user.employee, "新建员工", "employee", employee.id, after={"employee_no": employee_no, "role": role.name}, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {
        "ok": True,
        "employee": employee_payload(db, employee),
        "temporary_password": temporary_password,
        "must_change_password": True,
    }


@router.get("/hr/employee-number-targets")
def employee_number_targets(
    keyword: str = "",
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Search current or historic employee numbers without exposing HR accounts."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以变更员工号")
    normalized = str(keyword or "").strip()
    if not normalized:
        return {"items": [], "total": 0}
    escaped = normalized.replace("!", "!!").replace("%", "!%").replace("_", "!_")
    pattern = f"%{escaped}%"
    candidate_rows = (
        db.query(Employee, EmployeeNumberHistory.old_employee_no)
        .join(UserAccount, UserAccount.employee_id == Employee.id)
        .outerjoin(EmployeeNumberHistory, EmployeeNumberHistory.employee_id == Employee.id)
        .filter(
            or_(
                Employee.employee_no.like(pattern, escape="!"),
                Employee.name.like(pattern, escape="!"),
                EmployeeNumberHistory.old_employee_no.like(pattern, escape="!"),
            )
        )
        .order_by(Employee.name, Employee.employee_no, EmployeeNumberHistory.id.desc())
        .limit(max(1, min(int(limit or 30) * 3, 150)))
        .all()
    )
    items: list[dict] = []
    seen: set[int] = set()
    for employee, matched_old_number in candidate_rows:
        if employee.id in seen:
            continue
        role = role_at(db, employee.id)
        try:
            ensure_employee_number_change_target(db, user, employee, role)
        except HTTPException:
            continue
        seen.add(employee.id)
        items.append(
            {
                "id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": employee.attraction.name if employee.attraction else "未分配景点圈",
                "matched_historical_no": matched_old_number or "",
            }
        )
        if len(items) >= max(1, min(int(limit or 30), 50)):
            break
    return {"items": items, "total": len(items)}


@router.post("/hr/employees/{employee_id}/employee-number")
def change_employee_number(
    employee_id: int,
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Migrate a login number while preserving the employee primary key and all records."""
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    role = ensure_employee_number_change_target(db, user, employee)
    new_employee_no = str(payload.get("new_employee_no") or "").strip()
    reason = str(payload.get("reason") or "").strip()
    reset_password = bool(payload.get("reset_password"))
    if len(new_employee_no) != 7 or not new_employee_no.isdigit():
        raise HTTPException(400, "新员工号必须为7位纯数字")
    if not reason:
        raise HTTPException(400, "请填写员工号变更原因")
    if len(reason) > 300:
        raise HTTPException(400, "员工号变更原因不能超过300个字符")
    old_employee_no = employee.employee_no
    if new_employee_no == old_employee_no:
        return {"ok": True, "unchanged": True, "employee_id": employee.id, "employee_no": employee.employee_no}
    collision = db.query(Employee).filter(Employee.employee_no == new_employee_no, Employee.id != employee.id).first()
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    login_collision = db.query(UserAccount).filter(UserAccount.login_account == new_employee_no).first()
    if collision or (login_collision and (not account or login_collision.id != account.id)):
        raise HTTPException(400, "新员工号已存在，不能合并或重复创建账号")
    if not account:
        raise HTTPException(400, "该员工尚未开通登录账号")
    before = {"employee_no": old_employee_no, "login_account": account.login_account, "role": role.name}
    employee.employee_no = new_employee_no
    employee.updated_at = datetime.now()
    account.login_account = new_employee_no
    account.failed_attempts = 0
    account.locked_until = None
    if reset_password:
        account.password_hash = hash_password(default_initial_password(new_employee_no))
        account.must_change_password = True
        account.credential_initialized = True
        account.password_changed_at = None
    db.add(
        EmployeeNumberHistory(
            employee_id=employee.id,
            old_employee_no=old_employee_no,
            new_employee_no=new_employee_no,
            effective_on=date.today().isoformat(),
            reason=reason,
            changed_by=user.id,
            changed_by_name=user.name,
        )
    )
    revoked_sessions = db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "变更员工号",
        "employee_number_change",
        employee.id,
        before=before,
        after={
            "employee_no": new_employee_no,
            "login_account": new_employee_no,
            "role": role.name,
            "password_reset": reset_password,
            "sessions_revoked": int(revoked_sessions),
        },
        reason=reason,
        ip_address=client_ip(request),
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(400, "新员工号已存在，不能合并或重复创建账号")
    invalidate_data_caches()
    return {
        "ok": True,
        "unchanged": False,
        "employee_id": employee.id,
        "old_employee_no": old_employee_no,
        "employee_no": new_employee_no,
        "sessions_revoked": int(revoked_sessions),
        "must_change_password": bool(account.must_change_password),
        "password_reset": reset_password,
    }


@router.put("/hr/employees/{employee_id}")
def update_employee(employee_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    ensure_scoped_hr_employee(db, user, employee)
    existing_role = role_at(db, employee.id)
    if existing_role and existing_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES and "SYSTEM_ADMIN" not in user.permissions:
        raise HTTPException(403, "景点圈HR只能编辑LEAD及以下员工")
    before = employee_payload(db, employee)
    score_sensitive_fields = {"is_active", "employment_status", "role_code", "attraction_id", "leader_id", "group_id"}
    if score_sensitive_fields & payload.keys():
        current_month = date.today().strftime("%Y-%m")
        ensure_month_open(db, current_month, employee.attraction_id, "修改当月人员计分状态")
        if "attraction_id" in payload and payload.get("attraction_id") not in (None, ""):
            requested_attraction_id = employee_circle_id(db, payload.get("attraction_id"))
            if requested_attraction_id != employee.attraction_id:
                ensure_month_open(db, current_month, requested_attraction_id, "调入员工")
        if str(payload.get("employment_status") or "").strip().lower() == "loa":
            loa_start = parse_iso_date(str(payload.get("loa_start_date") or date.today().isoformat()), "LOA开始日期")
            for loa_month in months_between(loa_start.replace(day=1), date.today().replace(day=1)):
                ensure_month_open(db, loa_month, employee.attraction_id, "设置LOA")
    attendance_state_changed = "is_active" in payload or "employment_status" in payload
    if "name" in payload:
        employee.name = str(payload["name"]).strip()
    if "attraction_id" in payload:
        employee.attraction_id = employee_circle_id(db, payload["attraction_id"])
        ensure_scoped_hr_attraction(db, user, employee.attraction_id)
    requested_employment_status = str(payload.get("employment_status") or "").strip().lower()
    if requested_employment_status:
        if requested_employment_status not in {"active", "loa", "terminated"}:
            raise HTTPException(400, "人员状态无效")
        current_loa = (
            db.query(EmployeeLOAPeriod)
            .filter(
                EmployeeLOAPeriod.employee_id == employee.id,
                EmployeeLOAPeriod.status == "active",
                EmployeeLOAPeriod.ends_on.is_(None),
            )
            .order_by(EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
            .first()
        )
        if requested_employment_status == "loa":
            if not existing_role or existing_role.code not in FRONTLINE_CODES:
                raise HTTPException(400, "仅可将CM/TR演职人员设置为LOA")
            starts_on = str(payload.get("loa_start_date") or date.today().isoformat())
            parse_iso_date(starts_on, "LOA开始日期")
            if not current_loa:
                employee.is_active = True
                employee.terminated_on = None
                db.add(
                    EmployeeLOAPeriod(
                        employee_id=employee.id,
                        starts_on=starts_on,
                        status="active",
                        note=str(payload.get("reason") or "HR设置LOA").strip() or None,
                        created_by=user.id,
                        created_by_name=user.name,
                    )
                )
                write_audit(
                    db,
                    user.employee,
                    "设置LOA（长期病假）",
                    "employee_loa",
                    employee.id,
                    after={"starts_on": starts_on, "status": "LOA（长期病假）"},
                    reason=str(payload.get("reason") or ""),
                    ip_address=client_ip(request),
                )
        else:
            if current_loa:
                end_value = date.today() - timedelta(days=1)
                if end_value.isoformat() < current_loa.starts_on:
                    current_loa.status = "cancelled"
                    current_loa.ends_on = current_loa.starts_on
                else:
                    current_loa.status = "ended"
                    current_loa.ends_on = end_value.isoformat()
                current_loa.ended_by = user.id
                current_loa.ended_by_name = user.name
                current_loa.ended_at = datetime.now()
                write_audit(
                    db,
                    user.employee,
                    "结束LOA（长期病假）",
                    "employee_loa",
                    employee.id,
                    after={"ends_on": current_loa.ends_on, "next_status": requested_employment_status},
                    reason=str(payload.get("reason") or ""),
                    ip_address=client_ip(request),
                )
            if requested_employment_status == "terminated" and groups_led_by(db, employee.id):
                raise HTTPException(400, "该员工仍在带组，请先整组移交")
            employee.is_active = requested_employment_status == "active"
            employee.terminated_on = None if employee.is_active else date.today().isoformat()
    elif "is_active" in payload:
        if not bool(payload["is_active"]) and groups_led_by(db, employee.id):
            raise HTTPException(400, "该员工仍在带组，请先整组移交")
        employee.is_active = bool(payload["is_active"])
        employee.terminated_on = None if employee.is_active else date.today().isoformat()
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if account and "account_enabled" in payload:
        next_enabled = bool(payload["account_enabled"])
        if next_enabled != account.enabled:
            account.enabled = next_enabled
            account.disabled_at = None if next_enabled else datetime.now()
    new_role_code = str(payload.get("role_code") or "")
    current_role = existing_role
    resulting_role = current_role
    if new_role_code and current_role and new_role_code != current_role.code:
        attendance_state_changed = True
        new_role = db.query(Role).filter(Role.code == new_role_code).first()
        if not new_role:
            raise HTTPException(400, "角色不存在")
        ensure_hr_role_allowed(user, new_role.code)
        led_groups = groups_led_by(db, employee.id)
        if new_role.code not in LEADER_CODES and led_groups:
            member_ids = {
                membership.employee_id
                for group in led_groups
                for membership in active_group_memberships(db, group.id)
            }
            if member_ids:
                raise HTTPException(400, "该员工仍有组员，请先整组移交")
            pending_review_count = db.query(RecognitionRecord).filter(
                RecognitionRecord.assigned_reviewer_id == employee.id,
                RecognitionRecord.status == "pending",
            ).count()
            if pending_review_count:
                raise HTTPException(400, f"该员工还有{pending_review_count}条待复核记录，请先完成审批后再更改角色")
            for group in led_groups:
                leader_assignment = active_group_leader(db, group.id)
                if leader_assignment:
                    leader_assignment.status = "ended"
                    leader_assignment.ends_on = date.today().isoformat()
                group.status = "closed"
                group.revision += 1
                db.query(SystemAlert).filter(
                    SystemAlert.group_id == group.id,
                    SystemAlert.status == "open",
                ).update(
                    {SystemAlert.status: "handled", SystemAlert.handled_by: user.id, SystemAlert.handled_at: datetime.now()},
                    synchronize_session=False,
                )
                write_audit(
                    db,
                    user.employee,
                    "关闭空工作组",
                    "work_group",
                    group.id,
                    before={"leader": employee.name, "status": "active", "member_count": 0},
                    after={"leader": "", "status": "closed", "member_count": 0},
                    reason=f"组长角色由{current_role.name}变更为{new_role.name}",
                    ip_address=client_ip(request),
                )
        current_assignment = (
            db.query(EmployeeRoleAssignment)
            .filter(EmployeeRoleAssignment.employee_id == employee.id, EmployeeRoleAssignment.status == "active")
            .order_by(EmployeeRoleAssignment.starts_on.desc())
            .first()
        )
        if current_assignment:
            current_assignment.ends_on = date.today().isoformat()
            current_assignment.status = "expired"
        ends_on = str(payload.get("role_ends_on") or "").strip() or None
        return_role = db.query(Role).filter(Role.code == str(payload.get("return_role_code") or "")).first() if ends_on else None
        if ends_on:
            if not return_role:
                raise HTTPException(400, "临时角色必须选择有效的到期恢复角色")
            ensure_hr_role_allowed(user, return_role.code, "到期恢复角色")
        db.add(
            EmployeeRoleAssignment(
                employee_id=employee.id,
                role_id=new_role.id,
                starts_on=date.today().isoformat(),
                ends_on=ends_on,
                assignment_type="temporary" if ends_on else "permanent",
                return_role_id=return_role.id if return_role else None,
                status="active",
                reason=str(payload.get("reason") or "HR变更角色"),
                created_by=user.id,
            )
        )
        resulting_role = new_role

    synchronize_gsm_management_scope(db, employee, resulting_role.code if resulting_role else None)

    group_change_requested = "leader_id" in payload or "group_id" in payload
    if group_change_requested:
        requested_leader_id = int(payload["leader_id"]) if payload.get("leader_id") else None
        requested_group_id = int(payload["group_id"]) if payload.get("group_id") else None
        current_membership = (
            db.query(GroupMembership)
            .filter(
                GroupMembership.employee_id == employee.id,
                GroupMembership.status == "active",
                GroupMembership.starts_on <= date.today().isoformat(),
                or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= date.today().isoformat()),
            )
            .order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc())
            .first()
        )
        if not resulting_role or resulting_role.code not in FRONTLINE_CODES:
            if requested_leader_id or requested_group_id:
                raise HTTPException(400, "只有CM/TR可以设置组长")
            if current_membership:
                current_membership.status = "ended"
                current_membership.ends_on = date.today().isoformat()
        else:
            if (requested_leader_id or requested_group_id) and not employee.is_active:
                raise HTTPException(400, "离职员工不能设置组长")
            target_group = db.get(WorkGroup, requested_group_id) if requested_group_id else None
            target_leader = None
            if target_group:
                target_assignment = active_group_leader(db, target_group.id)
                target_leader = target_assignment.leader if target_assignment else None
                if requested_leader_id and (not target_leader or target_leader.id != requested_leader_id):
                    raise HTTPException(409, "工作组组长已经变化，请刷新后重试")
            elif requested_leader_id:
                target_leader = db.get(Employee, requested_leader_id)
                candidate_groups = [group for group in groups_led_by(db, requested_leader_id) if group.attraction_id == employee.attraction_id]
                if len(candidate_groups) > 1:
                    raise HTTPException(409, "该组长有多个工作组，请选择具体工作组")
                if candidate_groups:
                    target_group = candidate_groups[0]
            if target_leader:
                target_role = role_at(db, target_leader.id)
                if not target_leader.is_active or not target_role or target_role.code not in LEADER_CODES:
                    raise HTTPException(400, "新组长必须是在职TA主管或主管")
                if target_leader.attraction_id != employee.attraction_id:
                    raise HTTPException(400, "新组长必须与员工属于同一景点圈")
                if not target_group:
                    target_group = WorkGroup(name=f"{target_leader.name}工作组", attraction_id=employee.attraction_id, status="active")
                    db.add(target_group)
                    db.flush()
                    db.add(GroupLeaderAssignment(group_id=target_group.id, leader_employee_id=target_leader.id, starts_on=date.today().isoformat(), status="active"))
            elif requested_group_id:
                raise HTTPException(400, "所选工作组当前没有有效组长")

            # Returning a former TA主管/主管 to CM/TR should restore the
            # latest viable historic group in the selected circle when that
            # choice is unambiguous.  We never guess between two historical
            # groups: HR must select the intended LEAD in that case.
            if not target_group and not requested_leader_id and new_role_code and current_role and current_role.code in LEADER_CODES:
                candidates: list[tuple[GroupMembership, WorkGroup, Employee]] = []
                history = db.query(GroupMembership).filter(
                    GroupMembership.employee_id == employee.id,
                    GroupMembership.status.in_(("ended", "active")),
                ).order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc()).all()
                seen_groups: set[int] = set()
                for historic in history:
                    if historic.group_id in seen_groups:
                        continue
                    seen_groups.add(historic.group_id)
                    group = db.get(WorkGroup, historic.group_id)
                    assignment = active_group_leader(db, group.id) if group and group.status == "active" and group.attraction_id == employee.attraction_id else None
                    leader = assignment.leader if assignment else None
                    leader_role = role_at(db, leader.id) if leader else None
                    if leader and leader.is_active and leader_role and leader_role.code in LEADER_CODES:
                        candidates.append((historic, group, leader))
                if len(candidates) == 1:
                    _, target_group, target_leader = candidates[0]
                elif len(candidates) > 1:
                    raise HTTPException(409, "该员工在目标景点圈有多个历史工作组，请选择明确组长后再保存")

            pending_query = db.query(RecognitionRecord).filter(RecognitionRecord.employee_id == employee.id, RecognitionRecord.status == "pending")
            current_group_id = current_membership.group_id if current_membership else None
            target_group_id = target_group.id if target_group else None
            if target_group_id != current_group_id:
                if not target_group and pending_query.count():
                    raise HTTPException(400, "该员工还有待复核记录，必须选择新组长")
                if current_membership:
                    current_membership.status = "ended"
                    current_membership.ends_on = date.today().isoformat()
                if target_group:
                    db.add(GroupMembership(group_id=target_group.id, employee_id=employee.id, starts_on=date.today().isoformat(), status="active", reason=str(payload.get("reason") or "HR调整组长")))
                    pending_query.update({RecognitionRecord.assigned_reviewer_id: target_leader.id}, synchronize_session=False)
                write_audit(
                    db,
                    user.employee,
                    "调整员工组长",
                    "employee_group",
                    employee.id,
                    before={"group_id": current_group_id, "leader_id": before.get("leader_id"), "leader_name": before.get("leader_name")},
                    after={"group_id": target_group_id, "leader_id": target_leader.id if target_leader else None, "leader_name": target_leader.name if target_leader else "未分配"},
                    reason=str(payload.get("reason") or "HR员工管理页面调整"),
                    ip_address=client_ip(request),
                )
    employee.updated_at = datetime.now()
    db.flush()
    if attendance_state_changed:
        recalculate_attendance(db, employee, date.today().strftime("%Y-%m"))
    write_audit(db, user.employee, "修改员工", "employee", employee.id, before=before, after=employee_payload(db, employee), reason=str(payload.get("reason") or ""), ip_address=client_ip(request))
    db.commit()
    return {"ok": True}

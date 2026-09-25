"""Read-only, cross-circle statistics for registered disciplinary events."""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.excel_export_utils import content_disposition, excel_row
from app.v2_auth import V2User, current_user, require_permissions
from app.v2_database import get_db
from app.v2_models import Attraction, DeductionLevel, DeductionRecord, DeductionType, DeductionUpgradeRequest, StoredFile
from app.v2_services import write_audit
from app.v2_watermark import watermark_workbook
from app.routers._shared import client_ip

router = APIRouter()

STATUS_LABELS = {"active": "已生效", "pending_material": "待补材料", "material_processing": "待补材料", "material_failed": "待补材料", "pending_upgrade": "审核中", "void": "已作废"}
MATERIAL_LABELS = {"ready": "材料已就绪", "missing": "待补材料", "processing": "材料处理中", "failed": "材料处理失败"}
UPGRADE_LABELS = {"pending": "审核中", "material_processing": "材料处理中", "source_first": "原声明", "source_second": "升级来源", "result": "升级结果", "rejected": "审核不通过", "eligible": "可升级"}


def checked_month(month: str) -> str:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise HTTPException(400, "月份格式无效")
    return month


def declaration_payload(
    db: Session, month: str, attraction_id: int | None, level_id: int | None,
    type_id: int | None, status: str, keyword: str,
) -> dict:
    checked_month(month)
    if status and status not in set(STATUS_LABELS.values()):
        raise HTTPException(400, "业务状态无效")
    attraction_rows = db.query(Attraction).all()
    attractions = {row.id: row.name for row in attraction_rows}
    levels = {row.id: row for row in db.query(DeductionLevel).all()}
    types = {row.id: row for row in db.query(DeductionType).all()}
    if attraction_id is not None and attraction_id not in attractions:
        raise HTTPException(400, "景点圈无效")
    if level_id is not None and level_id not in levels:
        raise HTTPException(400, "处分等级无效")
    if type_id is not None and type_id not in types:
        raise HTTPException(400, "处分类型无效")

    # Approved upgrades store both the second statement and its final result.
    # The result represents the second event; retaining both would double count.
    query = db.query(DeductionRecord).filter(DeductionRecord.deduction_month == month)
    if attraction_id is not None:
        query = query.filter(DeductionRecord.attraction_id_snapshot == attraction_id)
    if level_id is not None:
        query = query.filter(DeductionRecord.deduction_level_id == level_id)
    if type_id is not None:
        query = query.filter(DeductionRecord.deduction_type_id == type_id)
    needle = keyword.strip().casefold()
    selected = query.order_by(DeductionRecord.occurred_on.desc(), DeductionRecord.id.desc()).all()
    approved_seconds = {
        second_id for (second_id,) in db.query(DeductionUpgradeRequest.second_deduction_id)
        .join(DeductionRecord, DeductionRecord.id == DeductionUpgradeRequest.second_deduction_id)
        .filter(DeductionUpgradeRequest.status == "approved", DeductionUpgradeRequest.result_deduction_id.isnot(None),
                DeductionRecord.deduction_month == month).all()
    }
    file_ids = {row.document_file_id for row in selected if row.document_file_id}
    files = {row.id: row for row in db.query(StoredFile).filter(StoredFile.id.in_(file_ids)).all()} if file_ids else {}
    records = []
    for row in selected:
        if row.id in approved_seconds:
            continue
        business_status = STATUS_LABELS.get(row.status, row.status)
        if status and business_status != status:
            continue
        if needle and needle not in row.employee_name.casefold() and needle not in row.employee_no.casefold():
            continue
        material = files.get(row.document_file_id)
        document_url = f"/api/files/{row.document_file_id}" if material and material.status == "active" and row.material_status == "ready" and row.upgrade_role != "result" else ""
        level = levels.get(row.deduction_level_id)
        records.append({
            "id": row.id, "employee_id": row.employee_id, "employee_name": row.employee_name,
            "employee_no": row.employee_no, "attraction_id": row.attraction_id_snapshot,
            "attraction_name": attractions.get(row.attraction_id_snapshot, "未记录原景点"),
            "occurred_on": row.occurred_on, "type_id": row.deduction_type_id,
            "type_name": row.deduction_type_name, "level_id": row.deduction_level_id,
            "level_code": level.code if level else "OTHER", "level_name": row.deduction_level_name,
            "description": row.description, "submitter_name": row.submitter_name,
            "submitted_at": row.submitted_at.strftime("%Y-%m-%d %H:%M:%S"),
            "material_status": MATERIAL_LABELS.get(row.material_status, row.material_status),
            "upgrade_status": UPGRADE_LABELS.get(row.upgrade_state or "", row.upgrade_state or "无"),
            "business_status": business_status, "raw_status": row.status,
            "upgrade_request_id": row.upgrade_request_id or 0, "document_url": document_url,
        })

    level_counts = Counter(row["level_id"] for row in records)
    circle_rows = defaultdict(list)
    for row in records:
        circle_rows[row["attraction_id"]].append(row)
    circles = []
    for circle_id, rows in sorted(circle_rows.items(), key=lambda item: (item[0] is None, attractions.get(item[0], ""))):
        circles.append({
            "attraction_id": circle_id, "attraction_name": attractions.get(circle_id, "未记录原景点"),
            "record_count": len(rows), "employee_count": len({row["employee_id"] for row in rows}),
            "levels": [{"id": item.id, "code": item.code, "name": item.name, "count": sum(row["level_id"] == item.id for row in rows)} for item in levels.values() if any(row["level_id"] == item.id for row in rows)],
            "statuses": {name: sum(row["business_status"] == name for row in rows) for name in dict.fromkeys(STATUS_LABELS.values())},
        })
    return {
        "month": month, "total_records": len(records),
        "employee_count": len({row["employee_id"] for row in records}),
        "levels": [{"id": item.id, "code": item.code, "name": item.name, "count": level_counts[item.id]} for item in sorted(levels.values(), key=lambda item: (item.points, item.id))],
        "statuses": {name: sum(row["business_status"] == name for row in records) for name in dict.fromkeys(STATUS_LABELS.values())},
        "attractions": [{"id": item.id, "name": item.name} for item in sorted(attraction_rows, key=lambda item: item.name)
                        if item.employee_circle or item.id in circle_rows],
        "types": [{"id": item.id, "name": item.name} for item in sorted(types.values(), key=lambda item: item.name)],
        "circles": circles, "records": records,
    }


@router.get("/declaration-statistics")
def get_declaration_statistics(
    month: str | None = None, attraction_id: int | None = None, level_id: int | None = None,
    type_id: int | None = None, status: str = "", keyword: str = "",
    db: Session = Depends(get_db), user: V2User = Depends(require_permissions("DECLARATION_STATS_VIEW")),
):
    return declaration_payload(db, month or date.today().strftime("%Y-%m"), attraction_id, level_id, type_id, status, keyword)


def build_declaration_workbook(data: dict) -> Workbook:
    wb = Workbook()
    summary = wb.active
    summary.title = "景点圈与等级汇总"
    summary.append(["声明登记统计", data["month"]])
    summary.append(["总记录数", data["total_records"], "涉及人数", data["employee_count"]])
    summary.append(["景点圈", "处分等级", "记录数", "涉及人数", "已生效", "待补材料", "审核中", "已作废"])
    for circle in data["circles"]:
        rows = [row for row in data["records"] if row["attraction_id"] == circle["attraction_id"]]
        for level in circle["levels"]:
            subset = [row for row in rows if row["level_id"] == level["id"]]
            summary.append(excel_row([circle["attraction_name"], level["name"], len(subset), len({row["employee_id"] for row in subset}), *[sum(row["business_status"] == state for row in subset) for state in ("已生效", "待补材料", "审核中", "已作废")]]))
    detail = wb.create_sheet("逐条明细")
    detail.append(["员工", "工号", "原景点圈", "事件日期", "处分类型", "处分等级", "内容", "登记人", "登记时间", "材料状态", "审核或升级状态", "业务状态", "升级工单号", "材料入口"])
    for row in data["records"]:
        detail.append(excel_row([row["employee_name"], row["employee_no"], row["attraction_name"], row["occurred_on"], row["type_name"], row["level_name"], row["description"], row["submitter_name"], row["submitted_at"], row["material_status"], row["upgrade_status"], row["business_status"], row["upgrade_request_id"] or "", "系统内查看" if row["document_url"] else "无可查看材料"]))
    for ws in wb:
        ws.freeze_panes = "A4" if ws is summary else "A2"
        ws.auto_filter.ref = f"A{3 if ws is summary else 1}:{get_column_letter(ws.max_column)}{ws.max_row}"
        for cell in ws[3 if ws is summary else 1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(45, max(13, max(len(str(cell.value or "")) for cell in col) + 2))
            for cell in col:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    return wb


@router.get("/declaration-statistics/export")
def export_declaration_statistics(
    request: Request, month: str | None = None, attraction_id: int | None = None,
    level_id: int | None = None, type_id: int | None = None, status: str = "", keyword: str = "",
    db: Session = Depends(get_db), user: V2User = Depends(current_user),
):
    if not {"DECLARATION_STATS_VIEW", "DECLARATION_STATS_EXPORT"}.issubset(user.permissions):
        raise HTTPException(403, "没有权限执行该操作")
    data = declaration_payload(db, month or date.today().strftime("%Y-%m"), attraction_id, level_id, type_id, status, keyword)
    workbook = build_declaration_workbook(data)
    watermark_workbook(workbook, user.employee.employee_no)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    write_audit(db, user.employee, "导出声明登记统计", "declaration_statistics_export", data["month"], after={"attraction_id": attraction_id, "level_id": level_id, "type_id": type_id, "status": status, "keyword": keyword, "record_count": data["total_records"]}, ip_address=client_ip(request))
    db.commit()
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": content_disposition(f"declaration_statistics_{data['month']}.xlsx", f"声明登记统计_{data['month']}.xlsx")})

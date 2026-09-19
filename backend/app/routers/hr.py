"""Employee, account and HR administration endpoints."""
from __future__ import annotations

import json
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import or_, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.v2_auth import V2User, current_user, require_permissions
from app.v2_crypto import account_reset_password, default_initial_password, hash_password
from app.v2_database import CIRCLE_HR_ACCOUNTS, get_db, synchronize_gsm_management_scope
from app.v2_models import Attraction, AuditLog, CircleTransferRequest, DeductionFollowUp, DeductionRecord, Employee, EmployeeNumberHistory, EmployeeLOAPeriod, EmployeeRoleAssignment, GroupLeaderAssignment, GroupMembership, GroupTransfer, GroupTransferMember, GovernanceCase, RecognitionRecord, RecognitionScoreRule, Role, SickLeaveRecord, StoredFile, SystemAlert, UserAccount, UserSession, WorkGroup
from app.v2_services import FRONTLINE_CODES, LEADER_CODES, RECOGNIZER_CODES, active_group_leader, active_group_leaders_bulk, active_group_memberships, active_group_memberships_bulk, current_group_for_employee, current_leader_for_employee, groups_led_by, groups_led_by_bulk, gsm_candidates_for_attractions_bulk, managed_attraction_ids, process_role_expirations, recalculate_attendance, role_at, roles_at, write_audit
from app.v2_watermark import watermark_workbook
from app.excel_export import build_employee_import_template
from app.routers._shared import (
    CIRCLE_HR_MANAGED_ROLE_CODES,
    EMPLOYEE_TARGET_PERMISSIONS,
    REGULAR_ACCOUNT_ROLE_CODES,
    SCOPED_HR_ROLE_CODE,
    backup_health_payload,
    client_ip,
    ensure_month_open,
    ensure_scoped_hr_attraction,
    ensure_scoped_hr_employee,
    group_display_metadata_bulk,
    invalidate_data_caches,
    like_escaped_pattern,
    month_closure_payload,
    months_between,
    parse_iso_date,
    scoped_hr_attraction_ids,
    search_employee_targets,
    visible_system_alerts,
)

router = APIRouter()


def ensure_employee_number_change_target(db: Session, user: V2User, employee: Employee, role: Role | None = None) -> Role:
    """Apply the employee-number migration scope without changing org data."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以变更员工号")
    target_role = role or role_at(db, employee.id)
    if not target_role:
        raise HTTPException(400, "该员工当前没有有效角色")
    if target_role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR和最高管理员账号不能在此处变更员工号")
    ensure_scoped_hr_employee(db, user, employee)
    if user.role.code == SCOPED_HR_ROLE_CODE and target_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能变更本圈CM、TR、TA主管或主管的员工号")
    return target_role


def ensure_hr_role_allowed(user: V2User, role_code: str, label: str = "角色") -> None:
    if "SYSTEM_ADMIN" not in user.permissions and role_code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, f"景点圈HR只能设置{label}为CM、TR、TA主管或主管")


def employee_payload(db: Session, employee: Employee) -> dict:
    return employee_payloads(db, [employee])[employee.id]


def login_account_archive_state(employee: Employee, account: UserAccount | None, role: Role | None) -> dict:
    """Return non-sensitive eligibility for removing only a login account.

    The employee primary key and every business row remain untouched.  Both
    the directory and the write endpoint use this same policy so the seven-day
    rule cannot be bypassed by calling the API directly.
    """
    result = {"deleted": bool(employee.account_deleted_at), "eligible": False, "reason": "", "eligible_on": ""}
    if employee.account_deleted_at:
        result["reason"] = "账号已删除，业务档案已保留"
        return result
    if not account:
        result["reason"] = "该员工尚未开通登录账号"
        return result
    if not role or role.code not in REGULAR_ACCOUNT_ROLE_CODES:
        result["reason"] = "HR和最高管理员账号不支持在此删除"
        return result
    reference_day: date | None = None
    if not employee.is_active:
        try:
            reference_day = date.fromisoformat(str(employee.terminated_on or ""))
        except ValueError:
            result["reason"] = "离职日期未记录，暂不能删除登录账号"
            return result
    elif not account.enabled:
        if not account.disabled_at:
            result["reason"] = "停用时间未记录，暂不能删除登录账号"
            return result
        reference_day = account.disabled_at.date()
    else:
        result["reason"] = "需先离职或停用账号"
        return result
    eligible_on = reference_day + timedelta(days=7)
    result["eligible_on"] = eligible_on.isoformat()
    if date.today() < eligible_on:
        result["reason"] = f"{eligible_on.isoformat()} 起可删除登录账号"
        return result
    result["eligible"] = True
    result["reason"] = "可删除登录账号；员工及所有业务档案会保留"
    return result


def employee_payloads(db: Session, employees: list[Employee], on_date: str | None = None) -> dict[int, dict]:
    """Build the HR employee directory with a bounded set of batch queries."""
    if not employees:
        return {}
    value = on_date or date.today().isoformat()
    employee_ids = [employee.id for employee in employees]
    role_map = roles_at(db, employee_ids, value)
    accounts = {
        account.employee_id: account
        for account in db.query(UserAccount).filter(UserAccount.employee_id.in_(employee_ids)).all()
    }
    loa_periods = (
        db.query(EmployeeLOAPeriod)
        .filter(
            EmployeeLOAPeriod.employee_id.in_(employee_ids),
            EmployeeLOAPeriod.status != "cancelled",
            EmployeeLOAPeriod.starts_on <= value,
            or_(EmployeeLOAPeriod.ends_on.is_(None), EmployeeLOAPeriod.ends_on >= value),
        )
        .order_by(EmployeeLOAPeriod.employee_id, EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
        .all()
    )
    loa_by_employee: dict[int, EmployeeLOAPeriod] = {}
    for period in loa_periods:
        loa_by_employee.setdefault(period.employee_id, period)
    attraction_ids = {employee.attraction_id for employee in employees if employee.attraction_id}
    attractions = {
        attraction.id: attraction
        for attraction in db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).all()
    } if attraction_ids else {}
    memberships = (
        db.query(GroupMembership)
        .filter(
            GroupMembership.employee_id.in_(employee_ids),
            GroupMembership.status == "active",
            GroupMembership.starts_on <= value,
            or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= value),
        )
        .order_by(GroupMembership.employee_id, GroupMembership.starts_on.desc(), GroupMembership.id.desc())
        .all()
    )
    membership_by_employee: dict[int, GroupMembership] = {}
    for membership in memberships:
        membership_by_employee.setdefault(membership.employee_id, membership)
    group_ids = {membership.group_id for membership in membership_by_employee.values()}
    groups = {
        group.id: group
        for group in db.query(WorkGroup).filter(WorkGroup.id.in_(group_ids)).all()
    } if group_ids else {}
    leader_assignments = (
        db.query(GroupLeaderAssignment)
        .filter(
            GroupLeaderAssignment.group_id.in_(group_ids),
            GroupLeaderAssignment.status == "active",
            GroupLeaderAssignment.starts_on <= value,
            or_(GroupLeaderAssignment.ends_on.is_(None), GroupLeaderAssignment.ends_on >= value),
        )
        .order_by(GroupLeaderAssignment.group_id, GroupLeaderAssignment.starts_on.desc(), GroupLeaderAssignment.id.desc())
        .all()
    ) if group_ids else []
    leader_assignment_by_group: dict[int, GroupLeaderAssignment] = {}
    for assignment in leader_assignments:
        leader_assignment_by_group.setdefault(assignment.group_id, assignment)
    leader_ids = {assignment.leader_employee_id for assignment in leader_assignment_by_group.values()}
    leaders = {
        leader.id: leader
        for leader in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()
    } if leader_ids else {}
    group_display = group_display_metadata_bulk(db, group_ids, value)
    result: dict[int, dict] = {}
    for employee in employees:
        role = role_map.get(employee.id)
        account = accounts.get(employee.id)
        membership = membership_by_employee.get(employee.id)
        group = groups.get(membership.group_id) if membership else None
        leader_assignment = leader_assignment_by_group.get(group.id) if group else None
        leader = leaders.get(leader_assignment.leader_employee_id) if leader_assignment else None
        display = group_display.get(group.id, {}) if group else {}
        attraction = attractions.get(employee.attraction_id)
        loa_period = loa_by_employee.get(employee.id)
        archive_state = login_account_archive_state(employee, account, role)
        result[employee.id] = {
        "id": employee.id,
        "employee_no": employee.employee_no,
        "name": employee.name,
        "role_code": role.code if role else "",
        "role_name": role.name if role else "未配置",
        "attraction_id": employee.attraction_id,
        "attraction_name": attraction.name if attraction else "",
        "group_id": group.id if group else None,
        "group_name": display.get("name", group.name if group else ""),
        "previous_group_leader_name": display.get("previous_leader_name", ""),
        "previous_group_leader_until": display.get("previous_leader_until", ""),
        "leader_id": leader.id if leader else None,
        "leader_name": leader.name if leader else "",
        "is_active": employee.is_active,
        "employment_status": "loa" if employee.is_active and loa_period else "active" if employee.is_active else "terminated",
        "loa_start_date": loa_period.starts_on if loa_period else "",
        "account_enabled": bool(account and account.enabled),
        "login_account": account.login_account if account else "",
        "account_deleted_at": employee.account_deleted_at.strftime("%Y-%m-%d %H:%M:%S") if employee.account_deleted_at else "",
        "account_deleted_by_name": employee.account_deleted_by_name or "",
        "account_deletion_eligible": archive_state["eligible"],
        "account_deletion_reason": archive_state["reason"],
        "account_deletion_eligible_on": archive_state["eligible_on"],
        }
    return result


def account_password_status(account: UserAccount | None) -> tuple[str, str]:
    """Return a non-sensitive, evidence-based password state for HR views."""
    if not account:
        return "未开通账号", "unprovisioned"
    if account.must_change_password:
        return "待本人修改初始/重置密码", "pending_change"
    if account.password_changed_at:
        return "已修改密码", "changed"
    return "历史状态未记录", "unknown"


def account_login_state(account: UserAccount | None, employee: Employee) -> tuple[str, str]:
    if employee.account_deleted_at:
        return "账号已删除·留档", "archived"
    if not account:
        return "未开通账号", "unprovisioned"
    if not employee.is_active:
        return "员工已离职（不可登录）", "employee_inactive"
    if not account.enabled:
        return "账号已停用", "disabled"
    if account.locked_until and account.locked_until > datetime.now():
        return f"临时锁定至 {account.locked_until.strftime('%Y-%m-%d %H:%M')}", "locked"
    return "账号已启用", "enabled"


def employee_circle_id(db: Session, value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        attraction_id = int(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "员工景点圈无效") from exc
    circle = db.get(Attraction, attraction_id)
    if not circle or not circle.active or not circle.employee_circle:
        raise HTTPException(400, "员工只能归属有效景点圈")
    return circle.id


@router.get("/frontline-employees")
def frontline_employees(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if not ({"EMPLOYEE_ADD", "SICK_REGISTER", "DEDUCTION_ALL", "DEDUCTION_DIRECT"} & user.permissions):
        raise HTTPException(403, "没有权限查询CM/TR")
    allowed = scoped_hr_attraction_ids(db, user)
    attraction_id = next(iter(allowed)) if allowed else None
    return search_employee_targets(db, attraction_id=attraction_id, limit=500)["items"]


@router.get("/employee-targets")
def employee_targets(
    usage: str,
    keyword: str = "",
    attraction_id: int | None = None,
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    required = EMPLOYEE_TARGET_PERMISSIONS.get(usage)
    if not required:
        raise HTTPException(400, "员工查询用途无效")
    if not (required & user.permissions):
        raise HTTPException(403, "没有对应的员工登记权限")
    if usage == "poc":
        if user.role.code not in {"TA_GSM", "GSM", "AM"}:
            raise HTTPException(403, "仅TA GSM、GSM、AM可以查询POC被认可员工")
        value = like_escaped_pattern(keyword)
        query = db.query(Employee).join(UserAccount, UserAccount.employee_id == Employee.id).filter(
            Employee.is_active.is_(True), UserAccount.enabled.is_(True),
            or_(Employee.name.like(value, escape="\\"), Employee.employee_no.like(value, escape="\\")),
        )
        if attraction_id is not None:
            query = query.filter(Employee.attraction_id == attraction_id)
        rows = []
        for item in query.order_by(Employee.name, Employee.employee_no).limit(max(1, min(limit, 50)) * 4).all():
            role = role_at(db, item.id)
            if role and role.code in (FRONTLINE_CODES | LEADER_CODES):
                circle = db.get(Attraction, item.attraction_id) if item.attraction_id else None
                rows.append({"id": item.id, "employee_no": item.employee_no, "name": item.name, "role_code": role.code, "role_name": role.name, "attraction_id": item.attraction_id, "attraction_name": circle.name if circle else "", "group_id": None, "group_name": ""})
        return {"items": rows[:max(1, min(limit, 50))], "total": len(rows), "limit": max(1, min(limit, 50))}
    # Deduction/absence target searches are intentionally keyword-only for
    # cross-circle managers.  TA主管 retains its home-circle boundary, while
    # 主管、TA GSM、GSM may locate active CM/TR across all circles.
    global_target_search = usage in {"deduction", "attendance"} and user.role.code in {"SUPERVISOR", "TA_GSM", "GSM"}
    if global_target_search and not keyword.strip():
        return {"items": [], "total": 0, "limit": max(1, min(limit, 50)), "search_scope": "全部景点圈在职CM/TR（请输入姓名或员工号）"}
    if user.role.code == "TA_SUPERVISOR" and usage in {"deduction", "attendance"}:
        attraction_id = user.employee.attraction_id
    if attraction_id is not None:
        circle = db.get(Attraction, attraction_id)
        if not circle or not circle.active or not circle.employee_circle:
            raise HTTPException(400, "请选择有效员工景点圈")
    elif not global_target_search:
        allowed = scoped_hr_attraction_ids(db, user)
        attraction_id = next(iter(allowed)) if allowed else None
    if not global_target_search:
        ensure_scoped_hr_attraction(db, user, attraction_id)
    result = search_employee_targets(
        db,
        keyword=keyword,
        attraction_id=attraction_id,
        limit=max(1, min(limit, 50)),
    )
    result["search_scope"] = "全部景点圈在职CM/TR（请输入姓名或员工号）" if global_target_search else "本景点圈在职CM/TR"
    return result


@router.get("/admin/operations-health")
def operations_health(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SYSTEM_ADMIN"))):
    health = backup_health_payload()
    month = date.today().strftime("%Y-%m")
    circles = db.query(Attraction).filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True)).order_by(Attraction.name).all()
    closures = [month_closure_payload(db, month, circle.id, user) for circle in circles]
    return {
        "backup": health,
        "month": month,
        "month_closures": closures,
        "open_system_alerts": db.query(SystemAlert).filter(SystemAlert.status == "open").count(),
        "pending_circle_transfers": db.query(CircleTransferRequest).filter(CircleTransferRequest.status == "pending").count(),
        "governance": {
            "open_cases": db.query(GovernanceCase).filter(GovernanceCase.status == "open").count(),
            "overdue_cases": db.query(GovernanceCase).filter(GovernanceCase.status == "open", GovernanceCase.due_at < datetime.now()).count(),
            "overdue_recognition_reviews": db.query(RecognitionRecord).filter(RecognitionRecord.status == "pending", RecognitionRecord.submitted_at < datetime.now() - timedelta(hours=48)).count(),
            "retention_review_files": db.query(StoredFile).filter(StoredFile.status == "active", StoredFile.uploaded_at < datetime.now() - timedelta(days=730)).count(),
        },
    }


@router.get("/hr/employees")
def hr_employees(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee)
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.employee_no).all()
    payloads = employee_payloads(db, employees)
    return [payloads[employee.id] for employee in employees]


@router.get("/hr/account-status")
def hr_account_status(request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    """Return approved account-status metadata without exposing passwords.

    This is deliberately separate from employee-management edit scope: a
    scoped HR can read CM/TR/TALEAD/Lead status across all three circles, but
    their existing edit endpoints remain limited to their own circle.
    """
    if user.role.code == SCOPED_HR_ROLE_CODE:
        permitted_roles = CIRCLE_HR_MANAGED_ROLE_CODES
        scope_label = "全部景点圈的 CM/TR、TALEAD、Lead"
    elif user.role.code == "SYSTEM_ADMIN":
        permitted_roles = REGULAR_ACCOUNT_ROLE_CODES
        scope_label = "全部常规账号"
    else:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以查看账号状态")

    employees = db.query(Employee).order_by(Employee.name, Employee.employee_no).all()
    role_map = roles_at(db, [employee.id for employee in employees])
    target_employees = [employee for employee in employees if role_map.get(employee.id) and role_map[employee.id].code in permitted_roles]
    target_ids = [employee.id for employee in target_employees]
    accounts = {
        account.employee_id: account
        for account in db.query(UserAccount).filter(UserAccount.employee_id.in_(target_ids)).all()
    } if target_ids else {}
    attraction_ids = {employee.attraction_id for employee in target_employees if employee.attraction_id}
    attractions = {
        row.id: row.name
        for row in db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).all()
    } if attraction_ids else {}
    items = []
    for employee in target_employees:
        account = accounts.get(employee.id)
        role = role_map[employee.id]
        account_status, account_status_code = account_login_state(account, employee)
        password_status, password_status_code = account_password_status(account)
        items.append(
            {
                "employee_id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": attractions.get(employee.attraction_id, "未分配景点圈"),
                "login_account": account.login_account if account else "",
                "account_status": account_status,
                "account_status_code": account_status_code,
                "password_status": password_status,
                "password_status_code": password_status_code,
                "password_changed_at": account.password_changed_at.strftime("%Y-%m-%d %H:%M:%S") if account and account.password_changed_at else "",
                "login_status": "已登录" if account and account.last_login_at else "从未登录",
                "last_login_at": account.last_login_at.strftime("%Y-%m-%d %H:%M:%S") if account and account.last_login_at else "",
            }
        )
    write_audit(
        db,
        user.employee,
        "查看账号状态",
        "account_status",
        after={"scope": scope_label, "returned_count": len(items)},
        ip_address=client_ip(request),
    )
    db.commit()
    return {"scope_label": scope_label, "items": items}


def gsm_candidates_for_attraction(db: Session, attraction_id: int, on_date: str | None = None) -> list[tuple[Employee, Role]]:
    return gsm_candidates_for_attractions_bulk(db, [attraction_id], on_date).get(attraction_id, [])


@router.get("/hr/organization")
def hr_organization(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    today_value = date.today().isoformat()
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee)
    attraction_query = db.query(Attraction)
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
        attraction_query = attraction_query.filter(Attraction.id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.name, Employee.employee_no).all()
    payloads = employee_payloads(db, employees, today_value)
    roles = roles_at(db, [employee.id for employee in employees], today_value)
    rows = []
    represented: set[int] = set()

    for attraction in attraction_query.order_by(Attraction.name).all():
        attraction_employees = [employee for employee in employees if employee.attraction_id == attraction.id and employee.is_active]
        if not attraction_employees:
            continue
        attraction_node = f"hr-attraction-{attraction.id}"
        rows.append({"node_id": attraction_node, "parent_id": "", "level": 0, "node_type": "attraction", "name": attraction.name})
        gsm_candidates = gsm_candidates_for_attraction(db, attraction.id, today_value)
        gsm_managers = [(employee, role) for employee, role in gsm_candidates if role.code == "GSM"]
        ta_gsm_managers = [(employee, role) for employee, role in gsm_candidates if role.code == "TA_GSM"]
        management_parent = attraction_node
        if len(gsm_managers) == 1:
            primary_gsm, _primary_gsm_role = gsm_managers[0]
            management_parent = f"hr-gsm-{attraction.id}-{primary_gsm.id}"
            rows.append({"node_id": management_parent, "parent_id": attraction_node, "level": 1, "node_type": "employee", "hierarchy_role": "gsm", "employee": payloads[primary_gsm.id]})
            represented.add(primary_gsm.id)
        elif len(gsm_managers) > 1:
            # Parallel GSMs jointly own every supervisor group in the circle.
            # A single team node prevents duplicating supervisors under each GSM.
            for gsm, _gsm_role in gsm_managers:
                gsm_node = f"hr-gsm-{attraction.id}-{gsm.id}"
                rows.append({"node_id": gsm_node, "parent_id": attraction_node, "level": 1, "node_type": "employee", "hierarchy_role": "gsm", "employee": payloads[gsm.id]})
                represented.add(gsm.id)
            management_parent = f"hr-gsm-team-{attraction.id}"
            manager_names = "、".join(gsm.name for gsm, _gsm_role in gsm_managers)
            rows.append({"node_id": management_parent, "parent_id": attraction_node, "level": 1, "node_type": "gsm_team", "name": f"主管组（由GSM共同承接：{manager_names}）"})
        else:
            management_parent = f"hr-gsm-team-{attraction.id}"
            rows.append({"node_id": management_parent, "parent_id": attraction_node, "level": 1, "node_type": "gsm_team", "name": "主管组（未配置GSM）"})

        leaders = [employee for employee in attraction_employees if roles[employee.id] and roles[employee.id].code in LEADER_CODES]
        assigned_frontline: set[int] = set()
        for leader in sorted(leaders, key=lambda employee: (employee.name, employee.employee_no)):
            leader_node = f"hr-leader-{leader.id}"
            rows.append({"node_id": leader_node, "parent_id": management_parent, "level": 2, "node_type": "employee", "hierarchy_role": "supervisor", "employee": payloads[leader.id]})
            represented.add(leader.id)
            leader_members = [
                member for member in attraction_employees
                if payloads[member.id]["leader_id"] == leader.id
                and roles.get(member.id)
                and roles[member.id].code in FRONTLINE_CODES
            ]
            for member in sorted(leader_members, key=lambda employee: (employee.name, employee.employee_no)):
                rows.append({"node_id": f"hr-employee-{member.id}", "parent_id": leader_node, "level": 3, "node_type": "employee", "hierarchy_role": "frontline", "employee": payloads[member.id]})
                assigned_frontline.add(member.id)
                represented.add(member.id)

        unassigned = [
            employee for employee in attraction_employees
            if roles[employee.id] and roles[employee.id].code in FRONTLINE_CODES and employee.id not in assigned_frontline
        ]
        if unassigned:
            placeholder = f"hr-unassigned-{attraction.id}"
            rows.append({"node_id": placeholder, "parent_id": management_parent, "level": 2, "node_type": "placeholder", "name": "未分配主管", "member_count": len(unassigned)})
            for employee in sorted(unassigned, key=lambda item: (item.name, item.employee_no)):
                rows.append({"node_id": f"hr-employee-{employee.id}", "parent_id": placeholder, "level": 3, "node_type": "employee", "hierarchy_role": "frontline", "employee": payloads[employee.id]})
                represented.add(employee.id)

        # TA GSM is a peer in the GSM layer and never owns a supervisor branch.
        for gsm, _gsm_role in ta_gsm_managers:
            gsm_node = f"hr-gsm-{attraction.id}-{gsm.id}"
            rows.append({"node_id": gsm_node, "parent_id": attraction_node, "level": 1, "node_type": "employee", "hierarchy_role": "gsm", "employee": payloads[gsm.id]})
            represented.add(gsm.id)

    other_employees = [employee for employee in employees if employee.id not in represented]
    if other_employees:
        other_node = "hr-other-employees"
        rows.append({"node_id": other_node, "parent_id": "", "level": 0, "node_type": "other", "name": "其他管理人员 / 未归属员工"})
        for employee in other_employees:
            rows.append({"node_id": f"hr-other-{employee.id}", "parent_id": other_node, "level": 1, "node_type": "employee", "hierarchy_role": "other", "employee": payloads[employee.id]})
    return {"rows": rows}


@router.get("/hr/import-template")
def employee_import_template(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    wb = build_employee_import_template()
    watermark_workbook(wb, user.employee.employee_no)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    write_audit(db, user.employee, "下载员工导入模板", "employee_import_template", None)
    db.commit()
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="employee_import_template_v2.xlsx"'})


@router.post("/hr/import-employees")
async def import_employees(request: Request, workbook: UploadFile = File(...), db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    if Path(workbook.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(400, "请上传xlsx文件")
    content = await workbook.read(5 * 1024 * 1024 + 1)
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "导入文件不能为空且不能超过5MB")
    try:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:
        raise HTTPException(400, "Excel文件无法读取") from exc
    expected = ["员工号", "姓名", "角色代码", "景点圈", "初始密码", "在职", "账号启用", "任职开始日", "任职结束日", "到期恢复角色代码"]
    header = [str(cell or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if header[:len(expected)] != expected:
        raise HTTPException(400, "表头与V2导入模板不一致")
    roles = {row.code: row for row in db.query(Role).all()}
    circles = {
        row.name: row
        for row in db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .all()
    }
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    allowed_circle_names = {
        row.name for row in db.query(Attraction).filter(Attraction.id.in_(allowed_attractions)).all()
    } if allowed_attractions is not None else None
    role_codes_allowed = set(roles) if "SYSTEM_ADMIN" in user.permissions else CIRCLE_HR_MANAGED_ROLE_CODES
    parsed, seen, errors = [], set(), []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values[:len(expected)]) + [None] * max(0, len(expected) - len(values))
        employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code = [str(v).strip() if v is not None else "" for v in values]
        if not any((employee_no, name, role_code, attraction_name)):
            continue
        if not employee_no or not name or role_code not in roles:
            errors.append(f"第{row_number}行：员工号、姓名必填且角色代码必须有效")
            continue
        if role_code not in role_codes_allowed:
            errors.append(f"第{row_number}行：景点圈HR只能创建LEAD及以下账号")
            continue
        if len(employee_no) != 7 or not employee_no.isdigit():
            errors.append(f"第{row_number}行：员工号必须为7位纯数字")
            continue
        if attraction_name and attraction_name not in circles:
            errors.append(f"第{row_number}行：景点圈只能填写热力追踪、矮人迷宫或小熊罐子")
            continue
        if allowed_circle_names is not None and attraction_name not in allowed_circle_names:
            errors.append(f"第{row_number}行：景点圈HR只能导入所属景点圈员工")
            continue
        if (
            employee_no in seen
            or db.query(Employee).filter(Employee.employee_no == employee_no).first()
            or db.query(UserAccount).filter(UserAccount.login_account == employee_no).first()
        ):
            errors.append(f"第{row_number}行：员工号{employee_no}重复")
            continue
        if ends_on and return_code not in roles:
            errors.append(f"第{row_number}行：临时角色必须填写有效的到期恢复角色")
            continue
        if ends_on and return_code not in role_codes_allowed:
            errors.append(f"第{row_number}行：景点圈HR只能将到期恢复角色设置为LEAD及以下")
            continue
        seen.add(employee_no)
        parsed.append((employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code))
    if errors:
        raise HTTPException(400, "；".join(errors[:20]))
    created = 0
    try:
        for employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code in parsed:
            attraction = None
            if attraction_name:
                attraction = circles[attraction_name]
            employee = Employee(employee_no=employee_no, name=name, attraction_id=attraction.id if attraction else None, is_active=active not in {"否", "0", "false", "False"}, hired_on=starts_on or date.today().isoformat())
            db.add(employee)
            db.flush()
            db.add(EmployeeRoleAssignment(employee_id=employee.id, role_id=roles[role_code].id, starts_on=starts_on or date.today().isoformat(), ends_on=ends_on or None, assignment_type="temporary" if ends_on else "permanent", return_role_id=roles[return_code].id if ends_on else None, status="active", reason="HR Excel导入", created_by=user.id))
            # Keep Excel imports consistent with the documented employee-account
            # rule. An explicitly supplied initial password still takes priority.
            initial_password = password or default_initial_password(employee_no)
            db.add(
                UserAccount(
                    employee_id=employee.id,
                    login_account=employee_no,
                    password_hash=hash_password(initial_password),
                    enabled=enabled not in {"否", "0", "false", "False"},
                    must_change_password=True,
                    credential_initialized=True,
                )
            )
            created += 1
        write_audit(db, user.employee, "Excel导入员工", "employee_import", after={"created": created, "filename": workbook.filename}, ip_address=client_ip(request))
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True, "created": created}


@router.post("/hr/employees")
def create_employee(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee_no = str(payload.get("employee_no") or "").strip()
    name = str(payload.get("name") or "").strip()
    role = db.query(Role).filter(Role.code == str(payload.get("role_code") or "")).first()
    if not employee_no or not name or not role:
        raise HTTPException(400, "员工号、姓名和角色必填")
    if len(employee_no) != 7 or not employee_no.isdigit():
        raise HTTPException(400, "员工号必须为7位纯数字")
    ensure_hr_role_allowed(user, role.code)
    if db.query(Employee).filter(Employee.employee_no == employee_no).first() or db.query(UserAccount).filter(UserAccount.login_account == employee_no).first():
        raise HTTPException(400, "该员工号已存在，不能重复创建账号")
    attraction_id = employee_circle_id(db, payload.get("attraction_id"))
    ensure_scoped_hr_attraction(db, user, attraction_id)
    employee = Employee(employee_no=employee_no, name=name, attraction_id=attraction_id, is_active=True, hired_on=str(payload.get("hired_on") or date.today().isoformat()))
    db.add(employee)
    db.flush()
    ends_on = str(payload.get("ends_on") or "").strip() or None
    return_role = db.query(Role).filter(Role.code == str(payload.get("return_role_code") or "")).first() if ends_on else None
    if ends_on:
        if not return_role:
            raise HTTPException(400, "临时角色必须选择有效的到期恢复角色")
        ensure_hr_role_allowed(user, return_role.code, "到期恢复角色")
    db.add(
        EmployeeRoleAssignment(
            employee_id=employee.id,
            role_id=role.id,
            starts_on=str(payload.get("starts_on") or date.today().isoformat()),
            ends_on=ends_on,
            assignment_type="temporary" if ends_on else "permanent",
            return_role_id=return_role.id if return_role else None,
            status="active",
            reason="HR新建员工",
            created_by=user.id,
        )
    )
    db.flush()
    synchronize_gsm_management_scope(db, employee, role.code)
    requested_password = str(payload.get("password") or "").strip()
    # New accounts use the employee-number suffix by default. The user must
    # change it on the first successful login; HR may still explicitly provide
    # a different initial password when needed.
    temporary_password = requested_password or default_initial_password(employee_no)
    db.add(
        UserAccount(
            employee_id=employee.id,
            login_account=employee_no,
            password_hash=hash_password(temporary_password),
            enabled=True,
            must_change_password=True,
            credential_initialized=True,
        )
    )
    write_audit(db, user.employee, "新建员工", "employee", employee.id, after={"employee_no": employee_no, "role": role.name}, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {
        "ok": True,
        "employee": employee_payload(db, employee),
        "temporary_password": temporary_password,
        "must_change_password": True,
    }


def resolved_reset_password(account: UserAccount) -> str:
    """Apply the documented "登录账号后四位" reset rule with the policy length guard."""
    try:
        return account_reset_password(account.login_account)
    except ValueError as exc:
        raise HTTPException(
            400,
            f"该账号的登录账号“{account.login_account}”不足四位，无法按“登录账号后四位”规则重置密码，"
            "否则新密码会低于系统密码长度下限、本人也无法自行改回。请管理员先规范该账号的登录账号，"
            "或改用其他方式处理该账号。",
        ) from exc


@router.post("/accounts/reset-password")
def reset_employee_password(
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("PASSWORD_RESET")),
):
    employee_no = str(payload.get("employee_no") or "").strip()
    name = str(payload.get("name") or "").strip()
    if not employee_no or not name:
        raise HTTPException(400, "员工号和姓名必填")
    employee = db.query(Employee).filter(Employee.employee_no == employee_no).first()
    if not employee or employee.name != name:
        raise HTTPException(400, "员工号或姓名不匹配")
    target_role = role_at(db, employee.id)
    if not target_role:
        raise HTTPException(400, "该员工当前没有有效角色")
    ensure_scoped_hr_employee(db, user, employee)
    if target_role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR管理员、景点圈HR和系统管理员账号不能在此处重置")
    if user.role.code == SCOPED_HR_ROLE_CODE and target_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能重置LEAD及以下员工密码")
    if user.role.code == "GSM" and employee.attraction_id not in managed_attraction_ids(db, user.id):
        raise HTTPException(403, "只能重置本人管理景点圈内的员工密码")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if not account:
        raise HTTPException(400, "该员工尚未开通登录账号")
    reset_password = resolved_reset_password(account)
    account.password_hash = hash_password(reset_password)
    account.failed_attempts = 0
    account.locked_until = None
    account.must_change_password = True
    account.credential_initialized = True
    account.password_changed_at = None
    db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "重置账号密码",
        "user_account",
        account.id,
        after={
            "employee_no": employee.employee_no,
            "employee_name": employee.name,
            "role": target_role.name,
            "account_enabled": account.enabled,
            "password_rule": "登录账号后四位",
            "must_change_password": True,
            "sessions_revoked": True,
        },
        ip_address=client_ip(request),
    )
    db.commit()
    return {
        "ok": True,
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "account_enabled": account.enabled,
        "password_rule": "登录账号后四位",
        "temporary_password": reset_password,
        "must_change_password": True,
        "sessions_revoked": True,
    }


@router.get("/accounts/name-targets")
def account_name_targets(
    keyword: str = "",
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Return only accounts whose current display name this HR user may correct."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以修改账号中文姓名")
    normalized = str(keyword or "").strip()
    if not normalized:
        return {"items": [], "total": 0}
    escaped = normalized.replace("!", "!!").replace("%", "!%").replace("_", "!_")
    pattern = f"%{escaped}%"
    candidates = (
        db.query(Employee)
        .join(UserAccount, UserAccount.employee_id == Employee.id)
        .filter(
            or_(Employee.employee_no.like(pattern, escape="!"), Employee.name.like(pattern, escape="!")),
        )
        .order_by(Employee.name, Employee.employee_no)
        .limit(max(1, min(int(limit or 30), 50)))
        .all()
    )
    roles = roles_at(db, [employee.id for employee in candidates])
    allowed = []
    for employee in candidates:
        role = roles.get(employee.id)
        if not role or role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
            continue
        if user.role.code == SCOPED_HR_ROLE_CODE:
            if employee.attraction_id != user.employee.attraction_id or role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
                continue
        allowed.append(
            {
                "id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": employee.attraction.name if employee.attraction else "未分配景点圈",
            }
        )
    return {"items": allowed, "total": len(allowed)}


@router.get("/hr/employee-number-targets")
def employee_number_targets(
    keyword: str = "",
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Search current or historic employee numbers without exposing HR accounts."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以变更员工号")
    normalized = str(keyword or "").strip()
    if not normalized:
        return {"items": [], "total": 0}
    escaped = normalized.replace("!", "!!").replace("%", "!%").replace("_", "!_")
    pattern = f"%{escaped}%"
    candidate_rows = (
        db.query(Employee, EmployeeNumberHistory.old_employee_no)
        .join(UserAccount, UserAccount.employee_id == Employee.id)
        .outerjoin(EmployeeNumberHistory, EmployeeNumberHistory.employee_id == Employee.id)
        .filter(
            or_(
                Employee.employee_no.like(pattern, escape="!"),
                Employee.name.like(pattern, escape="!"),
                EmployeeNumberHistory.old_employee_no.like(pattern, escape="!"),
            )
        )
        .order_by(Employee.name, Employee.employee_no, EmployeeNumberHistory.id.desc())
        .limit(max(1, min(int(limit or 30) * 3, 150)))
        .all()
    )
    items: list[dict] = []
    seen: set[int] = set()
    for employee, matched_old_number in candidate_rows:
        if employee.id in seen:
            continue
        role = role_at(db, employee.id)
        try:
            ensure_employee_number_change_target(db, user, employee, role)
        except HTTPException:
            continue
        seen.add(employee.id)
        items.append(
            {
                "id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": employee.attraction.name if employee.attraction else "未分配景点圈",
                "matched_historical_no": matched_old_number or "",
            }
        )
        if len(items) >= max(1, min(int(limit or 30), 50)):
            break
    return {"items": items, "total": len(items)}


@router.post("/hr/employees/{employee_id}/employee-number")
def change_employee_number(
    employee_id: int,
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Migrate a login number while preserving the employee primary key and all records."""
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    role = ensure_employee_number_change_target(db, user, employee)
    new_employee_no = str(payload.get("new_employee_no") or "").strip()
    reason = str(payload.get("reason") or "").strip()
    reset_password = bool(payload.get("reset_password"))
    if len(new_employee_no) != 7 or not new_employee_no.isdigit():
        raise HTTPException(400, "新员工号必须为7位纯数字")
    if not reason:
        raise HTTPException(400, "请填写员工号变更原因")
    if len(reason) > 300:
        raise HTTPException(400, "员工号变更原因不能超过300个字符")
    old_employee_no = employee.employee_no
    if new_employee_no == old_employee_no:
        return {"ok": True, "unchanged": True, "employee_id": employee.id, "employee_no": employee.employee_no}
    collision = db.query(Employee).filter(Employee.employee_no == new_employee_no, Employee.id != employee.id).first()
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    login_collision = db.query(UserAccount).filter(UserAccount.login_account == new_employee_no).first()
    if collision or (login_collision and (not account or login_collision.id != account.id)):
        raise HTTPException(400, "新员工号已存在，不能合并或重复创建账号")
    if not account:
        raise HTTPException(400, "该员工尚未开通登录账号")
    before = {"employee_no": old_employee_no, "login_account": account.login_account, "role": role.name}
    employee.employee_no = new_employee_no
    employee.updated_at = datetime.now()
    account.login_account = new_employee_no
    account.failed_attempts = 0
    account.locked_until = None
    if reset_password:
        account.password_hash = hash_password(default_initial_password(new_employee_no))
        account.must_change_password = True
        account.credential_initialized = True
        account.password_changed_at = None
    db.add(
        EmployeeNumberHistory(
            employee_id=employee.id,
            old_employee_no=old_employee_no,
            new_employee_no=new_employee_no,
            effective_on=date.today().isoformat(),
            reason=reason,
            changed_by=user.id,
            changed_by_name=user.name,
        )
    )
    revoked_sessions = db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "变更员工号",
        "employee_number_change",
        employee.id,
        before=before,
        after={
            "employee_no": new_employee_no,
            "login_account": new_employee_no,
            "role": role.name,
            "password_reset": reset_password,
            "sessions_revoked": int(revoked_sessions),
        },
        reason=reason,
        ip_address=client_ip(request),
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(400, "新员工号已存在，不能合并或重复创建账号")
    invalidate_data_caches()
    return {
        "ok": True,
        "unchanged": False,
        "employee_id": employee.id,
        "old_employee_no": old_employee_no,
        "employee_no": new_employee_no,
        "sessions_revoked": int(revoked_sessions),
        "must_change_password": bool(account.must_change_password),
        "password_reset": reset_password,
    }


@router.post("/accounts/update-name")
def update_account_name(
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以修改账号中文姓名")
    employee_id = int(payload.get("employee_id") or 0)
    new_name = str(payload.get("name") or "").strip()
    if not employee_id or not new_name:
        raise HTTPException(400, "请选择账号并填写中文姓名")
    if len(new_name) > 100:
        raise HTTPException(400, "中文姓名不能超过100个字符")
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "账号不存在")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    role = role_at(db, employee.id)
    if not account or not role:
        raise HTTPException(400, "该员工尚未配置有效登录账号或角色")
    ensure_scoped_hr_employee(db, user, employee)
    if role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR和系统管理员账号不能在此处修改姓名")
    if user.role.code == SCOPED_HR_ROLE_CODE and role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能修改本圈CM、TR、TA主管或主管的姓名")
    old_name = employee.name
    if old_name == new_name:
        return {"ok": True, "employee_id": employee.id, "employee_no": employee.employee_no, "name": employee.name, "unchanged": True}
    employee.name = new_name
    employee.updated_at = datetime.now()
    write_audit(
        db,
        user.employee,
        "修改账号中文姓名",
        "employee",
        employee.id,
        before={"employee_no": employee.employee_no, "name": old_name, "role": role.name},
        after={"employee_no": employee.employee_no, "name": new_name, "role": role.name},
        reason="HR账号姓名纠错",
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "employee_id": employee.id, "employee_no": employee.employee_no, "name": employee.name, "unchanged": False}


@router.put("/hr/employees/{employee_id}")
def update_employee(employee_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    ensure_scoped_hr_employee(db, user, employee)
    existing_role = role_at(db, employee.id)
    if existing_role and existing_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES and "SYSTEM_ADMIN" not in user.permissions:
        raise HTTPException(403, "景点圈HR只能编辑LEAD及以下员工")
    before = employee_payload(db, employee)
    score_sensitive_fields = {"is_active", "employment_status", "role_code", "attraction_id", "leader_id", "group_id"}
    if score_sensitive_fields & payload.keys():
        current_month = date.today().strftime("%Y-%m")
        ensure_month_open(db, current_month, employee.attraction_id, "修改当月人员计分状态")
        if "attraction_id" in payload and payload.get("attraction_id") not in (None, ""):
            requested_attraction_id = employee_circle_id(db, payload.get("attraction_id"))
            if requested_attraction_id != employee.attraction_id:
                ensure_month_open(db, current_month, requested_attraction_id, "调入员工")
        if str(payload.get("employment_status") or "").strip().lower() == "loa":
            loa_start = parse_iso_date(str(payload.get("loa_start_date") or date.today().isoformat()), "LOA开始日期")
            for loa_month in months_between(loa_start.replace(day=1), date.today().replace(day=1)):
                ensure_month_open(db, loa_month, employee.attraction_id, "设置LOA")
    attendance_state_changed = "is_active" in payload or "employment_status" in payload
    if "name" in payload:
        employee.name = str(payload["name"]).strip()
    if "attraction_id" in payload:
        employee.attraction_id = employee_circle_id(db, payload["attraction_id"])
        ensure_scoped_hr_attraction(db, user, employee.attraction_id)
    requested_employment_status = str(payload.get("employment_status") or "").strip().lower()
    if requested_employment_status:
        if requested_employment_status not in {"active", "loa", "terminated"}:
            raise HTTPException(400, "人员状态无效")
        current_loa = (
            db.query(EmployeeLOAPeriod)
            .filter(
                EmployeeLOAPeriod.employee_id == employee.id,
                EmployeeLOAPeriod.status == "active",
                EmployeeLOAPeriod.ends_on.is_(None),
            )
            .order_by(EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
            .first()
        )
        if requested_employment_status == "loa":
            if not existing_role or existing_role.code not in FRONTLINE_CODES:
                raise HTTPException(400, "仅可将CM/TR演职人员设置为LOA")
            starts_on = str(payload.get("loa_start_date") or date.today().isoformat())
            parse_iso_date(starts_on, "LOA开始日期")
            if not current_loa:
                employee.is_active = True
                employee.terminated_on = None
                db.add(
                    EmployeeLOAPeriod(
                        employee_id=employee.id,
                        starts_on=starts_on,
                        status="active",
                        note=str(payload.get("reason") or "HR设置LOA").strip() or None,
                        created_by=user.id,
                        created_by_name=user.name,
                    )
                )
                write_audit(
                    db,
                    user.employee,
                    "设置LOA（长期病假）",
                    "employee_loa",
                    employee.id,
                    after={"starts_on": starts_on, "status": "LOA（长期病假）"},
                    reason=str(payload.get("reason") or ""),
                    ip_address=client_ip(request),
                )
        else:
            if current_loa:
                end_value = date.today() - timedelta(days=1)
                if end_value.isoformat() < current_loa.starts_on:
                    current_loa.status = "cancelled"
                    current_loa.ends_on = current_loa.starts_on
                else:
                    current_loa.status = "ended"
                    current_loa.ends_on = end_value.isoformat()
                current_loa.ended_by = user.id
                current_loa.ended_by_name = user.name
                current_loa.ended_at = datetime.now()
                write_audit(
                    db,
                    user.employee,
                    "结束LOA（长期病假）",
                    "employee_loa",
                    employee.id,
                    after={"ends_on": current_loa.ends_on, "next_status": requested_employment_status},
                    reason=str(payload.get("reason") or ""),
                    ip_address=client_ip(request),
                )
            if requested_employment_status == "terminated" and groups_led_by(db, employee.id):
                raise HTTPException(400, "该员工仍在带组，请先整组移交")
            employee.is_active = requested_employment_status == "active"
            employee.terminated_on = None if employee.is_active else date.today().isoformat()
    elif "is_active" in payload:
        if not bool(payload["is_active"]) and groups_led_by(db, employee.id):
            raise HTTPException(400, "该员工仍在带组，请先整组移交")
        employee.is_active = bool(payload["is_active"])
        employee.terminated_on = None if employee.is_active else date.today().isoformat()
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if account and "account_enabled" in payload:
        next_enabled = bool(payload["account_enabled"])
        if next_enabled != account.enabled:
            account.enabled = next_enabled
            account.disabled_at = None if next_enabled else datetime.now()
    new_role_code = str(payload.get("role_code") or "")
    current_role = existing_role
    resulting_role = current_role
    if new_role_code and current_role and new_role_code != current_role.code:
        attendance_state_changed = True
        new_role = db.query(Role).filter(Role.code == new_role_code).first()
        if not new_role:
            raise HTTPException(400, "角色不存在")
        ensure_hr_role_allowed(user, new_role.code)
        led_groups = groups_led_by(db, employee.id)
        if new_role.code not in LEADER_CODES and led_groups:
            member_ids = {
                membership.employee_id
                for group in led_groups
                for membership in active_group_memberships(db, group.id)
            }
            if member_ids:
                raise HTTPException(400, "该员工仍有组员，请先整组移交")
            pending_review_count = db.query(RecognitionRecord).filter(
                RecognitionRecord.assigned_reviewer_id == employee.id,
                RecognitionRecord.status == "pending",
            ).count()
            if pending_review_count:
                raise HTTPException(400, f"该员工还有{pending_review_count}条待复核记录，请先完成审批后再更改角色")
            for group in led_groups:
                leader_assignment = active_group_leader(db, group.id)
                if leader_assignment:
                    leader_assignment.status = "ended"
                    leader_assignment.ends_on = date.today().isoformat()
                group.status = "closed"
                group.revision += 1
                db.query(SystemAlert).filter(
                    SystemAlert.group_id == group.id,
                    SystemAlert.status == "open",
                ).update(
                    {SystemAlert.status: "handled", SystemAlert.handled_by: user.id, SystemAlert.handled_at: datetime.now()},
                    synchronize_session=False,
                )
                write_audit(
                    db,
                    user.employee,
                    "关闭空工作组",
                    "work_group",
                    group.id,
                    before={"leader": employee.name, "status": "active", "member_count": 0},
                    after={"leader": "", "status": "closed", "member_count": 0},
                    reason=f"组长角色由{current_role.name}变更为{new_role.name}",
                    ip_address=client_ip(request),
                )
        current_assignment = (
            db.query(EmployeeRoleAssignment)
            .filter(EmployeeRoleAssignment.employee_id == employee.id, EmployeeRoleAssignment.status == "active")
            .order_by(EmployeeRoleAssignment.starts_on.desc())
            .first()
        )
        if current_assignment:
            current_assignment.ends_on = date.today().isoformat()
            current_assignment.status = "expired"
        ends_on = str(payload.get("role_ends_on") or "").strip() or None
        return_role = db.query(Role).filter(Role.code == str(payload.get("return_role_code") or "")).first() if ends_on else None
        if ends_on:
            if not return_role:
                raise HTTPException(400, "临时角色必须选择有效的到期恢复角色")
            ensure_hr_role_allowed(user, return_role.code, "到期恢复角色")
        db.add(
            EmployeeRoleAssignment(
                employee_id=employee.id,
                role_id=new_role.id,
                starts_on=date.today().isoformat(),
                ends_on=ends_on,
                assignment_type="temporary" if ends_on else "permanent",
                return_role_id=return_role.id if return_role else None,
                status="active",
                reason=str(payload.get("reason") or "HR变更角色"),
                created_by=user.id,
            )
        )
        resulting_role = new_role

    synchronize_gsm_management_scope(db, employee, resulting_role.code if resulting_role else None)

    group_change_requested = "leader_id" in payload or "group_id" in payload
    if group_change_requested:
        requested_leader_id = int(payload["leader_id"]) if payload.get("leader_id") else None
        requested_group_id = int(payload["group_id"]) if payload.get("group_id") else None
        current_membership = (
            db.query(GroupMembership)
            .filter(
                GroupMembership.employee_id == employee.id,
                GroupMembership.status == "active",
                GroupMembership.starts_on <= date.today().isoformat(),
                or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= date.today().isoformat()),
            )
            .order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc())
            .first()
        )
        if not resulting_role or resulting_role.code not in FRONTLINE_CODES:
            if requested_leader_id or requested_group_id:
                raise HTTPException(400, "只有CM/TR可以设置组长")
            if current_membership:
                current_membership.status = "ended"
                current_membership.ends_on = date.today().isoformat()
        else:
            if (requested_leader_id or requested_group_id) and not employee.is_active:
                raise HTTPException(400, "离职员工不能设置组长")
            target_group = db.get(WorkGroup, requested_group_id) if requested_group_id else None
            target_leader = None
            if target_group:
                target_assignment = active_group_leader(db, target_group.id)
                target_leader = target_assignment.leader if target_assignment else None
                if requested_leader_id and (not target_leader or target_leader.id != requested_leader_id):
                    raise HTTPException(409, "工作组组长已经变化，请刷新后重试")
            elif requested_leader_id:
                target_leader = db.get(Employee, requested_leader_id)
                candidate_groups = [group for group in groups_led_by(db, requested_leader_id) if group.attraction_id == employee.attraction_id]
                if len(candidate_groups) > 1:
                    raise HTTPException(409, "该组长有多个工作组，请选择具体工作组")
                if candidate_groups:
                    target_group = candidate_groups[0]
            if target_leader:
                target_role = role_at(db, target_leader.id)
                if not target_leader.is_active or not target_role or target_role.code not in LEADER_CODES:
                    raise HTTPException(400, "新组长必须是在职TA主管或主管")
                if target_leader.attraction_id != employee.attraction_id:
                    raise HTTPException(400, "新组长必须与员工属于同一景点圈")
                if not target_group:
                    target_group = WorkGroup(name=f"{target_leader.name}工作组", attraction_id=employee.attraction_id, status="active")
                    db.add(target_group)
                    db.flush()
                    db.add(GroupLeaderAssignment(group_id=target_group.id, leader_employee_id=target_leader.id, starts_on=date.today().isoformat(), status="active"))
            elif requested_group_id:
                raise HTTPException(400, "所选工作组当前没有有效组长")

            # Returning a former TA主管/主管 to CM/TR should restore the
            # latest viable historic group in the selected circle when that
            # choice is unambiguous.  We never guess between two historical
            # groups: HR must select the intended LEAD in that case.
            if not target_group and not requested_leader_id and new_role_code and current_role and current_role.code in LEADER_CODES:
                candidates: list[tuple[GroupMembership, WorkGroup, Employee]] = []
                history = db.query(GroupMembership).filter(
                    GroupMembership.employee_id == employee.id,
                    GroupMembership.status.in_(("ended", "active")),
                ).order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc()).all()
                seen_groups: set[int] = set()
                for historic in history:
                    if historic.group_id in seen_groups:
                        continue
                    seen_groups.add(historic.group_id)
                    group = db.get(WorkGroup, historic.group_id)
                    assignment = active_group_leader(db, group.id) if group and group.status == "active" and group.attraction_id == employee.attraction_id else None
                    leader = assignment.leader if assignment else None
                    leader_role = role_at(db, leader.id) if leader else None
                    if leader and leader.is_active and leader_role and leader_role.code in LEADER_CODES:
                        candidates.append((historic, group, leader))
                if len(candidates) == 1:
                    _, target_group, target_leader = candidates[0]
                elif len(candidates) > 1:
                    raise HTTPException(409, "该员工在目标景点圈有多个历史工作组，请选择明确组长后再保存")

            pending_query = db.query(RecognitionRecord).filter(RecognitionRecord.employee_id == employee.id, RecognitionRecord.status == "pending")
            current_group_id = current_membership.group_id if current_membership else None
            target_group_id = target_group.id if target_group else None
            if target_group_id != current_group_id:
                if not target_group and pending_query.count():
                    raise HTTPException(400, "该员工还有待复核记录，必须选择新组长")
                if current_membership:
                    current_membership.status = "ended"
                    current_membership.ends_on = date.today().isoformat()
                if target_group:
                    db.add(GroupMembership(group_id=target_group.id, employee_id=employee.id, starts_on=date.today().isoformat(), status="active", reason=str(payload.get("reason") or "HR调整组长")))
                    pending_query.update({RecognitionRecord.assigned_reviewer_id: target_leader.id}, synchronize_session=False)
                write_audit(
                    db,
                    user.employee,
                    "调整员工组长",
                    "employee_group",
                    employee.id,
                    before={"group_id": current_group_id, "leader_id": before.get("leader_id"), "leader_name": before.get("leader_name")},
                    after={"group_id": target_group_id, "leader_id": target_leader.id if target_leader else None, "leader_name": target_leader.name if target_leader else "未分配"},
                    reason=str(payload.get("reason") or "HR员工管理页面调整"),
                    ip_address=client_ip(request),
                )
    employee.updated_at = datetime.now()
    db.flush()
    if attendance_state_changed:
        recalculate_attendance(db, employee, date.today().strftime("%Y-%m"))
    write_audit(db, user.employee, "修改员工", "employee", employee.id, before=before, after=employee_payload(db, employee), reason=str(payload.get("reason") or ""), ip_address=client_ip(request))
    db.commit()
    return {"ok": True}


def ensure_login_account_archive_target(db: Session, user: V2User, employee: Employee) -> tuple[Role, UserAccount, dict]:
    """Authorize credential removal without ever deleting the employee archive."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以删除登录账号")
    ensure_scoped_hr_employee(db, user, employee)
    role = role_at(db, employee.id)
    if not role or role.code not in REGULAR_ACCOUNT_ROLE_CODES:
        raise HTTPException(403, "HR和最高管理员账号不能在此删除")
    if user.role.code == SCOPED_HR_ROLE_CODE and role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能删除本圈CM、TR、TA主管或主管的登录账号")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    state = login_account_archive_state(employee, account, role)
    if not account:
        raise HTTPException(400, state["reason"] or "该员工没有可删除的登录账号")
    if not state["eligible"]:
        raise HTTPException(400, state["reason"])
    return role, account, state


@router.delete("/hr/employees/{employee_id}/account")
def archive_employee_login_account(
    employee_id: int,
    request: Request,
    payload: dict | None = None,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Remove only credentials and sessions after seven days; keep all data."""
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    role, account, _state = ensure_login_account_archive_target(db, user, employee)
    before = {
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "role": role.name,
        "login_account": account.login_account,
        "employee_active": employee.is_active,
        "account_enabled": account.enabled,
    }
    revoked_sessions = db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    db.delete(account)
    employee.account_deleted_at = datetime.now()
    employee.account_deleted_by_name = user.name
    employee.updated_at = datetime.now()
    write_audit(
        db,
        user.employee,
        "删除停用登录账号",
        "employee_login_archive",
        employee.id,
        before=before,
        after={
            "employee_no": employee.employee_no,
            "employee_name": employee.name,
            "role": role.name,
            "login_removed": True,
            "sessions_revoked": int(revoked_sessions),
            "employee_archive_retained": True,
            "business_history_retained": True,
        },
        reason=str((payload or {}).get("reason") or "离职/停用满7天后删除登录账号，保留员工与业务档案"),
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "employee_id": employee.id, "sessions_revoked": int(revoked_sessions), "archive_retained": True}


@router.post("/hr/employees/batch-leaders")
def batch_assign_unclassified_leaders(
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Assign multiple currently unclassified CM/TR employees in one atomic transaction."""
    raw_items = payload.get("items")
    if not isinstance(raw_items, list) or not raw_items:
        raise HTTPException(400, "请至少选择一名未分类组员的组长")
    if len(raw_items) > 200:
        raise HTTPException(400, "单次最多保存200名组员")
    reason = str(payload.get("reason") or "HR未分类组员批量分组").strip()[:200]
    today_value = date.today().isoformat()
    seen_employee_ids: set[int] = set()
    plans: list[dict] = []

    # Validate the full batch before creating or ending any relationship.
    for index, item in enumerate(raw_items, start=1):
        if not isinstance(item, dict):
            raise HTTPException(400, f"第{index}项格式无效")
        try:
            employee_id = int(item.get("employee_id") or 0)
            leader_id = int(item.get("leader_id") or 0)
            requested_group_id = int(item.get("group_id") or 0) or None
        except (TypeError, ValueError):
            raise HTTPException(400, f"第{index}项员工、组长或工作组编号无效")
        if not employee_id or not leader_id:
            raise HTTPException(400, f"第{index}项必须选择员工和组长")
        if employee_id in seen_employee_ids:
            raise HTTPException(400, f"第{index}项员工重复")
        seen_employee_ids.add(employee_id)

        employee = db.get(Employee, employee_id)
        employee_role = role_at(db, employee_id) if employee else None
        if not employee:
            raise HTTPException(404, f"第{index}项员工不存在")
        ensure_scoped_hr_employee(db, user, employee)
        if not employee.is_active or not employee_role or employee_role.code not in FRONTLINE_CODES:
            raise HTTPException(400, f"{employee.name}不是在职CM/TR，不能批量分组")
        if not employee.attraction_id:
            raise HTTPException(400, f"{employee.name}未配置景点圈")
        ensure_month_open(db, date.today().strftime("%Y-%m"), employee.attraction_id, "批量调整组长")

        current_membership = (
            db.query(GroupMembership)
            .filter(
                GroupMembership.employee_id == employee.id,
                GroupMembership.status == "active",
                GroupMembership.starts_on <= today_value,
                or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= today_value),
            )
            .order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc())
            .first()
        )
        current_group = db.get(WorkGroup, current_membership.group_id) if current_membership else None
        current_assignment = active_group_leader(db, current_group.id) if current_group else None
        if current_assignment:
            raise HTTPException(409, f"{employee.name}已分配组长，请刷新页面后重试")

        leader = db.get(Employee, leader_id)
        leader_role = role_at(db, leader_id) if leader else None
        if not leader or not leader.is_active or not leader_role or leader_role.code not in LEADER_CODES:
            raise HTTPException(400, f"{employee.name}选择的组长不是在职TA主管或主管")
        if leader.attraction_id != employee.attraction_id:
            raise HTTPException(400, f"{employee.name}与所选组长不属于同一景点圈")

        target_group = db.get(WorkGroup, requested_group_id) if requested_group_id else None
        if target_group:
            target_assignment = active_group_leader(db, target_group.id)
            if target_group.status == "closed" or target_group.attraction_id != employee.attraction_id:
                raise HTTPException(400, f"{employee.name}选择的工作组无效")
            if not target_assignment or target_assignment.leader_employee_id != leader.id:
                raise HTTPException(409, f"{employee.name}选择的工作组组长已经变化，请刷新后重试")
        else:
            existing_groups = [group for group in groups_led_by(db, leader.id) if group.attraction_id == employee.attraction_id]
            if len(existing_groups) > 1:
                raise HTTPException(409, f"{leader.name}有多个工作组，请刷新后选择具体工作组")
            target_group = existing_groups[0] if existing_groups else None

        plans.append(
            {
                "employee": employee,
                "leader": leader,
                "current_membership": current_membership,
                "current_group": current_group,
                "target_group": target_group,
            }
        )

    created_groups: dict[int, WorkGroup] = {}
    changes = []
    try:
        for plan in plans:
            employee = plan["employee"]
            leader = plan["leader"]
            target_group = plan["target_group"]
            if not target_group:
                target_group = created_groups.get(leader.id)
                if not target_group:
                    target_group = WorkGroup(
                        name=f"{leader.name}工作组",
                        attraction_id=employee.attraction_id,
                        status="active",
                    )
                    db.add(target_group)
                    db.flush()
                    db.add(
                        GroupLeaderAssignment(
                            group_id=target_group.id,
                            leader_employee_id=leader.id,
                            starts_on=today_value,
                            status="active",
                        )
                    )
                    created_groups[leader.id] = target_group

            current_membership = plan["current_membership"]
            if current_membership:
                current_membership.status = "ended"
                current_membership.ends_on = today_value
            db.add(
                GroupMembership(
                    group_id=target_group.id,
                    employee_id=employee.id,
                    starts_on=today_value,
                    status="active",
                    reason=reason,
                )
            )
            db.query(RecognitionRecord).filter(
                RecognitionRecord.employee_id == employee.id,
                RecognitionRecord.status == "pending",
            ).update({RecognitionRecord.assigned_reviewer_id: leader.id}, synchronize_session=False)
            change = {
                "employee_id": employee.id,
                "employee_no": employee.employee_no,
                "employee_name": employee.name,
                "group_id": target_group.id,
                "group_name": target_group.name,
                "leader_id": leader.id,
                "leader_name": leader.name,
            }
            changes.append(change)
            write_audit(
                db,
                user.employee,
                "批量调整未分类组员组长",
                "employee_group",
                employee.id,
                before={
                    "group_id": plan["current_group"].id if plan["current_group"] else None,
                    "leader_id": None,
                    "leader_name": "未分配",
                },
                after={"group_id": target_group.id, "leader_id": leader.id, "leader_name": leader.name},
                reason=reason,
                ip_address=client_ip(request),
            )
        write_audit(
            db,
            user.employee,
            "批量保存未分类组员",
            "employee_group_batch",
            None,
            after={"count": len(changes), "changes": changes},
            reason=reason,
            ip_address=client_ip(request),
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True, "updated": len(changes), "changes": changes}


def circle_transfer_payload(
    db: Session,
    row: CircleTransferRequest,
    *,
    groups: dict[int, WorkGroup] | None = None,
    leaders: dict[int, Employee] | None = None,
) -> dict:
    group_map = groups or {}
    leader_map = leaders or {}
    source_group = group_map.get(row.source_group_id) if row.source_group_id else None
    source_leader = leader_map.get(row.source_leader_id) if row.source_leader_id else None
    target_group = group_map.get(row.target_group_id) if row.target_group_id else None
    target_leader = leader_map.get(row.target_leader_id) if row.target_leader_id else None
    group_display = group_display_metadata_bulk(
        db,
        [group.id for group in (source_group, target_group) if group],
    )
    try:
        migrated_counts = json.loads(row.migrated_record_counts_json or "{}")
    except (TypeError, ValueError):
        migrated_counts = {}
    return {
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": row.employee_no,
        "employee_name": row.employee_name,
        "source_attraction_id": row.source_attraction_id,
        "source_attraction_name": row.source_attraction_name,
        "target_attraction_id": row.target_attraction_id,
        "target_attraction_name": row.target_attraction_name,
        "source_group_name": group_display.get(source_group.id, {}).get("name", source_group.name) if source_group else "未分组",
        "source_leader_name": source_leader.name if source_leader else "未分配",
        "target_group_id": row.target_group_id,
        "target_group_name": group_display.get(target_group.id, {}).get("name", target_group.name) if target_group else "",
        "target_leader_id": row.target_leader_id,
        "target_leader_name": target_leader.name if target_leader else "",
        "reason": row.reason,
        "status": row.status,
        "status_name": {"pending": "待目标HR确认", "completed": "已生效", "rejected": "已拒绝", "cancelled": "已撤回"}.get(row.status, row.status),
        "requested_by_name": row.requested_by_name,
        "requested_at": row.requested_at.strftime("%Y-%m-%d %H:%M:%S"),
        "reviewed_by_name": row.reviewed_by_name or "",
        "reviewed_at": row.reviewed_at.strftime("%Y-%m-%d %H:%M:%S") if row.reviewed_at else "",
        "review_note": row.review_note or "",
        "completed_at": row.completed_at.strftime("%Y-%m-%d %H:%M:%S") if row.completed_at else "",
        "migrated_record_counts": migrated_counts,
    }


@router.get("/hr/circle-transfers")
def circle_transfers(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    transfer_rows = db.query(CircleTransferRequest).order_by(CircleTransferRequest.requested_at.desc(), CircleTransferRequest.id.desc()).limit(500).all()
    if allowed_attractions is not None:
        transfer_rows = [
            row for row in transfer_rows
            if row.source_attraction_id in allowed_attractions or row.target_attraction_id in allowed_attractions
        ]
    circles = (
        db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .order_by(Attraction.name)
        .all()
    )
    group_ids = {row.source_group_id for row in transfer_rows if row.source_group_id} | {
        row.target_group_id for row in transfer_rows if row.target_group_id
    }
    leader_ids = {row.source_leader_id for row in transfer_rows if row.source_leader_id} | {
        row.target_leader_id for row in transfer_rows if row.target_leader_id
    }
    groups = (
        {group.id: group for group in db.query(WorkGroup).filter(WorkGroup.id.in_(group_ids)).all()}
        if group_ids
        else {}
    )
    leaders = (
        {leader.id: leader for leader in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()}
        if leader_ids
        else {}
    )
    return {
        "items": [circle_transfer_payload(db, row, groups=groups, leaders=leaders) for row in transfer_rows],
        "target_circles": [{"id": row.id, "name": row.name} for row in circles],
    }


@router.post("/hr/circle-transfers")
def create_circle_transfer(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee = db.get(Employee, int(payload.get("employee_id") or 0))
    employee_role = role_at(db, employee.id) if employee else None
    if not employee or not employee.is_active or not employee_role or employee_role.code not in FRONTLINE_CODES:
        raise HTTPException(400, "只能为在职CM/TR发起景点圈调动")
    ensure_scoped_hr_employee(db, user, employee)
    source = db.get(Attraction, employee.attraction_id) if employee.attraction_id else None
    target = db.get(Attraction, int(payload.get("target_attraction_id") or 0))
    if not source or not source.employee_circle or not target or not target.active or not target.employee_circle:
        raise HTTPException(400, "原景点圈或目标景点圈无效")
    if source.id == target.id:
        raise HTTPException(400, "目标景点圈不能与当前景点圈相同")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "调动原因必填")
    existing = db.query(CircleTransferRequest).filter(
        CircleTransferRequest.employee_id == employee.id,
        CircleTransferRequest.status == "pending",
    ).first()
    if existing:
        raise HTTPException(409, "该员工已有待确认的景点圈调动申请")
    source_group = current_group_for_employee(db, employee.id)
    source_leader = current_leader_for_employee(db, employee.id)
    row = CircleTransferRequest(
        employee_id=employee.id,
        employee_no=employee.employee_no,
        employee_name=employee.name,
        source_attraction_id=source.id,
        source_attraction_name=source.name,
        target_attraction_id=target.id,
        target_attraction_name=target.name,
        source_group_id=source_group.id if source_group else None,
        source_leader_id=source_leader.id if source_leader else None,
        reason=reason,
        status="pending",
        requested_by=user.id,
        requested_by_name=user.name,
    )
    db.add(row)
    db.flush()
    write_audit(
        db,
        user.employee,
        "发起跨景点圈调动",
        "circle_transfer",
        row.id,
        after=circle_transfer_payload(db, row),
        reason=reason,
        ip_address=client_ip(request),
    )
    db.commit()
    return {"ok": True, "transfer": circle_transfer_payload(db, row)}


def complete_circle_transfer(db: Session, row: CircleTransferRequest, target_leader: Employee, target_group: WorkGroup, reviewer: V2User, request: Request) -> dict:
    employee = db.get(Employee, row.employee_id)
    if not employee or employee.attraction_id != row.source_attraction_id:
        raise HTTPException(409, "员工景点圈已经变化，请撤回申请后重新发起")
    today_value = date.today().isoformat()
    current_month = today_value[:7]
    ensure_month_open(db, current_month, row.source_attraction_id, "迁出员工及当月数据")
    ensure_month_open(db, current_month, row.target_attraction_id, "迁入员工及当月数据")
    old_memberships = (
        db.query(GroupMembership)
        .filter(GroupMembership.employee_id == employee.id, GroupMembership.status == "active")
        .all()
    )
    for membership in old_memberships:
        membership.status = "ended"
        membership.ends_on = max(membership.starts_on, today_value)
        membership.reason = f"跨景点圈调动至{row.target_attraction_name}"
    db.add(
        GroupMembership(
            group_id=target_group.id,
            employee_id=employee.id,
            starts_on=today_value,
            status="active",
            reason=f"由{row.source_attraction_name}调入",
        )
    )
    employee.attraction_id = row.target_attraction_id
    employee.updated_at = datetime.now()
    pending_reviews = db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id == employee.id,
        RecognitionRecord.status == "pending",
    ).update({RecognitionRecord.assigned_reviewer_id: target_leader.id}, synchronize_session=False)
    pending_follow_ups = db.query(DeductionFollowUp).filter(
        DeductionFollowUp.employee_id == employee.id,
        DeductionFollowUp.status == "pending",
    ).update(
        {DeductionFollowUp.supervisor_id: target_leader.id, DeductionFollowUp.supervisor_name: target_leader.name},
        synchronize_session=False,
    )
    recognition_count = db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id == employee.id,
        RecognitionRecord.recognition_month == current_month,
    ).update(
        {RecognitionRecord.home_attraction_id: row.target_attraction_id, RecognitionRecord.home_attraction_name: row.target_attraction_name},
        synchronize_session=False,
    )
    deduction_count = db.query(DeductionRecord).filter(
        DeductionRecord.employee_id == employee.id,
        DeductionRecord.deduction_month == current_month,
    ).update(
        {
            DeductionRecord.attraction_id_snapshot: row.target_attraction_id,
            DeductionRecord.employee_group_id_snapshot: target_group.id,
        },
        synchronize_session=False,
    )
    sick_leave_count = db.query(SickLeaveRecord).filter(
        SickLeaveRecord.employee_id == employee.id,
        SickLeaveRecord.attendance_month == current_month,
    ).update({SickLeaveRecord.attraction_id_snapshot: row.target_attraction_id}, synchronize_session=False)
    recalculate_attendance(db, employee, current_month)
    migrated_counts = {
        "recognitions": recognition_count,
        "deductions": deduction_count,
        "sick_leaves": sick_leave_count,
        "pending_reviews": pending_reviews,
        "pending_follow_ups": pending_follow_ups,
    }
    row.target_group_id = target_group.id
    row.target_leader_id = target_leader.id
    row.status = "completed"
    row.reviewed_by = reviewer.id
    row.reviewed_by_name = reviewer.name
    row.reviewed_at = datetime.now()
    row.completed_at = datetime.now()
    row.migrated_record_counts_json = json.dumps(migrated_counts, ensure_ascii=False)
    write_audit(
        db,
        reviewer.employee,
        "确认跨景点圈调动并同步迁移数据",
        "circle_transfer",
        row.id,
        before={"employee_no": employee.employee_no, "attraction": row.source_attraction_name, "group_id": row.source_group_id},
        after={"attraction": row.target_attraction_name, "group_id": target_group.id, "leader": target_leader.name, "migrated": migrated_counts},
        reason=row.reason,
        ip_address=client_ip(request),
    )
    invalidate_data_caches()
    return migrated_counts


@router.post("/hr/circle-transfers/{transfer_id}/review")
def review_circle_transfer(transfer_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    row = db.get(CircleTransferRequest, transfer_id)
    if not row or row.status != "pending":
        raise HTTPException(404, "待确认调动申请不存在")
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is not None and row.target_attraction_id not in allowed_attractions:
        raise HTTPException(403, "只有目标景点圈HR可以确认该调动")
    action = str(payload.get("action") or "").strip().lower()
    note = str(payload.get("review_note") or "").strip()
    if action == "reject":
        if not note:
            raise HTTPException(400, "拒绝原因必填")
        row.status = "rejected"
        row.reviewed_by = user.id
        row.reviewed_by_name = user.name
        row.reviewed_at = datetime.now()
        row.review_note = note
        write_audit(db, user.employee, "拒绝跨景点圈调动", "circle_transfer", row.id, after=circle_transfer_payload(db, row), reason=note, ip_address=client_ip(request))
        db.commit()
        return {"ok": True, "transfer": circle_transfer_payload(db, row)}
    if action != "accept":
        raise HTTPException(400, "审批动作无效")
    target_leader = db.get(Employee, int(payload.get("target_leader_id") or 0))
    target_role = role_at(db, target_leader.id) if target_leader else None
    if not target_leader or not target_leader.is_active or target_leader.attraction_id != row.target_attraction_id or not target_role or target_role.code not in LEADER_CODES:
        raise HTTPException(400, "请选择目标景点圈内在职TA主管或主管")
    target_group_id = int(payload.get("target_group_id") or 0)
    target_group = db.get(WorkGroup, target_group_id) if target_group_id else None
    if target_group:
        if target_group.status == "closed" or target_group.attraction_id != row.target_attraction_id:
            raise HTTPException(400, "目标工作组无效")
        assignment = active_group_leader(db, target_group.id)
        if not assignment or assignment.leader_employee_id != target_leader.id:
            raise HTTPException(400, "目标工作组与所选LEAD不匹配")
    else:
        target_group = WorkGroup(name=f"{target_leader.name}工作组", attraction_id=row.target_attraction_id, status="active")
        db.add(target_group)
        db.flush()
        db.add(GroupLeaderAssignment(group_id=target_group.id, leader_employee_id=target_leader.id, starts_on=date.today().isoformat(), status="active"))
    row.review_note = note or None
    migrated_counts = complete_circle_transfer(db, row, target_leader, target_group, user, request)
    db.commit()
    return {"ok": True, "transfer": circle_transfer_payload(db, row), "migrated_record_counts": migrated_counts}


@router.post("/hr/circle-transfers/{transfer_id}/cancel")
def cancel_circle_transfer(transfer_id: int, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    row = db.get(CircleTransferRequest, transfer_id)
    if not row or row.status != "pending":
        raise HTTPException(404, "待确认调动申请不存在")
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is not None and row.source_attraction_id not in allowed_attractions:
        raise HTTPException(403, "只有原景点圈HR可以撤回该申请")
    row.status = "cancelled"
    write_audit(db, user.employee, "撤回跨景点圈调动", "circle_transfer", row.id, after=circle_transfer_payload(db, row), reason=row.reason, ip_address=client_ip(request))
    db.commit()
    return {"ok": True, "transfer": circle_transfer_payload(db, row)}


@router.get("/hr/groups")
def hr_groups(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    group_query = db.query(WorkGroup).filter(WorkGroup.status != "closed")
    if allowed_attractions is not None:
        group_query = group_query.filter(WorkGroup.attraction_id.in_(allowed_attractions))
    groups = group_query.order_by(WorkGroup.name).all()
    group_ids = [group.id for group in groups]
    group_display = group_display_metadata_bulk(db, group_ids)
    leader_assignments = active_group_leaders_bulk(db, group_ids)
    memberships = active_group_memberships_bulk(db, group_ids)
    referenced_employees = {assignment.leader_employee_id for assignment in leader_assignments.values() if assignment}
    for member_rows in memberships.values():
        referenced_employees.update(row.employee_id for row in member_rows)
    employees = (
        {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(referenced_employees)).all()}
        if referenced_employees
        else {}
    )
    attractions = {attraction.id: attraction.name for attraction in db.query(Attraction).all()}
    rows = []
    for group in groups:
        leader_assignment = leader_assignments.get(group.id)
        members = memberships.get(group.id, [])
        leader = employees.get(leader_assignment.leader_employee_id) if leader_assignment else None
        display = group_display.get(group.id, {})
        rows.append(
            {
                "id": group.id,
                "name": display.get("name", group.name),
                "stored_name": group.name,
                "previous_leader_name": display.get("previous_leader_name", ""),
                "previous_leader_until": display.get("previous_leader_until", ""),
                "attraction_id": group.attraction_id,
                "attraction_name": attractions.get(group.attraction_id, ""),
                "status": group.status,
                "revision": group.revision,
                "leader_id": leader_assignment.leader_employee_id if leader_assignment else None,
                "leader_name": leader.name if leader else "待接管",
                "member_count": len(members),
                "members": [
                    {
                        "id": employee.id,
                        "employee_no": employee.employee_no,
                        "name": employee.name,
                    }
                    for employee in (employees.get(row.employee_id) for row in members)
                    if employee
                ],
            }
        )
    return rows


@router.get("/hr/leader-options")
def leader_options(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee).filter(Employee.is_active.is_(True))
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.name).all()
    roles = roles_at(db, [employee.id for employee in employees])
    leaders = [employee for employee in employees if (roles.get(employee.id) or None) and roles[employee.id].code in LEADER_CODES]
    led_groups = groups_led_by_bulk(db, [leader.id for leader in leaders])
    group_display = group_display_metadata_bulk(db, {group.id for rows in led_groups.values() for group in rows})
    gsm_by_attraction = gsm_candidates_for_attractions_bulk(db, {leader.attraction_id for leader in leaders if leader.attraction_id})
    attraction_names = {attraction.id: attraction.name for attraction in db.query(Attraction).all()}
    result = []
    for employee in leaders:
        role = roles[employee.id]
        groups = [group for group in led_groups.get(employee.id, []) if group.attraction_id == employee.attraction_id]
        if not groups:
            groups = [None]
        candidates = gsm_by_attraction.get(employee.attraction_id) if employee.attraction_id else None
        gsm, gsm_role = candidates[0] if candidates else (None, None)
        for group in groups:
            result.append(
                {
                    "id": employee.id,
                    "name": employee.name,
                    "employee_no": employee.employee_no,
                    "role_name": role.name,
                    "attraction_id": employee.attraction_id,
                    "attraction_name": attraction_names.get(employee.attraction_id, "") if employee.attraction_id else "",
                    "group_id": group.id if group else None,
                    "group_name": group_display.get(group.id, {}).get("name", group.name) if group else "新建工作组",
                    "gsm_id": gsm.id if gsm else None,
                    "gsm_name": gsm.name if gsm else "未配置GSM",
                    "gsm_role_name": gsm_role.name if gsm_role else "",
                }
            )
    return result


@router.post("/hr/groups/{group_id}/transfer")
def transfer_group(group_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    group = db.get(WorkGroup, group_id)
    if not group:
        raise HTTPException(404, "工作组不存在")
    ensure_scoped_hr_attraction(db, user, group.attraction_id)
    if int(payload.get("revision") or 0) != group.revision:
        raise HTTPException(409, "工作组已被其他操作修改，请刷新后重试")
    new_leader = db.get(Employee, int(payload.get("new_leader_id") or 0))
    new_role = role_at(db, new_leader.id) if new_leader else None
    if not new_leader or not new_leader.is_active or not new_role or new_role.code not in LEADER_CODES:
        raise HTTPException(400, "新组长必须是在职TA主管或主管")
    if new_leader.attraction_id != group.attraction_id:
        raise HTTPException(400, "新组长必须与工作组属于同一景点圈")
    effective_date = str(payload.get("effective_date") or date.today().isoformat())
    if effective_date != date.today().isoformat():
        raise HTTPException(400, "当前版本整组移交仅支持当天生效")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "移交原因必填")
    old_assignment = active_group_leader(db, group.id)
    if old_assignment and old_assignment.leader_employee_id == new_leader.id:
        raise HTTPException(400, "新旧组长不能相同")
    members = active_group_memberships(db, group.id)
    pending_count = db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id.in_([member.employee_id for member in members]) if members else RecognitionRecord.id == -1,
        RecognitionRecord.status == "pending",
    ).count()
    transfer = GroupTransfer(
        group_id=group.id,
        old_leader_id=old_assignment.leader_employee_id if old_assignment else None,
        new_leader_id=new_leader.id,
        attraction_id=group.attraction_id,
        effective_date=effective_date,
        member_count=len(members),
        pending_review_count=pending_count,
        reason=reason,
        operator_id=user.id,
    )
    db.add(transfer)
    db.flush()
    if old_assignment:
        old_assignment.status = "ended"
        old_assignment.ends_on = effective_date
    db.add(GroupLeaderAssignment(group_id=group.id, leader_employee_id=new_leader.id, starts_on=effective_date, status="active", transfer_id=transfer.id))
    for member in members:
        db.add(GroupTransferMember(transfer_id=transfer.id, employee_id=member.employee_id, employee_no=member.employee.employee_no, employee_name=member.employee.name))
    member_ids = [member.employee_id for member in members]
    if member_ids:
        db.query(RecognitionRecord).filter(RecognitionRecord.employee_id.in_(member_ids), RecognitionRecord.status == "pending").update({RecognitionRecord.assigned_reviewer_id: new_leader.id}, synchronize_session=False)
    group.status = "active"
    group.revision += 1
    db.query(SystemAlert).filter(SystemAlert.group_id == group.id, SystemAlert.status == "open").update({SystemAlert.status: "handled", SystemAlert.handled_by: user.id, SystemAlert.handled_at: datetime.now()}, synchronize_session=False)
    write_audit(db, user.employee, "整组移交", "work_group", group.id, before={"leader": old_assignment.leader.name if old_assignment else "待接管"}, after={"leader": new_leader.name, "member_count": len(members)}, reason=reason, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "transfer_id": transfer.id}


@router.get("/hr/alerts")
def hr_alerts(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    process_role_expirations(db)
    rows = visible_system_alerts(db, user)
    return [
        {"id": row.id, "type": row.alert_type, "message": row.message, "due_date": row.due_date or "", "status": row.status, "created_at": row.created_at.strftime("%Y-%m-%d %H:%M:%S")}
        for row in rows
    ]


@router.get("/hr/score-rules")
def hr_score_rules(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    result = []
    for role_code in RECOGNIZER_CODES:
        role = db.query(Role).filter(Role.code == role_code).one()
        rule = db.query(RecognitionScoreRule).filter_by(role_id=role.id, active=True).order_by(RecognitionScoreRule.effective_date.desc()).first()
        result.append({"role_id": role.id, "role_code": role.code, "role_name": role.name, "score": float(rule.score) if rule else 0, "effective_date": rule.effective_date if rule else ""})
    return sorted(result, key=lambda item: item["role_id"])


@router.post("/hr/score-rules")
def update_score_rule(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    if user.role.code != "SYSTEM_ADMIN":
        raise HTTPException(403, "认可角色分值是全局规则，仅最高管理员可调整")
    role = db.query(Role).filter(Role.code == str(payload.get("role_code") or "")).first()
    if not role or role.code not in RECOGNIZER_CODES:
        raise HTTPException(400, "认可人角色无效")
    try:
        score = Decimal(str(payload.get("score"))).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise HTTPException(400, "分值格式错误") from exc
    if not score.is_finite() or score < 0:
        raise HTTPException(400, "分值必须是大于或等于0的有效数字")
    effective_date = str(payload.get("effective_date") or date.today().isoformat())
    parse_iso_date(effective_date, "生效日期")
    rule = db.query(RecognitionScoreRule).filter_by(role_id=role.id, effective_date=effective_date).first()
    before = {"score": float(rule.score), "active": rule.active} if rule else None
    if rule:
        rule.score = score
        rule.active = True
    else:
        rule = RecognitionScoreRule(role_id=role.id, score=score, effective_date=effective_date, active=True)
        db.add(rule)
    write_audit(db, user.employee, "修改认可角色分值", "recognition_score_rule", role.id, before=before, after={"role": role.name, "score": float(score), "effective_date": effective_date}, ip_address=client_ip(request))
    db.commit()
    return {"ok": True}


@router.get("/admin/logs")
def admin_logs(limit: int = Query(300, ge=1, le=1000), db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SYSTEM_ADMIN"))):
    return [
        {"id": row.id, "time": row.created_at.strftime("%Y-%m-%d %H:%M:%S"), "operator": row.operator_name or "系统", "action": row.action, "entity": f"{row.entity_type}:{row.entity_id or ''}", "reason": row.reason or ""}
        for row in db.query(AuditLog).order_by(AuditLog.created_at.desc(), AuditLog.id.desc()).limit(limit).all()
    ]


@router.get("/admin/circle-hr-accounts")
def circle_hr_accounts(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SYSTEM_ADMIN"))):
    """Show the three scoped HR accounts without exposing stored passwords."""
    login_accounts = [item[0] for item in CIRCLE_HR_ACCOUNTS]
    employees = {
        employee.employee_no: employee
        for employee in db.query(Employee).filter(Employee.employee_no.in_(login_accounts)).all()
    }
    employee_ids = [employee.id for employee in employees.values()]
    accounts = {
        account.employee_id: account
        for account in db.query(UserAccount).filter(UserAccount.employee_id.in_(employee_ids)).all()
    } if employee_ids else {}
    items = []
    for login_account, name, circle_name in CIRCLE_HR_ACCOUNTS:
        employee = employees.get(login_account)
        account = accounts.get(employee.id) if employee else None
        items.append(
            {
                "employee_id": employee.id if employee else None,
                "login_account": login_account,
                "name": name,
                "attraction_name": circle_name,
                "account_enabled": bool(account and account.enabled),
                "must_change_password": bool(account.must_change_password) if account else True,
                "last_login_at": account.last_login_at.strftime("%Y-%m-%d %H:%M:%S") if account and account.last_login_at else "",
                "password_status": "待首次修改" if not account or account.must_change_password else "已设置（不可读取）",
            }
        )
    return {"items": items}


@router.post("/admin/circle-hr-accounts/{employee_id}/reset-password")
def reset_circle_hr_password(
    employee_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("SYSTEM_ADMIN")),
):
    employee = db.get(Employee, employee_id)
    role = role_at(db, employee.id) if employee else None
    if not employee or not role or role.code != "HR_CIRCLE":
        raise HTTPException(404, "景点圈HR账号不存在")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if not account:
        raise HTTPException(400, "该景点圈HR尚未开通登录账号")
    reset_password = resolved_reset_password(account)
    account.password_hash = hash_password(reset_password)
    account.failed_attempts = 0
    account.locked_until = None
    account.must_change_password = True
    account.credential_initialized = True
    account.password_changed_at = None
    account.enabled = True
    db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "重置景点圈HR密码",
        "user_account",
        account.id,
        after={
            "employee_no": employee.employee_no,
            "employee_name": employee.name,
            "attraction_name": employee.attraction.name if employee.attraction else "",
            "password_rule": "登录账号后四位",
            "must_change_password": True,
        },
        reason="最高管理员重置景点圈HR账号密码",
        ip_address=client_ip(request),
    )
    db.commit()
    return {
        "ok": True,
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "temporary_password": reset_password,
        "must_change_password": True,
        "message": "密码已重置为登录账号后四位，请要求账号本人登录后立即修改。",
    }

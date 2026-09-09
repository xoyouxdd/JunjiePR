"""Workbook builders for ranking, statistics and HR import templates.

Routes in routers.v2 fetch data, call these builders, then watermark and return.
"""
from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.excel_export_utils import (
    add_banded_table,
    apply_date_format,
    apply_score_format,
    configure_sheet,
    excel_row,
    excel_safe_value,
    style_sheet,
    write_kpi_block,
    write_label_value_pairs,
    write_merged_title,
)
from app.v2_models import Attraction, Employee, EmployeeNumberHistory, Role
from app.v2_services import roles_at


def ordered_statistics_score_rows(data: dict) -> list[dict]:
    """Return score rows in the same order as the on-screen organization tree."""
    score_by_employee_id = {int(row["employee_id"]): row for row in data["scores"]}
    ordered_ids = [
        int(node["employee_id"])
        for node in data.get("hierarchy", [])
        if node.get("node_type") == "employee" and node.get("employee_id") is not None
    ]
    ordered = [score_by_employee_id.pop(employee_id) for employee_id in ordered_ids if employee_id in score_by_employee_id]
    # Keep a deterministic fallback for any historical row which no longer has a
    # current organization membership. It must remain exportable rather than vanish.
    ordered.extend(sorted(score_by_employee_id.values(), key=lambda row: (str(row.get("attraction_name") or ""), str(row.get("employee_name") or ""), str(row.get("employee_no") or ""))))
    return ordered


def employee_number_history_context(db: Session, employee_ids: list[int]) -> dict[int, list[EmployeeNumberHistory]]:
    if not employee_ids:
        return {}
    rows = (
        db.query(EmployeeNumberHistory)
        .filter(EmployeeNumberHistory.employee_id.in_(employee_ids))
        .order_by(EmployeeNumberHistory.employee_id, EmployeeNumberHistory.effective_on, EmployeeNumberHistory.id)
        .all()
    )
    result: dict[int, list[EmployeeNumberHistory]] = {}
    for row in rows:
        result.setdefault(row.employee_id, []).append(row)
    return result


def employee_number_status(employee_no: str, history: list[EmployeeNumberHistory] | None) -> str:
    if not history:
        return "当前工号"
    old_values = "、".join(dict.fromkeys(row.old_employee_no for row in history))
    return f"已变更：{old_values} → {employee_no}"


def employee_number_at_date(current_no: str, history: list[EmployeeNumberHistory] | None, occurred_on: str | None) -> str:
    """Resolve a visible historical number for exports; identity remains employee_id."""
    if not history or not occurred_on:
        return current_no
    value = str(occurred_on)[:10]
    number = current_no
    for change in reversed(history):
        if value < change.effective_on:
            number = change.old_employee_no
        else:
            break
    return number


def write_employee_number_history_sheet(wb: Workbook, db: Session, employee_ids: list[int]) -> None:
    """Add a review-only mapping sheet without exposing HR-only change reasons."""
    ws = wb.create_sheet("员工号变更对照")
    ws.append(["当前员工号", "历史员工号", "姓名", "生效日期", "变更状态", "说明"])
    histories = employee_number_history_context(db, employee_ids)
    employees = {row.id: row for row in db.query(Employee).filter(Employee.id.in_(list(histories))).all()} if histories else {}
    for employee_id, changes in histories.items():
        employee = employees.get(employee_id)
        if not employee:
            continue
        for change in changes:
            ws.append(excel_row([employee.employee_no, change.old_employee_no, employee.name, change.effective_on, "已变更", "历史记录与当前档案为同一员工"]))
    if ws.max_row == 1:
        ws.append(["", "", "", "", "无", "本次导出范围内无员工号变更记录"])
    style_sheet(ws, landscape=True)


def write_monthly_score_row(ws, row: dict, role, status: str = "在职") -> int:
    status = "账号已删除·留档" if row.get("account_deleted_at") else status
    participating = not status.startswith("LOA")
    ws.append(excel_row([
        row["employee_no"],
        row["employee_name"],
        role.code if role else "",
        row.get("attraction_name") or "",
        row.get("recognition_score") if participating else "",
        row.get("attendance_score") if participating else "",
        row.get("deduction_score") if participating else "",
        row.get("total_score") if participating else "",
        status,
        row.get("employee_number_status", "当前工号"),
    ]))
    return ws.max_row


def write_organization_score_sheet(ws, data: dict, export_roles: dict[int, Role], loa_roles: dict[int, Role]) -> None:
    """Write the monthly score view as a collapsible Excel organization tree."""
    ws.append(["组织 / 员工", "员工号", "Title", "景点圈", "签卡加分", "全勤分", "扣分", "综合分", "人员状态", "工号状态"])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    attraction_fill = PatternFill("solid", fgColor="EAF3FE")
    gsm_fill = PatternFill("solid", fgColor="F0F5FA")
    supervisor_fill = PatternFill("solid", fgColor="EBF8F0")
    employee_fill = PatternFill("solid", fgColor="FFF9E6")
    loa_fill = PatternFill("solid", fgColor="F5F5F5")
    score_by_employee_id = {int(row["employee_id"]): row for row in data["scores"]}
    attraction_names = {str(node["node_id"]): node["name"] for node in data.get("hierarchy", []) if node.get("node_type") == "attraction"}
    supervisor_rows: list[tuple[int, list[int]]] = []
    current_supervisor_row: int | None = None
    current_employee_rows: list[int] = []

    def finish_supervisor() -> None:
        nonlocal current_supervisor_row, current_employee_rows
        if current_supervisor_row and current_employee_rows:
            ws.row_dimensions[current_supervisor_row].collapsed = True
            for employee_row in current_employee_rows:
                dimensions = ws.row_dimensions[employee_row]
                dimensions.outlineLevel = 3
                dimensions.hidden = True
            supervisor_rows.append((current_supervisor_row, current_employee_rows))
        current_supervisor_row = None
        current_employee_rows = []

    for node in data.get("hierarchy", []):
        node_type = node.get("node_type")
        if node_type == "attraction":
            finish_supervisor()
            ws.append([f"景点圈：{node['name']}"])
            row_number = ws.max_row
            ws.row_dimensions[row_number].outlineLevel = 0
            for cell in ws[row_number]:
                cell.fill = attraction_fill
                cell.font = Font(bold=True, color="1F4E78")
            continue
        if node_type == "gsm":
            finish_supervisor()
            ws.append([f"GSM层：{node['name']}（{node.get('role_name') or 'GSM'}）"])
            row_number = ws.max_row
            ws.row_dimensions[row_number].outlineLevel = 1
            for cell in ws[row_number]:
                cell.fill = gsm_fill
                cell.font = Font(bold=True, color="38536D")
            continue
        if node_type == "gsm_team":
            finish_supervisor()
            ws.append([f"主管组（由GSM共同承接）：{node['name']}"])
            row_number = ws.max_row
            ws.row_dimensions[row_number].outlineLevel = 1
            for cell in ws[row_number]:
                cell.fill = gsm_fill
                cell.font = Font(bold=True, color="38536D")
            continue
        if node_type == "supervisor":
            finish_supervisor()
            ws.append([
                f"主管组：{node['name']}（{node.get('role_name') or '主管'}；共{node.get('member_count') or 0}人）",
                "",
                "",
                "",
                node.get("recognition_score") or 0,
                node.get("attendance_score") or 0,
                node.get("deduction_score") or 0,
                node.get("total_score") or 0,
                "小组汇总",
                "",
            ])
            current_supervisor_row = ws.max_row
            ws.row_dimensions[current_supervisor_row].outlineLevel = 2
            for cell in ws[current_supervisor_row]:
                cell.fill = supervisor_fill
                cell.font = Font(bold=True, color="276749")
            continue
        if node_type != "employee":
            continue
        score = score_by_employee_id.get(int(node["employee_id"]))
        if not score:
            continue
        row_number = write_monthly_score_row(ws, score, export_roles.get(int(node["employee_id"])))
        current_employee_rows.append(row_number)
        for cell in ws[row_number]:
            cell.fill = employee_fill

    finish_supervisor()

    if data.get("loa_rows"):
        ws.append(["整月 LOA（未参与计分）"])
        loa_header_row = ws.max_row
        for cell in ws[loa_header_row]:
            cell.fill = loa_fill
            cell.font = Font(bold=True, color="666666")
        for row in sorted(data["loa_rows"], key=lambda item: (str(item.get("attraction_name") or ""), str(item.get("employee_name") or ""), str(item.get("employee_no") or ""))):
            row_number = write_monthly_score_row(ws, row, loa_roles.get(int(row["employee_id"])), "LOA（长期病假）")
            ws.row_dimensions[row_number].outlineLevel = 1
            for cell in ws[row_number]:
                cell.fill = loa_fill

    for column in (5, 6, 7, 8):
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "0.00"
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_view.showOutlineSymbols = True
    style_sheet(ws, landscape=True)
    # The tree headings are navigation rows, not records. Filtering remains available
    # in the flat detail sheet where it cannot hide structural context.
    ws.auto_filter.ref = None
    # Reapply the tree heading styling because style_sheet owns the ordinary header only.
    for cell in ws[1]:
        cell.fill = header_fill


def write_monthly_score_detail_sheet(ws, rows: list[dict], loa_rows: list[dict], export_roles: dict[int, Role], loa_roles: dict[int, Role]) -> None:
    """Keep the previous flat, filterable worksheet for editing and ad-hoc searches."""
    ws.append(["员工号", "姓名", "Title", "景点圈", "签卡加分", "全勤分", "扣分", "综合分", "人员状态", "工号状态"])
    for row in rows:
        write_monthly_score_row(ws, row, export_roles.get(int(row["employee_id"])))
    for row in loa_rows:
        write_monthly_score_row(ws, row, loa_roles.get(int(row["employee_id"])), "LOA（长期病假）")
    for col in (5, 6, 7, 8):
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            cell[0].number_format = "0.00"
    style_sheet(ws, landscape=True)


def write_hierarchical_performance_sheet(ws, data: dict) -> None:
    """Write an editable, collapsible tree whose upper rows are formulas over record rows."""
    ws.append(["层级 / 明细", "日期", "类别", "对象", "详情", "加分", "全勤分", "扣分", "综合分", "景点圈键", "GSM键", "组长键", "员工ID", "行类型", "明细类型"])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    fills = {
        "attraction": PatternFill("solid", fgColor="EAF3FE"),
        "gsm": PatternFill("solid", fgColor="F0F5FA"),
        "supervisor": PatternFill("solid", fgColor="EBF8F0"),
        "employee": PatternFill("solid", fgColor="FFF9E6"),
        "category": PatternFill("solid", fgColor="FFFDF2"),
        "detail": PatternFill("solid", fgColor="FFFFFF"),
    }
    duplicate_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
    contexts: dict[int, dict[str, str]] = {}
    attraction_key = gsm_key = leader_key = ""
    for node in data.get("hierarchy", []):
        kind = node.get("node_type")
        if kind == "attraction":
            attraction_key = f"ATTRACTION:{node.get('node_id')}"
            gsm_key = leader_key = ""
        elif kind in {"gsm", "gsm_team"}:
            gsm_key = f"GSM:{node.get('node_id')}"
            leader_key = ""
        elif kind == "supervisor":
            leader_key = f"LEADER:{node.get('node_id')}"
        elif kind == "employee" and node.get("employee_id") is not None:
            contexts[int(node["employee_id"])] = {"attraction": attraction_key, "gsm": gsm_key, "leader": leader_key}

    records: dict[int, list[dict]] = {}
    def add_record(employee_id: int, record: dict) -> None:
        records.setdefault(int(employee_id), []).append(record)

    for row in data.get("recognitions", []):
        cap_note = row.get("monthly_cap_reason") or ""
        add_record(row["employee_id"], {"date": row["recognition_date"], "type": "加分", "object": row["recognition_type"], "detail": f"认可人：{row['recognizer_name']}；{row['content']}" + (f"；{cap_note}" if cap_note else ""), "add": float(row.get("credited_fraction", row["fraction"])) if row["status"] == "confirmed" else 0, "attendance": 0, "deduction": 0, "status": row["status_name"], "duplicate": bool(row.get("same_day_duplicate")), "duplicate_note": (f"同日重复：第{int(row.get('same_day_duplicate_sequence') or 0)}次" if row.get("same_day_duplicate") else "")})
    for row in data.get("deductions", []):
        upgrade_role = row.get("upgrade_role") or ""
        upgrade_state = row.get("upgrade_state") or ""
        add_record(row["employee_id"], {"date": row["occurred_on"], "type": "扣分", "object": f"{row['deduction_type']} · {row['deduction_level']}", "detail": row["description"], "add": 0, "attendance": 0, "deduction": -abs(float(row["points"])) if row["status"] == "active" else 0, "status": row["status_name"], "duplicate": False, "duplicate_note": "", "upgrade_role": upgrade_role, "upgrade_state": upgrade_state, "upgrade_request_id": row.get("upgrade_request_id") or 0})
    for row in data.get("sick_leaves", []):
        add_record(row["employee_id"], {"date": f"{row['leave_start_date']} 至 {row['leave_end_date']}", "type": "缺勤", "object": f"{float(row['leave_days']):.1f}天", "detail": row.get("note") or "无备注", "add": 0, "attendance": 0, "deduction": 0, "status": row["status_name"], "duplicate": False, "duplicate_note": ""})
    for row in data.get("scores", []):
        add_record(row["employee_id"], {"date": row.get("month") or data.get("month") or "", "type": "全勤", "object": "当月全勤分", "detail": "可编辑全勤分后将联动上层汇总", "add": 0, "attendance": float(row.get("attendance_score") or 0), "deduction": 0, "status": "计入综合分", "duplicate": False, "duplicate_note": ""})

    def formula_for(row_number: int, scope: str) -> None:
        column = {"attraction": "J", "gsm": "K", "leader": "L", "employee": "M"}[scope]
        for output_column, detail_type in (("F", "加分"), ("G", "全勤"), ("H", "扣分")):
            ws[f"{output_column}{row_number}"] = f'=SUMIFS(${output_column}:${output_column},${column}:${column},${column}{row_number},$N:$N,"明细",$O:$O,"{detail_type}")'
        ws[f"I{row_number}"] = f"=SUM(F{row_number}:H{row_number})"

    def append_header(label: str, level: int, scope: str, context: dict[str, str], kind: str, hidden: bool = False) -> int:
        ws.append([label, "", "", "", "", "", "", "", "", context["attraction"], context["gsm"], context["leader"], context.get("employee", ""), "汇总", ""])
        row_number = ws.max_row
        formula_for(row_number, scope)
        ws.row_dimensions[row_number].outlineLevel = level
        ws.row_dimensions[row_number].hidden = hidden
        for cell in ws[row_number]:
            cell.fill = fills[kind]
            cell.font = Font(bold=True)
        return row_number

    current_context = {"attraction": "", "gsm": "", "leader": "", "employee": ""}
    for node in data.get("hierarchy", []):
        kind = node.get("node_type")
        if kind == "attraction":
            current_context = {"attraction": f"ATTRACTION:{node.get('node_id')}", "gsm": "", "leader": "", "employee": ""}
            append_header(f"景点圈：{node['name']}", 0, "attraction", current_context, "attraction")
        elif kind in {"gsm", "gsm_team"}:
            current_context["gsm"] = f"GSM:{node.get('node_id')}"
            current_context["leader"] = current_context["employee"] = ""
            append_header(f"GSM层：{node['name']}（{node.get('role_name') or 'GSM'}）", 1, "gsm", current_context, "gsm")
        elif kind == "supervisor":
            current_context["leader"] = f"LEADER:{node.get('node_id')}"
            current_context["employee"] = ""
            row_number = append_header(f"组长：{node['name']}（{node.get('role_name') or '主管'}；{node.get('member_count') or 0}人）", 2, "leader", current_context, "supervisor")
            ws.row_dimensions[row_number].collapsed = True
        elif kind == "employee" and node.get("employee_id") is not None:
            employee_id = int(node["employee_id"])
            context = dict(contexts.get(employee_id, current_context))
            context["employee"] = str(employee_id)
            employee_row = append_header(f"组员：{node['employee_name']}（{node.get('role_name') or node.get('role_code') or ''}）", 3, "employee", context, "employee", hidden=True)
            ws.row_dimensions[employee_row].collapsed = True
            for category in ("加分", "扣分", "缺勤", "全勤"):
                category_rows = [item for item in records.get(employee_id, []) if item["type"] == category]
                ws.append([f"{category}明细（{len(category_rows)}条）", "", "", "", "", "", "", "", "", context["attraction"], context["gsm"], context["leader"], str(employee_id), "分类", category])
                category_row = ws.max_row
                for output_column, detail_type in (("F", "加分"), ("G", "全勤"), ("H", "扣分")):
                    ws[f"{output_column}{category_row}"] = f'=SUMIFS(${output_column}:${output_column},$M:$M,$M{category_row},$N:$N,"明细",$O:$O,"{detail_type}")'
                ws[f"I{category_row}"] = f"=SUM(F{category_row}:H{category_row})"
                ws.row_dimensions[category_row].outlineLevel = 4
                ws.row_dimensions[category_row].hidden = True
                ws.row_dimensions[category_row].collapsed = True
                for cell in ws[category_row]:
                    cell.fill = fills["category"]
                    cell.font = Font(bold=True)
                for item in category_rows:
                    detail = item["detail"] + (f"；{item['duplicate_note']}" if item["duplicate_note"] else "")
                    ws.append(["记录详情", item["date"], item["type"], excel_safe_value(item["object"]), excel_safe_value(detail), item["add"], item["attendance"], item["deduction"], f"=SUM(F{ws.max_row + 1}:H{ws.max_row + 1})", context["attraction"], context["gsm"], context["leader"], str(employee_id), "明细", item["type"]])
                    detail_row = ws.max_row
                    ws.row_dimensions[detail_row].outlineLevel = 5
                    ws.row_dimensions[detail_row].hidden = True
                    if item["duplicate"]:
                        for cell in ws[detail_row]:
                            cell.fill = duplicate_fill
                    if item.get("upgrade_role") == "source_first":
                        for cell in ws[detail_row]: cell.fill = PatternFill("solid", fgColor="DDEBF7")
                    elif item.get("upgrade_role") == "source_second":
                        for cell in ws[detail_row]: cell.fill = PatternFill("solid", fgColor="FFF2CC")
                    elif item.get("upgrade_role") == "result":
                        for cell in ws[detail_row]: cell.fill = PatternFill("solid", fgColor="E2F0D9")
                    elif item.get("upgrade_state") in {"pending", "rejected"}:
                        for cell in ws[detail_row]: cell.fill = PatternFill("solid", fgColor="FCE4D6" if item.get("upgrade_state") == "pending" else "E7E6E6")

    for column in range(6, 10):
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "0.00"
    for column in ("J", "K", "L", "M", "N", "O"):
        ws.column_dimensions[column].hidden = True
    ws.freeze_panes = "A2"
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_view.showOutlineSymbols = True
    for column in ws.columns:
        letter = column[0].column_letter
        if letter not in {"J", "K", "L", "M", "N", "O"}:
            ws.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
    configure_sheet(ws, freeze="A2", landscape=True)

def build_pr_rankings_workbook(db: Session, data: dict, category: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "PR排名数据"
    ranking_employee_ids = [int(row["employee_id"]) for row in data["rows"] if row.get("employee_id") is not None]
    ranking_employees = {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(ranking_employee_ids)).all()} if ranking_employee_ids else {}
    ranking_histories = employee_number_history_context(db, ranking_employee_ids)
    for row in data["rows"]:
        employee_id = row.get("employee_id")
        if employee_id is None:
            row["employee_number_status"] = "当前工号"
            continue
        employee = ranking_employees.get(int(employee_id))
        current_no = employee.employee_no if employee else row["employee_no"]
        row["employee_no"] = current_no
        row["employee_number_status"] = employee_number_status(current_no, ranking_histories.get(int(employee_id)))
    if category == "overall":
        headers = ["排名", "员工号", "工号状态", "姓名", "角色", "主管", "加分", "扣分", "全勤分", "综合分"]
        ws.append(headers)
        for row in data["rows"]:
            ws.append(excel_row([row["rank"], row["employee_no"], row["employee_number_status"], row["employee_name"], row["role_name"], row["leader_name"], row["recognition_score"], row["deduction_score"], row["attendance_score"], row["total_score"]]))
    elif category == "absence":
        headers = ["排名", "员工号", "工号状态", "姓名", "角色", "主管", "缺勤类型", "登记次数", "缺勤天数", "计费天数", "扣减全勤分", "最近一次缺勤"]
        ws.append(headers)
        for row in data["rows"]:
            ws.append(excel_row([row["rank"], row["employee_no"], row["employee_number_status"], row["employee_name"], row["role_name"], row["leader_name"], data["subtype_name"], row["count"], row["leave_days"], row["charged_days"], row["score"], row["recent_date"]]))
    else:
        person_label = "GSM/TA GSM" if category == "gsm_leader" else ("组长" if category == "leader" else "员工")
        headers = ["排名", "员工号", "工号状态", person_label, "角色", "主管", "类型", "登记次数", "累计分值", "最近一次登记"]
        ws.append(headers)
        for row in data["rows"]:
            ws.append(excel_row([row["rank"], row["employee_no"], row["employee_number_status"], row["employee_name"], row["role_name"], row["leader_name"], data["subtype_name"], row["count"], row["score"], row["recent_date"]]))
    style_sheet(ws, landscape=True)
    add_banded_table(ws, ws.dimensions, "PrRankingData")
    if category == "overall":
        apply_score_format(ws, (7, 8, 9, 10))
    elif category == "absence":
        apply_score_format(ws, (11,))
        apply_date_format(ws, (12,))
    else:
        apply_score_format(ws, (9,))
        apply_date_format(ws, (10,))
    write_employee_number_history_sheet(wb, db, ranking_employee_ids)
    return wb


def build_statistics_workbook(
    db: Session,
    data: dict,
    *,
    month: str,
    attraction_id: int | None,
    keyword: str | None,
    exporter_no: str,
    exporter_name: str,
    exporter_role_name: str,
    exporter_role_code: str,
) -> tuple[Workbook, str, str]:
    month_start = date.fromisoformat(f"{month}-01")
    month_end = month_start.replace(day=monthrange(month_start.year, month_start.month)[1]).isoformat()
    export_roles = roles_at(db, [int(row["employee_id"]) for row in data["scores"]], month_end)
    recognition_counts: dict[int, int] = {}
    for row in data["recognitions"]:
        if row["status"] == "confirmed":
            recognition_counts[row["employee_id"]] = recognition_counts.get(row["employee_id"], 0) + 1
    wb = Workbook()
    # The hierarchy sheet deliberately rolls up editable record rows with formulas.
    # Ask Excel/WPS to recalculate those formulas immediately after opening the file.
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    exported_at = datetime.now()
    selected_attraction = db.get(Attraction, attraction_id) if attraction_id else None
    attraction_label = selected_attraction.name if selected_attraction else "全部可见景点圈"
    title_label = data["title"] or "CM/TR全部"
    keyword_label = (keyword or "").strip() or "无"
    ws0 = wb.active
    ws0.title = "导出说明"
    write_merged_title(ws0, f"认可数据 · {month} · {attraction_label}", 2)
    next_row = write_label_value_pairs(
        ws0,
        [
            ("统计月份", month),
            ("景点圈", attraction_label),
            ("Title", title_label),
            ("员工搜索", keyword_label),
            ("导出账号", f"{exporter_no} · {exporter_name}"),
            ("导出账号角色", f"{exporter_role_name}（{exporter_role_code}）"),
            ("导出时间", exported_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("在职计分人数", len(data["scores"])),
            ("整月LOA人数", len(data["loa_rows"])),
            ("小组平均分口径", f"{title_label}筛选后小组综合分 ÷ 筛选后人数"),
        ],
        2,
    )
    scores = data["scores"]
    recognition_total = round(sum(float(row.get("recognition_score") or 0) for row in scores), 2)
    deduction_total = round(sum(float(row.get("deduction_score") or 0) for row in scores), 2)
    if not scores:
        recognition_total = round(sum(float(row.get("credited_fraction", row.get("fraction") or 0)) for row in data["recognitions"] if row.get("status") == "confirmed"), 2)
        deduction_total = round(sum(float(row.get("points") or 0) for row in data["deductions"] if row.get("status") == "active"), 2)
    attendance_values = [float(row.get("attendance_score") or 0) for row in scores]
    perfect_threshold = None
    for row in scores:
        base = row.get("base_score")
        bonus = row.get("perfect_bonus")
        if base is not None and bonus is not None:
            perfect_threshold = max(perfect_threshold or 0.0, float(base) + float(bonus))
        elif bonus is not None and float(bonus) > 0:
            max_att = max(attendance_values) if attendance_values else 0.0
            perfect_threshold = max(perfect_threshold or 0.0, max_att)
    if perfect_threshold is None:
        max_att = max(attendance_values) if attendance_values else 0.0
        if max_att >= 11.99:
            perfect_threshold = 11.99
    if perfect_threshold is not None:
        perfect_count = sum(1 for value in attendance_values if value >= perfect_threshold)
    else:
        sick_ids = {int(row["employee_id"]) for row in data["sick_leaves"] if row.get("status") != "void" and row.get("employee_id") is not None}
        perfect_count = sum(1 for row in scores if float(row.get("attendance_score") or 0) > 0 and int(row["employee_id"]) not in sick_ids)
    top3 = sorted(scores, key=lambda row: (-float(row.get("total_score") or 0), str(row.get("employee_no") or "")))[:3]
    top3_label = "、".join(f"{row.get('employee_name')} {float(row.get('total_score') or 0):.2f}" for row in top3) or "无"
    write_kpi_block(
        ws0,
        [
            ("加分合计", recognition_total),
            ("扣分合计", deduction_total),
            ("全勤人数", perfect_count),
            ("综合分 Top3", top3_label),
        ],
        next_row + 1,
    )
    configure_sheet(ws0, freeze="A2", landscape=False)
    loa_role_ids = [int(row["employee_id"]) for row in data["loa_rows"]]
    loa_roles = roles_at(db, loa_role_ids, month_end) if loa_role_ids else {}
    ordered_scores = ordered_statistics_score_rows(data)
    export_employee_ids = sorted(
        {
            int(row["employee_id"])
            for bucket in (data["scores"], data["loa_rows"], data["recognitions"], data["deductions"], data["sick_leaves"])
            for row in bucket
            if row.get("employee_id") is not None
        }
    )
    export_employees = {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(export_employee_ids)).all()} if export_employee_ids else {}
    number_histories = employee_number_history_context(db, export_employee_ids)
    for row in [*data["scores"], *data["loa_rows"]]:
        current_no = export_employees.get(int(row["employee_id"]), None)
        row["employee_number_status"] = employee_number_status(current_no.employee_no if current_no else row["employee_no"], number_histories.get(int(row["employee_id"])))
    for bucket, date_key in ((data["recognitions"], "recognition_date"), (data["deductions"], "occurred_on"), (data["sick_leaves"], "leave_start_date")):
        for row in bucket:
            employee = export_employees.get(int(row["employee_id"]))
            current_no = employee.employee_no if employee else row["employee_no"]
            row["current_employee_no"] = current_no
            row["export_employee_no"] = employee_number_at_date(current_no, number_histories.get(int(row["employee_id"])), row.get(date_key))
            row["employee_number_status"] = employee_number_status(current_no, number_histories.get(int(row["employee_id"])))
    ws = wb.create_sheet("月度综合分")
    write_organization_score_sheet(ws, data, export_roles, loa_roles)
    ws_monthly_detail = wb.create_sheet("月度综合分明细")
    write_monthly_score_detail_sheet(ws_monthly_detail, ordered_scores, data["loa_rows"], export_roles, loa_roles)
    ws_hierarchy = wb.create_sheet("层级绩效明细")
    write_hierarchical_performance_sheet(ws_hierarchy, data)
    write_employee_number_history_sheet(wb, db, export_employee_ids)
    wb.active = wb.sheetnames.index(ws_hierarchy.title)
    ws2 = wb.create_sheet("签卡明细")
    ws2.append(["记录编号", "日期", "员工号", "姓名", "类型", "认可人", "原始分值", "实际计入分值", "封顶状态/原因", "POC周期", "POC原因", "内容", "状态", "录入人", "当月已确认认可次数", "同日重复登记", "撤回人", "撤回时间", "撤回原因", "发生时员工号", "当前员工号", "工号状态"])
    detail_locations: dict[str, tuple[str, int]] = {}
    for row in data["recognitions"]:
        record_no = f"REC-{row['id']}"
        duplicate_label = f"是（第{int(row.get('same_day_duplicate_sequence') or 0)}次）" if row.get("same_day_duplicate") else "否"
        cap_label = row.get("monthly_cap_reason") or ({"within_limit": "单项上限内", "legacy_not_limited": "生效日前历史记录"}.get(row.get("monthly_cap_status"), "不适用"))
        ws2.append(excel_row([record_no, row["recognition_date"], row["current_employee_no"], row["employee_name"], row["recognition_type"], row["recognizer_name"], row["fraction"], row.get("credited_fraction", row["fraction"]), cap_label, row.get("poc_period_key") or "", row.get("poc_reason") or "", row["content"], row["status_name"], row["operator_name"], recognition_counts.get(row["employee_id"], 0), duplicate_label, row["voided_by_name"], row["voided_at"], row["void_reason"], row["export_employee_no"], row["current_employee_no"], row["employee_number_status"]]))
        detail_locations[record_no] = (ws2.title, ws2.max_row)
    style_sheet(ws2, landscape=True)
    add_banded_table(ws2, ws2.dimensions, "RecognitionDetails")
    apply_score_format(ws2, (7, 8))
    apply_date_format(ws2, (2,))
    for row in data["recognitions"]:
        if row.get("same_day_duplicate"):
            detail_row = detail_locations[f"REC-{row['id']}"][1]
            for cell in ws2[detail_row]:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    ws3 = wb.create_sheet("扣分明细")
    ws3.append(["记录编号", "日期", "员工号", "姓名", "类型", "等级", "实际扣分", "登记人", "状态", "说明", "升级关联", "升级角色", "升级状态", "作废人", "作废时间", "作废原因", "发生时员工号", "当前员工号", "工号状态"])
    for row in data["deductions"]:
        record_no = f"DED-{row['id']}"
        upgrade_no = f"UPG-{int(row['upgrade_request_id']):06d}" if row.get("upgrade_request_id") else ""
        ws3.append(excel_row([record_no, row["occurred_on"], row["current_employee_no"], row["employee_name"], row["deduction_type"], row["deduction_level"], row["points"], row["submitter_name"], row["status_name"], row["description"], upgrade_no, row.get("upgrade_role") or "", row.get("upgrade_state_name") or "", row["voided_by_name"], row["voided_at"], row["void_reason"], row["export_employee_no"], row["current_employee_no"], row["employee_number_status"]]))
        detail_locations[record_no] = (ws3.title, ws3.max_row)
    style_sheet(ws3, landscape=True)
    add_banded_table(ws3, ws3.dimensions, "DeductionDetails")
    apply_score_format(ws3, (7,))
    apply_date_format(ws3, (2,))
    for row in data["deductions"]:
        detail_row = detail_locations[f"DED-{row['id']}"][1]
        color = {"source_first": "DDEBF7", "source_second": "FFF2CC", "result": "E2F0D9"}.get(row.get("upgrade_role") or "")
        if not color and (row.get("upgrade_state") or "") == "pending":
            color = "FCE4D6"
        if not color and (row.get("upgrade_state") or "") == "rejected":
            color = "E7E6E6"
        if color:
            for cell in ws3[detail_row]:
                cell.fill = PatternFill("solid", fgColor=color)
    ws4 = wb.create_sheet("病假明细")
    ws4.append(["记录编号", "开始日期", "结束日期", "员工号", "姓名", "实际天数", "计费天数", "登记人", "状态", "备注", "作废人", "作废时间", "作废原因", "发生时员工号", "当前员工号", "工号状态"])
    for row in data["sick_leaves"]:
        record_no = f"SICK-{row['id']}"
        ws4.append(excel_row([record_no, row["leave_start_date"], row["leave_end_date"], row["current_employee_no"], row["employee_name"], row["leave_days"], row["charged_days"], row["submitter_name"], row["status_name"], row["note"], row["voided_by_name"], row["voided_at"], row["void_reason"], row["export_employee_no"], row["current_employee_no"], row["employee_number_status"]]))
        detail_locations[record_no] = (ws4.title, ws4.max_row)
    style_sheet(ws4, landscape=True)
    add_banded_table(ws4, ws4.dimensions, "SickLeaveDetails")
    apply_date_format(ws4, (2, 3))
    ws5 = wb.create_sheet("作废操作记录")
    ws5.append(["记录类型", "记录编号", "原记录", "业务日期", "员工号", "姓名", "原登记人", "操作人", "操作时角色", "操作时管理范围", "操作时间", "操作原因", "作废前状态", "原分值/天数", "原始内容"])
    void_locations: dict[str, int] = {}
    for row in data["recognitions"]:
        if row["status"] == "void":
            record_no = f"REC-{row['id']}"
            ws5.append(excel_row(["签卡撤回", record_no, "查看原记录", row["recognition_date"], row["employee_no"], row["employee_name"], row["operator_name"], row["voided_by_name"], row["voided_by_role_name"] or "历史数据未记录", row["void_permission_scope"] or "历史数据未记录", row["voided_at"], row["void_reason"], row["voided_from_status"] or "历史数据未记录", f"{row['fraction']:.2f}分", f"{row['recognition_type']} · {row['content']}"]))
            void_locations[record_no] = ws5.max_row
    for row in data["deductions"]:
        if row["status"] == "void":
            record_no = f"DED-{row['id']}"
            ws5.append(excel_row(["扣分作废", record_no, "查看原记录", row["occurred_on"], row["employee_no"], row["employee_name"], row["submitter_name"], row["voided_by_name"], row["voided_by_role_name"] or "历史数据未记录", row["void_permission_scope"] or "历史数据未记录", row["voided_at"], row["void_reason"], row["voided_from_status"] or "历史数据未记录", f"-{row['points']:.2f}分", f"{row['deduction_type']} · {row['deduction_level']} · {row['description']}"]))
            void_locations[record_no] = ws5.max_row
    for row in data["sick_leaves"]:
        if row["status"] == "void":
            record_no = f"SICK-{row['id']}"
            ws5.append(excel_row(["病假作废", record_no, "查看原记录", row["leave_start_date"], row["employee_no"], row["employee_name"], row["submitter_name"], row["voided_by_name"], row["voided_by_role_name"] or "历史数据未记录", row["void_permission_scope"] or "历史数据未记录", row["voided_at"], row["void_reason"], row["voided_from_status"] or "历史数据未记录", f"{row['leave_days']:.1f}天（计费{row['charged_days']:.1f}天）", f"{row['leave_start_date']}至{row['leave_end_date']} · {row['note'] or '无备注'}"]))
            void_locations[record_no] = ws5.max_row
    style_sheet(ws5, landscape=True)
    void_fill = PatternFill("solid", fgColor="FDECEC")
    void_font = Font(color="9F1D1D")
    link_font = Font(color="0563C1", underline="single")
    for record_no, void_row_number in void_locations.items():
        detail_location = detail_locations.get(record_no)
        if detail_location:
            detail_sheet_name, detail_row_number = detail_location
            detail_sheet = wb[detail_sheet_name]
            for cell in detail_sheet[detail_row_number]:
                cell.fill = void_fill
                cell.font = void_font
            detail_sheet.cell(detail_row_number, 1).hyperlink = f"#'作废操作记录'!B{void_row_number}"
            detail_sheet.cell(detail_row_number, 1).font = link_font
            ws5.cell(void_row_number, 3).hyperlink = f"#'{detail_sheet_name}'!A{detail_row_number}"
            ws5.cell(void_row_number, 3).font = link_font
    return wb, attraction_label, title_label


def build_employee_import_template() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "员工导入"
    ws.append(["员工号", "姓名", "角色代码", "景点圈", "初始密码", "在职", "账号启用", "任职开始日", "任职结束日", "到期恢复角色代码"])
    ws.append(["1000001", "示例员工", "CM", "热力追踪", "0001", "是", "是", date.today().isoformat(), "", ""])
    ref = wb.create_sheet("填写说明")
    ref.append(["可用角色代码", "CM, TR, TA_SUPERVISOR, SUPERVISOR, TA_GSM, GSM, AM, OM"])
    ref.append(["可用景点圈", "热力追踪、矮人迷宫、小熊罐子"])
    ref.append(["说明", "员工号必须为7位纯数字，姓名必填；员工只能归属三个景点圈之一；密码留空默认员工号后4位。"])
    style_sheet(ws)
    style_sheet(ref)
    return wb

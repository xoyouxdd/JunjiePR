"""Fill approved OOXML templates without a desktop Office/runtime dependency.

Native DrawingML tables and charts remain editable, including chart workbooks.
Only parts reachable from the final presentation survive (no hidden sample data).
"""
from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import posixpath
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships", "c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "rel": "http://schemas.openxmlformats.org/package/2006/relationships", "ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)
ROOT = Path(__file__).parent / "report_templates" / "hr_monthly"


def tag(prefix, name):
    return f"{{{NS[prefix]}}}{name}"


def xml(root):
    value = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    # OPC's package readers expect unprefixed Relationships/Types roots.
    for prefix in ("ct", "rel"):
        if root.tag.startswith("{" + NS[prefix] + "}"):
            value = value.replace(f"xmlns:{prefix}=".encode(), b"xmlns=").replace(f"<{prefix}:".encode(), b"<").replace(f"</{prefix}:".encode(), b"</")
    return value


def set_text(node, text):
    if "\n" in str(text):
        body = node.find("p:txBody", NS)
        if body is None:
            body = node.find("a:txBody", NS)
        if body is not None:
            paragraphs = body.findall("a:p", NS)
            if paragraphs:
                prototype = deepcopy(paragraphs[0])
                for paragraph in paragraphs:
                    body.remove(paragraph)
                for line in str(text).splitlines():
                    paragraph = deepcopy(prototype)
                    set_text(paragraph, line)
                    body.append(paragraph)
                return
    values = node.findall(".//a:t", NS)
    if values:
        values[0].text = str(text)
        for t in values[1:]:
            t.text = ""


def fill_table(slide, headers, rows):
    table = slide.find(".//a:tbl", NS)
    existing = table.findall("a:tr", NS)
    header, sample = deepcopy(existing[0]), deepcopy(existing[1])
    for r in existing:
        table.remove(r)
    for cell, value in zip(header.findall("a:tc", NS), headers):
        set_text(cell, value)
    table.append(header)
    for values in rows:
        row = deepcopy(sample)
        for cell, value in zip(row.findall("a:tc", NS), values):
            set_text(cell, value)
            # Keep lengthy Chinese names in their original cell, without truncation.
            for style in cell.findall(".//a:rPr", NS):
                if len(str(value)) > 12:
                    style.set("sz", "1100")
        table.append(row)


def chunks(rows, size=8):
    return [rows[i:i + size] for i in range(0, len(rows), size)]


def text_pages(value, width=48, lines=9):
    wrapped = []
    for paragraph in value.splitlines():
        wrapped.extend([paragraph[i:i + width] for i in range(0, len(paragraph), width)] or [""])
    return ["\n".join(page) for page in chunks(wrapped, lines)]


def resize_text(shape, *, size, height=None, y=None):
    for style in shape.findall(".//a:rPr", NS):
        style.set("sz", str(size))
    transform = shape.find("p:spPr/a:xfrm", NS)
    if height:
        transform.find("a:ext", NS).set("cy", str(height * 9525))
    if y:
        transform.find("a:off", NS).set("y", str(y * 9525))


def build_pptx(data, template, sections, excellent_ids, notes, birthday, photos, photo_caption, label):
    with zipfile.ZipFile(ROOT / f"{template}.pptx") as source:
        parts = {name: source.read(name) for name in source.namelist()}
    content_types = ET.fromstring(parts["[Content_Types].xml"])
    overrides = {x.get("PartName").lstrip("/"): x.get("ContentType") for x in content_types if x.tag == tag("ct", "Override")}
    slides = []
    status = "草稿（未全部月结）" if data["draft"] else "已月结"
    footer = f"{status}  {label}"
    month_label = data["month"].replace("-", "年") + "月"

    def add(source_number, *, heading=None, subtitle=None, table=None, replacements=None):
        old = f"ppt/slides/slide{source_number}.xml"
        root = ET.fromstring(parts[old])
        sp = [s for s in root.findall(".//p:sp", NS) if s.findall(".//a:t", NS)]
        if heading is not None:
            set_text(sp[0], heading)
        if subtitle is not None:
            set_text(sp[1], subtitle)
        for shape in sp:
            for text in shape.findall(".//a:t", NS):
                if text.text == "示例数据 / 非正式报告":
                    text.text = footer
                elif text.text and "示例" in text.text:
                    text.text = text.text.replace("示例", "")
        # Replace footer page numbering after optional chapters/pagination.
        for shape in sp:
            if "".join(t.text or "" for t in shape.findall(".//a:t", NS)).isdigit():
                set_text(shape, f"{len(slides) + 1:02d}")
        for index, value in (replacements or {}).items():
            set_text(sp[index], value)
        if source_number == 20:
            prototype = ET.fromstring(parts["ppt/slides/slide2.xml"])
            watermark = next(shape for shape in prototype.findall(".//p:sp", NS) if any(t.text == "示例数据 / 非正式报告" for t in shape.findall(".//a:t", NS)))
            set_text(watermark, footer)
            watermark.find("p:nvSpPr/p:cNvPr", NS).set("id", "950")
            root.find("p:cSld/p:spTree", NS).append(watermark)
        if table:
            fill_table(root, *table)
        relations = ET.fromstring(parts[f"ppt/slides/_rels/slide{source_number}.xml.rels"])
        for rel in list(relations):
            if rel.get("Type", "").endswith("/notesSlide"):
                relations.remove(rel)
            elif rel.get("TargetMode") != "External":
                target = rel.get("Target")
                rel.set("Target", "/" + posixpath.normpath(posixpath.join(posixpath.dirname(old), target)) if not target.startswith("/") else target)
        slides.append((root, relations))
        return root, relations

    def table_pages(number, rows, title, headers, subtitle=""):
        template_slide = ET.fromstring(parts[f"ppt/slides/slide{number}.xml"])
        capacity = min(8, len(template_slide.findall(".//a:tbl/a:tr", NS)) - 1)
        pages = chunks(rows, capacity)
        for i, page in enumerate(pages):
            add(number, heading=title, subtitle=f"{subtitle}  {i + 1}/{len(pages)}", table=(headers, page))

    def chart(number, categories, series):
        path = f"ppt/slides/charts/chart{number}.xml"
        root = ET.fromstring(parts[path])
        book = Workbook()
        sheet = book.active
        sheet.title = "Chart Data"
        chart_series = root.findall(".//c:ser", NS)
        for unused in chart_series[len(series):]:
            next(parent for parent in root.iter() if unused in list(parent)).remove(unused)
        for index, ser in enumerate(chart_series[:len(series)]):
            values = series[index][1] if index < len(series) else [None] * len(categories)
            title = series[index][0] if index < len(series) else ""
            tx = ser.find("c:tx", NS)
            for child in list(tx):
                tx.remove(child)
            ET.SubElement(tx, tag("c", "v")).text = title
            from openpyxl.utils import get_column_letter
            xcol, ycol = get_column_letter(index * 2 + 1), get_column_letter(index * 2 + 2)
            sheet.cell(1, index * 2 + 1, "类别")
            sheet.cell(1, index * 2 + 2, title)
            for kind, refkind, cachekind, vals, col in (("cat", "strRef", "strCache", categories, xcol), ("val", "numRef", "numCache", values, ycol)):
                holder = ser.find(f"c:{kind}", NS)
                for child in list(holder):
                    holder.remove(child)
                ref = ET.SubElement(holder, tag("c", refkind))
                ET.SubElement(ref, tag("c", "f")).text = f"'Chart Data'!${col}$2:${col}${len(categories) + 1}"
                cache = ET.SubElement(ref, tag("c", cachekind))
                if kind == "val":
                    ET.SubElement(cache, tag("c", "formatCode")).text = "General"
                ET.SubElement(cache, tag("c", "ptCount"), val=str(len(vals)))
                for i, value in enumerate(vals):
                    sheet.cell(i + 2, index * 2 + (1 if kind == "cat" else 2), value)
                    if value is not None:
                        point = ET.SubElement(cache, tag("c", "pt"), idx=str(i))
                        ET.SubElement(point, tag("c", "v")).text = str(value)
            # Freeze no category/value links to the original sample workbook.
        chart_rels_path = "ppt/slides/charts/_rels/" + posixpath.basename(path) + ".rels"
        chart_rels = ET.fromstring(parts[chart_rels_path])
        target = next(r.get("Target") for r in chart_rels if r.get("Type", "").endswith("/package"))
        workbook_path = posixpath.normpath(posixpath.join(posixpath.dirname(path), target)).lstrip("/")
        output = BytesIO()
        book.save(output)
        parts[workbook_path] = output.getvalue()
        parts[path] = xml(root)

    add(1, replacements={2: month_label, 3: "报告范围：" + data["scope_name"], 4: status + " / " + label})
    summary = data["summary"]
    if "overview" in sections:
        add(2, subtitle=month_label + "  " + status, replacements={5: f'{summary["employee_count"]} 人', 7: f'{summary["recognition_count"]} 条', 9: f'{summary["deduction_count"]} 条', 10: "认可率 " + (f'{summary["recognition_rate"]}%' if summary["recognition_rate"] is not None else "不适用"), 12: f'实际加分 {summary["recognition_score"]}，实际扣分 {summary["deduction_score"]}。LOA排除 {data["loa_count"]} 人。'})
    for section, chart_slide, table_slide, chart_number in (("deductions", 3, 4, 1), ("recognitions", 5, 6, 2)):
        records = data[section]
        if section not in sections or not records:
            continue
        totals = Counter()
        for r in records:
            totals[r["type"]] += r["count"]
        ordered = [k for k, _ in totals.most_common()]
        categories = ordered[:7] + (["其他类别合计"] if len(ordered) > 7 else [])
        values = []
        for role, codes in (("CM/TR", {"CM", "TR"}), ("LEAD/其他", {"LEAD/其他"})):
            values.append((role, [sum(r["count"] for r in records if r["category"] in codes and (r["type"] == category or category == "其他类别合计" and r["type"] in ordered[7:])) for category in categories]))
        chart(chart_number, categories, values)
        add(chart_slide)
        rows = [[c["name"], sum(r["count"] for r in records if r["attraction_name"] == c["name"] and r["category"] in {"CM", "TR"}), sum(r["count"] for r in records if r["attraction_name"] == c["name"] and r["category"] not in {"CM", "TR"}), sum(r["count"] for r in records if r["attraction_name"] == c["name"])] for c in data["circles"]]
        table_pages(table_slide, rows, "减分分类：景点圈" if section == "deductions" else "加分分类：景点圈", ["景点圈", "CM/TR次数", "LEAD/其他次数", "总次数"], "有效记录次数")
        table_pages(table_slide, [[r["attraction_name"], r["category"], r["type"], r["count"]] for r in records], "分类明细（次数）", ["景点圈", "人员类别", "类型", "次数"])
    if "issuers" in sections:
        table_pages(7, [[i + 1, r["name"], "选定报告范围", r["count"]] for i, r in enumerate(data["issuers"])], "认可发放排名", ["排名", "签卡人", "发放范围", "发放数量"])
    if "recognition_stats" in sections:
        table_pages(8, [[r["name"], r["recognition_count"], r["employee_count"], f'{r["recognition_rate"]}%' if r["recognition_rate"] is not None else "不适用"] for r in data["circles"]], "认可统计", ["景点圈", "认可数量", "统计人数", "认可率"])
        if any(r["rate"] is not None for r in data["trend"]):
            chart(3, [r["month"] for r in data["trend"]], [("认可率", [r["rate"] for r in data["trend"]])])
            add(9, subtitle="同一报告范围，缺失月份留空，未月结为草稿  单位：%")
    if "groups" in sections:
        table_pages(10, [[i + 1, r["name"], r["attraction_name"], r["employee_count"], r["average"]] for i, r in enumerate(data["groups"])], "小组分排名", ["排名", "小组", "景点圈", "计分人数", "平均分"])
    if "excellent" in sections:
        selected = [r for r in data["candidates"] if r["employee_id"] in excellent_ids]
        table_pages(11, [[r["category"], r["employee_name"], r["attraction_name"], "人工确认名单"] for r in selected], "优秀员工", ["人员类别", "员工", "景点圈", "名单来源"])
        for root, _ in slides:
            for t in root.findall(".//a:t", NS):
                if t.text and ("名单状态：" in t.text or "事迹说明" in t.text):
                    t.text = "名单已经制作者人工确认；不自动生成事迹。"
    if "rankings" in sections:
        for ranking in data["rankings"]:
            table_pages(12 if ranking["category"] in {"CM", "TR"} else 14, [[r["rank"], r["employee_name"], r["employee_no"], r["recognition_score"], r["deduction_score"], r["attendance_score"], r["total_score"]] for r in ranking["rows"]], "PR排名：" + ranking["category"], ["排名", "员工", "工号", "实际加分", "实际扣分", "全勤分", "综合分"], ranking["attraction_name"])
    if "notes" in sections and notes.strip():
        for page in text_pages(notes):
            root, _ = add(15, replacements={4: "工作与规则说明", 5: page, 6: "", 7: "", 8: ""})
            shapes = [s for s in root.findall(".//p:sp", NS) if s.findall(".//a:t", NS)]
            resize_text(shapes[5], size=1350, height=300, y=300)
    if "photos" in sections:
        for page in chunks(photos, 4):
            number = 16 if len(page) == 1 else 17 if len(page) == 2 else 18
            root, relations = add(number, replacements={1: "用户提供的团队活动照片"})
            slots = [s for s in root.findall(".//p:sp", NS) if any((t.text or "").startswith("照片插入区") for t in s.findall(".//a:t", NS))]
            parent = root.find("p:cSld/p:spTree", NS)
            for slot_index, slot in enumerate(slots):
                if slot_index < len(page):
                    photo = page[slot_index]
                    previous = list(parent)[list(parent).index(slot) - 1]
                    # Template has a background rectangle followed by its label.
                    # Fill the rectangle's box, not the much smaller label box.
                    rectangle = previous if previous.tag == tag("p", "sp") and not previous.findall(".//a:t", NS) else slot
                    transform = rectangle.find("p:spPr/a:xfrm", NS)
                    off, ext = transform.find("a:off", NS), transform.find("a:ext", NS)
                    x, y, w, h = [int(n) for n in (off.get("x"), off.get("y"), ext.get("cx"), ext.get("cy"))]
                    scale = min(w / photo["width"], h / photo["height"])
                    pw, ph = int(photo["width"] * scale), int(photo["height"] * scale)
                    media = f"ppt/media/hr-photo-{len(slides)}-{slot_index}.png"
                    parts[media] = photo["bytes"]
                    rid = f"hrPhoto{slot_index}"
                    ET.SubElement(relations, tag("rel", "Relationship"), Id=rid, Type=NS["r"] + "/image", Target="/" + media)
                    pic = ET.SubElement(parent, tag("p", "pic"))
                    nv = ET.SubElement(pic, tag("p", "nvPicPr"))
                    ET.SubElement(nv, tag("p", "cNvPr"), id=str(900 + slot_index), name="活动照片")
                    ET.SubElement(nv, tag("p", "cNvPicPr"))
                    ET.SubElement(nv, tag("p", "nvPr"))
                    fill = ET.SubElement(pic, tag("p", "blipFill"))
                    ET.SubElement(fill, tag("a", "blip"), {tag("r", "embed"): rid})
                    ET.SubElement(ET.SubElement(fill, tag("a", "stretch")), tag("a", "fillRect"))
                    props = ET.SubElement(pic, tag("p", "spPr"))
                    xfrm = ET.SubElement(props, tag("a", "xfrm"))
                    ET.SubElement(xfrm, tag("a", "off"), x=str(x + (w - pw) // 2), y=str(y + (h - ph) // 2))
                    ET.SubElement(xfrm, tag("a", "ext"), cx=str(pw), cy=str(ph))
                    ET.SubElement(ET.SubElement(props, tag("a", "prstGeom"), prst="rect"), tag("a", "avLst"))
                parent.remove(slot)
            for t in root.findall(".//a:t", NS):
                if t.text and ("活动名称" in t.text or "日期、地点" in t.text):
                    t.text = "团队活动" if number == 16 and t.text == "活动名称" else (photo_caption or "团队活动")
            for shape in root.findall(".//p:sp", NS):
                if photo_caption and any(t.text == photo_caption for t in shape.findall(".//a:t", NS)):
                    resize_text(shape, size=1200 if number == 16 else 1050, height=220 if number == 16 else 60)
    if "birthdays" in sections and birthday.strip():
        for page in text_pages(birthday, width=48, lines=7):
            root, _ = add(19, replacements={1: "人工提供的生日祝福", 5: page, 6: "生日内容由制作者提供。"})
            shapes = [s for s in root.findall(".//p:sp", NS) if s.findall(".//a:t", NS)]
            resize_text(shapes[5], size=1500, height=200, y=320)
    if "closing" in sections:
        add(20)
    presentation = ET.fromstring(parts["ppt/presentation.xml"])
    ids = presentation.find("p:sldIdLst", NS)
    ids.clear()
    rels = ET.fromstring(parts["ppt/_rels/presentation.xml.rels"])
    for rel in list(rels):
        if rel.get("Type", "").endswith("/slide"):
            rels.remove(rel)
    for i, (slide, slide_rels) in enumerate(slides, 1):
        name = f"ppt/slides/report{i}.xml"
        parts[name] = xml(slide)
        parts[f"ppt/slides/_rels/report{i}.xml.rels"] = xml(slide_rels)
        overrides[name] = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
        rid = f"hrSlide{i}"
        ET.SubElement(ids, tag("p", "sldId"), {"id": str(256 + i), tag("r", "id"): rid})
        ET.SubElement(rels, tag("rel", "Relationship"), Id=rid, Type=NS["r"] + "/slide", Target="/" + name)
    parts["ppt/presentation.xml"] = xml(presentation)
    parts["ppt/_rels/presentation.xml.rels"] = xml(rels)
    if "docProps/app.xml" in parts:
        properties = ET.fromstring(parts["docProps/app.xml"])
        for node in properties.iter():
            if node.tag.endswith("}Slides"):
                node.text = str(len(slides))
        parts["docProps/app.xml"] = xml(properties)
    # Export no template sample text, embedded sample workbooks, or unused slides.
    reachable = set()
    def visit(name):
        if name in reachable or name not in parts:
            return
        reachable.add(name)
        relpath = posixpath.join(posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels") if name else "_rels/.rels"
        if relpath in parts:
            reachable.add(relpath)
            for rel in ET.fromstring(parts[relpath]):
                if rel.get("TargetMode") != "External":
                    target = rel.get("Target")
                    visit(target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(name), target)))
    reachable.add("_rels/.rels")
    for rel in ET.fromstring(parts["_rels/.rels"]):
        visit(rel.get("Target").lstrip("/"))
    for child in list(content_types):
        if child.tag == tag("ct", "Override"):
            content_types.remove(child)
    for path in sorted(reachable):
        if path in overrides:
            ET.SubElement(content_types, tag("ct", "Override"), PartName="/" + path, ContentType=overrides[path])
    parts["[Content_Types].xml"] = xml(content_types)
    reachable.add("[Content_Types].xml")
    result = BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as output:
        for path in sorted(reachable):
            output.writestr(path, parts[path])
    return result.getvalue()

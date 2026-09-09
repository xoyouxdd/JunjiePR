"""Excel 导出的安全值、列宽与基础样式原语。"""

from __future__ import annotations

from datetime import date, datetime
from urllib.parse import quote

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FONT = Font(bold=True, size=16, color="FFFFFF")
LABEL_FILL = PatternFill("solid", fgColor="F3F7FB")
KPI_FILL = PatternFill("solid", fgColor="D9EAF7")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
BODY_ALIGNMENT = Alignment(vertical="center", wrap_text=True)


def excel_safe_value(value: object) -> object:
    """Keep numeric types; prefix formula-like text so Excel will not execute it."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def excel_row(values) -> list:
    return [excel_safe_value(value) for value in values]


def display_width(text: object) -> int:
    """Excel column units: CJK (ord > 0x2E80) counts as 2, everything else as 1."""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text or ""))


def style_sheet(ws, *, landscape: bool = False) -> None:
    """Apply JunjiePR header styling, CJK-aware column widths, borders and wrap."""
    fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = max(display_width(cell.value) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 42)
        for cell in column:
            cell.border = THIN_BORDER
            cell.alignment = BODY_ALIGNMENT
    for row_idx in range(1, (ws.max_row or 1) + 1):
        ws.row_dimensions[row_idx].height = 24
    configure_sheet(ws, freeze="A2", landscape=landscape)


def configure_sheet(ws, *, freeze="A2", landscape=False) -> None:
    """Hide gridlines, set zoom/print fit, orientation and margins."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def add_banded_table(ws, ref, name) -> None:
    """Attach an Excel Table. Failures are ignored so exports never 500."""
    previous_filter = getattr(ws.auto_filter, "ref", None)
    try:
        from openpyxl.worksheet.table import Table, TableStyleInfo

        if not ref or (ws.max_row or 0) < 1:
            return
        ws.auto_filter.ref = None
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    except Exception:
        if previous_filter:
            ws.auto_filter.ref = previous_filter


def write_merged_title(ws, title, columns) -> None:
    """Merge row 1 across `columns` with JunjiePR navy title styling."""
    end_column = max(int(columns or 1), 1)
    if end_column > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = ws.cell(1, 1, excel_safe_value(title))
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28


def write_label_value_pairs(ws, pairs, start_row) -> int:
    """Write vertical label-value pairs in columns A/B. Returns the next empty row."""
    row_number = start_row
    for label, value in pairs:
        label_cell = ws.cell(row_number, 1, excel_safe_value(label))
        value_cell = ws.cell(row_number, 2, excel_safe_value(value))
        label_cell.font = Font(bold=True)
        label_cell.fill = LABEL_FILL
        for cell in (label_cell, value_cell):
            cell.border = THIN_BORDER
            cell.alignment = BODY_ALIGNMENT
        ws.row_dimensions[row_number].height = 24
        row_number += 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 22)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 56)
    return row_number


def write_kpi_block(ws, kpis, start_row) -> int:
    """Write a 关键指标 block of name-value pairs. Returns the next empty row."""
    header = ws.cell(start_row, 1, "关键指标")
    header.font = Font(bold=True, color="1F4E78")
    header.fill = KPI_FILL
    header.alignment = BODY_ALIGNMENT
    header.border = THIN_BORDER
    ws.cell(start_row, 2).fill = KPI_FILL
    ws.cell(start_row, 2).border = THIN_BORDER
    ws.row_dimensions[start_row].height = 24
    row_number = start_row + 1
    for name, value in kpis:
        label_cell = ws.cell(row_number, 1, excel_safe_value(name))
        value_cell = ws.cell(row_number, 2, excel_safe_value(value))
        label_cell.font = Font(bold=True)
        label_cell.fill = LABEL_FILL
        if isinstance(value, float):
            value_cell.number_format = "0.00"
        for cell in (label_cell, value_cell):
            cell.border = THIN_BORDER
            cell.alignment = BODY_ALIGNMENT
        ws.row_dimensions[row_number].height = 24
        row_number += 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 22)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 56)
    return row_number


def apply_score_format(ws, columns) -> None:
    for column in columns:
        for row_idx in range(2, (ws.max_row or 1) + 1):
            ws.cell(row_idx, column).number_format = "0.00"


def apply_date_format(ws, columns) -> None:
    """Format date/datetime cells only; leave ISO strings untouched."""
    for column in columns:
        for row_idx in range(2, (ws.max_row or 1) + 1):
            cell = ws.cell(row_idx, column)
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"


def content_disposition(ascii_filename: str, display_filename: str) -> str:
    """RFC 5987 Content-Disposition: ASCII filename= plus UTF-8 filename*."""
    return f'attachment; filename="{ascii_filename}"; filename*=UTF-8\'\'{quote(display_filename)}'

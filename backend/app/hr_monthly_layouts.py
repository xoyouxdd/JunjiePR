"""Named layout filling; no business scoring, Office, Node, or network at runtime."""
from collections import Counter
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import math
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from app.hr_monthly_pptx import NS, ROOT, tag, xml, set_text, pack_report

EMU = 9525
WIDTHS = {4: [110, 530, 230, 282], 5: [80, 350, 350, 160, 212], 7: [75, 265, 140, 155, 155, 155, 207]}


def units(value):
    return sum(1 if unicodedata.east_asian_width(c) in "WF" else .56 for c in str(value))


def wrap(value, width, point=18):
    lines = []
    for paragraph in str(value).split("\n"):
        line = ""
        for char in paragraph:
            if ord(char) < 32 and char != "\t":
                continue
            if line and units(line + char) > width / (point * 4 / 3):
                lines.append(line)
                line = ""
            line += char
        lines.append(line)
    return lines


def chunks(rows, size):
    return [rows[i:i + size] for i in range(0, len(rows), size)]


def shape(slide, name):
    for item in slide.findall(".//p:sp", NS):
        node = item.find("p:nvSpPr/p:cNvPr", NS)
        if node is not None and node.get("name") == name:
            return item
    raise ValueError(f"月报模板缺少命名区域：{name}")


def put(slide, name, value):
    set_text(shape(slide, name), str(value))


def paginate(rows, widths):
    pages, page, used = [], [], 0
    for row in rows:
        cells = [wrap(value, width - 28) for value, width in zip(row, widths)]
        # Long fields continue on subsequent rows without truncation or tiny text.
        for segment in range(math.ceil(max(map(len, cells)) / 9)):
            values = ["\n".join(c[segment * 9:(segment + 1) * 9]) for c in cells]
            height = max(48, max(v.count("\n") + 1 for v in values) * 30 + 14)
            if page and used + height > 300:
                pages.append(page)
                page, used = [], 0
            page.append((values, height))
            used += height
    if page:
        pages.append(page)
    return pages


class LayoutReport:
    def __init__(self, theme, footer, month):
        with zipfile.ZipFile(ROOT / f"{theme}.pptx") as source:
            self.parts = {n: source.read(n) for n in source.namelist()}
        self.layouts = json.loads((ROOT / f"{theme}-layouts.json").read_text(encoding="utf8"))["layouts"]
        self.types = ET.fromstring(self.parts["[Content_Types].xml"])
        self.overrides = {r.get("PartName").lstrip("/"): r.get("ContentType") for r in self.types if r.tag == tag("ct", "Override")}
        self.slides, self.media = [], {}
        self.footer, self.month = footer, month

    def add(self, layout, title, subtitle="", context="", section=""):
        source = f'ppt/slides/slide{self.layouts[layout]}.xml'
        root = ET.fromstring(self.parts[source])
        rels = ET.fromstring(self.parts[f'ppt/slides/_rels/slide{self.layouts[layout]}.xml.rels'])
        import posixpath
        for rel in list(rels):
            if rel.get("Type", "").endswith("/notesSlide"):
                rels.remove(rel)
            elif rel.get("TargetMode") != "External":
                rel.set("Target", "/" + posixpath.normpath(posixpath.join(posixpath.dirname(source), rel.get("Target"))).lstrip("/"))
        root.find("p:cSld/p:spTree", NS).remove(shape(root, "layout-key"))
        title_width = int(shape(root, "title").find("p:spPr/a:xfrm/a:ext", NS).get("cx")) / EMU
        chapter = "小组积分" if layout.startswith('group-') else "活动记录" if layout.startswith('photo') else "人员数据" if title.startswith(('PR排名', '优秀员工', '认可发放')) else "数据明细" if layout.startswith('table') else "月度数据" if layout in {'overview', 'trend', 'classification'} else "团队记录"
        for name, value in {"title": "\n".join(wrap(title, title_width, 48 if layout == "cover" else 33)), "subtitle": "\n".join(wrap(subtitle, 1140, 18)), "context": "\n".join(wrap(context, 1140, 11)), "section": section or chapter, "footer": self.footer, "month": self.month, "page": str(len(self.slides) + 1).zfill(2)}.items():
            put(root, name, value)
        self.slides.append((root, rels))
        return root, rels

    def tables(self, headers, rows, title, subtitle="", context=""):
        if not rows:
            return
        widths = WIDTHS[len(headers)]
        if len(headers) == 4:
            widths = [90, 360, 470, 232] if headers[0] == "排名" else [120, 380, 400, 252] if headers[0] == "类别" else [330, 230, 360, 232]
        pages = paginate(rows, widths)
        for index, page in enumerate(pages):
            root, _ = self.add("table" + str(len(headers)), title, subtitle + (f"  {index + 1}/{len(pages)}" if len(pages) > 1 else ""), context)
            self.fill_table(root, headers, page, widths)

    def fill_table(self, root, headers, page, widths):
        table = root.find(".//a:tbl", NS)
        for col, width in zip(table.findall("a:tblGrid/a:gridCol", NS), widths):
            col.set("w", str(width * EMU))
        old = table.findall("a:tr", NS)
        header, prototype = deepcopy(old[0]), deepcopy(old[1])
        for row in old:
            table.remove(row)
        header.set("h", str(46 * EMU))
        for cell, value in zip(header.findall("a:tc", NS), headers):
            set_text(cell, value)
        table.append(header)
        for values, height in page:
            row = deepcopy(prototype)
            row.set("h", str(height * EMU))
            for cell, value in zip(row.findall("a:tc", NS), values):
                set_text(cell, value)
            table.append(row)
        frame = next(f for f in root.findall(".//p:graphicFrame", NS) if f.find(".//a:tbl", NS) is not None)
        frame.find("p:xfrm/a:ext", NS).set("cy", str((46 + sum(h for _, h in page)) * EMU))

    def chart(self, rels, categories, series, fmt="0", reverse=False):
        rel = next(r for r in rels if r.get("Type", "").endswith("/chart"))
        root = ET.fromstring(self.parts[rel.get("Target").lstrip("/")])
        # Keep role buckets and their colors stable, even when one bucket is zero.
        if series[0][0] not in {'CM/TR', 'LEAD/其他'} and any(any(v not in (None, 0) for v in values) for _, values in series):
            series = [(name, values) for name, values in series if any(v not in (None, 0) for v in values)]
        if reverse:
            categories, series = list(reversed(categories)), [(n, list(reversed(v))) for n, v in series]
        nodes = root.findall(".//c:ser", NS)
        parent = next(p for p in root.iter() if nodes[0] in list(p))
        for node in nodes[len(series):]:
            parent.remove(node)
        book = Workbook()
        sheet = book.active
        sheet.title = "Chart Data"
        from openpyxl.utils import get_column_letter
        for i, ((name, values), node) in enumerate(zip(series, nodes)):
            tx = node.find("c:tx", NS)
            tx.clear()
            ET.SubElement(tx, tag("c", "v")).text = name
            for kind, refkind, cachekind, vals, col in (("cat", "strRef", "strCache", categories, i * 2 + 1), ("val", "numRef", "numCache", values, i * 2 + 2)):
                holder = node.find(f"c:{kind}", NS)
                holder.clear()
                ref = ET.SubElement(holder, tag("c", refkind))
                letter = get_column_letter(col)
                ET.SubElement(ref, tag("c", "f")).text = f"'Chart Data'!${letter}$2:${letter}${len(categories) + 1}"
                cache = ET.SubElement(ref, tag("c", cachekind))
                if kind == "val":
                    ET.SubElement(cache, tag("c", "formatCode")).text = fmt
                ET.SubElement(cache, tag("c", "ptCount"), val=str(len(vals)))
                sheet.cell(1, col, "类别" if kind == "cat" else name)
                for j, value in enumerate(vals):
                    cell = sheet.cell(j + 2, col, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    if value is not None:
                        point = ET.SubElement(cache, tag("c", "pt"), idx=str(j))
                        ET.SubElement(point, tag("c", "v")).text = str(value)
        nums = [v for _, values in series for v in values if v is not None]
        high, low = max([0, *nums]), min([0, *nums])
        span = high - low or 1
        exponent = 10 ** math.floor(math.log10(span / 4))
        step = next(n * exponent for n in (1, 2, 5, 10) if n * exponent >= span / 4)
        if fmt == '0':
            step = max(1, step)
        for axis in root.findall(".//c:valAx", NS):
            scaling = axis.find("c:scaling", NS)
            for key, value in (("min", math.floor(low / step) * step), ("max", math.ceil((high + .1 * span) / step) * step or step)):
                element = scaling.find(f"c:{key}", NS)
                if element is None:
                    element = ET.SubElement(scaling, tag("c", key))
                element.set("val", str(value))
            unit = axis.find("c:majorUnit", NS)
            if unit is None:
                unit = ET.SubElement(axis, tag("c", "majorUnit"))
            unit.set("val", str(step))
            axis_format = axis.find('c:numFmt', NS)
            if axis_format is None:
                axis_format = ET.SubElement(axis, tag('c', 'numFmt'))
            decimals = max(0, math.ceil(-math.log10(step)))
            axis_format.set('formatCode', ('0' + ('.' + '0' * decimals if decimals else '')) + ('"%"' if '%' in fmt else ''))
            axis_format.set('sourceLinked', '0')
        for item in root.findall(".//c:dLbls/c:numFmt", NS):
            item.set("formatCode", '0;-0;;' if fmt == '0' else fmt)
            item.set('sourceLinked', '0')
        chart = root.find("c:chart", NS)
        if len(series) == 1 and chart.find("c:legend", NS) is not None:
            chart.remove(chart.find("c:legend", NS))
        elif chart.find('c:legend', NS) is not None:
            legend = chart.find('c:legend', NS)
            old = legend.find('c:txPr', NS)
            if old is not None:
                legend.remove(old)
            text_style = deepcopy(root.find('.//c:catAx/c:txPr', NS))
            for style in text_style.findall('.//a:defRPr', NS):
                style.set('sz', '1800')
            legend.append(text_style)
        blanks = chart.find("c:dispBlanksAs", NS)
        if blanks is None:
            blanks = ET.SubElement(chart, tag("c", "dispBlanksAs"))
        blanks.set("val", "gap")
        i = len(self.slides)
        chart_path, book_path = f"ppt/charts/report{i}.xml", f"ppt/embeddings/report{i}.xlsx"
        root.find("c:externalData", NS).set(tag("r", "id"), "reportData")
        self.parts[chart_path] = xml(root)
        out = BytesIO()
        book.save(out)
        self.parts[book_path] = out.getvalue()
        relations = ET.Element(tag("rel", "Relationships"))
        ET.SubElement(relations, tag("rel", "Relationship"), Id="reportData", Type=NS["r"] + "/package", Target="/" + book_path)
        self.parts[f"ppt/charts/_rels/report{i}.xml.rels"] = xml(relations)
        rel.set("Target", "/" + chart_path)
        self.overrides[chart_path] = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
        self.overrides[book_path] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def photo(self, root, rels, photo, index):
        slot = shape(root, f"photo-slot-{index}")
        trans = slot.find("p:spPr/a:xfrm", NS)
        off, ext = trans.find("a:off", NS), trans.find("a:ext", NS)
        x, y, w, h = map(int, (off.get("x"), off.get("y"), ext.get("cx"), ext.get("cy")))
        factor = min(w / photo["width"], h / photo["height"])
        pw, ph = int(photo["width"] * factor), int(photo["height"] * factor)
        key = hashlib.sha256(photo["bytes"]).hexdigest()
        if key not in self.media:
            self.media[key] = f'activity-{len(self.media) + 1}.{photo.get("extension", "png")}'
            path = "ppt/media/" + self.media[key]
            self.parts[path] = photo["bytes"]
            self.overrides[path] = photo.get("content_type", "image/png")
        rid = f"activity{index}"
        ET.SubElement(rels, tag("rel", "Relationship"), Id=rid, Type=NS["r"] + "/image", Target="/ppt/media/" + self.media[key])
        parent = root.find("p:cSld/p:spTree", NS)
        pic = ET.SubElement(parent, tag("p", "pic"))
        nv = ET.SubElement(pic, tag("p", "nvPicPr"))
        ET.SubElement(nv, tag("p", "cNvPr"), id=str(9000 + index), name="活动照片")
        ET.SubElement(nv, tag("p", "cNvPicPr"))
        ET.SubElement(nv, tag("p", "nvPr"))
        fill = ET.SubElement(pic, tag("p", "blipFill"))
        ET.SubElement(fill, tag("a", "blip"), {tag("r", "embed"): rid})
        ET.SubElement(ET.SubElement(fill, tag("a", "stretch")), tag("a", "fillRect"))
        props = ET.SubElement(pic, tag("p", "spPr"))
        trans = ET.SubElement(props, tag("a", "xfrm"))
        ET.SubElement(trans, tag("a", "off"), x=str(x + (w - pw) // 2), y=str(y + (h - ph) // 2))
        ET.SubElement(trans, tag("a", "ext"), cx=str(pw), cy=str(ph))
        ET.SubElement(ET.SubElement(props, tag("a", "prstGeom"), prst="rect"), tag("a", "avLst"))
        parent.remove(slot)

    def finish(self):
        return pack_report(self.parts, self.slides, self.overrides, self.types)


def export_report(data, theme, sections, excellent_ids, notes, birthday, photos, caption, label):
    status = "草稿（未全部月结）" if data["draft"] else "已月结"
    report = LayoutReport(theme, f"{status}  {label}", data["month"])
    scope, summary = data["scope_name"], data["summary"]
    root, _ = report.add("cover", {"forest": "团队运营月报", "minimal": "团队运营数据简报", "warm": "团队月刊"}[theme], section="")
    put(root, "section", "")
    put(root, "scope", "\n".join(wrap(scope, 535, 24)))
    put(root, "status", status)
    if "overview" in sections:
        root, rels = report.add("overview", f'本月确认认可{summary["recognition_count"]}条' if theme == 'forest' else '月度数据概览', f'{scope}  实际加分{summary["recognition_score"]}，实际扣分{summary["deduction_score"]}，LOA排除{data["loa_count"]}人', "认可率按记录数计算，可超过100%，不代表人员覆盖率。")
        put(root, "hero", summary["recognition_count"])
        put(root, "metrics", f'{summary["employee_count"]} 人\n{summary["deduction_count"]} 条有效减分\n认可率 ' + (f'{summary["recognition_rate"]}%' if summary["recognition_rate"] is not None else "不适用"))
        totals = Counter()
        for row in data["recognitions"]:
            totals[row["type"]] += row["count"]
        items = totals.most_common(5)
        if len(totals) > 5:
            items.append(("其他类型", sum(totals.values()) - sum(v for _, v in items)))
        report.chart(rels, [k for k, _ in items] or ["暂无认可"], [("确认认可次数", [v for _, v in items] or [0])], reverse=theme == "minimal")
    for section, title in (("deductions", "减分"), ("recognitions", "加分")):
        records = data[section]
        if section not in sections or not records:
            continue
        totals = Counter()
        for row in records:
            totals[row["type"]] += row["count"]
        for page in chunks(totals.most_common(), 6):
            categories = [k for k, _ in page]
            root, rels = report.add("classification", title + "类型分布", "有效记录次数", "分类按记录次数展示，不代表实际计分；景点圈与人员类别见明细。")
            series = [(role, [sum(r["count"] for r in records if r["type"] == c and (r["category"] in {"CM", "TR"}) == cm) for c in categories]) for role, cm in (("CM/TR", True), ("LEAD/其他", False))]
            report.chart(rels, categories, series, reverse=theme == "minimal")
        report.tables(["景点圈", "人员类别", "类型", "次数"], [[r["attraction_name"], r["category"], r["type"], r["count"]] for r in records], title + "分类明细", "有效记录次数")
    if "issuers" in sections:
        issuers = data['issuers']
        if all(units(r['name']) < 15 for r in issuers):
            for index, page in enumerate(chunks(issuers, 6)):
                _, rels = report.add('issuer-chart', '认可发放排名', scope, '按签卡人归属，不重复累计代录人。')
                report.chart(rels, [f'{index * 6 + i + 1} {r["name"]}' for i, r in enumerate(page)], [('认可发放次数', [r['count'] for r in page])], reverse=True)
        else:
            report.tables(["排名", "签卡人", "发放范围", "次数"], [[i + 1, r["name"], scope, r["count"]] for i, r in enumerate(issuers)], "认可发放排名", "按签卡人归属，不重复累计代录人")
    if "recognition_stats" in sections:
        headers = ["景点圈", "认可数量", "统计人数", "认可率"]
        rows = [[r["name"], r["recognition_count"], r["employee_count"], f'{r["recognition_rate"]}%' if r["recognition_rate"] is not None else "不适用"] for r in data["circles"] if r["employee_count"] or r["recognition_count"]]
        trend = any(r["rate"] is not None for r in data["trend"])
        pages = paginate(rows, [330,230,360,232])
        combined = bool(trend and len(pages) == 1 and sum(h for _,h in pages[0]) <= 96)
        if combined:
            root, rels = report.add('stats-trend', '认可统计与趋势', scope, '认可率可超过100%；缺失月份留空，未月结月份标为草稿。')
            report.fill_table(root, headers, pages[0], [330,230,360,232])
        else:
            report.tables(headers, rows, "认可统计", "非LOA统计人数；认可率可超过100%")
        if trend:
            if not combined:
                _, rels = report.add("trend", "认可率趋势", scope, "同一报告范围，缺失值留空，未月结月份为草稿。")
            report.chart(rels, [r["month"] + ("（草稿）" if r.get("draft") else "") for r in data["trend"]], [("认可率（%）", [r["rate"] for r in data["trend"]])], '0.00"%"')
    if "groups" in sections:
        rows = data["groups"]
        context = "展示顺序沿用系统规则；月度积分不代表员工综合能力。"
        if theme == "minimal" and all(units(r["name"] + r["attraction_name"]) < 18 for r in rows):
            for index, page in enumerate(chunks(rows, 6)):
                _, rels = report.add("group-chart", "小组分排名", f"{scope}  {index + 1}/{math.ceil(len(rows) / 6)}", context)
                report.chart(rels, [f'{index * 6 + i + 1} {r["name"]} / {r["attraction_name"]} / {r["employee_count"]}人' for i, r in enumerate(page)], [("平均分", [r["average"] for r in page])], "0.00", True)
        elif theme == "warm" and all(units(r["name"]) < 15 and units(r["attraction_name"]) < 14 for r in rows):
            for index, page in enumerate(chunks(rows, 6)):
                root, _ = report.add("group-editorial", "小组积分一览", f"{scope}  {index + 1}/{math.ceil(len(rows) / 6)}", context)
                for i in range(6):
                    r = page[i] if i < len(page) else None
                    values = {"order": str(index * 6 + i + 1).zfill(2), "name": r["name"] + "\n" + r["attraction_name"] if r else "", "score": f'{r["average"]:.2f}' if r else "", "people": str(r["employee_count"]) + "人计分" if r else ""}
                    for key, value in values.items():
                        put(root, f"group-{key}-{i}", value if r else "")
        else:
            report.tables(["顺序", "小组", "景点圈", "人数", "均分"], [[i + 1, r["name"], r["attraction_name"], r["employee_count"], f'{r["average"]:.2f}'] for i, r in enumerate(rows)], "小组分排名", scope, context)
    if "excellent" in sections:
        selected = [r for r in data['candidates'] if r['employee_id'] in excellent_ids]
        if all(units(r['employee_name']) < 17 and units(r['attraction_name']) < 14 for r in selected):
            for page in chunks(selected, 6):
                root, _ = report.add('excellent-editorial', '优秀员工', '本月人工确认名单', '名单由制作者确认，不自动生成获奖事迹。')
                for i in range(6):
                    r = page[i] if i < len(page) else None
                    for key, value in {'category': r['category'] if r else '', 'name': '\n'.join(wrap(r['employee_name'],360,26)) if r else '', 'scope': r['attraction_name'] if r else ''}.items():
                        put(root, f'excellent-{key}-{i}', value)
        else:
            report.tables(["类别", "员工", "景点圈", "名单来源"], [[r["category"], r["employee_name"], r["attraction_name"], "人工确认"] for r in selected], "优秀员工", "由制作者确认名单，不自动生成事迹")
    if "rankings" in sections:
        for ranking in data["rankings"]:
            report.tables(["排名", "员工", "工号", "实际加分", "实际扣分", "全勤分", "综合分"], [[r["rank"], r["employee_name"], r["employee_no"], *[f'{r[key]:.2f}' for key in ('recognition_score', 'deduction_score', 'attendance_score', 'total_score')]] for r in ranking["rows"]], "PR排名：" + ranking["category"], ranking["attraction_name"], "沿用当月计分与排序，保留现有同分顺序。")
    for section, content, title in (("notes", notes, "工作与规则说明"), ("birthdays", birthday, "生日祝福")):
        if section in sections and content.strip():
            lines = wrap(content, 1120, 23)
            for index, page in enumerate(chunks(lines, 8)):
                root, _ = report.add("text", title, f"{index + 1}/{math.ceil(len(lines) / 8)}")
                put(root, "body", "\n".join(page))
    if "photos" in sections and photos:
        capacity = 4 if len(photos) > 8 else 2
        for index, page in enumerate(chunks(photos, capacity)):
            layout = "photo1" if len(page) == 1 else "photo2" if len(page) == 2 else "photo4"
            root, rels = report.add(layout, "团队活动", f"{index + 1}/{math.ceil(len(photos) / capacity)}", caption)
            for i, photo in enumerate(page):
                report.photo(root, rels, photo, i)
            if layout == "photo1" and theme != "warm":
                put(root, "photo-side", "\n".join(wrap(caption or "团队活动记录", 275, 23)))
            parent = root.find("p:cSld/p:spTree", NS)
            for item in list(parent):
                nv = item.find("p:nvSpPr/p:cNvPr", NS)
                if nv is not None and nv.get("name", "").startswith("photo-slot-"):
                    parent.remove(item)
    if "closing" in sections:
        root, _ = report.add("closing", "本月报告结束")
        put(root, "closing-scope", "\n".join(wrap(scope, 1120, 30)))
        put(root, "closing-status", status)
    return report.finish()

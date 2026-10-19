"""Presentation-only changes: source rows, scoring and permission scopes stay owned by callers."""
from __future__ import annotations

from io import BytesIO
from copy import copy
import hashlib
import posixpath
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl.styles import Alignment, Font

from app.excel_export_utils import configure_sheet, excel_safe_value, fit_row_heights, fold_columns


HISTORY_FIELDS = {"发生时员工号", "当前员工号", "工号状态", "撤回人", "撤回时间", "撤回原因",
                  "作废人", "作废时间", "作废原因", "升级关联", "升级角色", "升级状态",
                  "当月已确认认可次数", "同日重复登记", "POC周期", "POC原因"}


def add_personnel_groups(ws, data) -> None:
    """Expose the existing hierarchy association, without new grouping/scoring rules."""
    current_group = ""
    group_by_id = {}
    for node in data.get("hierarchy", []):
        if node.get("node_type") in {"attraction", "gsm", "gsm_team"}:
            current_group = ""
        elif node.get("node_type") == "supervisor":
            current_group = node.get("name") or ""
        elif node.get("node_type") == "employee":
            group_by_id[int(node["employee_id"])] = current_group
    by_no = {row["employee_no"]: group_by_id.get(int(row["employee_id"]), "")
             for row in [*data.get("scores", []), *data.get("loa_rows", [])]}
    column = ws.max_column + 1
    ws.cell(1, column, "小组")
    for row in ws.iter_rows(min_row=2, max_col=1):
        ws.cell(row[0].row, column, excel_safe_value(by_no.get(row[0].value, "")))
    from app.excel_export_utils import style_sheet
    style_sheet(ws, landscape=True)


def present_detail(ws, *, personnel: bool = False) -> None:
    """Do not move/delete columns: existing references and import consumers keep working."""
    headers = {cell.value: cell.column for cell in ws[1]}
    low_frequency = HISTORY_FIELDS if not personnel else {"工号状态", "备注"}
    fold_columns(ws, low_frequency)
    ws.sheet_format.outlineLevelCol = 1
    for label in ("姓名", "员工", "组织 / 员工"):
        if label in headers:
            ws.column_dimensions[ws.cell(1, headers[label]).column_letter].width = 26
    # Keep the original order and freeze the identity columns alongside the header.
    identity_end = max((headers.get(label, 0) for label in ("员工号", "工号", "姓名", "员工")), default=0)
    if identity_end:
        ws.freeze_panes = f"{ws.cell(1, identity_end + 1).column_letter}2"
        ws.print_title_cols = f"A:{ws.cell(1, identity_end).column_letter}"
    fit_row_heights(ws)


def finalize_statistics_workbook(wb) -> None:
    """One simple entry point; preserve every existing sheet, column and formula."""
    order = ["导出说明", "月度综合分明细", "月度综合分", "主管月度综合分明细", "主管月度综合分",
             "签卡明细", "扣分明细", "病假明细", "LOA明细", "层级绩效明细", "作废操作记录",
             "员工留档标记", "员工号变更对照"]
    for name in reversed(order):
        if name in wb.sheetnames:
            wb.move_sheet(name, offset=-wb.sheetnames.index(name))
    for ws in wb:
        if ws.title in {"月度综合分明细", "主管月度综合分明细"}:
            present_detail(ws, personnel=True)
        elif ws.title in {"签卡明细", "扣分明细", "病假明细", "LOA明细"}:
            present_detail(ws)
    cover = wb["导出说明"]
    # Replace only our own directory when a route finalizes again after adding supervisors.
    for row in cover.iter_rows(min_col=4, max_col=5):
        for cell in row:
            cell.value = None
            cell.hyperlink = None
    cover.cell(1, 4, "工作表目录").font = Font(bold=True, color="1F4E78")
    for row_number, ws in enumerate((ws for ws in wb if ws is not cover), 2):
        cell = cover.cell(row_number, 4, ws.title)
        cell.hyperlink = f"#'{ws.title.replace(chr(39), chr(39)*2)}'!A1"
        cell.font = Font(color="0563C1", underline="single")
        cover.cell(row_number, 5, "＋展开低频字段" if ws.title in {
            "月度综合分明细", "主管月度综合分明细", "签卡明细", "扣分明细", "病假明细"} else "查看完整记录")
    cover.column_dimensions["A"].width = 23
    cover.column_dimensions["B"].width = 58
    cover.column_dimensions["C"].width = 3
    cover.column_dimensions["D"].width = 26
    cover.column_dimensions["E"].width = 22
    for row in cover:
        for cell in row:
            font = copy(cell.font)
            font.name = "Microsoft YaHei"
            font.scheme = None
            cell.font = font
    fit_row_heights(cover)
    for dimension in cover.row_dimensions.values():
        if dimension.height == 24:
            dimension.height = 20
    configure_sheet(cover, freeze="A2", landscape=True)
    cover.page_setup.fitToHeight = 1
    cover.print_area = cover.dimensions
    wb.active = wb.sheetnames.index("导出说明")
    wb.views[0].firstSheet = 0


def append_export_context(ws, pairs) -> None:
    """Metadata goes below the existing source table, never inside its filter range."""
    row = ws.max_row + 2
    for label, value in pairs:
        ws.cell(row, 1, excel_safe_value(label)).font = Font(bold=True, color="1F4E78")
        ws.cell(row, 2, excel_safe_value(value)).alignment = Alignment(wrap_text=True, vertical="center")
        row += 1
    fit_row_heights(ws)


def save_export_workbook(wb, output) -> None:
    """Serialize normally, then share byte-identical media without removing watermarks.

    All drawing anchors remain. Only relationship targets and redundant media
    package entries change; cell XML, formulas and styles are copied verbatim.
    """
    source = BytesIO()
    wb.save(source)
    source.seek(0)
    with ZipFile(source) as archive:
        redirects, hashes = {}, {}
        for item in archive.infolist():
            if item.filename.startswith("xl/media/"):
                content = archive.read(item.filename)
                key = (posixpath.splitext(item.filename)[1].lower(), hashlib.sha256(content).digest())
                if key in hashes:
                    redirects[item.filename] = hashes[key]
                else:
                    hashes[key] = item.filename
        if not redirects:
            output.write(source.getvalue())
            return
        with ZipFile(output, "w", ZIP_DEFLATED) as packed:
            for item in archive.infolist():
                if item.filename in redirects:
                    continue
                content = archive.read(item.filename)
                if item.filename.endswith(".rels"):
                    tree = ET.fromstring(content)
                    base = posixpath.dirname(posixpath.dirname(item.filename))
                    changed = False
                    for relationship in tree:
                        if relationship.get("TargetMode") == "External":
                            continue
                        target = relationship.get("Target", "")
                        resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base, target))
                        if resolved in redirects:
                            relationship.set("Target", "/" + redirects[resolved])
                            changed = True
                    if changed:
                        content = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
                elif item.filename == "[Content_Types].xml":
                    tree = ET.fromstring(content)
                    changed = False
                    for child in list(tree):
                        if child.get("PartName", "").lstrip("/") in redirects:
                            tree.remove(child)
                            changed = True
                    if changed:
                        content = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
                packed.writestr(item, content)

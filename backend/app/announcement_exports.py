"""Compact, permission-scoped announcement reports with verifiable signatures."""
from datetime import date, datetime
from functools import lru_cache
from io import BytesIO
import hashlib
import json
import math
import re

import fitz
from PIL import Image, ImageChops
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.pagebreak import Break

from app.v2_database import FILE_DIR, SessionLocal
from app.v2_models import StoredFile

STATUS = {"pending": "待查收", "completed": "已查收", "suspended": "暂缓", "exited": "退出范围",
          "superseded": "已替换", "withdrawn": "已撤下", "expired": "已过期"}
INK, MUTED, TEAL = "173F52", "526779", "176152"


def text(value):
    if isinstance(value, str) and len(value) >= 19:
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            pass
    return value.isoformat(sep=" ", timespec="seconds") if isinstance(value, datetime) else str(value or "—")


def excel_date(value):
    if not value or isinstance(value, (date, datetime)):
        return value or None
    try:
        return datetime.fromisoformat(value) if len(value) > 10 else date.fromisoformat(value)
    except (TypeError, ValueError):
        return str(value)


def status_label(row):
    return "逾期未签" if row.get("overdue") else STATUS.get(row["status"], row["status"])


def status_note(row):
    if row.get("status_note"):
        return row["status_note"]
    if row.get("confirmed_at"):
        return "本人账号签收"
    if row["status"] == "exited":
        return "已退出适用范围；具体原因未记录"
    if row["status"] == "suspended":
        return "LOA暂缓" if row.get("reason") == "LOA暂缓" else "暂缓；具体原因未记录"
    return ("已阅读，尚未签收" if row.get("seen_at") else "尚无阅读及签收记录") + "；未签原因未记录"


def report_groups(data):
    metadata = {g.get("key"): g for g in data.get("groups", [])}
    groups = {}
    for row in data["items"]:
        key = row.get("group_key") or (row["attraction_name"], row.get("group_name") or "无小组")
        group = groups.setdefault(key, {**metadata.get(key, {}), "circle": row["attraction_name"],
            "name": row.get("group_name") or "无小组", "items": [],
            "required": 0, "completed": 0, "pending": 0, "suspended": 0, "exited": 0})
        group["items"].append(row)
        if row.get("export_reference"):
            continue
        if row["status"] in {"suspended", "exited"}:
            group[row["status"]] += 1
        else:
            group["required"] += 1
            group["completed" if row.get("confirmed_at") else "pending"] += 1
    for group in groups.values():
        group["items"].sort(key=lambda r: (0 if r.get("export_role","").startswith("组长") else 1 if r.get("export_role")=="代理组长" else 2, r["employee_no"]))
    return list(groups.values())


def signature_id(row):
    url = row.get("signature_url")
    return int(url.split("/")[-1].split("?")[0]) if url else None


def signature_previews(data):
    """Verify originals once, then crop/resize report copies; never rewrite evidence."""
    ids = {signature_id(row) for row in data["items"] if row.get("signature_url")}
    result = {}
    with SessionLocal() as db:
        files = {row.id: row for row in db.query(StoredFile).filter(StoredFile.id.in_(ids))} if ids else {}
        for file_id in ids:
            stored = files.get(file_id)
            if not stored:
                raise ValueError("签字凭据文件缺失，导出已停止")
            path = (FILE_DIR / stored.storage_key).resolve()
            if FILE_DIR.resolve() not in path.parents or not path.is_file():
                raise ValueError("签字凭据文件缺失，导出已停止")
            raw = path.read_bytes()
            if hashlib.sha256(raw).hexdigest() != stored.sha256:
                raise ValueError("签字凭据摘要不匹配，导出已停止")
            try:
                with Image.open(BytesIO(raw)) as source:
                    if source.width * source.height > 2_000_000:
                        raise ValueError("签字凭据尺寸异常，导出已停止")
                    source.load()
                    white = Image.new("RGBA", source.size, "white")
                    white.alpha_composite(source.convert("RGBA"))
                    gray = white.convert("L")
                    bounds = ImageChops.difference(gray, Image.new("L", gray.size, 255)).getbbox()
                    if not bounds:
                        raise ValueError("签字凭据内容为空，导出已停止")
                    left, top, right, bottom = bounds
                    gray = gray.crop((max(0,left-6), max(0,top-6), min(gray.width,right+6), min(gray.height,bottom+6)))
                    gray.thumbnail((600, 180), Image.Resampling.LANCZOS)
                    output = BytesIO()
                    gray.save(output, format="PNG", optimize=True)
                    result[file_id] = {"content": output.getvalue(), "sha256": stored.sha256, "size": gray.size}
            except (OSError, SyntaxError) as exc:
                raise ValueError("签字凭据无法读取，导出已停止") from exc
    return result


def evidence_rows(data, signatures):
    seen = set()
    result = []
    for row in data["items"]:
        key = row.get("id", (row["employee_no"], row.get("confirmed_at")))
        if key in seen:
            continue
        seen.add(key)
        file_id = signature_id(row)
        result.append({"工号": row["employee_no"], "姓名": row["employee_name"], "签字时间": text(row.get("confirmed_at")),
            "签字声明": row.get("signing_statement") or "", "签字文件ID": file_id,
            "SHA256": signatures[file_id]["sha256"] if file_id else "",
            "接收原因": row.get("reason"), "接收时间": text(row.get("delivered_at")),
            "阅读时间": text(row.get("seen_at")), "状态说明": status_note(row),
            "跟进培训师": row.get("trainer_name"), "培训师工号": row.get("trainer_no"),
            "培训批次": row.get("training_session_id"), "验码时间": text(row.get("training_verified_at"))})
    return result


def expected_statement(row, announcement):
    return f"本人{row['employee_name']}（{row['employee_no']}）已阅读并知悉《{announcement['title']}》V{announcement['version']}的要求。"


def compact_statement(row, announcement):
    statement = row.get("signing_statement") or ""
    return f"已阅读并知悉V{announcement['version']}要求" if statement == expected_statement(row, announcement) else statement


def scope_note(data):
    return "当前管理范围内人员" if data["scoped"] else "全部有权查看人员"


def count_line(counts, historical=False):
    required, pending = ("原应签", "未签") if historical else ("应签", "待签")
    return f"{required} {counts['required']} · 已签 {counts['completed']} · {pending} {counts['pending']} · 暂缓 {counts['suspended']} · 退出 {counts['exited']}"


def export_records(data, format, user, followup=False):
    return spreadsheet(data, user, followup) if format == "xlsx" else pdf(data, user)


def spreadsheet(data, user, followup=False):
    signatures = signature_previews(data)
    announcement, groups = data["announcement"], report_groups(data)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "公告与总览"
    label = f"ANN-{announcement['id']:06d} / V{announcement['version']}"
    stamp = f"导出：{user.name}（{user.employee.employee_no}） · {datetime.now():%Y-%m-%d %H:%M:%S}"

    def setup(sheet, widths):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.page_margins.left = sheet.page_margins.right = .3
        sheet.page_margins.top = sheet.page_margins.bottom = .4
        sheet.oddFooter.center.text = f"{label} / {user.employee.employee_no} / &P / &N"
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width

    def merged(sheet, value, height=24, bold=False, fill=None):
        row = sheet.max_row + 1 if sheet.cell(1,1).value is not None else 1
        sheet.cell(row,1,value)
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8)
        sheet.row_dimensions[row].height = height
        cell = sheet.cell(row,1)
        cell.font = Font(name="微软雅黑",size=16 if bold else 11,bold=bold,color=INK)
        if fill:
            cell.fill = PatternFill("solid",fgColor=fill)
        return row

    setup(overview,[18,22,16,18,20,22,24,24])
    merged(overview,"FZ团队工作台 / 公告签收总览",30,True)
    merged(overview,announcement["title"],max(32,math.ceil(len(announcement["title"])/70)*20),True)
    merged(overview,f"{label} · {announcement['category']} · {scope_note(data)}")
    merged(overview,f"本版发布人：{announcement.get('author') or '未记录'} · 首次发布：{text(announcement.get('published_at'))} · 本版更新：{text(announcement.get('updated_at'))}")
    merged(overview,f"范围：{announcement['scope_label']} · 生效：{announcement['effective_on']} · 截止：{announcement.get('due_on') or '未设置'} · 失效：{announcement.get('expires_on') or '持续有效'}",30)
    mode = "培训码＋本人签字" if announcement.get("training_required") else {1:"阅读",2:"明确确认",3:"本人手写签字"}[announcement["confirmation_level"]]
    merged(overview,f"确认方式：{mode} · 状态：{'有效' if announcement['effective'] else '历史 / 已停止适用'}")
    merged(overview,"发布内容",26,True)
    # Explicit wrapping preserves paragraphs and avoids Excel's 409pt row cap.
    for title, value in [("正文",announcement.get("body")),("重点",announcement.get("summary")),("变更",announcement.get("change_summary"))]:
        if value is None:
            continue
        if title != "正文":
            merged(overview,title,24,True)
        for paragraph in value.split("\n"):
            for start in range(0,max(1,len(paragraph)),100):
                part = paragraph[start:start+100]
                merged(overview,part,max(20,math.ceil(len(part)/70)*18+8))
    merged(overview,count_line(data["counts"],not announcement["effective"]),30,False,"EDF4F7")
    ratio = f"{data['counts']['completed']/data['counts']['required']:.0%}" if data["counts"]["required"] else "—"
    merged(overview,f"签收率 {ratio} · 暂缓、退出不计入应签；新建账号不自动免签")
    overview.append(["景点圈","小组","应签" if announcement["effective"] else "原应签","已签","待签" if announcement["effective"] else "未签","暂缓","退出","签收率"])
    summary_header = overview.max_row
    for group in groups:
        overview.append([group["circle"],group["name"],group["required"],group["completed"],group["pending"],group["suspended"],group["exited"],group["completed"]/group["required"] if group["required"] else None])
        overview.cell(overview.max_row,8).number_format = "0%"
    merged(overview,stamp)
    overview.freeze_panes = "A7"
    overview.print_area = f"A1:H{overview.max_row}"
    # One readable cover for normal announcements, natural continuation for long ones.
    if sum(overview.row_dimensions[i].height or 20 for i in range(1,overview.max_row+1)) < 490:
        overview.page_setup.fitToHeight = 1
    overview.print_title_rows = "1:6"

    def safe_name(name):
        base = re.sub(r"[\\/*?:\[\]]","-",name).strip("'") or "小组"
        candidate, index = base[:31], 1
        while candidate in workbook.sheetnames:
            index += 1
            suffix = f"-{index}"
            candidate = base[:31-len(suffix)] + suffix
        return candidate

    for group in groups:
        sheet = workbook.create_sheet(safe_name(group["name"] if group["name"].startswith(group["circle"]) else f"{group['circle']}-{group['name']}"))
        setup(sheet,[18,14,16,28,23,46,22,23])
        merged(sheet,f"{group['circle']} · {group['name']}",30,True)
        merged(sheet,f"{label} · {announcement['title']}",max(26,math.ceil(len(announcement["title"])/65)*18+8))
        merged(sheet,f"当前组长：{group.get('formal_leader_name') or '未设置'} · 当前代理组长：{group.get('acting_leader_name') or '未设置'}")
        merged(sheet,count_line(group,not announcement["effective"]),26,False,"EDF4F7")
        merged(sheet,"当前负责人职务；参照记录不重复计数。状态说明为系统事实，不推定未签原因。",24)
        sheet.append(["工号 / 姓名","当前身份","签收状态","签名及确认内容","签字时间","状态说明 / 可核实情况","跟进培训师","验码时间"])
        height_on_page = 0
        for row in group["items"]:
            reference = "（参照，不重复计数）" if row.get("export_reference") else ""
            statement = compact_statement(row,announcement)
            sheet.append([f"{row['employee_no']}\n{row['employee_name']}",row.get("export_role") or row["role_label"],status_label(row)+reference,
                statement or ("无手写签名要求" if row.get("confirmed_at") else "—"),excel_date(row.get("confirmed_at")),status_note(row),
                row.get("trainer_name") or ("未验码" if announcement.get("training_required") else "—"),excel_date(row.get("training_verified_at"))])
            index = sheet.max_row
            # Keep the image within its own row, with the declaration below it.
            file_id = signature_id(row)
            statement_lines = max(1,math.ceil(len(statement)/16))
            note_lines = max(1,math.ceil(len(status_note(row))/23))
            height = max(54, note_lines*14+12, statement_lines*14+(38 if file_id else 12))
            if height > 360:
                raise ValueError("签字声明或状态说明过长，无法在人员行中清晰展示")
            if height_on_page + height > 440 and index > 7:
                sheet.row_breaks.append(Break(id=index-1))
                height_on_page = 0
            height_on_page += height
            sheet.row_dimensions[index].height = height
            if file_id:
                signature = signatures[file_id]
                image = ExcelImage(BytesIO(signature["content"]))
                width, img_height = signature["size"]
                factor = min(155/width,36/img_height)
                image.width, image.height = width*factor, img_height*factor
                image.anchor = TwoCellAnchor(editAs="twoCell",
                    _from=AnchorMarker(col=3,row=index-1,colOff=pixels_to_EMU(4),rowOff=pixels_to_EMU(2)),
                    to=AnchorMarker(col=3,row=index-1,colOff=pixels_to_EMU(4+image.width),rowOff=pixels_to_EMU(2+image.height)))
                sheet.add_image(image)
                sheet.cell(index,4).alignment = Alignment(vertical="bottom",wrap_text=True)
            sheet.cell(index,1).number_format = "@"
        statement = f"统一签字声明模板：本人〔本行姓名、工号〕已阅读并知悉《{announcement['title']}》V{announcement['version']}的要求。非标准声明逐行原文展示。"
        for offset in range(0,len(statement),100):
            merged(sheet,statement[offset:offset+100],32)
        merged(sheet,"原始签名保留；报表为裁边缩放副本。完整声明、原件SHA256及培训凭据见隐藏工作表「凭据核对」。",30)
        merged(sheet,stamp)
        sheet.freeze_panes = "C7"
        sheet.auto_filter.ref = f"A6:H{6+len(group['items'])}"
        sheet.print_title_rows = "1:6"
        sheet.print_area = f"A1:H{sheet.max_row}"
        if not sheet.row_breaks.brk:
            sheet.page_setup.fitToHeight = 1
        for row in sheet.iter_rows(min_row=7,max_row=6+len(group["items"])):
            for cell in (row[4],row[7]):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

    audit = workbook.create_sheet("凭据核对")
    rows = evidence_rows(data,signatures)
    headers = ["工号","姓名","签字时间","签字声明","签字文件ID","SHA256","接收原因","接收时间","阅读时间","状态说明","跟进培训师","培训师工号","培训批次","验码时间"]
    audit.append(headers)
    for record in rows:
        audit.append([excel_date(record.get(key)) if key in {"签字时间","验码时间"} and record.get(key) != "—" else record.get(key) for key in headers])
    audit.sheet_state = "hidden"
    audit.freeze_panes = "C2"
    audit.auto_filter.ref = f"A1:N{max(1,audit.max_row)}"
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = "s"  # Includes user text beginning with =, +, - or @.
                if cell.alignment.wrap_text is not True:
                    cell.alignment = Alignment(vertical="center",wrap_text=True)
                if not cell.font.name:
                    cell.font = Font(name="微软雅黑",size=11,color=INK)
        if sheet is audit:
            continue
        header = summary_header if sheet is overview else 6
        for cell in sheet[header]:
            cell.fill = PatternFill("solid",fgColor=TEAL if sheet is overview else INK)
            cell.font = Font(name="微软雅黑",size=11,bold=True,color="FFFFFF")
        start, end = (summary_header+1,summary_header+len(groups)) if sheet is overview else (7,int(sheet.auto_filter.ref.split(":")[-1][1:]))
        for row in sheet.iter_rows(min_row=start,max_row=end):
            for cell in row:
                cell.font = Font(name="微软雅黑",size=10,color=INK)
                cell.fill = PatternFill("solid",fgColor="F0F5F8" if cell.row % 2 else "FFFFFF")
                cell.border = Border(bottom=Side(style="hair",color="D7E3EA"))
                if sheet is not overview and cell.column == 3:
                    value = str(cell.value)
                    color = TEAL if value.startswith("已查收") else "A25C11" if value.startswith(("待查收","逾期未签")) else MUTED
                    cell.font = Font(name="微软雅黑",size=10,bold=True,color=color)
    if followup and groups:
        workbook.active = 1
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@lru_cache(maxsize=4096)
def char_width(char, size):
    return fitz.get_text_length(char,fontname="helv" if ord(char)<128 else "china-s",fontsize=size)


def wrap_lines(value, width, size):
    lines = []
    for paragraph in str(value).split("\n"):
        line, length = "", 0
        for char in paragraph:
            advance = char_width(char,size)
            if line and length+advance > width:
                lines.append(line)
                line, length = "", 0
            line += char
            length += advance
        lines.append(line)
    return lines


def pdf(data, user):
    signatures = signature_previews(data)
    announcement, groups = data["announcement"], report_groups(data)
    document = fitz.open()
    label = f"ANN-{announcement['id']:06d} / V{announcement['version']}"
    stamp = f"导出：{user.name}（{user.employee.employee_no}） | {datetime.now():%Y-%m-%d %H:%M:%S}"
    ink, muted, teal = (.09,.25,.32), (.32,.40,.47), (.09,.38,.32)

    def draw(p,value,x,y,width,size=9,color=ink):
        lines = wrap_lines(value,width,size)
        for index,line in enumerate(lines):
            left = x
            for run in re.findall(r"[\x00-\x7f]+|[^\x00-\x7f]+",line):
                font = "helv" if ord(run[0])<128 else "china-s"
                p.insert_text((left,y+size+index*size*1.4),run,fontname=font,fontsize=size,color=color)
                left += fitz.get_text_length(run,fontname=font,fontsize=size)
        return len(lines)*size*1.4

    def page(title, landscape=False):
        p = document.new_page(width=842 if landscape else 595,height=595 if landscape else 842)
        draw(p,f"FZ团队工作台 / {title}",32,20,p.rect.width-280,10)
        draw(p,label,p.rect.width-210,20,178,9,muted)
        p.draw_line((32,45),(p.rect.width-32,45),color=teal,width=1)
        return p

    p, y = page("公告签收总览"), 62
    def flow(value,size=10,gap=6):
        nonlocal p,y
        for line in wrap_lines(value,531,size):
            if y+size*1.4 > 772:
                p,y = page("公告内容（续）"),62
            draw(p,line,32,y,531,size)
            y += size*1.4
        y += gap

    flow(announcement["title"],18,10)
    flow(f"{announcement['category']} · {label} · {scope_note(data)}",9)
    flow(f"本版发布人：{announcement.get('author') or '未记录'} | 首次发布：{text(announcement.get('published_at'))}",9)
    flow(f"本版更新：{text(announcement.get('updated_at'))} | 状态：{'有效' if announcement['effective'] else '历史 / 已停止适用'}",9)
    flow(f"范围：{announcement['scope_label']}",9)
    flow(f"生效：{announcement['effective_on']} | 截止：{announcement.get('due_on') or '未设置'} | 失效：{announcement.get('expires_on') or '持续有效'}",9)
    mode = "培训码＋本人签字" if announcement.get("training_required") else {1:"阅读",2:"明确确认",3:"本人手写签字"}[announcement["confirmation_level"]]
    flow(f"确认方式：{mode}",9,12)
    flow("发布内容 / 公告正文",12)
    flow(announcement.get("body") or "（无正文）",10,10)
    for title,key in [("重点","summary"),("变更说明","change_summary")]:
        if announcement.get(key):
            flow(f"{title}：{announcement[key]}",9,8)
    flow(count_line(data["counts"],not announcement["effective"]),11)
    ratio = f"{data['counts']['completed']/data['counts']['required']:.0%}" if data["counts"]["required"] else "—"
    flow(f"签收率 {ratio}；暂缓、退出不计入应签；新建账号不自动免签。",9)
    for group in groups:
        flow(f"{group['circle']} · {group['name']}：{count_line(group,not announcement['effective'])}",9,4)
    if not groups:
        flow("本版本尚无接收记录。",10)

    training = announcement.get("training_required") or any(r.get("trainer_name") for r in data["items"])
    widths = [123,64,166,112,205,108] if training else [130,66,182,118,282]
    headings = ["工号 / 姓名","状态","签名及确认内容","签字时间","状态说明 / 可核实情况"]+(["培训师 / 验码"] if training else [])
    bounds = [32]
    for width in widths:
        bounds.append(bounds[-1]+width)
    for group in groups:
        def group_page(continued=False):
            p = page("小组签收记录",True)
            title = f"{group['circle']} · {group['name']}" + ("（续）" if continued else "")
            top = 58+draw(p,title,32,58,778,13)+4
            top += draw(p,f"当前组长：{group.get('formal_leader_name') or '未设置'} | 当前代理组长：{group.get('acting_leader_name') or '未设置'}",32,top,778,8)+3
            top += draw(p,count_line(group,not announcement["effective"]),32,top,778,8)+5
            p.draw_rect(fitz.Rect(32,top,810,top+23),color=None,fill=ink)
            for i,title in enumerate(headings):
                draw(p,title,bounds[i]+5,top+5,widths[i]-10,8,(1,1,1))
            return p,top+23
        p,y = group_page()
        for index,row in enumerate(group["items"]):
            file_id = signature_id(row)
            declaration = compact_statement(row,announcement)
            digest = signatures[file_id]["sha256"] if file_id else ""
            note = status_note(row)+(f"\nSHA256: {digest[:12]}…" if digest else "")
            cells = [row["employee_name"]+"\n"+row["employee_no"]+("\n"+row["export_role"] if row.get("export_role") else ""),
                     status_label(row)+("\n参照，不重复计数" if row.get("export_reference") else ""),declaration or ("无手写签名要求" if row.get("confirmed_at") else "—"),
                     text(row.get("confirmed_at")).replace(" ","\n",1),note]
            if training:
                cells.append((row.get("trainer_name") or ("未验码" if announcement.get("training_required") else "—"))+"\n"+text(row.get("training_verified_at")).replace(" ","\n",1))
            heights = [len(wrap_lines(value,widths[i]-12,8.5))*11.9+10+(22 if i==2 and file_id else 0) for i,value in enumerate(cells)]
            height = max(44,*heights)
            if y+height > 495:
                p,y = group_page(True)
            if y+height > 495:
                raise ValueError("签字声明或状态说明过长，无法在人员行中清晰展示")
            if index % 2 == 0:
                p.draw_rect(fitz.Rect(32,y,810,y+height),color=None,fill=(.95,.97,.98))
            for i,value in enumerate(cells):
                top = y+5
                if i==2 and file_id:
                    p.insert_image(fitz.Rect(bounds[i]+6,top,bounds[i+1]-6,top+19),stream=signatures[file_id]["content"],keep_proportion=True)
                    top += 22
                color = ink
                if i==1:
                    color = teal if row["status"]=="completed" else (.63,.35,.07) if row["status"]=="pending" else muted
                draw(p,value,bounds[i]+6,top,widths[i]-12,8.5,color)
            y += height
            p.draw_line((32,y),(810,y),color=(.85,.9,.91),width=.4)
        # A standard statement is shortened per row only when it matches exactly.
        note = f"统一签字声明模板：本人〔本行姓名、工号〕已阅读并知悉《{announcement['title']}》V{announcement['version']}的要求。"
        note += "\n签字凭据：原件保留；此处为裁边缩放副本。完整声明及SHA256见PDF内附签字凭据核对.json。"
        note += "\n当前职务不代表签收时职务；参照行不重复计数。可核实情况不代表未签动机，不推定代签。"
        # Very long titles continue as readable notes rather than overlapping the footer.
        for line in wrap_lines(note,778,7):
            if y+10 > 552:
                p,y = group_page(True)
            draw(p,line,32,y+4,778,7,muted)
            y += 10
    proof = {"announcement_id":announcement["id"],"version":announcement["version"],"scope":scope_note(data),
             "note":"SHA256校验对应系统原始签名；页面图片为报表副本。","records":evidence_rows(data,signatures)}
    document.embfile_add("signature-evidence.json",json.dumps(proof,ensure_ascii=False,default=str).encode("utf-8"),filename="signature-evidence.json",ufilename="签字凭据核对.json")
    document.set_metadata({"title":announcement["title"],"author":"FZ团队工作台","subject":label})
    for index,p in enumerate(document):
        draw(p,stamp,32,p.rect.height-33,p.rect.width-190,7,muted)
        draw(p,f"受控签收记录 | {index+1} / {document.page_count}",p.rect.width-158,p.rect.height-33,126,7,muted)
    content = document.tobytes(garbage=4,deflate=True)
    document.close()
    return content

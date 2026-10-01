from __future__ import annotations

import os
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import tempfile

from openpyxl import load_workbook
from PIL import Image


TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-statement-upgrade-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402
from app.score_queries import employee_month_scores  # noqa: E402
from app.v2_database import SessionLocal  # noqa: E402
from app.routers._shared import upgrade_type_codes  # noqa: E402
from sqlalchemy import text  # noqa: E402


def login(client: TestClient, account: str) -> None:
    response = client.post("/api/login", json={"employee_no": account, "password": "1234"})
    assert response.status_code == 200, response.text


def statement_form(target_id: int, type_id: int, level_id: int, description: str) -> dict:
    return {
        "employee_id": str(target_id),
        "deduction_type_id": str(type_id),
        "deduction_level_id": str(level_id),
        "occurred_on": date.today().isoformat(),
        "description": description,
    }


def pdf_file(name: str = "statement.pdf"):
    stream = BytesIO()
    Image.new("RGB", (8, 8), "white").save(stream, format="PDF")
    return {"document": (name, stream.getvalue(), "application/pdf")}


def test_attendance_upgrade_categories_do_not_mix() -> None:
    assert upgrade_type_codes("ATT_EARLY_CLOCK") == {"ATT_EARLY_CLOCK", "ATT_LATE_CLOCK"}
    assert upgrade_type_codes("ATT_LATE_WITHIN_30") == {"ATT_LATE_WITHIN_30", "ATT_EARLY_LEAVE_WITHIN_30"}
    assert upgrade_type_codes("ATT_EARLY_CLOCK").isdisjoint(upgrade_type_codes("ATT_LATE_WITHIN_30"))
    assert upgrade_type_codes("SICK_LEAVE_VIOLATION") == {"SICK_LEAVE_VIOLATION"}


def test_deduction_picker_labels_upgrade_groups_without_changing_type_ids() -> None:
    source = (Path(__file__).parents[1] / "app" / "static" / "js" / "app.js").read_text(encoding="utf-8")
    assert "${opt(types,'id',deductionUpgradeOptionName)}" in source
    assert "${row.name}（${group.category}3个月内第2次触发升级）" in source
    assert "['ATT_EARLY_CLOCK','ATT_LATE_CLOCK']" in source
    assert "['ATT_LATE_WITHIN_30','ATT_EARLY_LEAVE_WITHIN_30']" in source
    assert "bindDeductionUpgradeHint();bindDeductionV2244()" in source
    assert 'aria-describedby="deductionUpgradeHint"' in source
    assert "对应升级项：${group.items}" in source


def test_late_and_early_leave_share_upgrade_history() -> None:
    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        target = next(row for row in client.get("/api/employee-targets", params={"usage": "deduction", "keyword": "TRTEST01"}).json()["items"] if row["employee_no"] == "TRTEST01")
        types = {row["code"]: row["id"] for row in options["deduction_types"]}
        statement = next(row for row in options["deduction_levels"] if row["code"] == "STATEMENT")
        first = client.post("/api/deductions", data=statement_form(target["id"], types["ATT_LATE_WITHIN_30"], statement["id"], "迟到"), files=pdf_file())
        assert first.status_code == 200, first.text
        params = {"employee_id": target["id"], "occurred_on": date.today().isoformat()}
        preview = client.get("/api/deduction-upgrades/preview", params={**params, "deduction_type_id": types["ATT_EARLY_LEAVE_WITHIN_30"]})
        assert preview.status_code == 200 and preview.json()["first_record"]["id"] == first.json()["record"]["id"]
        separate = client.get("/api/deduction-upgrades/preview", params={**params, "deduction_type_id": types["ATT_EARLY_CLOCK"]})
        assert separate.status_code == 200 and not separate.json()["eligible"]


def test_upgrade_matches_registered_statement_on_either_side_of_event_date() -> None:
    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        target = next(row for row in client.get("/api/employee-targets", params={"usage": "deduction", "keyword": "CMTEST02"}).json()["items"] if row["employee_no"] == "CMTEST02")
        types = {row["code"]: row["id"] for row in options["deduction_types"]}
        statement = next(row for row in options["deduction_levels"] if row["code"] == "STATEMENT")
        first = client.post("/api/deductions", data=statement_form(target["id"], types["ATT_EARLY_CLOCK"], statement["id"], "已登记的声明A"), files=pdf_file())
        assert first.status_code == 200, first.text
        params = {"employee_id": target["id"], "deduction_type_id": types["ATT_LATE_CLOCK"]}
        before_a = (date.today() - timedelta(days=20)).isoformat()
        preview = client.get("/api/deduction-upgrades/preview", params={**params, "occurred_on": before_a})
        assert preview.status_code == 200 and preview.json()["eligible"], preview.text
        assert preview.json()["first_record"]["id"] == first.json()["record"]["id"]
        repeat = client.get("/api/deductions/attendance-repeat-check", params={**params, "occurred_on": before_a})
        assert repeat.status_code == 200 and repeat.json()["has_repeat"], repeat.text
        assert repeat.json()["previous_records"][0]["id"] == first.json()["record"]["id"]
        outside = client.get("/api/deduction-upgrades/preview", params={**params, "occurred_on": (date.today() - timedelta(days=120)).isoformat()})
        assert outside.status_code == 200 and not outside.json()["eligible"], outside.text


def test_statement_upgrade_submit_transfer_approve_and_export_marking() -> None:
    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        target = next(row for row in client.get("/api/employee-targets", params={"usage": "deduction", "keyword": "CMTEST01"}).json()["items"] if row["employee_no"] == "CMTEST01")
        deduction_type = next(row for row in options["deduction_types"] if row["code"] == "ATT_EARLY_CLOCK")
        second_type = next(row for row in options["deduction_types"] if row["code"] == "ATT_LATE_CLOCK")
        statement = next(row for row in options["deduction_levels"] if row["code"] == "STATEMENT")
        first = client.post("/api/deductions", data=statement_form(target["id"], deduction_type["id"], statement["id"], "A1"), files=pdf_file())
        assert first.status_code == 200, first.text
        preview = client.get("/api/deduction-upgrades/preview", params={"employee_id": target["id"], "deduction_type_id": second_type["id"], "occurred_on": date.today().isoformat()})
        assert preview.status_code == 200 and preview.json()["eligible"], preview.text
        assert preview.json()["first_record"]["id"] == first.json()["record"]["id"]
        tagsm = next(row for row in preview.json()["reviewers"] if row["role_code"] == "TA_GSM")
        data = statement_form(target["id"], second_type["id"], statement["id"], "A2")
        data["reviewer_id"] = str(tagsm["id"])
        submitted = client.post("/api/deduction-upgrades", data=data, files=pdf_file("second.pdf"))
        assert submitted.status_code == 200, submitted.text
        request_id = submitted.json()["request"]["id"]
        assert submitted.json()["request"]["second_record"]["status"] == "pending_upgrade"
        client.post("/api/logout")
        login(client, "TAGSMTEST01")
        pending = client.get("/api/deduction-upgrades/pending").json()["items"]
        assert [row["id"] for row in pending] == [request_id]
        reviewers = client.get("/api/deduction-upgrades/reviewers").json()["items"]
        gsm = next(row for row in reviewers if row["role_code"] == "GSM")
        moved = client.post(f"/api/deduction-upgrades/{request_id}/transfer", json={"reviewer_id": gsm["id"], "reason": "由GSM审核"})
        assert moved.status_code == 200, moved.text
        client.post("/api/logout")
        login(client, "GSMTEST01")
        pending = client.get("/api/deduction-upgrades/pending").json()["items"]
        assert pending[0]["transfers"][0]["to_name"] == "测试GSM"
        memo = next(row for row in client.get("/api/options").json()["deduction_levels"] if row["code"] == "MEMO")
        approved = client.post(f"/api/deduction-upgrades/{request_id}/resolve", json={"decision": "approve", "result_level_id": memo["id"], "handling_note": "已完成真实备忘录开具", "issued_confirmed": True})
        assert approved.status_code == 200, approved.text
        request = approved.json()["request"]
        assert request["first_record"]["actual_points"] == statement["points"]
        assert request["second_record"]["actual_points"] == 0
        assert request["result_points"] == memo["points"]
        month = date.today().strftime("%Y-%m")
        with SessionLocal() as db:
            score = employee_month_scores(db, month, [target["id"]])[0]
            view_score = db.execute(text("SELECT deduction_score FROM v_employee_month_scores WHERE employee_id=:id AND score_month=:month"), {"id": target["id"], "month": month}).scalar_one()
        expected_deduction = round(float(statement["points"]) + float(memo["points"]), 2)
        assert score["deduction_score"] == expected_deduction
        assert view_score == expected_deduction
        ranking = client.get("/api/pr-rankings", params={"start_date": date.today().isoformat(), "end_date": date.today().isoformat(), "category": "deduction", "keyword": "CMTEST01"})
        assert ranking.status_code == 200, ranking.text
        ranked = next(row for row in ranking.json()["rows"] if row["employee_id"] == target["id"])
        assert ranked["score"] == expected_deduction
        response = client.get("/api/statistics/export", params={"month": date.today().strftime("%Y-%m")})
        assert response.status_code == 200, response.text
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        sheet = workbook["扣分明细"]
        roles = [sheet.cell(row=index, column=12).value for index in range(2, sheet.max_row + 1)]
        assert {"source_first", "source_second", "result"}.issubset(set(roles))
        second_rows = [index for index in range(2, sheet.max_row + 1) if sheet.cell(index, 12).value == "source_second"]
        assert all(sheet.cell(index, 7).value == 0 for index in second_rows)
        colors = [sheet.cell(row=index, column=1).fill.fgColor.rgb for index in range(2, sheet.max_row + 1)]
        assert any(color and color.endswith("DDEBF7") for color in colors)
        assert any(color and color.endswith("FFF2CC") for color in colors)
        assert any(color and color.endswith("E2F0D9") for color in colors)

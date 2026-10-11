"""Training codes exercise real sessions, signing evidence and publisher export."""
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from test_announcements import client, login, ok, form, publish, receipt, signature, new_employee
from app.main import app
from app.v2_database import SessionLocal
from app.v2_models import (Employee, EmployeeRoleAssignment, Role, AnnouncementTrainingCode as Code,
    AnnouncementTrainingReceipt as Evidence, AnnouncementTrainingAttempt as Attempt, AnnouncementVersion, UserAccount,
    AnnouncementDelivery)


def person(no):
    with SessionLocal() as db:
        return db.query(Employee).filter_by(employee_no=no).one().id


def extra_tr(circle):
    id, no = new_employee(circle)
    with SessionLocal() as db:
        assignment = db.query(EmployeeRoleAssignment).filter_by(employee_id=id).one()
        assignment.role_id = db.query(Role).filter_by(code="TR").one().id
        db.commit()
    return id, no


def setup(client, *, all_circles=False):
    login(client, "AMTEST01" if all_circles else "GSMTEST01")
    trainers = ok(client.get("/api/announcements/training-options"))
    tr = next(e for e in trainers["items"] if e["employee_no"] == "TRTEST01")
    a = publish(client, confirmation_level=3, training_required=True, trainer_scope="all" if all_circles else "circle", trainer_ids=[tr["id"]])
    return a, tr


def generate(client, a, no="TRTEST01", hours=2):
    login(client, no)
    return ok(client.post(f"/api/announcements/{a['id']}/training/codes", json={"version": a["version"], "hours": hours, "request_key": uuid4().hex}))


def verify(client, a, code):
    return client.post(f"/api/announcements/{a['id']}/training/verify", json={"version": a["version"], "code": code})


def sign(client, a, code):
    return receipt(client, a, action="confirm", acknowledged=True, signature=signature(), training_code=code)


def test_training_selection_and_access_boundaries(client):
    login(client)
    assert client.post("/api/announcements", json=form(client, training_required=True, trainer_ids=[person("TRTEST01")])).status_code == 422
    assert client.post("/api/announcements", json=form(client, confirmation_level=3, training_required=True)).status_code == 422
    opts = ok(client.get("/api/announcements/training-options"))
    other_circle = next(c["id"] for c in ok(client.get("/api/announcements/options"))["circles"] if c["id"] != opts["own_circle_id"])
    tr_id, tr_no = extra_tr(other_circle)
    assert client.post("/api/announcements", json=form(client, confirmation_level=3, training_required=True, trainer_ids=[tr_id])).status_code == 400
    a = publish(client, confirmation_level=3, training_required=True, trainer_scope="all", trainer_ids=[tr_id])
    code = generate(client, a, tr_no, 6)
    assert len(code["code"]) == 6 and code["code"].isdigit()
    login(client, "TRTEST01")
    assert client.get(f"/api/announcements/{a['id']}/training").status_code == 403
    login(client, tr_no)
    assert ok(client.get("/api/announcements?view=training"))["items"][0]["id"] == a["id"]
    assert client.get(f"/api/announcements/{a['id']}").status_code == 200  # Trainer may teach outside receiver scope.
    assert verify(client, a, code["code"]).status_code == 403  # Not an announcement recipient.


def test_multiple_trainers_independent_codes_and_immutable_signing(client):
    a, tr = setup(client)
    extra_id, extra_no = extra_tr(tr["attraction_id"])
    login(client)
    ok(client.post(f"/api/announcements/{a['id']}/training/trainers", json={"expected_version":1,"trainer_scope":"circle","trainer_ids":[tr["id"],extra_id]}))
    first = generate(client, a)
    second = generate(client, a, extra_no, 6)
    assert first["code"] != second["code"]
    replaced = generate(client, a)
    login(client, "CMTEST01")
    assert verify(client, a, first["code"]).status_code == 400
    assert verify(client, a, second["code"]).status_code == 200
    login(client)
    pending = next(r for r in ok(client.get(f"/api/announcements/{a['id']}/records"))["items"] if r["employee_id"] == person("CMTEST01"))
    assert pending["status"] == "pending" and pending["trainer_id"] == extra_id and not pending["confirmed_at"]
    login(client, "CMTEST01")
    signed = ok(sign(client, a, second["code"]))
    assert signed["trainer_id"] == extra_id and signed["signature_url"]
    original_time = signed["training_verified_at"]
    duplicate = ok(sign(client, a, replaced["code"]))
    assert duplicate["trainer_id"] == extra_id and duplicate["signature_url"] == signed["signature_url"]
    assert duplicate["training_verified_at"] == original_time
    login(client, extra_no)
    progress = ok(client.get(f"/api/announcements/{a['id']}/training"))
    assert any(r["employee_id"] == person("CMTEST01") and r["confirmed_at"] for r in progress["items"])
    assert all("signature_url" not in r for r in progress["items"])
    assert second["code"] not in str(progress)  # No code disclosure via polling.


@pytest.mark.parametrize("change", ["expire", "revoke", "disable", "role", "withdraw", "update"])
def test_code_is_rechecked_when_signature_is_submitted(client, change):
    a, tr = setup(client)
    code = generate(client, a)
    login(client, "CMTEST01")
    assert verify(client, a, code["code"]).status_code == 200
    if change == "expire":
        with SessionLocal() as db:
            db.get(Code, code["id"]).expires_at = datetime.now()-timedelta(seconds=1)
            db.commit()
    elif change in {"disable", "role"}:
        with SessionLocal() as db:
            if change == "disable":
                db.query(UserAccount).filter_by(employee_id=tr["id"]).one().enabled = False
            else:
                role = db.query(EmployeeRoleAssignment).filter_by(employee_id=tr["id"]).first()
                old_role = role.role_id
                role.role_id = db.query(Role).filter_by(code="CM").one().id
            db.commit()
    else:
        login(client)
        if change == "revoke":
            ok(client.post(f"/api/announcements/{a['id']}/training/trainers", json={"expected_version":1,"trainer_scope":"circle","trainer_ids":[]}))
        elif change == "withdraw":
            ok(client.post(f"/api/announcements/{a['id']}/withdraw", json={"expected_version":1,"reason":"撤下测试"}))
        else:
            ok(client.post(f"/api/announcements/{a['id']}/versions", json=form(client, confirmation_level=3, training_required=True,
                trainer_ids=[tr["id"]], expected_version=1, change_summary="要求更新")))
    try:
        login(client, "CMTEST01")
        assert sign(client, a, code["code"]).status_code in {400,409}
        login(client)
        rows = ok(client.get(f"/api/announcements/{a['id']}/records?version=1"))["items"]
        assert not next(r for r in rows if r["employee_id"] == person("CMTEST01"))["confirmed_at"]
    finally:
        if change in {"disable", "role"}:
            with SessionLocal() as db:
                if change == "disable":
                    db.query(UserAccount).filter_by(employee_id=tr["id"]).one().enabled = True
                else:
                    db.query(EmployeeRoleAssignment).filter_by(employee_id=tr["id"]).first().role_id = old_role
                db.commit()


def test_wrong_code_is_throttled_and_end_only_affects_own_code(client):
    a, tr = setup(client)
    code = generate(client, a)
    login(client, "CMTEST01")
    wrong = "999999" if code["code"] != "999999" else "000000"
    for i in range(5):
        assert verify(client, a, wrong).status_code == (429 if i==4 else 400)
    assert verify(client, a, code["code"]).status_code == 429
    with SessionLocal() as db:
        attempt = db.query(Attempt).filter_by(employee_id=person("CMTEST01")).one()
        attempt.locked_until = datetime.now()-timedelta(seconds=1)
        db.commit()
    assert verify(client, a, code["code"]).status_code == 200
    login(client, "TRTEST01")
    ok(client.post(f"/api/announcements/{a['id']}/training/end", json={"expected_version":1,"reason":"结束"}))
    login(client, "CMTEST01")
    assert sign(client, a, code["code"]).status_code == 400


def test_publisher_export_keeps_followup_snapshots_and_group_order(client):
    a, tr = setup(client, all_circles=True)
    code = generate(client, a)
    login(client, "CMTEST01")
    ok(sign(client, a, code["code"]))
    with SessionLocal() as db:
        db.get(Employee, tr["id"]).name = "后来改名的培训师"
        db.commit()
    login(client, "GSMTEST01")
    assert client.get(f"/api/announcements/{a['id']}/followup-export").status_code == 403
    login(client, "AMTEST01")
    data = ok(client.get(f"/api/announcements/{a['id']}/records"))
    response = client.get(f"/api/announcements/{a['id']}/followup-export")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    sheet = book["跟进记录"]
    headers = [c.value for c in sheet[6]]
    assert headers[:4] == ["景点圈","小组","工号","姓名"]
    rows = list(sheet.iter_rows(min_row=7,values_only=True))
    assert [(r[0],r[1],r[2]) for r in rows] == [(r["attraction_name"],r["group_name"],r["employee_no"]) for r in data["items"]]
    signed = next(r for r in rows if r[2]=="CMTEST01")
    assert signed[12] == tr["name"] and signed[13] == "TRTEST01" and signed[14] == code["id"]
    assert signed[15] and signed[9] and sheet.freeze_panes == "A7"
    assert code["code"] not in str(rows)
    from pathlib import Path
    target = Path(__file__).resolve().parents[2]/".handover_review"/"announcement-training-followup-20261011.xlsx"
    target.parent.mkdir(exist_ok=True)
    target.write_bytes(response.content)


def test_concurrent_duplicate_signs_record_one_evidence(client):
    a, tr = setup(client)
    code = generate(client, a)
    def submit(_):
        with TestClient(app) as c:
            login(c,"CMTEST01")
            return sign(c,a,code["code"])
    with ThreadPoolExecutor(max_workers=3) as pool:
        results = list(pool.map(submit, range(3)))
    assert all(r.status_code == 200 for r in results)
    assert len({r.json()["signature_url"] for r in results}) == 1
    with SessionLocal() as db:
        assert db.query(Evidence).count() == 1


def test_handover_preserves_codes_publisher_export_and_successor_control(client):
    a, tr = setup(client)
    code = generate(client,a)
    login(client)
    h = ok(client.post(f"/api/announcements/{a['id']}/handover",json={"expected_version":1,"successor_id":person("AMTEST01"),"reason":"正式交接"}))
    login(client,"AMTEST01")
    ok(client.post(f"/api/announcements/handovers/{h['handover_id']}/accept",json={}))
    login(client,"CMTEST01")
    assert verify(client,a,code["code"]).status_code == 200
    login(client)
    assert client.get(f"/api/announcements/{a['id']}/followup-export").status_code == 200
    assert client.post(f"/api/announcements/{a['id']}/training/trainers",json={"expected_version":1,"trainer_scope":"circle","trainer_ids":[]}).status_code == 403
    login(client,"AMTEST01")
    assert client.get(f"/api/announcements/{a['id']}/followup-export").status_code == 200
    ok(client.post(f"/api/announcements/{a['id']}/training/trainers",json={"expected_version":1,"trainer_scope":"circle","trainer_ids":[]}))
    login(client,"CMTEST01")
    assert sign(client,a,code["code"]).status_code == 400


def test_late_supplement_and_overdue_signing_do_not_need_republication(client):
    a, tr = setup(client)
    expired = generate(client,a)
    with SessionLocal() as db:
        db.get(Code,expired["id"]).expires_at = datetime.now()-timedelta(minutes=1)
        db.commit()
    id, no = new_employee(tr["attraction_id"])
    login(client,no)
    assert ok(client.get("/api/announcements?view=pending"))["items"][0]["id"] == a["id"]
    with SessionLocal() as db:
        delivery = db.query(AnnouncementDelivery).filter_by(employee_id=id,version_id=a["version_id"]).one()
        delivery.due_on = (datetime.now()-timedelta(days=1)).date().isoformat()
        db.commit()
    code = generate(client,a)
    login(client,no)
    assert ok(client.get(f"/api/announcements/{a['id']}"))["receipt"]["overdue"]
    assert sign(client,a,code["code"]).status_code == 200

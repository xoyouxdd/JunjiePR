from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app.excel_export import write_organization_score_sheet
from app.excel_export_utils import style_sheet
from app.excel_presentation import finalize_statistics_workbook, present_detail, save_export_workbook
from app.v2_watermark import watermark_workbook


def values(ws):
    return list(ws.values)


def test_long_wrapped_text_remains_complete_and_formula_does_not_drive_height():
    wb = Workbook()
    ws = wb.active
    text = "完整说明\n" + "长内容原样保留。" * 70
    ws.append(["员工号", "姓名", "内容", "综合分"])
    ws.append(["0012345", "长姓名 Alexandria", text, "=SUM(A2:A99999)"])
    style_sheet(ws)
    before = values(ws)
    present_detail(ws)
    assert values(ws) == before
    assert ws.row_dimensions[2].height > 24
    assert ws.print_title_rows == "$1:$1"
    assert ws.freeze_panes == "C2"
    assert ws["C2"].value == text
    assert ws["D2"].data_type == "f"


def test_folded_audit_fields_survive_save_and_are_expandable():
    wb = Workbook()
    ws = wb.active
    ws.append(["记录编号", "日期", "员工号", "姓名", "内容", "撤回人", "撤回时间", "撤回原因"])
    ws.append(["REC-1", "2026-09-12", "0012345", "模拟甲", "完整内容", "模拟主管", "2026-09-13 10:20:31", "重复记录"])
    style_sheet(ws)
    before = values(ws)
    present_detail(ws)
    output = BytesIO()
    save_export_workbook(wb, output)
    restored = load_workbook(BytesIO(output.getvalue())).active
    assert values(restored) == before
    assert restored.column_dimensions["F"].hidden
    assert restored.column_dimensions["F"].outlineLevel == 1
    assert restored.sheet_format.outlineLevelCol == 1
    assert not restored.column_dimensions["E"].hidden
    assert restored["G2"].value.endswith(":31")


def test_summary_order_preserves_all_sheets_rows_and_formula_links():
    wb = Workbook()
    wb.active.title = "导出说明"
    wb.active.append(["月份", "2026-09"])
    for name in ["月度综合分", "层级绩效明细", "签卡明细", "月度综合分明细", "员工留档标记", "主管月度综合分明细"]:
        ws = wb.create_sheet(name)
        ws.append(["员工号", "姓名", "综合分", "工号状态", "备注"])
        ws.append(["0012345", "模拟甲", "='签卡明细'!C2", "曾变更", "留档说明"])
    before = {ws.title: values(ws) for ws in wb if ws.title != "导出说明"}
    finalize_statistics_workbook(wb)
    finalize_statistics_workbook(wb)
    assert wb.sheetnames[:3] == ["导出说明", "月度综合分明细", "月度综合分"]
    assert wb.active.title == "导出说明"
    assert {ws.title: values(ws) for ws in wb if ws.title != "导出说明"} == before
    assert sum(c.value == "工作表目录" for row in wb["导出说明"] for c in row) == 1
    assert len([c for row in wb["导出说明"] for c in row if c.hyperlink]) == len(before)


def test_organization_identity_columns_match_headers_including_loa():
    wb = Workbook()
    ws = wb.active
    score = dict(employee_id=1, employee_no="0012345", employee_name="模拟甲", attraction_name="模拟圈",
                 recognition_score=1.5, attendance_score=12, deduction_score=0, total_score=13.5)
    data = {"scores": [score], "loa_rows": [{**score, "employee_id": 2, "employee_no": "0012346", "employee_name": "模拟乙"}],
            "hierarchy": [{"node_type": "employee", "employee_id": 1}]}
    write_organization_score_sheet(ws, deepcopy(data), {1: SimpleNamespace(code="CM")}, {2: SimpleNamespace(code="CM")})
    assert list(ws.values)[1][:2] == ("模拟甲", "0012345")
    assert list(ws.values)[3][:2] == ("模拟乙", "0012346")


def test_identical_watermark_media_is_shared_but_all_anchors_and_cells_remain():
    wb = Workbook()
    for i in range(5):
        ws = wb.create_sheet(f"明细{i}")
        ws.append(["编号", "分值", "公式"])
        ws.append([f"REC-{i}", -0.5, "=B2+1"])
    watermark_workbook(wb, "0000000", datetime(2026, 9, 30, 10, 20))
    original = BytesIO()
    wb.save(original)
    # Watermark image streams are consumed on save: build a fresh fixture for serialization.
    fixture = load_workbook(BytesIO(original.getvalue()))
    output = BytesIO()
    save_export_workbook(fixture, output)
    with ZipFile(original) as before, ZipFile(BytesIO(output.getvalue())) as after:
        assert len([n for n in after.namelist() if n.startswith("xl/media/")]) == 1
        for name in before.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                assert before.read(name) == after.read(name)
    restored = load_workbook(BytesIO(output.getvalue()))
    assert all(len(ws._images) == 1 for ws in restored)
    assert values(restored["明细0"]) == values(wb["明细0"])
    assert restored["明细0"]["B2"].value == -0.5
    assert len(output.getvalue()) < len(original.getvalue()) * 0.5


def test_ranking_export_fetches_all_rows_beyond_old_5000_limit(monkeypatch):
    from app.routers import statistics
    calls = []
    received = []

    def payload(*args):
        calls.append(args[9])  # page_size, after db/user/date/filter/page arguments
        return {"total": 5002, "rows": list(range(min(args[9], 5002)))}

    def builder(db, data, category):
        received.extend(data["rows"])
        return Workbook()

    monkeypatch.setattr(statistics, "pr_ranking_payload", payload)
    monkeypatch.setattr(statistics, "build_pr_rankings_workbook", builder)
    monkeypatch.setattr(statistics, "watermark_workbook", lambda *args: None)
    monkeypatch.setattr(statistics, "write_audit", lambda *args, **kwargs: None)
    user = SimpleNamespace(role=SimpleNamespace(code="GSM"), employee=SimpleNamespace(employee_no="0000000"))
    statistics.export_pr_rankings("2026-09-01", "2026-09-30", db=SimpleNamespace(commit=lambda: None), user=user)
    assert calls == [5000, 5002]
    assert received == list(range(5002))

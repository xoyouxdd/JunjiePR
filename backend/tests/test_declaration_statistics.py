from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from hashlib import sha256
import os
from pathlib import Path
import tempfile
from urllib.parse import unquote

TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-declaration-statistics-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from fastapi.testclient import TestClient  # noqa: E402
import fitz  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402
from app.v2_database import FILE_DIR, ROLE_PERMISSION_CODES, SessionLocal  # noqa: E402
from app.v2_crypto import hash_password  # noqa: E402
from app.v2_models import Attraction, AuditLog, DeductionLevel, DeductionRecord, DeductionType, DeductionUpgradeRequest, Employee, EmployeeRoleAssignment, Role, StoredFile, UserAccount  # noqa: E402


def login(client: TestClient, account: str, password: str = "1234") -> None:
    response = client.post("/api/login", json={"employee_no": account, "password": password})
    assert response.status_code == 200, response.text


def test_role_matrix_and_cross_circle_read_export() -> None:
    allowed = {"TA_SUPERVISOR", "SUPERVISOR", "TA_GSM", "GSM", "AM", "OM", "HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}
    for code, permissions in ROLE_PERMISSION_CODES.items():
        assert ("DECLARATION_STATS_VIEW" in permissions) == (code in allowed)
        assert ("DECLARATION_STATS_EXPORT" in permissions) == (code in allowed)
    with TestClient(app) as client:
        month = date.today().strftime("%Y-%m")
        for account in ("TATEST01", "SUPTEST01", "TAGSMTEST01", "GSMTEST01", "AMTEST01", "OMTEST01"):
            login(client, account)
            assert client.get("/api/declaration-statistics", params={"month": month}).status_code == 200
            assert client.get("/api/declaration-statistics/export", params={"month": month}).status_code == 200
            client.post("/api/logout")
        with SessionLocal() as db:
            scoped_hr = db.query(Employee).filter_by(employee_no="HR-HEAT").one()
            scoped_account = db.query(UserAccount).filter_by(employee_id=scoped_hr.id).one()
            scoped_account.password_hash = hash_password("1234")
            scoped_account.must_change_password = False
            scoped_account.enabled = True
            om = db.query(Employee).filter_by(employee_no="OMTEST01").one()
            hr_admin = db.query(Role).filter_by(code="HR_ADMIN").one()
            assignment = db.query(EmployeeRoleAssignment).filter_by(employee_id=om.id, status="active").one()
            assignment.role_id = hr_admin.id
            db.commit()
        for account in ("HR-HEAT", "OMTEST01"):
            login(client, account)
            assert client.get("/api/declaration-statistics", params={"month": month}).status_code == 200
            assert client.get("/api/declaration-statistics/export", params={"month": month}).status_code == 200
            client.post("/api/logout")
        login(client, "HR01", "HR123")
        assert client.get("/api/declaration-statistics").status_code == 200
        assert client.get("/api/declaration-statistics/export").status_code == 200
        client.post("/api/logout")
        login(client, "CMTEST01")
        assert client.get("/api/declaration-statistics").status_code == 403
        assert client.get("/api/declaration-statistics/export").status_code == 403


def test_event_snapshot_status_upgrade_and_export_match() -> None:
    with TestClient(app) as client:
        with SessionLocal() as db:
            employee = db.query(Employee).filter_by(employee_no="CMTEST01").one()
            circles = db.query(Attraction).filter(Attraction.employee_circle.is_(True)).order_by(Attraction.id).all()
            other_circle = next(row for row in circles if row.id != employee.attraction_id)
            level_statement = db.query(DeductionLevel).filter_by(code="STATEMENT").one()
            level_memo = db.query(DeductionLevel).filter_by(code="MEMO").one()
            dtype = db.query(DeductionType).first()
            file_row = StoredFile(storage_key="test/not-present.pdf", original_filename="测试.pdf", extension=".pdf", mime_type="application/pdf", file_size=1, sha256="0" * 64, uploaded_by=employee.id, status="active")
            db.add(file_row)
            db.flush()
            today = date.today().isoformat()

            def add(status: str, level=level_statement, upgrade_role=None):
                row = DeductionRecord(employee_id=employee.id, employee_no=employee.employee_no, employee_name=employee.name,
                    employee_role_snapshot="CM", attraction_id_snapshot=other_circle.id,
                    deduction_type_id=dtype.id, deduction_type_name=dtype.name,
                    deduction_level_id=level.id, deduction_level_name=level.name, points=Decimal("1.00"),
                    occurred_on=today, deduction_month=today[:7], description="统计测试", document_file_id=file_row.id,
                    submitter_id=employee.id, submitter_name="测试登记人", submitter_role_snapshot="CM",
                    permission_scope_snapshot="test", status=status, material_status="ready", upgrade_role=upgrade_role)
                db.add(row)
                db.flush()
                return row

            first = add("active", upgrade_role="source_first")
            second = add("active", upgrade_role="source_second")
            result = add("active", level_memo, "result")
            add("pending_material")
            add("pending_upgrade")
            add("void")
            upgrade = DeductionUpgradeRequest(employee_id=employee.id, deduction_type_id=dtype.id,
                first_deduction_id=first.id, second_deduction_id=second.id, reviewer_id=employee.id,
                reviewer_name="审核人", status="approved", result_level_id=level_memo.id,
                result_deduction_id=result.id, submitted_by=employee.id, submitted_by_name="测试登记人")
            db.add(upgrade)
            db.commit()
            second_id, result_id, file_id = second.id, result.id, file_row.id
            other_circle_id, other_circle_name = other_circle.id, other_circle.name
            employee_no, employee_attraction_id = employee.employee_no, employee.attraction_id
            memo_level_id = level_memo.id

        login(client, "TATEST01")
        params = {"month": date.today().strftime("%Y-%m"), "attraction_id": other_circle_id}
        response = client.get("/api/declaration-statistics", params=params)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["total_records"] == 5
        assert body["employee_count"] == 1
        assert {row["id"] for row in body["records"]}.isdisjoint({second_id})
        assert result_id in {row["id"] for row in body["records"]}
        assert {row["attraction_id"] for row in body["records"]} == {other_circle_id}
        assert body["statuses"] == {"已生效": 2, "待补材料": 1, "审核中": 1, "已作废": 1}
        assert sum(row["count"] for row in body["levels"]) == 5
        assert body["circles"][0]["attraction_name"] == other_circle_name
        assert client.get("/api/declaration-statistics", params={**params, "status": "已作废"}).json()["total_records"] == 1
        assert client.get("/api/declaration-statistics", params={**params, "level_id": memo_level_id}).json()["total_records"] == 1
        assert client.get("/api/declaration-statistics", params={**params, "keyword": employee_no[-4:]}).json()["total_records"] == 5
        assert client.get("/api/declaration-statistics", params={"month": params["month"], "attraction_id": employee_attraction_id}).json()["total_records"] == 0
        assert client.get("/api/files/" + str(file_id), params={"preview": "true"}).status_code == 404  # authorized; test file has no disk bytes
        exported = client.get("/api/declaration-statistics/export", params=params)
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content))
        assert workbook.sheetnames == ["景点圈与等级汇总", "逐条明细"]
        assert workbook.worksheets[0]["B2"].value == body["total_records"]
        assert workbook.worksheets[1].max_row - 1 == body["total_records"]
        assert {workbook.worksheets[1].cell(row, 2).value for row in range(2, workbook.worksheets[1].max_row + 1)} == {employee_no}
        with SessionLocal() as db:
            assert db.query(AuditLog).filter_by(action="导出声明登记统计").count() >= 1


def test_ui_keeps_level_text_and_mobile_cards() -> None:
    root = Path(__file__).resolve().parents[1] / "app" / "static"
    script = (root / "js" / "app.js").read_text(encoding="utf-8")
    css = (root / "css" / "style.css").read_text(encoding="utf-8")
    assert "声明登记统计" in script
    assert "${esc(row.level_name)}" in script
    assert "declaration-mobile" in script and "declaration-desktop" in script
    assert "data-declaration-circle" in script and "data-declaration-level" in script
    assert 'href="${esc(portalPath(row.document_url))}" download>下载材料</a>' in script
    for tone in range(6):
        assert f".tone-{tone}" in css
    assert ".declaration-filter { grid-template-columns: 1fr; }" in css


def test_direct_material_download_uses_event_snapshot_name_and_keeps_authorization() -> None:
    with TestClient(app) as client:
        document = fitz.open()
        document.new_page().insert_text((50, 50), "Declaration material")
        content = document.tobytes()
        document.close()
        path = FILE_DIR / "declaration-download-test.pdf"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        with SessionLocal() as db:
            employee = db.query(Employee).filter_by(employee_no="CMTEST01").one()
            dtype = db.query(DeductionType).first()
            level = db.query(DeductionLevel).filter_by(code="STATEMENT").one()
            stored = StoredFile(storage_key=path.name, original_filename="原始材料.pdf", extension=".pdf",
                mime_type="application/pdf", file_size=len(content), sha256=sha256(content).hexdigest(),
                uploaded_by=employee.id, status="active")
            db.add(stored)
            db.flush()
            db.add(DeductionRecord(employee_id=employee.id, employee_no=employee.employee_no,
                employee_name="凌元富", employee_role_snapshot="CM", attraction_id_snapshot=employee.attraction_id,
                deduction_type_id=dtype.id, deduction_type_name="考勤-未打卡", deduction_level_id=level.id,
                deduction_level_name=level.name, points=Decimal("1.00"), occurred_on="2026-09-24",
                deduction_month="2026-09", description="下载命名测试", document_file_id=stored.id,
                submitter_id=employee.id, submitter_name=employee.name, submitter_role_snapshot="CM",
                permission_scope_snapshot="test", status="active", material_status="ready"))
            db.commit()
            file_id = stored.id
        login(client, "TATEST01")
        response = client.get(f"/api/files/{file_id}")
        assert response.status_code == 200, response.text
        disposition = unquote(response.headers["content-disposition"])
        assert disposition == "attachment; filename*=UTF-8''2026-09-24_凌元富_考勤-未打卡.pdf"
        downloaded = fitz.open(stream=response.content, filetype="pdf")
        assert "TATEST01" in "".join(page.get_text() for page in downloaded)
        downloaded.close()
        preview = client.get(f"/api/files/{file_id}", params={"preview": "true"})
        assert preview.status_code == 200
        assert unquote(preview.headers["content-disposition"]) == "inline; filename*=UTF-8''原始材料.pdf"
        with SessionLocal() as db:
            assert db.get(StoredFile, file_id).original_filename == "原始材料.pdf"
            assert db.query(AuditLog).filter_by(action="下载带水印材料", entity_id=file_id).count() >= 1
        client.post("/api/logout")
        login(client, "CMTEST02")
        assert client.get(f"/api/files/{file_id}").status_code == 403


def test_material_download_filename_removes_path_and_control_characters() -> None:
    from app.routers.files import deduction_download_name
    row = DeductionRecord(occurred_on="2026-09-24", employee_name="姓名/测试", deduction_type_name='类型:未打卡\\测试\n')
    name = deduction_download_name(row, ".pdf")
    assert name == "2026-09-24_姓名-测试_类型-未打卡-测试-.pdf"

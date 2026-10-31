"""Leading-zero aliases use the same persisted identity and account controls."""
import json
from datetime import datetime, timedelta, date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from test_employee_types import setup, app, TestClient, SessionLocal, login, book, circle
from app.employee_numbers import employee_number, employee_number_aliases, employee_number_cell
from app.v2_models import Employee, UserAccount, EmployeeNumberHistory
from app.rotation.roster import parse_any, parse_roster


@pytest.mark.parametrize('raw,want',[
    ('1234567','1234567'),('01234567','1234567'),(' 01234567 ','1234567'),
    ('00000001','0000001'),('001234567','001234567'),('11234567','11234567'),
    ('0123456','0123456'),('HR01','HR01'),('HR-HEAT','HR-HEAT'),
    ('０1234567','０1234567'),('0123456A','0123456A'),('', ''), (None,''),
])
def test_alias_is_strictly_one_zero_before_seven_ascii_digits(raw,want):
    assert employee_number(raw)==want
    assert employee_number_cell(1234567.0)=='1234567'
    assert employee_number_aliases('01234567')==('1234567','01234567')


def create(c, number, name='合成兼容员工'):
    response=c.post('/api/hr/employees',json={'employee_no':number,'name':name,'role_code':'CM','attraction_id':circle(c)})
    assert response.status_code==200,response.text
    return response.json()['employee']['id']


def test_creation_login_search_and_number_change_share_one_identity():
    with TestClient(app) as c:
        login(c);eid=create(c,'09871001')
        with SessionLocal() as db:
            e=db.get(Employee,eid);a=db.query(UserAccount).filter_by(employee_id=eid).one()
            assert e.employee_no==a.login_account=='9871001'
            a.must_change_password=False;db.commit()
        duplicate=c.post('/api/hr/employees',json={'employee_no':'9871001','name':'重复','role_code':'CM','attraction_id':circle(c)})
        assert duplicate.status_code==400
        for endpoint in ['/api/hr/employee-number-targets','/api/accounts/reset-targets','/api/accounts/name-targets']:
            seven=c.get(endpoint,params={'keyword':'9871001'});eight=c.get(endpoint,params={'keyword':'09871001'})
            assert seven.status_code==eight.status_code==200
            assert seven.json()==eight.json() and any(r['id']==eid for r in eight.json()['items'])
        unchanged=c.post(f'/api/hr/employees/{eid}/employee-number',json={'new_employee_no':'09871001','reason':'仅前导零'})
        assert unchanged.status_code==200 and unchanged.json()['unchanged']
        with SessionLocal() as db:assert not db.query(EmployeeNumberHistory).filter_by(employee_id=eid).first()
        for no in ['9871001','09871001']:
            c.post('/api/logout');r=c.post('/api/login',json={'employee_no':no,'password':'1001'})
            assert r.status_code==200,r.text
            assert c.get('/api/me').json()['id']==eid
        with SessionLocal() as db:
            a=db.query(UserAccount).filter_by(employee_id=eid).one();a.locked_until=datetime.now()+timedelta(minutes=5);db.commit()
        assert c.post('/api/login',json={'employee_no':'09871001','password':'1001'}).status_code==423
        with SessionLocal() as db:
            a=db.query(UserAccount).filter_by(employee_id=eid).one();a.locked_until=None;a.enabled=False;db.commit()
        assert c.post('/api/login',json={'employee_no':'09871001','password':'1001'}).status_code==401


def test_alias_import_updates_existing_and_rejects_duplicate_normalized_numbers():
    with TestClient(app) as c:
        login(c);eid=create(c,'9871002')
        files=book([['09871002','合成兼容员工','CM','热力追踪','Comp40','']])
        choice=json.dumps([{'row':2,'action':'update'}])
        preview=c.post('/api/hr/import-employees',files=files,data={'preview':'true','decisions':choice})
        assert preview.status_code==200,preview.text
        batch=preview.json();assert batch['would_update']==1 and batch['would_create']==0
        confirm=c.post('/api/hr/import-employees',files=files,data={'expected_digest':batch['digest'],'decisions':choice})
        assert confirm.status_code==200,confirm.text
        with SessionLocal() as db:
            e=db.get(Employee,eid);assert e.employee_no=='9871002' and e.employee_type=='Comp40'
            assert db.query(Employee).filter(Employee.employee_no.in_(['9871002','09871002'])).count()==1
        mixed=book([['9871003','重复人','CM','热力追踪','',''],['09871003','重复人','CM','热力追踪','','']])
        response=c.post('/api/hr/import-employees',files=mixed,data={'preview':'true'})
        assert response.status_code==400 and '重复' in response.text
        numeric=book([[9871004.0,'数值单元格','CM','热力追踪','','']])
        response=c.post('/api/hr/import-employees',files=numeric)
        assert response.status_code==200,response.text


def roster(rows):
    wb=Workbook();wb.active.append(['工号','姓名','日期','班次','备注'])
    for r in rows:wb.active.append(r)
    out=BytesIO();wb.save(out);out.seek(0);return out


def test_standard_roster_alias_and_duplicate_day_guard():
    day=date.today().isoformat();tomorrow=(date.today()+timedelta(days=1)).isoformat()
    result=parse_any(roster([['09871005','班表员工',day,'08:00-16:00',''],['9871005','班表员工',tomorrow,'08:00-16:00','']]))
    assert set(result['people'])=={'9871005'}
    assert all(set(people)=={'9871005'} for people in result['days'].values())
    with pytest.raises(ValueError,match='重复'):
        parse_any(roster([['09871005','班表员工',day,'08:00-16:00',''],['9871005','班表员工',day,'08:00-16:00','']]))


def test_ztp_roster_alias_preserves_shift_and_days():
    wb=Workbook();ws=wb.active
    ws.append(['','','',46000,46001,46002,46003,46004,46005,46006])
    ws.append(['Name','ID','Type','Sun','Mon','Tue','Wed','Thu','Fri','Sat'])
    ws.append(['班表员工','09871006','CM',*(['08:00-16:00']*7)])
    out=BytesIO();wb.save(out);out.seek(0)
    parsed=parse_roster(out);assert set(parsed['people'])=={'9871006'}
    assert len(parsed['days'])==7 and all(set(p)=={'9871006'} for p in parsed['days'].values())


def test_export_filter_alias_returns_same_employee_and_canonical_workbook_values():
    with TestClient(app) as c:
        login(c);eid=create(c,'09871007');month=date.today().strftime('%Y-%m')
        seven=c.get('/api/statistics',params={'month':month,'keyword':'9871007'})
        eight=c.get('/api/statistics',params={'month':month,'keyword':'09871007'})
        assert seven.status_code==eight.status_code==200
        assert seven.json()['summary']==eight.json()['summary']
        assert seven.json()['hierarchy']==eight.json()['hierarchy']
        assert eight.json()['summary']['employee_count']==1
        assert '9871007' in json.dumps(eight.json()['hierarchy'])
        result=c.get('/api/statistics/export',params={'month':month,'keyword':'09871007'})
        assert result.status_code==200,result.text
        wb=load_workbook(BytesIO(result.content),read_only=True,data_only=True)
        cells=[v for ws in wb for row in ws.iter_rows(values_only=True) for v in row]
        assert '9871007' in cells
        wb.close()


def test_alias_cannot_bypass_failure_counter_or_select_an_ambiguous_account():
    with TestClient(app) as c:
        login(c);first=create(c,'9871008');second=create(c,'9871009')
        for i in range(5):
            no='09871008' if i%2 else '9871008'
            assert c.post('/api/login',json={'employee_no':no,'password':'wrong'}).status_code==401
        assert c.post('/api/login',json={'employee_no':'09871008','password':'1008'}).status_code==423
        with SessionLocal() as db:
            a=db.query(UserAccount).filter_by(employee_id=first).one();a.locked_until=None
            b=db.query(UserAccount).filter_by(employee_id=second).one();b.login_account='09871008';db.commit()
        for no in ['9871008','09871008']:
            assert c.post('/api/login',json={'employee_no':no,'password':'1008'}).status_code==401

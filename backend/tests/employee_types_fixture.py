"""Isolated synthetic employee-Type UI backend; never opens user data."""
import sys
from pathlib import Path
from io import BytesIO
from datetime import date,timedelta
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from test_employee_types import setup,app
from fastapi.responses import Response
from openpyxl import Workbook
generator=setup.__wrapped__();next(generator)

@app.get('/__fixture/workbook')
def workbook(existing: str, new: str):
    wb=Workbook();ws=wb.active
    ws.append(['员工号','姓名','角色代码','景点圈','Type','实习结束日期'])
    ws.append([existing,'不应覆盖旧姓名','CM','热力追踪','Comp28',''])
    ws.append([new,'导入合成实习生','tr','热力追踪','Intern',(date.today()+timedelta(days=60)).isoformat()])
    out=BytesIO();wb.save(out);return Response(out.getvalue(),media_type='application/octet-stream')


@app.get('/__fixture/occupied-workbook')
def occupied_workbook(index: int = 0):
    from test_employee_import_conflicts import seed,circle_id
    foreign=str(9852000+index*10);own=str(int(foreign)+1);fresh=str(int(foreign)+2)
    seed(foreign,circle_id('热力追踪'))
    seed(own,circle_id('矮人迷宫'))
    wb=Workbook();ws=wb.active
    ws.append(['员工号','姓名','角色代码','景点圈','Type','实习结束日期'])
    for no in (foreign,own,fresh):
        ws.append([no,'导入文件名字'+no,'TR','矮人迷宫','Comp40',''])
    out=BytesIO();wb.save(out);return Response(out.getvalue(),media_type='application/octet-stream')

@app.get('/__fixture/simple-workbook')
def simple_workbook(index: int = 0):
    from test_employee_import_conflicts import seed,circle_id
    from test_employee_types import SessionLocal
    from app.v2_models import Employee
    wb=Workbook();ws=wb.active
    ws.append(['员工号','姓名','角色代码','景点圈','Type','实习结束日期'])
    names=['张林','李青','王晨','刘洋','陈雨','周宁','赵晴']
    for n,name in enumerate(names):
        no=str(9860000+index*100+n)
        eid=seed(no,circle_id('热力追踪' if n==5 else '矮人迷宫'))
        with SessionLocal() as db:
            e=db.get(Employee,eid);e.name=name
            e.employee_type=None if n==2 else 'Intern' if n==1 else 'Comp40' if n==3 else 'Comp28'
            e.internship_ends_on='2099-11-30' if n==1 else None
            db.commit()
        kind='Intern' if n in (1,2) else '' if n==6 else 'Comp40'
        ws.append([no,'陈宇' if n==4 else name,'CM','矮人迷宫',kind,'2099-12-31' if n in (1,2) else ''])
    out=BytesIO();wb.save(out);return Response(out.getvalue(),media_type='application/octet-stream')

import uvicorn
uvicorn.run(app,host='127.0.0.1',port=int(sys.argv[1]),log_level='warning')

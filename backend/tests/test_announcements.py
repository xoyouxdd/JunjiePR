"""Operational announcement integration contracts; isolated database and files."""
from __future__ import annotations

import base64
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from io import BytesIO
import os
from pathlib import Path
import tempfile
from uuid import uuid4

import fitz
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
import pytest

DATA = Path(tempfile.mkdtemp(prefix="recognition-announcements-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(DATA)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from app.main import app
from app.v2_database import SessionLocal, FILE_DIR
from app.v2_models import (Announcement, AnnouncementVersion, AnnouncementDelivery, AnnouncementNotification,
    AnnouncementTrainingConfig, AnnouncementTrainer, AnnouncementTrainingCode, AnnouncementTrainingReceipt, AnnouncementTrainingAttempt,
    AnnouncementHandover, AnnouncementResponsibility, AnnouncementResponsibilityHandover, AnnouncementFavorite, AnnouncementGrant, AnnouncementProjectMember,
    AnnouncementProject, AnnouncementAsset, Employee, EmployeeLOAPeriod, EmployeeRoleAssignment,
    GroupMembership, Role, StoredFile, UserAccount, WorkGroup)
from app.v2_crypto import hash_password
from app import announcements as domain


def login(client, name="GSMTEST01"):
    response = client.post("/api/login", json={"employee_no": name, "password": "HR123" if name == "HR01" else "1234"})
    assert response.status_code == 200, response.text
    return response.json()


def ok(response):
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture
def client():
    with TestClient(app) as client:
        with SessionLocal() as db:
            for model in [AnnouncementResponsibilityHandover, AnnouncementResponsibility, AnnouncementTrainingReceipt, AnnouncementTrainingAttempt, AnnouncementTrainingCode, AnnouncementTrainer, AnnouncementTrainingConfig,
                          AnnouncementHandover, AnnouncementFavorite, AnnouncementNotification, AnnouncementDelivery,
                          AnnouncementVersion, Announcement, AnnouncementGrant, AnnouncementProjectMember,
                          AnnouncementProject, AnnouncementAsset]:
                db.query(model).delete()
            originals = {e.id: (e.name, e.attraction_id, e.is_active, e.hired_on, e.terminated_on)
                         for e in db.query(Employee)}
            db.commit()
        yield client
        with SessionLocal() as db:
            for id, values in originals.items():
                e = db.get(Employee, id)
                e.name, e.attraction_id, e.is_active, e.hired_on, e.terminated_on = values
            db.query(EmployeeLOAPeriod).filter_by(note="announcement-test").delete()
            for e in db.query(Employee).filter(Employee.employee_no.like("NOTICE-%")):
                e.is_active = False
            db.commit()


def form(client, **changes):
    options = ok(client.get("/api/announcements/options"))
    result = {"title": "安全检查要求", "summary": "当班开始前完成巡查", "body": "保留行\n检查护栏\n发现异常及时报告",
              "category": options["categories"][0], "scope_kind": "circle", "attraction_ids": options["publisher"]["circles"],
              "tiers": ["frontline", "lead", "gsm", "am"], "confirmation_level": 2,
              "effective_on": date.today().isoformat(), "request_key": uuid4().hex}
    return {**result, **changes}


def publish(client, **changes):
    return ok(client.post("/api/announcements", json=form(client, **changes)))


def receipt(client, a, **changes):
    return client.post(f"/api/announcements/{a['id']}/receipt", json={"version": a["version"], **changes})


def signature(blank=False):
    image = Image.new("RGB", (640, 160), "white")
    if not blank:
        ImageDraw.Draw(image).line([(60, 110), (160, 25), (130, 130), (260, 30), (340, 115), (420, 40), (510, 100)], fill="black", width=5)
    stream = BytesIO()
    image.save(stream, format="PNG")
    return "data:image/png;base64,"+base64.b64encode(stream.getvalue()).decode()


def config(client):
    login(client, "HR01")
    return ok(client.get("/api/announcements/admin/config"))


def employee(config, number):
    return next(e for e in config["employees"] if e["employee_no"] == number)


def new_employee(circle, *, future=False, disabled=False, terminated=False):
    with SessionLocal() as db:
        e = Employee(employee_no="NOTICE-"+uuid4().hex[:12], name="新伙伴", attraction_id=circle,
                     hired_on=(date.today()+timedelta(days=1)).isoformat() if future else date.today().isoformat(),
                     terminated_on=date.today().isoformat() if terminated else None)
        db.add(e)
        db.flush()
        db.add(UserAccount(employee_id=e.id, login_account=e.employee_no, password_hash=hash_password("1234"), enabled=not disabled, must_change_password=False))
        role = db.query(Role).filter_by(code="CM").one()
        db.add(EmployeeRoleAssignment(employee_id=e.id, role_id=role.id, starts_on="2020-01-01"))
        db.commit()
        return e.id, e.employee_no


def test_circle_defaults_and_admin_grants_cannot_be_self_escalated(client):
    login(client)
    options = ok(client.get("/api/announcements/options"))
    assert len(options["publisher"]["circles"]) == 1
    circles = [c["id"] for c in options["circles"]]
    assert len(circles) == 3
    assert client.post("/api/announcements", json=form(client, attraction_ids=circles)).status_code == 403
    assert client.get("/api/announcements/admin/config").status_code == 403
    cfg = config(client)
    gsm = employee(cfg, "GSMTEST01")
    grant = {"employee_id": gsm["id"], "attraction_ids": circles, "starts_on": date.today().isoformat(), "reason": "批准三个圈发布"}
    saved = ok(client.post("/api/announcements/admin/grants", json=grant))
    login(client)
    assert len(ok(client.get("/api/announcements/options"))["publisher"]["circles"]) == 3
    a = publish(client)
    login(client, "HR01")
    updated = ok(client.post("/api/announcements/admin/grants", json={**grant, "enabled": False}))
    assert updated["id"] == saved["id"]
    rows = [g for g in ok(client.get("/api/announcements/admin/config"))["grants"] if g["employee_id"] == gsm["id"]]
    assert len(rows) == 1 and rows[0]["enabled"] is False
    login(client)
    assert not ok(client.get("/api/announcements/options"))["can_publish"]
    assert client.post("/api/announcements", json=form(client)).status_code in {400, 403}
    assert client.get(f"/api/announcements/{a['id']}/records").status_code == 200
    for number in ["TAGSMTEST01", "AMTEST01"]:
        login(client, number)
        assert len(ok(client.get("/api/announcements/options"))["publisher"]["circles"]) == 3


@pytest.mark.parametrize("account", ["CMTEST01", "SUPTEST01", "TATEST01", "OMTEST01", "HR01"])
def test_ordinary_nonpublishers_receive_but_cannot_publish_without_project_grant(client, account):
    login(client)
    values = form(client)
    login(client, account)
    assert client.post("/api/announcements", json=values).status_code == 403
    assert not ok(client.get("/api/announcements/options"))["can_publish"]


def test_expired_or_future_explicit_grant_does_not_fall_back_to_defaults(client):
    cfg = config(client)
    gsm = employee(cfg, "GSMTEST01")
    today = date.today()
    values = {"employee_id": gsm["id"], "attraction_ids": [gsm["attraction_id"]], "reason": "限期授权",
              "starts_on": (today-timedelta(days=3)).isoformat(), "ends_on": (today-timedelta(days=1)).isoformat()}
    ok(client.post("/api/announcements/admin/grants", json=values))
    login(client)
    assert not ok(client.get("/api/announcements/options"))["can_publish"]
    login(client, "HR01")
    ok(client.post("/api/announcements/admin/grants", json={**values, "starts_on": (today+timedelta(days=2)).isoformat(), "ends_on": None}))
    login(client)
    assert not ok(client.get("/api/announcements/options"))["can_publish"]


@pytest.mark.parametrize("level", [1, 2, 3])
def test_three_receipt_levels_archive_and_do_not_overwrite_evidence(client, level):
    login(client)
    a = publish(client, confirmation_level=level)
    login(client, "CMTEST01")
    assert any(x["id"] == a["id"] for x in ok(client.get("/api/announcements"))["items"])
    read = ok(receipt(client, a, action="read"))
    assert read["seen_at"]
    assert bool(read["confirmed_at"]) == (level == 1)
    if level > 1:
        assert receipt(client, a, action="confirm").status_code == 400
        if level == 3:
            assert receipt(client, a, action="confirm", acknowledged=True, signature=signature(True)).status_code == 400
        confirmed = ok(receipt(client, a, action="confirm", acknowledged=True, signature=signature() if level == 3 else None))
        repeat = ok(receipt(client, a, action="confirm", acknowledged=True, signature=signature(True)))
        assert repeat["confirmed_at"] == confirmed["confirmed_at"]
        assert repeat["signature_url"] == confirmed["signature_url"]
        if level == 3:
            assert client.get(confirmed["signature_url"]).status_code == 200
    assert not any(x["id"] == a["id"] for x in ok(client.get("/api/announcements"))["items"])
    archive = ok(client.get("/api/announcements?view=history&q=护栏"))["items"]
    assert archive[0]["effective"] and archive[0]["version"] == 1
    ok(client.post(f"/api/announcements/{a['id']}/favorite"))
    assert ok(client.get("/api/announcements?view=favorites"))["total"] == 1


def test_update_binds_receipt_to_new_version_and_yellows_only_changed_lines(client):
    login(client)
    first = form(client, confirmation_level=3)
    a = ok(client.post("/api/announcements", json=first))
    login(client, "CMTEST01")
    old = ok(receipt(client, a, action="confirm", acknowledged=True, signature=signature()))
    login(client)
    revision = {**first, "body": "保留行\n新增检查门锁\n发现异常及时报告", "change_summary": "检查项改为门锁",
                "expected_version": 1, "request_key": uuid4().hex}
    updated = ok(client.post(f"/api/announcements/{a['id']}/versions", json=revision))
    assert updated["version"] == 2
    assert [x["changed"] for x in updated["body_lines"]] == [False, True, False]
    assert ok(client.post(f"/api/announcements/{a['id']}/versions", json=revision))["version"] == 2
    assert client.post(f"/api/announcements/{a['id']}/versions", json={**revision, "request_key": uuid4().hex}).status_code == 409
    login(client, "CMTEST01")
    pending = ok(client.get("/api/announcements"))["items"]
    assert pending[0]["version"] == 2 and not pending[0]["receipt"]["confirmed_at"]
    assert receipt(client, a, action="confirm", acknowledged=True, signature=signature()).status_code == 409
    historic = ok(client.get(f"/api/announcements/{a['id']}?version=1"))
    assert historic["receipt"]["signature_url"] == old["signature_url"]
    notices = ok(client.get("/api/announcements/notifications"))["items"]
    assert any(n["kind"] == "updated" for n in notices)
    ok(receipt(client, updated, action="read"))
    assert all(n["read_at"] for n in ok(client.get("/api/announcements/notifications"))["items"])


def test_idempotent_publication_rejects_changed_payload(client):
    login(client)
    values = form(client)
    a = ok(client.post("/api/announcements", json=values))
    assert ok(client.post("/api/announcements", json=values))["id"] == a["id"]
    assert client.post("/api/announcements", json={**values, "title": "冲突正文"}).status_code == 409
    with SessionLocal() as db:
        assert db.query(Announcement).count() == 1
        assert db.query(AnnouncementVersion).count() == 1


def test_catchup_new_hire_return_transfer_and_exclusion_rules(client):
    cfg = config(client)
    cm = employee(cfg, "CMTEST01")
    login(client)
    yesterday = (date.today()-timedelta(days=1)).isoformat()
    a = publish(client, effective_on=yesterday, due_on=yesterday)
    with SessionLocal() as db:
        row = db.get(Announcement, a["id"])
        row.published_at = datetime.now()-timedelta(days=2)
        db.add(EmployeeLOAPeriod(employee_id=cm["id"], starts_on=yesterday, ends_on=None,
                                created_by=cm["id"], created_by_name="测试", note="announcement-test"))
        db.commit()
    login(client, "CMTEST01")
    assert ok(client.get("/api/announcements"))["total"] == 0
    assert ok(client.get(f"/api/announcements/{a['id']}"))["receipt"]["status"] == "suspended"
    with SessionLocal() as db:
        db.query(EmployeeLOAPeriod).filter_by(note="announcement-test").update({"status": "ended"})
        db.commit()
    resumed = ok(client.get("/api/announcements"))["items"][0]["receipt"]
    assert resumed["reason"] == "返岗补收"
    assert resumed["due_on"] == (date.today()+timedelta(days=3)).isoformat()
    new_id, number = new_employee(cm["attraction_id"])
    future, _ = new_employee(cm["attraction_id"], future=True)
    disabled, _ = new_employee(cm["attraction_id"], disabled=True)
    terminated, _ = new_employee(cm["attraction_id"], terminated=True)
    login(client, number)
    catchup = ok(client.get("/api/announcements"))["items"][0]["receipt"]
    assert catchup["reason"] == "适用公告补收"
    with SessionLocal() as db:
        assert not db.query(AnnouncementDelivery).filter(AnnouncementDelivery.employee_id.in_([future, disabled, terminated])).count()
        db.get(Employee, new_id).attraction_id = next(id for id in domain.live_circles(db) if id != cm["attraction_id"])
        db.commit()
    assert ok(client.get("/api/announcements"))["total"] == 0
    assert ok(client.get(f"/api/announcements/{a['id']}"))["receipt"]["status"] == "exited"


def test_future_and_expired_notices_never_create_new_required_tasks(client):
    login(client)
    tomorrow = (date.today()+timedelta(days=1)).isoformat()
    yesterday = (date.today()-timedelta(days=1)).isoformat()
    publish(client, effective_on=tomorrow)
    publish(client, effective_on=yesterday, expires_on=yesterday)
    login(client, "CMTEST01")
    assert ok(client.get("/api/announcements"))["total"] == 0
    with SessionLocal() as db:
        assert db.query(AnnouncementDelivery).count() == 0


def test_project_poc_is_explicit_and_cannot_publish_to_circles(client):
    cfg = config(client)
    cm = employee(cfg, "CMTEST01")
    other = employee(cfg, "CMTEST02")
    p = ok(client.post("/api/announcements/admin/projects", json={"name": "新演出项目", "employee_ids": [cm["id"]]}))
    grant = {"employee_id": cm["id"], "scope_kind": "project", "project_id": p["id"], "starts_on": date.today().isoformat(), "reason": "项目POC"}
    ok(client.post("/api/announcements/admin/grants", json=grant))
    login(client, "CMTEST01")
    options = ok(client.get("/api/announcements/options"))
    assert options["can_publish"] and not options["publisher"]["circles"]
    assert client.get("/api/announcement-media/options").status_code == 200
    assert client.post("/api/announcements", json=form(client, attraction_ids=[cm["attraction_id"]])).status_code == 403
    a = publish(client, scope_kind="project", attraction_ids=[], project_id=p["id"], tiers=["frontline"])
    login(client, "CMTEST02")
    assert client.get(f"/api/announcements/{a['id']}").status_code == 404
    login(client, "HR01")
    ok(client.post("/api/announcements/admin/projects", json={"id": p["id"], "name": "新演出项目", "employee_ids": [cm["id"], other["id"]]}))
    login(client, "CMTEST02")
    assert ok(client.get("/api/announcements"))["total"] == 1


@pytest.mark.parametrize("account", ["GSMTEST01", "TAGSMTEST01", "AMTEST01"])
def test_manager_circle_and_training_project_grants_are_independent(client, account):
    cfg = config(client)
    manager = employee(cfg, account)
    member = employee(cfg, "CMTEST01")
    project = ok(client.post("/api/announcements/admin/projects", json={
        "name": "培训项目", "employee_ids": [member["id"]]}))
    circle_values = {"employee_id": manager["id"], "attraction_ids": [manager["attraction_id"]],
                     "categories": list(domain.CATEGORIES), "starts_on": date.today().isoformat(), "reason": "圈负责人"}
    circle_grant = ok(client.post("/api/announcements/admin/grants", json=circle_values))
    project_values = {"employee_id": manager["id"], "scope_kind": "project", "project_id": project["id"],
                      "categories": ["培训"], "tiers": ["frontline"],
                      "starts_on": date.today().isoformat(), "reason": "兼任培训项目POC"}
    project_grant = ok(client.post("/api/announcements/admin/grants", json=project_values))
    assert circle_grant["id"] != project_grant["id"]
    login(client, account)
    options = ok(client.get("/api/announcements/options"))
    assert "培训" in options["categories"] and len(options["categories"]) == 6
    assert options["publisher"]["circles"] == [manager["attraction_id"]]
    assert options["publisher"]["projects"][0]["categories"] == ["培训"]
    circle_notice = publish(client, category="培训")
    project_notice = publish(client, scope_kind="project", attraction_ids=[], project_id=project["id"],
                             category="培训", tiers=["frontline"])
    assert circle_notice["scope_kind"] == "circle" and project_notice["scope_kind"] == "project"
    assert client.post("/api/announcements", json=form(client, scope_kind="project", attraction_ids=[],
        project_id=project["id"], category=domain.CATEGORIES[0], tiers=["frontline"])).status_code == 403
    assert client.post("/api/announcements", json=form(client, scope_kind="project", attraction_ids=[],
        project_id=project["id"], category="培训", tiers=["gsm"])).status_code == 403
    login(client, "CMTEST01")
    assert client.get(f"/api/announcements/{project_notice['id']}").status_code == 200
    assert all(a["category"] == "培训" for a in ok(client.get("/api/announcements?category=培训"))["items"])
    login(client, "HR01")
    ok(client.post("/api/announcements/admin/grants", json={**project_values, "enabled": False}))
    login(client, account)
    options = ok(client.get("/api/announcements/options"))
    assert options["can_publish"] and options["publisher"]["circles"] and not options["publisher"]["projects"]
    login(client, "HR01")
    ok(client.post("/api/announcements/admin/grants", json=project_values))
    # Existing explicit restrictions must not inherit the new category.
    ok(client.post("/api/announcements/admin/grants", json={**circle_values, "categories": domain.CATEGORIES[:-1]}))
    login(client, account)
    assert client.post("/api/announcements", json=form(client, category="培训")).status_code == 403
    login(client, "HR01")
    ok(client.post("/api/announcements/admin/grants", json={**circle_values, "enabled": False}))
    login(client, account)
    options = ok(client.get("/api/announcements/options"))
    assert options["can_publish"] and not options["publisher"]["circles"] and options["publisher"]["projects"]


def test_handover_requires_acceptance_and_admin_can_force_qualified_successor(client):
    cfg = config(client)
    ta, am = employee(cfg, "TAGSMTEST01"), employee(cfg, "AMTEST01")
    login(client)
    a = publish(client)
    h = ok(client.post(f"/api/announcements/{a['id']}/handover", json={"successor_id": ta["id"], "reason": "岗位交接", "expected_version": 1}))
    overview = ok(client.get(f"/api/announcements/{a['id']}/records"))["handover"]
    assert overview["owner"]["employee_no"] == "GSMTEST01"
    assert overview["history"][0]["status"] == "pending"
    assert overview["history"][0]["to"]["id"] == ta["id"]
    assert overview["history"][0]["requested_at"] and not overview["history"][0]["completed_at"]
    login(client, "AMTEST01")
    assert client.post(f"/api/announcements/handovers/{h['handover_id']}/accept").status_code == 404
    login(client, "TAGSMTEST01")
    incoming = ok(client.get("/api/announcements/handovers"))
    assert incoming[0]["version"] == 1 and incoming[0]["to_name"] == ta["name"]
    assert incoming[0]["scope_label"] and incoming[0]["requested_at"]
    ok(client.post(f"/api/announcements/handovers/{h['handover_id']}/accept"))
    assert ok(client.get(f"/api/announcements/{a['id']}"))["owner_id"] == ta["id"]
    login(client)
    assert client.post(f"/api/announcements/{a['id']}/withdraw", json={"reason": "过期", "expected_version": 1}).status_code == 403
    assert client.post(f"/api/announcements/{a['id']}/handover", json={"successor_id": am["id"], "reason": "兜底", "expected_version": 1, "force": True}).status_code == 403
    login(client, "HR01")
    ok(client.post(f"/api/announcements/{a['id']}/handover", json={"successor_id": am["id"], "reason": "原岗位未交接，管理员兜底", "expected_version": 1, "force": True}))
    assert ok(client.get(f"/api/announcements/{a['id']}"))["owner_id"] == am["id"]
    overview = ok(client.get(f"/api/announcements/{a['id']}/records"))["handover"]
    assert overview["owner"]["id"] == am["id"]
    assert [h["status"] for h in overview["history"]] == ["forced", "accepted"]
    assert all(h["completed_at"] for h in overview["history"])


def test_withdraw_notifies_every_historical_recipient_and_preserves_body(client):
    login(client)
    values = form(client)
    a = ok(client.post("/api/announcements", json=values))
    login(client, "CMTEST01")
    ok(receipt(client, a, action="confirm", acknowledged=True))
    login(client)
    revised = {**values, "tiers": ["gsm"], "request_key": uuid4().hex, "expected_version": 1, "change_summary": "仅适用管理层"}
    updated = ok(client.post(f"/api/announcements/{a['id']}/versions", json=revised))
    ok(client.post(f"/api/announcements/{a['id']}/withdraw", json={"reason": "要求撤销", "expected_version": 2}))
    login(client, "CMTEST01")
    history = ok(client.get(f"/api/announcements/{a['id']}?version=1"))
    assert history["body"] == values["body"] and history["receipt"]["confirmed_at"]
    assert any(n["kind"] == "withdrawn" for n in ok(client.get("/api/announcements/notifications"))["items"])
    assert ok(client.get("/api/announcements"))["total"] == 0
    login(client)
    assert client.post(f"/api/announcements/{a['id']}/versions", json={**revised, "expected_version": 2, "request_key": uuid4().hex}).status_code == 409


def test_record_scope_and_signature_privacy_are_enforced_server_side(client):
    cfg = config(client)
    own_circle = employee(cfg, "GSMTEST01")["attraction_id"]
    login(client, "TAGSMTEST01")
    a = publish(client, confirmation_level=3)
    login(client, "CMTEST01")
    signed = ok(receipt(client, a, action="confirm", acknowledged=True, signature=signature()))
    assert client.get(f"/api/announcements/{a['id']}/records").status_code == 403
    login(client, "CMTEST02")
    assert client.get(signed["signature_url"]).status_code == 404
    login(client)
    scoped = ok(client.get(f"/api/announcements/{a['id']}/records"))
    assert scoped["scoped"] and scoped["items"] and all(r["attraction_id"] == own_circle for r in scoped["items"])
    assert client.get(signed["signature_url"]).status_code == 200
    login(client, "SUPTEST01")
    lead = ok(client.get(f"/api/announcements/{a['id']}/records"))
    with SessionLocal() as db:
        expected = domain.direct_member_ids(db, employee(cfg, "SUPTEST01")["id"], include_overseen=True)
    assert all(r["employee_id"] in expected for r in lead["items"])


def test_attachments_private_until_published_and_protected_after_publication(client):
    login(client)
    png = base64.b64decode(signature().split(",")[1])
    item = ok(client.post("/api/announcement-media/images", data={"title": "配图", "source": "photo", "request_key": uuid4().hex},
                          files={"image": ("test.png", png, "image/png")}))["item"]
    login(client, "CMTEST01")
    assert client.get(f"/api/files/{item['file_id']}?preview=true").status_code == 404
    login(client)
    a = publish(client, attachments=[{"file_id": item["file_id"], "kind": "cover", "caption": "检查示意"}])
    login(client, "CMTEST01")
    assert client.get(f"/api/files/{item['file_id']}?preview=true").status_code == 200
    login(client, "AMTEST01")
    assert client.post("/api/announcements", json=form(client, attachments=[{"file_id": item["file_id"], "kind": "cover"}])).status_code == 400
    assert client.post("/api/announcements/assets", files={"file": ("fake.pdf", b"not pdf", "application/pdf")}).status_code == 400


def test_reports_preserve_snapshot_formula_safety_and_actual_signature_evidence(client):
    login(client)
    a = publish(client, confirmation_level=3, title="公告"*60)
    login(client, "CMTEST01")
    ok(receipt(client, a, action="confirm", acknowledged=True, signature=signature()))
    with SessionLocal() as db:
        row = db.query(AnnouncementDelivery).filter_by(version_id=a["version_id"]).first()
        old_name = row.employee_name
        db.get(Employee, row.employee_id).name = "新名字"
        row.employee_no = '=HYPERLINK("https://example.com")'
        db.commit()
    login(client)
    report = client.get(f"/api/announcements/{a['id']}/export?format=pdf")
    assert report.status_code == 200, report.text[:1000]
    with fitz.open(stream=report.content, filetype="pdf") as document:
        full = "".join(page.get_text() for page in document)
        assert "签字凭据" in full and "SHA256" in full and "V1" in full and old_name in full
        assert "公告正文" in full and "检查护栏" in full
        assert document.page_count >= 3
        # Retain a locally inspectable QA artifact outside production data.
        output = Path(__file__).resolve().parents[2]/"output"/"pdf"
        output.mkdir(parents=True, exist_ok=True)
        (output/"announcement-receipt-integration.pdf").write_bytes(report.content)
        for index, page in enumerate(document):
            page.get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(output/f"announcement-receipt-integration-{index+1}.png")
    excel = client.get(f"/api/announcements/{a['id']}/export?format=xlsx")
    assert excel.status_code == 200
    workbook = load_workbook(BytesIO(excel.content))
    assert workbook.active.title == "签收总览"
    assert workbook.sheetnames == ["签收总览","签收记录","跟进记录","公告内容"]
    sheet = workbook["签收记录"]
    dangerous = [c for row in sheet for c in row if isinstance(c.value,str) and c.value.startswith("=")]
    assert dangerous and all(c.data_type == "s" for c in dangerous)
    assert old_name in [c.value for row in sheet for c in row]
    assert sheet.freeze_panes == "E7"
    assert sheet.column_dimensions["J"].hidden
    assert "$A$1:$I$" + str(sheet.max_row) in str(sheet.print_area)
    (output/"announcement-receipt-integration.xlsx").write_bytes(excel.content)
    assert "公告内容" in workbook.sheetnames


def test_concurrent_repeated_publication_and_signature_have_one_evidence_row(client):
    login(client)
    values = form(client, confirmation_level=3)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: client.post("/api/announcements", json=values), range(2)))
    ids = [ok(r)["id"] for r in responses]
    assert ids[0] == ids[1]
    a = ok(client.get(f"/api/announcements/{ids[0]}"))
    login(client, "CMTEST01")
    with ThreadPoolExecutor(max_workers=2) as pool:
        signed = list(pool.map(lambda _: receipt(client, a, action="confirm", acknowledged=True, signature=signature()), range(2)))
    urls = [ok(r)["signature_url"] for r in signed]
    assert urls[0] == urls[1]
    with SessionLocal() as db:
        assert db.query(AnnouncementVersion).count() == 1
        assert db.query(AnnouncementDelivery).filter_by(signature_file_id=int(urls[0].split('/')[-1].split('?')[0])).count() == 1


def test_document_attachment_inherits_to_successor_without_becoming_public(client):
    cfg = config(client)
    ta = employee(cfg, "TAGSMTEST01")
    with SessionLocal() as db:
        circles = list(domain.live_circles(db))
    outside, outside_number = new_employee(next(c for c in circles if c != employee(cfg, "GSMTEST01")["attraction_id"]))
    login(client)
    doc = fitz.open()
    doc.new_page().insert_text((50, 50), "Announcement attachment")
    raw = doc.tobytes()
    doc.close()
    asset = ok(client.post("/api/announcements/assets", files={"file": ("requirements.pdf", raw, "application/pdf")}))
    values = form(client, attachments=[{"file_id": asset["file_id"], "kind": "attachment"}])
    a = ok(client.post("/api/announcements", json=values))
    h = ok(client.post(f"/api/announcements/{a['id']}/handover", json={"successor_id": ta["id"], "reason": "交接附件", "expected_version": 1}))
    login(client, "TAGSMTEST01")
    ok(client.post(f"/api/announcements/handovers/{h['handover_id']}/accept"))
    revised = {**values, "expected_version": 1, "request_key": uuid4().hex, "change_summary": "接任后修订", "body": "新要求"}
    ok(client.post(f"/api/announcements/{a['id']}/versions", json=revised))
    assert client.get(asset["url"]).status_code == 200
    login(client, outside_number)
    assert client.get(asset["url"]).status_code == 404


def test_office_upload_validates_real_document_and_rolls_back_invalid_archive(client):
    import zipfile
    login(client)
    workbook = Workbook()
    workbook.active.append(["要求", "执行"])
    output = BytesIO()
    workbook.save(output)
    asset = ok(client.post("/api/announcements/assets", files={"file": ("requirements.xlsx", output.getvalue(), "application/octet-stream")}))
    assert asset["file_id"]
    invalid = BytesIO()
    with zipfile.ZipFile(invalid, "w") as archive:
        archive.writestr("xl/workbook.xml", "not XML")
    with SessionLocal() as db:
        count = db.query(StoredFile).count()
    before = set(FILE_DIR.iterdir())
    response = client.post("/api/announcements/assets", files={"file": ("broken.xlsx", invalid.getvalue(), "application/octet-stream")})
    assert response.status_code == 400
    with SessionLocal() as db:
        assert db.query(StoredFile).count() == count
    assert set(FILE_DIR.iterdir()) == before


def test_same_day_supplement_gets_grace_and_reminders_are_deduplicated(client):
    login(client)
    values = form(client, effective_on=(date.today()-timedelta(days=1)).isoformat(), due_on=(date.today()-timedelta(days=1)).isoformat())
    a = ok(client.post("/api/announcements", json=values))
    id, number = new_employee(values["attraction_ids"][0])
    login(client, number)
    own = ok(client.get("/api/announcements"))["items"][0]
    assert own["receipt"]["due_on"] == (date.today()+timedelta(days=3)).isoformat()
    login(client)
    request = {"expected_version": 1, "reason": "提醒查收"}
    ok(client.post(f"/api/announcements/{a['id']}/remind", json=request))
    ok(client.post(f"/api/announcements/{a['id']}/remind", json=request))
    with SessionLocal() as db:
        assert db.query(AnnouncementNotification).filter_by(employee_id=id, kind="reminder").count() == 1


def test_return_to_scope_gets_new_deadline_without_extending_validity(client):
    cfg = config(client)
    cm = employee(cfg, "CMTEST01")
    login(client)
    yesterday = (date.today()-timedelta(days=1)).isoformat()
    a = publish(client, effective_on=yesterday, due_on=yesterday, expires_on=date.today().isoformat())
    with SessionLocal() as db:
        db.get(Employee, cm["id"]).attraction_id = next(c for c in domain.live_circles(db) if c != cm["attraction_id"])
        db.commit()
    login(client, "CMTEST01")
    assert ok(client.get("/api/announcements"))["total"] == 0
    with SessionLocal() as db:
        db.get(Employee, cm["id"]).attraction_id = cm["attraction_id"]
        db.commit()
    returned = ok(client.get("/api/announcements"))["items"][0]["receipt"]
    assert returned["reason"] == "重新适用补收"
    assert returned["due_on"] == date.today().isoformat()


def test_restricted_category_and_tier_are_enforced_for_authorized_publisher(client):
    cfg = config(client)
    gsm = employee(cfg, "GSMTEST01")
    ok(client.post("/api/announcements/admin/grants", json={"employee_id": gsm["id"], "attraction_ids": [gsm["attraction_id"]],
        "tiers": ["lead"], "categories": ["演出 & 5S"], "starts_on": date.today().isoformat(), "reason": "限定范围"}))
    login(client)
    assert client.post("/api/announcements", json=form(client)).status_code == 403
    publish(client, tiers=["lead"], category="演出 & 5S")


def test_missing_or_changed_signature_blocks_report_instead_of_omitting_evidence(client):
    login(client)
    a = publish(client, confirmation_level=3)
    login(client, "CMTEST01")
    signed = ok(receipt(client, a, action="confirm", acknowledged=True, signature=signature()))
    file_id = int(signed["signature_url"].split('/')[-1].split('?')[0])
    with SessionLocal() as db:
        path = FILE_DIR / db.get(StoredFile, file_id).storage_key
    path.write_bytes(base64.b64decode(signature(True).split(',')[1]))
    login(client)
    result = client.get(f"/api/announcements/{a['id']}/export?format=pdf")
    assert result.status_code == 409 and "摘要不匹配" in result.json()["detail"]
    path.unlink()
    result = client.get(f"/api/announcements/{a['id']}/export?format=pdf")
    assert result.status_code == 409 and "文件缺失" in result.json()["detail"]



def test_group_distribution_is_ordered_scoped_and_uses_current_valid_memberships(client):
    login(client, "TAGSMTEST01")
    circles = ok(client.get("/api/announcements/options"))["publisher"]["circles"]
    with SessionLocal() as db:
        own_circle = db.query(Employee).filter_by(employee_no="GSMTEST01").one().attraction_id
        other_circle = next(id for id in circles if id != own_circle)
        group_a = db.query(WorkGroup).filter_by(attraction_id=own_circle, code="A").one()
        group_b = db.query(WorkGroup).filter_by(attraction_id=own_circle, code="B").one()
        other_group = db.query(WorkGroup).filter_by(attraction_id=other_circle).first()
        if not other_group:
            other_group = WorkGroup(name="其他圈测试小组", code="Z", attraction_id=other_circle)
            db.add(other_group)
            db.flush()
        group_ids = group_a.id, group_b.id, other_group.id
        db.commit()
    person_a, account_a = new_employee(own_circle)
    person_b, _ = new_employee(own_circle)
    other_person, _ = new_employee(other_circle)
    with SessionLocal() as db:
        for person, group in zip([person_a, person_b, other_person], group_ids):
            db.add(GroupMembership(employee_id=person, group_id=group, starts_on=date.today().isoformat()))
        db.add(GroupMembership(employee_id=person_a, group_id=group_ids[1],
                               starts_on=(date.today()+timedelta(days=1)).isoformat()))
        db.commit()
    a = publish(client)
    login(client, account_a)
    ok(receipt(client, a, action="confirm", acknowledged=True))
    assert client.get(f"/api/announcements/{a['id']}/records").status_code == 403
    login(client)
    data = ok(client.get(f"/api/announcements/{a['id']}/records"))
    assert data["scoped"] and data["group_basis"] == "当前小组"
    assert other_person not in {r["employee_id"] for r in data["items"]}
    assert all(g["attraction_id"] == own_circle for g in data["groups"])
    own_groups = [g["id"] for g in data["groups"] if g["id"]]
    assert own_groups.index(group_ids[0]) < own_groups.index(group_ids[1])
    signed = next(r for r in data["items"] if r["employee_id"] == person_a)
    assert signed["group_id"] == group_ids[0] and signed["confirmed_at"]
    assert sum(g["required"] for g in data["groups"]) == data["counts"]["required"]
    assert sum(g["completed"] for g in data["groups"]) == data["counts"]["completed"]
    with SessionLocal() as db:
        db.get(Employee, person_b).attraction_id = other_circle
        db.query(GroupMembership).filter_by(employee_id=person_b).update({"status": "ended"})
        db.add(GroupMembership(employee_id=person_b, group_id=group_ids[2], starts_on=date.today().isoformat()))
        db.commit()
    data = ok(client.get(f"/api/announcements/{a['id']}/records"))
    transferred = next(r for r in data["items"] if r["employee_id"] == person_b)
    # This announcement covers both circles, so the task remains applicable.
    # Its original circle snapshot must not disclose the later group's identity.
    assert transferred["group_id"] is None and transferred["status"] == "pending"
    assert transferred["attraction_id"] == own_circle
    assert group_ids[2] not in {g["id"] for g in data["groups"]}


def test_report_layout_keeps_scoped_counts_and_full_content(client):
    from types import SimpleNamespace
    from app.announcement_exports import spreadsheet, pdf, report_groups
    rows = []
    for i in range(24):
        status = ["completed", "pending", "suspended", "exited"][i % 4]
        rows.append({"attraction_name": "模拟热力追踪" if i < 12 else "模拟矮人迷宫",
            "group_name": f"模拟小组 {i//6+1}", "employee_no": f"009{i:04d}", "employee_name": f"模拟人员{i+1}",
            "role_label": "CM / TR", "status": status, "reason": "首次接收", "delivered_at": "2026-10-13T08:00:00",
            "seen_at": "2026-10-13T08:30:00" if status == "completed" else None,
            "confirmed_at": "2026-10-13T09:00:00" if status == "completed" else None,
            "due_on": "2026-10-16", "signing_statement": None, "signature_url": None,
            "trainer_name": "模拟培训师" if status == "completed" else None, "trainer_no": "0001001" if status == "completed" else None,
            "training_session_id": 9 if status == "completed" else None,
            "training_verified_at": "2026-10-13T08:55:00" if status == "completed" else None})
    announcement = {"id": 99, "version": 2, "title": "模拟数据 · 岗位培训签收", "category": "培训", "effective": True,
        "scope_label": "模拟热力追踪、模拟矮人迷宫", "effective_on": "2026-10-13", "due_on": "2026-10-16", "expires_on": None,
        "confirmation_level": 3, "training_required": True, "summary": "先完成现场培训，再验码并本人签字。",
        "change_summary": "新增岗位操作要点及培训跟进要求。",
        "body": "第一部分：培训要求\n完成现场讲解、操作示范与本人确认。\n\n第二部分：操作要点\n核对岗位安排，发现异常立即报告。"}
    data = {"announcement": announcement, "items": rows, "counts": {"required": 12, "completed": 6, "pending": 6, "suspended": 6, "exited": 6}, "scoped": True}
    user = SimpleNamespace(name="模拟导出人", employee=SimpleNamespace(employee_no="0000001"))
    raw = spreadsheet(data, user)
    book = load_workbook(BytesIO(raw))
    assert book.active.title == "签收总览"
    groups = report_groups(data)
    assert len(groups) == 4
    totals = [sum(book.active.cell(r,c).value for r in range(7,11)) for c in range(3,8)]
    assert totals == [12,6,6,6,6]
    sheet = book["签收记录"]
    assert sheet.max_row == 30 and sheet["C7"].value == "0090000" and sheet["C7"].data_type == "s"
    assert isinstance(sheet["F7"].value, datetime) and isinstance(sheet["H7"].value, datetime)
    assert sheet["F7"].number_format == "yyyy-mm-dd hh:mm"
    assert sheet.column_dimensions["J"].hidden and book["跟进记录"].max_column == 11
    assert [c.value for c in book["跟进记录"][6]] != [c.value for c in sheet[6]]
    body_values = [row[1].value or "" for row in book["公告内容"] if row[0].value in {"正文","续"}]
    assert "\n".join(body_values) == announcement["body"]
    content = pdf(data, user)
    target = Path(__file__).resolve().parents[2]/".handover_review"/"announcement-export-20261013"
    target.mkdir(parents=True, exist_ok=True)
    (target/"receipt-demo.xlsx").write_bytes(raw)
    (target/"receipt-demo.pdf").write_bytes(content)
    with fitz.open(stream=content,filetype="pdf") as document:
        full = "".join(page.get_text() for page in document)
        assert "签收总览" in full and "当前管理范围" in full and "操作示范" in full
        assert "新增岗位操作要点" in full
        for i,page in enumerate(document):
            page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(target/f"receipt-demo-{i+1}.png")
    # Empty reports remain printable. Many short paragraphs must paginate without loss.
    empty = {**data,"items": [], "counts": dict.fromkeys(data["counts"],0)}
    with fitz.open(stream=pdf(empty,user),filetype="pdf") as document:
        assert "尚无接收记录" in "".join(p.get_text() for p in document)
    long_data = {**empty,"announcement": {**announcement,"body": "\n".join(f"段落{i:04d}：逐项核对培训要求。" for i in range(120))}}
    long_book = load_workbook(BytesIO(spreadsheet({**long_data,"announcement": {**long_data["announcement"],"change_summary": "长变更"*300}},user)))
    content_sheet = long_book["公告内容"]
    assert all(d.height <= 409 for d in content_sheet.row_dimensions.values())
    assert "".join(row[1].value for row in content_sheet if str(row[0].value).startswith("变更")) == "长变更"*300
    with fitz.open(stream=pdf(long_data,user),filetype="pdf") as document:
        full = "".join(p.get_text() for p in document)
        assert all(f"段落{i:04d}" in full for i in range(120))


def test_handover_center_and_history_are_private_to_publishers_owners_and_admins(client):
    cfg = config(client)
    ta = employee(cfg, "TAGSMTEST01")
    login(client)
    mine = publish(client, category="培训")
    login(client, "AMTEST01")
    other = publish(client, category="礼仪 & MM")
    assert {a["id"] for a in ok(client.get("/api/announcements/handover-center"))["items"]} == {other["id"]}
    records = ok(client.get(f"/api/announcements/{mine['id']}/records"))
    assert records["handover"] is None and records["announcement"]["can_view_handover"] is False
    assert client.get(f"/api/announcements/successors/{mine['id']}").status_code == 403
    assert client.post(f"/api/announcements/{mine['id']}/handover", json={"successor_id": ta["id"],"reason":"越权","expected_version":1}).status_code == 403
    login(client)
    center = ok(client.get("/api/announcements/handover-center?all=true"))
    assert not center["all_visible"] and {a["id"] for a in center["items"]} == {mine["id"]}
    assert ok(client.get(f"/api/announcements/{other['id']}/records"))["handover"] is None
    request = ok(client.post(f"/api/announcements/{mine['id']}/handover", json={"successor_id":ta["id"],"reason":"本人交接","expected_version":1}))
    assert ok(client.get("/api/announcements/handover-center"))["items"][0]["pending"] == 1
    login(client,"TAGSMTEST01")
    assert ok(client.get("/api/announcements/handover-center"))["items"] == []
    assert ok(client.get(f"/api/announcements/{mine['id']}/records"))["handover"] is None
    assert ok(client.get("/api/announcements/handovers"))[0]["id"] == request["handover_id"]
    ok(client.post(f"/api/announcements/handovers/{request['handover_id']}/accept"))
    center = ok(client.get("/api/announcements/handover-center"))
    assert {a["id"] for a in center["items"]} == {mine["id"]}
    assert center["items"][0]["owner"]["id"] == ta["id"] and center["items"][0]["can_handover"]
    assert ok(client.get(f"/api/announcements/{mine['id']}/records"))["handover"]["owner"]["id"] == ta["id"]
    login(client)
    center = ok(client.get("/api/announcements/handover-center"))
    assert center["items"][0]["owner"]["id"] == ta["id"] and not center["items"][0]["can_handover"]
    assert ok(client.get(f"/api/announcements/{mine['id']}/records"))["handover"]["history"][0]["status"] == "accepted"
    # Two independent administrator accounts use the same permission boundary.
    admin_id, admin_no = new_employee(ta["attraction_id"])
    with SessionLocal() as db:
        assignment = db.query(EmployeeRoleAssignment).filter_by(employee_id=admin_id).one()
        assignment.role_id = db.query(Role).filter_by(code="SYSTEM_ADMIN").one().id
        db.commit()
    for account in ["HR01",admin_no]:
        login(client,account)
        center = ok(client.get("/api/announcements/handover-center"))
        assert center["all_visible"] and {a["id"] for a in center["items"]} == {mine["id"],other["id"]}
        item = next(a for a in center["items"] if a["id"] == mine["id"])
        assert item["category"] == "培训" and item["owner"]["id"] == ta["id"]
        assert ok(client.get(f"/api/announcements/{mine['id']}/records"))["handover"]["owner"]["id"] == ta["id"]


def shared_setup(client):
    cfg = config(client)
    gsm, ta = employee(cfg, "GSMTEST01"), employee(cfg, "TAGSMTEST01")
    partner_id, partner_no = new_employee(gsm["attraction_id"])
    with SessionLocal() as db:
        assignment = db.query(EmployeeRoleAssignment).filter_by(employee_id=partner_id).one()
        assignment.role_id = db.query(Role).filter_by(code="GSM").one().id
        db.commit()
    login(client)
    a, b = publish(client), publish(client, title="同板块第二份公告")
    login(client, "HR01")
    group = next(g for g in ok(client.get("/api/announcements/handover-center"))["groups"]
        if a["id"] in {x["id"] for x in g["announcements"]})
    ok(client.post("/api/announcements/responsibilities", json={"key":group["key"],"expected_revision":0,
        "member_ids":[gsm["id"],partner_id],"reason":"两位景点圈负责人共同管理"}))
    return gsm, ta, partner_id, partner_no, a, b, group["key"]


def batch_input(client, key, from_id, to_id, **changes):
    group = ok(client.get("/api/announcements/responsibility", params={"key":key}))
    return {**{"key":key,"expected_revision":group["revision"],"from_id":from_id,"successor_id":to_id,
        "announcements":sorted([{"id":a["id"],"version":a["version"]} for a in group["announcements"]],key=lambda a:a["id"]),
        "reason":"整组岗位交接","force":False}, **changes}


def test_shared_managers_update_export_withdraw_train_and_transfer_entire_board(client):
    gsm, ta, partner_id, partner_no, a, b, key = shared_setup(client)
    login(client, partner_no)
    assert ok(client.get(f"/api/announcements/{a['id']}"))["can_edit"]
    assert {p["id"] for p in ok(client.get(f"/api/announcements/{a['id']}/records"))["handover"]["owners"]} == {gsm["id"],partner_id}
    assert client.get(f"/api/announcements/{a['id']}/followup-export").status_code == 200
    assert client.get(f"/api/announcements/{a['id']}/export?format=pdf").status_code == 200
    ok(client.post(f"/api/announcements/{b['id']}/withdraw",json={"expected_version":1,"reason":"共同负责人撤下"}))
    # A later publication in the same unit inherits both explicitly assigned managers.
    tr = next(e for e in ok(client.get("/api/announcements/training-options"))["items"] if e["employee_no"]=="TRTEST01")
    c = publish(client,title="同板块培训公告",confirmation_level=3,training_required=True,trainer_ids=[tr["id"]])
    login(client)
    ok(client.post(f"/api/announcements/{c['id']}/training/trainers",json={"expected_version":1,"trainer_scope":"circle","trainer_ids":[tr["id"]]}))
    login(client,"TRTEST01")
    code = ok(client.post(f"/api/announcements/{c['id']}/training/codes",json={"version":1,"hours":2,"request_key":uuid4().hex}))
    login(client,"CMTEST01")
    ok(receipt(client,a,action="confirm",acknowledged=True))
    with SessionLocal() as db:
        before = [(d.id,d.confirmed_at,d.version_id) for d in db.query(AnnouncementDelivery).filter_by(version_id=a["version_id"])]
    login(client)
    requested = ok(client.post("/api/announcements/responsibility-handover",json=batch_input(client,key,gsm["id"],ta["id"])))
    assert requested["count"]==3
    login(client,partner_no)
    assert client.post(f"/api/announcements/handovers/{requested['handover_id']}/accept").status_code==404
    assert client.post("/api/announcements/responsibility-handover",json=batch_input(client,key,partner_id,ta["id"])).status_code==409
    login(client,"TAGSMTEST01")
    incoming = ok(client.get("/api/announcements/handovers"))
    assert len(incoming)==1 and incoming[0]["count"]==3
    assert ok(client.get("/api/announcements/summary"))["handovers"]==1
    assert ok(client.post(f"/api/announcements/handovers/{requested['handover_id']}/accept"))["count"]==3
    group = ok(client.get("/api/announcements/responsibility",params={"key":key}))
    assert {p["id"] for p in group["owners"]}=={partner_id,ta["id"]}
    for item in [a,b,c]:
        assert ok(client.get(f"/api/announcements/{item['id']}"))["can_edit"]
    with SessionLocal() as db:
        assert before==[(d.id,d.confirmed_at,d.version_id) for d in db.query(AnnouncementDelivery).filter_by(version_id=a["version_id"])]
        assert db.query(AnnouncementTrainingCode).filter_by(version_id=c["version_id"]).one().ended_at is None
        assert db.query(AnnouncementHandover).filter_by(status="accepted").count()==3
    login(client)
    assert not ok(client.get(f"/api/announcements/{a['id']}"))["can_edit"]
    assert client.post(f"/api/announcements/{a['id']}/withdraw",json={"expected_version":1,"reason":"原发布者"}).status_code==403
    assert client.get(f"/api/announcements/{a['id']}/followup-export").status_code==200
    login(client,partner_no)
    revised=ok(client.post(f"/api/announcements/{a['id']}/versions",json=form(client,expected_version=1,change_summary="共同负责人更新")))
    assert revised["version"]==2


def test_shared_scope_is_explicit_and_batch_acceptance_rejects_stale_versions(client):
    gsm,ta,partner_id,partner_no,a,b,key=shared_setup(client)
    login(client,"AMTEST01")
    other_circle=next(c["id"] for c in ok(client.get("/api/announcements/options"))["circles"] if c["id"]!=gsm["attraction_id"])
    elsewhere=publish(client,attraction_ids=[other_circle])
    separate=publish(client,category="培训",attraction_ids=[gsm["attraction_id"]])
    login(client,partner_no)
    assert client.get(f"/api/announcements/{elsewhere['id']}").status_code == 404
    assert ok(client.get(f"/api/announcements/{separate['id']}/records"))["handover"] is None
    assert client.post("/api/announcements/responsibilities",json={"key":key,"expected_revision":1,"member_ids":[partner_id],"reason":"越权设置"}).status_code==403
    assert client.post("/api/announcements/responsibility-handover",json=batch_input(client,key,gsm["id"],ta["id"])).status_code==403
    login(client)
    requested=ok(client.post("/api/announcements/responsibility-handover",json=batch_input(client,key,gsm["id"],ta["id"])))
    ok(client.post(f"/api/announcements/{a['id']}/versions",json=form(client,expected_version=1,change_summary="交接期间更新")))
    login(client,"TAGSMTEST01")
    assert client.post(f"/api/announcements/handovers/{requested['handover_id']}/accept").status_code==409
    with SessionLocal() as db:
        assert {gsm["id"],partner_id}==set(domain.unpack(db.query(AnnouncementResponsibility).filter_by(scope_key=key).one().member_ids_json))
        assert db.query(AnnouncementHandover).filter_by(status="accepted").count()==0
        assert db.get(Announcement,b["id"]).owner_id==gsm["id"]


def test_admin_shared_roster_validation_and_force_transfer_preserve_other_manager(client):
    gsm,ta,partner_id,partner_no,a,b,key=shared_setup(client)
    login(client,"HR01")
    base={"key":key,"expected_revision":1,"member_ids":[gsm["id"],partner_id],"reason":"调整"}
    assert client.post("/api/announcements/responsibilities",json={**base,"member_ids":[gsm["id"],partner_id,ta["id"]]}).status_code==422
    cm=employee(ok(client.get("/api/announcements/admin/config")),"CMTEST01")
    assert client.post("/api/announcements/responsibilities",json={**base,"member_ids":[gsm["id"],cm["id"]]}).status_code==403
    assert client.post("/api/announcements/responsibilities",json={**base,"expected_revision":0}).status_code==409
    forced=ok(client.post("/api/announcements/responsibility-handover",json=batch_input(client,key,gsm["id"],ta["id"],force=True)))
    assert forced["count"]==2
    group=ok(client.get("/api/announcements/responsibility",params={"key":key}))
    assert {p["id"] for p in group["owners"]}=={partner_id,ta["id"]}
    assert ok(client.get(f"/api/announcements/{b['id']}/records"))["handover"]["history"][0]["status"]=="forced"


def test_batch_handover_rechecks_successor_grants_and_rejects_single_notice_route(client):
    gsm,ta,partner_id,partner_no,a,b,key=shared_setup(client)
    login(client)
    assert client.post(f"/api/announcements/{a['id']}/handover",json={"expected_version":1,"successor_id":ta["id"],"reason":"未核对整组"}).status_code==409
    requested=ok(client.post("/api/announcements/responsibility-handover",json=batch_input(client,key,gsm["id"],ta["id"])))
    login(client,"HR01")
    ok(client.post("/api/announcements/admin/grants",json={"employee_id":ta["id"],"attraction_ids":[gsm["attraction_id"]],
        "enabled":False,"starts_on":date.today().isoformat(),"reason":"接受前停用发布授权"}))
    login(client,"TAGSMTEST01")
    assert client.post(f"/api/announcements/handovers/{requested['handover_id']}/accept").status_code==403
    with SessionLocal() as db:
        assert db.get(Announcement,a["id"]).owner_id==gsm["id"]
        assert db.get(Announcement,b["id"]).owner_id==gsm["id"]
        assert db.query(AnnouncementHandover).filter_by(status="accepted").count()==0


def test_project_coowners_share_only_explicit_project_and_category(client):
    cfg=config(client)
    gsm,am=employee(cfg,"GSMTEST01"),employee(cfg,"AMTEST01")
    project=ok(client.post("/api/announcements/admin/projects",json={"name":"礼仪项目","employee_ids":[gsm["id"]]}))
    for person in [gsm,am]:
        ok(client.post("/api/announcements/admin/grants",json={"employee_id":person["id"],"scope_kind":"project","project_id":project["id"],
            "starts_on":date.today().isoformat(),"reason":"项目POC发布"}))
    login(client)
    a=publish(client,scope_kind="project",project_id=project["id"],attraction_ids=[],category="礼仪 & MM")
    b=publish(client,category="礼仪 & MM")
    c=publish(client,scope_kind="project",project_id=project["id"],attraction_ids=[],category="培训")
    login(client,"HR01")
    group=next(g for g in ok(client.get("/api/announcements/handover-center"))["groups"] if a["id"] in {x["id"] for x in g["announcements"]})
    ok(client.post("/api/announcements/responsibilities",json={"key":group["key"],"expected_revision":0,"member_ids":[gsm["id"],am["id"]],"reason":"两位项目POC"}))
    login(client,"AMTEST01")
    assert ok(client.get(f"/api/announcements/{a['id']}"))["can_edit"]
    assert not ok(client.get(f"/api/announcements/{b['id']}"))["can_edit"]
    assert not ok(client.get(f"/api/announcements/{c['id']}"))["can_edit"]
    records=ok(client.get(f"/api/announcements/{a['id']}/records"))
    assert not records["scoped"] and {p["id"] for p in records["handover"]["owners"]}=={gsm["id"],am["id"]}


def test_unconfigured_board_with_separate_publishers_has_no_personal_transfer_entry(client):
    cfg = config(client)
    gsm = employee(cfg, "GSMTEST01")
    login(client)
    mine = publish(client)
    login(client, "AMTEST01")
    other = publish(client, attraction_ids=[gsm["attraction_id"]])
    login(client)
    group = next(g for g in ok(client.get("/api/announcements/handover-center"))["groups"]
        if mine["id"] in {a["id"] for a in g["announcements"]})
    assert not group["can_handover"]
    assert not ok(client.get("/api/announcements/handover-center"))["items"][0]["can_handover"]
    assert {a["id"] for a in group["announcements"]} == {mine["id"]}
    assert not ok(client.get(f"/api/announcements/{mine['id']}/records"))["handover"]["can_handover"]
    assert client.get("/api/announcements/responsibility", params={"key": group["key"]}).status_code == 403
    login(client, "HR01")
    group = ok(client.get("/api/announcements/responsibility", params={"key": group["key"]}))
    assert group["can_handover"] and {a["id"] for a in group["announcements"]} == {mine["id"], other["id"]}


def test_managed_board_scope_change_is_rejected_without_mutating_receipts_or_training(client):
    gsm, ta, partner_id, partner_no, a, b, key = shared_setup(client)
    login(client, partner_no)
    trainer = next(e for e in ok(client.get("/api/announcements/training-options"))["items"] if e["employee_no"] == "TRTEST01")
    notice = publish(client, confirmation_level=3, training_required=True, trainer_ids=[trainer["id"]])
    assert notice["responsibility_scope_locked"]
    login(client, "TRTEST01")
    ok(client.post(f"/api/announcements/{notice['id']}/training/codes",
        json={"version":1,"hours":2,"request_key":uuid4().hex}))
    with SessionLocal() as db:
        before = [(d.id,d.status,d.version_id) for d in db.query(AnnouncementDelivery).filter_by(version_id=notice["version_id"])]
    login(client, partner_no)
    attempted = client.post(f"/api/announcements/{notice['id']}/versions",
        json=form(client,category="培训",expected_version=1,change_summary="尝试移出共同管理"))
    assert attempted.status_code == 409
    with SessionLocal() as db:
        assert db.get(Announcement,notice["id"]).current_version == 1
        assert before == [(d.id,d.status,d.version_id) for d in db.query(AnnouncementDelivery).filter_by(version_id=notice["version_id"])]
        assert db.query(AnnouncementTrainingCode).filter_by(version_id=notice["version_id"]).one().ended_at is None
        assert db.query(AnnouncementVersion).filter_by(announcement_id=notice["id"]).count() == 1
    requested = ok(client.post("/api/announcements/responsibility-handover",
        json=batch_input(client,key,partner_id,ta["id"])))
    login(client, "TAGSMTEST01")
    assert ok(client.post(f"/api/announcements/handovers/{requested['handover_id']}/accept"))["count"] == 3

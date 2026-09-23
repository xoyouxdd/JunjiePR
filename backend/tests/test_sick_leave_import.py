from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import tempfile

TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-sick-import-test-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from fastapi import HTTPException  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
import pytest  # noqa: E402
import xlrd  # noqa: E402

from app.main import app  # noqa: E402
from app.routers import sick_leave_import  # noqa: E402
from app.routers.sick_leave_import import _read_workbook, _reconcile, _system_number  # noqa: E402
from app.v2_database import ROLE_PERMISSION_CODES, SessionLocal  # noqa: E402
from app.v2_models import SickLeaveRecord  # noqa: E402


def test_legacy_manual_sick_registration_stays_disabled_for_production_roles():
    assert all("SICK_REGISTER" not in permissions for permissions in ROLE_PERMISSION_CODES.values())


def test_source_id_drops_exactly_one_leading_digit():
    assert _system_number("01727264") == "1727264"
    assert _system_number("01278490") == "1278490"


def test_reconcile_combines_paid_sick_leave_types():
    records = [
        {"source_id": "01727264", "leave_type": "法定病假", "hours": Decimal("8")},
        {"source_id": "01727264", "leave_type": "全薪病假", "hours": Decimal("4")},
        {"source_id": "01727264", "leave_type": "无薪病假", "hours": Decimal("8")},
    ]
    summaries = {
        ("01727264", "病假时间总计"): Decimal("12"),
        ("01727264", "无薪病假"): Decimal("8"),
    }

    assert _reconcile(records, summaries) == []


def test_reconcile_blocks_a_summary_difference():
    records = [{"source_id": "01727264", "leave_type": "法定病假", "hours": Decimal("8")}]

    errors = _reconcile(records, {("01727264", "病假时间总计"): Decimal("4")})

    assert len(errors) == 1
    assert "总数区对账不一致" in errors[0]


def test_commit_rechecks_loa_added_after_preview():
    with TestClient(app) as client:
        login = client.post("/api/login", json={"employee_no": "GSMTEST01", "password": "1234"})
        assert login.status_code == 200, login.text
        target = next(
            row for row in client.get("/api/employee-targets", params={"usage": "loa", "keyword": "测试CM甲"}).json()["items"]
            if row["employee_no"] == "CMTEST01"
        )
        user_id = client.get("/api/me").json()["id"]
        token = "stale-loa-review"
        sick_leave_import._store_preview(token, {
            "created_at": datetime.now(), "filename": "review.xls", "content": b"",
            "month": "2099-08", "matched": [{
                "employee_id": target["id"], "date": "2099-08-20", "days": Decimal("1"), "leave_type": "法定病假",
            }],
            "unmatched": [], "loa_protected": [], "errors": [], "user_id": user_id,
        })
        created = client.post("/api/loa-periods", json={
            "employee_id": target["id"], "starts_on": "2099-08-20", "ends_on": "2099-08-20",
        })
        assert created.status_code == 200, created.text
        committed = client.post("/api/sick-leave-imports/commit", json={"token": token})
        assert committed.status_code == 409, committed.text
        assert "重新上传文件预检" in committed.json()["detail"]
        with SessionLocal() as db:
            assert db.query(SickLeaveRecord).filter(
                SickLeaveRecord.employee_id == target["id"],
                SickLeaveRecord.leave_start_date == "2099-08-20",
                SickLeaveRecord.status == "active",
                SickLeaveRecord.import_source == "monthly_transaction_import",
            ).count() == 0
        sick_leave_import._drop_preview(token)


def test_preview_rejects_workbooks_above_size_limit():
    with TestClient(app) as client:
        login = client.post("/api/login", json={"employee_no": "GSMTEST01", "password": "1234"})
        assert login.status_code == 200, login.text
        response = client.post("/api/sick-leave-imports/preview", files={
            "workbook": ("too-large.xls", b"x" * (sick_leave_import._MAX_WORKBOOK_BYTES + 1)),
        })
        assert response.status_code == 413, response.text


def test_preview_cache_prunes_expired_and_oldest_entries(monkeypatch):
    monkeypatch.setattr(sick_leave_import, "_MAX_CACHED_PREVIEWS", 2)
    monkeypatch.setattr(sick_leave_import, "_MAX_CACHED_BYTES", 8)
    sick_leave_import._PREVIEWS.clear()
    now = datetime.now()
    for token, age in (("expired", timedelta(minutes=21)), ("first", timedelta(seconds=3)),
                       ("second", timedelta(seconds=2)), ("third", timedelta(seconds=1))):
        sick_leave_import._store_preview(token, {
            "created_at": now - age, "content": b"1234", "user_id": 1,
        })
    assert set(sick_leave_import._PREVIEWS) == {"second", "third"}
    assert sick_leave_import._get_preview("expired", 1) is None
    assert sick_leave_import._get_preview("second", 2) is None
    sick_leave_import._PREVIEWS.clear()


class _FakeSheet:
    """Minimal stand-in for an xlrd sheet; rows are padded to a common width."""

    def __init__(self, rows: list[list], name: str = "员工事务"):
        self.name = name
        self.ncols = max(len(row) for row in rows)
        self.nrows = len(rows)
        self._rows = [list(row) + [""] * (self.ncols - len(row)) for row in rows]

    def row_values(self, index: int) -> list:
        return list(self._rows[index])


class _FakeBook:
    datemode = 0

    def __init__(self, sheet: _FakeSheet):
        self._sheet = sheet

    def sheet_by_index(self, index: int) -> _FakeSheet:
        assert index == 0
        return self._sheet

    def sheets(self) -> list[_FakeSheet]:
        return [self._sheet]


def _serial(year: int, month: int, day: int) -> float:
    return xlrd.xldate.xldate_from_date_tuple((year, month, day), 0)


def _workbook_rows(transactions: list[list], summaries: list[list], *, colon: str = ":") -> list[list]:
    """Build the HR layout: header two rows below "事务:", three rows below "总数:"."""
    return [
        ["员工事务报表"],
        [f"事务{colon}"],
        [],
        ["员工", "ID", "日期", "工资代码", "时数"],
        *transactions,
        [f"总数{colon}"],
        [],
        [],
        ["员工", "ID", "工资代码", "时数"],
        *summaries,
    ]


def _install_book(monkeypatch, rows: list[list]) -> None:
    book = _FakeBook(_FakeSheet(rows))
    monkeypatch.setattr(sick_leave_import.xlrd, "open_workbook", lambda *args, **kwargs: book)


def test_read_workbook_normalizes_text_and_numeric_ids(monkeypatch):
    _install_book(monkeypatch, _workbook_rows(
        [
            ["ZHANG, 张三", "01727264", _serial(2099, 8, 20), "法定病假", 8.0],
            ["LI, 李四", 1278490.0, "2099/08/21", "全薪病假", "4"],
            ["LI, 李四", 1278490.0, _serial(2099, 8, 22), "无薪病假", 8.0],
            ["LI, 李四", 1278490.0, _serial(2099, 8, 23), "年假", 8.0],
            ["WANG, 王五", "1727265", _serial(2099, 8, 24), "法定病假", 8.0],
        ],
        [
            ["ZHANG, 张三", "01727264", "病假时间总计", 8.0],
            ["WANG, 王五", "01727265", "病假时间总计", 8.0],
            ["LI, 李四", 1278490.0, "病假时间总计", 4.0],
            ["LI, 李四", 1278490.0, "无薪病假", 8.0],
        ],
    ))

    records, summaries, errors = _read_workbook(b"ignored")

    assert errors == []
    assert [(row["source_id"], row["employee_no"], row["name"], row["date"], row["hours"]) for row in records] == [
        ("01727264", "1727264", "张三", "2099-08-20", Decimal("8.00")),
        ("01278490", "1278490", "李四", "2099-08-21", Decimal("4.00")),
        ("01278490", "1278490", "李四", "2099-08-22", Decimal("8.00")),
        ("01727265", "1727265", "王五", "2099-08-24", Decimal("8.00")),
    ]
    assert records[0]["days"] == Decimal("1")
    assert summaries == {
        ("01727264", "病假时间总计"): Decimal("8.00"),
        ("01727265", "病假时间总计"): Decimal("8.00"),
        ("01278490", "病假时间总计"): Decimal("4.00"),
        ("01278490", "无薪病假"): Decimal("8.00"),
    }
    assert _reconcile(records, summaries) == []


def test_read_workbook_reports_row_errors_and_reconcile_mismatch(monkeypatch):
    _install_book(monkeypatch, _workbook_rows(
        [
            ["ZHANG, 张三", "01727264", _serial(2099, 8, 20), "法定病假", 8.0],
            ["ZHANG, 张三", "01727264", _serial(2099, 8, 21), "法定病假", 6.0],
            ["WANG, 王五", "117272650", _serial(2099, 8, 20), "法定病假", 8.0],
            ["ZHAO, 赵六", 1727266.5, _serial(2099, 8, 20), "法定病假", 8.0],
            ["QIAN, 钱七", "01727267", "not-a-date", "法定病假", 8.0],
        ],
        [["ZHANG, 张三", "01727264", "病假时间总计", 16.0]],
    ))

    records, summaries, errors = _read_workbook(b"ignored")

    assert [row["row"] for row in records] == [5]
    assert len(errors) == 4
    assert errors[0].startswith("事务区第 6 行") and "4 小时倍数" in errors[0]
    assert errors[1].startswith("事务区第 7 行") and "8 位数字" in errors[1]
    assert errors[2].startswith("事务区第 8 行") and "8 位数字" in errors[2]
    assert errors[3].startswith("事务区第 9 行") and "事务日期无效" in errors[3]
    mismatch = _reconcile(records, summaries)
    assert len(mismatch) == 1
    assert "总数区对账不一致" in mismatch[0] and "01727264" in mismatch[0]


def test_read_workbook_accepts_full_width_colon_markers(monkeypatch):
    _install_book(monkeypatch, _workbook_rows(
        [["ZHANG, 张三", 1727264.0, _serial(2099, 8, 20), "法定病假", 8.0]],
        [["ZHANG, 张三", "01727264", "病假时间总计", 8.0]],
        colon="：",
    ))

    records, summaries, errors = _read_workbook(b"ignored")

    assert errors == []
    assert [row["source_id"] for row in records] == ["01727264"]
    assert _reconcile(records, summaries) == []


def test_read_workbook_requires_block_markers(monkeypatch):
    _install_book(monkeypatch, [["员工", "ID", "日期", "工资代码", "时数"]])

    with pytest.raises(HTTPException) as excinfo:
        _read_workbook(b"ignored")

    assert excinfo.value.status_code == 400
    assert "事务:" in excinfo.value.detail


def test_preview_end_to_end_reports_unmatched_and_serves_marked_file(monkeypatch):
    _install_book(monkeypatch, _workbook_rows(
        [["NOBODY, 查无此人", 9999999.0, _serial(2099, 8, 20), "法定病假", 8.0]],
        [["NOBODY, 查无此人", "09999999", "病假时间总计", 8.0]],
    ))
    with TestClient(app) as client:
        login = client.post("/api/login", json={"employee_no": "GSMTEST01", "password": "1234"})
        assert login.status_code == 200, login.text
        response = client.post("/api/sick-leave-imports/preview", files={"workbook": ("transactions.xls", b"fake-xls")})
        assert response.status_code == 200, response.text
        body = response.json()
        token = body["token"]
        try:
            assert body["month"] == "2099-08"
            assert body["blocking_errors"] == []
            assert body["matched_record_count"] == 0
            assert body["can_commit"] is False
            assert [(row["source_id"], row["employee_no"], row["reason"]) for row in body["unmatched"]] == [
                ("09999999", "9999999", "系统中不存在该员工号"),
            ]

            download = client.get(f"/api/sick-leave-imports/{token}/unmatched-file")
            assert download.status_code == 200, download.text
            assert download.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            exported = load_workbook(BytesIO(download.content))
            assert exported.sheetnames == ["员工事务", "未覆盖说明"]
            notes = list(exported["未覆盖说明"].iter_rows(min_row=2, values_only=True))
            assert [(row[0], row[2], row[3], row[8]) for row in notes] == [(5, "09999999", "9999999", "系统中不存在该员工号")]
            assert exported["员工事务"].cell(row=5, column=1).fill.fgColor.rgb.endswith("FFC7CE")

            committed = client.post("/api/sick-leave-imports/commit", json={"token": token})
            assert committed.status_code == 400, committed.text
            assert "预检未通过" in committed.json()["detail"]
        finally:
            sick_leave_import._drop_preview(token)


def test_expired_preview_token_explains_how_to_recover():
    with TestClient(app) as client:
        login = client.post("/api/login", json={"employee_no": "GSMTEST01", "password": "1234"})
        assert login.status_code == 200, login.text
        committed = client.post("/api/sick-leave-imports/commit", json={"token": "missing-token"})
        assert committed.status_code == 400, committed.text
        assert "服务重启" in committed.json()["detail"]
        assert "重新上传文件预检" in committed.json()["detail"]
        download = client.get("/api/sick-leave-imports/missing-token/unmatched-file")
        assert download.status_code == 404, download.text
        assert "重新上传文件预检" in download.json()["detail"]

"""Batch registration remains independent recognitions, not a new business table."""
import json
import os
import tempfile
from datetime import date
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
import pytest

os.environ["RECOGNITION_V2_DATA_DIR"] = tempfile.mkdtemp(prefix="pr-batch-")
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from PIL import Image
from fastapi.testclient import TestClient
from app.main import app
from app.v2_database import SessionLocal, FILE_DIR
from app.v2_models import RecognitionRecord, RecognitionAttachment, StoredFile


def login(client, account):
    client.post("/api/logout")
    assert client.post("/api/login", json={"employee_no": account, "password": "1234"}).status_code == 200


def picture():
    out = BytesIO()
    Image.new("RGB", (4, 4), "white").save(out, format="PNG")
    return out.getvalue()


def entry(client, content="独立表现", code="SAFETY"):
    opts = client.get("/api/options").json()
    return {"recognition_date": date.today().isoformat(), "content": content,
            "occurred_attraction_id": next(x["id"] for x in opts["recognition_venues"] if x["name"] == "热力追踪"),
            "recognition_type_id": next(x["id"] for x in opts["recognition_types"] if x["code"] == code),
            "recognizer_employee_id": next(x["id"] for x in client.get("/api/recognizers", params={"attraction_id": client.get("/api/me").json()["attraction_id"], "recognition_date": date.today().isoformat()}).json() if x["employee_no"] == "GSMTEST01")}


def send(client, items, key):
    return client.post("/api/recognitions/batch", data={"entries": json.dumps(items), "idempotency_key": key},
                       files={"image": ("evidence.png", picture(), "image/png")})


@pytest.mark.parametrize("batch", [False, True])
@pytest.mark.parametrize("format", ["HEIF", "MPO"])
def test_iphone_heif_with_jpeg_filename_can_be_registered_and_previewed(batch, format):
    out = BytesIO()
    extra = {"save_all": True, "append_images": [Image.new("RGB", (24, 32), "white")]} if format == "MPO" else {}
    Image.new("RGB", (24, 32), "blue").save(out, format=format, **extra)
    photo = out.getvalue()
    with TestClient(app) as client:
        login(client, "CMTEST01")
        item = dict(entry(client, "手机照片兼容"), same_day_duplicate_confirmed=True)
        key = f"iphone-photo-{format}-{batch}"
        data = {"entries": json.dumps([item]), "idempotency_key": key} if batch else {**item, "idempotency_key": key}
        response = client.post("/api/recognitions/batch" if batch else "/api/recognitions", data=data,
                               files={"image": ("IMG_4845.jpeg", photo, "image/jpeg")})
        assert response.status_code == 200, response.text
        record = response.json()["records"][0] if batch else response.json()["record"]
        with SessionLocal() as db:
            attachment = db.query(RecognitionAttachment).filter_by(recognition_id=record["id"]).one()
            stored = db.get(StoredFile, attachment.file_id)
            extension, mime = (".heic", "image/heic") if format == "HEIF" else (".jpg", "image/jpeg")
            assert stored.extension == extension and stored.mime_type == mime
            assert stored.original_filename == "IMG_4845.jpeg"
            assert (FILE_DIR / stored.storage_key).read_bytes() == photo
            file_id = stored.id
        preview = client.get(f"/api/files/{file_id}?preview=true")
        assert preview.status_code == 200, preview.text
        assert preview.headers["content-type"].startswith("image/jpeg")
        with Image.open(BytesIO(preview.content)) as image:
            image.load()
            assert image.size == (24, 32)


def test_five_rows_one_photo_and_retry_conflict():
    with TestClient(app) as client:
        login(client, "CMTEST01")
        items = [dict(entry(client, f"表现{i}"), same_day_duplicate_confirmed=True) for i in range(5)]
        response = send(client, items, "five")
        assert response.status_code == 200, response.text
        records = response.json()["records"]
        assert len({r["id"] for r in records}) == 5
        assert all(r["status"] == "pending" for r in records)
        with SessionLocal() as db:
            attachments = db.query(RecognitionAttachment).filter(RecognitionAttachment.recognition_id.in_([r["id"] for r in records])).all()
            assert len(attachments) == 5
            assert len({r.file_id for r in attachments}) == 1
        retry = send(client, items, "five")
        assert retry.status_code == 200, retry.text
        assert [r["id"] for r in retry.json()["records"]] == [r["id"] for r in records]
        items[0]["content"] = "修改内容"
        assert send(client, items, "five").status_code == 409
        mine = client.get("/api/dashboard").json()["records"]
        assert all(r["id"] in {x["id"] for x in mine} for r in records)


def test_invalid_second_row_rolls_back_rows_photo_and_receipts():
    with TestClient(app) as client:
        login(client, "TRTEST01")
        with SessionLocal() as db:
            before = (db.query(RecognitionRecord).count(), db.query(StoredFile).count())
        files = set(FILE_DIR.iterdir())
        bad = dict(entry(client), content="字" * 21)
        result = send(client, [entry(client), bad], "invalid")
        assert result.status_code == 400, result.text
        assert result.json()["detail"]["entry_index"] == 2
        with SessionLocal() as db:
            assert before == (db.query(RecognitionRecord).count(), db.query(StoredFile).count())
        assert set(FILE_DIR.iterdir()) == files
        assert send(client, [entry(client)], "invalid").status_code == 200


def test_size_permission_and_no_target_override():
    with TestClient(app) as client:
        login(client, "TATEST01")
        item = entry(client, "TA本人")
        assert send(client, [item], "ta").status_code == 200
        assert send(client, [], "empty").status_code == 400
        assert send(client, [item] * 6, "six").status_code == 400
        assert send(client, [dict(item, employee_id=1)], "override").status_code == 400
        login(client, "SUPTEST01")
        assert send(client, [item], "supervisor").status_code == 403


def test_in_batch_duplicates_require_confirmation_and_quota_is_atomic():
    with TestClient(app) as client:
        login(client, "CMTEST02")
        items = [entry(client, "同批1", "EFFICIENCY"), entry(client, "同批2", "EFFICIENCY")]
        result = send(client, items, "duplicate")
        assert result.status_code == 409
        assert result.json()["detail"]["code"] == "SAME_DAY_RECOGNITION_DUPLICATE"
        assert result.json()["detail"]["entry_index"] == 2
        items[1]["same_day_duplicate_confirmed"] = True
        assert send(client, items, "duplicate").status_code == 200
        letter = entry(client, "表扬信", "COMMENDATION_LETTER")
        letter["recognizer_employee_id"] = "special:COMMENDATION_LETTER"
        letter["same_day_duplicate_confirmed"] = True
        result = send(client, [letter, letter], "quota")
        assert result.status_code == 409, result.text
        assert result.json()["detail"]["entry_index"] == 2
        assert send(client, [letter], "quota").status_code == 200


def test_concurrent_reviews_cap_and_statistics():
    with TestClient(app) as client:
        login(client, "TRTEST01")
        items = [dict(entry(client, f"上限{i}", "COURTESY"), same_day_duplicate_confirmed=True) for i in range(5)]
        result = send(client, items, "cap")
        assert result.status_code == 200, result.text
        ids = [r["id"] for r in result.json()["records"]]
        def confirm(record_id):
            with TestClient(app) as reviewer:
                login(reviewer, "TATEST01")
                return reviewer.post(f"/api/reviews/{record_id}", json={"action": "confirm"})
        with ThreadPoolExecutor(max_workers=5) as pool:
            responses = list(pool.map(confirm, ids))
        assert all(r.status_code == 200 for r in responses), [r.text for r in responses]
        with SessionLocal() as db:
            rows = db.query(RecognitionRecord).filter(RecognitionRecord.id.in_(ids)).all()
            assert all(r.status == "confirmed" for r in rows)
            assert sum(r.credited_fraction for r in rows) == min(sum(r.fraction for r in rows), 5)
            assert sum(r.credited_fraction for r in rows) <= 5
        login(client, "TRTEST01")
        dashboard = client.get("/api/dashboard").json()
        mine = dashboard["records"]
        assert len([r for r in mine if r["id"] in ids]) == 5
        assert dashboard["category_scores"]["礼仪"] <= 5
        login(client, "GSMTEST01")
        params = {"start_date": date.today().replace(day=1).isoformat(), "end_date": date.today().isoformat(),
                  "category": "recognition", "subtype_id": items[0]["recognition_type_id"], "keyword": "TRTEST01"}
        ranking = client.get("/api/pr-rankings", params=params)
        assert ranking.status_code == 200, ranking.text
        ranked = ranking.json()["rows"][0]
        assert ranked["count"] == 5
        assert ranked["score"] == dashboard["category_scores"]["礼仪"]
        exported = client.get("/api/pr-rankings/export", params=params)
        assert exported.status_code == 200, exported.text
        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(exported.content))["PR排名数据"]
        assert sheet.cell(2, 8).value == ranked["count"]
        assert sheet.cell(2, 9).value == ranked["score"]
        assert sheet.cell(2, 11).value == ranked["uncapped_score"]


def test_later_link_reviewer_can_view_shared_photo():
    with TestClient(app) as client:
        login(client, "CMTEST01")
        result = send(client, [dict(entry(client, f"材料{i}", "SHOW"), same_day_duplicate_confirmed=True) for i in range(2)], "photo")
        assert result.status_code == 200, result.text
        records = result.json()["records"]
        login(client, "CMTEST02")
        other_id = client.get("/api/me").json()["id"]
        with SessionLocal() as db:
            row = db.get(RecognitionRecord, records[1]["id"])
            row.reviewed_by = other_id
            file_id = row.attachments[0].file_id
            db.commit()
        assert client.get(f"/api/files/{file_id}?preview=true").status_code == 200
        login(client, "TRTEST01")
        assert client.get(f"/api/files/{file_id}?preview=true").status_code == 403


def test_legacy_attachment_migration_preserves_rows_and_is_repeatable():
    from sqlalchemy import create_engine, text
    from sqlalchemy.orm import Session
    from app.v2_database import ensure_recognition_shared_evidence
    engine = create_engine("sqlite:///:memory:")
    with Session(engine) as db:
        db.execute(text("CREATE TABLE recognition_records (id INTEGER PRIMARY KEY)"))
        db.execute(text("CREATE TABLE stored_files (id INTEGER PRIMARY KEY)"))
        db.execute(text("INSERT INTO recognition_records VALUES (1),(2)"))
        db.execute(text("INSERT INTO stored_files VALUES (1)"))
        db.execute(text("""CREATE TABLE recognition_attachments (
            id INTEGER PRIMARY KEY, recognition_id INTEGER NOT NULL,
            file_id INTEGER NOT NULL UNIQUE, attachment_type VARCHAR(30) NOT NULL,
            sort_order INTEGER NOT NULL, created_at DATETIME NOT NULL,
            UNIQUE(recognition_id, attachment_type))"""))
        db.execute(text("INSERT INTO recognition_attachments VALUES(9,1,1,'evidence',1,'2026-10-01')"))
        before = db.execute(text("SELECT * FROM recognition_attachments")).all()
        ensure_recognition_shared_evidence(db)
        assert db.execute(text("SELECT * FROM recognition_attachments")).all() == before
        ensure_recognition_shared_evidence(db)
        db.execute(text("INSERT INTO recognition_attachments VALUES(10,2,1,'evidence',1,'2026-10-09')"))
        assert db.execute(text("SELECT COUNT(*) FROM recognition_attachments")).scalar() == 2
        assert not db.execute(text("PRAGMA foreign_key_check")).all()


def test_concurrent_retries_and_independent_withdrawal():
    with TestClient(app) as client:
        login(client, "CMTEST02")
        items = [dict(entry(client, f"并发{i}", "INCLUSION"), same_day_duplicate_confirmed=True) for i in range(2)]
        def submit(_):
            with TestClient(app) as browser:
                login(browser, "CMTEST02")
                return send(browser, items, "concurrent-retry")
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(submit, range(2)))
        assert all(x.status_code == 200 for x in results), [x.text for x in results]
        ids = [r["id"] for r in results[0].json()["records"]]
        assert ids == [r["id"] for r in results[1].json()["records"]]
        response = client.delete(f"/api/recognitions/{ids[0]}")
        assert response.status_code == 200, response.text
        with SessionLocal() as db:
            assert db.get(RecognitionRecord, ids[0]).status == "void"
            other = db.get(RecognitionRecord, ids[1])
            assert other.status == "pending"
            file_id = other.attachments[0].file_id
        assert client.get(f"/api/files/{file_id}?preview=true").status_code == 200


def test_rejection_reason_is_validated_before_any_state_change():
    with TestClient(app) as client:
        login(client, "CMTEST01")
        response = send(client, [entry(client, "校验原因", "EFFICIENCY")], "review-reason-validation")
        assert response.status_code == 200, response.text
        record_id = response.json()["records"][0]["id"]
        login(client, "TATEST01")
        for note in ("", "  \n\t", "原" * 1001):
            result = client.post(f"/api/reviews/{record_id}", json={"action": "reject", "note": note})
            assert result.status_code == 400, result.text
            with SessionLocal() as db:
                assert db.get(RecognitionRecord, record_id).status == "pending"
        result = client.post(f"/api/reviews/{record_id}", json={"action": "reject", "note": "  图片与内容不符  "})
        assert result.status_code == 200, result.text
        with SessionLocal() as db:
            row = db.get(RecognitionRecord, record_id)
            assert row.status == "rejected" and row.review_note == "图片与内容不符"


def test_forged_picture_rejected_before_creating_business_rows_or_files():
    with TestClient(app) as client:
        login(client, "CMTEST01")
        item = entry(client, "损坏材料", "EFFICIENCY")
        before_files = set(FILE_DIR.iterdir())
        with SessionLocal() as db:
            before = db.query(RecognitionRecord).count(), db.query(StoredFile).count()
        result = client.post("/api/recognitions/batch", data={"entries": json.dumps([item]), "idempotency_key": "forged-picture"},
            files={"image": ("photo.png", b"\x89PNG\r\n\x1a\nnot-a-real-picture", "image/png")})
        assert result.status_code == 400, result.text
        assert set(FILE_DIR.iterdir()) == before_files
        with SessionLocal() as db:
            assert (db.query(RecognitionRecord).count(), db.query(StoredFile).count()) == before


@pytest.mark.parametrize("browser", ["BaiduBrowser", "baiduboxapp", "UCBrowser", "MicroMessenger", "MQQBrowser", "CriOS", "Safari"])
@pytest.mark.parametrize("batch", [False, True])
def test_branded_user_agent_does_not_change_upload_rules(browser, batch):
    # Server boundary test with synthetic user agents, not physical client certification.
    with TestClient(app) as client:
        login(client, "CMTEST01")
        item = dict(entry(client, "浏览器上传边界"), same_day_duplicate_confirmed=True)
        key = f"browser-upload-{browser}-{batch}"
        data = {"entries": json.dumps([item]), "idempotency_key": key} if batch else {**item, "idempotency_key": key}
        data["image_client_ready"] = "1"
        photo = picture()
        response = client.post("/api/recognitions/batch" if batch else "/api/recognitions", data=data,
            files={"image": ("合成认可.png", photo, "image/png")}, headers={"User-Agent": f"Mozilla/5.0 {browser}/fixture"})
        assert response.status_code == 200, response.text
        record = response.json()["records"][0] if batch else response.json()["record"]
        assert record["status"] == "pending"
        with SessionLocal() as db:
            attachment = db.query(RecognitionAttachment).filter_by(recognition_id=record["id"]).one()
            stored = db.get(StoredFile, attachment.file_id)
            assert stored.original_filename == "合成认可.png"
            assert (FILE_DIR / stored.storage_key).read_bytes() == photo


@pytest.mark.parametrize("batch", [False, True])
def test_confirmed_missing_image_does_not_consume_key_or_create_records(batch):
    with TestClient(app) as client:
        login(client, "CMTEST01")
        item = dict(entry(client, "缺图重传边界"), same_day_duplicate_confirmed=True)
        key = f"missing-image-retry-{batch}"
        data = {"entries": json.dumps([item]), "idempotency_key": key} if batch else {**item, "idempotency_key": key}
        data["image_client_ready"] = "1"
        target = "/api/recognitions/batch" if batch else "/api/recognitions"
        with SessionLocal() as db:
            before = db.query(RecognitionRecord).count(), db.query(StoredFile).count()
            original_groups = [(r.id, r.same_day_duplicate_group, r.same_day_duplicate_sequence) for r in db.query(RecognitionRecord).order_by(RecognitionRecord.id)]
        rejected = client.post(target, data=data)
        assert rejected.status_code == (422 if batch else 400)
        if batch:
            assert rejected.json()["detail"]["code"] == "RECOGNITION_IMAGE_BROWSER_INCOMPATIBLE"
        else:
            assert rejected.json()["detail"] == "CM/TR本人登记必须上传1张认可图片"
        with SessionLocal() as db:
            assert (db.query(RecognitionRecord).count(), db.query(StoredFile).count()) == before
            assert [(r.id, r.same_day_duplicate_group, r.same_day_duplicate_sequence) for r in db.query(RecognitionRecord).order_by(RecognitionRecord.id)] == original_groups
        photo = picture()
        first = client.post(target, data=data, files={"image": ("fixture.png", photo, "image/png")})
        retry = client.post(target, data=data, files={"image": ("fixture.png", photo, "image/png")})
        assert first.status_code == retry.status_code == 200
        first_row = first.json()["records"][0] if batch else first.json()["record"]
        retry_row = retry.json()["records"][0] if batch else retry.json()["record"]
        assert first_row["id"] == retry_row["id"]
        with SessionLocal() as db:
            assert db.query(RecognitionRecord).count() == before[0] + 1
            assert db.query(StoredFile).count() == before[1] + 1

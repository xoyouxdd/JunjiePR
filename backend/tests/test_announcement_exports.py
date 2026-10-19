"""Report layout, evidence integrity, scope-safe facts, and reproducible QA samples."""
import hashlib
import importlib.util
import json
import os
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
import tempfile
from time import perf_counter
from types import SimpleNamespace
from uuid import uuid4

os.environ["RECOGNITION_V2_DATA_DIR"] = tempfile.mkdtemp(prefix="announcement-export-")
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw
import pytest

from app.main import app
from app.announcement_exports import pdf, spreadsheet, report_groups, signature_previews, expected_statement
from app.announcement_export_data import prepare_export_data
from app.v2_database import SessionLocal, FILE_DIR
from app.v2_models import Employee, GroupLeaderAssignment, StoredFile, UserAccount, WorkGroup

OUTPUT = Path(__file__).resolve().parents[2] / "output" / "announcement-export-20261018"
USER = SimpleNamespace(name="模拟导出人",employee=SimpleNamespace(employee_no="0000001"))


@pytest.fixture(scope="module",autouse=True)
def startup():
    with TestClient(app):
        yield


def sample():
    announcement = {"id":128,"version":2,"title":"模拟公告：开班检查与岗位交接要求","category":"安全 & 合规",
        "author":"模拟发布人","published_at":"2026-10-17T09:00:00","updated_at":"2026-10-17T09:00:00",
        "scope_label":"热力追踪F组、E组","effective_on":"2026-10-17","due_on":"2026-10-20","expires_on":None,
        "effective":True,"confirmation_level":3,"training_required":False,"summary":"开班检查、交接说明、异常报告。",
        "change_summary":"补充交接事项说明。","body":"开班前核对设备状态与现场通道。\n交接时说明未完成事项，确认接班人员知悉。\n发现异常及时报告，按现场流程处理。"}
    rows = []
    with SessionLocal() as db:
        uploader = db.query(Employee).first().id
        for i in range(13):
            group = "F组" if i < 8 else "E组"
            status = "pending" if i in {6,11} else "exited" if i == 7 else "suspended" if i == 12 else "completed"
            name = f"模拟人员{i+1}"
            row = {"id":900+i,"employee_id":900+i,"employee_no":f"009{i:04d}","employee_name":name,
                "attraction_name":"热力追踪","attraction_id":1,"group_name":group,"group_id":1 if i<8 else 2,"group_key":f"group:{1 if i<8 else 2}",
                "role_label":"CM / TR","status":status,"reason":"适用公告补收" if i==6 else "首次接收",
                "confirmed_at":"2026-10-18T09:20:35" if status=="completed" else None,"seen_at":None,
                "delivered_at":"2026-10-17T09:00:00","due_on":"2026-10-20","signature_url":None,"signing_statement":None}
            if i in {0,1,8}:
                row["export_role"] = "代理组长" if i==1 else "组长"
            if i==6:row["status_note"]="本版发布后账号新建（2026-10-18 11:00）；11:05补收；尚无阅读及签收记录"
            if i==7:row["status_note"]="当前账号停用（2026-10-18 10:20）；已退出适用范围"
            if i==11:row["status_note"]="已阅读，尚未签收；未签原因未记录"
            if i==12:row["status_note"]="当前有有效LOA记录；签收暂缓"
            if status=="completed":
                image = Image.new("RGB",(1600,600),"white")
                ImageDraw.Draw(image).line([(300,360),(460,190),(400,410),(640,210),(780,360),(940,180),(1100,350+i)],fill="black",width=7)
                out=BytesIO();image.save(out,format="PNG");raw=out.getvalue()
                stored=StoredFile(storage_key=f"sample-{uuid4().hex}.png",original_filename="模拟签名.png",extension=".png",mime_type="image/png",
                    uploaded_by=uploader,file_size=len(raw),sha256=hashlib.sha256(raw).hexdigest(),status="active")
                db.add(stored);db.flush();(FILE_DIR/stored.storage_key).write_bytes(raw)
                row["signature_url"]=f"/api/files/{stored.id}?preview=true"
                row["signing_statement"]=expected_statement(row,announcement)
            rows.append(row)
        db.commit()
    return {"announcement":announcement,"items":rows,"groups":[
        {"id":1,"key":"group:1","name":"F组","attraction_id":1,"formal_leader_name":"模拟人员1","acting_leader_name":"模拟人员2"},
        {"id":2,"key":"group:2","name":"E组","attraction_id":1,"formal_leader_name":"模拟人员9","acting_leader_name":""}],
        "counts":{"required":11,"completed":9,"pending":2,"suspended":1,"exited":1},"scoped":True}


def test_compact_reports_keep_signatures_together_and_generate_review_samples():
    data=sample()
    OUTPUT.mkdir(parents=True,exist_ok=True)
    started=perf_counter();content=pdf(data,USER);pdf_seconds=perf_counter()-started
    started=perf_counter();excel=spreadsheet(data,USER);excel_seconds=perf_counter()-started
    (OUTPUT/"receipt-sample.pdf").write_bytes(content)
    (OUTPUT/"receipt-sample.xlsx").write_bytes(excel)
    with fitz.open(stream=content,filetype="pdf") as document:
        assert document.page_count==3
        assert "发布内容" in document[0].get_text() and "开班前核对" in document[0].get_text()
        assert "如何阅读" not in "".join(p.get_text() for p in document)
        for page_index,group in enumerate(report_groups(data),1):
            page=document[page_index]
            page_text=page.get_text()
            assert all(r["employee_name"] in page_text for r in group["items"])
            assert "2026-10-18" in page_text and "09:20:35" in page_text
            assert len(page.get_images())==sum(bool(r.get("signature_url")) for r in group["items"])
        proof=json.loads(document.embfile_get(0))
        assert len(proof["records"])==13
        assert len([r for r in proof["records"] if len(r["SHA256"])==64])==9
        for i,page in enumerate(document):
            page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(OUTPUT/f"receipt-sample-{i+1}.png")
    book=load_workbook(BytesIO(excel))
    assert len(book.worksheets)==4 and book.active.title=="公告与总览"
    assert sum(len(s._images) for s in book)==9
    assert all(image.anchor.editAs=="twoCell" for s in book for image in s._images)
    assert book.worksheets[1]["B7"].value=="组长" and book.worksheets[1]["B8"].value=="代理组长"
    assert book["凭据核对"].sheet_state=="hidden"
    previews=signature_previews(data)
    assert all(s["size"][0]<=600 and s["size"][1]<=180 for s in previews.values())
    with SessionLocal() as db:
        for stored in db.query(StoredFile).filter(StoredFile.id.in_(previews)):
            assert hashlib.sha256((FILE_DIR/stored.storage_key).read_bytes()).hexdigest()==stored.sha256
    measurements={"people":13,"signatures":9,"pdf_pages":3,"pdf_bytes":len(content),"pdf_seconds":round(pdf_seconds,3),
                  "excel_bytes":len(excel),"excel_seconds":round(excel_seconds,3),"production_sample":False}
    previous=OUTPUT/"previous_export.py"
    if previous.exists():
        spec=importlib.util.spec_from_file_location("previous_export",previous)
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        started=perf_counter();before=module.pdf(data,USER);elapsed=perf_counter()-started
        with fitz.open(stream=before,filetype="pdf") as old:
            measurements.update(previous_pdf_pages=old.page_count,previous_pdf_bytes=len(before),previous_pdf_seconds=round(elapsed,3))
        (OUTPUT/"receipt-before.pdf").write_bytes(before)
    (OUTPUT/"measurements.json").write_text(json.dumps(measurements,ensure_ascii=False,indent=2),encoding="utf-8")


@pytest.mark.parametrize("format",["pdf","xlsx"])
def test_tampered_signature_stops_both_formats(format):
    data=sample()
    file_id=int(data["items"][0]["signature_url"].split("/")[-1].split("?")[0])
    with SessionLocal() as db:path=FILE_DIR/db.get(StoredFile,file_id).storage_key
    path.write_bytes(b"tampered")
    try:
        with pytest.raises(ValueError,match="摘要不匹配"):
            (pdf if format=="pdf" else spreadsheet)(data,USER)
        path.unlink()
        with pytest.raises(ValueError,match="文件缺失"):
            (pdf if format=="pdf" else spreadsheet)(data,USER)
    finally:
        path.unlink(missing_ok=True)


def test_export_facts_only_use_authorized_ids_and_do_not_change_counts():
    with SessionLocal() as db:
        base=db.query(Employee).first()
        circle=base.attraction_id
        groups=[]
        for index in range(2):
            group=WorkGroup(name=f"模拟导出组{index}",code=f"X{index}",attraction_id=circle)
            db.add(group);db.flush();groups.append(group)
        people=[]
        for index in range(4):
            employee=Employee(employee_no=f"EXPORT-{index}",name=f"导出测试{index}",attraction_id=circle)
            db.add(employee);db.flush();people.append(employee)
            db.add(UserAccount(employee_id=employee.id,login_account=employee.employee_no,password_hash="unused",
                enabled=index!=1,disabled_at=datetime.now() if index==1 else None,created_at=datetime.now() if index==2 else datetime(2020,1,1)))
        # The ungrouped leader has an actual receipt; a second group's leader has no authorized receipt.
        db.add_all([GroupLeaderAssignment(group_id=groups[0].id,leader_employee_id=people[0].id,leader_type="formal",starts_on="2020-01-01"),
                    GroupLeaderAssignment(group_id=groups[1].id,leader_employee_id=people[0].id,leader_type="acting",starts_on="2020-01-01"),
                    GroupLeaderAssignment(group_id=groups[1].id,leader_employee_id=people[3].id,leader_type="formal",starts_on="2020-01-01")])
        db.flush()
        records=[]
        for i,person in enumerate(people[:3]):
            records.append({"id":person.id,"employee_id":person.id,"employee_no":person.employee_no,"employee_name":person.name,
                "attraction_id":circle,"attraction_name":"模拟圈","group_id":None if i==0 else groups[i-1].id,
                "group_key":"ungrouped" if i==0 else f"group:{groups[i-1].id}","group_name":"未分组" if i==0 else groups[i-1].name,
                "status":"exited" if i==1 else "pending","seen_at":None,"confirmed_at":None,"reason":"首次接收"})
        data={"announcement":{"published_at":datetime(2021,1,1)},"items":records,"counts":{"required":2,"completed":0,"pending":2,"suspended":0,"exited":1},
              "groups":[{"id":g.id,"key":f"group:{g.id}","name":g.name,"attraction_id":circle} for g in groups]}
        result=prepare_export_data(db,data)
        assert {r["employee_id"] for r in result["items"]}=={p.id for p in people[:3]}
        assert result["counts"]==data["counts"] and records[0]["group_id"] is None
        leader=[r for r in result["items"] if r["employee_id"]==people[0].id]
        assert leader[0]["group_id"]==groups[0].id and leader[0]["export_role"]=="组长"
        assert sum(g["required"] for g in report_groups(result))==2
        assert "当前账号停用" in next(r for r in result["items"] if r["employee_id"]==people[1].id)["status_note"]
        assert "本版发布后账号新建" in next(r for r in result["items"] if r["employee_id"]==people[2].id)["status_note"]
        db.rollback()


def test_large_group_continues_without_omitting_people_or_nonstandard_statement():
    data=sample()
    row=data["items"][0]
    data["items"]=[{**row,"id":10000+i,"employee_no":f"LONG-{i:03d}","employee_name":f"继续人员{i:03d}",
                    "signing_statement":"特殊声明原文：本人已知悉本版要求。"} for i in range(35)]
    data["counts"]={"required":35,"completed":35,"pending":0,"suspended":0,"exited":0}
    with fitz.open(stream=pdf(data,USER),filetype="pdf") as document:
        full="".join(p.get_text() for p in document)
        assert document.page_count>2 and "（续）" in full
        assert all(f"LONG-{i:03d}" in full for i in range(35))
        assert "特殊声明原文" in full
    book=load_workbook(BytesIO(spreadsheet(data,USER)))
    assert book.worksheets[1].row_breaks.brk
    assert len(book.worksheets[1]._images)==35


def test_replaced_or_out_of_scope_leader_is_not_treated_as_current():
    with SessionLocal() as db:
        people=db.query(Employee).limit(2).all()
        circle=people[0].attraction_id
        group=WorkGroup(name="模拟更替组",code="XC",attraction_id=circle)
        db.add(group);db.flush()
        db.add_all([GroupLeaderAssignment(group_id=group.id,leader_employee_id=people[0].id,leader_type="formal",starts_on="2020-01-01"),
                    GroupLeaderAssignment(group_id=group.id,leader_employee_id=people[1].id,leader_type="formal",starts_on="2021-01-01")])
        db.flush()
        item={"id":10001,"employee_id":people[0].id,"employee_no":people[0].employee_no,"employee_name":"授权人员",
              "attraction_id":circle,"attraction_name":"模拟圈","group_id":None,"group_key":"none","group_name":"未分组",
              "status":"pending","confirmed_at":None,"seen_at":None,"reason":"首次接收"}
        data={"announcement":{},"items":[item],"groups":[{"id":group.id,"key":f"group:{group.id}","name":group.name,"attraction_id":circle}]}
        result=prepare_export_data(db,data)
        assert len(result["items"])==1 and result["items"][0]["group_id"] is None
        assert "export_role" not in result["items"][0]
        db.rollback()


def test_deleted_account_and_loa_facts_do_not_export_private_notes():
    from app.v2_models import EmployeeLOAPeriod
    with SessionLocal() as db:
        base=db.query(Employee).first()
        person=Employee(employee_no="EXPORT-ARCHIVED",name="模拟留档",attraction_id=base.attraction_id,
                        account_deleted_at=datetime.now(),is_active=False)
        db.add(person);db.flush()
        db.add(EmployeeLOAPeriod(employee_id=person.id,starts_on=date.today().isoformat(),status="active",
            note="不可导出的私人说明",created_by=base.id,created_by_name="测试"))
        db.flush()
        item={"employee_id":person.id,"status":"suspended","reason":"LOA暂缓","seen_at":None,"confirmed_at":None}
        result=prepare_export_data(db,{"announcement":{},"items":[item],"groups":[]})
        note=result["items"][0]["status_note"]
        assert "账号已删除" in note and "有效LOA" in note and "不可导出" not in note
        db.rollback()

from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from PIL import Image


TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-v219-test-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.v2_database import SessionLocal, ensure_gsm_management_scopes, ensure_highest_admin_account  # noqa: E402
from app.v2_models import (  # noqa: E402
    AttendanceMonthlyScore,
    Attraction,
    AuditLog,
    CircleTransferRequest,
    DeductionFollowUp,
    DeductionLevel,
    DeductionRecord,
    DeductionType,
    Employee,
    EmployeeRoleAssignment,
    EmployeeLOAPeriod,
    GroupLeaderAssignment,
    GroupMembership,
    ManagementScope,
    RecognitionMonthlyQuota,
    RecognitionRecord,
    RecognitionType,
    Role,
    SickLeaveRecord,
    StoredFile,
    UserAccount,
    WorkGroup,
)


def login(client: TestClient, employee_no: str, password: str = "1234") -> None:
    response = client.post("/api/login", json={"employee_no": employee_no, "password": password})
    assert response.status_code == 200, response.text


def test_sick_leave_validation_returns_field_message_and_writes_metadata_audit() -> None:
    with TestClient(app) as client:
        login(client, "TATEST01")
        targets = client.get("/api/employee-targets", params={"usage": "attendance", "keyword": "CMTEST01"}).json()["items"]
        target = next(row for row in targets if row["employee_no"] == "CMTEST01")
        response = client.post(
            "/api/sick-leaves",
            data={
                "employee_id": str(target["id"]),
                "leave_start_date": date.today().isoformat(),
                "leave_end_date": date.today().isoformat(),
                "leave_days": "1",
            },
        )
        assert response.status_code == 422, response.text
        detail = response.json()["detail"]
        assert detail["code"] == "SICK_LEAVE_VALIDATION_ERROR"
        assert detail["fields"] == [{"field": "proof", "message": "缺勤证明未成功上传，请重新选择文件后提交。"}]
        with SessionLocal() as db:
            audit = db.query(AuditLog).filter(AuditLog.action == "缺勤登记校验失败").order_by(AuditLog.id.desc()).first()
            assert audit is not None
            assert audit.operator_name == "测试TA主管"
            assert '"fields": ["proof"]' in audit.after_json
            assert "filename" not in audit.after_json


def test_baidu_absence_submission_explicitly_sets_all_critical_multipart_fields() -> None:
    script = (Path(__file__).parents[1] / "app" / "static" / "js" / "app.js").read_text(encoding="utf-8")
    for field in ("employee_id", "leave_start_date", "leave_end_date", "leave_days", "note"):
        assert f"data.set('{field}'" in script
    assert "data.set('proof',file,file.name)" in script


def login_circle_hr(client: TestClient, employee_no: str) -> str:
    """Exercise the production-safe reset/forced-change flow for a test HR account."""
    login(client, "HR01", "HR123")
    items = client.get("/api/admin/circle-hr-accounts").json()["items"]
    account = next(item for item in items if item["login_account"] == employee_no)
    reset = client.post(f"/api/admin/circle-hr-accounts/{account['employee_id']}/reset-password")
    assert reset.status_code == 200, reset.text
    reset_password = reset.json()["temporary_password"]
    assert reset_password == employee_no[-4:]
    client.post("/api/logout")
    login(client, employee_no, reset_password)
    assert client.get("/api/options").status_code == 403
    replacement = f"{employee_no.replace('-', '')}Test1"
    changed = client.post("/api/password", json={"current_password": reset_password, "new_password": replacement, "confirm_password": replacement})
    assert changed.status_code == 200, changed.text
    return replacement


def test_void_filter_export_and_security_watermarks() -> None:
    month = date.today().strftime("%Y-%m")
    today = date.today().isoformat()
    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        me = client.get("/api/me").json()
        circle_id = next(row["id"] for row in options["employee_circles"] if row["name"] == "热力追踪")
        venue_id = next(row["id"] for row in options["recognition_venues"] if row["name"] == "热力追踪")
        type_id = next(row["id"] for row in options["recognition_types"] if row["code"] == "SAFETY")
        cm_id = next(row["id"] for row in client.get("/api/employee-targets", params={"usage": "recognition", "keyword": "CMTEST01"}).json()["items"])

        created = client.post(
            "/api/recognitions",
            data={
                "recognition_date": today,
                "occurred_attraction_id": str(venue_id),
                "recognition_type_id": str(type_id),
                "recognizer_employee_id": str(me["id"]),
                "content": "V219撤回留痕测试",
                "employee_id": str(cm_id),
                "idempotency_key": "v219-soft-void-test",
            },
        )
        assert created.status_code == 200, created.text
        record_id = created.json()["record"]["id"]
        withdrawn = client.delete(f"/api/recognitions/{record_id}")
        assert withdrawn.status_code == 200, withdrawn.text
        assert withdrawn.json()["deducted_score"] > 0

        with SessionLocal() as db:
            record = db.get(RecognitionRecord, record_id)
            assert record is not None
            assert record.status == "void"
            assert record.voided_by_name == "测试TA主管"
            assert record.voided_by_role_code == "TA_SUPERVISOR"
            assert record.voided_by_role_name == "TA主管"
            assert "直属小组" in record.void_permission_scope_snapshot
            assert record.voided_from_status == "confirmed"
            assert record.voided_at is not None
            assert db.query(RecognitionMonthlyQuota).filter_by(recognition_id=record_id).count() == 0
            assert db.query(AuditLog).filter_by(entity_type="recognition", entity_id=str(record_id), action="撤回签卡（留档）").count() == 1

        client.post("/api/logout")
        login(client, "GSMTEST01")

        cm_stats = client.get("/api/statistics", params={"month": month, "attraction_id": circle_id, "title": "CM"})
        assert cm_stats.status_code == 200, cm_stats.text
        cm_data = cm_stats.json()
        assert cm_data["data_updated_at"]
        employee_rows = [row for row in cm_data["hierarchy"] if row["node_type"] == "employee"]
        assert employee_rows and all(row["role_code"] == "CM" for row in employee_rows)
        supervisor_rows = [row for row in cm_data["hierarchy"] if row["node_type"] == "supervisor" and row["member_count"]]
        assert supervisor_rows
        for row in supervisor_rows:
            assert row["average_score"] == round(row["total_score"] / row["member_count"], 2)

        all_stats = client.get("/api/statistics", params={"month": month, "attraction_id": circle_id})
        assert all_stats.status_code == 200, all_stats.text
        team_employee_rows = [row for row in all_stats.json()["hierarchy"] if row["node_type"] == "employee"]
        assert team_employee_rows and all(row["role_code"] in {"CM", "TR"} for row in team_employee_rows)

        tr_stats = client.get("/api/statistics", params={"month": month, "attraction_id": circle_id, "title": "TR"})
        assert tr_stats.status_code == 200, tr_stats.text
        assert all(row["role_code"] == "TR" for row in tr_stats.json()["hierarchy"] if row["node_type"] == "employee")
        assert client.get("/api/statistics", params={"month": month, "title": "GSM"}).status_code == 400

        exported = client.get("/api/statistics/export", params={"month": month, "attraction_id": circle_id, "title": "CM"})
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content))
        assert workbook.sheetnames == ["导出说明", "月度综合分", "月度综合分明细", "层级绩效明细", "员工号变更对照", "签卡明细", "扣分明细", "病假明细", "作废操作记录"]
        assert workbook.active.title == "层级绩效明细"
        assert "认可数据" in str(workbook["导出说明"]["A1"].value)
        export_context = {row[0].value: row[1].value for row in workbook["导出说明"].iter_rows(min_row=2)}
        assert export_context["统计月份"] == month
        assert export_context["景点圈"] == "热力追踪"
        assert export_context["Title"] == "CM"
        assert "Title" in [cell.value for cell in workbook["月度综合分"][1]]
        assert workbook["月度综合分"].sheet_properties.outlinePr.summaryBelow is False
        assert workbook["月度综合分"].auto_filter.ref is None
        hierarchy_sheet = workbook["层级绩效明细"]
        assert hierarchy_sheet.sheet_properties.outlinePr.summaryBelow is False
        assert hierarchy_sheet.sheet_view.showOutlineSymbols is True
        assert hierarchy_sheet["F2"].data_type == "f"
        assert "SUMIFS" in hierarchy_sheet["F2"].value
        assert any(row[13].value == "明细" for row in hierarchy_sheet.iter_rows(min_row=2))
        supervisor_export_row = next(row for row in workbook["月度综合分"].iter_rows(min_row=2) if isinstance(row[0].value, str) and row[0].value.startswith("主管组："))
        assert supervisor_export_row[8].value == "小组汇总"
        assert workbook["月度综合分"].row_dimensions[supervisor_export_row[0].row].collapsed is True
        detail_headers = [cell.value for cell in workbook["月度综合分明细"][1]]
        assert detail_headers == ["员工号", "姓名", "Title", "景点圈", "签卡加分", "全勤分", "扣分", "综合分", "人员状态", "工号状态"]
        assert workbook["月度综合分明细"].auto_filter.ref
        assert "当月已确认认可次数" in [cell.value for cell in workbook["签卡明细"][1]]
        assert "记录编号" in [cell.value for cell in workbook["签卡明细"][1]]
        assert "撤回时间" in [cell.value for cell in workbook["签卡明细"][1]]
        assert "同日重复登记" in [cell.value for cell in workbook["签卡明细"][1]]
        void_row = next(row for row in workbook["作废操作记录"].iter_rows(min_row=2) if row[0].value == "签卡撤回" and row[1].value == f"REC-{record_id}")
        assert void_row[8].value == "TA主管"
        assert "直属小组" in void_row[9].value
        detail_row = next(row for row in workbook["签卡明细"].iter_rows(min_row=2) if row[0].value == f"REC-{record_id}")
        assert detail_row[0].hyperlink is not None
        assert void_row[2].hyperlink is not None
        assert workbook["月度综合分"]._images
        assert "GSMTEST01" in (workbook.properties.description or "")
        for sheet in workbook.worksheets:
            assert sheet.protection.sheet is False
            assert sheet.protection.objects is False
            assert sheet.oddHeader.center.text and "GSMTEST01" in sheet.oddHeader.center.text
            assert sheet.oddFooter.center.text and "GSMTEST01" in sheet.oddFooter.center.text
        recent_exports = client.get("/api/statistics/my-exports")
        assert recent_exports.status_code == 200
        assert recent_exports.json()["items"][0]["month"] == month
        assert recent_exports.json()["items"][0]["attraction_name"] == "热力追踪"
        assert recent_exports.json()["items"][0]["title"] == "CM"
        active_details = client.get("/api/statistics/details", params={"month": month, "employee_ids": cm_id, "effective_only": "true"})
        assert active_details.status_code == 200, active_details.text
        assert all(record["included"] for record in active_details.json()["details"][str(cm_id)]["all_records"])

        screenshot = client.post("/api/security/screenshot-event")
        assert screenshot.status_code == 200
        assert screenshot.json()["message"] == "系统已记录截图按键事件；页面访问和导出均受审计，请勿分享敏感数据。"
        with SessionLocal() as db:
            assert db.query(AuditLog).filter(AuditLog.action == "导出统计数据").count() == 1
            assert db.query(AuditLog).filter(AuditLog.action == "检测到截图按键").count() == 1


def test_v220_sick_loa_follow_up_and_password_permissions() -> None:
    today = date.today()
    month = today.strftime("%Y-%m")
    month_start = today.replace(day=1)
    next_month_start = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
    next_month = next_month_start.strftime("%Y-%m")
    first_three_days = [month_start + timedelta(days=index) for index in range(3)]
    pdf_stream = BytesIO()
    Image.new("RGB", (8, 8), "white").save(pdf_stream, format="PDF")
    proof = {"proof": ("proof.pdf", pdf_stream.getvalue(), "application/pdf")}
    document = {"document": ("statement.pdf", pdf_stream.getvalue(), "application/pdf")}

    with TestClient(app) as client:
        login(client, "TATEST01")
        targets = client.get("/api/employee-targets", params={"usage": "attendance", "keyword": "CMTEST"}).json()["items"]
        cm1 = next(row for row in targets if row["employee_no"] == "CMTEST01")
        cm2 = next(row for row in targets if row["employee_no"] == "CMTEST02")
        options = client.get("/api/options").json()
        me = client.get("/api/me").json()
        venue_id = next(row["id"] for row in options["recognition_venues"] if row["name"] == "热力追踪")
        type_id = next(row["id"] for row in options["recognition_types"] if row["code"] == "SAFETY")
        loa_audit_record = client.post(
            "/api/recognitions",
            data={
                "recognition_date": next_month_start.isoformat(),
                "occurred_attraction_id": str(venue_id),
                "recognition_type_id": str(type_id),
                "recognizer_employee_id": str(me["id"]),
                "content": "LOA作废审计测试",
                "employee_id": str(cm1["id"]),
                "idempotency_key": "v221-loa-void-audit",
            },
        )
        assert loa_audit_record.status_code == 200, loa_audit_record.text
        loa_audit_record_id = loa_audit_record.json()["record"]["id"]
        assert client.delete(f"/api/recognitions/{loa_audit_record_id}").status_code == 200

        half_day = client.post(
            "/api/sick-leaves",
            data={
                "employee_id": cm1["id"],
                "leave_start_date": today.isoformat(),
                "leave_end_date": today.isoformat(),
                "leave_days": "0.5",
                "idempotency_key": "v220-half-day",
            },
            files=proof,
        )
        assert half_day.status_code == 200, half_day.text
        assert half_day.json()["attendance_score"] == 9.75

        needs_confirmation = client.post(
            "/api/sick-leaves",
            data={
                "employee_id": cm2["id"],
                "leave_start_date": first_three_days[0].isoformat(),
                "leave_end_date": first_three_days[1].isoformat(),
                "leave_days": "2",
                "idempotency_key": "v220-rest-warning",
            },
            files=proof,
        )
        assert needs_confirmation.status_code == 409
        assert needs_confirmation.json()["detail"]["code"] == "SICK_LEAVE_REST_DAY_CONFIRMATION_REQUIRED"

        first_range = client.post(
            "/api/sick-leaves",
            data={
                "employee_id": cm2["id"],
                "leave_start_date": first_three_days[0].isoformat(),
                "leave_end_date": first_three_days[1].isoformat(),
                "leave_days": "2",
                "rest_day_confirmed": "true",
                "idempotency_key": "v220-range-one",
            },
            files=proof,
        )
        assert first_range.status_code == 200, first_range.text
        first_range_id = client.get("/api/sick-leaves", params={"month": month}).json()[0]["id"]
        assert client.post(f"/api/sick-leaves/{first_range_id}/void", json={"reason": "延长缺勤区间前先作废原记录"}).status_code == 200
        extended = client.post(
            "/api/sick-leaves",
            data={
                "employee_id": cm2["id"],
                "leave_start_date": first_three_days[0].isoformat(),
                "leave_end_date": first_three_days[2].isoformat(),
                "leave_days": "3",
                "rest_day_confirmed": "true",
                "idempotency_key": "v220-range-extended",
            },
            files=proof,
        )
        assert extended.status_code == 200, extended.text
        assert extended.json()["attendance_score"] == 8.65
        with SessionLocal() as db:
            attendance = db.query(AttendanceMonthlyScore).filter_by(employee_id=cm2["id"], attendance_month=month).one()
            assert float(attendance.charged_sick_days) == 3.0

        options = client.get("/api/options").json()
        sick_type_id = next(row["id"] for row in options["deduction_types"] if row["code"] == "SICK_LEAVE_VIOLATION")
        statement_id = next(row["id"] for row in options["deduction_levels"] if row["code"] == "STATEMENT")
        tr = client.get("/api/employee-targets", params={"usage": "deduction", "keyword": "TRTEST01"}).json()["items"][0]
        first_deduction = client.post(
            "/api/deductions",
            data={
                "employee_id": tr["id"],
                "deduction_type_id": sick_type_id,
                "deduction_level_id": statement_id,
                "occurred_on": today.isoformat(),
                "description": "首次违规病假",
                "idempotency_key": "v220-first-violation",
            },
            files=document,
        )
        assert first_deduction.status_code == 200, first_deduction.text
        repeat = client.get(
            "/api/deductions/attendance-repeat-check",
            params={"employee_id": tr["id"], "deduction_type_id": sick_type_id, "occurred_on": today.isoformat()},
        )
        assert repeat.status_code == 200, repeat.text
        assert repeat.json()["blocked"] is True
        entries = client.get("/api/my-entries", params={"record_type": "follow_up"}).json()
        assert entries["items"][0]["status"] == "pending"

        client.post("/api/logout")
        login(client, "GSMTEST01")
        gsm_options = client.get("/api/options").json()
        memo_id = next(row["id"] for row in gsm_options["deduction_levels"] if row["code"] == "MEMO")
        issued = client.post(
            "/api/deductions",
            data={
                "employee_id": tr["id"],
                "deduction_type_id": sick_type_id,
                "deduction_level_id": memo_id,
                "occurred_on": today.isoformat(),
                "description": "经理开具备忘录",
                "repeat_confirmed": "true",
                "idempotency_key": "v220-issued-memo",
            },
            files=document,
        )
        assert issued.status_code == 200, issued.text

        client.post("/api/logout")
        login(client, "TATEST01")
        follow_up = client.get("/api/my-entries", params={"record_type": "follow_up"}).json()["items"][0]
        assert follow_up["status"] == "issued"
        assert follow_up["issued_by_name"] == "测试GSM"
        with SessionLocal() as db:
            assert db.query(DeductionFollowUp).filter_by(status="issued").count() == 1

        client.post("/api/logout")
        login_circle_hr(client, "HR-HEAT")
        employees = client.get("/api/hr/employees").json()
        cm1_hr = next(row for row in employees if row["employee_no"] == "CMTEST01")
        set_loa = client.put(
            f"/api/hr/employees/{cm1_hr['id']}",
            json={
                "employment_status": "loa",
                "loa_start_date": month_start.isoformat(),
                "reason": "长期病假测试",
            },
        )
        assert set_loa.status_code == 200, set_loa.text
        assert client.post("/api/accounts/reset-password", json={"employee_no": "CMTEST02", "name": "测试CM乙"}).status_code == 200
        with SessionLocal() as db:
            assert db.query(EmployeeLOAPeriod).filter_by(employee_id=cm1_hr["id"], status="active").count() == 1

        client.post("/api/logout")
        login(client, "GSMTEST01")
        current_stats = client.get("/api/statistics", params={"month": month}).json()
        assert any(row["employee_id"] == cm1_hr["id"] for row in current_stats["hierarchy"] if row["node_type"] == "employee")
        next_stats = client.get("/api/statistics", params={"month": next_month}).json()
        assert all(row.get("employee_id") != cm1_hr["id"] for row in next_stats["hierarchy"])
        assert any(row["employee_id"] == cm1_hr["id"] for row in next_stats["loa_rows"])
        exported = client.get("/api/statistics/export", params={"month": next_month})
        workbook = load_workbook(BytesIO(exported.content))
        status_col = [cell.value for cell in workbook["月度综合分"][1]].index("人员状态") + 1
        assert any(row[status_col - 1].value == "LOA（长期病假）" for row in workbook["月度综合分"].iter_rows(min_row=2))
        assert any(row[1].value == f"REC-{loa_audit_record_id}" for row in workbook["作废操作记录"].iter_rows(min_row=2))
        assert any(row[0].value == f"REC-{loa_audit_record_id}" for row in workbook["签卡明细"].iter_rows(min_row=2))
        assert client.post("/api/accounts/reset-password", json={"employee_no": "CMTEST02", "name": "测试CM乙"}).status_code == 200

        client.post("/api/logout")
        login(client, "TAGSMTEST01")
        assert client.post("/api/accounts/reset-password", json={"employee_no": "CMTEST02", "name": "测试CM乙"}).status_code == 403


def test_four_role_groups_and_statistics_reads_are_business_read_only() -> None:
    month = (date.today().replace(day=28) + timedelta(days=35)).replace(day=1).strftime("%Y-%m")
    with TestClient(app) as client:
        for employee_no in ("CMTEST01", "TRTEST01"):
            login(client, employee_no)
            assert client.get("/api/statistics", params={"month": month}).status_code == 403
            assert client.get("/api/member-score-summary", params={"month": month}).status_code == 403
            assert client.get("/api/hr/employees").status_code == 403
            client.post("/api/logout")

        for employee_no in ("TATEST01", "SUPTEST01"):
            login(client, employee_no)
            assert client.get("/api/statistics", params={"month": month}).status_code == 403
            assert client.get("/api/member-score-summary", params={"month": month}).status_code == 200
            assert client.get("/api/hr/employees").status_code == 403
            client.post("/api/logout")

        with SessionLocal() as db:
            attendance_before = db.query(AttendanceMonthlyScore).filter_by(attendance_month=month).count()

        login(client, "TAGSMTEST01")
        ta_gsm_me = client.get("/api/me").json()
        assert "DATA_VIEW" in ta_gsm_me["permissions"]
        assert "DATA_EXPORT" not in ta_gsm_me["permissions"]
        circle_ids = [row["id"] for row in client.get("/api/options").json()["employee_circles"]]
        assert len(circle_ids) == 3
        for circle_id in circle_ids:
            assert client.get("/api/statistics", params={"month": month, "attraction_id": circle_id}).status_code == 200
        assert client.get("/api/statistics/export", params={"month": month}).status_code == 403
        assert client.get("/api/statistics/my-exports").status_code == 403
        pr_base = {"start_date": date.today().isoformat(), "end_date": date.today().isoformat()}
        for circle_id in [None, *circle_ids]:
            params = dict(pr_base)
            if circle_id is not None:
                params["attraction_id"] = circle_id
            response = client.get("/api/pr-rankings", params=params)
            assert response.status_code == 200, response.text
            assert response.json()["attraction_id"] == circle_id
        assert client.get("/api/pr-rankings/export", params=pr_base).status_code == 403
        assert client.get("/api/hr/employees").status_code == 403
        client.post("/api/logout")

        login(client, "GSMTEST01")
        gsm_me = client.get("/api/me").json()
        assert {"DATA_VIEW", "DATA_EXPORT"} <= set(gsm_me["permissions"])
        for circle_id in circle_ids:
            assert client.get("/api/statistics", params={"month": month, "attraction_id": circle_id}).status_code == 200
        assert client.get("/api/statistics/export", params={"month": month}).status_code == 200
        gsm_exports = client.get("/api/statistics/my-exports")
        assert gsm_exports.status_code == 200
        assert gsm_exports.json()["items"]
        for circle_id in [None, *circle_ids]:
            params = dict(pr_base)
            if circle_id is not None:
                params["attraction_id"] = circle_id
            response = client.get("/api/pr-rankings", params=params)
            assert response.status_code == 200, response.text
            assert response.json()["attraction_id"] == circle_id
            assert client.get("/api/pr-rankings/export", params=params).status_code == 200
        assert client.get("/api/hr/employees").status_code == 403
        client.post("/api/logout")

        with SessionLocal() as db:
            assert db.query(AttendanceMonthlyScore).filter_by(attendance_month=month).count() == attendance_before

        login(client, "HR01", "HR123")
        assert client.get("/api/statistics", params={"month": month}).status_code == 200
        assert client.get("/api/member-score-summary", params={"month": month}).status_code == 403
        assert client.get("/api/hr/employees").status_code == 200


def test_pr_leader_ranking_counts_each_confirmed_record_once_per_supervisor() -> None:
    """Self-submitted and manager-entered recognitions both credit their supervisor."""
    today_value = date.today().isoformat()
    with SessionLocal() as db:
        supervisor = db.query(Employee).filter_by(employee_no="SUPTEST01").one()
        cm = db.query(Employee).filter_by(employee_no="CMTEST01").one()
        supervisor_id, cm_attraction_id = supervisor.id, cm.attraction_id
        recognition_types = db.query(RecognitionType).filter_by(active=True).order_by(RecognitionType.id).all()
        primary_type, alternate_type = recognition_types[:2]
        primary_type_id, primary_type_name = primary_type.id, primary_type.name

        def add_record(*, recognizer: Employee, operator: Employee, recognition_type: RecognitionType, fraction: str, source: str) -> None:
            db.add(
                RecognitionRecord(
                    employee_id=cm.id,
                    employee_no=cm.employee_no,
                    employee_name=cm.name,
                    employee_role_snapshot="CM",
                    home_attraction_id=cm.attraction_id,
                    home_attraction_name="测试景点",
                    occurred_attraction_id=cm.attraction_id,
                    recognition_date=today_value,
                    recognition_month=today_value[:7],
                    recognition_type_id=recognition_type.id,
                    recognition_type_name=recognition_type.name,
                    content="主管排名测试",
                    recognizer_employee_id=recognizer.id,
                    recognizer_name=recognizer.name,
                    recognizer_role_snapshot="主管" if recognizer.id == supervisor.id else "CM",
                    operator_employee_id=operator.id,
                    operator_name=operator.name,
                    source=source,
                    fraction=Decimal(fraction),
                    credited_fraction=Decimal(fraction),
                    status="confirmed",
                )
            )

        # Same supervisor is signer and operator in the second row: it remains one count.
        add_record(recognizer=supervisor, operator=cm, recognition_type=primary_type, fraction="1.00", source="self")
        add_record(recognizer=supervisor, operator=supervisor, recognition_type=primary_type, fraction="1.50", source="manager")
        add_record(recognizer=cm, operator=supervisor, recognition_type=primary_type, fraction="2.00", source="manager")
        add_record(recognizer=cm, operator=supervisor, recognition_type=alternate_type, fraction="0.50", source="manager")
        db.commit()

    with TestClient(app) as client:
        login(client, "GSMTEST01")
        base_params = {
            "category": "leader",
            "start_date": today_value,
            "end_date": today_value,
            "attraction_id": cm_attraction_id,
        }
        response = client.get("/api/pr-rankings", params=base_params)
        assert response.status_code == 200, response.text
        row = next(item for item in response.json()["rows"] if item["employee_id"] == supervisor_id)
        assert row["count"] == 4
        assert row["score"] == 5.0

        filtered = client.get("/api/pr-rankings", params={**base_params, "subtype_id": primary_type_id})
        assert filtered.status_code == 200, filtered.text
        filtered_row = next(item for item in filtered.json()["rows"] if item["employee_id"] == supervisor_id)
        assert filtered.json()["subtype_name"] == primary_type_name
        assert filtered_row["count"] == 3
        assert filtered_row["score"] == 4.5


def test_circle_hr_transfer_acceptance_is_immediate_and_migrates_current_data() -> None:
    today = date.today()
    today_value = today.isoformat()
    current_month = today.strftime("%Y-%m")
    previous_date = today.replace(day=1) - timedelta(days=1)
    previous_month = previous_date.strftime("%Y-%m")

    with TestClient(app) as client:
        login_circle_hr(client, "HR-DWARF")
        options = client.get("/api/options").json()
        dwarf_id = next(row["id"] for row in options["employee_circles"] if row["name"] == "矮人迷宫")
        created_leader = client.post(
            "/api/hr/employees",
            json={
                "employee_no": "2000001",
                "name": "矮人调动接收主管",
                "role_code": "SUPERVISOR",
                "attraction_id": dwarf_id,
                "hired_on": today_value,
                "password": "2001",
            },
        )
        assert created_leader.status_code == 200, created_leader.text
        target_leader_id = created_leader.json()["employee"]["id"]
        assert all(row["employee_no"] != "1000001" for row in client.get("/api/hr/employees").json())
        client.post("/api/logout")

        login_circle_hr(client, "HR-HEAT")
        heat_options = client.get("/api/options").json()
        heat_id = next(row["id"] for row in heat_options["employee_circles"] if row["name"] == "热力追踪")
        created_employee = client.post(
            "/api/hr/employees",
            json={
                "employee_no": "1000001",
                "name": "跨圈调动测试CM",
                "role_code": "CM",
                "attraction_id": heat_id,
                "hired_on": today_value,
                "password": "0001",
            },
        )
        assert created_employee.status_code == 200, created_employee.text
        employee = created_employee.json()["employee"]
        searchable = client.get(
            "/api/employee-targets",
            params={"usage": "circle_transfer", "keyword": "1000001"},
        )
        assert searchable.status_code == 200, searchable.text
        assert [row["id"] for row in searchable.json()["items"]] == [employee["id"]]

        with SessionLocal() as db:
            employee_row = db.get(Employee, employee["id"])
            source_leader = db.query(Employee).filter_by(employee_no="TATEST01").one()
            recognition_type = db.query(RecognitionType).filter_by(code="SAFETY").one()
            deduction_type = db.query(DeductionType).filter_by(code="ATT_LATE_WITHIN_30").one()
            deduction_level = db.query(DeductionLevel).filter_by(code="STATEMENT").one()
            source_group = db.query(WorkGroup).filter_by(name="V2-A-TA组").one()
            stored_file = StoredFile(
                storage_key="circle-transfer-fixture.pdf",
                original_filename="circle-transfer-fixture.pdf",
                extension=".pdf",
                mime_type="application/pdf",
                file_size=4,
                sha256="a" * 64,
                uploaded_by=source_leader.id,
            )
            db.add(stored_file)
            db.flush()

            def recognition_record(record_date: str, record_month: str, status: str) -> RecognitionRecord:
                return RecognitionRecord(
                    employee_id=employee_row.id,
                    employee_no=employee_row.employee_no,
                    employee_name=employee_row.name,
                    employee_role_snapshot="CM",
                    home_attraction_id=heat_id,
                    home_attraction_name="热力追踪",
                    occurred_attraction_id=heat_id,
                    recognition_date=record_date,
                    recognition_month=record_month,
                    recognition_type_id=recognition_type.id,
                    recognition_type_name=recognition_type.name,
                    content="跨圈迁移测试",
                    recognizer_employee_id=source_leader.id,
                    recognizer_name=source_leader.name,
                    recognizer_role_snapshot="TA_SUPERVISOR",
                    operator_employee_id=source_leader.id,
                    operator_name=source_leader.name,
                    source="supervisor",
                    fraction=Decimal("1.00"),
                    status=status,
                    assigned_reviewer_id=source_leader.id,
                )

            def deduction_record(record_date: str, record_month: str) -> DeductionRecord:
                return DeductionRecord(
                    employee_id=employee_row.id,
                    employee_no=employee_row.employee_no,
                    employee_name=employee_row.name,
                    employee_role_snapshot="CM",
                    attraction_id_snapshot=heat_id,
                    deduction_type_id=deduction_type.id,
                    deduction_type_name=deduction_type.name,
                    deduction_level_id=deduction_level.id,
                    deduction_level_name=deduction_level.name,
                    points=deduction_level.points,
                    occurred_on=record_date,
                    deduction_month=record_month,
                    description="跨圈迁移测试",
                    document_file_id=stored_file.id,
                    submitter_id=source_leader.id,
                    submitter_name=source_leader.name,
                    submitter_role_snapshot="TA_SUPERVISOR",
                    permission_scope_snapshot="直属小组",
                    status="active",
                )

            def sick_record(record_date: str, record_month: str) -> SickLeaveRecord:
                return SickLeaveRecord(
                    employee_id=employee_row.id,
                    employee_no_snapshot=employee_row.employee_no,
                    employee_name_snapshot=employee_row.name,
                    employee_role_snapshot="CM",
                    attraction_id_snapshot=heat_id,
                    attendance_month=record_month,
                    leave_start_date=record_date,
                    leave_end_date=record_date,
                    leave_days=Decimal("0.5"),
                    charged_days=Decimal("0.5"),
                    proof_file_id=stored_file.id,
                    status="active",
                    submitted_by=source_leader.id,
                    submitted_by_name=source_leader.name,
                )

            current_recognition = recognition_record(today_value, current_month, "pending")
            historical_recognition = recognition_record(previous_date.isoformat(), previous_month, "confirmed")
            current_deduction = deduction_record(today_value, current_month)
            historical_deduction = deduction_record(previous_date.isoformat(), previous_month)
            current_sick = sick_record(today_value, current_month)
            historical_sick = sick_record(previous_date.isoformat(), previous_month)
            follow_up = DeductionFollowUp(
                supervisor_id=source_leader.id,
                supervisor_name=source_leader.name,
                employee_id=employee_row.id,
                employee_no=employee_row.employee_no,
                employee_name=employee_row.name,
                deduction_type_id=deduction_type.id,
                deduction_type_name=deduction_type.name,
                occurred_on=today_value,
                previous_record_ids="[]",
                status="pending",
            )
            source_membership = GroupMembership(
                group_id=source_group.id,
                employee_id=employee_row.id,
                starts_on=today_value,
                status="active",
                reason="跨圈测试源小组",
            )
            db.add_all(
                [
                    source_membership,
                    current_recognition,
                    historical_recognition,
                    current_deduction,
                    historical_deduction,
                    current_sick,
                    historical_sick,
                    follow_up,
                ]
            )
            db.commit()
            fixture_ids = {
                "current_recognition": current_recognition.id,
                "historical_recognition": historical_recognition.id,
                "current_deduction": current_deduction.id,
                "historical_deduction": historical_deduction.id,
                "current_sick": current_sick.id,
                "historical_sick": historical_sick.id,
                "follow_up": follow_up.id,
            }

        created = client.post(
            "/api/hr/circle-transfers",
            json={"employee_id": employee["id"], "target_attraction_id": dwarf_id, "reason": "跨圈调动闭环测试"},
        )
        assert created.status_code == 200, created.text
        transfer_id = created.json()["transfer"]["id"]
        duplicate = client.post(
            "/api/hr/circle-transfers",
            json={"employee_id": employee["id"], "target_attraction_id": dwarf_id, "reason": "重复申请"},
        )
        assert duplicate.status_code == 409
        source_approval = client.post(
            f"/api/hr/circle-transfers/{transfer_id}/review",
            json={"action": "accept", "target_leader_id": target_leader_id},
        )
        assert source_approval.status_code == 403
        client.post("/api/logout")

        login_circle_hr(client, "HR-DWARF")
        accepted = client.post(
            f"/api/hr/circle-transfers/{transfer_id}/review",
            json={"action": "accept", "target_leader_id": target_leader_id, "review_note": "接收并立即生效"},
        )
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["transfer"]["status"] == "completed"
        assert accepted.json()["migrated_record_counts"] == {
            "recognitions": 1,
            "deductions": 1,
            "sick_leaves": 1,
            "pending_reviews": 1,
            "pending_follow_ups": 1,
        }
        assert any(row["employee_no"] == "1000001" for row in client.get("/api/hr/employees").json())

        with SessionLocal() as db:
            moved_employee = db.get(Employee, employee["id"])
            assert moved_employee.attraction_id == dwarf_id
            membership = db.query(GroupMembership).filter_by(employee_id=moved_employee.id, status="active").one()
            target_group = db.get(WorkGroup, membership.group_id)
            assert target_group.attraction_id == dwarf_id
            assert db.query(GroupLeaderAssignment).filter_by(group_id=target_group.id, leader_employee_id=target_leader_id, status="active").count() == 1
            assert db.get(RecognitionRecord, fixture_ids["current_recognition"]).home_attraction_id == dwarf_id
            assert db.get(RecognitionRecord, fixture_ids["current_recognition"]).assigned_reviewer_id == target_leader_id
            assert db.get(DeductionRecord, fixture_ids["current_deduction"]).attraction_id_snapshot == dwarf_id
            assert db.get(DeductionRecord, fixture_ids["current_deduction"]).employee_group_id_snapshot == target_group.id
            assert db.get(SickLeaveRecord, fixture_ids["current_sick"]).attraction_id_snapshot == dwarf_id
            assert db.get(DeductionFollowUp, fixture_ids["follow_up"]).supervisor_id == target_leader_id
            assert db.get(RecognitionRecord, fixture_ids["historical_recognition"]).home_attraction_id == heat_id
            assert db.get(DeductionRecord, fixture_ids["historical_deduction"]).attraction_id_snapshot == heat_id
            assert db.get(DeductionRecord, fixture_ids["historical_deduction"]).employee_group_id_snapshot is None
            assert db.get(SickLeaveRecord, fixture_ids["historical_sick"]).attraction_id_snapshot == heat_id
            assert db.query(CircleTransferRequest).filter_by(id=transfer_id, status="completed").count() == 1
            assert db.query(AuditLog).filter_by(
                entity_type="circle_transfer",
                entity_id=str(transfer_id),
                action="确认跨景点圈调动并同步迁移数据",
            ).count() == 1

        client.post("/api/logout")
        login_circle_hr(client, "HR-HEAT")
        assert all(row["employee_no"] != "1000001" for row in client.get("/api/hr/employees").json())

        cm2 = next(row for row in client.get("/api/hr/employees").json() if row["employee_no"] == "CMTEST02")
        rejected_request = client.post(
            "/api/hr/circle-transfers",
            json={"employee_id": cm2["id"], "target_attraction_id": dwarf_id, "reason": "拒绝路径测试"},
        ).json()["transfer"]
        client.post("/api/logout")
        login_circle_hr(client, "HR-DWARF")
        rejected = client.post(
            f"/api/hr/circle-transfers/{rejected_request['id']}/review",
            json={"action": "reject", "review_note": "目标圈暂不接收"},
        )
        assert rejected.status_code == 200 and rejected.json()["transfer"]["status"] == "rejected"
        client.post("/api/logout")
        login_circle_hr(client, "HR-HEAT")
        cancelled_request = client.post(
            "/api/hr/circle-transfers",
            json={"employee_id": cm2["id"], "target_attraction_id": dwarf_id, "reason": "撤回路径测试"},
        ).json()["transfer"]
        cancelled = client.post(f"/api/hr/circle-transfers/{cancelled_request['id']}/cancel")
        assert cancelled.status_code == 200 and cancelled.json()["transfer"]["status"] == "cancelled"


def test_circle_hr_accounts_are_scoped_and_can_change_password() -> None:
    month = date.today().strftime("%Y-%m")
    credentials = (
        ("HR-HEAT", "热力追踪"),
        ("HR-DWARF", "矮人迷宫"),
        ("HR-BEAR", "小熊罐子"),
    )
    with TestClient(app) as client:
        login(client, "HR01", "HR123")
        circle_ids = {row["name"]: row["id"] for row in client.get("/api/options").json()["employee_circles"]}
        for account, circle_name in credentials:
            password = login_circle_hr(client, account)
            me = client.get("/api/me").json()
            assert me["role_code"] == "HR_CIRCLE"
            assert me["attraction_name"] == circle_name
            assert me["must_change_password"] is False
            permissions = set(me["permissions"])
            assert {"HR_MANAGE", "DATA_VIEW", "DATA_EXPORT", "PASSWORD_RESET"} <= permissions
            assert not {"EMPLOYEE_ADD", "SICK_REGISTER", "DEDUCTION_ALL", "REVIEW_DIRECT", "MEMBER_RECORDS"} & permissions
            scoped_options = client.get("/api/options").json()
            assert {row["name"] for row in scoped_options["employee_circles"]} == {circle_name}
            assert {row["code"] for row in scoped_options["roles"]} == {"CM", "TR", "TA_SUPERVISOR", "SUPERVISOR"}
            assert client.get("/api/hr/employees").status_code == 200
            own_stats = client.get("/api/statistics", params={"month": month, "attraction_id": circle_ids[circle_name]})
            assert own_stats.status_code == 200, own_stats.text
            assert client.get("/api/member-records", params={"month": month}).status_code == 403
            assert client.get("/api/member-score-summary", params={"month": month}).status_code == 403
            assert client.get("/api/reviews").status_code == 403
            assert client.get("/api/sick-leaves", params={"month": month}).status_code == 403
            assert client.get("/api/employee-targets", params={"usage": "recognition", "keyword": "CM"}).status_code == 403
            other_circle = next(name for name in circle_ids if name != circle_name)
            assert client.get("/api/statistics", params={"month": month, "attraction_id": circle_ids[other_circle]}).status_code == 403

            new_password = f"{account.replace('-', '')}New2026"
            changed = client.post("/api/password", json={"current_password": password, "new_password": new_password, "confirm_password": new_password})
            assert changed.status_code == 200, changed.text
            client.post("/api/logout")
            assert client.post("/api/login", json={"employee_no": account, "password": password}).status_code == 401
            login(client, account, new_password)
            client.post("/api/logout")

        login_circle_hr(client, "HR-HEAT")
        employees = client.get("/api/hr/employees").json()
        tagsm = next(row for row in employees if row["employee_no"] == "TAGSMTEST01")
        assert client.put(
            f"/api/hr/employees/{tagsm['id']}",
            json={"role_code": "GSM", "reason": "景点圈HR越权验证"},
        ).status_code == 403
        assert client.post(
            "/api/accounts/reset-password",
            json={"employee_no": tagsm["employee_no"], "name": tagsm["name"]},
        ).status_code == 403
        assert client.post(
            "/api/hr/employees",
            json={
                "employee_no": "9990030",
                "name": "越权GSM测试",
                "role_code": "GSM",
                "attraction_id": circle_ids["热力追踪"],
            },
        ).status_code == 403
        import_book = Workbook()
        import_sheet = import_book.active
        import_sheet.append(["员工号", "姓名", "角色代码", "景点圈", "初始密码", "在职", "账号启用", "任职开始日", "任职结束日", "到期恢复角色代码"])
        import_sheet.append(["9990031", "越权导入GSM", "GSM", "热力追踪", "1234", "是", "是", date.today().isoformat(), "", ""])
        import_content = BytesIO()
        import_book.save(import_content)
        assert client.post(
            "/api/hr/import-employees",
            files={"workbook": ("senior-role.xlsx", import_content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).status_code == 400

        rules = client.get("/api/hr/score-rules")
        assert rules.status_code == 200 and len(rules.json()) >= 4
        assert client.post(
            "/api/hr/score-rules",
            json={"role_code": "GSM", "score": "9.99", "effective_date": date.today().isoformat()},
        ).status_code == 403
        scoped_rule = client.post(
            "/api/hr/score-rules",
            json={"role_code": "SUPERVISOR", "score": "1.10", "effective_date": date.today().isoformat()},
        )
        assert scoped_rule.status_code == 403, scoped_rule.text
        client.post("/api/logout")
        login(client, "HR01", "HR123")
        admin_rule = client.post(
            "/api/hr/score-rules",
            json={"role_code": "GSM", "score": "0.50", "effective_date": date.today().isoformat()},
        )
        assert admin_rule.status_code == 200, admin_rule.text
        assert client.get("/api/admin/logs", params={"limit": -1}).status_code == 422
        assert client.get("/api/admin/logs", params={"limit": 0}).status_code == 422
        assert client.get("/api/admin/logs", params={"limit": 1001}).status_code == 422
        assert client.get("/api/admin/logs", params={"limit": 1}).status_code == 200
        assert client.post(
            "/api/hr/score-rules",
            json={"role_code": "SUPERVISOR", "score": "-0.01", "effective_date": date.today().isoformat()},
        ).status_code == 400
        # Restore the seeded values so later tests in this file keep default scores.
        client.post("/api/hr/score-rules", json={"role_code": "SUPERVISOR", "score": "0.50", "effective_date": date.today().isoformat()})
        client.post("/api/hr/score-rules", json={"role_code": "GSM", "score": "1.00", "effective_date": date.today().isoformat()})
        client.post("/api/logout")


def test_gsm_scopes_are_backfilled_and_parallel_gsms_jointly_own_supervisor_branch() -> None:
    today = date.today().isoformat()
    with TestClient(app) as client:
        login(client, "HR01", "HR123")
        options = client.get("/api/options").json()
        heat_id = next(row["id"] for row in options["employee_circles"] if row["name"] == "热力追踪")

        def create(employee_no: str, name: str, role_code: str) -> dict:
            response = client.post(
                "/api/hr/employees",
                json={
                    "employee_no": employee_no,
                    "name": name,
                    "role_code": role_code,
                    "attraction_id": heat_id,
                    "password": "1234",
                },
            )
            assert response.status_code == 200, response.text
            return response.json()["employee"]

        gsm_one = create("9990101", "并行GSM甲", "GSM")
        gsm_two = create("9990102", "并行GSM乙", "GSM")
        tagsm = create("9990103", "同级TAGSM", "TA_GSM")
        cm = create("9990104", "待晋升CM", "CM")

        with SessionLocal() as db:
            for employee_id in (gsm_one["id"], gsm_two["id"], tagsm["id"]):
                assert db.query(ManagementScope).filter_by(employee_id=employee_id, attraction_id=heat_id).count() == 1

            # Simulate an existing GSM imported before this version.  Startup
            # backfill must repair it without any HR-page re-save.
            imported = Employee(employee_no="9990105", name="存量GSM", attraction_id=heat_id, is_active=True, hired_on=today)
            db.add(imported)
            db.flush()
            imported_id = imported.id
            role = db.query(Role).filter_by(code="GSM").one()
            db.add(EmployeeRoleAssignment(employee_id=imported.id, role_id=role.id, starts_on=today, status="active"))
            db.commit()
            assert db.query(ManagementScope).filter_by(employee_id=imported_id).count() == 0
            ensure_gsm_management_scopes(db)
            assert db.query(ManagementScope).filter_by(employee_id=imported_id, attraction_id=heat_id).count() == 1

        promoted = client.put(
            f"/api/hr/employees/{cm['id']}",
            json={"role_code": "GSM", "attraction_id": heat_id, "reason": "组织树自动同步测试"},
        )
        assert promoted.status_code == 200, promoted.text
        with SessionLocal() as db:
            assert db.query(ManagementScope).filter_by(employee_id=cm["id"], attraction_id=heat_id).count() == 1

        organization = client.get("/api/hr/organization")
        assert organization.status_code == 200, organization.text
        rows = organization.json()["rows"]
        heat_node = next(row["node_id"] for row in rows if row["node_type"] == "attraction" and row["name"] == "热力追踪")
        direct_gsm_rows = [
            row for row in rows
            if row["node_type"] == "employee" and row["parent_id"] == heat_node and row.get("hierarchy_role") == "gsm"
        ]
        direct_gsm_ids = {row["employee"]["id"] for row in direct_gsm_rows}
        assert {gsm_one["id"], gsm_two["id"], cm["id"], tagsm["id"], imported_id} <= direct_gsm_ids
        team_row = next(row for row in rows if row["node_type"] == "gsm_team" and row["parent_id"] == heat_node)
        assert "共同承接" in team_row["name"]
        assert gsm_one["name"] in team_row["name"] and gsm_two["name"] in team_row["name"]
        assert not any(row["parent_id"] == f"hr-gsm-{heat_id}-{tagsm['id']}" for row in rows)


def test_hr_unclassified_frontline_batch_leader_assignment_is_atomic() -> None:
    today = date.today().isoformat()
    temporary_ids: list[int] = []
    with TestClient(app) as client:
        login(client, "HR01", "HR123")
        options = client.get("/api/options").json()
        heat_id = next(row["id"] for row in options["employee_circles"] if row["name"] == "热力追踪")
        leader_options = client.get("/api/hr/leader-options")
        assert leader_options.status_code == 200, leader_options.text
        leader = next(row for row in leader_options.json() if row["attraction_id"] == heat_id and row.get("group_id"))

        with SessionLocal() as db:
            role = db.query(Role).filter_by(code="CM").one()
            operator = db.query(Employee).filter_by(employee_no="HR01").one()
            for suffix in ("81", "82"):
                employee = Employee(
                    employee_no=f"99800{suffix}",
                    name=f"批量分组测试{suffix}",
                    attraction_id=heat_id,
                    is_active=True,
                    hired_on=today,
                )
                db.add(employee)
                db.flush()
                temporary_ids.append(employee.id)
                db.add(
                    EmployeeRoleAssignment(
                        employee_id=employee.id,
                        role_id=role.id,
                        starts_on=today,
                        status="active",
                        reason="批量分组隔离测试",
                        created_by=operator.id,
                    )
                )
            db.commit()

        invalid = client.post(
            "/api/hr/employees/batch-leaders",
            json={
                "items": [
                    {"employee_id": temporary_ids[0], "group_id": leader["group_id"], "leader_id": leader["id"]},
                    {"employee_id": temporary_ids[1], "group_id": leader["group_id"], "leader_id": 999999999},
                ]
            },
        )
        assert invalid.status_code == 400, invalid.text
        with SessionLocal() as db:
            assert db.query(GroupMembership).filter(GroupMembership.employee_id.in_(temporary_ids)).count() == 0

        saved = client.post(
            "/api/hr/employees/batch-leaders",
            json={
                "items": [
                    {"employee_id": employee_id, "group_id": leader["group_id"], "leader_id": leader["id"]}
                    for employee_id in temporary_ids
                ],
                "reason": "批量分组接口测试",
            },
        )
        assert saved.status_code == 200, saved.text
        assert saved.json()["updated"] == 2
        with SessionLocal() as db:
            memberships = db.query(GroupMembership).filter(GroupMembership.employee_id.in_(temporary_ids), GroupMembership.status == "active").all()
            assert len(memberships) == 2
            assert {row.group_id for row in memberships} == {leader["group_id"]}
            assert db.query(AuditLog).filter_by(entity_type="employee_group_batch").count() >= 1

    with SessionLocal() as db:
        db.query(AuditLog).filter(
            (AuditLog.entity_type == "employee_group_batch")
            | ((AuditLog.entity_type == "employee_group") & AuditLog.entity_id.in_([str(row) for row in temporary_ids]))
        ).delete(synchronize_session=False)
        db.query(GroupMembership).filter(GroupMembership.employee_id.in_(temporary_ids)).delete(synchronize_session=False)
        db.query(EmployeeRoleAssignment).filter(EmployeeRoleAssignment.employee_id.in_(temporary_ids)).delete(synchronize_session=False)
        db.query(Employee).filter(Employee.id.in_(temporary_ids)).delete(synchronize_session=False)
        db.commit()


def test_highest_admin_can_manage_circle_hr_passwords_and_senior_roles() -> None:
    with SessionLocal() as db:
        highest = db.query(Employee).filter_by(employee_no="HR01").one()
        account = db.query(UserAccount).filter_by(employee_id=highest.id).one()
        password_hash_before = account.password_hash
        hr_role = db.query(Role).filter_by(code="HR_ADMIN").one()
        assignment = db.query(EmployeeRoleAssignment).filter_by(employee_id=highest.id, status="active").one()
        assignment.role_id = hr_role.id
        db.commit()
        ensure_highest_admin_account(db)
        system_role = db.query(Role).filter_by(code="SYSTEM_ADMIN").one()
        assert db.query(EmployeeRoleAssignment).filter_by(employee_id=highest.id, role_id=system_role.id, status="active").count() == 1
        assert db.query(UserAccount).filter_by(employee_id=highest.id).one().password_hash == password_hash_before

    with TestClient(app) as client:
        login(client, "HR01", "HR123")
        me = client.get("/api/me").json()
        assert me["role_code"] == "SYSTEM_ADMIN"
        assert me["role_name"] == "最高管理员"
        accounts = client.get("/api/admin/circle-hr-accounts")
        assert accounts.status_code == 200, accounts.text
        items = accounts.json()["items"]
        assert {item["login_account"] for item in items} == {"HR-HEAT", "HR-DWARF", "HR-BEAR"}
        assert all("temporary_password" not in item for item in items)

        heat = next(item for item in items if item["login_account"] == "HR-HEAT")
        reset = client.post(f"/api/admin/circle-hr-accounts/{heat['employee_id']}/reset-password")
        assert reset.status_code == 200, reset.text
        reset_password = reset.json()["temporary_password"]
        assert reset_password == "HEAT" and reset.json()["must_change_password"] is True
        client.post("/api/logout")
        login(client, "HR-HEAT", reset_password)
        client.post("/api/logout")

        login(client, "HR01", "HR123")
        employees = client.get("/api/hr/employees").json()
        tagsm = next(row for row in employees if row["employee_no"] == "TAGSMTEST01")
        changed = client.put(f"/api/hr/employees/{tagsm['id']}", json={"role_code": "GSM", "reason": "最高管理员级别调整测试"})
        assert changed.status_code == 200, changed.text
        with SessionLocal() as db:
            gsm_role = db.query(Role).filter_by(code="GSM").one()
            assert db.query(EmployeeRoleAssignment).filter_by(employee_id=tagsm["id"], status="active", role_id=gsm_role.id).count() == 1


def test_v2231_global_grouped_recognizers_exclude_hr_and_all_roles_have_home_password_entry() -> None:
    today = date.today().isoformat()
    with TestClient(app) as client:
        with SessionLocal() as db:
            circles = {
                row.name: row
                for row in db.query(Attraction)
                .filter(Attraction.employee_circle.is_(True), Attraction.active.is_(True))
                .all()
            }
            ta_role = db.query(Role).filter_by(code="TA_SUPERVISOR").one()
            for employee_no, name, circle_name in (
                ("DWARFLEAD01", "矮人测试TALEAD", "矮人迷宫"),
                ("BEARLEAD01", "小熊测试TALEAD", "小熊罐子"),
            ):
                employee = Employee(
                    employee_no=employee_no,
                    name=name,
                    attraction_id=circles[circle_name].id,
                    is_active=True,
                )
                db.add(employee)
                db.flush()
                db.add(
                    EmployeeRoleAssignment(
                        employee_id=employee.id,
                        role_id=ta_role.id,
                        starts_on=today,
                        status="active",
                    )
                )
            tagsm_role = db.query(Role).filter_by(code="TA_GSM").one()
            tagsm = Employee(
                employee_no="TAGSMSELECT01",
                name="认可人列表TAGSM",
                attraction_id=circles["热力追踪"].id,
                is_active=True,
            )
            db.add(tagsm)
            db.flush()
            db.add(
                EmployeeRoleAssignment(
                    employee_id=tagsm.id,
                    role_id=tagsm_role.id,
                    starts_on=today,
                    status="active",
                )
            )
            db.commit()

        login(client, "TATEST01")
        options = client.get("/api/options").json()
        heat_id = next(row["id"] for row in options["employee_circles"] if row["name"] == "热力追踪")
        rows_response = client.get("/api/recognizers", params={"attraction_id": heat_id})
        assert rows_response.status_code == 200, rows_response.text
        rows = [row for row in rows_response.json() if not row.get("special")]

        assert {row["group_label"] for row in rows} == {
            "热力追踪",
            "矮人迷宫",
            "小熊罐子",
            "TAGSM及以上",
        }
        assert {row["role_code"] for row in rows} == {
            "TA_SUPERVISOR",
            "SUPERVISOR",
            "TA_GSM",
            "GSM",
            "AM",
            "OM",
        }
        assert all(row["role_code"] not in {"HR_CIRCLE", "SYSTEM_ADMIN"} for row in rows)
        assert all(
            row["group_label"] == "TAGSM及以上"
            for row in rows
            if row["role_code"] in {"TA_GSM", "GSM", "AM", "OM"}
        )

        cross_circle_recognizer = next(row for row in rows if row["employee_no"] == "DWARFLEAD01")
        cm_id = next(
            row["id"]
            for row in client.get(
                "/api/employee-targets",
                params={"usage": "recognition", "keyword": "CMTEST01"},
            ).json()["items"]
        )
        venue_id = next(row["id"] for row in options["recognition_venues"] if row["name"] == "热力追踪")
        type_id = next(row["id"] for row in options["recognition_types"] if row["code"] == "SAFETY")
        created = client.post(
            "/api/recognitions",
            data={
                "recognition_date": today,
                "occurred_attraction_id": str(venue_id),
                "recognition_type_id": str(type_id),
                "recognizer_employee_id": str(cross_circle_recognizer["id"]),
                "content": "跨景点圈认可人测试",
                "employee_id": str(cm_id),
                "idempotency_key": "v2231-global-recognizer",
            },
        )
        assert created.status_code == 200, created.text
        assert created.json()["record"]["recognizer_name"] == "矮人测试TALEAD"

    script = (Path(__file__).parents[1] / "app" / "static" / "js" / "app.js").read_text(encoding="utf-8")
    assert "const r=state.me.role_code, items=[];" in script
    assert "if (['CM','TR'].includes(r)) items.push(['home','首页'],['register','登记'],['governance','申诉']);" in script
    assert "items.push(['review','复核'],['members','组员记录'],['register','绩效登记'],['absence','缺勤登记'],['entries','主管登记记录']);" in script
    assert "items.push(['hrEmployees','员工管理'],['monthClose','月结'],['circleHrAccounts','景点圈HR账号'],['hrGroups','整组移交']" in script
    assert "async function renderMonthClose" in script
    assert "items.push(['password',has('PASSWORD_RESET')?'密码管理':'修改密码']);" in script
    assert "function renderPasswordPage" in script
    assert "const canCorrectName=['HR_CIRCLE','SYSTEM_ADMIN'].includes(state.me.role_code);" in script
    assert '<h2>账号姓名修改</h2>' in script
    assert 'id="passwordBtn"' not in script
    assert "function recognizerGroups(rows)" in script
    assert "recognitionSubmissionData(e.target,has('SELF_RECOGNITION')&&!has('EMPLOYEE_ADD'))" in script
    assert 'id="hrBatchLeaderSave"' in script
    assert "const canEdit=state.me.role_code==='SYSTEM_ADMIN'" in script
    assert "label:'热力追踪主管'" in script
    assert "label:'矮人迷宫主管'" in script
    assert "label:'小熊罐子主管'" in script
    assert 'data-recognizer-picker' in script
    assert 'recognizer-group-toggle' in script
    assert 'recognizer-group-list' in script
    stylesheet = (Path(__file__).parents[1] / "app" / "static" / "css" / "style.css").read_text(encoding="utf-8")
    assert ".recognizer-menu[hidden]" in stylesheet
    assert ".recognizer-group-list[hidden]" in stylesheet


def test_v2234_imported_recognizers_are_available_from_policy_start() -> None:
    with SessionLocal() as db:
        circle = db.query(Attraction).filter_by(name="热力追踪", employee_circle=True, active=True).one()
        role = db.query(Role).filter_by(code="TA_SUPERVISOR").one()
        employee = Employee(
            employee_no="DATELEAD0801",
            name="生效日TALEAD",
            attraction_id=circle.id,
            is_active=True,
        )
        db.add(employee)
        db.flush()
        recognizer_id = employee.id
        db.add(EmployeeRoleAssignment(employee_id=employee.id, role_id=role.id, starts_on="2026-08-09", status="active"))
        db.commit()

    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        circle_id = next(row["id"] for row in options["employee_circles"] if row["name"] == "热力追踪")
        before = client.get("/api/recognizers", params={"attraction_id": circle_id, "recognition_date": "2026-07-31"})
        assert before.status_code == 200, before.text
        assert not any(not str(item["id"]).startswith("special:") for item in before.json())
        after = client.get("/api/recognizers", params={"attraction_id": circle_id, "recognition_date": "2026-08-01"})
        assert after.status_code == 200, after.text
        row = next(item for item in after.json() if item.get("employee_no") == "DATELEAD0801")
        assert row["role_code"] == "TA_SUPERVISOR"
        assert all(item["role_code"] not in {"HR_CIRCLE", "SYSTEM_ADMIN"} for item in after.json())

        venue_id = next(row["id"] for row in options["recognition_venues"] if row["name"] == "热力追踪")
        type_id = next(row["id"] for row in options["recognition_types"] if row["code"] == "SAFETY")
        cm_id = next(
            row["id"]
            for row in client.get(
                "/api/employee-targets",
                params={"usage": "recognition", "keyword": "CMTEST01"},
            ).json()["items"]
        )
        created = client.post(
            "/api/recognitions",
            data={
                "recognition_date": "2026-08-01",
                "occurred_attraction_id": str(venue_id),
                "recognition_type_id": str(type_id),
                "recognizer_employee_id": str(recognizer_id),
                "content": "认可人生效日回归",
                "employee_id": str(cm_id),
                "idempotency_key": "v2234-recognizer-policy-start",
            },
        )
        assert created.status_code == 200, created.text
        assert created.json()["record"]["recognizer_name"] == "生效日TALEAD"

    script = (Path(__file__).parents[1] / "app" / "static" / "js" / "app.js").read_text(encoding="utf-8")
    assert 'data-evidence-camera' in script
    assert 'data-evidence-album' in script
    assert 'data-evidence-camera-input type="file" accept="image/jpeg,image/png,image/webp,.jpg,.jpeg,.png,.webp" capture="environment"' in script
    assert 'data-evidence-album-input type="file" accept="image/jpeg,image/png,image/webp,.jpg,.jpeg,.png,.webp"' in script
    assert 'data-image-preview' in script
    assert 'openImagePreview' in script
    assert "data.append('image',file,file.name)" in script
    stylesheet = (Path(__file__).parents[1] / "app" / "static" / "css" / "style.css").read_text(encoding="utf-8")
    assert ".image-preview-dialog" in stylesheet
    assert ".image-preview-stage img[hidden]" in stylesheet
    assert 'id="recognitionDate"' in script
    assert "RECOGNIZER_ELIGIBILITY_START" in (Path(__file__).parents[1] / "app" / "v2_services.py").read_text(encoding="utf-8")
    assert "recognizer_role_for_date" in (Path(__file__).parents[1] / "app" / "routers" / "v2.py").read_text(encoding="utf-8")


def test_same_day_recognition_warning_and_editable_hierarchy_export() -> None:
    """A second active same-day recognition needs confirmation and marks both rows."""
    recognition_date = (date.today() + timedelta(days=1)).isoformat()
    month = recognition_date[:7]
    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        me = client.get("/api/me").json()
        venue_id = next(row["id"] for row in options["recognition_venues"] if row["name"] == "热力追踪")
        type_id = next(row["id"] for row in options["recognition_types"] if row["code"] == "SAFETY")
        target = next(
            row for row in client.get(
                "/api/employee-targets",
                params={"usage": "recognition", "keyword": "CMTEST02"},
            ).json()["items"]
            if row["employee_no"] == "CMTEST02"
        )
        base = {
            "recognition_date": recognition_date,
            "occurred_attraction_id": str(venue_id),
            "recognition_type_id": str(type_id),
            "recognizer_employee_id": str(me["id"]),
            "employee_id": str(target["id"]),
        }
        first = client.post("/api/recognitions", data={**base, "content": "同日重复第一项", "idempotency_key": "same-day-first"})
        assert first.status_code == 200, first.text
        second = client.post("/api/recognitions", data={**base, "content": "同日重复第二项", "idempotency_key": "same-day-second"})
        assert second.status_code == 409, second.text
        assert second.json()["detail"]["code"] == "SAME_DAY_RECOGNITION_DUPLICATE"
        confirmed = client.post(
            "/api/recognitions",
            data={**base, "content": "同日重复第二项", "same_day_duplicate_confirmed": "true", "idempotency_key": "same-day-second-confirmed"},
        )
        assert confirmed.status_code == 200, confirmed.text
        entered = client.get("/api/entered-recognitions", params={"month": month})
        matching = [row for row in entered.json() if row["id"] in {first.json()["record"]["id"], confirmed.json()["record"]["id"]}]
        assert sorted(row["same_day_duplicate_sequence"] for row in matching) == [1, 2]
        assert all(row["same_day_duplicate"] for row in matching)

        client.post("/api/logout")
        login(client, "GSMTEST01")
        exported = client.get("/api/statistics/export", params={"month": month})
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content), data_only=False)
        hierarchy = workbook["层级绩效明细"]
        assert workbook.active.title == "层级绩效明细"
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
        detail_rows = [row for row in hierarchy.iter_rows(min_row=2) if row[13].value == "明细"]
        assert detail_rows
        assert any(row[5].value for row in hierarchy.iter_rows(min_row=2) if row[13].value == "汇总")
        assert any(row[5].data_type == "f" and "SUMIFS" in row[5].value for row in hierarchy.iter_rows(min_row=2) if row[13].value == "汇总")
        assert all(sheet.protection.sheet is False and sheet.protection.objects is False for sheet in workbook.worksheets)
        details = workbook["签卡明细"]
        duplicate_column = [cell.value for cell in details[1]].index("同日重复登记") + 1
        matched_rows = [row for row in details.iter_rows(min_row=2) if row[0].value in {f"REC-{first.json()['record']['id']}", f"REC-{confirmed.json()['record']['id']}"}]
        assert len(matched_rows) == 2
        assert all(str(row[duplicate_column - 1].value).startswith("是（第") for row in matched_rows)
        assert all(row[0].fill.fgColor.rgb in {"00FFF2CC", "FFF2CC"} for row in matched_rows)

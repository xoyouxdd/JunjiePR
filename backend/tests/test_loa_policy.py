from __future__ import annotations

from datetime import date
from io import BytesIO
import os
from pathlib import Path
import tempfile


TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-loa-policy-test-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402


def login(client: TestClient, employee_no: str) -> None:
    response = client.post("/api/login", json={"employee_no": employee_no, "password": "1234"})
    assert response.status_code == 200, response.text


def test_loa_any_touched_month_is_excluded_and_target_search_is_global() -> None:
    with TestClient(app) as client:
        login(client, "GSMTEST01")
        targets = client.get("/api/employee-targets", params={"usage": "loa", "keyword": "测试CM甲"})
        assert targets.status_code == 200, targets.text
        target = next(row for row in targets.json()["items"] if row["employee_no"] == "CMTEST01")
        reset_targets = client.get("/api/accounts/reset-targets", params={"keyword": "测试CM甲"})
        assert reset_targets.status_code == 200, reset_targets.text
        assert any(row["id"] == target["id"] for row in reset_targets.json()["items"])
        created = client.post(
            "/api/loa-periods",
            json={"employee_id": target["id"], "starts_on": "2099-08-31", "ends_on": "2099-09-01", "note": "跨月LOA"},
        )
        assert created.status_code == 200, created.text
        assert created.json()["excluded_months"] == ["2099-08", "2099-09"]
        for month in ("2099-08", "2099-09"):
            payload = client.get("/api/statistics", params={"month": month}).json()
            row = next(item for item in payload["loa_rows"] if item["employee_id"] == target["id"])
            assert row["total_score"] == 0
            assert row["recognition_score"] == 0
            assert row["employment_status"] == "LOA（当月不参与计分）"
        exported = client.get("/api/statistics/export", params={"month": "2099-09"})
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content))
        assert "LOA明细" in workbook.sheetnames
        loa_sheet = workbook["LOA明细"]
        row = next(item for item in loa_sheet.iter_rows(min_row=2) if item[0].value == "CMTEST01")
        assert row[7].value == "当月不参与计分"
        assert row[4].fill.fgColor.rgb.endswith("D9D9D9")

        second = next(row for row in client.get("/api/employee-targets", params={"usage": "loa", "keyword": "测试CM乙"}).json()["items"] if row["employee_no"] == "CMTEST02")
        entered = client.post("/api/loa-periods", json={"employee_id": second["id"], "starts_on": "2099-10-05"})
        assert entered.status_code == 200, entered.text
        assert entered.json()["completed"] is False
        marked = next(row for row in client.get("/api/employee-targets", params={"usage": "loa", "keyword": "测试CM乙"}).json()["items"] if row["id"] == second["id"])
        assert marked["loa_active"] is True
        assert marked["loa_starts_on"] == "2099-10-05"
        completed = client.post("/api/loa-periods", json={"employee_id": second["id"], "starts_on": "2099-10-05", "ends_on": "2099-10-20"})
        assert completed.status_code == 200, completed.text
        assert completed.json()["completed"] is True
        scheduled = next(row for row in client.get("/api/employee-targets", params={"usage": "loa", "keyword": "测试CM乙"}).json()["items"] if row["id"] == second["id"])
        assert scheduled["loa_active"] is True
        assert scheduled["loa_starts_on"] == "2099-10-05"
        assert scheduled["loa_ends_on"] == "2099-10-20"

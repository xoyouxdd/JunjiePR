"""轮岗接口端到端：名单上传、草稿、发布、测试时钟、大屏操作、个人待办、记录与配置。"""
from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import tempfile

from fastapi.testclient import TestClient
from openpyxl import Workbook


TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-rotation-flow-test-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from app.main import app  # noqa: E402


DAY = "2026-10-04"
LINES = [
    {"id": "A", "group": "A", "posts": [{"name": "A入口"}, {"name": "迎宾", "seven": 1}, {"name": "A出口"}]},
    {"id": "B", "group": "B", "posts": [{"name": "B入口"}, {"name": "B中"}, {"name": "B出口"}]},
]
ROSTER = [
    ("9000001", "早班甲", "07:00-13:30", ""),
    ("9000002", "开园甲", "07:15-16:15", ""),
    ("9000003", "开园乙", "07:15-16:15", ""),
    ("9000004", "开园丙", "07:15-16:15", ""),
    ("9000005", "开园丁", "07:15-16:15", ""),
    ("9000006", "开园戊", "07:15-16:15", ""),
    ("9000007", "推七甲", "07:15-17:15", ""),
    ("9000008", "中班甲", "09:00-18:00", "1000-1100 SSEI"),
    ("9000009", "放休甲", "放休", ""),
]


def roster_file(day=DAY) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.append(["工号", "姓名", "日期", "班次", "备注"])
    for no, name, shift, note in ROSTER:
        ws.append([no, name, day, shift, note])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def login(client: TestClient, account: str, password: str = "1234") -> None:
    response = client.post("/api/login", json={"employee_no": account, "password": password})
    assert response.status_code == 200, response.text


def enter(client: TestClient, account: str) -> None:
    response = client.post("/api/rotation/enter", json={"account": account})
    assert response.status_code == 200, response.text


def ok(response):
    assert response.status_code == 200, response.text
    return response.json()


def clock(client: TestClient, **body):
    return client.post("/api/rotation/test-clock", json=body)


def person(board: dict, no: str) -> dict:
    return next(p for p in board["day"]["persons"] if p["pid"] == no)


def test_future_opening_waits_for_departure_window_in_every_api():
    day = '2026-10-08'
    wb = Workbook()
    wb.active.append(['工号', '姓名', '日期', '班次', '备注'])
    wb.active.append(['9100001', '定时开岗员工', day, '07:15-16:15', ''])
    buf = BytesIO()
    wb.save(buf)
    with TestClient(app) as manager, TestClient(app) as screen, TestClient(app) as member:
        login(manager, 'GSMTEST01')
        enter(manager, '7777777')
        login(screen, '6666666', '1243')
        login(member, 'TRTEST01')
        ok(clock(manager, action='set', date=day, time='03:59'))
        ok(manager.put('/api/rotation/config', json={'lines': [{'id': 'C2', 'group': 'C',
            'posts': [{'name': '单人迎宾', 'openAt': '08:15'}]}], 'settings': {'departEarly': 1}}))
        ok(manager.post('/api/rotation/roster/upload', files={'file': ('名单.xlsx', buf.getvalue(),
            'application/octet-stream')}, data={'scope': 'day'}))
        enter(member, '9100001')
        ok(clock(manager, action='jump', minutes=240))
        ok(clock(manager, action='jump', minutes=6))
        for client, path in ((manager, '/api/rotation/board'), (screen, '/api/rotation/screen/state')):
            p = person(ok(client.get(path)), '9100001')
            assert p['assign']['departAt'] == 495 and p['readyAt'] <= 485
            assert p['canDepart'] is False
        me = ok(member.get('/api/rotation/me'))
        assert me['settings']['departEarly'] == 1 and me['person']['canDepart'] is False
        assert any(n['kind'] == 'rest' and '08:15' in n['title'] for n in me['notices'] if n['status'] == 'open')
        for client, path in ((manager, '/api/rotation/act'), (screen, '/api/rotation/screen/act')):
            response = client.post(path, json={'action': 'depart', 'pid': '9100001'})
            assert response.status_code == 400 and '还没到出发时间' in response.text
        ok(clock(manager, action='jump', minutes=9))
        assert person(ok(manager.get('/api/rotation/board')), '9100001')['canDepart'] is True
        ok(manager.put('/api/rotation/config', json={'settings': {'departEarly': 0}}))
        assert person(ok(screen.get('/api/rotation/screen/state')), '9100001')['canDepart'] is False
        assert ok(member.get('/api/rotation/me'))['settings']['departEarly'] == 0
        ok(clock(manager, action='jump', minutes=1))
        assert person(ok(manager.get('/api/rotation/board')), '9100001')['canDepart'] is True
        ok(screen.post('/api/rotation/screen/act', json={'action': 'depart', 'pid': '9100001'}))
        assert person(ok(manager.get('/api/rotation/board')), '9100001')['state'] == 'onpost'


def test_manual_opening_default_edit_cancel_and_screen_departure():
    day = '2026-10-09'
    wb = Workbook()
    wb.active.append(['工号', '姓名', '日期', '班次', '备注'])
    wb.active.append(['9200001', '加岗演示', day, '07:15-16:15', ''])
    buf = BytesIO()
    wb.save(buf)
    with TestClient(app) as manager, TestClient(app) as screen:
        login(manager, 'GSMTEST01')
        enter(manager, '7777777')
        login(screen, '6666666', '1243')
        ok(clock(manager, action='set', date=day, time='03:59'))
        ok(manager.put('/api/rotation/config', json={'lines': [{'id': 'C1', 'group': 'C',
            'posts': [{'name': '单人队末', 'openAt': '09:00'}]}], 'settings': {'departEarly': 1}}))
        ok(manager.post('/api/rotation/roster/upload', files={'file': ('名单.xlsx', buf.getvalue(),
            'application/octet-stream')}, data={'scope': 'day'}))
        ok(clock(manager, action='jump', minutes=240))
        ok(clock(manager, action='jump', minutes=6))
        request = {'action': 'post', 'line': 'C1', 'i': 0, 'open': True}
        assert ok(manager.get('/api/rotation/board'))['day']['status'] == 'live'
        ok(manager.post('/api/rotation/act', json=request))
        board = ok(manager.get('/api/rotation/board'))
        assert board['day']['lines'][0]['posts'][0]['openAt'] == '08:08'
        assert person(board, '9200001')['assign']['mode'] == 'opening'
        ok(manager.post('/api/rotation/act', json={**request, 'openAt': '08:12'}))
        assert person(ok(manager.get('/api/rotation/board')), '9200001')['assign']['departAt'] == 492
        for value in ('08:04', '24:00', 'no-time'):
            assert manager.post('/api/rotation/act', json={**request, 'openAt': value}).status_code == 400
        assert person(ok(manager.get('/api/rotation/board')), '9200001')['assign']['departAt'] == 492
        ok(manager.post('/api/rotation/act', json={**request, 'open': False}))
        assert person(ok(manager.get('/api/rotation/board')), '9200001')['assign'] is None
        ok(manager.post('/api/rotation/act', json=request))
        ok(clock(manager, action='jump', minutes=2))
        assert person(ok(screen.get('/api/rotation/screen/state')), '9200001')['canDepart'] is False
        assert screen.post('/api/rotation/screen/act', json={'action': 'depart', 'pid': '9200001'}).status_code == 400
        next_step = ok(clock(manager, action='next'))
        assert next_step['clock']['minute'] == 488 and next_step['waiting']
        ok(screen.post('/api/rotation/screen/act', json={'action': 'depart', 'pid': '9200001'}))
        assert person(ok(manager.get('/api/rotation/board')), '9200001')['state'] == 'onpost'
        assert manager.post('/api/rotation/act', json={**request, 'openAt': '08:12'}).status_code == 400


def test_full_day_flow_through_the_api() -> None:
    with TestClient(app) as manager, TestClient(app) as screen, TestClient(app) as member:
        login(manager, "GSMTEST01")
        enter(manager, "7777777")
        login(screen, "6666666", "1243")
        login(member, "TRTEST01")

        # 配置两条线，上传名单
        ok(manager.put("/api/rotation/config", json={"lines": LINES}))
        assert manager.put("/api/rotation/config", json={"lines": [{"id": "X", "posts": []}]}).status_code == 400
        assert manager.put("/api/rotation/config", json={"settings": {"offLead": "abc"}}).status_code == 400
        bad = manager.post("/api/rotation/roster/upload", files={"file": ("名单.txt", b"x", "text/plain")}, data={"scope": "week"})
        assert bad.status_code == 400
        uploaded = ok(manager.post("/api/rotation/roster/upload", files={"file": ("名单.xlsx", roster_file(), "application/octet-stream")}, data={"scope": "day"}))
        assert uploaded["entry_count"] == len(ROSTER) and uploaded["start_date"] == DAY
        summary = ok(manager.get("/api/rotation/roster", params={"date": DAY}))
        assert summary["available"] and {p["kind"] for p in summary["people"]} == {"rotation", "off"}

        # 模拟一个 CM 工号：名单有了才能进
        enter(member, "9000001")

        # 草稿只能在模拟时钟到这一天后发布
        ok(clock(manager, action="set", date=DAY, time="06:30"))
        ok(manager.post("/api/rotation/draft", json={"action": "draft_generate", "date": DAY}))
        board = ok(manager.get("/api/rotation/board"))
        assert board["draft"]["plan"]["seven"] == {"A#1": "9000001"}
        assert board["draft"]["plan"]["push7"] == ["9000007"]
        assert manager.post("/api/rotation/draft", json={"action": "publish", "date": "2026-10-05"}).status_code == 400
        ok(manager.post("/api/rotation/draft", json={"action": "publish", "date": DAY}))
        assert manager.post("/api/rotation/draft", json={"action": "draft_generate", "date": DAY}).status_code == 400

        # 下一步：推进到第一个需要在大屏操作的时刻
        step = ok(clock(manager, action="next"))
        assert step["clock"]["paused"] is True
        assert step["waiting"] == ["推七甲 点「去轮岗」（A 线）"]
        assert step["clock"]["minute"] == 7 * 60 + 15
        early = person(ok(manager.get("/api/rotation/board")), "9000001")
        assert early["state"] == "onpost" and early["preparing"] == "推7点"
        assert "7点岗" in early["flags"] and early["lineStart"] == 7 * 60
        assert clock(manager, action="next").status_code == 400  # 还没处理，不能下一步

        # 大屏只能点去休息、去轮岗
        assert screen.post("/api/rotation/screen/act", json={"action": "undo", "pid": "9000007"}).status_code == 403
        assert screen.post("/api/rotation/act", json={"action": "depart", "pid": "9000007"}).status_code == 403
        ok(screen.post("/api/rotation/screen/act", json={"action": "depart", "pid": "9000007"}))
        assert screen.post("/api/rotation/screen/act", json={"action": "depart", "pid": "9000007"}).status_code == 400

        assert clock(manager, action="next").status_code == 400  # 点击即替下，已有去休息待办
        ok(clock(manager, action="jump", minutes=3))
        board = ok(manager.get("/api/rotation/board"))
        assert person(board, "9000001")["downReason"] == "推7点下来"
        assert person(board, "9000007")["post"] == "迎宾"

        # 员工个人页：待办标测试，随状态更新
        me = ok(member.get("/api/rotation/me"))
        assert me["test_mode"] is True and me["person"]["state"] == "walkback"
        assert [n["title"] for n in me["notices"] if n["status"] == "open"] == ["请到休息室大屏点「去休息」"]
        ok(screen.post("/api/rotation/screen/act", json={"action": "arrive", "pid": "9000001"}))
        me = ok(member.get("/api/rotation/me"))
        open_titles = [n["title"] for n in me["notices"] if n["status"] == "open"]
        assert open_titles == ["休息至 07:33"]
        assert any(n["status"] == "done" and "去休息" in n["title"] for n in me["notices"])
        assert me["line"] is not None or me["person"]["assign"] is None

        # 主管撤回（大屏不能撤回），再到达时休息仍从第一次到达算起
        ok(manager.post("/api/rotation/act", json={"action": "undo", "pid": "9000001"}))
        ok(screen.post("/api/rotation/screen/act", json={"action": "arrive", "pid": "9000001"}))
        assert person(ok(manager.get("/api/rotation/board")), "9000001")["readyAt"] == 7 * 60 + 33
        assert manager.post("/api/rotation/act", json={"action": "fix_remove", "pid": "9000002"}).status_code == 400

        # 快进与倍速
        before = ok(manager.get("/api/rotation/board"))["clock"]["minute"]
        jumped = ok(clock(manager, action="jump", minutes=30))
        assert jumped["clock"]["minute"] == before + 30
        assert clock(manager, action="speed", speed=500).status_code == 400
        ok(clock(manager, action="speed", speed=30))

        # 记录
        log = ok(manager.get("/api/rotation/log", params={"date": DAY}))["items"]
        types = {item["type"] for item in log}
        assert {"publish", "crew_start", "push7", "depart", "arrive", "undo_arrive"} <= types
        assert any(item["actor"] == "热力追踪休息室大屏" for item in log)
        record = ok(manager.get("/api/rotation/person", params={"employee_no": "9000001", "date": DAY}))
        assert record["segments"][0]["line"] == "A" and record["duties"] == []
        assert ok(manager.get("/api/rotation/person", params={"employee_no": "9000007", "date": DAY}))["duties"] == [{"date": DAY, "kind": "push7"}]

        # 权限：大屏、模拟员工都不能看主管看板
        assert screen.get("/api/rotation/board").status_code == 403
        assert member.get("/api/rotation/board").status_code == 403
        assert member.post("/api/rotation/test-clock", json={"action": "jump"}).status_code == 403
        assert ok(screen.get("/api/rotation/screen/state"))["day"]["status"] == "live"

        ok(clock(manager, action="real"))


def test_automatic_four_am_prearrangement_and_full_test_reset(monkeypatch):
    from app.rotation import service
    original_action = service.do_live_action

    def tick_before_drag(db, attraction_id, actor, body):
        # Force the previously flaky scheduling: the background ticker runs
        # after capturing the original plan, but before the first manual edit.
        if body.get('action') == 'plan_place':
            service.tick_attraction(attraction_id)
        return original_action(db, attraction_id, actor, body)

    monkeypatch.setattr(service, 'do_live_action', tick_before_drag)
    day = '2026-10-06'
    with TestClient(app) as manager, TestClient(app) as screen:
        login(manager, 'GSMTEST01')
        enter(manager, '8888888')
        login(screen, '6666666', '1243')
        ok(manager.put('/api/rotation/config', json={'lines': LINES}))
        ok(manager.post('/api/rotation/roster/upload', files={'file': ('名单.xlsx', roster_file(day), 'application/octet-stream')}, data={'scope':'day'}))
        ok(clock(manager, action='set', date=day, time='03:59'))
        assert ok(manager.get('/api/rotation/board'))['day'] is None
        ok(clock(manager, action='jump', minutes=1))
        board = ok(manager.get('/api/rotation/board'))
        assert board['day']['status'] == 'live'
        original_plan = board['day']['plan']
        from app.v2_database import SessionLocal
        from app.rotation.models import RotationDay
        with SessionLocal() as db:
            previous = db.query(RotationDay).filter(RotationDay.work_date == DAY).first()
            if previous:
                assert previous.status == 'ended'
        assert person(board, '9000001')['lineStart'] is None
        assert board['day']['lines'][0]['posts'][1]['occ'] == '9000001'
        ok(manager.post('/api/rotation/act', json={'action':'plan_place','pid':'9000001','line':'A','i':0}))
        ok(manager.post('/api/rotation/act', json={'action':'post','line':'A','i':2,'open':False}))
        assert screen.post('/api/rotation/test-clock', json={'action':'reset'}).status_code == 403
        ok(clock(manager, action='reset'))
        reset = ok(manager.get('/api/rotation/board'))
        assert reset['clock']['minute'] == 240 and reset['clock']['paused']
        assert reset['day']['plan'] == original_plan
        assert reset['day']['lines'][0]['posts'][2]['open']
        assert all(p['lineStart'] is None for p in reset['day']['persons'])
        log = ok(manager.get('/api/rotation/log', params={'date':day}))['items']
        assert [entry['type'] for entry in log] == ['publish']
        ok(clock(manager, action='jump', minutes=180))
        started = ok(manager.get('/api/rotation/board'))
        assert person(started, '9000001')['lineStart'] == 420
        assert person(started, '9000002')['lineStart'] is None
        ok(clock(manager, action='jump', minutes=15))
        ok(screen.post('/api/rotation/screen/act', json={'action':'depart','pid':original_plan['push7'][0]}))
        before_reset = ok(manager.get('/api/rotation/person', params={'employee_no':'9000001','date':day}))
        assert any(segment['date'] == day for segment in before_reset['segments'])
        ok(clock(manager, action='reset'))
        assert ok(manager.get('/api/rotation/board'))['day']['plan'] == original_plan
        after_reset = ok(manager.get('/api/rotation/person', params={'employee_no':'9000001','date':day}))
        assert not any(segment['date'] == day for segment in after_reset['segments'])
        assert after_reset['segments'] == [segment for segment in before_reset['segments'] if segment['date'] != day]


def test_reset_restores_persisted_snapshot_not_new_week_or_configuration(monkeypatch):
    from app.rotation import service
    from app.rotation.models import RotationSegment
    from app.v2_database import SessionLocal
    from app.v2_models import Attraction

    day = '2026-10-07'
    with TestClient(app) as manager:
        login(manager, 'GSMTEST01')
        enter(manager, '8888888')
        ok(clock(manager, action='set', date=day, time='03:59'))
        ok(manager.put('/api/rotation/config', json={'lines': LINES}))
        ok(manager.post('/api/rotation/roster/upload', files={'file': ('名单.xlsx', roster_file(day), 'application/octet-stream')}, data={'scope':'day'}))
        ok(clock(manager, action='jump', minutes=1))
        original_plan = ok(manager.get('/api/rotation/board'))['day']['plan']
        with SessionLocal() as db:
            attraction_id = db.query(Attraction).filter_by(name='热力追踪').one().id
            db.add(RotationSegment(attraction_id=attraction_id, work_date='2026-10-05',
                employee_no='9000007', line='A', start_min=420,
                end_min=1020, minutes=600))
            db.commit()
        ok(manager.put('/api/rotation/config', json={'lines': list(reversed(LINES))}))
        service.RUNTIME.cache.clear()  # Prove recovery uses the database snapshot.

        def no_recalculation(*args, **kwargs):
            raise AssertionError('reset must not recalculate the initial plan')

        monkeypatch.setattr(service.E, 'build_draft', no_recalculation)
        ok(clock(manager, action='reset'))
        first = ok(manager.get('/api/rotation/board'))
        assert first['day']['plan'] == original_plan
        assert [line['id'] for line in first['day']['lines']] == ['A', 'B']
        assert [line['id'] for line in first['lines']] == ['B', 'A']
        ok(clock(manager, action='reset'))
        assert ok(manager.get('/api/rotation/board'))['day']['plan'] == original_plan
        with SessionLocal() as db:
            assert db.query(RotationSegment).filter_by(attraction_id=attraction_id,
                work_date='2026-10-05', employee_no='9000007', minutes=600).count() == 1

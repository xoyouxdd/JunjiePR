from __future__ import annotations
from frontend_source import read_frontend_source

import os
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import pytest
from types import SimpleNamespace
from openpyxl import load_workbook


TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-login-archive-test-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "isolated-test-only"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "isolated-admin-only"

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.v2_database import SessionLocal  # noqa: E402
from app.v2_models import AuditLog, Employee, GroupLeaderAssignment, MonthClosure, RecognitionRecord, RecognitionType, UserAccount, WorkGroup  # noqa: E402


ROOT = Path(__file__).parents[1]


def login(client: TestClient, employee_no: str, password: str) -> None:
    response = client.post("/api/login", json={"employee_no": employee_no, "password": password})
    assert response.status_code == 200, response.text


def organization_employee_ids(client: TestClient) -> set[int]:
    response = client.get("/api/hr/organization")
    assert response.status_code == 200, response.text
    organization = response.json()
    employees = list(organization["others"])
    for circle in organization["circles"]:
        for key in ("supervisors", "unassigned", "managers", "inactive"):
            employees.extend(circle[key])
        for group in circle["groups"]:
            employees.extend(group["members"])
    return {employee["id"] for employee in employees}


def mark_eligible_with_business_history() -> tuple[int, int]:
    """Create a retained recognition row and an account eligible for removal."""
    with SessionLocal() as db:
        employee = db.query(Employee).filter_by(employee_no="CMTEST01").one()
        account = db.query(UserAccount).filter_by(employee_id=employee.id).one()
        admin = db.query(Employee).filter_by(employee_no="HR01").one()
        recognition_type = db.query(RecognitionType).first()
        assert recognition_type is not None
        employee.is_active = False
        employee.terminated_on = (date.today() - timedelta(days=8)).isoformat()
        employee.updated_at = datetime.now() - timedelta(days=8)
        account.enabled = False
        account.disabled_at = datetime.now() - timedelta(days=3, seconds=1)
        db.add(
            RecognitionRecord(
                employee_id=employee.id,
                employee_no=employee.employee_no,
                employee_name=employee.name,
                employee_role_snapshot="CM",
                employee_role_code_snapshot="CM",
                home_attraction_id=employee.attraction_id,
                home_attraction_name=employee.attraction.name,
                occurred_attraction_id=employee.attraction_id,
                recognition_date=date.today().isoformat(),
                recognition_month=date.today().strftime("%Y-%m"),
                recognition_type_id=recognition_type.id,
                recognition_type_name=recognition_type.name,
                content="隔离留档测试",
                recognizer_employee_id=admin.id,
                recognizer_name=admin.name,
                recognizer_role_snapshot="最高管理员",
                recognizer_role_code_snapshot="SYSTEM_ADMIN",
                operator_employee_id=admin.id,
                operator_name=admin.name,
                operator_role_snapshot="最高管理员",
                operator_role_code_snapshot="SYSTEM_ADMIN",
                source="admin",
                fraction=Decimal("0.50"),
                credited_fraction=Decimal("0.50"),
                status="confirmed",
            )
        )
        db.commit()
        return employee.id, account.id


def test_remove_login_after_three_days_keeps_employee_and_business_archive() -> None:
    with TestClient(app) as client:
        employee_id, account_id = mark_eligible_with_business_history()
        login(client, "HR01", "isolated-admin-only")
        directory = client.get("/api/hr/employees")
        assert directory.status_code == 200, directory.text
        item = next(row for row in directory.json() if row["id"] == employee_id)
        assert item["account_deletion_eligible"] is True
        assert employee_id in organization_employee_ids(client)
        removed = client.request(
            "DELETE",
            f"/api/hr/employees/{employee_id}/account",
            json={"reason": "隔离测试：删除登录账号，保留业务档案"},
        )
        assert removed.status_code == 200, removed.text
        assert removed.json()["archive_retained"] is True
        with SessionLocal() as db:
            employee = db.get(Employee, employee_id)
            assert employee is not None
            assert employee.account_deleted_at is not None
            archived_at = employee.account_deleted_at.strftime("%Y-%m-%d %H:%M:%S")
            assert db.get(UserAccount, account_id) is None
            assert db.query(RecognitionRecord).filter_by(employee_id=employee_id).count() == 1
            assert db.query(AuditLog).filter_by(action="删除人员（留档）", entity_id=str(employee_id)).count() == 1
        rejected_login = client.post("/api/login", json={"employee_no": "CMTEST01", "password": "isolated-test-only"})
        assert rejected_login.status_code == 401
        refreshed = client.get("/api/hr/employees")
        item = next(row for row in refreshed.json() if row["id"] == employee_id)
        assert item["account_deleted_at"] == archived_at
        assert item["account_deletion_eligible"] is False
        assert "已删除·留档" in item["account_deletion_reason"]
        assert employee_id not in organization_employee_ids(client)
        exported = client.get('/api/statistics/export', params={'month':date.today().strftime('%Y-%m')})
        assert exported.status_code == 200, exported.text
        book = load_workbook(BytesIO(exported.content))
        assert any('隔离留档测试' in str(cell.value) for row in book['签卡明细'] for cell in row)
        assert any('已删除·留档' == cell.value for row in book['员工留档标记'] for cell in row)
        assert client.put(f'/api/hr/employees/{employee_id}', json={'employment_status':'active','account_enabled':True}).status_code == 409


def test_delete_endpoint_enforces_the_full_three_day_wait() -> None:
    with TestClient(app) as client:
        with SessionLocal() as db:
            employee = db.query(Employee).filter_by(employee_no="CMTEST02").one()
            employee.is_active = False
            employee.terminated_on = (date.today() - timedelta(days=6)).isoformat()
            account = db.query(UserAccount).filter_by(employee_id=employee.id).one()
            account.enabled = False
            account.disabled_at = datetime.now() - timedelta(days=2)
            db.commit()
            employee_id = employee.id
        login(client, "HR01", "isolated-admin-only")
        response = client.request("DELETE", f"/api/hr/employees/{employee_id}/account", json={})
        assert response.status_code == 400
        assert "起可删除" in response.json()["detail"]
        assert employee_id in organization_employee_ids(client)
        with SessionLocal() as db:
            assert db.query(UserAccount).filter_by(employee_id=employee_id).count() == 1


def test_legacy_active_employee_with_archived_login_is_hidden_from_management_tree() -> None:
    with TestClient(app) as client:
        with SessionLocal() as db:
            employee = db.query(Employee).filter_by(employee_no="CMTEST02").one()
            employee.is_active = True
            employee.account_deleted_at = datetime.now()
            db.commit()
            employee_id = employee.id
        login(client, "HR01", "isolated-admin-only")
        assert employee_id not in organization_employee_ids(client)


def test_hr_archive_ui_and_export_marker_are_present() -> None:
    script = read_frontend_source()
    router = "\n".join(_p.read_text(encoding="utf-8") for _p in sorted((ROOT / "app" / "routers").glob("*.py")))
    assert "data-delete-login-account" in script
    assert "确认删除" in script
    assert "已删除·留档" in script
    assert '"/hr/employees/{employee_id}/account"' in router
    assert "business_history_retained" in router
    # The archive timestamp is verified through the directory API above;
    # its projection may live outside the HTTP routing module.


def test_eligibility_requires_full_continuous_disabled_duration(monkeypatch):
    from app import organization_queries
    instant = datetime(2026,10,17,12,0,0)
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return instant
    monkeypatch.setattr(organization_queries,'datetime',Clock)
    employee=SimpleNamespace(account_deleted_at=None,is_active=False,terminated_on='2020-01-01')
    account=SimpleNamespace(enabled=True,disabled_at=instant-timedelta(days=100))
    role=SimpleNamespace(code='CM')
    state=lambda:organization_queries.login_account_archive_state(employee,account,role)
    assert state()['eligible'] is False  # Even a long-ago departure cannot bypass disabling.
    account.enabled=False;account.disabled_at=instant-timedelta(days=3)+timedelta(microseconds=1)
    assert state()['eligible'] is False
    account.disabled_at=instant-timedelta(days=3)
    assert state()['eligible'] is True
    account.disabled_at=None
    assert state()['eligible'] is False


def test_active_employee_disable_reenable_resets_wait_and_delete_archives_profile():
    with TestClient(app) as client:
        login(client,'HR01','isolated-admin-only')
        circle=client.get('/api/options').json()['employee_circles'][0]['id']
        created=client.post('/api/hr/employees',json={'employee_no':'9817011','name':'停用删除合成测试','role_code':'CM','attraction_id':circle})
        assert created.status_code==200,created.text
        target=created.json()['employee']['id']
        update=lambda body:client.put(f'/api/hr/employees/{target}',json=body)
        remove=lambda:client.request('DELETE',f'/api/hr/employees/{target}/account',json={})
        assert update({'employment_status':'terminated'}).status_code==400
        assert update({'is_active':False}).status_code==400
        assert update({'account_enabled':False}).status_code==200
        assert remove().status_code==400
        with SessionLocal() as db:
            account=db.query(UserAccount).filter_by(employee_id=target).one()
            account.disabled_at=datetime.now()-timedelta(days=3, seconds=1);db.commit()
        assert update({'account_enabled':True}).status_code==200
        with SessionLocal() as db:
            assert db.query(UserAccount).filter_by(employee_id=target).one().disabled_at is None
        assert remove().status_code==400
        assert update({'account_enabled':False}).status_code==200
        assert remove().status_code==400
        with SessionLocal() as db:
            account=db.query(UserAccount).filter_by(employee_id=target).one()
            account.disabled_at=datetime.now()-timedelta(days=3, seconds=1);db.commit()
        result=remove()
        assert result.status_code==200,result.text
        with SessionLocal() as db:
            employee=db.get(Employee,target)
            assert employee.is_active is False and employee.account_deleted_at
            assert employee.terminated_on==date.today().isoformat()
            assert db.query(UserAccount).filter_by(employee_id=target).count()==0
        assert target not in organization_employee_ids(client)


@pytest.mark.parametrize('blocker',['leader','closed_month'])
def test_archive_failures_keep_account_person_and_history_atomic(blocker):
    with TestClient(app) as client:
        login(client,'HR01','isolated-admin-only')
        circle=client.get('/api/options').json()['employee_circles'][0]['id']
        created=client.post('/api/hr/employees',json={'employee_no':'9817012' if blocker=='leader' else '9817013',
            'name':'删除约束测试','role_code':'SUPERVISOR' if blocker=='leader' else 'CM','attraction_id':circle})
        assert created.status_code==200,created.text
        target=created.json()['employee']['id']
        with SessionLocal() as db:
            account=db.query(UserAccount).filter_by(employee_id=target).one()
            account.enabled=False;account.disabled_at=datetime.now()-timedelta(days=3, seconds=1)
            if blocker=='leader':
                group=WorkGroup(name='删除约束合成小组',code='DELETE',attraction_id=circle)
                db.add(group);db.flush()
                db.add(GroupLeaderAssignment(group_id=group.id,leader_employee_id=target,
                    leader_type='formal',starts_on=date.today().isoformat()))
            else:
                db.add(MonthClosure(closure_month=date.today().strftime('%Y-%m'),attraction_id=circle,status='closed'))
            db.commit()
        result=client.request('DELETE',f'/api/hr/employees/{target}/account',json={})
        assert result.status_code==(400 if blocker=='leader' else 423),result.text
        with SessionLocal() as db:
            employee=db.get(Employee,target)
            assert employee.is_active and employee.account_deleted_at is None
            assert db.query(UserAccount).filter_by(employee_id=target).count()==1
            assert db.query(AuditLog).filter_by(action='删除人员（留档）',entity_id=str(target)).count()==0
            if blocker=='closed_month':
                db.query(MonthClosure).filter_by(closure_month=date.today().strftime('%Y-%m'),attraction_id=circle).delete()
                db.commit()


def test_scoped_hr_cannot_delete_other_circle_or_privileged_accounts():
    from app.v2_crypto import hash_password
    with TestClient(app) as client:
        login(client,'HR01','isolated-admin-only')
        with SessionLocal() as db:
            hr=db.query(Employee).filter_by(employee_no='HR-HEAT').one()
            hr_account=db.query(UserAccount).filter_by(employee_id=hr.id).one()
            hr_account.password_hash=hash_password('isolated-circle-test')
            hr_account.enabled=True;hr_account.must_change_password=False
            admin_id=db.query(Employee).filter_by(employee_no='HR01').one().id
            hr_id=hr.id;own_circle=hr.attraction_id;db.commit()
        other_circle=next(r['id'] for r in client.get('/api/options').json()['employee_circles'] if r['id']!=own_circle)
        created=client.post('/api/hr/employees',json={'employee_no':'9817014','name':'跨圈删除合成测试','role_code':'CM','attraction_id':other_circle})
        assert created.status_code==200,created.text
        target=created.json()['employee']['id']
        with SessionLocal() as db:
            account=db.query(UserAccount).filter_by(employee_id=target).one()
            account.enabled=False;account.disabled_at=datetime.now()-timedelta(days=3, seconds=1);db.commit()
        client.post('/api/logout');login(client,'HR-HEAT','isolated-circle-test')
        for denied in (target,admin_id,hr_id):
            assert client.request('DELETE',f'/api/hr/employees/{denied}/account',json={}).status_code==403
        with SessionLocal() as db:
            assert db.query(UserAccount).filter_by(employee_id=target).count()==1
            assert db.get(Employee,target).account_deleted_at is None


@pytest.mark.parametrize('legacy_archive', [False, True])
def test_deleted_employee_can_register_original_number_without_losing_history(legacy_archive):
    from app.services.employment import employment_history
    from app.services.identity import employee_active_on, role_at
    from app.v2_models import EmployeeActingDuty, GroupMembership, ManagementScope, Role
    from app.excel_export import write_employee_number_history_sheet
    from openpyxl import Workbook
    number = '9817022' if legacy_archive else '9817021'
    today = date.today()
    earlier = (today - timedelta(days=20)).isoformat()
    departed = (today - timedelta(days=8)).isoformat()
    with TestClient(app) as client:
        login(client, 'HR01', 'isolated-admin-only')
        circles = client.get('/api/options').json()['employee_circles']
        body = {'employee_no': number, 'name': '重新入职合成测试', 'role_code': 'CM', 'attraction_id': circles[0]['id']}
        created = client.post('/api/hr/employees', json=body)
        assert created.status_code == 200, created.text
        target = created.json()['employee']['id']
        with SessionLocal() as db:
            employee = db.get(Employee, target)
            employee.hired_on = earlier
            employee.is_active = False
            employee.terminated_on = departed
            account = db.query(UserAccount).filter_by(employee_id=target).one()
            account.enabled = False
            account.disabled_at = datetime.now() - timedelta(days=3, seconds=1)
            admin = db.query(Employee).filter_by(employee_no='HR01').one()
            recognition = RecognitionRecord(employee_id=target, employee_no=number, employee_name=employee.name,
                home_attraction_id=employee.attraction_id, occurred_attraction_id=employee.attraction_id,
                employee_role_snapshot='CM', recognition_type_id=db.query(RecognitionType).first().id,
                recognition_type_name=db.query(RecognitionType).first().name,
                recognition_date=earlier, recognition_month=earlier[:7], content='重新入职保留原记录', fraction=Decimal('1'),
                recognizer_employee_id=admin.id, recognizer_name=admin.name, recognizer_role_snapshot='SYSTEM_ADMIN',
                operator_employee_id=admin.id, operator_name=admin.name, source='admin', status='confirmed')
            db.add(recognition)
            db.commit()
            original_record = recognition.id
        # Merely disabled is insufficient: the existing account still owns its number.
        assert client.post('/api/hr/employees', json=body).status_code == 400
        removed = client.request('DELETE', f'/api/hr/employees/{target}/account', json={})
        assert removed.status_code == 200, removed.text
        with SessionLocal() as db:
            employee = db.get(Employee, target)
            if legacy_archive:
                employee.employment_history = '[]'
            group = db.query(WorkGroup).filter_by(attraction_id=circles[0]['id']).first()
            if group:
                db.add(GroupMembership(employee_id=target, group_id=group.id, starts_on=earlier))
            duty = EmployeeActingDuty(employee_id=target, role_id=db.query(Role).filter_by(code='TA_SUPERVISOR').one().id, starts_on=earlier)
            db.add(duty)
            db.add(ManagementScope(employee_id=target, attraction_id=circles[0]['id'], starts_on=earlier))
            db.commit()
        wrong_name = client.post('/api/hr/employees', json={**body, 'name':'另一人'})
        assert wrong_name.status_code == 409, wrong_name.text
        # A forbidden change must leave the archived identity and login untouched.
        future = client.post('/api/hr/employees', json={**body, 'hired_on':(today+timedelta(days=1)).isoformat()})
        assert future.status_code == 400, future.text
        with SessionLocal() as db:
            db.add(MonthClosure(closure_month=today.strftime('%Y-%m'), attraction_id=circles[0]['id'], status='closed', closed_by_name='rehire gate test'))
            db.commit()
        blocked = client.post('/api/hr/employees', json=body)
        assert blocked.status_code == 423, blocked.text
        with SessionLocal() as db:
            assert db.get(Employee, target).account_deleted_at is not None
            assert db.query(UserAccount).filter_by(employee_id=target).count() == 0
            db.query(MonthClosure).filter_by(closure_month=today.strftime('%Y-%m'), attraction_id=circles[0]['id']).delete()
            db.commit()
        body.update(role_code='TR', attraction_id=circles[-1]['id'])
        rehired = client.post('/api/hr/employees', json=body)
        assert rehired.status_code == 200, rehired.text
        assert rehired.json()['rehired'] is True
        assert rehired.json()['employee']['id'] == target
        assert target in organization_employee_ids(client)
        assert client.post('/api/hr/employees', json=body).status_code == 400
        with SessionLocal() as db:
            employee = db.get(Employee, target)
            assert employee.account_deleted_at is None and employee.is_active
            assert employee.hired_on == today.isoformat()
            assert employee.attraction_id == circles[-1]['id']
            assert db.query(UserAccount).filter_by(employee_id=target).one().enabled
            assert db.get(RecognitionRecord, original_record).content == '重新入职保留原记录'
            assert role_at(db, target).code == 'TR'
            assert not db.query(EmployeeActingDuty).filter_by(employee_id=target).filter(EmployeeActingDuty.ends_on >= today.isoformat()).count()
            assert all(row.ends_on < today.isoformat() for row in db.query(GroupMembership).filter_by(employee_id=target))
            assert all(row.ends_on < today.isoformat() for row in db.query(ManagementScope).filter_by(employee_id=target))
            assert employment_history(employee)[-1]['rehired_on'] == today.isoformat()
            assert employee_active_on(employee, earlier)
            assert not employee_active_on(employee, (today-timedelta(days=4)).isoformat())
            assert employee_active_on(employee, today.isoformat())
            book = Workbook()
            write_employee_number_history_sheet(book, db, [target])
            values = list(book['员工留档标记'].values)
            assert values[1][2] == '曾删除·已重新注册'
            assert values[1][5] == today.isoformat()
            assert db.query(AuditLog).filter_by(action='重新入职注册', entity_id=str(target)).count() == 1


def test_multiple_deleted_employment_cycles_preserve_dates_and_gaps():
    from app.services.employment import employment_history, retain_deleted_employment
    from app.services.identity import employee_active_on, employed_on
    employee = SimpleNamespace(employee_no='9817030', name='历次任职测试', attraction_id=1,
        hired_on='2024-01-01', terminated_on='2024-06-01', is_active=False,
        account_deleted_at=datetime(2024,6,1), account_deleted_by_name='测试', employment_history='[]')
    retain_deleted_employment(employee)
    retain_deleted_employment(employee)
    assert len(employment_history(employee)) == 1
    employee.hired_on='2025-01-01'; employee.terminated_on='2025-06-01'
    employee.account_deleted_at=datetime(2025,6,1)
    retain_deleted_employment(employee)
    employee.hired_on='2026-01-01'; employee.terminated_on=None
    employee.account_deleted_at=None; employee.is_active=True
    assert len(employment_history(employee)) == 2
    for predicate in (employee_active_on, employed_on):
        for day in ('2024-03-01','2025-03-01','2026-01-01'):
            assert predicate(employee, day)
        for day in ('2023-12-31','2024-06-01','2024-12-31','2025-08-01'):
            assert not predicate(employee, day)

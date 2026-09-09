from __future__ import annotations

from datetime import date, datetime
from urllib.parse import quote

from openpyxl import Workbook

from app.excel_export_utils import (
    add_banded_table,
    apply_date_format,
    content_disposition,
    display_width,
    excel_safe_value,
    style_sheet,
    write_kpi_block,
    write_label_value_pairs,
    write_merged_title,
)


def test_excel_safe_value_escapes_formula_like_text_only() -> None:
    assert excel_safe_value("=cmd|'/c calc'!A0") == "'=cmd|'/c calc'!A0"
    assert excel_safe_value("+SUM(A1:A2)") == "'+SUM(A1:A2)"
    assert excel_safe_value("@foo") == "'@foo"
    assert excel_safe_value("认可内容") == "认可内容"
    value = excel_safe_value(-0.5)
    assert value == -0.5
    assert type(value) is float
    # 工具层会给公式形文本加引号；业务公式单元格不得走 excel_safe_value。
    assert excel_safe_value("=SUM(A1:A2)") == "'=SUM(A1:A2)"


def test_style_sheet_column_width_uses_cjk_display_width() -> None:
    header = "当月已确认认可次数"
    wb = Workbook()
    ws = wb.active
    ws.append([header, "姓名"])
    ws.append(["内容", "张三"])
    style_sheet(ws)
    width = ws.column_dimensions["A"].width
    assert display_width(header) == 18
    assert width >= display_width(header)
    assert width <= 42
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws.sheet_view.showGridLines is False
    assert ws.sheet_view.zoomScale == 90
    assert ws.page_setup.fitToWidth == 1
    assert ws["A1"].fill.fgColor.rgb in ("00D9EAF7", "D9EAF7")
    assert ws["A1"].alignment.wrap_text is True
    assert ws.row_dimensions[1].height == 24
    assert ws.protection.sheet is False


def test_content_disposition_includes_ascii_and_rfc5987_filename() -> None:
    ascii_filename = "pr_rankings_overall_2026-01-01_2026-01-31.xlsx"
    display_filename = "PR排名_综合.xlsx"
    header = content_disposition(ascii_filename, display_filename)
    assert f'filename="{ascii_filename}"' in header
    assert "filename*=UTF-8''" in header
    assert quote(display_filename) in header


def test_cover_title_and_kpi_rows() -> None:
    wb = Workbook()
    ws = wb.active
    write_merged_title(ws, "认可数据 · 2026-01 · 全部可见景点圈", 2)
    next_row = write_label_value_pairs(ws, [("统计月份", "2026-01")], 2)
    write_kpi_block(
        ws,
        [("加分合计", 3.5), ("扣分合计", 1.0), ("全勤人数", 2), ("综合分 Top3", "张三 12.00")],
        next_row + 1,
    )
    labels = [ws.cell(row, 1).value for row in range(1, (ws.max_row or 1) + 1)]
    assert "认可数据" in str(ws["A1"].value)
    assert "加分合计" in labels
    assert "扣分合计" in labels
    assert "全勤人数" in labels
    assert "综合分 Top3" in labels
    assert ws["A1"].fill.fgColor.rgb in ("001F4E78", "1F4E78")


def test_add_banded_table_does_not_raise() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "分值"])
    ws.append(["张三", 1.5])
    style_sheet(ws)
    add_banded_table(ws, ws.dimensions, "TestBanded")
    add_banded_table(ws, "A1:B2", "TestBanded")
    apply_date_format(ws, (1,))
    ws["A3"] = date(2026, 1, 15)
    ws["A4"] = "2026-01-16"
    ws["A5"] = datetime(2026, 1, 17, 8, 0)
    apply_date_format(ws, (1,))
    assert ws["A3"].number_format == "yyyy-mm-dd"
    assert ws["A4"].number_format != "yyyy-mm-dd"
    assert ws["A5"].number_format == "yyyy-mm-dd"

from __future__ import annotations

import os
from datetime import date
from io import BytesIO
from pathlib import Path
import tempfile

from openpyxl import load_workbook
from PIL import Image


TEST_DATA_DIR = Path(tempfile.mkdtemp(prefix="recognition-statement-upgrade-"))
os.environ["RECOGNITION_V2_DATA_DIR"] = str(TEST_DATA_DIR)
os.environ["RECOGNITION_ENABLE_TEST_ACCOUNTS"] = "1"
os.environ["RECOGNITION_TEST_DEFAULT_PASSWORD"] = "1234"
os.environ["RECOGNITION_TEST_ADMIN_PASSWORD"] = "HR123"

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402


def login(client: TestClient, account: str) -> None:
    response = client.post("/api/login", json={"employee_no": account, "password": "1234"})
    assert response.status_code == 200, response.text


def statement_form(target_id: int, type_id: int, level_id: int, description: str) -> dict:
    return {
        "employee_id": str(target_id),
        "deduction_type_id": str(type_id),
        "deduction_level_id": str(level_id),
        "occurred_on": date.today().isoformat(),
        "description": description,
    }


def pdf_file(name: str = "statement.pdf"):
    stream = BytesIO()
    Image.new("RGB", (8, 8), "white").save(stream, format="PDF")
    return {"document": (name, stream.getvalue(), "application/pdf")}


def test_statement_upgrade_submit_transfer_approve_and_export_marking() -> None:
    with TestClient(app) as client:
        login(client, "TATEST01")
        options = client.get("/api/options").json()
        target = next(row for row in client.get("/api/employee-targets", params={"usage": "deduction", "keyword": "CMTEST01"}).json()["items"] if row["employee_no"] == "CMTEST01")
        deduction_type = next(row for row in options["deduction_types"] if row["repeat_check"])
        statement = next(row for row in options["deduction_levels"] if row["code"] == "STATEMENT")
        first = client.post("/api/deductions", data=statement_form(target["id"], deduction_type["id"], statement["id"], "A1"), files=pdf_file())
        assert first.status_code == 200, first.text
        preview = client.get("/api/deduction-upgrades/preview", params={"employee_id": target["id"], "deduction_type_id": deduction_type["id"], "occurred_on": date.today().isoformat()})
        assert preview.status_code == 200 and preview.json()["eligible"], preview.text
        tagsm = next(row for row in preview.json()["reviewers"] if row["role_code"] == "TA_GSM")
        data = statement_form(target["id"], deduction_type["id"], statement["id"], "A2")
        data["reviewer_id"] = str(tagsm["id"])
        submitted = client.post("/api/deduction-upgrades", data=data, files=pdf_file("second.pdf"))
        assert submitted.status_code == 200, submitted.text
        request_id = submitted.json()["request"]["id"]
        assert submitted.json()["request"]["second_record"]["status"] == "pending_upgrade"
        client.post("/api/logout")
        login(client, "TAGSMTEST01")
        pending = client.get("/api/deduction-upgrades/pending").json()["items"]
        assert [row["id"] for row in pending] == [request_id]
        reviewers = client.get("/api/deduction-upgrades/reviewers").json()["items"]
        gsm = next(row for row in reviewers if row["role_code"] == "GSM")
        moved = client.post(f"/api/deduction-upgrades/{request_id}/transfer", json={"reviewer_id": gsm["id"], "reason": "由GSM审核"})
        assert moved.status_code == 200, moved.text
        client.post("/api/logout")
        login(client, "GSMTEST01")
        pending = client.get("/api/deduction-upgrades/pending").json()["items"]
        assert pending[0]["transfers"][0]["to_name"] == "测试GSM"
        memo = next(row for row in client.get("/api/options").json()["deduction_levels"] if row["code"] == "MEMO")
        approved = client.post(f"/api/deduction-upgrades/{request_id}/resolve", json={"decision": "approve", "result_level_id": memo["id"], "handling_note": "已完成真实备忘录开具", "issued_confirmed": True})
        assert approved.status_code == 200, approved.text
        response = client.get("/api/statistics/export", params={"month": date.today().strftime("%Y-%m")})
        assert response.status_code == 200, response.text
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        sheet = workbook["扣分明细"]
        roles = [sheet.cell(row=index, column=12).value for index in range(2, sheet.max_row + 1)]
        assert {"source_first", "source_second", "result"}.issubset(set(roles))
        colors = [sheet.cell(row=index, column=1).fill.fgColor.rgb for index in range(2, sheet.max_row + 1)]
        assert any(color and color.endswith("DDEBF7") for color in colors)
        assert any(color and color.endswith("FFF2CC") for color in colors)
        assert any(color and color.endswith("E2F0D9") for color in colors)

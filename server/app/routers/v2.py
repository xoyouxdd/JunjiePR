from __future__ import annotations

import hashlib
import json
import os
import secrets
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from math import ceil
from pathlib import Path
from threading import Lock
from time import monotonic
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import func, or_, text, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.v2_auth import V2User, current_user, require_permissions
from app.v2_crypto import default_initial_password, hash_password, new_session_token, new_temporary_password, token_hash, verify_password
from app.v2_database import CIRCLE_HR_ACCOUNTS, EMPLOYEE_CIRCLES, EXPORT_DIR, FILE_DIR, LEGACY_CIRCLE_BY_VENUE, RECOGNITION_VENUES, get_db, synchronize_gsm_management_scope
from app.v2_models import (
    Attraction,
    AttendanceMonthlyScore,
    AuditLog,
    CircleTransferRequest,
    DeductionLevel,
    DeductionFollowUp,
    DeductionMaterialJob,
    DeductionUpgradeRequest,
    DeductionUpgradeTransfer,
    DeductionRecord,
    DeductionType,
    Employee,
    EmployeeNumberHistory,
    EmployeeMonthOrganizationSnapshot,
    EmployeeLOAPeriod,
    EmployeeRoleAssignment,
    GroupLeaderAssignment,
    GroupMembership,
    GroupTransfer,
    GroupTransferMember,
    GovernanceCase,
    ManagementScope,
    MonthClosure,
    RecognitionRecord,
    RecognitionAttachment,
    RecognitionMonthlyQuota,
    RecognitionReview,
    RecognitionScoreRule,
    RecognitionType,
    Role,
    SickLeaveRecord,
    StoredFile,
    SubmissionRequest,
    SystemAlert,
    UserAccount,
    UserSession,
    WorkGroup,
)
from app.recognition_encouragement import encouragement_options
from app.v2_services import (
    FRONTLINE_CODES,
    GSM_CODES,
    LEADER_CODES,
    RECOGNIZER_CODES,
    active_group_leader,
    active_group_leaders_bulk,
    active_group_memberships,
    active_group_memberships_bulk,
    current_group_for_employee,
    current_leader_for_employee,
    direct_member_ids,
    ensure_month_attendance,
    groups_led_by,
    groups_led_by_bulk,
    gsm_candidates_for_attractions_bulk,
    managed_attraction_ids,
    full_month_loa,
    process_role_expirations,
    recalculate_attendance,
    recognition_score_for_role,
    recognizer_role_for_date,
    recognizer_options,
    remove_upload_file,
    role_at,
    roles_at,
    role_permissions,
    save_upload,
    save_image_upload,
    write_audit,
)
from app.deduction_materials import PDF_HIGH_QUALITY_OPTIMIZATION_THRESHOLD, create_pdf_placeholder, queue_photo_material_job, stage_pdf_material, stage_photo_materials
from app.v2_watermark import watermark_image, watermark_pdf, watermark_workbook
from app.excel_export_utils import content_disposition
from app.excel_export import (
    build_employee_import_template,
    build_pr_rankings_workbook,
    build_statistics_workbook,
)
from app.changelog import visible_releases
from app.version import APP_VERSION
from app.v2_preview_cache import watermarked_preview_cache
from app.security import password_policy_error, request_is_https


router = APIRouter(prefix="/api", tags=["v2"])

STATISTICS_CACHE_SECONDS = 2.0
_statistics_cache_lock = Lock()
_statistics_response_cache: dict[tuple, tuple[float, bytes]] = {}
PR_RANKING_CACHE_SECONDS = 5.0
_pr_ranking_cache_lock = Lock()
_pr_ranking_response_cache: dict[tuple, tuple[float, bytes]] = {}
PREVIEW_IMAGE_EXTENSIONS = frozenset({".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"})
PREVIEW_PDF_EXTENSIONS = frozenset({".pdf"})


def invalidate_data_caches() -> None:
    with _statistics_cache_lock:
        _statistics_response_cache.clear()
    with _pr_ranking_cache_lock:
        _pr_ranking_response_cache.clear()


def is_previewable_image(file_row: StoredFile | None) -> bool:
    """Only active image files use the protected in-page viewer."""
    return bool(file_row and file_row.status == "active" and file_row.extension.lower() in PREVIEW_IMAGE_EXTENSIONS)


def preview_kind(file_row: StoredFile | None) -> str:
    """Return the safe, browser-supported in-page preview type for a stored file."""
    if not file_row or file_row.status != "active":
        return ""
    suffix = file_row.extension.lower()
    if suffix in PREVIEW_IMAGE_EXTENSIONS:
        return "image"
    if suffix in PREVIEW_PDF_EXTENSIONS:
        return "pdf"
    return ""

ATTENDANCE_DEDUCTION_CODES = {
    "ATT_EARLY_CLOCK",
    "ATT_LATE_WITHIN_30",
    "ATT_LATE_OVER_30",
    "ATT_LATE_CLOCK",
    "ATT_EARLY_LEAVE_WITHIN_30",
    "ATT_EARLY_LEAVE_OVER_30",
    "ATT_MISSING_CLOCK",
    "ATT_REMOTE_CLOCK",
}
REPEAT_CONTROLLED_DEDUCTION_CODES = ATTENDANCE_DEDUCTION_CODES | {"SICK_LEAVE_VIOLATION"}
DIRECT_HIDDEN_DEDUCTION_CODES = {"ATT_LATE_OVER_30", "ATT_EARLY_LEAVE_OVER_30"}
DEDUCTION_TYPE_ORDER = (
    "SAFETY",
    "AUDIT",
    "COURTESY",
    "INCLUSION",
    "SHOW",
    "EFFICIENCY",
    "MSP",
    "COMPLAINT",
    "SICK_LEAVE_VIOLATION",
    "OTHER",
    "ATT_EARLY_CLOCK",
    "ATT_LATE_CLOCK",
    "ATT_LATE_OVER_30",
    "ATT_LATE_WITHIN_30",
    "ATT_EARLY_LEAVE_OVER_30",
    "ATT_EARLY_LEAVE_WITHIN_30",
    "ATT_MISSING_CLOCK",
    "ATT_REMOTE_CLOCK",
)
SPECIAL_RECOGNITION_TYPES = {
    "MSP": {"option_id": "special:MSP", "name": "MSP", "score": Decimal("3.00"), "monthly_limit": None},
    "COMMENDATION_LETTER": {
        "option_id": "special:COMMENDATION_LETTER",
        "name": "表扬信",
        "score": Decimal("3.00"),
        "monthly_limit": 1,
    },
}
# Five ordinary recognition categories are capped from the agreed effective
# date.  Existing August history keeps its already-confirmed score unchanged.
MONTHLY_CATEGORY_CAP_CODES = {"SAFETY", "COURTESY", "INCLUSION", "EFFICIENCY", "SHOW"}
MONTHLY_CATEGORY_CAP_EFFECTIVE_DATE = date(2026, 9, 1)
MONTHLY_CATEGORY_CAP_LIMIT = Decimal("5.00")
DEDUCTION_LEVEL_ORDER = {"STATEMENT": 1, "MEMO": 2, "WARNING_1": 3, "WARNING_2": 4}
NEXT_DEDUCTION_LEVEL = {"STATEMENT": "MEMO", "MEMO": "WARNING_1", "WARNING_1": "WARNING_2", "WARNING_2": "WARNING_2"}
CIRCLE_HR_MANAGED_ROLE_CODES = {"CM", "TR", "TA_SUPERVISOR", "SUPERVISOR"}
# Account-status visibility is intentionally broader than employee-edit
# authority: scoped HR may read the four frontline/leader roles across all
# circles, while the highest administrator may read every regular account.
REGULAR_ACCOUNT_ROLE_CODES = CIRCLE_HR_MANAGED_ROLE_CODES | {"TA_GSM", "GSM", "AM", "OM"}
SCOPED_HR_ROLE_CODE = "HR_CIRCLE"
# Circle HR may only retune recognizer roles it manages; GSM/TA_GSM/AM/OM rules are global.
CIRCLE_HR_SCORE_RULE_ROLE_CODES = CIRCLE_HR_MANAGED_ROLE_CODES & RECOGNIZER_CODES
# Record status families shared by the list endpoints' filters.
RECOGNITION_FILTER_STATUSES = frozenset({"pending", "rejected", "confirmed"})
ATTENDANCE_FILTER_STATUSES = frozenset({"active", "void"})
DEDUCTION_FILTER_STATUSES = frozenset({"active", "void", "pending_upgrade", "pending_material", "material_processing", "material_failed"})
MONTH_CLOSE_ROLE_CODES = {"HR_CIRCLE", "SYSTEM_ADMIN"}
MATERIAL_COLLABORATOR_CODES = LEADER_CODES | GSM_CODES
MONTH_CLOSE_EFFECTIVE_DATE = "2026-09-01"
GROUP_DISPLAY_EFFECTIVE_DATE = "2026-09-01"
BACKUP_HEALTH_STATUS_PATH = Path(
    os.environ.get(
        "RECOGNITION_BACKUP_HEALTH_STATUS_FILE",
        r"C:\Server\zhaojunjie\backups\recognition-card-system-sqlite\backup-health-status.json",
    )
)


def scoped_hr_attraction_ids(db: Session, user: V2User) -> set[int] | None:
    """Return the circle restriction for a scoped HR account; global roles remain unrestricted."""
    if user.role.code != SCOPED_HR_ROLE_CODE:
        return None
    if not user.employee.attraction_id:
        raise HTTPException(403, "景点圈HR账号未配置景点圈")
    return {user.employee.attraction_id}


def ensure_scoped_hr_attraction(db: Session, user: V2User, attraction_id: int | None) -> None:
    allowed = scoped_hr_attraction_ids(db, user)
    if allowed is not None and attraction_id not in allowed:
        raise HTTPException(403, "景点圈HR只能操作所属景点圈")


def ensure_scoped_hr_employee(db: Session, user: V2User, employee: Employee) -> None:
    ensure_scoped_hr_attraction(db, user, employee.attraction_id)


def ensure_operational_target_scope(user: V2User, employee: Employee) -> None:
    """Backend authority for deduction/absence target selection.

    Search UI is only a convenience: TA主管 is confined to its own circle;
    主管、TA GSM、GSM may support active frontline staff across circles.
    """
    if user.role.code == "TA_SUPERVISOR" and employee.attraction_id != user.employee.attraction_id:
        raise HTTPException(403, "TA主管仅可登记本景点圈CM/TR")
    if user.role.code not in {"TA_SUPERVISOR", "SUPERVISOR", "TA_GSM", "GSM", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "当前角色无权登记该员工")


def ensure_employee_number_change_target(db: Session, user: V2User, employee: Employee, role: Role | None = None) -> Role:
    """Apply the employee-number migration scope without changing org data."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以变更员工号")
    target_role = role or role_at(db, employee.id)
    if not target_role:
        raise HTTPException(400, "该员工当前没有有效角色")
    if target_role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR和最高管理员账号不能在此处变更员工号")
    ensure_scoped_hr_employee(db, user, employee)
    if user.role.code == SCOPED_HR_ROLE_CODE and target_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能变更本圈CM、TR、TA主管或主管的员工号")
    return target_role


def visible_system_alerts(db: Session, user: V2User) -> list[SystemAlert]:
    """Return only alerts inside the caller's existing HR scope."""
    rows = db.query(SystemAlert).order_by(SystemAlert.status.asc(), SystemAlert.created_at.desc()).limit(200).all()
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is None:
        return rows
    group_ids = {
        group.id
        for group in db.query(WorkGroup).filter(WorkGroup.attraction_id.in_(allowed_attractions)).all()
    }
    employee_ids = {
        employee.id
        for employee in db.query(Employee).filter(Employee.attraction_id.in_(allowed_attractions)).all()
    }
    return [row for row in rows if (row.group_id in group_ids if row.group_id else row.employee_id in employee_ids)]


def backup_health_payload() -> dict:
    """Read the fixed server health report without exposing paths or webhook data."""
    default = {
        "available": False,
        "ok": False,
        "status": "未发现健康报告",
        "checked_at_utc": "",
        "task": {},
        "latest_backup": None,
        "issues": [{"code": "health_report_missing", "message": "未发现备份健康报告，请检查每日备份巡检任务。"}],
        "alert": {"configured": False, "attempted": False, "delivered": False, "last_alert_at_utc": ""},
    }
    try:
        raw = json.loads(BACKUP_HEALTH_STATUS_PATH.read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        return default
    except (OSError, ValueError, TypeError):
        default["status"] = "健康报告无法读取"
        default["issues"] = [{"code": "health_report_invalid", "message": "备份健康报告无法读取，请检查巡检任务。"}]
        return default
    if not isinstance(raw, dict):
        default["status"] = "健康报告格式无效"
        default["issues"] = [{"code": "health_report_invalid", "message": "备份健康报告格式无效，请检查巡检任务。"}]
        return default
    task = raw.get("task") if isinstance(raw.get("task"), dict) else {}
    latest = raw.get("latest_backup") if isinstance(raw.get("latest_backup"), dict) else None
    alert = raw.get("alert") if isinstance(raw.get("alert"), dict) else {}
    issues = raw.get("issues") if isinstance(raw.get("issues"), list) else []
    return {
        "available": True,
        "ok": bool(raw.get("ok")),
        "status": "正常" if bool(raw.get("ok")) else "需处理",
        "checked_at_utc": str(raw.get("checked_at_utc") or ""),
        "task": {
            "exists": bool(task.get("exists")),
            "enabled": bool(task.get("enabled")),
            "state": str(task.get("state") or "未知"),
            "last_run_time_utc": str(task.get("last_run_time_utc") or ""),
            "next_run_time_utc": str(task.get("next_run_time_utc") or ""),
            "last_task_result": task.get("last_task_result"),
        },
        "latest_backup": None if latest is None else {
            "file_name": Path(str(latest.get("file_name") or "")).name,
            "created_at_utc": str(latest.get("created_at_utc") or ""),
            "age_hours": latest.get("age_hours"),
            "size_bytes": latest.get("size_bytes"),
            "manifest_verified": bool(latest.get("manifest_verified")),
            "sha256_verified": bool(latest.get("sha256_verified")),
            "size_verified": bool(latest.get("size_verified")),
            "quick_check": str(latest.get("quick_check") or ""),
        },
        "issues": [
            {"code": str(item.get("code") or "unknown"), "message": str(item.get("message") or "")}
            for item in issues if isinstance(item, dict)
        ][:20],
        "alert": {
            "configured": bool(alert.get("configured")),
            "attempted": bool(alert.get("attempted")),
            "delivered": bool(alert.get("delivered")),
            "last_alert_at_utc": str(alert.get("last_alert_at_utc") or ""),
        },
    }


def client_ip(request: Request) -> str | None:
    return request.client.host if request.client else None


def parse_iso_date(value: str, label: str = "日期") -> date:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, f"{label}格式必须为YYYY-MM-DD") from exc


def parse_score_month(value: str) -> str:
    month = str(value or "").strip()
    if len(month) != 7:
        raise HTTPException(400, "月份格式必须为YYYY-MM")
    try:
        parsed = date.fromisoformat(f"{month}-01")
    except ValueError as exc:
        raise HTTPException(400, "月份格式必须为YYYY-MM") from exc
    normalized = parsed.strftime("%Y-%m")
    if month != normalized:
        raise HTTPException(400, "月份格式必须为YYYY-MM")
    return normalized


def month_closure_scope(db: Session, month: str, attraction_id: int | None) -> MonthClosure | None:
    query = db.query(MonthClosure).filter(
        MonthClosure.closure_month == month,
        MonthClosure.status == "closed",
    )
    if attraction_id is None:
        return query.filter(MonthClosure.attraction_id.is_(None)).first()
    return query.filter(or_(MonthClosure.attraction_id.is_(None), MonthClosure.attraction_id == attraction_id)).order_by(
        MonthClosure.attraction_id.is_(None).desc()
    ).first()


def ensure_month_open(db: Session, month: str, attraction_id: int | None, operation: str) -> None:
    month = parse_score_month(month)
    closure = month_closure_scope(db, month, attraction_id)
    if not closure:
        return
    attraction = db.get(Attraction, closure.attraction_id) if closure.attraction_id else None
    scope_name = attraction.name if attraction else "全部景点圈"
    raise HTTPException(
        423,
        detail={
            "code": "MONTH_CLOSED",
            "message": f"{month} · {scope_name}已月结，不能{operation}。如需更正，请由该景点圈HR填写原因后临时开放。",
            "month": month,
            "attraction_id": closure.attraction_id,
            "attraction_name": scope_name,
            "closed_at": closure.closed_at.strftime("%Y-%m-%d %H:%M:%S") if closure.closed_at else "",
            "closed_by_name": closure.closed_by_name or "",
        },
    )


def ensure_month_close_scope(db: Session, user: V2User, attraction_id: int | None) -> Attraction | None:
    if user.role.code not in MONTH_CLOSE_ROLE_CODES:
        raise HTTPException(403, "仅景点圈HR或最高管理员可以关闭月结")
    if attraction_id is None:
        if user.role.code != "SYSTEM_ADMIN":
            raise HTTPException(400, "请选择要关闭的景点圈")
        return None
    attraction = db.get(Attraction, attraction_id)
    if not attraction or not attraction.active or not attraction.employee_circle:
        raise HTTPException(400, "请选择有效景点圈")
    if user.role.code != "SYSTEM_ADMIN" and attraction.id not in managed_attraction_ids(db, user.id):
        raise HTTPException(403, "只能关闭自己管理范围内的景点圈")
    return attraction


def month_close_checklist(db: Session, month: str, attraction_id: int | None) -> list[dict]:
    """Only unresolved work that can alter this month blocks the close."""
    month_start = date.fromisoformat(f"{month}-01")
    month_end = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
    employee_query = db.query(Employee.id)
    if attraction_id is not None:
        employee_query = employee_query.filter(Employee.attraction_id == attraction_id)
    employee_ids = [row[0] for row in employee_query.all()]
    pending_recognitions = db.query(RecognitionRecord).filter(
        RecognitionRecord.recognition_month == month,
        RecognitionRecord.status == "pending",
        *(() if attraction_id is None else (RecognitionRecord.home_attraction_id == attraction_id,)),
    ).count()
    unresolved_materials = db.query(DeductionRecord).filter(
        DeductionRecord.deduction_month == month,
        DeductionRecord.status.in_({"pending_material", "material_processing", "material_failed"}),
        *(() if attraction_id is None else (DeductionRecord.attraction_id_snapshot == attraction_id,)),
    ).count()
    upgrade_query = db.query(DeductionUpgradeRequest).join(
        DeductionRecord, DeductionRecord.id == DeductionUpgradeRequest.second_deduction_id
    ).filter(
        DeductionRecord.deduction_month == month,
        DeductionUpgradeRequest.status == "pending",
    )
    if attraction_id is not None:
        upgrade_query = upgrade_query.filter(DeductionRecord.attraction_id_snapshot == attraction_id)
    upgrade_requests = upgrade_query.count()
    follow_up_query = db.query(DeductionFollowUp).filter(
        DeductionFollowUp.occurred_on.like(f"{month}%"),
        DeductionFollowUp.status == "pending",
    )
    if attraction_id is not None:
        follow_up_query = follow_up_query.filter(
            DeductionFollowUp.employee_id.in_(employee_ids) if employee_ids else DeductionFollowUp.id == -1
        )
    follow_ups = follow_up_query.count()
    transfer_query = db.query(CircleTransferRequest).filter(
        CircleTransferRequest.status == "pending",
        CircleTransferRequest.requested_at >= datetime.combine(month_start, datetime.min.time()),
        CircleTransferRequest.requested_at < datetime.combine(month_end, datetime.min.time()),
    )
    if attraction_id is not None:
        transfer_query = transfer_query.filter(
            or_(CircleTransferRequest.source_attraction_id == attraction_id, CircleTransferRequest.target_attraction_id == attraction_id)
        )
    pending_transfers = transfer_query.count()
    return [
        {"code": "recognition_review", "name": "待复核签卡", "count": pending_recognitions, "target": "review"},
        {"code": "deduction_material", "name": "待补/生成失败材料", "count": unresolved_materials, "target": "entries"},
        {"code": "deduction_upgrade", "name": "待升级工单", "count": upgrade_requests, "target": "actionCenter"},
        {"code": "deduction_follow_up", "name": "重复违规待跟进", "count": follow_ups, "target": "entries"},
        {"code": "circle_transfer", "name": "待确认跨圈调动", "count": pending_transfers, "target": "circleTransfers"},
    ]


def month_closure_payload(db: Session, month: str, attraction_id: int | None, user: V2User) -> dict:
    global_row = db.query(MonthClosure).filter_by(
        closure_month=month, attraction_id=None, status="closed"
    ).first()
    scope_row = None
    if attraction_id is not None:
        scope_row = db.query(MonthClosure).filter_by(
            closure_month=month, attraction_id=attraction_id, status="closed"
        ).first()
    effective = global_row or scope_row
    attraction = db.get(Attraction, attraction_id) if attraction_id else None
    effective_attraction = db.get(Attraction, effective.attraction_id) if effective and effective.attraction_id else None
    can_close = False
    if not effective and user.role.code in MONTH_CLOSE_ROLE_CODES:
        if user.role.code == "SYSTEM_ADMIN":
            can_close = True
        elif attraction_id is not None and attraction_id in managed_attraction_ids(db, user.id):
            can_close = True
    return {
        "month": month,
        "attraction_id": attraction_id,
        "attraction_name": attraction.name if attraction else "全部景点圈",
        "status": "closed" if effective else "open",
        "is_closed": bool(effective),
        "effective_scope": effective_attraction.name if effective_attraction else ("全部景点圈" if effective else ""),
        "closed_by_name": effective.closed_by_name if effective else "",
        "closed_at": effective.closed_at.strftime("%Y-%m-%d %H:%M:%S") if effective and effective.closed_at else "",
        "close_reason": effective.close_reason if effective else "",
        "can_close": can_close,
        "can_reopen": bool(effective and (user.role.code == "SYSTEM_ADMIN" or (user.role.code == "HR_CIRCLE" and attraction_id is not None and attraction_id in scoped_hr_attraction_ids(db, user)))),
        "reopen_target_attraction_id": effective.attraction_id if effective else attraction_id,
        "checklist": month_close_checklist(db, month, attraction_id),
    }


def ensure_hr_role_allowed(user: V2User, role_code: str, label: str = "角色") -> None:
    if "SYSTEM_ADMIN" not in user.permissions and role_code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, f"景点圈HR只能设置{label}为CM、TR、TA主管或主管")


def normalize_request_key(value: str | None) -> str | None:
    key = str(value or "").strip()
    if not key:
        return None
    if len(key) > 100:
        raise HTTPException(400, "重复提交标识过长")
    return key


def submission_payload_digest(payload: dict) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def existing_submission(
    db: Session,
    actor_id: int,
    operation: str,
    request_key: str | None,
    model,
    payload_digest: str | None = None,
):
    if not request_key:
        return None
    request_row = db.query(SubmissionRequest).filter_by(
        actor_id=actor_id,
        operation=operation,
        request_key=request_key,
    ).first()
    if not request_row:
        return None
    stored = (request_row.payload_digest or "").strip()
    if payload_digest and stored and stored != payload_digest:
        raise HTTPException(
            409,
            {
                "code": "IDEMPOTENCY_PAYLOAD_CONFLICT",
                "message": "同一提交键不能用于不同内容，请刷新页面后重新提交。",
            },
        )
    return db.get(model, request_row.entity_id)


def remember_submission(
    db: Session,
    actor_id: int,
    operation: str,
    request_key: str | None,
    entity_id: int,
    payload_digest: str | None = None,
) -> None:
    if request_key:
        db.add(
            SubmissionRequest(
                actor_id=actor_id,
                operation=operation,
                request_key=request_key,
                entity_id=entity_id,
                payload_digest=payload_digest,
            )
        )
        db.flush()


def subtract_calendar_months(value: date, months: int) -> date:
    month_index = value.year * 12 + value.month - 1 - months
    year, month_zero = divmod(month_index, 12)
    month = month_zero + 1
    return date(year, month, min(value.day, monthrange(year, month)[1]))


def add_calendar_months(value: date, months: int) -> date:
    month_index = value.year * 12 + value.month - 1 + months
    year, month_zero = divmod(month_index, 12)
    month = month_zero + 1
    return date(year, month, min(value.day, monthrange(year, month)[1]))


def group_display_metadata_bulk(db: Session, group_ids: list[int] | set[int], on_date: str | None = None) -> dict[int, dict]:
    """Resolve presentation-only group labels without ever renaming historic groups.

    Since 2026-09-01, a workgroup is displayed under its current leader.  The
    prior leader remains a one-month visual aid only; IDs, memberships,
    snapshots and historic exports all continue to use their original records.
    """
    ids = list(dict.fromkeys(int(group_id) for group_id in group_ids if group_id))
    if not ids:
        return {}
    value = on_date or date.today().isoformat()
    try:
        as_of = date.fromisoformat(value)
    except ValueError:
        as_of = date.today()
        value = as_of.isoformat()
    groups = {row.id: row for row in db.query(WorkGroup).filter(WorkGroup.id.in_(ids)).all()}
    active = active_group_leaders_bulk(db, ids, value)
    # The old leader only matters during the 1-calendar-month transition
    # window.  It is intentionally calculated on read and never stored on the
    # WorkGroup itself.
    transition_start = add_calendar_months(as_of, -1).isoformat()
    ended_rows = (
        db.query(GroupLeaderAssignment)
        .filter(
            GroupLeaderAssignment.group_id.in_(ids),
            GroupLeaderAssignment.status == "ended",
            GroupLeaderAssignment.ends_on.is_not(None),
            GroupLeaderAssignment.ends_on >= transition_start,
            GroupLeaderAssignment.ends_on <= value,
        )
        .order_by(GroupLeaderAssignment.group_id, GroupLeaderAssignment.ends_on.desc(), GroupLeaderAssignment.id.desc())
        .all()
    )
    latest_ended: dict[int, GroupLeaderAssignment] = {}
    for row in ended_rows:
        latest_ended.setdefault(row.group_id, row)
    leader_ids = {
        assignment.leader_employee_id
        for assignment in [*active.values(), *latest_ended.values()]
        if assignment
    }
    employees = {
        row.id: row for row in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()
    } if leader_ids else {}
    enabled = as_of >= date.fromisoformat(GROUP_DISPLAY_EFFECTIVE_DATE)
    result: dict[int, dict] = {}
    for group_id in ids:
        group = groups.get(group_id)
        current_assignment = active.get(group_id)
        current = employees.get(current_assignment.leader_employee_id) if current_assignment else None
        previous_assignment = latest_ended.get(group_id) if enabled else None
        previous = employees.get(previous_assignment.leader_employee_id) if previous_assignment else None
        previous_until = ""
        if previous_assignment and previous_assignment.ends_on:
            previous_until = add_calendar_months(date.fromisoformat(previous_assignment.ends_on), 1).isoformat()
            if as_of > date.fromisoformat(previous_until) or (current and previous and current.id == previous.id):
                previous = None
                previous_until = ""
        result[group_id] = {
            "name": f"{current.name}工作组" if enabled and current else (group.name if group else ""),
            "leader_name": current.name if current else "",
            "previous_leader_name": previous.name if previous else "",
            "previous_leader_until": previous_until if previous else "",
        }
    return result


def user_payload(db: Session, user: V2User) -> dict:
    leader = current_leader_for_employee(db, user.id)
    group = current_group_for_employee(db, user.id)
    display = group_display_metadata_bulk(db, [group.id]).get(group.id, {}) if group else {}
    return {
        "id": user.id,
        "employee_no": user.employee.employee_no,
        "name": user.name,
        "role_code": user.role.code,
        "role_name": user.role.name,
        "attraction_id": user.employee.attraction_id,
        "attraction_name": user.employee.attraction.name if user.employee.attraction else "",
        "leader_name": leader.name if leader else ("待接管" if group and group.status == "pending_takeover" else "未分配"),
        "group_name": display.get("name", group.name if group else ""),
        "previous_group_leader_name": display.get("previous_leader_name", ""),
        "previous_group_leader_until": display.get("previous_leader_until", ""),
        "permissions": sorted(user.permissions),
        "member_count": len(direct_member_ids(db, user.id)),
        "must_change_password": user.account.must_change_password,
    }


def employee_payload(db: Session, employee: Employee) -> dict:
    return employee_payloads(db, [employee])[employee.id]


def login_account_archive_state(employee: Employee, account: UserAccount | None, role: Role | None) -> dict:
    """Return non-sensitive eligibility for removing only a login account.

    The employee primary key and every business row remain untouched.  Both
    the directory and the write endpoint use this same policy so the seven-day
    rule cannot be bypassed by calling the API directly.
    """
    result = {"deleted": bool(employee.account_deleted_at), "eligible": False, "reason": "", "eligible_on": ""}
    if employee.account_deleted_at:
        result["reason"] = "账号已删除，业务档案已保留"
        return result
    if not account:
        result["reason"] = "该员工尚未开通登录账号"
        return result
    if not role or role.code not in REGULAR_ACCOUNT_ROLE_CODES:
        result["reason"] = "HR和最高管理员账号不支持在此删除"
        return result
    reference_day: date | None = None
    if not employee.is_active:
        try:
            reference_day = date.fromisoformat(str(employee.terminated_on or ""))
        except ValueError:
            result["reason"] = "离职日期未记录，暂不能删除登录账号"
            return result
    elif not account.enabled:
        if not account.disabled_at:
            result["reason"] = "停用时间未记录，暂不能删除登录账号"
            return result
        reference_day = account.disabled_at.date()
    else:
        result["reason"] = "需先离职或停用账号"
        return result
    eligible_on = reference_day + timedelta(days=7)
    result["eligible_on"] = eligible_on.isoformat()
    if date.today() < eligible_on:
        result["reason"] = f"{eligible_on.isoformat()} 起可删除登录账号"
        return result
    result["eligible"] = True
    result["reason"] = "可删除登录账号；员工及所有业务档案会保留"
    return result


def employee_payloads(db: Session, employees: list[Employee], on_date: str | None = None) -> dict[int, dict]:
    """Build the HR employee directory with a bounded set of batch queries."""
    if not employees:
        return {}
    value = on_date or date.today().isoformat()
    employee_ids = [employee.id for employee in employees]
    role_map = roles_at(db, employee_ids, value)
    accounts = {
        account.employee_id: account
        for account in db.query(UserAccount).filter(UserAccount.employee_id.in_(employee_ids)).all()
    }
    loa_periods = (
        db.query(EmployeeLOAPeriod)
        .filter(
            EmployeeLOAPeriod.employee_id.in_(employee_ids),
            EmployeeLOAPeriod.status != "cancelled",
            EmployeeLOAPeriod.starts_on <= value,
            or_(EmployeeLOAPeriod.ends_on.is_(None), EmployeeLOAPeriod.ends_on >= value),
        )
        .order_by(EmployeeLOAPeriod.employee_id, EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
        .all()
    )
    loa_by_employee: dict[int, EmployeeLOAPeriod] = {}
    for period in loa_periods:
        loa_by_employee.setdefault(period.employee_id, period)
    attraction_ids = {employee.attraction_id for employee in employees if employee.attraction_id}
    attractions = {
        attraction.id: attraction
        for attraction in db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).all()
    } if attraction_ids else {}
    memberships = (
        db.query(GroupMembership)
        .filter(
            GroupMembership.employee_id.in_(employee_ids),
            GroupMembership.status == "active",
            GroupMembership.starts_on <= value,
            or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= value),
        )
        .order_by(GroupMembership.employee_id, GroupMembership.starts_on.desc(), GroupMembership.id.desc())
        .all()
    )
    membership_by_employee: dict[int, GroupMembership] = {}
    for membership in memberships:
        membership_by_employee.setdefault(membership.employee_id, membership)
    group_ids = {membership.group_id for membership in membership_by_employee.values()}
    groups = {
        group.id: group
        for group in db.query(WorkGroup).filter(WorkGroup.id.in_(group_ids)).all()
    } if group_ids else {}
    leader_assignments = (
        db.query(GroupLeaderAssignment)
        .filter(
            GroupLeaderAssignment.group_id.in_(group_ids),
            GroupLeaderAssignment.status == "active",
            GroupLeaderAssignment.starts_on <= value,
            or_(GroupLeaderAssignment.ends_on.is_(None), GroupLeaderAssignment.ends_on >= value),
        )
        .order_by(GroupLeaderAssignment.group_id, GroupLeaderAssignment.starts_on.desc(), GroupLeaderAssignment.id.desc())
        .all()
    ) if group_ids else []
    leader_assignment_by_group: dict[int, GroupLeaderAssignment] = {}
    for assignment in leader_assignments:
        leader_assignment_by_group.setdefault(assignment.group_id, assignment)
    leader_ids = {assignment.leader_employee_id for assignment in leader_assignment_by_group.values()}
    leaders = {
        leader.id: leader
        for leader in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()
    } if leader_ids else {}
    group_display = group_display_metadata_bulk(db, group_ids, value)
    result: dict[int, dict] = {}
    for employee in employees:
        role = role_map.get(employee.id)
        account = accounts.get(employee.id)
        membership = membership_by_employee.get(employee.id)
        group = groups.get(membership.group_id) if membership else None
        leader_assignment = leader_assignment_by_group.get(group.id) if group else None
        leader = leaders.get(leader_assignment.leader_employee_id) if leader_assignment else None
        display = group_display.get(group.id, {}) if group else {}
        attraction = attractions.get(employee.attraction_id)
        loa_period = loa_by_employee.get(employee.id)
        archive_state = login_account_archive_state(employee, account, role)
        result[employee.id] = {
        "id": employee.id,
        "employee_no": employee.employee_no,
        "name": employee.name,
        "role_code": role.code if role else "",
        "role_name": role.name if role else "未配置",
        "attraction_id": employee.attraction_id,
        "attraction_name": attraction.name if attraction else "",
        "group_id": group.id if group else None,
        "group_name": display.get("name", group.name if group else ""),
        "previous_group_leader_name": display.get("previous_leader_name", ""),
        "previous_group_leader_until": display.get("previous_leader_until", ""),
        "leader_id": leader.id if leader else None,
        "leader_name": leader.name if leader else "",
        "is_active": employee.is_active,
        "employment_status": "loa" if employee.is_active and loa_period else "active" if employee.is_active else "terminated",
        "loa_start_date": loa_period.starts_on if loa_period else "",
        "account_enabled": bool(account and account.enabled),
        "login_account": account.login_account if account else "",
        "account_deleted_at": employee.account_deleted_at.strftime("%Y-%m-%d %H:%M:%S") if employee.account_deleted_at else "",
        "account_deleted_by_name": employee.account_deleted_by_name or "",
        "account_deletion_eligible": archive_state["eligible"],
        "account_deletion_reason": archive_state["reason"],
        "account_deletion_eligible_on": archive_state["eligible_on"],
        }
    return result


def account_password_status(account: UserAccount | None) -> tuple[str, str]:
    """Return a non-sensitive, evidence-based password state for HR views."""
    if not account:
        return "未开通账号", "unprovisioned"
    if account.must_change_password:
        return "待本人修改初始/重置密码", "pending_change"
    if account.password_changed_at:
        return "已修改密码", "changed"
    return "历史状态未记录", "unknown"


def account_login_state(account: UserAccount | None, employee: Employee) -> tuple[str, str]:
    if employee.account_deleted_at:
        return "账号已删除·留档", "archived"
    if not account:
        return "未开通账号", "unprovisioned"
    if not employee.is_active:
        return "员工已离职（不可登录）", "employee_inactive"
    if not account.enabled:
        return "账号已停用", "disabled"
    if account.locked_until and account.locked_until > datetime.now():
        return f"临时锁定至 {account.locked_until.strftime('%Y-%m-%d %H:%M')}", "locked"
    return "账号已启用", "enabled"


EMPLOYEE_TARGET_PERMISSIONS = {
    "recognition": {"EMPLOYEE_ADD"},
    "deduction": {"DEDUCTION_DIRECT", "DEDUCTION_ALL"},
    "attendance": {"SICK_REGISTER"},
    "circle_transfer": {"HR_MANAGE"},
    "poc": {"POC_ISSUE"},
}


def like_escaped_pattern(keyword: str, escape_char: str = "\\") -> str:
    """Wrap a user keyword in a LIKE pattern with wildcards escaped."""
    value = str(keyword or "").strip()
    escaped = value.replace(escape_char, escape_char * 2).replace("%", escape_char + "%").replace("_", escape_char + "_")
    return f"%{escaped}%"


def search_employee_targets(
    db: Session,
    *,
    keyword: str = "",
    attraction_id: int | None = None,
    limit: int = 30,
) -> dict:
    """Return active, enabled CM/TR targets in one bounded SQLite query."""
    today_value = date.today().isoformat()
    normalized = str(keyword or "").strip()
    escaped = normalized.replace("!", "!!").replace("%", "!%").replace("_", "!_")
    params = {
        "today": today_value,
        "keyword": normalized,
        "pattern": f"%{escaped}%",
        "prefix": f"{escaped}%",
        "attraction_id": attraction_id,
        "limit": max(1, min(int(limit or 30), 500)),
    }
    rows = db.execute(
        text(
            """
            WITH ranked_roles AS (
                SELECT era.employee_id, era.role_id,
                       ROW_NUMBER() OVER (
                           PARTITION BY era.employee_id
                           ORDER BY era.starts_on DESC, era.id DESC
                       ) AS row_number
                FROM employee_role_assignments era
                WHERE era.status != 'cancelled'
                  AND era.starts_on <= :today
                  AND (era.ends_on IS NULL OR era.ends_on >= :today)
            ), ranked_memberships AS (
                SELECT gm.employee_id, gm.group_id,
                       ROW_NUMBER() OVER (
                           PARTITION BY gm.employee_id
                           ORDER BY gm.starts_on DESC, gm.id DESC
                       ) AS row_number
                FROM group_memberships gm
                WHERE gm.status = 'active'
                  AND gm.starts_on <= :today
                  AND (gm.ends_on IS NULL OR gm.ends_on >= :today)
            )
            SELECT e.id, e.employee_no, e.name,
                   r.code AS role_code, r.name AS role_name,
                   e.attraction_id, COALESCE(a.name, '') AS attraction_name,
                   wg.id AS group_id, COALESCE(wg.name, '') AS group_name,
                   COUNT(*) OVER() AS total_count
            FROM employees e
            JOIN ranked_roles rr ON rr.employee_id = e.id AND rr.row_number = 1
            JOIN roles r ON r.id = rr.role_id AND r.active = 1 AND r.code IN ('CM', 'TR')
            JOIN user_accounts ua ON ua.employee_id = e.id AND ua.enabled = 1
            LEFT JOIN attractions a ON a.id = e.attraction_id
            LEFT JOIN ranked_memberships rm ON rm.employee_id = e.id AND rm.row_number = 1
            LEFT JOIN work_groups wg ON wg.id = rm.group_id
            WHERE e.is_active = 1
              AND (:attraction_id IS NULL OR e.attraction_id = :attraction_id)
              AND (
                  :keyword = ''
                  OR e.name LIKE :pattern ESCAPE '!'
                  OR e.employee_no LIKE :pattern ESCAPE '!'
              )
            ORDER BY
                CASE
                    WHEN e.employee_no = :keyword THEN 0
                    WHEN e.name = :keyword THEN 1
                    WHEN e.employee_no LIKE :prefix ESCAPE '!' THEN 2
                    WHEN e.name LIKE :prefix ESCAPE '!' THEN 3
                    ELSE 4
                END,
                e.name ASC, e.employee_no ASC
            LIMIT :limit
            """
        ),
        params,
    ).mappings().all()
    total = int(rows[0]["total_count"]) if rows else 0
    group_display = group_display_metadata_bulk(db, {row["group_id"] for row in rows if row["group_id"]}, today_value)
    return {
        "items": [
            {
                "id": row["id"],
                "employee_no": row["employee_no"],
                "name": row["name"],
                "role_code": row["role_code"],
                "role_name": row["role_name"],
                "attraction_id": row["attraction_id"],
                "attraction_name": row["attraction_name"],
                "group_id": row["group_id"],
                "group_name": group_display.get(row["group_id"], {}).get("name", row["group_name"]),
            }
            for row in rows
        ],
        "total": total,
        "limit": params["limit"],
    }


def ensure_enabled_frontline_target(db: Session, employee_id: int, action_name: str) -> tuple[Employee, Role]:
    target = db.get(Employee, employee_id)
    target_role = role_at(db, employee_id) if target else None
    account_enabled = bool(
        target
        and db.query(UserAccount.id)
        .filter(UserAccount.employee_id == target.id, UserAccount.enabled.is_(True))
        .first()
    )
    if not target or not target.is_active or not account_enabled or not target_role or target_role.code not in FRONTLINE_CODES:
        raise HTTPException(400, f"只能为在职、账号启用的CM/TR登记{action_name}")
    return target, target_role


def ensure_enabled_poc_target(db: Session, employee_id: int) -> tuple[Employee, Role]:
    target = db.get(Employee, employee_id)
    target_role = role_at(db, employee_id) if target else None
    account_enabled = bool(target and db.query(UserAccount.id).filter(UserAccount.employee_id == target.id, UserAccount.enabled.is_(True)).first())
    if not target or not target.is_active or not account_enabled or not target_role or target_role.code not in (FRONTLINE_CODES | LEADER_CODES):
        raise HTTPException(400, "POC仅可认可在职、账号启用的CM/TR、TA主管或主管")
    return target, target_role


def apply_monthly_category_cap(db: Session, row: RecognitionRecord, recognition_type: RecognitionType) -> None:
    """Freeze the actually credited score when a record becomes confirmed."""
    row.credited_fraction = Decimal("0.00")
    row.monthly_cap_rule_code = None
    row.monthly_cap_status = "not_applicable"
    row.monthly_cap_limit = None
    row.monthly_cap_confirmed_before = None
    row.monthly_cap_reason = None
    row.monthly_cap_evaluated_at = datetime.now()
    if recognition_type.code not in MONTHLY_CATEGORY_CAP_CODES or date.fromisoformat(row.recognition_date) < MONTHLY_CATEGORY_CAP_EFFECTIVE_DATE:
        row.credited_fraction = Decimal(row.fraction or 0)
        row.monthly_cap_status = "legacy_not_limited" if date.fromisoformat(row.recognition_date) < MONTHLY_CATEGORY_CAP_EFFECTIVE_DATE else "not_applicable"
        return
    previous = db.query(func.coalesce(func.sum(RecognitionRecord.credited_fraction), 0)).filter(
        RecognitionRecord.employee_id == row.employee_id,
        RecognitionRecord.recognition_month == row.recognition_month,
        RecognitionRecord.recognition_type_id == row.recognition_type_id,
        RecognitionRecord.status == "confirmed",
        RecognitionRecord.id != (row.id or -1),
    ).scalar() or Decimal("0.00")
    previous = Decimal(previous)
    raw = Decimal(row.fraction or 0)
    credit = min(raw, max(Decimal("0.00"), MONTHLY_CATEGORY_CAP_LIMIT - previous))
    row.credited_fraction = credit
    row.monthly_cap_rule_code = "five_category_monthly_5_points"
    row.monthly_cap_limit = MONTHLY_CATEGORY_CAP_LIMIT
    row.monthly_cap_confirmed_before = previous
    if credit == raw:
        row.monthly_cap_status = "within_limit"
    elif credit > 0:
        row.monthly_cap_status = "partially_capped"
        row.monthly_cap_reason = f"本月{recognition_type.name}已确认 {previous:.2f} 分，本条仅计入 {credit:.2f} 分（单项上限5分）"
    else:
        row.monthly_cap_status = "capped_zero"
        row.monthly_cap_reason = f"本月{recognition_type.name}已达5分上限，本条保留认可记录但不再计分"


def effective_recognition_credit(row: RecognitionRecord) -> Decimal:
    """Legacy records before the 2026-09-01 cap retain their raw score."""
    if row.recognition_date and date.fromisoformat(row.recognition_date) < MONTHLY_CATEGORY_CAP_EFFECTIVE_DATE:
        return Decimal(row.fraction or 0)
    return Decimal(row.credited_fraction or 0)


def void_operator_snapshot(db: Session, user: V2User) -> tuple[str, str, str]:
    if user.role.code in FRONTLINE_CODES:
        scope = f"本人账号（{user.employee.employee_no}）"
    elif user.role.code in LEADER_CODES:
        group_names = [group.name for group in groups_led_by(db, user.id)]
        scope = f"直属小组：{'、'.join(group_names)}" if group_names else "直属小组：未配置"
    elif user.role.code in GSM_CODES:
        attraction_ids = sorted(managed_attraction_ids(db, user.id))
        attraction_names = [
            row.name
            for row in db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).order_by(Attraction.name).all()
        ] if attraction_ids else []
        scope = f"景点圈：{'、'.join(attraction_names)}" if attraction_names else "景点圈：未配置"
    elif user.role.code == "HR_ADMIN":
        scope = "HR人员与组织管理"
    else:
        scope = "全局系统权限"
    return user.role.code, user.role.name, scope


def recognition_payload(row: RecognitionRecord) -> dict:
    status_names = {"pending": "待复核", "confirmed": "已确认", "rejected": "不通过", "void": "已撤回"}
    attachment = row.attachments[0] if row.attachments else None
    attachment_file = attachment.file if attachment else None
    return {
        "record_type": "recognition",
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": row.employee_no,
        "employee_name": row.employee_name,
        "recognition_date": row.recognition_date,
        "recognition_type": row.recognition_type_name,
        "content": row.content,
        "recognizer_name": row.recognizer_name,
        "recognizer_role": row.recognizer_role_snapshot,
        "operator_id": row.operator_employee_id,
        "operator_name": row.operator_name,
        "source": row.source,
        "entry_label": f"{row.operator_name}录入" if row.source == "manager" else "",
        "fraction": float(row.fraction or 0),
        "credited_fraction": float(effective_recognition_credit(row)),
        "monthly_cap_status": row.monthly_cap_status or "not_applicable",
        "monthly_cap_reason": row.monthly_cap_reason or "",
        "monthly_cap_limit": float(row.monthly_cap_limit or 0),
        "monthly_cap_confirmed_before": float(row.monthly_cap_confirmed_before or 0),
        "poc_period_type": row.poc_period_type or "",
        "poc_period_key": row.poc_period_key or "",
        "poc_reason": row.poc_reason or "",
        "status": row.status,
        "status_name": status_names.get(row.status, row.status),
        "submitted_at": row.submitted_at.strftime("%Y-%m-%d %H:%M:%S"),
        "review_note": row.review_note or "",
        "same_day_duplicate": bool(row.same_day_duplicate_group),
        "same_day_duplicate_sequence": int(row.same_day_duplicate_sequence or 0),
        "same_day_duplicate_label": (f"今日已有同类登记（第{int(row.same_day_duplicate_sequence or 0)}次）" if row.same_day_duplicate_group else ""),
        "voided_by_name": row.voided_by_name or "",
        "voided_by_role_code": row.voided_by_role_code or "",
        "voided_by_role_name": row.voided_by_role_name or "",
        "void_permission_scope": row.void_permission_scope_snapshot or "",
        "voided_from_status": row.voided_from_status or "",
        "voided_at": row.voided_at.strftime("%Y-%m-%d %H:%M:%S") if row.voided_at else "",
        "void_reason": row.void_reason or "",
        "has_image": bool(attachment),
        "image_url": f"/api/files/{attachment.file_id}" if attachment else "",
        "image_is_previewable": is_previewable_image(attachment_file),
        "image_preview_kind": preview_kind(attachment_file),
        "available_actions": ["withdraw"] if row.status != "void" else [],
    }


def deduction_payload(row: DeductionRecord) -> dict:
    status_names = {"active": "已扣分", "void": "已作废", "pending_upgrade": "待升级审核", "pending_material": "待补充材料（未扣分）", "material_processing": "材料生成中", "material_failed": "材料生成失败"}
    upgrade_state_names = {"pending": "待升级审核", "source_first": "已参与升级", "source_second": "已用于升级，不计分", "result": "升级结果", "rejected": "审核不通过"}
    return {
        "record_type": "deduction",
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": row.employee_no,
        "employee_name": row.employee_name,
        "deduction_type": row.deduction_type_name,
        "deduction_level": row.deduction_level_name,
        "points": float(row.points),
        "actual_points": float(row.points),
        "occurred_on": row.occurred_on,
        "description": row.description,
        "submitter_id": row.submitter_id,
        "submitter_name": row.submitter_name,
        "status": row.status,
        "status_name": status_names.get(row.status, row.status),
        "submitted_at": row.submitted_at.strftime("%Y-%m-%d %H:%M:%S"),
        "void_reason": row.void_reason or "",
        "voided_by_name": row.voided_by_name or "",
        "voided_by_role_code": row.voided_by_role_code or "",
        "voided_by_role_name": row.voided_by_role_name or "",
        "void_permission_scope": row.void_permission_scope_snapshot or "",
        "voided_from_status": row.voided_from_status or "",
        "voided_at": row.voided_at.strftime("%Y-%m-%d %H:%M:%S") if row.voided_at else "",
        "material_status": row.material_status or "ready",
        "material_error": row.material_error or "",
        "material_source_type": row.material_source_type or "pdf",
        "material_job_id": row.material_job_id or 0,
        "material_uploaded_by_name": row.material_uploaded_by_name or "",
        "material_uploaded_at": row.material_uploaded_at.strftime("%Y-%m-%d %H:%M:%S") if row.material_uploaded_at else "",
        "material_revision": int(row.material_revision or 1),
        "legacy_upgrade_excluded": bool(row.legacy_upgrade_excluded),
        "legacy_upgrade_note": row.legacy_upgrade_note or "",
        "document_url": "" if row.upgrade_role == "result" or (row.material_status or "ready") != "ready" else f"/api/files/{row.document_file_id}",
        # Declarations are accepted only as PDF documents when the record is created.
        "document_preview_kind": "" if row.upgrade_role == "result" or (row.material_status or "ready") != "ready" else "pdf",
        "upgrade_request_id": row.upgrade_request_id or 0,
        "upgrade_role": row.upgrade_role or "",
        "upgrade_state": row.upgrade_state or "",
        "upgrade_state_name": upgrade_state_names.get(row.upgrade_state or "", ""),
        "available_actions": (["supplement_material"] if row.status in {"pending_material", "material_failed"} else []) + (["void"] if row.status == "active" and not row.upgrade_request_id else []),
    }


def pending_material_conflict(
    db: Session,
    *,
    employee_id: int,
    deduction_type_id: int,
    deduction_level_id: int,
    occurred_on: str,
) -> DeductionRecord | None:
    """Return the one unfinished declaration that must be completed before another is entered."""
    return (
        db.query(DeductionRecord)
        .filter(
            DeductionRecord.employee_id == employee_id,
            DeductionRecord.deduction_type_id == deduction_type_id,
            DeductionRecord.deduction_level_id == deduction_level_id,
            DeductionRecord.occurred_on == occurred_on,
            DeductionRecord.status.in_({"pending_material", "material_failed"}),
        )
        .order_by(DeductionRecord.submitted_at.asc(), DeductionRecord.id.asc())
        .first()
    )


def raise_pending_material_conflict(row: DeductionRecord) -> None:
    raise HTTPException(
        409,
        detail={
            "code": "PENDING_MATERIAL_EXISTS",
            "message": "该员工同日、同类型、同等级的声明已有待补材料记录，请先补充该记录材料后再登记第二条。",
            "record": deduction_payload(row),
        },
    )


def deduction_follow_up_payload(row: DeductionFollowUp) -> dict:
    status_names = {"pending": "待经理跟进", "issued": "已开具"}
    return {
        "record_type": "follow_up",
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": row.employee_no,
        "employee_name": row.employee_name,
        "deduction_type": row.deduction_type_name,
        "occurred_on": row.occurred_on,
        "status": row.status,
        "status_name": status_names.get(row.status, row.status),
        "submitted_at": row.created_at.strftime("%Y-%m-%d %H:%M:%S"),
        "issued_by_name": row.issued_by_name or "",
        "issued_at": row.issued_at.strftime("%Y-%m-%d %H:%M:%S") if row.issued_at else "",
        "issued_deduction_id": row.issued_deduction_id,
        "available_actions": [],
    }


def governance_case_payload(row: GovernanceCase, *, can_resolve: bool = False) -> dict:
    status_names = {"open": "待处理", "resolved": "已处理", "withdrawn": "已撤回"}
    decision_names = {"uphold": "维持原记录", "correction_required": "需要按受控流程更正", "month_reopened": "已重开月结"}
    return {
        "id": row.id,
        "case_type": row.case_type,
        "record_type": row.record_type or "",
        "record_id": row.record_id or 0,
        "attraction_id": row.attraction_id or 0,
        "score_month": row.score_month or "",
        "reason": row.reason,
        "status": row.status,
        "status_name": status_names.get(row.status, row.status),
        "decision": row.decision or "",
        "decision_name": decision_names.get(row.decision, row.decision or ""),
        "resolution": row.resolution or "",
        "submitted_by_name": row.submitted_by_name,
        "submitted_at": row.submitted_at.strftime("%Y-%m-%d %H:%M:%S"),
        "due_at": row.due_at.strftime("%Y-%m-%d %H:%M:%S"),
        "resolved_by_name": row.resolved_by_name or "",
        "resolved_at": row.resolved_at.strftime("%Y-%m-%d %H:%M:%S") if row.resolved_at else "",
        "can_resolve": can_resolve and row.status == "open",
    }


def governance_scope_allows_case(db: Session, user: V2User, row: GovernanceCase) -> bool:
    if user.role.code in {"SYSTEM_ADMIN", "HR_ADMIN"}:
        return True
    allowed = scoped_hr_attraction_ids(db, user)
    return allowed is not None and row.attraction_id in allowed


def appeal_record_for_employee(db: Session, employee_id: int, record_type: str, record_id: int):
    if record_type == "recognition":
        row = db.get(RecognitionRecord, record_id)
        eligible = row and row.employee_id == employee_id and row.status in {"rejected", "void"}
        if not eligible:
            return None
        return row, row.home_attraction_id, row.recognition_month, {row.operator_employee_id, row.reviewed_by, row.voided_by}
    if record_type == "deduction":
        row = db.get(DeductionRecord, record_id)
        eligible = row and row.employee_id == employee_id and row.status == "active"
        if not eligible:
            return None
        return row, row.attraction_id_snapshot, row.deduction_month, {row.submitter_id, row.voided_by}
    return None


def capture_month_organization_snapshots(db: Session, month: str, attraction_id: int | None) -> int:
    """Freeze organization context once; later transfers cannot rewrite closed reports."""
    employee_ids = [
        int(row[0])
        for row in db.execute(text("SELECT employee_id FROM v_employee_month_scores WHERE score_month=:month"), {"month": month}).all()
    ]
    if not employee_ids:
        return 0
    employees = db.query(Employee).filter(Employee.id.in_(employee_ids)).all()
    captured = 0
    for employee in employees:
        if attraction_id is not None and employee.attraction_id != attraction_id:
            continue
        existing = db.query(EmployeeMonthOrganizationSnapshot.id).filter_by(employee_id=employee.id, score_month=month).first()
        if existing:
            continue
        attraction = db.get(Attraction, employee.attraction_id) if employee.attraction_id else None
        group = current_group_for_employee(db, employee.id)
        leader = current_leader_for_employee(db, employee.id)
        db.add(
            EmployeeMonthOrganizationSnapshot(
                employee_id=employee.id,
                score_month=month,
                attraction_id=employee.attraction_id,
                attraction_name=attraction.name if attraction else "未设置景点圈",
                group_id=group.id if group else None,
                group_name=group.name if group else "未分组",
                leader_employee_id=leader.id if leader else None,
                leader_name=leader.name if leader else "未配置主管",
            )
        )
        captured += 1
    return captured


def ensure_repeat_follow_up(
    db: Session,
    user: V2User,
    employee: Employee,
    deduction_type: DeductionType,
    occurred_on: str,
    repeat_context: dict,
) -> DeductionFollowUp | None:
    if not repeat_context.get("has_repeat") or user.role.code not in LEADER_CODES:
        return None
    row = db.query(DeductionFollowUp).filter_by(
        supervisor_id=user.id,
        employee_id=employee.id,
        deduction_type_id=deduction_type.id,
        occurred_on=occurred_on,
    ).first()
    if row:
        return row
    row = DeductionFollowUp(
        supervisor_id=user.id,
        supervisor_name=user.name,
        employee_id=employee.id,
        employee_no=employee.employee_no,
        employee_name=employee.name,
        deduction_type_id=deduction_type.id,
        deduction_type_name=deduction_type.name,
        occurred_on=occurred_on,
        previous_record_ids=json.dumps([item["id"] for item in repeat_context["previous_records"]]),
        status="pending",
    )
    db.add(row)
    db.flush()
    write_audit(
        db,
        user.employee,
        "生成重复处分跟进",
        "deduction_follow_up",
        row.id,
        after=deduction_follow_up_payload(row),
    )
    return row


def direct_only_deduction_user(user: V2User) -> bool:
    return "DEDUCTION_DIRECT" in user.permissions and "DEDUCTION_ALL" not in user.permissions


def ensure_deduction_type_allowed(user: V2User, deduction_type: DeductionType) -> None:
    if direct_only_deduction_user(user) and deduction_type.code in DIRECT_HIDDEN_DEDUCTION_CODES:
        raise HTTPException(403, "TA主管、主管不能登记迟到30分钟以上或早退30分钟以上")


def attendance_repeat_context(db: Session, employee_id: int, deduction_type: DeductionType, occurred_on: str, user: V2User) -> dict:
    if deduction_type.code not in REPEAT_CONTROLLED_DEDUCTION_CODES:
        return {"has_repeat": False, "blocked": False, "requires_confirmation": False, "previous_records": []}
    event_date = parse_iso_date(occurred_on, "事件日期")
    window_start = subtract_calendar_months(event_date, 3).isoformat()
    rows = (
        db.query(DeductionRecord)
        .filter(
            DeductionRecord.employee_id == employee_id,
            DeductionRecord.deduction_type_id == deduction_type.id,
            DeductionRecord.status == "active",
            DeductionRecord.legacy_upgrade_excluded.is_(False),
            DeductionRecord.occurred_on >= window_start,
            DeductionRecord.occurred_on <= occurred_on,
        )
        .order_by(DeductionRecord.occurred_on.desc(), DeductionRecord.id.desc())
        .all()
    )
    if not rows:
        return {
            "has_repeat": False,
            "blocked": False,
            "requires_confirmation": False,
            "window_start": window_start,
            "previous_records": [],
        }
    highest = max(rows, key=lambda row: DEDUCTION_LEVEL_ORDER.get(db.get(DeductionLevel, row.deduction_level_id).code, 0))
    highest_level = db.get(DeductionLevel, highest.deduction_level_id)
    minimum_code = NEXT_DEDUCTION_LEVEL.get(highest_level.code, "MEMO")
    minimum_level = db.query(DeductionLevel).filter(DeductionLevel.code == minimum_code, DeductionLevel.active.is_(True)).one()
    direct_only = direct_only_deduction_user(user)
    history = [
        {
            "id": row.id,
            "occurred_on": row.occurred_on,
            "deduction_type": row.deduction_type_name,
            "deduction_level": row.deduction_level_name,
            "points": float(row.points),
        }
        for row in rows
    ]
    category_name = "违规病假" if deduction_type.code == "SICK_LEAVE_VIOLATION" else "考勤"
    message = (
        f"此员工在3个月内已有同类{category_name}登记，本次需要升级为备忘录或警告，"
        + ("TA主管/主管不能继续登记，请由GSM登记。" if direct_only else f"本次最低处分等级为{minimum_level.name}，请确认历史记录后继续登记。")
    )
    return {
        "has_repeat": True,
        "blocked": direct_only,
        "requires_confirmation": not direct_only,
        "window_start": window_start,
        "message": message,
        "minimum_level_id": minimum_level.id,
        "minimum_level_code": minimum_level.code,
        "minimum_level_name": minimum_level.name,
        "previous_records": history,
    }


UPGRADE_REVIEWER_CODES = {"GSM", "TA_GSM"}


def statement_upgrade_candidate(db: Session, employee_id: int, deduction_type_id: int, occurred_on: str) -> DeductionRecord | None:
    """Return the oldest still-eligible same-type statement in the 3-month window."""
    statement = db.query(DeductionLevel).filter_by(code="STATEMENT", active=True).first()
    if not statement:
        return None
    window_start = subtract_calendar_months(parse_iso_date(occurred_on, "事件日期"), 3).isoformat()
    return (
        db.query(DeductionRecord)
        .filter(
            DeductionRecord.employee_id == employee_id,
            DeductionRecord.deduction_type_id == deduction_type_id,
            DeductionRecord.deduction_level_id == statement.id,
            DeductionRecord.status == "active",
            DeductionRecord.legacy_upgrade_excluded.is_(False),
            or_(DeductionRecord.upgrade_state.is_(None), DeductionRecord.upgrade_state == "eligible"),
            DeductionRecord.occurred_on >= window_start,
            DeductionRecord.occurred_on <= occurred_on,
        )
        .order_by(DeductionRecord.occurred_on.asc(), DeductionRecord.id.asc())
        .first()
    )


def statement_upgrade_reviewer_options(db: Session) -> list[dict]:
    rows = (
        db.query(Employee, UserAccount)
        .join(UserAccount, UserAccount.employee_id == Employee.id)
        .filter(Employee.is_active.is_(True), UserAccount.enabled.is_(True))
        .order_by(Employee.name, Employee.employee_no)
        .all()
    )
    result = []
    for employee, _account in rows:
        role = role_at(db, employee.id)
        if role and role.code in UPGRADE_REVIEWER_CODES:
            result.append({"id": employee.id, "name": employee.name, "employee_no": employee.employee_no, "role_code": role.code, "role_name": role.name})
    return result


def deduction_upgrade_payload(db: Session, row: DeductionUpgradeRequest, *, include_details: bool = True) -> dict:
    first = db.get(DeductionRecord, row.first_deduction_id)
    second = db.get(DeductionRecord, row.second_deduction_id)
    result = db.get(DeductionRecord, row.result_deduction_id) if row.result_deduction_id else None
    transfer_rows = db.query(DeductionUpgradeTransfer).filter_by(request_id=row.id).order_by(DeductionUpgradeTransfer.created_at.asc()).all()
    status_names = {"pending": "待审核", "approved": "已升级", "rejected": "审核不通过"}
    payload = {
        "id": row.id,
        "record_type": "deduction_upgrade",
        "employee_id": row.employee_id,
        "employee_name": first.employee_name if first else "",
        "employee_no": first.employee_no if first else "",
        "deduction_type": first.deduction_type_name if first else "",
        "status": row.status,
        "status_name": status_names.get(row.status, row.status),
        "reviewer_id": row.reviewer_id,
        "reviewer_name": row.reviewer_name,
        "submitted_by": row.submitted_by_name,
        "created_at": row.created_at.strftime("%Y-%m-%d %H:%M:%S"),
        "handling_note": row.handling_note or "",
        "result_level": result.deduction_level_name if result else "",
        "result_points": float(result.points) if result else 0,
        "transfers": [{"from_name": item.from_reviewer_name, "to_name": item.to_reviewer_name, "reason": item.reason, "created_at": item.created_at.strftime("%Y-%m-%d %H:%M:%S")} for item in transfer_rows],
    }
    if include_details:
        payload["first_record"] = deduction_payload(first) if first else None
        payload["second_record"] = deduction_payload(second) if second else None
    return payload


def sick_leave_payload(db: Session, row: SickLeaveRecord, *, employees: dict[int, Employee] | None = None, files: dict[int, StoredFile] | None = None) -> dict:
    employee = (employees or {}).get(row.employee_id) or db.get(Employee, row.employee_id)
    proof_file = (files or {}).get(row.proof_file_id) or db.get(StoredFile, row.proof_file_id)
    status_names = {"active": "已生效", "void": "已作废"}
    return {
        "record_type": "sick_leave",
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": row.employee_no_snapshot or (employee.employee_no if employee else ""),
        "employee_name": row.employee_name_snapshot or (employee.name if employee else "未知员工"),
        "attendance_month": row.attendance_month,
        "leave_start_date": row.leave_start_date,
        "leave_end_date": row.leave_end_date,
        "leave_days": float(row.leave_days),
        "charged_days": float(row.charged_days),
        "note": row.note or "",
        "submitter_id": row.submitted_by,
        "submitter_name": row.submitted_by_name,
        "status": row.status,
        "status_name": status_names.get(row.status, row.status),
        "submitted_at": row.submitted_at.strftime("%Y-%m-%d %H:%M:%S"),
        "void_reason": row.void_reason or "",
        "voided_by_name": row.voided_by_name or "",
        "voided_by_role_code": row.voided_by_role_code or "",
        "voided_by_role_name": row.voided_by_role_name or "",
        "void_permission_scope": row.void_permission_scope_snapshot or "",
        "voided_from_status": row.voided_from_status or "",
        "voided_at": row.voided_at.strftime("%Y-%m-%d %H:%M:%S") if row.voided_at else "",
        "proof_url": f"/api/files/{row.proof_file_id}",
        "proof_is_previewable": is_previewable_image(proof_file),
        "proof_preview_kind": preview_kind(proof_file),
        "is_violation": bool(row.is_violation),
        "violation_deduction_id": row.violation_deduction_id or 0,
        "available_actions": ["void"] if row.status == "active" else [],
    }


def sick_leave_payloads(db: Session, rows: list[SickLeaveRecord]) -> list[dict]:
    """Serialize many sick-leave rows with batched employee/proof lookups."""
    employee_ids = {row.employee_id for row in rows}
    file_ids = {row.proof_file_id for row in rows}
    employees = (
        {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(employee_ids)).all()}
        if employee_ids
        else {}
    )
    files = (
        {file_row.id: file_row for file_row in db.query(StoredFile).filter(StoredFile.id.in_(file_ids)).all()}
        if file_ids
        else {}
    )
    return [sick_leave_payload(db, row, employees=employees, files=files) for row in rows]


@router.post("/login")
def login(payload: dict, request: Request, response: Response, db: Session = Depends(get_db)):
    login_name = str(payload.get("employee_no") or "").strip()
    password = str(payload.get("password") or "")
    account = (
        db.query(UserAccount)
        .join(Employee, Employee.id == UserAccount.employee_id)
        .filter(or_(UserAccount.login_account == login_name, Employee.employee_no == login_name))
        .first()
    )
    if not account or not account.enabled:
        raise HTTPException(401, "账号或密码/PIN不正确")
    if account.locked_until and account.locked_until > datetime.now():
        raise HTTPException(423, "登录失败次数过多，请稍后再试")
    if not verify_password(password, account.password_hash):
        account.failed_attempts += 1
        if account.failed_attempts >= 5:
            account.locked_until = datetime.now() + timedelta(minutes=15)
            account.failed_attempts = 0
        db.commit()
        raise HTTPException(401, "账号或密码/PIN不正确")
    employee = account.employee
    if not employee.is_active:
        raise HTTPException(403, "员工账号已停用")
    role = role_at(db, employee.id)
    if not role:
        raise HTTPException(403, "当前未配置有效角色")
    account.failed_attempts = 0
    account.locked_until = None
    account.last_login_at = datetime.now()
    raw_token = new_session_token()
    db.add(
        UserSession(
            account_id=account.id,
            token_hash=token_hash(raw_token),
            expires_at=datetime.now() + timedelta(hours=12),
        )
    )
    write_audit(db, employee, "登录", "user_account", account.id, ip_address=client_ip(request))
    db.commit()
    response.set_cookie(
        "rc_v2_session",
        raw_token,
        httponly=True,
        samesite="lax",
        secure=request_is_https(request),
        max_age=12 * 3600,
    )
    return {"ok": True, "role": role.name}


@router.post("/logout")
def logout(request: Request, response: Response, db: Session = Depends(get_db)):
    raw_token = request.cookies.get("rc_v2_session")
    if raw_token:
        db.query(UserSession).filter(UserSession.token_hash == token_hash(raw_token)).delete(synchronize_session=False)
        db.commit()
    response.delete_cookie("rc_v2_session", secure=request_is_https(request), httponly=True, samesite="lax")
    return {"ok": True}


@router.get("/me")
def me(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    return user_payload(db, user)


@router.get("/changelog")
def changelog(user: V2User = Depends(current_user)):
    return {
        "app_version": APP_VERSION,
        "role_code": user.role.code,
        "role_name": user.role.name,
        "releases": visible_releases(user.role.code, user.permissions),
    }


@router.post("/password")
def change_password(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    current = str(payload.get("current_password") or "")
    new = str(payload.get("new_password") or "")
    confirm = str(payload.get("confirm_password") or "")
    if not verify_password(current, user.account.password_hash):
        raise HTTPException(400, "当前密码/PIN不正确")
    if new != confirm:
        raise HTTPException(400, "两次输入的新密码必须一致")
    policy_error = password_policy_error(new, user.role.code)
    if policy_error:
        raise HTTPException(400, policy_error)
    user.account.password_hash = hash_password(new)
    user.account.must_change_password = False
    user.account.credential_initialized = True
    user.account.password_changed_at = datetime.now()
    current_raw_token = request.cookies.get("rc_v2_session")
    current_token_hash = token_hash(current_raw_token) if current_raw_token else ""
    db.query(UserSession).filter(
        UserSession.account_id == user.account.id,
        UserSession.token_hash != current_token_hash,
    ).delete(synchronize_session=False)
    write_audit(db, user.employee, "修改密码", "user_account", user.account.id, ip_address=client_ip(request))
    db.commit()
    return {"ok": True}


@router.get("/options")
def options(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    recognition_types = db.query(RecognitionType).filter(RecognitionType.active.is_(True)).order_by(RecognitionType.id).all()
    deduction_types = db.query(DeductionType).filter(DeductionType.active.is_(True)).all()
    deduction_order = {code: index for index, code in enumerate(DEDUCTION_TYPE_ORDER)}
    deduction_types.sort(key=lambda row: (deduction_order.get(row.code, len(deduction_order)), row.id))
    if direct_only_deduction_user(user):
        deduction_types = [row for row in deduction_types if row.code not in DIRECT_HIDDEN_DEDUCTION_CODES]
    role_query = db.query(Role).filter(Role.active.is_(True))
    if "HR_MANAGE" in user.permissions and "SYSTEM_ADMIN" not in user.permissions:
        role_query = role_query.filter(Role.code.in_(CIRCLE_HR_MANAGED_ROLE_CODES))
    circle_rows = (
        db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .all()
    )
    venue_rows = (
        db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.recognition_venue.is_(True))
        .all()
    )
    circles_by_name = {row.name: row for row in circle_rows}
    venues_by_name = {row.name: row for row in venue_rows}
    circles = [circles_by_name[name] for name in EMPLOYEE_CIRCLES if name in circles_by_name]
    venues = [venues_by_name[name] for name in RECOGNITION_VENUES if name in venues_by_name]
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is not None:
        circles = [row for row in circles if row.id in allowed_attractions]
        allowed_circle_names = {row.name for row in circles}
        venues = [row for row in venues if LEGACY_CIRCLE_BY_VENUE.get(row.name, row.name) in allowed_circle_names]
    circle_payload = [{"id": row.id, "name": row.name} for row in circles]
    return {
        "roles": [{"id": row.id, "code": row.code, "name": row.name} for row in role_query.order_by(Role.rank).all()],
        # Keep the legacy key for existing management screens; it now contains circles only.
        "attractions": circle_payload,
        "employee_circles": circle_payload,
        "recognition_venues": [{"id": row.id, "name": row.name} for row in venues],
        "recognition_types": [
            {
                "id": row.id,
                "code": row.code,
                "name": row.name,
                "fixed_score": float(SPECIAL_RECOGNITION_TYPES[row.code]["score"]) if row.code in SPECIAL_RECOGNITION_TYPES else None,
                "monthly_limit": SPECIAL_RECOGNITION_TYPES[row.code]["monthly_limit"] if row.code in SPECIAL_RECOGNITION_TYPES else None,
            }
            for row in recognition_types
        ],
        "deduction_types": [
            {
                "id": row.id,
                "code": row.code,
                "name": row.name,
                "repeat_check": row.code in REPEAT_CONTROLLED_DEDUCTION_CODES,
            }
            for row in deduction_types
        ],
        "deduction_levels": [{"id": row.id, "code": row.code, "name": row.name, "points": float(row.points)} for row in db.query(DeductionLevel).filter(DeductionLevel.active.is_(True)).order_by(DeductionLevel.points).all()],
    }


def employee_circle_id(db: Session, value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        attraction_id = int(value)
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "员工景点圈无效") from exc
    circle = db.get(Attraction, attraction_id)
    if not circle or not circle.active or not circle.employee_circle:
        raise HTTPException(400, "员工只能归属有效景点圈")
    return circle.id


@router.get("/recognizers")
def get_recognizers(attraction_id: int, recognition_date: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    circle = db.get(Attraction, attraction_id)
    if not circle or not circle.active or not circle.employee_circle:
        raise HTTPException(400, "请选择有效员工景点圈")
    ensure_scoped_hr_attraction(db, user, attraction_id)
    if recognition_date:
        parse_iso_date(recognition_date, "认可日期")
    rows = recognizer_options(db, attraction_id, recognition_date)
    rows.extend(
        {
            "id": rule["option_id"],
            "employee_no": "",
            "name": rule["name"],
            "role_code": code,
            "role_name": "系统固定分值",
            "score": float(rule["score"]),
            "special": True,
        }
        for code, rule in SPECIAL_RECOGNITION_TYPES.items()
    )
    return rows


@router.get("/frontline-employees")
def frontline_employees(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if not ({"EMPLOYEE_ADD", "SICK_REGISTER", "DEDUCTION_ALL", "DEDUCTION_DIRECT"} & user.permissions):
        raise HTTPException(403, "没有权限查询CM/TR")
    allowed = scoped_hr_attraction_ids(db, user)
    attraction_id = next(iter(allowed)) if allowed else None
    return search_employee_targets(db, attraction_id=attraction_id, limit=500)["items"]


@router.get("/employee-targets")
def employee_targets(
    usage: str,
    keyword: str = "",
    attraction_id: int | None = None,
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    required = EMPLOYEE_TARGET_PERMISSIONS.get(usage)
    if not required:
        raise HTTPException(400, "员工查询用途无效")
    if not (required & user.permissions):
        raise HTTPException(403, "没有对应的员工登记权限")
    if usage == "poc":
        if user.role.code not in {"TA_GSM", "GSM", "AM"}:
            raise HTTPException(403, "仅TA GSM、GSM、AM可以查询POC被认可员工")
        value = like_escaped_pattern(keyword)
        query = db.query(Employee).join(UserAccount, UserAccount.employee_id == Employee.id).filter(
            Employee.is_active.is_(True), UserAccount.enabled.is_(True),
            or_(Employee.name.like(value, escape="\\"), Employee.employee_no.like(value, escape="\\")),
        )
        if attraction_id is not None:
            query = query.filter(Employee.attraction_id == attraction_id)
        rows = []
        for item in query.order_by(Employee.name, Employee.employee_no).limit(max(1, min(limit, 50)) * 4).all():
            role = role_at(db, item.id)
            if role and role.code in (FRONTLINE_CODES | LEADER_CODES):
                circle = db.get(Attraction, item.attraction_id) if item.attraction_id else None
                rows.append({"id": item.id, "employee_no": item.employee_no, "name": item.name, "role_code": role.code, "role_name": role.name, "attraction_id": item.attraction_id, "attraction_name": circle.name if circle else "", "group_id": None, "group_name": ""})
        return {"items": rows[:max(1, min(limit, 50))], "total": len(rows), "limit": max(1, min(limit, 50))}
    # Deduction/absence target searches are intentionally keyword-only for
    # cross-circle managers.  TA主管 retains its home-circle boundary, while
    # 主管、TA GSM、GSM may locate active CM/TR across all circles.
    global_target_search = usage in {"deduction", "attendance"} and user.role.code in {"SUPERVISOR", "TA_GSM", "GSM"}
    if global_target_search and not keyword.strip():
        return {"items": [], "total": 0, "limit": max(1, min(limit, 50)), "search_scope": "全部景点圈在职CM/TR（请输入姓名或员工号）"}
    if user.role.code == "TA_SUPERVISOR" and usage in {"deduction", "attendance"}:
        attraction_id = user.employee.attraction_id
    if attraction_id is not None:
        circle = db.get(Attraction, attraction_id)
        if not circle or not circle.active or not circle.employee_circle:
            raise HTTPException(400, "请选择有效员工景点圈")
    elif not global_target_search:
        allowed = scoped_hr_attraction_ids(db, user)
        attraction_id = next(iter(allowed)) if allowed else None
    if not global_target_search:
        ensure_scoped_hr_attraction(db, user, attraction_id)
    result = search_employee_targets(
        db,
        keyword=keyword,
        attraction_id=attraction_id,
        limit=max(1, min(limit, 50)),
    )
    result["search_scope"] = "全部景点圈在职CM/TR（请输入姓名或员工号）" if global_target_search else "本景点圈在职CM/TR"
    return result


@router.get("/dashboard")
def dashboard(month: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    month = month or date.today().strftime("%Y-%m")
    if user.role.code not in FRONTLINE_CODES:
        return {"month": month, "role": user.role.name}
    recalculate_attendance(db, user.employee, month)
    db.commit()
    score = db.execute(
        text("SELECT recognition_score, attendance_score, deduction_score, total_score FROM v_employee_month_scores WHERE employee_id=:employee_id AND score_month=:month"),
        {"employee_id": user.id, "month": month},
    ).mappings().first()
    score = score or {"recognition_score": 0, "attendance_score": 0, "deduction_score": 0, "total_score": 0}
    categories = {
        name: float(total or 0)
        for name, total in db.query(RecognitionRecord.recognition_type_name, text("SUM(CASE WHEN recognition_date < '2026-09-01' THEN fraction ELSE credited_fraction END)"))
        .filter(
            RecognitionRecord.employee_id == user.id,
            RecognitionRecord.recognition_month == month,
            RecognitionRecord.status == "confirmed",
        )
        .group_by(RecognitionRecord.recognition_type_name)
        .all()
    }
    records = (
        db.query(RecognitionRecord)
        .filter(
            RecognitionRecord.employee_id == user.id,
            RecognitionRecord.recognition_month == month,
            RecognitionRecord.status != "void",
        )
        .order_by(RecognitionRecord.submitted_at.desc(), RecognitionRecord.id.desc())
        .all()
    )
    records.sort(key=lambda row: row.status == "confirmed")
    return {
        "month": month,
        "category_scores": categories,
        "recognition_score": float(score["recognition_score"] or 0),
        "attendance_score": float(score["attendance_score"] or 0),
        "deduction_score": float(score["deduction_score"] or 0),
        "total_score": float(score["total_score"] or 0),
        "records": [recognition_payload(row) for row in records],
    }


@router.post("/recognitions")
async def create_recognition(
    request: Request,
    recognition_date: str = Form(...),
    occurred_attraction_id: int = Form(...),
    recognition_type_id: int = Form(...),
    recognizer_employee_id: str = Form(...),
    content: str = Form(...),
    employee_id: int | None = Form(None),
    idempotency_key: str | None = Form(None),
    same_day_duplicate_confirmed: bool = Form(False),
    image: UploadFile | None = File(None),
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    request_key = normalize_request_key(idempotency_key)
    payload_digest = submission_payload_digest(
        {
            "employee_id": int(employee_id or 0),
            "recognition_date": str(recognition_date or "").strip(),
            "occurred_attraction_id": int(occurred_attraction_id or 0),
            "recognition_type_id": int(recognition_type_id or 0),
            "recognizer_employee_id": str(recognizer_employee_id or "").strip(),
            "content": str(content or "").strip(),
            "same_day_duplicate_confirmed": bool(same_day_duplicate_confirmed),
        }
    )
    duplicate_row = existing_submission(db, user.id, "recognition", request_key, RecognitionRecord, payload_digest)
    if duplicate_row:
        return {"ok": True, "record": recognition_payload(duplicate_row), "duplicate": True}
    recognition_date = str(recognition_date or "")
    parse_iso_date(recognition_date, "认可日期")
    content = str(content or "").strip()
    if not content or len(content) > 20:
        raise HTTPException(400, "认可内容必填且不能超过20字")
    target_id = int(employee_id or user.id)
    target, target_role = ensure_enabled_frontline_target(db, target_id, "加分")
    ensure_scoped_hr_employee(db, user, target)
    is_self = target.id == user.id and user.role.code in FRONTLINE_CODES
    if not is_self and "EMPLOYEE_ADD" not in user.permissions:
        raise HTTPException(403, "没有员工加分权限")
    attraction = db.get(Attraction, int(occurred_attraction_id or 0))
    if not attraction or not attraction.active or not attraction.recognition_venue:
        raise HTTPException(400, "请选择有效认可发生景点")
    target_circle = db.get(Attraction, target.attraction_id) if target.attraction_id else None
    if not target_circle or not target_circle.active or not target_circle.employee_circle:
        raise HTTPException(400, "被加分员工未配置有效景点圈")
    ensure_month_open(db, recognition_date[:7], target_circle.id, "新增签卡")
    if user.role.code == SCOPED_HR_ROLE_CODE and LEGACY_CIRCLE_BY_VENUE.get(attraction.name, attraction.name) != target_circle.name:
        raise HTTPException(403, "景点圈HR只能登记所属景点圈发生的认可")
    recognition_type = db.get(RecognitionType, int(recognition_type_id or 0))
    if not recognition_type or not recognition_type.active:
        raise HTTPException(400, "请选择有效认可类型")
    selected_recognizer = str(recognizer_employee_id or "").strip()
    special_rule = SPECIAL_RECOGNITION_TYPES.get(recognition_type.code)
    recognizer_role = None
    if special_rule:
        if selected_recognizer != special_rule["option_id"]:
            raise HTTPException(400, f"认可类型为{recognition_type.name}时，认可人/签卡人必须选择{recognition_type.name}")
        recognizer = user.employee
        recognizer_name = special_rule["name"]
        recognizer_role_name = "系统固定分值"
        fraction = special_rule["score"]
    else:
        if selected_recognizer.startswith("special:"):
            raise HTTPException(400, "普通认可类型必须选择实际认可人/签卡人")
        try:
            recognizer_id = int(selected_recognizer)
        except (TypeError, ValueError) as exc:
            raise HTTPException(400, "请选择有效认可人") from exc
        recognizer = db.get(Employee, recognizer_id)
        recognizer_role = recognizer_role_for_date(db, recognizer.id, recognition_date) if recognizer else None
        if not recognizer or not recognizer.is_active or not recognizer_role or recognizer_role.code not in RECOGNIZER_CODES:
            raise HTTPException(400, "请选择有效认可人")
        allowed_ids = {row["id"] for row in recognizer_options(db, target_circle.id, recognition_date)}
        if recognizer.id not in allowed_ids:
            raise HTTPException(400, "认可人必须是TALEAD/LEAD或TAGSM及以上的非HR在职人员")
        recognizer_name = recognizer.name
        recognizer_role_name = recognizer_role.name
        fraction = recognition_score_for_role(db, recognizer_role.id, recognition_date)
    quota_code = recognition_type.code if special_rule and special_rule["monthly_limit"] else None
    active_same_day_rows = (
        db.query(RecognitionRecord)
        .filter(
            RecognitionRecord.employee_id == target.id,
            RecognitionRecord.recognition_date == recognition_date,
            RecognitionRecord.recognition_type_id == recognition_type.id,
            RecognitionRecord.recognizer_employee_id == recognizer.id,
            RecognitionRecord.status.in_(("pending", "confirmed")),
        )
        .order_by(RecognitionRecord.submitted_at.asc(), RecognitionRecord.id.asc())
        .all()
    )
    duplicate_group = None
    duplicate_sequence = None
    if active_same_day_rows:
        duplicate_group = f"{recognition_date}:{target.id}:{recognition_type.id}:{recognizer.id}"
        if not same_day_duplicate_confirmed:
            # A concurrent retry can reach the same-day reminder after the
            # first request has committed.  Re-check the idempotency receipt so
            # the caller receives its original record rather than a false
            # duplicate warning.
            duplicate_row = existing_submission(db, user.id, "recognition", request_key, RecognitionRecord)
            if duplicate_row:
                return {"ok": True, "record": recognition_payload(duplicate_row), "duplicate": True}
            raise HTTPException(
                409,
                {
                    "code": "SAME_DAY_RECOGNITION_DUPLICATE",
                    "message": "该员工今日已由该认可人登记过同类认可，请确认是否为另一项独立表现。",
                    "previous_records": [
                        {
                            "id": item.id,
                            "submitted_at": item.submitted_at.strftime("%Y-%m-%d %H:%M"),
                            "status": item.status,
                            "content": item.content,
                        }
                        for item in active_same_day_rows
                    ],
                },
            )
        for index, item in enumerate(active_same_day_rows, start=1):
            item.same_day_duplicate_group = duplicate_group
            item.same_day_duplicate_sequence = index
        duplicate_sequence = len(active_same_day_rows) + 1
    if quota_code and db.query(RecognitionMonthlyQuota).filter_by(
        employee_id=target.id, quota_code=quota_code, quota_month=recognition_date[:7]
    ).first():
        raise HTTPException(409, "该员工本月已经登记过表扬信，每名员工每月只能获得一次表扬信加分")
    reviewer = current_leader_for_employee(db, target.id) if is_self else None
    if is_self and (not image or not image.filename):
        raise HTTPException(400, "CM/TR本人登记必须上传1张认可图片")
    file_row = None
    row = RecognitionRecord(
        employee_id=target.id,
        employee_no=target.employee_no,
        employee_name=target.name,
        employee_role_snapshot=target_role.name,
        employee_role_code_snapshot=target_role.code,
        home_attraction_id=target.attraction_id,
        home_attraction_name=target.attraction.name if target.attraction else "",
        occurred_attraction_id=attraction.id,
        recognition_date=recognition_date,
        recognition_month=recognition_date[:7],
        recognition_type_id=recognition_type.id,
        recognition_type_name=recognition_type.name,
        content=content,
        recognizer_employee_id=recognizer.id,
        recognizer_name=recognizer_name,
        recognizer_role_snapshot=recognizer_role_name,
        recognizer_role_code_snapshot=recognizer_role.code if recognizer_role else None,
        operator_employee_id=user.id,
        operator_name=user.name,
        operator_role_snapshot=user.role.name,
        operator_role_code_snapshot=user.role.code,
        source="self" if is_self else "manager",
        fraction=fraction,
        status="pending" if is_self else "confirmed",
        same_day_duplicate_group=duplicate_group,
        same_day_duplicate_sequence=duplicate_sequence,
        assigned_reviewer_id=reviewer.id if reviewer else None,
    )
    try:
        db.add(row)
        db.flush()
        if row.status == "confirmed":
            apply_monthly_category_cap(db, row, recognition_type)
            db.flush()
        if quota_code:
            db.add(
                RecognitionMonthlyQuota(
                    employee_id=target.id,
                    quota_code=quota_code,
                    quota_month=recognition_date[:7],
                    recognition_id=row.id,
                )
            )
            db.flush()
        if is_self and image:
            file_row = await save_image_upload(db, image, user.id)
            row.attachments.append(RecognitionAttachment(file_id=file_row.id, attachment_type="evidence", sort_order=1))
            db.flush()
        remember_submission(db, user.id, "recognition", request_key, row.id, payload_digest)
        audit_payload = recognition_payload(row)
        if file_row:
            audit_payload["image"] = {"sha256": file_row.sha256, "size": file_row.file_size, "type": file_row.mime_type}
        write_audit(db, user.employee, "登记签卡", "recognition", row.id, after=audit_payload, ip_address=client_ip(request))
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        if file_row:
            path = (FILE_DIR / file_row.storage_key).resolve()
            if FILE_DIR.resolve() in path.parents:
                path.unlink(missing_ok=True)
        duplicate_row = existing_submission(db, user.id, "recognition", request_key, RecognitionRecord, payload_digest)
        if duplicate_row:
            return {"ok": True, "record": recognition_payload(duplicate_row), "duplicate": True}
        if quota_code:
            raise HTTPException(409, "该员工本月已经登记过表扬信，每名员工每月只能获得一次表扬信加分") from exc
        raise
    except Exception:
        db.rollback()
        if file_row:
            path = (FILE_DIR / file_row.storage_key).resolve()
            if FILE_DIR.resolve() in path.parents:
                path.unlink(missing_ok=True)
        raise
    invalidate_data_caches()
    response = {"ok": True, "record": recognition_payload(row)}
    if is_self:
        response["encouragement_options"] = encouragement_options(
            record_id=row.id,
            recognition_type_name=recognition_type.name,
            content=row.content,
            stage="submitted",
        )
    return response


@router.post("/recognitions/poc")
def create_poc_recognition(
    request: Request,
    recognition_date: str = Form(...),
    employee_id: int = Form(...),
    points: str = Form(...),
    poc_period_type: str = Form(...),
    poc_reason: str = Form(...),
    idempotency_key: str | None = Form(None),
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("POC_ISSUE")),
):
    if user.role.code not in {"TA_GSM", "GSM", "AM"}:
        raise HTTPException(403, "仅TA GSM、GSM、AM可以开具POC特别贡献")
    request_key = normalize_request_key(idempotency_key)
    payload_digest = submission_payload_digest(
        {
            "employee_id": int(employee_id),
            "recognition_date": str(recognition_date or "").strip(),
            "points": str(points or "").strip(),
            "poc_period_type": str(poc_period_type or "").strip(),
            "poc_reason": str(poc_reason or "").strip(),
        }
    )
    duplicate_row = existing_submission(db, user.id, "poc", request_key, RecognitionRecord, payload_digest)
    if duplicate_row:
        return {"ok": True, "record": recognition_payload(duplicate_row), "duplicate": True}
    recognition_date = str(recognition_date or "")
    parse_iso_date(recognition_date, "认可日期")
    try:
        amount = Decimal(str(points))
    except InvalidOperation as exc:
        raise HTTPException(400, "POC分值必须为1至5的整数") from exc
    if amount not in {Decimal("1"), Decimal("2"), Decimal("3"), Decimal("4"), Decimal("5")}:
        raise HTTPException(400, "POC分值必须为1至5分")
    if poc_period_type not in {"month", "quarter"}:
        raise HTTPException(400, "请选择月度或季度认可周期")
    reason = str(poc_reason or "").strip()
    if not reason or len(reason) > 100:
        raise HTTPException(400, "特别贡献原因必填且不能超过100字")
    target, target_role = ensure_enabled_poc_target(db, int(employee_id))
    target_circle = db.get(Attraction, target.attraction_id) if target.attraction_id else None
    if not target_circle or not target_circle.employee_circle:
        raise HTTPException(400, "被认可员工未配置有效景点圈")
    ensure_month_open(db, recognition_date[:7], target_circle.id, "新增POC特别贡献")
    recognition_type = db.query(RecognitionType).filter_by(code="POC", active=True).first()
    if not recognition_type:
        raise HTTPException(500, "POC认可类型尚未初始化")
    parsed = date.fromisoformat(recognition_date)
    period_key = f"{parsed.year}-{parsed.month:02d}" if poc_period_type == "month" else f"{parsed.year}-Q{((parsed.month - 1) // 3) + 1}"
    row = RecognitionRecord(
        employee_id=target.id, employee_no=target.employee_no, employee_name=target.name,
        employee_role_snapshot=target_role.name, employee_role_code_snapshot=target_role.code,
        home_attraction_id=target.attraction_id, home_attraction_name=target.attraction.name if target.attraction else "",
        occurred_attraction_id=target_circle.id, recognition_date=recognition_date, recognition_month=recognition_date[:7],
        recognition_type_id=recognition_type.id, recognition_type_name=recognition_type.name, content="POC特别贡献",
        recognizer_employee_id=user.id, recognizer_name=user.name, recognizer_role_snapshot=user.role.name,
        recognizer_role_code_snapshot=user.role.code, operator_employee_id=user.id, operator_name=user.name,
        operator_role_snapshot=user.role.name, operator_role_code_snapshot=user.role.code, source="manager",
        fraction=amount, credited_fraction=amount, monthly_cap_status="not_applicable", status="confirmed",
        poc_period_type=poc_period_type, poc_period_key=period_key, poc_reason=reason,
    )
    try:
        db.add(row)
        db.flush()
        remember_submission(db, user.id, "poc", request_key, row.id, payload_digest)
        write_audit(db, user.employee, "登记POC特别贡献", "recognition", row.id, after=recognition_payload(row), ip_address=client_ip(request))
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        duplicate_row = existing_submission(db, user.id, "poc", request_key, RecognitionRecord, payload_digest)
        if duplicate_row:
            return {"ok": True, "record": recognition_payload(duplicate_row), "duplicate": True}
        raise HTTPException(409, "POC特别贡献提交冲突，请刷新后重试") from exc
    invalidate_data_caches()
    return {"ok": True, "record": recognition_payload(row)}


@router.delete("/recognitions/{record_id}")
def withdraw_recognition(record_id: int, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    row = db.get(RecognitionRecord, record_id)
    if not row:
        raise HTTPException(404, "签卡记录不存在")
    if user.id not in (row.employee_id, row.operator_employee_id) and "SYSTEM_ADMIN" not in user.permissions:
        raise HTTPException(403, "只有CM/TR本人、原代录人或管理员可以撤回")
    if row.status == "void":
        raise HTTPException(400, "签卡记录已经撤回")
    employee = db.get(Employee, row.employee_id)
    ensure_month_open(db, row.recognition_month, row.home_attraction_id or (employee.attraction_id if employee else None), "撤回签卡")
    before = recognition_payload(row)
    role_code, role_name, permission_scope = void_operator_snapshot(db, user)
    row.voided_from_status = row.status
    row.status = "void"
    row.voided_by = user.id
    row.voided_by_name = user.name
    row.voided_by_role_code = role_code
    row.voided_by_role_name = role_name
    row.void_permission_scope_snapshot = permission_scope
    row.voided_at = datetime.now()
    row.void_reason = "用户撤回"
    db.query(RecognitionMonthlyQuota).filter(RecognitionMonthlyQuota.recognition_id == row.id).delete(synchronize_session=False)
    db.flush()
    write_audit(
        db,
        user.employee,
        "撤回签卡（留档）",
        "recognition",
        row.id,
        before=before,
        after=recognition_payload(row),
        reason=row.void_reason,
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "deducted_score": before["credited_fraction"] if before["status"] == "confirmed" else 0}


@router.get("/entered-recognitions")
def entered_recognitions(month: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if "EMPLOYEE_ADD" not in user.permissions and "SYSTEM_ADMIN" not in user.permissions:
        raise HTTPException(403, "没有代录查询权限")
    query = db.query(RecognitionRecord).filter(
        RecognitionRecord.operator_employee_id == user.id,
        RecognitionRecord.status != "void",
    )
    if month:
        query = query.filter(RecognitionRecord.recognition_month == month)
    return [recognition_payload(row) for row in query.order_by(RecognitionRecord.submitted_at.desc()).limit(500).all()]


def action_center_item(item_type: str, title: str, count: int, tab: str, severity: str, description: str) -> dict:
    return {
        "type": item_type,
        "title": title,
        "count": int(count),
        "tab": tab,
        "severity": severity,
        "description": description,
    }


@router.get("/action-center")
def action_center(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    """A read-only queue built from existing workflow records and role scopes."""
    items: list[dict] = []
    role_code = user.role.code
    current_month = date.today().strftime("%Y-%m")
    month_to_close = (date.today().replace(day=1) - timedelta(days=1)).strftime("%Y-%m")

    if role_code in FRONTLINE_CODES:
        open_appeals = db.query(GovernanceCase).filter(
            GovernanceCase.case_type == "appeal",
            GovernanceCase.submitted_by == user.id,
            GovernanceCase.status == "open",
        ).count()
        if open_appeals:
            items.append(action_center_item("appeal", "我的申诉待处理", open_appeals, "governance", "warning", "申诉不会改变原记录，处理结果会以站内消息告知。"))

    if role_code in LEADER_CODES:
        member_ids = direct_member_ids(db, user.id)
        pending_reviews = db.query(RecognitionRecord).filter(
            RecognitionRecord.employee_id.in_(member_ids) if member_ids else RecognitionRecord.id == -1,
            RecognitionRecord.status == "pending",
        ).count()
        if pending_reviews:
            items.append(action_center_item("recognition_review", "待复核签卡", pending_reviews, "review", "warning", "直属组员提交的签卡等待复核。"))
        follow_ups = db.query(DeductionFollowUp).filter(
            DeductionFollowUp.supervisor_id == user.id,
            DeductionFollowUp.status == "pending",
        ).count()
        if follow_ups:
            items.append(action_center_item("deduction_follow_up", "重复违规待跟进", follow_ups, "entries", "warning", "请完成声明、备忘录或警告的管理闭环。"))
    if role_code in MATERIAL_COLLABORATOR_CODES:
        pending_materials = db.query(DeductionRecord).filter(
            DeductionRecord.status.in_({"pending_material", "material_failed"}),
        ).count()
        if pending_materials:
            items.append(action_center_item("deduction_material", "待补充声明材料", pending_materials, "entries", "warning", "TA主管、主管、TA GSM和GSM均可协作补充；材料处理成功后该扣分才会正式生效。"))

    if role_code in GSM_CODES:
        attraction_ids = managed_attraction_ids(db, user.id)
        scoped_employee_ids = {
            employee_id for (employee_id,) in db.query(Employee.id).filter(
                Employee.attraction_id.in_(attraction_ids), Employee.is_active.is_(True)
            ).all()
        } if attraction_ids else set()
        pending_follow_ups = db.query(DeductionFollowUp).filter(
            DeductionFollowUp.employee_id.in_(scoped_employee_ids) if scoped_employee_ids else DeductionFollowUp.id == -1,
            DeductionFollowUp.status == "pending",
        ).count()
        if pending_follow_ups:
            items.append(action_center_item("deduction_follow_up", "范围内重复违规待处理", pending_follow_ups, "entries", "warning", "管理范围内存在尚未形成闭环的重复违规。"))

    if role_code in {"HR_CIRCLE", "HR_ADMIN", "SYSTEM_ADMIN"}:
        governance_cases = [
            row for row in db.query(GovernanceCase).filter(GovernanceCase.status == "open").all()
            if governance_scope_allows_case(db, user, row)
        ]
        if governance_cases:
            overdue = sum(row.due_at < datetime.now() for row in governance_cases)
            items.append(action_center_item("governance_case", "申诉与更正待处理", len(governance_cases), "governance", "critical" if overdue else "warning", "处理人不得是原登记、复核或作废操作人；超时事项应优先处理。"))
        open_alerts = [row for row in visible_system_alerts(db, user) if row.status == "open"]
        if open_alerts:
            items.append(action_center_item("system_alert", "系统告警待处理", len(open_alerts), "hrEmployees", "critical", "员工、工作组或规则存在需要核对的告警。"))
        allowed_attractions = scoped_hr_attraction_ids(db, user)
        employee_query = db.query(Employee).filter(Employee.is_active.is_(True), Employee.attraction_id.is_not(None))
        if allowed_attractions is not None:
            employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
        ungrouped = 0
        for employee in employee_query.all():
            role = role_at(db, employee.id)
            if role and role.code in FRONTLINE_CODES and not current_group_for_employee(db, employee.id):
                ungrouped += 1
        if ungrouped:
            items.append(action_center_item("ungrouped_employee", "在职CM/TR待分组", ungrouped, "hrEmployees", "warning", "人员尚未归入主管组，影响直属复核和管理范围。"))
        transfers = db.query(CircleTransferRequest).filter(CircleTransferRequest.status == "pending")
        if allowed_attractions is not None:
            transfers = transfers.filter(CircleTransferRequest.target_attraction_id.in_(allowed_attractions))
        pending_transfers = transfers.count()
        if pending_transfers:
            items.append(action_center_item("circle_transfer", "待确认跨圈调动", pending_transfers, "circleTransfers", "warning", "确认后将按既有规则迁移归属、组关系和当月数据。"))
        if role_code == "HR_CIRCLE":
            attraction_ids = allowed_attractions or set()
            open_circles = sum(1 for attraction_id in attraction_ids if not month_closure_scope(db, month_to_close, attraction_id))
            if open_circles:
                items.append(action_center_item("month_close", "上月待月结景点圈", open_circles, "monthClose", "info", f"{month_to_close} 数据待核对；完成月结检查清单后即可关闭月结。"))
        if role_code == "SYSTEM_ADMIN":
            health = backup_health_payload()
            if not health["ok"]:
                items.append(action_center_item("backup_health", "备份健康异常", max(1, len(health["issues"])), "operations", "critical", "备份巡检报告存在异常，请立即进入系统运营核对。"))
            closed_count = db.query(MonthClosure).filter(
                MonthClosure.closure_month == current_month,
                MonthClosure.status == "closed",
            ).count()
            if closed_count:
                items.append(action_center_item("month_close", "本月已关闭月结", closed_count, "operations", "info", "最高管理员可在系统运营页汇总查看，并通过原月结流程进行复查。"))
            overdue_reviews = db.query(RecognitionRecord).filter(
                RecognitionRecord.status == "pending",
                RecognitionRecord.submitted_at < datetime.now() - timedelta(hours=48),
            ).count()
            if overdue_reviews:
                items.append(action_center_item("overdue_review", "超过48小时未复核签卡", overdue_reviews, "operations", "warning", "站内待复核已超出运营时限，请按现有复核边界跟进。"))

    order = {"critical": 0, "warning": 1, "info": 2}
    items.sort(key=lambda row: (order.get(row["severity"], 9), row["title"]))
    return {"items": items, "total": sum(row["count"] for row in items), "role": role_code, "month": current_month}


@router.get("/admin/operations-health")
def operations_health(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SYSTEM_ADMIN"))):
    health = backup_health_payload()
    month = date.today().strftime("%Y-%m")
    circles = db.query(Attraction).filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True)).order_by(Attraction.name).all()
    closures = [month_closure_payload(db, month, circle.id, user) for circle in circles]
    return {
        "backup": health,
        "month": month,
        "month_closures": closures,
        "open_system_alerts": db.query(SystemAlert).filter(SystemAlert.status == "open").count(),
        "pending_circle_transfers": db.query(CircleTransferRequest).filter(CircleTransferRequest.status == "pending").count(),
        "governance": {
            "open_cases": db.query(GovernanceCase).filter(GovernanceCase.status == "open").count(),
            "overdue_cases": db.query(GovernanceCase).filter(GovernanceCase.status == "open", GovernanceCase.due_at < datetime.now()).count(),
            "overdue_recognition_reviews": db.query(RecognitionRecord).filter(RecognitionRecord.status == "pending", RecognitionRecord.submitted_at < datetime.now() - timedelta(hours=48)).count(),
            "retention_review_files": db.query(StoredFile).filter(StoredFile.status == "active", StoredFile.uploaded_at < datetime.now() - timedelta(days=730)).count(),
        },
    }


@router.get("/reviews")
def reviews(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("REVIEW_DIRECT"))):
    member_ids = direct_member_ids(db, user.id)
    rows = (
        db.query(RecognitionRecord)
        .filter(
            RecognitionRecord.employee_id.in_(member_ids) if member_ids else RecognitionRecord.id == -1,
            RecognitionRecord.status != "void",
        )
        .order_by(RecognitionRecord.status == "confirmed", RecognitionRecord.submitted_at.desc())
        .all()
    )
    return [recognition_payload(row) for row in rows]


@router.post("/reviews/{record_id}")
def review_recognition(record_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("REVIEW_DIRECT"))):
    row = db.get(RecognitionRecord, record_id)
    if not row or row.employee_id not in direct_member_ids(db, user.id):
        raise HTTPException(404, "直属组员签卡不存在")
    employee = db.get(Employee, row.employee_id)
    ensure_month_open(db, row.recognition_month, row.home_attraction_id or (employee.attraction_id if employee else None), "复核签卡")
    action = str(payload.get("action") or "")
    mapping = {"confirm": "confirmed", "reject": "rejected", "restore": "pending"}
    if action not in mapping:
        raise HTTPException(400, "无效操作")
    before = row.status
    after = mapping[action]
    allowed_actions = {
        "pending": {"confirm", "reject"},
        "confirmed": {"restore"},
        "rejected": {"restore"},
    }
    if action not in allowed_actions.get(before, set()):
        raise HTTPException(409, "记录状态已经变化，请刷新后按当前状态操作")
    review_note = str(payload.get("note") or "").strip() or None
    claimed = db.execute(
        update(RecognitionRecord)
        .where(RecognitionRecord.id == row.id, RecognitionRecord.status == before)
        .values(
            status=after,
            reviewed_by=user.id,
            reviewed_by_name=user.name,
            reviewed_at=datetime.now(),
            review_note=review_note,
        )
        .execution_options(synchronize_session=False)
    )
    if claimed.rowcount != 1:
        db.rollback()
        raise HTTPException(409, "记录已被其他人处理，请刷新后重试")
    db.refresh(row)
    recognition_type = db.get(RecognitionType, row.recognition_type_id)
    if after == "confirmed" and recognition_type:
        apply_monthly_category_cap(db, row, recognition_type)
    elif after in {"pending", "rejected"}:
        row.credited_fraction = Decimal("0.00")
        row.monthly_cap_status = "not_applicable"
        row.monthly_cap_reason = None
    quota_code = recognition_type.code if recognition_type and recognition_type.code == "COMMENDATION_LETTER" else None
    quota = db.query(RecognitionMonthlyQuota).filter_by(recognition_id=row.id).first()
    if after == "rejected" and quota:
        db.delete(quota)
    elif after in {"pending", "confirmed"} and quota_code and not quota:
        occupied = db.query(RecognitionMonthlyQuota).filter_by(
            employee_id=row.employee_id, quota_code=quota_code, quota_month=row.recognition_month
        ).first()
        if occupied:
            raise HTTPException(409, "该员工本月已有其他表扬信记录，当前记录不能还原")
        db.add(
            RecognitionMonthlyQuota(
                employee_id=row.employee_id,
                quota_code=quota_code,
                quota_month=row.recognition_month,
                recognition_id=row.id,
            )
        )
    db.add(
        RecognitionReview(
            recognition_id=row.id,
            action=action,
            before_status=before,
            after_status=after,
            reviewer_id=user.id,
            reviewer_name=user.name,
            note=row.review_note,
        )
    )
    write_audit(db, user.employee, "复核签卡", "recognition", row.id, before={"status": before}, after={"status": after}, reason=row.review_note, ip_address=client_ip(request))
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "该员工本月已有其他表扬信记录，当前记录不能还原") from exc
    invalidate_data_caches()
    return {"ok": True, "record": recognition_payload(row)}


@router.get("/member-records")
def member_records(
    month: str | None = None,
    keyword: str | None = None,
    status: str | None = None,
    record_type: str = "all",
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("MEMBER_RECORDS")),
):
    if record_type not in {"all", "recognition", "deduction", "sick_leave"}:
        raise HTTPException(400, "记录类型无效")
    member_ids = direct_member_ids(db, user.id)
    empty_condition = RecognitionRecord.id == -1
    result: list[dict] = []
    if record_type in {"all", "recognition"}:
        query = db.query(RecognitionRecord).filter(RecognitionRecord.employee_id.in_(member_ids) if member_ids else empty_condition)
        if month:
            query = query.filter(RecognitionRecord.recognition_month == month)
        if keyword:
            value = like_escaped_pattern(keyword)
            query = query.filter(or_(RecognitionRecord.employee_name.like(value, escape="\\"), RecognitionRecord.employee_no.like(value, escape="\\")))
        if status in RECOGNITION_FILTER_STATUSES:
            query = query.filter(RecognitionRecord.status == status)
        elif status in ATTENDANCE_FILTER_STATUSES:
            query = query.filter(RecognitionRecord.id == -1)
        result.extend(recognition_payload(row) for row in query.order_by(RecognitionRecord.submitted_at.desc()).limit(1000).all())
    if record_type in {"all", "deduction"}:
        query = db.query(DeductionRecord).filter(DeductionRecord.employee_id.in_(member_ids) if member_ids else DeductionRecord.id == -1)
        if month:
            query = query.filter(DeductionRecord.deduction_month == month)
        if keyword:
            value = like_escaped_pattern(keyword)
            query = query.filter(or_(DeductionRecord.employee_name.like(value, escape="\\"), DeductionRecord.employee_no.like(value, escape="\\")))
        if status in DEDUCTION_FILTER_STATUSES:
            query = query.filter(DeductionRecord.status == status)
        elif status in RECOGNITION_FILTER_STATUSES or status in ATTENDANCE_FILTER_STATUSES:
            query = query.filter(DeductionRecord.id == -1)
        result.extend(deduction_payload(row) for row in query.order_by(DeductionRecord.submitted_at.desc()).limit(1000).all())
    if record_type in {"all", "sick_leave"}:
        query = db.query(SickLeaveRecord).filter(SickLeaveRecord.employee_id.in_(member_ids) if member_ids else SickLeaveRecord.id == -1)
        if month:
            query = query.filter(SickLeaveRecord.attendance_month == month)
        if keyword:
            value = like_escaped_pattern(keyword)
            query = query.join(Employee, Employee.id == SickLeaveRecord.employee_id).filter(or_(Employee.name.like(value, escape="\\"), Employee.employee_no.like(value, escape="\\")))
        if status in ATTENDANCE_FILTER_STATUSES:
            query = query.filter(SickLeaveRecord.status == status)
        elif status in RECOGNITION_FILTER_STATUSES:
            query = query.filter(SickLeaveRecord.id == -1)
        result.extend(sick_leave_payloads(db, query.order_by(SickLeaveRecord.submitted_at.desc()).limit(1000).all()))
    priority = {
        ("recognition", "pending"): 0,
        ("recognition", "rejected"): 0,
        ("deduction", "active"): 1,
        ("sick_leave", "active"): 1,
        ("recognition", "confirmed"): 2,
        ("deduction", "void"): 3,
        ("sick_leave", "void"): 3,
    }
    result.sort(key=lambda row: row["submitted_at"], reverse=True)
    result.sort(key=lambda row: priority.get((row["record_type"], row["status"]), 9))
    return result


def member_score_detail_payload(
    db: Session,
    employee: Employee,
    score: dict,
    recognitions: list[RecognitionRecord],
    deductions: list[DeductionRecord],
    sick_leaves: list[SickLeaveRecord],
    attendance: AttendanceMonthlyScore | None,
) -> dict:
    recognition_rows = [recognition_payload(row) for row in recognitions]
    deduction_rows = [deduction_payload(row) for row in deductions]
    sick_leave_rows = sick_leave_payloads(db, sick_leaves)
    all_records: list[dict] = []

    for row in recognition_rows:
        included = row["status"] == "confirmed"
        all_records.append(
            {
                "record_type": "recognition",
                "record_type_name": "加分",
                "business_date": row["recognition_date"],
                "submitted_at": row["submitted_at"],
                "title": row["recognition_type"],
                "content": f"{row['content']} · 认可人：{row['recognizer_name']}",
                "operator_name": row["operator_name"],
                "status": row["status"],
                "status_name": row["status_name"],
                "reason": row["review_note"],
                "score": row["credited_fraction"],
                "score_text": f"+{row['credited_fraction']:.2f}",
                "included": included,
                "included_name": "是" if included else "否",
                "attachment_url": row["image_url"],
                "attachment_name": "认可图片" if row["image_url"] else "",
                "attachment_is_previewable": row["image_is_previewable"],
                "attachment_preview_kind": row["image_preview_kind"],
            }
        )
    for row in deduction_rows:
        included = row["status"] == "active"
        all_records.append(
            {
                "record_type": "deduction",
                "record_type_name": "扣分",
                "business_date": row["occurred_on"],
                "submitted_at": row["submitted_at"],
                "title": f"{row['deduction_type']} · {row['deduction_level']}",
                "content": row["description"] + (f" · 作废原因：{row['void_reason']}" if row["void_reason"] else ""),
                "operator_name": row["submitter_name"],
                "status": row["status"],
                "status_name": row["status_name"],
                "reason": row["void_reason"],
                "score": -row["points"],
                "score_text": f"-{row['points']:.2f}",
                "included": included,
                "included_name": "是" if included else "否",
                "attachment_url": row["document_url"],
                "attachment_name": "声明PDF",
                "attachment_is_previewable": False,
                "attachment_preview_kind": row["document_preview_kind"],
            }
        )
    for row in sick_leave_rows:
        included = row["status"] == "active"
        all_records.append(
            {
                "record_type": "sick_leave",
                "record_type_name": "病假",
                "business_date": row["leave_start_date"],
                "submitted_at": row["submitted_at"],
                "title": f"{row['leave_start_date']} 至 {row['leave_end_date']}",
                "content": f"实际{row['leave_days']:.1f}天，计费{row['charged_days']}天" + (f" · {row['note']}" if row["note"] else "") + (f" · 作废原因：{row['void_reason']}" if row["void_reason"] else ""),
                "operator_name": row["submitter_name"],
                "status": row["status"],
                "status_name": row["status_name"],
                "reason": row["void_reason"],
                "score": None,
                "score_text": "影响全勤" if included else "不影响",
                "included": included,
                "included_name": "通过全勤分计算" if included else "否",
                "attachment_url": row["proof_url"],
                "attachment_name": "病假证明",
                "attachment_is_previewable": row["proof_is_previewable"],
                "attachment_preview_kind": row["proof_preview_kind"],
            }
        )

    attendance_payload = None
    if attendance:
        attendance_payload = {
            "eligible": attendance.eligible,
            "actual_sick_days": float(attendance.actual_sick_days),
            "charged_sick_days": attendance.charged_sick_days,
            "base_score": float(attendance.base_score),
            "perfect_bonus": float(attendance.perfect_bonus),
            "sick_deduction": float(attendance.sick_deduction),
            "final_score": float(attendance.final_score),
            "calculated_at": attendance.calculated_at.strftime("%Y-%m-%d %H:%M:%S"),
        }
        all_records.append(
            {
                "record_type": "attendance",
                "record_type_name": "全勤分",
                "business_date": attendance.attendance_month,
                "submitted_at": attendance_payload["calculated_at"],
                "title": "月度全勤分",
                "content": (
                    f"基础分{attendance_payload['base_score']:.2f}，全勤奖励{attendance_payload['perfect_bonus']:.2f}，"
                    f"病假扣减{attendance_payload['sick_deduction']:.2f}"
                ),
                "operator_name": "系统",
                "status": "calculated" if attendance.eligible else "ineligible",
                "status_name": "已计算" if attendance.eligible else "不适用",
                "reason": "",
                "score": attendance_payload["final_score"],
                "score_text": f"+{attendance_payload['final_score']:.2f}",
                "included": attendance.eligible,
                "included_name": "是" if attendance.eligible else "否",
                "attachment_url": "",
                "attachment_name": "",
                "attachment_is_previewable": False,
                "attachment_preview_kind": "",
            }
        )

    all_records.sort(key=lambda row: (row["business_date"], row["submitted_at"]), reverse=True)
    return {
        "employee_id": employee.id,
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "role_name": (role_at(db, employee.id).name if role_at(db, employee.id) else "未配置"),
        "recognition_score": float(score.get("recognition_score") or 0),
        "deduction_score": float(score.get("deduction_score") or 0),
        "attendance_score": float(score.get("attendance_score") or 0),
        "total_score": float(score.get("total_score") or 0),
        "details": {
            "recognitions": recognition_rows,
            "deductions": deduction_rows,
            "sick_leaves": sick_leave_rows,
            "attendance": attendance_payload,
            "all_records": all_records,
        },
    }


@router.get("/member-score-summary")
def member_score_summary(
    month: str | None = None,
    keyword: str | None = None,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("MEMBER_RECORDS")),
):
    month = month or date.today().strftime("%Y-%m")
    try:
        date.fromisoformat(f"{month}-01")
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "月份格式应为YYYY-MM") from exc
    member_ids = direct_member_ids(db, user.id)
    employee_query = db.query(Employee).filter(Employee.id.in_(member_ids) if member_ids else Employee.id == -1)
    if keyword and keyword.strip():
        value = f"%{keyword.strip()}%"
        employee_query = employee_query.filter(or_(Employee.name.like(value), Employee.employee_no.like(value)))
    employees = employee_query.order_by(Employee.name, Employee.employee_no).all()
    selected_ids = [employee.id for employee in employees]
    if not selected_ids:
        return {"month": month, "rows": []}

    ensure_month_attendance(db, month, selected_ids)
    db.commit()
    placeholders = []
    params: dict[str, object] = {"month": month}
    for index, employee_id in enumerate(selected_ids):
        key = f"employee_{index}"
        placeholders.append(f":{key}")
        params[key] = employee_id
    scores = {
        row["employee_id"]: dict(row)
        for row in db.execute(
            text(
                "SELECT employee_id, recognition_score, deduction_score, attendance_score, total_score "
                f"FROM v_employee_month_scores WHERE score_month=:month AND employee_id IN ({','.join(placeholders)})"
            ),
            params,
        ).mappings().all()
    }
    recognition_groups: dict[int, list[RecognitionRecord]] = {employee_id: [] for employee_id in selected_ids}
    deduction_groups: dict[int, list[DeductionRecord]] = {employee_id: [] for employee_id in selected_ids}
    sick_leave_groups: dict[int, list[SickLeaveRecord]] = {employee_id: [] for employee_id in selected_ids}
    for row in (
        db.query(RecognitionRecord)
        .filter(RecognitionRecord.employee_id.in_(selected_ids), RecognitionRecord.recognition_month == month)
        .order_by(RecognitionRecord.submitted_at.desc(), RecognitionRecord.id.desc())
        .all()
    ):
        recognition_groups[row.employee_id].append(row)
    for row in (
        db.query(DeductionRecord)
        .filter(DeductionRecord.employee_id.in_(selected_ids), DeductionRecord.deduction_month == month)
        .order_by(DeductionRecord.submitted_at.desc(), DeductionRecord.id.desc())
        .all()
    ):
        deduction_groups[row.employee_id].append(row)
    for row in (
        db.query(SickLeaveRecord)
        .filter(SickLeaveRecord.employee_id.in_(selected_ids), SickLeaveRecord.attendance_month == month)
        .order_by(SickLeaveRecord.submitted_at.desc(), SickLeaveRecord.id.desc())
        .all()
    ):
        sick_leave_groups[row.employee_id].append(row)
    attendance_rows = {
        row.employee_id: row
        for row in db.query(AttendanceMonthlyScore)
        .filter(AttendanceMonthlyScore.employee_id.in_(selected_ids), AttendanceMonthlyScore.attendance_month == month)
        .all()
    }
    empty_score = {"recognition_score": 0, "deduction_score": 0, "attendance_score": 0, "total_score": 0}
    return {
        "month": month,
        "rows": [
            member_score_detail_payload(
                db,
                employee,
                scores.get(employee.id, empty_score),
                recognition_groups[employee.id],
                deduction_groups[employee.id],
                sick_leave_groups[employee.id],
                attendance_rows.get(employee.id),
            )
            for employee in employees
        ],
    }


@router.get("/my-entries")
def my_entries(
    month: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    keyword: str | None = None,
    status: str | None = None,
    record_type: str = "all",
    scope: str = "mine",
    page: int = 1,
    page_size: int = 100,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    if not ({"EMPLOYEE_ADD", "SICK_REGISTER", "DEDUCTION_DIRECT", "DEDUCTION_ALL", "SYSTEM_ADMIN"} & user.permissions):
        raise HTTPException(403, "没有登记记录查询权限")
    if record_type not in {"all", "recognition", "deduction", "sick_leave", "follow_up"}:
        raise HTTPException(400, "记录类型无效")
    if scope not in {"mine", "supervisors"}:
        raise HTTPException(400, "查询范围无效")
    is_supervisor_entry_view = user.role.code in LEADER_CODES
    if scope == "supervisors" and not is_supervisor_entry_view:
        raise HTTPException(403, "仅TA主管或主管可查看全部主管登记记录")
    if scope == "supervisors" and record_type in {"recognition", "follow_up"}:
        raise HTTPException(400, "全部主管登记记录仅查询扣分和缺勤")
    parsed_start = parse_iso_date(start_date, "开始日期") if start_date else None
    parsed_end = parse_iso_date(end_date, "结束日期") if end_date else None
    if parsed_start and parsed_end and parsed_start > parsed_end:
        raise HTTPException(400, "开始日期不能晚于结束日期")
    page = max(1, page)
    page_size = min(max(10, page_size), 200)
    items: list[dict] = []
    supervisor_submitter_ids: set[int] = set()
    supervisor_role_names: set[str] = set()
    if scope == "supervisors":
        today_value = date.today().isoformat()
        supervisor_submitter_ids = {
            int(employee_id)
            for (employee_id,) in (
                db.query(EmployeeRoleAssignment.employee_id)
                .join(Role, Role.id == EmployeeRoleAssignment.role_id)
                .filter(
                    EmployeeRoleAssignment.status == "active",
                    EmployeeRoleAssignment.starts_on <= today_value,
                    or_(EmployeeRoleAssignment.ends_on.is_(None), EmployeeRoleAssignment.ends_on >= today_value),
                    Role.code.in_(LEADER_CODES),
                )
                .all()
            )
        }
        supervisor_role_names = {
            name for (name,) in db.query(Role.name).filter(Role.code.in_(LEADER_CODES), Role.active.is_(True)).all()
        }
    if scope != "supervisors" and record_type in {"all", "recognition"}:
        query = db.query(RecognitionRecord).filter(RecognitionRecord.operator_employee_id == user.id)
        if start_date:
            query = query.filter(RecognitionRecord.recognition_date >= start_date)
        if end_date:
            query = query.filter(RecognitionRecord.recognition_date <= end_date)
        if month and not (start_date or end_date):
            query = query.filter(RecognitionRecord.recognition_month == month)
        if keyword:
            value = like_escaped_pattern(keyword)
            query = query.filter(or_(RecognitionRecord.employee_name.like(value, escape="\\"), RecognitionRecord.employee_no.like(value, escape="\\")))
        if status in RECOGNITION_FILTER_STATUSES:
            query = query.filter(RecognitionRecord.status == status)
        elif status in ATTENDANCE_FILTER_STATUSES:
            query = query.filter(RecognitionRecord.id == -1)
        else:
            query = query.filter(RecognitionRecord.status != "void")
        items.extend(recognition_payload(row) for row in query.order_by(RecognitionRecord.submitted_at.desc()).all())
    if record_type in {"all", "deduction"}:
        query = db.query(DeductionRecord)
        if scope == "supervisors":
            if not supervisor_submitter_ids and not supervisor_role_names:
                query = query.filter(DeductionRecord.id == -1)
            else:
                scope_filters = []
                if supervisor_submitter_ids:
                    scope_filters.append(DeductionRecord.submitter_id.in_(supervisor_submitter_ids))
                if supervisor_role_names:
                    scope_filters.append(DeductionRecord.submitter_role_snapshot.in_(supervisor_role_names))
                query = query.filter(or_(*scope_filters))
        else:
            query = query.filter(DeductionRecord.submitter_id == user.id)
        if start_date:
            query = query.filter(DeductionRecord.occurred_on >= start_date)
        if end_date:
            query = query.filter(DeductionRecord.occurred_on <= end_date)
        if month and not (start_date or end_date):
            query = query.filter(DeductionRecord.deduction_month == month)
        if keyword:
            value = like_escaped_pattern(keyword)
            query = query.filter(or_(DeductionRecord.employee_name.like(value, escape="\\"), DeductionRecord.employee_no.like(value, escape="\\")))
        if status in DEDUCTION_FILTER_STATUSES:
            query = query.filter(DeductionRecord.status == status)
        elif status in RECOGNITION_FILTER_STATUSES or status in ATTENDANCE_FILTER_STATUSES:
            query = query.filter(DeductionRecord.id == -1)
        for row in query.order_by(DeductionRecord.submitted_at.desc()).all():
            payload = deduction_payload(row)
            if row.submitter_id == user.id and row.status in {"pending_material", "material_failed"}:
                payload["available_actions"] = [*payload["available_actions"], "void"]
            items.append(payload)
    if record_type in {"all", "sick_leave"}:
        query = db.query(SickLeaveRecord)
        if scope == "supervisors":
            query = query.filter(SickLeaveRecord.submitted_by.in_(supervisor_submitter_ids) if supervisor_submitter_ids else SickLeaveRecord.id == -1)
        else:
            query = query.filter(SickLeaveRecord.submitted_by == user.id)
        if start_date:
            query = query.filter(SickLeaveRecord.leave_end_date >= start_date)
        if end_date:
            query = query.filter(SickLeaveRecord.leave_start_date <= end_date)
        if month and not (start_date or end_date):
            query = query.filter(SickLeaveRecord.attendance_month == month)
        if keyword:
            value = like_escaped_pattern(keyword)
            query = query.join(Employee, Employee.id == SickLeaveRecord.employee_id).filter(or_(Employee.name.like(value, escape="\\"), Employee.employee_no.like(value, escape="\\")))
        if status in ATTENDANCE_FILTER_STATUSES:
            query = query.filter(SickLeaveRecord.status == status)
        elif status in RECOGNITION_FILTER_STATUSES:
            query = query.filter(SickLeaveRecord.id == -1)
        items.extend(sick_leave_payloads(db, query.order_by(SickLeaveRecord.submitted_at.desc()).all()))
    if scope != "supervisors" and record_type in {"all", "follow_up"}:
        query = db.query(DeductionFollowUp).filter(DeductionFollowUp.supervisor_id == user.id)
        if start_date:
            query = query.filter(DeductionFollowUp.occurred_on >= start_date)
        if end_date:
            query = query.filter(DeductionFollowUp.occurred_on <= end_date)
        if month and not (start_date or end_date):
            query = query.filter(DeductionFollowUp.occurred_on.like(f"{month}%"))
        if keyword:
            value = f"%{keyword.strip()}%"
            query = query.filter(or_(DeductionFollowUp.employee_name.like(value), DeductionFollowUp.employee_no.like(value)))
        if status in {"pending", "issued"}:
            query = query.filter(DeductionFollowUp.status == status)
        elif status in {"active", "void", "confirmed", "rejected"}:
            query = query.filter(DeductionFollowUp.id == -1)
        items.extend(deduction_follow_up_payload(row) for row in query.order_by(DeductionFollowUp.created_at.desc()).all())
    items.sort(key=lambda row: row["submitted_at"], reverse=True)
    total = len(items)
    start = (page - 1) * page_size
    page_items = items[start:start + page_size]
    active_deductions = sum(row["points"] for row in items if row["record_type"] == "deduction" and row["status"] == "active")
    recognition_score = sum(row.get("credited_fraction", row["fraction"]) for row in items if row["record_type"] == "recognition" and row["status"] == "confirmed")
    active_sick_days = sum(row["leave_days"] for row in items if row["record_type"] == "sick_leave" and row["status"] == "active")
    charged_sick_days = sum(row["charged_days"] for row in items if row["record_type"] == "sick_leave" and row["status"] == "active")
    return {
        "items": page_items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "scope": scope,
        "summary": {
            "recognition_count": sum(row["record_type"] == "recognition" for row in items),
            "recognition_score": round(recognition_score, 2),
            "deduction_count": sum(row["record_type"] == "deduction" for row in items),
            "deduction_score": round(active_deductions, 2),
            "sick_leave_count": sum(row["record_type"] == "sick_leave" for row in items),
            "active_sick_days": round(active_sick_days, 1),
            "charged_sick_days": charged_sick_days,
            "follow_up_count": sum(row["record_type"] == "follow_up" for row in items),
        },
    }


@router.get("/deduction-targets")
def deduction_targets(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if not ({"DEDUCTION_ALL", "DEDUCTION_DIRECT"} & user.permissions):
        raise HTTPException(403, "没有扣分权限")
    allowed = scoped_hr_attraction_ids(db, user)
    attraction_id = next(iter(allowed)) if allowed else None
    return search_employee_targets(db, attraction_id=attraction_id, limit=500)["items"]


@router.get("/deductions/attendance-repeat-check")
def check_attendance_repeat(
    employee_id: int,
    deduction_type_id: int,
    occurred_on: str,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    parse_iso_date(occurred_on, "事件日期")
    target, target_role = ensure_enabled_frontline_target(db, employee_id, "扣分")
    ensure_scoped_hr_employee(db, user, target)
    ensure_operational_target_scope(user, target)
    target_circle = db.get(Attraction, target.attraction_id) if target.attraction_id else None
    if not target_circle or not target_circle.active or not target_circle.employee_circle:
        raise HTTPException(400, "被扣分员工未配置有效景点圈")
    ensure_month_open(db, occurred_on[:7], target_circle.id, "生成重复处分跟进")
    if "DEDUCTION_ALL" in user.permissions:
        pass
    elif "DEDUCTION_DIRECT" in user.permissions:
        pass
    else:
        raise HTTPException(403, "没有扣分权限")
    deduction_type = db.get(DeductionType, deduction_type_id)
    if not deduction_type or not deduction_type.active:
        raise HTTPException(400, "扣分类型无效")
    ensure_deduction_type_allowed(user, deduction_type)
    result = attendance_repeat_context(db, employee_id, deduction_type, occurred_on, user)
    if result["has_repeat"] and user.role.code in LEADER_CODES:
        follow_up = ensure_repeat_follow_up(db, user, target, deduction_type, occurred_on, result)
        db.commit()
        result["follow_up_id"] = follow_up.id if follow_up else None
    result.update({"employee_id": target.id, "employee_name": target.name, "deduction_type_id": deduction_type.id, "deduction_type_name": deduction_type.name})
    return result


@router.post("/deductions")
async def create_deduction(
    request: Request,
    employee_id: int = Form(...),
    deduction_type_id: int = Form(...),
    deduction_level_id: int = Form(...),
    occurred_on: str = Form(...),
    description: str = Form(...),
    repeat_confirmed: bool = Form(False),
    idempotency_key: str | None = Form(None),
    document: UploadFile | None = File(None),
    document_images: list[UploadFile] | None = File(None),
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    request_key = normalize_request_key(idempotency_key)
    payload_digest = submission_payload_digest(
        {
            "employee_id": int(employee_id),
            "deduction_type_id": int(deduction_type_id),
            "deduction_level_id": int(deduction_level_id),
            "occurred_on": str(occurred_on or "").strip(),
            "description": str(description or "").strip(),
            "repeat_confirmed": bool(repeat_confirmed),
        }
    )
    duplicate_row = existing_submission(db, user.id, "deduction", request_key, DeductionRecord, payload_digest)
    if duplicate_row:
        return {"ok": True, "record": deduction_payload(duplicate_row), "duplicate": True}
    parse_iso_date(occurred_on, "事件日期")
    target, target_role = ensure_enabled_frontline_target(db, employee_id, "扣分")
    ensure_scoped_hr_employee(db, user, target)
    ensure_operational_target_scope(user, target)
    target_circle = db.get(Attraction, target.attraction_id) if target.attraction_id else None
    if not target_circle or not target_circle.active or not target_circle.employee_circle:
        raise HTTPException(400, "被扣分员工未配置有效景点圈")
    ensure_month_open(db, occurred_on[:7], target_circle.id, "新增扣分")
    if "DEDUCTION_ALL" in user.permissions:
        scope = "所有CM/TR"
    elif "DEDUCTION_DIRECT" in user.permissions:
        scope = "所有CM/TR（仅声明）"
    else:
        raise HTTPException(403, "没有扣分权限")
    level = db.get(DeductionLevel, deduction_level_id)
    deduction_type = db.get(DeductionType, deduction_type_id)
    if not level or not level.active or not deduction_type or not deduction_type.active:
        raise HTTPException(400, "扣分类型或等级无效")
    ensure_deduction_type_allowed(user, deduction_type)
    if "DEDUCTION_DIRECT" in user.permissions and "DEDUCTION_ALL" not in user.permissions and level.code != "STATEMENT":
        raise HTTPException(403, "TA主管、主管只能登记声明1分")
    existing_pending_material = pending_material_conflict(
        db,
        employee_id=target.id,
        deduction_type_id=deduction_type.id,
        deduction_level_id=level.id,
        occurred_on=occurred_on,
    )
    if existing_pending_material:
        raise_pending_material_conflict(existing_pending_material)
    photo_mode = bool([item for item in (document_images or []) if item and item.filename])
    pdf_mode = bool(document and document.filename)
    if direct_only_deduction_user(user) and level.code == "STATEMENT" and deduction_type.code in REPEAT_CONTROLLED_DEDUCTION_CODES and (photo_mode or pdf_mode):
        if statement_upgrade_candidate(db, employee_id, deduction_type.id, occurred_on):
            raise HTTPException(409, "三个月内已有可升级同类声明，请通过声明升级工单提交")
        repeat_context = {"has_repeat": False, "blocked": False, "requires_confirmation": False}
    elif photo_mode or pdf_mode:
        repeat_context = attendance_repeat_context(db, employee_id, deduction_type, occurred_on, user)
    else:
        repeat_context = {"has_repeat": False, "blocked": False, "requires_confirmation": False, "previous_records": []}
    if repeat_context["has_repeat"]:
        if repeat_context["blocked"]:
            raise HTTPException(409, detail={"code": "ATTENDANCE_REPEAT_BLOCKED", **repeat_context})
        if not repeat_confirmed:
            raise HTTPException(409, detail={"code": "ATTENDANCE_REPEAT_CONFIRMATION_REQUIRED", **repeat_context})
        minimum_rank = DEDUCTION_LEVEL_ORDER.get(repeat_context["minimum_level_code"], 0)
        if DEDUCTION_LEVEL_ORDER.get(level.code, 0) < minimum_rank:
            raise HTTPException(400, f"3个月内已有同类型处分，本次最低必须选择{repeat_context['minimum_level_name']}")
    description = description.strip()
    if not description:
        raise HTTPException(400, "事件说明必填")
    file_row = None
    source_rows: list[StoredFile] = []
    needs_pdf_processing = False
    if pdf_mode and photo_mode:
        raise HTTPException(400, "请在上传PDF和照片材料中选择一种方式")
    # A record without materials is a collaboration draft: it is visible in
    # supervisor to-dos but deliberately has no score and no repeat/escalation
    # effect until the material is supplied successfully.
    if not pdf_mode and not photo_mode:
        repeat_context = {"has_repeat": False, "blocked": False, "requires_confirmation": False, "previous_records": []}
    try:
        if photo_mode:
            source_rows = await stage_photo_materials(db, list(document_images or []), user.id)
            file_row = create_pdf_placeholder(db, user.id)
        elif pdf_mode:
            file_row = await stage_pdf_material(db, document, user.id)
            needs_pdf_processing = file_row.file_size > PDF_HIGH_QUALITY_OPTIMIZATION_THRESHOLD
            if needs_pdf_processing:
                file_row.status = "processing_source"
                source_rows = [file_row]
                file_row = create_pdf_placeholder(db, user.id)
            else:
                # A small, validated PDF is immediately usable.  Keep its
                # source row active instead of leaving an inaccessible
                # processing-source file behind.
                file_row.status = "active"
        else:
            file_row = create_pdf_placeholder(db, user.id)
            file_row.original_filename = "待补充声明材料.pdf"
            file_row.status = "pending_material"
        group = current_group_for_employee(db, target.id)
        row = DeductionRecord(
            employee_id=target.id,
            employee_no=target.employee_no,
            employee_name=target.name,
            employee_role_snapshot=target_role.name,
            employee_group_id_snapshot=group.id if group else None,
            attraction_id_snapshot=target.attraction_id,
            deduction_type_id=deduction_type.id,
            deduction_type_name=deduction_type.name,
            deduction_level_id=level.id,
            deduction_level_name=level.name,
            points=level.points,
            occurred_on=occurred_on,
            deduction_month=occurred_on[:7],
            description=description,
            document_file_id=file_row.id,
            submitter_id=user.id,
            submitter_name=user.name,
            submitter_role_snapshot=user.role.name,
            permission_scope_snapshot=scope,
            status="material_processing" if (photo_mode or needs_pdf_processing) else ("active" if pdf_mode else "pending_material"),
            material_status="processing" if (photo_mode or needs_pdf_processing) else ("ready" if pdf_mode else "missing"),
            material_source_type="photos" if photo_mode else ("pdf" if pdf_mode else "later"),
        )
        db.add(row)
        db.flush()
        if photo_mode or needs_pdf_processing:
            queue_photo_material_job(db, row, file_row, source_rows, mode="pdf_compress" if needs_pdf_processing else "deduction")
        elif level.code != "STATEMENT" and row.status == "active":
            follow_ups = db.query(DeductionFollowUp).filter_by(
                employee_id=target.id,
                deduction_type_id=deduction_type.id,
                occurred_on=occurred_on,
                status="pending",
            ).all()
            for follow_up in follow_ups:
                follow_up.status = "issued"
                follow_up.issued_deduction_id = row.id
                follow_up.issued_by = user.id
                follow_up.issued_by_name = user.name
                follow_up.issued_at = datetime.now()
        remember_submission(db, user.id, "deduction", request_key, row.id, payload_digest)
        audit_after = deduction_payload(row)
        if repeat_context["has_repeat"]:
            audit_after["repeat_warning_acknowledged"] = True
            audit_after["repeat_reference_ids"] = [item["id"] for item in repeat_context["previous_records"]]
            audit_after["minimum_level_name"] = repeat_context["minimum_level_name"]
        write_audit(db, user.employee, "登记扣分", "deduction", row.id, after=audit_after, ip_address=client_ip(request))
        db.commit()
    except IntegrityError:
        db.rollback()
        remove_upload_file(file_row)
        for source_row in source_rows:
            remove_upload_file(source_row)
        duplicate_row = existing_submission(db, user.id, "deduction", request_key, DeductionRecord, payload_digest)
        if duplicate_row:
            return {"ok": True, "record": deduction_payload(duplicate_row), "duplicate": True}
        raise
    except Exception:
        db.rollback()
        remove_upload_file(file_row)
        for source_row in source_rows:
            remove_upload_file(source_row)
        raise
    invalidate_data_caches()
    return {"ok": True, "record": deduction_payload(row)}


@router.get("/deduction-upgrades/preview")
def deduction_upgrade_preview(employee_id: int, deduction_type_id: int, occurred_on: str, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if not ({"DEDUCTION_DIRECT", "DEDUCTION_ALL"} & user.permissions):
        raise HTTPException(403, "没有声明升级登记权限")
    parse_iso_date(occurred_on, "事件日期")
    target, _role = ensure_enabled_frontline_target(db, employee_id, "扣分")
    deduction_type = db.get(DeductionType, deduction_type_id)
    statement = db.query(DeductionLevel).filter_by(code="STATEMENT", active=True).first()
    if deduction_type and statement:
        existing_pending_material = pending_material_conflict(
            db,
            employee_id=target.id,
            deduction_type_id=deduction_type.id,
            deduction_level_id=statement.id,
            occurred_on=occurred_on,
        )
        if existing_pending_material:
            raise_pending_material_conflict(existing_pending_material)
    if not deduction_type or not deduction_type.active or deduction_type.code not in REPEAT_CONTROLLED_DEDUCTION_CODES:
        return {"eligible": False}
    first = statement_upgrade_candidate(db, target.id, deduction_type.id, occurred_on)
    if not first:
        return {"eligible": False}
    return {"eligible": True, "first_record": deduction_payload(first), "reviewers": statement_upgrade_reviewer_options(db)}


@router.post("/deduction-upgrades")
async def create_deduction_upgrade(
    request: Request,
    employee_id: int = Form(...),
    deduction_type_id: int = Form(...),
    occurred_on: str = Form(...),
    description: str = Form(...),
    reviewer_id: int = Form(...),
    document: UploadFile | None = File(None),
    document_images: list[UploadFile] | None = File(None),
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("DEDUCTION_DIRECT")),
):
    parse_iso_date(occurred_on, "事件日期")
    description = description.strip()
    if not description:
        raise HTTPException(400, "事件说明必填")
    target, target_role = ensure_enabled_frontline_target(db, employee_id, "扣分")
    ensure_scoped_hr_employee(db, user, target)
    target_circle = db.get(Attraction, target.attraction_id) if target.attraction_id else None
    if not target_circle or not target_circle.active or not target_circle.employee_circle:
        raise HTTPException(400, "被扣分员工未配置有效景点圈")
    ensure_month_open(db, occurred_on[:7], target_circle.id, "提交声明升级工单")
    deduction_type = db.get(DeductionType, deduction_type_id)
    statement = db.query(DeductionLevel).filter_by(code="STATEMENT", active=True).first()
    reviewer = db.get(Employee, reviewer_id)
    reviewer_role = role_at(db, reviewer.id) if reviewer else None
    if not deduction_type or deduction_type.code not in REPEAT_CONTROLLED_DEDUCTION_CODES or not statement:
        raise HTTPException(400, "该扣分类型不支持声明升级")
    existing_pending_material = pending_material_conflict(
        db,
        employee_id=target.id,
        deduction_type_id=deduction_type.id,
        deduction_level_id=statement.id,
        occurred_on=occurred_on,
    )
    if existing_pending_material:
        raise_pending_material_conflict(existing_pending_material)
    if not reviewer or not reviewer.is_active or not reviewer_role or reviewer_role.code not in UPGRADE_REVIEWER_CODES:
        raise HTTPException(400, "请选择在职的GSM或TA GSM审核")
    reviewer_account = db.query(UserAccount).filter_by(employee_id=reviewer.id, enabled=True).first()
    if not reviewer_account:
        raise HTTPException(400, "所选审核人账号未启用")
    photo_mode = bool([item for item in (document_images or []) if item and item.filename])
    pdf_mode = bool(document and document.filename)
    if pdf_mode and photo_mode:
        raise HTTPException(400, "请在上传PDF和照片材料中选择一种方式")
    if not pdf_mode and not photo_mode:
        raise HTTPException(400, "请上传PDF，或拍照/从相册选择声明材料")
    first = statement_upgrade_candidate(db, target.id, deduction_type.id, occurred_on)
    if not first:
        raise HTTPException(409, "三个月内不存在可升级的同类声明，请按普通声明登记")
    claimed = db.execute(
        update(DeductionRecord)
        .where(DeductionRecord.id == first.id, DeductionRecord.status == "active", or_(DeductionRecord.upgrade_state.is_(None), DeductionRecord.upgrade_state == "eligible"))
        .values(upgrade_state="material_processing" if photo_mode else "pending", upgrade_role="source_first")
        .execution_options(synchronize_session=False)
    )
    if claimed.rowcount != 1:
        db.rollback()
        raise HTTPException(409, "该历史声明已被其他工单使用，请刷新后重试")
    file_row = None
    source_rows: list[StoredFile] = []
    needs_pdf_processing = False
    try:
        if photo_mode:
            source_rows = await stage_photo_materials(db, list(document_images or []), user.id)
            file_row = create_pdf_placeholder(db, user.id)
        else:
            file_row = await stage_pdf_material(db, document, user.id)
            needs_pdf_processing = file_row.file_size > PDF_HIGH_QUALITY_OPTIMIZATION_THRESHOLD
            if needs_pdf_processing:
                file_row.status = "processing_source"
                source_rows = [file_row]
                file_row = create_pdf_placeholder(db, user.id)
            else:
                file_row.status = "active"
        group = current_group_for_employee(db, target.id)
        second = DeductionRecord(
        employee_id=target.id, employee_no=target.employee_no, employee_name=target.name,
        employee_role_snapshot=target_role.name, employee_group_id_snapshot=group.id if group else None,
        attraction_id_snapshot=target.attraction_id, deduction_type_id=deduction_type.id,
        deduction_type_name=deduction_type.name, deduction_level_id=statement.id,
        deduction_level_name=statement.name, points=Decimal("0"), occurred_on=occurred_on,
        deduction_month=occurred_on[:7], description=description, document_file_id=file_row.id,
        submitter_id=user.id, submitter_name=user.name, submitter_role_snapshot=user.role.name,
        permission_scope_snapshot="所有CM/TR（声明升级待审核）", status="material_processing" if (photo_mode or needs_pdf_processing) else "pending_upgrade",
        material_status="processing" if (photo_mode or needs_pdf_processing) else "ready", material_source_type="photos" if photo_mode else "pdf",
        upgrade_role="source_second", upgrade_state="material_processing" if (photo_mode or needs_pdf_processing) else "pending",
        )
        db.add(second)
        db.flush()
        if photo_mode or needs_pdf_processing:
            queue_photo_material_job(db, second, file_row, source_rows, mode="upgrade_pdf_compress" if needs_pdf_processing else "upgrade", reviewer_id=reviewer.id, first_deduction_id=first.id)
            write_audit(db, user.employee, "提交声明升级材料", "deduction", second.id, after={"first_deduction_id": first.id, "reviewer": reviewer.name, "material_status": "processing"}, ip_address=client_ip(request))
            db.commit()
            invalidate_data_caches()
            return {"ok": True, "processing": True, "record": deduction_payload(second)}
        row = DeductionUpgradeRequest(
            employee_id=target.id, deduction_type_id=deduction_type.id, first_deduction_id=first.id,
            second_deduction_id=second.id, reviewer_id=reviewer.id, reviewer_name=reviewer.name,
            submitted_by=user.id, submitted_by_name=user.name, status="pending",
        )
        db.add(row)
        db.flush()
        first.upgrade_request_id = row.id
        second.upgrade_request_id = row.id
        write_audit(db, user.employee, "提交声明升级工单", "deduction_upgrade", row.id, after={"first_deduction_id": first.id, "second_deduction_id": second.id, "reviewer": reviewer.name}, ip_address=client_ip(request))
        db.commit()
        invalidate_data_caches()
        return {"ok": True, "request": deduction_upgrade_payload(db, row)}
    except Exception:
        db.rollback()
        remove_upload_file(file_row)
        for source_row in source_rows:
            remove_upload_file(source_row)
        raise


@router.get("/deduction-upgrades/pending")
def pending_deduction_upgrades(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if user.role.code not in UPGRADE_REVIEWER_CODES:
        raise HTTPException(403, "仅GSM或TA GSM可审核声明升级工单")
    rows = db.query(DeductionUpgradeRequest).filter_by(reviewer_id=user.id, status="pending").order_by(DeductionUpgradeRequest.created_at.desc()).all()
    return {"items": [deduction_upgrade_payload(db, row) for row in rows]}


@router.get("/deduction-upgrades/reviewers")
def deduction_upgrade_reviewers(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if user.role.code not in UPGRADE_REVIEWER_CODES:
        raise HTTPException(403, "仅GSM或TA GSM可转交声明升级工单")
    return {"items": statement_upgrade_reviewer_options(db)}


@router.post("/deduction-upgrades/{request_id}/transfer")
def transfer_deduction_upgrade(request_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    row = db.get(DeductionUpgradeRequest, request_id)
    if not row or row.status != "pending" or row.reviewer_id != user.id or user.role.code not in UPGRADE_REVIEWER_CODES:
        raise HTTPException(404, "待审核工单不存在或已转交")
    target_id = int(payload.get("reviewer_id") or 0)
    reason = str(payload.get("reason") or "").strip()
    target = db.get(Employee, target_id)
    target_role = role_at(db, target.id) if target else None
    if not reason:
        raise HTTPException(400, "转交说明必填")
    if not target or not target.is_active or not target_role or target_role.code not in UPGRADE_REVIEWER_CODES or not db.query(UserAccount).filter_by(employee_id=target.id, enabled=True).first():
        raise HTTPException(400, "请选择在职且账号启用的GSM或TA GSM")
    db.add(DeductionUpgradeTransfer(request_id=row.id, from_reviewer_id=user.id, from_reviewer_name=user.name, to_reviewer_id=target.id, to_reviewer_name=target.name, reason=reason))
    row.reviewer_id, row.reviewer_name = target.id, target.name
    write_audit(db, user.employee, "转交声明升级工单", "deduction_upgrade", row.id, after={"to": target.name, "reason": reason}, ip_address=client_ip(request))
    db.commit()
    return {"ok": True}


@router.post("/deduction-upgrades/{request_id}/resolve")
def resolve_deduction_upgrade(request_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    row = db.get(DeductionUpgradeRequest, request_id)
    if not row or row.status != "pending" or row.reviewer_id != user.id or user.role.code not in UPGRADE_REVIEWER_CODES:
        raise HTTPException(404, "待审核工单不存在或已被处理")
    decision = str(payload.get("decision") or "")
    note = str(payload.get("handling_note") or "").strip()
    if not note:
        raise HTTPException(400, "处理说明必填")
    first, second = db.get(DeductionRecord, row.first_deduction_id), db.get(DeductionRecord, row.second_deduction_id)
    if not first or not second:
        raise HTTPException(409, "来源声明不存在")
    ensure_month_open(db, second.deduction_month, second.attraction_id_snapshot, "处理声明升级工单")
    if decision == "reject":
        second.status, second.upgrade_state = "void", "rejected"
        second.void_reason = "声明升级审核不通过：" + note
        first.upgrade_request_id, first.upgrade_role, first.upgrade_state = None, None, "eligible"
        row.status, row.handling_note, row.resolved_by, row.resolved_by_name, row.resolved_at = "rejected", note, user.id, user.name, datetime.now()
    elif decision == "approve":
        level_id = int(payload.get("result_level_id") or 0)
        confirmed = bool(payload.get("issued_confirmed"))
        level = db.get(DeductionLevel, level_id)
        if not confirmed:
            raise HTTPException(400, "请先确认已完成真实备忘录或一级警告开具")
        if not level or level.code not in {"MEMO", "WARNING_1"}:
            raise HTTPException(400, "升级结果仅可选择备忘录或一级警告")
        placeholder = StoredFile(storage_key=f"system/no-document-{secrets.token_hex(12)}", original_filename="无需上传正式文书", extension=".none", mime_type="application/octet-stream", file_size=0, sha256="0" * 64, uploaded_by=user.id, status="not_required")
        db.add(placeholder)
        db.flush()
        result = DeductionRecord(employee_id=second.employee_id, employee_no=second.employee_no, employee_name=second.employee_name, employee_role_snapshot=second.employee_role_snapshot, employee_group_id_snapshot=second.employee_group_id_snapshot, attraction_id_snapshot=second.attraction_id_snapshot, deduction_type_id=second.deduction_type_id, deduction_type_name=second.deduction_type_name, deduction_level_id=level.id, deduction_level_name=level.name, points=level.points, occurred_on=second.occurred_on, deduction_month=second.deduction_month, description=note, document_file_id=placeholder.id, submitter_id=user.id, submitter_name=user.name, submitter_role_snapshot=user.role.name, permission_scope_snapshot="声明升级审核", status="active", upgrade_role="result", upgrade_state="result")
        db.add(result)
        db.flush()
        first.upgrade_role, first.upgrade_state = "source_first", "source_first"
        second.status, second.upgrade_role, second.upgrade_state = "active", "source_second", "source_second"
        row.status, row.result_level_id, row.result_deduction_id, row.handling_note, row.issued_confirmed = "approved", level.id, result.id, note, True
        row.resolved_by, row.resolved_by_name, row.resolved_at = user.id, user.name, datetime.now()
    else:
        raise HTTPException(400, "无效处理结论")
    write_audit(db, user.employee, "处理声明升级工单", "deduction_upgrade", row.id, after={"decision": decision, "note": note}, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "request": deduction_upgrade_payload(db, row)}


@router.get("/deductions")
def list_deductions(month: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if not ({"DEDUCTION_DIRECT", "DEDUCTION_ALL", "SYSTEM_ADMIN"} & user.permissions):
        raise HTTPException(403, "没有扣分查询权限")
    query = db.query(DeductionRecord)
    if "SYSTEM_ADMIN" not in user.permissions:
        query = query.filter(DeductionRecord.submitter_id == user.id)
    if month:
        query = query.filter(DeductionRecord.deduction_month == month)
    return [deduction_payload(row) for row in query.order_by(DeductionRecord.submitted_at.desc()).limit(500).all()]


@router.get("/deductions/pending-materials")
def pending_deduction_materials(
    scope: str = "mine",
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    """Shared material queue for LEAD and GSM collaboration.

    The default is the operator's own circle; the explicitly selected ``all``
    view preserves the agreed cross-circle support workflow.
    """
    if user.role.code not in MATERIAL_COLLABORATOR_CODES:
        raise HTTPException(403, "仅TA主管、主管、TA GSM或GSM可查看待补充材料")
    if scope not in {"mine", "all"}:
        raise HTTPException(400, "材料范围仅支持我的景点圈或全部景点圈")
    query = db.query(DeductionRecord).filter(DeductionRecord.status.in_({"pending_material", "material_failed"}))
    own_attraction_id = user.employee.attraction_id
    if scope == "mine":
        query = query.filter(DeductionRecord.attraction_id_snapshot == own_attraction_id) if own_attraction_id else query.filter(DeductionRecord.id == -1)
    rows = query.order_by(DeductionRecord.submitted_at.desc()).limit(200).all()
    return {
        "items": [deduction_payload(row) for row in rows],
        "scope": scope,
        "own_attraction_id": own_attraction_id,
    }


@router.get("/deductions/{record_id}/material-status")
def deduction_material_status(record_id: int, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    row = db.get(DeductionRecord, record_id)
    if not row:
        raise HTTPException(404, "扣分记录不存在")
    if row.submitter_id != user.id and "SYSTEM_ADMIN" not in user.permissions and user.role.code not in MATERIAL_COLLABORATOR_CODES:
        raise HTTPException(403, "无权查看该材料状态")
    return {"ok": True, "record": deduction_payload(row)}


@router.post("/deductions/{record_id}/material")
async def retry_deduction_material(
    record_id: int,
    request: Request,
    document: UploadFile | None = File(None),
    document_images: list[UploadFile] | None = File(None),
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    """Re-stage photos for a failed material job without creating another deduction."""
    row = db.get(DeductionRecord, record_id)
    if not row:
        raise HTTPException(404, "扣分记录不存在")
    can_collaborate = user.role.code in MATERIAL_COLLABORATOR_CODES
    if row.submitter_id != user.id and "SYSTEM_ADMIN" not in user.permissions and not can_collaborate:
        raise HTTPException(403, "仅原登记人、主管或管理员可以补充材料")
    if row.status not in {"pending_material", "material_failed"}:
        raise HTTPException(400, "该记录当前不需要补充材料")
    ensure_month_open(db, row.deduction_month, row.attraction_id_snapshot, "重新提交扣分材料")
    old_job = db.get(DeductionMaterialJob, row.material_job_id) if row.material_job_id else None
    upgrade_request = db.get(DeductionUpgradeRequest, row.upgrade_request_id) if row.upgrade_request_id else None
    first = None
    reviewer = None
    upgrade_material = False
    if upgrade_request and upgrade_request.status == "awaiting_material":
        first = db.get(DeductionRecord, upgrade_request.first_deduction_id)
        reviewer = db.get(Employee, upgrade_request.reviewer_id)
        if not first or not reviewer or first.upgrade_request_id != upgrade_request.id or row.upgrade_request_id != upgrade_request.id or first.status != "active":
            raise HTTPException(409, "关联升级声明状态已变化，无法重新提交材料")
        first.upgrade_state, first.upgrade_role = "material_processing", "source_first"
        row.upgrade_state, row.upgrade_role = "material_processing", "source_second"
        upgrade_material = True
    elif old_job and old_job.mode in {"upgrade", "upgrade_pdf_compress"}:
        first = db.get(DeductionRecord, old_job.first_deduction_id) if old_job.first_deduction_id else None
        reviewer = db.get(Employee, old_job.reviewer_id) if old_job.reviewer_id else None
        if not first or not reviewer or first.upgrade_request_id or first.status != "active":
            raise HTTPException(409, "关联升级声明状态已变化，无法重新提交材料")
        claimed = db.execute(
            update(DeductionRecord)
            .where(DeductionRecord.id == first.id, or_(DeductionRecord.upgrade_state.is_(None), DeductionRecord.upgrade_state == "eligible"))
            .values(upgrade_state="material_processing", upgrade_role="source_first")
            .execution_options(synchronize_session=False)
        )
        if claimed.rowcount != 1:
            db.rollback()
            raise HTTPException(409, "关联历史声明正在被其他工单使用，请刷新后重试")
        upgrade_material = True
    photo_mode = bool([item for item in (document_images or []) if item and item.filename])
    pdf_mode = bool(document and document.filename)
    if pdf_mode == photo_mode:
        raise HTTPException(400, "请上传一个PDF，或选择1至6张照片材料")
    source_rows: list[StoredFile] = []
    output = None
    try:
        source_rows = [await stage_pdf_material(db, document, user.id)] if pdf_mode else await stage_photo_materials(db, list(document_images or []), user.id)
        output = create_pdf_placeholder(db, user.id)
        old_output = db.get(StoredFile, row.document_file_id)
        if old_output and old_output.status == "pending_conversion":
            old_output.status = "superseded"
        row.document_file_id = output.id
        row.status = "material_processing"
        row.material_status = "processing"
        row.material_error = None
        row.material_source_type = "pdf" if pdf_mode else "photos"
        row.material_uploaded_by = user.id
        row.material_uploaded_by_name = user.name
        row.material_uploaded_at = datetime.now()
        row.material_revision = int(row.material_revision or 0) + 1
        # Records originally submitted without material have no job yet.  They
        # may be supplemented later, so never dereference a missing old job.
        row.upgrade_state = "material_processing" if upgrade_material else row.upgrade_state
        queue_photo_material_job(
            db,
            row,
            output,
            source_rows,
            mode=("upgrade_pdf_compress" if upgrade_material and pdf_mode else ("upgrade" if upgrade_material else ("pdf_compress" if pdf_mode else "deduction"))),
            reviewer_id=reviewer.id if reviewer else None,
            first_deduction_id=first.id if first else None,
        )
        write_audit(db, user.employee, "补充扣分材料", "deduction", row.id, after={"material_status": "processing", "contributor": user.name, "retry_of_job_id": old_job.id if old_job else None}, ip_address=client_ip(request))
        db.commit()
    except Exception:
        db.rollback()
        remove_upload_file(output)
        for source in source_rows:
            remove_upload_file(source)
        raise
    invalidate_data_caches()
    return {"ok": True, "record": deduction_payload(row)}


@router.post("/deductions/{record_id}/void")
def void_deduction(record_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    row = db.get(DeductionRecord, record_id)
    if not row:
        raise HTTPException(404, "扣分记录不存在")
    if row.submitter_id != user.id and "SYSTEM_ADMIN" not in user.permissions:
        raise HTTPException(403, "只有原登记人或管理员可以作废")
    if row.upgrade_request_id:
        raise HTTPException(400, "声明升级工单关联记录不可作废；请在升级工单中完成审核处理")
    if row.status not in {"active", "pending_material", "material_failed"}:
        raise HTTPException(400, "记录已经作废")
    employee = db.get(Employee, row.employee_id)
    ensure_month_open(db, row.deduction_month, row.attraction_id_snapshot or (employee.attraction_id if employee else None), "作废扣分")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "作废原因必填")
    before = deduction_payload(row)
    role_code, role_name, permission_scope = void_operator_snapshot(db, user)
    row.voided_from_status = row.status
    row.status = "void"
    row.voided_by = user.id
    row.voided_by_name = user.name
    row.voided_by_role_code = role_code
    row.voided_by_role_name = role_name
    row.void_permission_scope_snapshot = permission_scope
    row.voided_at = datetime.now()
    row.void_reason = reason
    write_audit(db, user.employee, "作废扣分", "deduction", row.id, before=before, after=deduction_payload(row), reason=reason, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {"ok": True}


def sick_leave_overlap_payload(rows: list[SickLeaveRecord]) -> list[dict]:
    return [
        {
            "id": row.id,
            "leave_start_date": row.leave_start_date,
            "leave_end_date": row.leave_end_date,
            "leave_days": float(row.leave_days),
            "submitted_at": row.submitted_at.strftime("%Y-%m-%d %H:%M:%S"),
        }
        for row in rows
    ]


def active_sick_leave_overlaps(db: Session, employee_id: int, start_date: str, end_date: str) -> list[SickLeaveRecord]:
    """Return any active absence that intersects an inclusive date range."""
    return (
        db.query(SickLeaveRecord)
        .filter(
            SickLeaveRecord.employee_id == employee_id,
            SickLeaveRecord.status == "active",
            SickLeaveRecord.leave_start_date <= end_date,
            SickLeaveRecord.leave_end_date >= start_date,
        )
        .order_by(SickLeaveRecord.leave_start_date, SickLeaveRecord.id)
        .all()
    )


def sick_leave_overlap_detail(rows: list[SickLeaveRecord]) -> dict:
    records = sick_leave_overlap_payload(rows)
    first = records[0]
    return {
        "code": "SICK_LEAVE_DATE_OVERLAP",
        "message": f"该员工已存在 {first['leave_start_date']} 至 {first['leave_end_date']} 的缺勤登记；日期有交集，不能重复提交。请先作废或更正原记录后再登记。",
        "records": records,
    }


@router.get("/sick-leaves/overlap-check")
def check_sick_leave_overlap(
    employee_id: int,
    leave_start_date: str,
    leave_end_date: str,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("SICK_REGISTER")),
):
    start = parse_iso_date(leave_start_date, "病假开始日期")
    end = parse_iso_date(leave_end_date, "病假结束日期")
    if end < start or start.strftime("%Y-%m") != end.strftime("%Y-%m"):
        raise HTTPException(400, "病假日期无效，跨月请分开登记")
    target, _ = ensure_enabled_frontline_target(db, employee_id, "病假")
    ensure_scoped_hr_employee(db, user, target)
    ensure_operational_target_scope(user, target)
    rows = active_sick_leave_overlaps(db, target.id, leave_start_date, leave_end_date)
    return {"conflict": bool(rows), "detail": sick_leave_overlap_detail(rows) if rows else None}


@router.get("/sick-leaves/violation-upgrade-preview")
def sick_leave_violation_upgrade_preview(
    employee_id: int,
    leave_start_date: str,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("SICK_REGISTER")),
):
    """Preview the statement-upgrade path used by the absence-page shortcut."""
    if not ({"DEDUCTION_DIRECT", "DEDUCTION_ALL"} & user.permissions):
        raise HTTPException(403, "当前账号没有登记违规病假声明的权限")
    parse_iso_date(leave_start_date, "病假开始日期")
    target, _ = ensure_enabled_frontline_target(db, employee_id, "违规病假")
    ensure_scoped_hr_employee(db, user, target)
    ensure_operational_target_scope(user, target)
    deduction_type = db.query(DeductionType).filter_by(code="SICK_LEAVE_VIOLATION", active=True).first()
    if not deduction_type:
        raise HTTPException(409, "违规病假声明配置缺失，请联系管理员")
    first = statement_upgrade_candidate(db, target.id, deduction_type.id, leave_start_date)
    return {
        "eligible": bool(first),
        "first_record": deduction_payload(first) if first else None,
        "reviewers": statement_upgrade_reviewer_options(db) if first else [],
    }


@router.post("/sick-leaves")
async def create_sick_leave(
    request: Request,
    employee_id: int = Form(...),
    leave_start_date: str = Form(...),
    leave_end_date: str = Form(...),
    leave_days: str | None = Form(None),
    rest_day_confirmed: bool = Form(False),
    note: str | None = Form(None),
    is_violation: bool = Form(False),
    violation_reviewer_id: int | None = Form(None),
    idempotency_key: str | None = Form(None),
    proof: UploadFile = File(...),
    violation_document: UploadFile | None = File(None),
    violation_document_images: list[UploadFile] | None = File(None),
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("SICK_REGISTER")),
):
    if not str(employee_id or "").strip():
        raise HTTPException(400, "请先通过搜索结果选择缺勤员工")
    request_key = normalize_request_key(idempotency_key)
    payload_digest = submission_payload_digest(
        {
            "employee_id": int(employee_id),
            "leave_start_date": str(leave_start_date or "").strip(),
            "leave_end_date": str(leave_end_date or "").strip(),
            "leave_days": str(leave_days or "").strip(),
            "rest_day_confirmed": bool(rest_day_confirmed),
            "note": str(note or "").strip(),
            "is_violation": bool(is_violation),
        }
    )
    duplicate_row = existing_submission(db, user.id, "sick_leave", request_key, SickLeaveRecord, payload_digest)
    if duplicate_row:
        attendance = ensure_month_attendance(db, duplicate_row.attendance_month, {duplicate_row.employee_id})[duplicate_row.employee_id]
        return {"ok": True, "attendance_score": float(attendance.final_score), "duplicate": True}
    start = parse_iso_date(leave_start_date, "病假开始日期")
    end = parse_iso_date(leave_end_date, "病假结束日期")
    if end < start or start.strftime("%Y-%m") != end.strftime("%Y-%m"):
        raise HTTPException(400, "病假日期无效，跨月请分开登记")
    if end > start and not rest_day_confirmed:
        raise HTTPException(
            409,
            detail={
                "code": "SICK_LEAVE_REST_DAY_CONFIRMATION_REQUIRED",
                "message": "请确认您所提交的病假日期中不含演职人员本休。",
            },
        )
    calendar_days = Decimal((end - start).days + 1)
    try:
        days = Decimal(leave_days) if leave_days not in (None, "") else calendar_days
    except InvalidOperation as exc:
        raise HTTPException(400, "病假天数必须是数字") from exc
    if days <= 0 or days * 2 != (days * 2).to_integral_value():
        raise HTTPException(400, "病假天数必须按0.5天递增")
    if days > calendar_days:
        raise HTTPException(400, "病假天数不能超过日期范围")
    target, target_role = ensure_enabled_frontline_target(db, employee_id, "病假")
    ensure_scoped_hr_employee(db, user, target)
    ensure_operational_target_scope(user, target)
    ensure_month_open(db, start.strftime("%Y-%m"), target.attraction_id, "新增病假")
    overlaps = active_sick_leave_overlaps(db, target.id, leave_start_date, leave_end_date)
    if overlaps:
        raise HTTPException(409, detail=sick_leave_overlap_detail(overlaps))
    if not proof.filename:
        raise HTTPException(400, "缺勤证明未上传，请重新选择图片或PDF文件")
    violation_type = statement = first_statement = reviewer = None
    violation_photo_mode = bool([item for item in (violation_document_images or []) if item and item.filename])
    violation_pdf_mode = bool(violation_document and violation_document.filename)
    if violation_pdf_mode and violation_photo_mode:
        raise HTTPException(400, "违规病假声明请在上传PDF和照片材料中选择一种方式")
    if is_violation:
        if not ({"DEDUCTION_DIRECT", "DEDUCTION_ALL"} & user.permissions):
            raise HTTPException(403, "当前账号没有登记违规病假声明的权限")
        violation_type = db.query(DeductionType).filter_by(code="SICK_LEAVE_VIOLATION", active=True).first()
        statement = db.query(DeductionLevel).filter_by(code="STATEMENT", active=True).first()
        if not violation_type or not statement:
            raise HTTPException(409, "违规病假声明配置缺失，请联系管理员")
        ensure_deduction_type_allowed(user, violation_type)
        existing_pending_material = pending_material_conflict(
            db,
            employee_id=target.id,
            deduction_type_id=violation_type.id,
            deduction_level_id=statement.id,
            occurred_on=leave_start_date,
        )
        if existing_pending_material:
            raise_pending_material_conflict(existing_pending_material)
        first_statement = statement_upgrade_candidate(db, target.id, violation_type.id, leave_start_date)
        if first_statement:
            reviewer = db.get(Employee, violation_reviewer_id) if violation_reviewer_id else None
            reviewer_role = role_at(db, reviewer.id) if reviewer else None
            reviewer_account = db.query(UserAccount).filter_by(employee_id=reviewer.id, enabled=True).first() if reviewer else None
            if not reviewer or not reviewer_role or reviewer_role.code not in UPGRADE_REVIEWER_CODES or not reviewer.is_active or not reviewer_account:
                raise HTTPException(400, "检测到3个月内可升级的违规病假声明，请选择在职且账号启用的GSM或TA GSM审核")
    file_row = violation_file_row = None
    violation_source_rows: list[StoredFile] = []
    try:
        file_row = await save_upload(db, proof, user.id, allowed_extensions={".jpg", ".jpeg", ".png", ".heic", ".webp", ".pdf"})
        row = SickLeaveRecord(
            employee_id=target.id,
            employee_no_snapshot=target.employee_no,
            employee_name_snapshot=target.name,
            employee_role_snapshot=target_role.name,
            attraction_id_snapshot=target.attraction_id,
            attendance_month=start.strftime("%Y-%m"),
            leave_start_date=leave_start_date,
            leave_end_date=leave_end_date,
            leave_days=days,
            charged_days=days,
            proof_file_id=file_row.id,
            note=(note or "").strip() or None,
            status="active",
            submitted_by=user.id,
            submitted_by_name=user.name,
            is_violation=bool(is_violation),
        )
        db.add(row)
        db.flush()
        if is_violation:
            violation_needs_processing = False
            if violation_photo_mode:
                violation_source_rows = await stage_photo_materials(db, list(violation_document_images or []), user.id)
                violation_file_row = create_pdf_placeholder(db, user.id)
            elif violation_pdf_mode:
                violation_file_row = await stage_pdf_material(db, violation_document, user.id)
                violation_needs_processing = violation_file_row.file_size > PDF_HIGH_QUALITY_OPTIMIZATION_THRESHOLD
                if violation_needs_processing:
                    violation_file_row.status = "processing_source"
                    violation_source_rows = [violation_file_row]
                    violation_file_row = create_pdf_placeholder(db, user.id)
                else:
                    violation_file_row.status = "active"
            else:
                violation_file_row = create_pdf_placeholder(db, user.id)
                violation_file_row.original_filename = "待补充违规病假声明材料.pdf"
                violation_file_row.status = "pending_material"
            group = current_group_for_employee(db, target.id)
            material_processing = violation_photo_mode or violation_needs_processing
            awaiting_material = not violation_photo_mode and not violation_pdf_mode
            violation = DeductionRecord(
                employee_id=target.id, employee_no=target.employee_no, employee_name=target.name,
                employee_role_snapshot=target_role.name, employee_group_id_snapshot=group.id if group else None,
                attraction_id_snapshot=target.attraction_id, deduction_type_id=violation_type.id,
                deduction_type_name=violation_type.name, deduction_level_id=statement.id,
                deduction_level_name=statement.name, points=Decimal("0") if first_statement else statement.points, occurred_on=leave_start_date,
                deduction_month=leave_start_date[:7], description=f"违规病假（关联缺勤：{leave_start_date} 至 {leave_end_date}）：" + ((note or "").strip() or "待补充处理说明"),
                document_file_id=violation_file_row.id, submitter_id=user.id, submitter_name=user.name,
                submitter_role_snapshot=user.role.name, permission_scope_snapshot="缺勤登记违规病假声明",
                status=("material_processing" if material_processing else ("pending_upgrade" if first_statement and not awaiting_material else ("active" if violation_pdf_mode else "pending_material"))),
                material_status="processing" if material_processing else ("ready" if violation_pdf_mode else "missing"),
                material_source_type="photos" if violation_photo_mode else ("pdf" if violation_pdf_mode else "later"),
                upgrade_role="source_second" if first_statement else None,
                upgrade_state=("material_processing" if material_processing else ("awaiting_material" if first_statement and awaiting_material else ("pending" if first_statement else None))),
            )
            db.add(violation)
            db.flush()
            row.violation_deduction_id = violation.id
            if first_statement:
                claimed = db.execute(
                    update(DeductionRecord)
                    .where(
                        DeductionRecord.id == first_statement.id,
                        DeductionRecord.status == "active",
                        or_(DeductionRecord.upgrade_state.is_(None), DeductionRecord.upgrade_state == "eligible"),
                    )
                    .values(
                        upgrade_role="source_first",
                        upgrade_state="awaiting_material" if awaiting_material else ("material_processing" if material_processing else "pending"),
                    )
                    .execution_options(synchronize_session=False)
                )
                if claimed.rowcount != 1:
                    raise HTTPException(409, "该历史违规病假声明已被其他升级工单使用，请刷新后重试")
                upgrade = DeductionUpgradeRequest(
                    employee_id=target.id, deduction_type_id=violation_type.id,
                    first_deduction_id=first_statement.id, second_deduction_id=violation.id,
                    reviewer_id=reviewer.id, reviewer_name=reviewer.name,
                    submitted_by=user.id, submitted_by_name=user.name,
                    status="awaiting_material" if (awaiting_material or material_processing) else "pending",
                )
                db.add(upgrade)
                db.flush()
                first_statement.upgrade_request_id = upgrade.id
                violation.upgrade_request_id = upgrade.id
                if material_processing:
                    queue_photo_material_job(
                        db, violation, violation_file_row, violation_source_rows,
                        mode="upgrade_pdf_compress" if violation_needs_processing else "upgrade",
                        reviewer_id=reviewer.id, first_deduction_id=first_statement.id,
                    )
            elif material_processing:
                queue_photo_material_job(
                    db, violation, violation_file_row, violation_source_rows,
                    mode="pdf_compress" if violation_needs_processing else "deduction",
                )
        remember_submission(db, user.id, "sick_leave", request_key, row.id, payload_digest)
        attendance = recalculate_attendance(db, target, row.attendance_month)
        write_audit(db, user.employee, "登记病假", "sick_leave", row.id, after={"employee": target.employee_no, "days": float(days), "score": float(attendance.final_score), "is_violation": bool(is_violation), "violation_deduction_id": row.violation_deduction_id}, ip_address=client_ip(request))
        db.commit()
    except IntegrityError:
        db.rollback()
        remove_upload_file(file_row)
        remove_upload_file(violation_file_row)
        for source_row in violation_source_rows:
            remove_upload_file(source_row)
        duplicate_row = existing_submission(db, user.id, "sick_leave", request_key, SickLeaveRecord, payload_digest)
        if duplicate_row:
            attendance = ensure_month_attendance(db, duplicate_row.attendance_month, {duplicate_row.employee_id})[duplicate_row.employee_id]
            return {"ok": True, "attendance_score": float(attendance.final_score), "duplicate": True}
        raise
    except Exception:
        db.rollback()
        remove_upload_file(file_row)
        remove_upload_file(violation_file_row)
        for source_row in violation_source_rows:
            remove_upload_file(source_row)
        raise
    invalidate_data_caches()
    return {
        "ok": True,
        "attendance_score": float(attendance.final_score),
        "violation": deduction_payload(violation) if is_violation else None,
        "violation_upgrade": bool(first_statement),
    }


@router.get("/sick-leaves")
def list_sick_leaves(month: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SICK_REGISTER"))):
    query = db.query(SickLeaveRecord).filter(SickLeaveRecord.submitted_by == user.id)
    if month:
        query = query.filter(SickLeaveRecord.attendance_month == month)
    rows = query.order_by(SickLeaveRecord.submitted_at.desc()).limit(500).all()
    return sick_leave_payloads(db, rows)


@router.post("/sick-leaves/{record_id}/void")
def void_sick_leave(record_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SICK_REGISTER"))):
    row = db.get(SickLeaveRecord, record_id)
    if not row or row.submitted_by != user.id:
        raise HTTPException(404, "只能作废自己登记的病假")
    if row.status != "active":
        raise HTTPException(400, "病假已经作废")
    employee = db.get(Employee, row.employee_id)
    ensure_month_open(db, row.attendance_month, row.attraction_id_snapshot or (employee.attraction_id if employee else None), "作废病假")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "作废原因必填")
    role_code, role_name, permission_scope = void_operator_snapshot(db, user)
    row.voided_from_status = row.status
    row.status = "void"
    row.voided_by = user.id
    row.voided_by_name = user.name
    row.voided_by_role_code = role_code
    row.voided_by_role_name = role_name
    row.void_permission_scope_snapshot = permission_scope
    row.voided_at = datetime.now()
    row.void_reason = reason
    db.flush()
    attendance = recalculate_attendance(db, db.get(Employee, row.employee_id), row.attendance_month)
    write_audit(db, user.employee, "作废病假", "sick_leave", row.id, after={"score": float(attendance.final_score)}, reason=reason, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "attendance_score": float(attendance.final_score)}


@router.get("/files/{file_id}")
def download_file(file_id: int, preview: bool = False, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    row = db.get(StoredFile, file_id)
    if not row or row.status != "active":
        raise HTTPException(404, "文件不存在")
    authorized = row.uploaded_by == user.id or "SYSTEM_ADMIN" in user.permissions
    if user.role.code == SCOPED_HR_ROLE_CODE:
        managed_attractions = scoped_hr_attraction_ids(db, user) or set()
    else:
        managed_attractions = {
            attraction_id
            for (attraction_id,) in db.query(Attraction.id)
            .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
            .all()
        } if {"DATA_VIEW", "DATA_EXPORT"} & user.permissions else set()
    recognition_link = db.query(RecognitionAttachment).filter(RecognitionAttachment.file_id == file_id).first()
    if recognition_link:
        recognition = db.get(RecognitionRecord, recognition_link.recognition_id)
        authorized = authorized or bool(
            recognition
            and (
                user.id in {recognition.employee_id, recognition.operator_employee_id, recognition.assigned_reviewer_id, recognition.reviewed_by}
                or recognition.employee_id in direct_member_ids(db, user.id)
                or recognition.home_attraction_id in managed_attractions
            )
        )
    deduction = db.query(DeductionRecord).filter(DeductionRecord.document_file_id == file_id).first()
    if deduction:
        authorized = authorized or user.id in {deduction.employee_id, deduction.submitter_id} or deduction.employee_id in direct_member_ids(db, user.id) or deduction.attraction_id_snapshot in managed_attractions
    sick_leave = db.query(SickLeaveRecord).filter(SickLeaveRecord.proof_file_id == file_id).first()
    if sick_leave:
        sick_employee = db.get(Employee, sick_leave.employee_id)
        authorized = authorized or user.id in {sick_leave.employee_id, sick_leave.submitted_by} or sick_leave.employee_id in direct_member_ids(db, user.id) or bool((sick_leave.attraction_id_snapshot or (sick_employee.attraction_id if sick_employee else None)) in managed_attractions)
    if not authorized:
        raise HTTPException(403, "没有权限查看该文件")
    path = (FILE_DIR / row.storage_key).resolve()
    if FILE_DIR.resolve() not in path.parents or not path.exists():
        raise HTTPException(404, "文件不存在")
    suffix = path.suffix.lower()
    image_suffixes = PREVIEW_IMAGE_EXTENSIONS
    preview_suffixes = image_suffixes | PREVIEW_PDF_EXTENSIONS
    if preview and suffix not in preview_suffixes:
        raise HTTPException(415, "该材料不是可预览图片或PDF")
    cache_hit = False
    try:
        if suffix == ".pdf":
            content = watermark_pdf(path, user.employee.employee_no)
            media_type = "application/pdf"
        elif suffix in image_suffixes:
            if preview:
                cache_key = f"{row.sha256}:{user.employee.employee_no}:image-preview-v1-1920"
                cached = watermarked_preview_cache.get(cache_key)
                if cached:
                    content, media_type = cached
                    cache_hit = True
                else:
                    content, media_type = watermark_image(path, user.employee.employee_no, max_dimension=1920)
                    watermarked_preview_cache.put(cache_key, content, media_type)
            else:
                content, media_type = watermark_image(path, user.employee.employee_no)
        else:
            raise HTTPException(415, "该文件类型暂不支持安全水印下载")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(500, "文件水印生成失败，请联系管理员") from exc
    context_type = "recognition" if recognition_link else "deduction" if deduction else "sick_leave" if sick_leave else "stored_file"
    write_audit(
        db,
        user.employee,
        "查看带水印材料" if preview else "下载带水印材料",
        context_type,
        file_id,
        after={
            "filename": row.original_filename,
            "watermark_account": user.employee.employee_no,
            "delivery_mode": "preview" if preview else "download",
            "preview_cache": "hit" if cache_hit else "miss" if preview else "not_used",
        },
    )
    db.commit()
    inline_preview = preview and (media_type.startswith("image/") or media_type == "application/pdf")
    headers = {
        "Content-Disposition": f"{'inline' if inline_preview else 'attachment'}; filename*=UTF-8''{quote(row.original_filename)}",
        "Cache-Control": "private, no-store",
        "X-Preview-Cache": "HIT" if cache_hit else "MISS" if preview else "BYPASS",
    }
    return StreamingResponse(BytesIO(content), media_type=media_type, headers=headers)


@router.post("/security/screenshot-event")
def record_screenshot_event(request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    write_audit(
        db,
        user.employee,
        "检测到截图按键",
        "user_account",
        user.account.id,
        after={"employee_no": user.employee.employee_no, "detection": "print_screen_key"},
        ip_address=client_ip(request),
    )
    db.commit()
    return {"ok": True, "message": "系统已记录截图按键事件；页面访问和导出均受审计，请勿分享敏感数据。"}


@router.get("/governance/appealable-records")
def appealable_records(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if user.role.code not in FRONTLINE_CODES:
        raise HTTPException(403, "仅CM/TR可提交本人记录申诉")
    records: list[dict] = []
    for row in db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id == user.id,
        RecognitionRecord.status.in_(("rejected", "void")),
    ).order_by(RecognitionRecord.submitted_at.desc()).limit(100).all():
        records.append({"record_type": "recognition", "record_id": row.id, "label": f"认可 · {row.recognition_date} · {row.recognition_type_name} · {row.status}"})
    for row in db.query(DeductionRecord).filter(
        DeductionRecord.employee_id == user.id,
        DeductionRecord.status == "active",
    ).order_by(DeductionRecord.submitted_at.desc()).limit(100).all():
        records.append({"record_type": "deduction", "record_id": row.id, "label": f"扣分 · {row.occurred_on} · {row.deduction_type_name} · -{Decimal(row.points):.2f}分"})
    return {"items": records}


@router.post("/governance/appeals")
def submit_appeal(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if user.role.code not in FRONTLINE_CODES:
        raise HTTPException(403, "仅CM/TR可提交本人记录申诉")
    record_type = str(payload.get("record_type") or "").strip()
    try:
        record_id = int(payload.get("record_id"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "请选择需要申诉的记录") from exc
    reason = str(payload.get("reason") or "").strip()
    if len(reason) < 5 or len(reason) > 500:
        raise HTTPException(400, "申诉说明需为5至500字")
    subject = appeal_record_for_employee(db, user.id, record_type, record_id)
    if not subject:
        raise HTTPException(404, "该记录不存在、当前不可申诉或不属于本人")
    _record, attraction_id, score_month, _conflicts = subject
    existing = db.query(GovernanceCase).filter(
        GovernanceCase.case_type == "appeal",
        GovernanceCase.record_type == record_type,
        GovernanceCase.record_id == record_id,
        GovernanceCase.status == "open",
    ).first()
    if existing:
        raise HTTPException(409, "该记录已有待处理申诉")
    row = GovernanceCase(
        case_type="appeal",
        record_type=record_type,
        record_id=record_id,
        attraction_id=attraction_id,
        score_month=score_month,
        subject_employee_id=user.id,
        submitted_by=user.id,
        submitted_by_name=user.name,
        reason=reason,
        due_at=datetime.now() + timedelta(hours=72),
    )
    db.add(row)
    db.flush()
    write_audit(db, user.employee, "提交记录申诉", "governance_case", row.id, after={"record_type": record_type, "record_id": record_id, "score_month": score_month}, reason=reason, ip_address=client_ip(request))
    db.commit()
    return {"ok": True, "case": governance_case_payload(row)}


@router.get("/governance/cases")
def governance_cases(db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if user.role.code in FRONTLINE_CODES:
        rows = db.query(GovernanceCase).filter(GovernanceCase.submitted_by == user.id).order_by(GovernanceCase.submitted_at.desc()).limit(100).all()
        return {"items": [governance_case_payload(row) for row in rows], "can_review": False}
    if user.role.code not in {"HR_CIRCLE", "HR_ADMIN", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "当前角色没有治理复核权限")
    rows = [
        row for row in db.query(GovernanceCase).order_by(GovernanceCase.status.asc(), GovernanceCase.due_at.asc(), GovernanceCase.id.desc()).limit(500).all()
        if governance_scope_allows_case(db, user, row)
    ]
    return {"items": [governance_case_payload(row, can_resolve=True) for row in rows], "can_review": True}


@router.post("/governance/cases/{case_id}/resolve")
def resolve_governance_case(case_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(current_user)):
    if user.role.code not in {"HR_CIRCLE", "HR_ADMIN", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "当前角色没有治理复核权限")
    row = db.get(GovernanceCase, case_id)
    if not row or not governance_scope_allows_case(db, user, row):
        raise HTTPException(404, "未找到可处理的治理事项")
    if row.status != "open":
        raise HTTPException(409, "该事项已处理")
    decision = str(payload.get("decision") or "").strip()
    if decision not in {"uphold", "correction_required"}:
        raise HTTPException(400, "请选择维持原记录或需要更正")
    resolution = str(payload.get("resolution") or "").strip()
    if len(resolution) < 5 or len(resolution) > 500:
        raise HTTPException(400, "处理说明需为5至500字")
    conflicts: set[int | None] = {row.submitted_by}
    if row.case_type == "appeal":
        subject = appeal_record_for_employee(db, row.subject_employee_id or -1, row.record_type or "", row.record_id or 0)
        if not subject:
            raise HTTPException(409, "原记录状态已变化，请由最高管理员在审计日志中复核")
        _record, _attraction_id, _score_month, record_conflicts = subject
        conflicts.update(record_conflicts)
    if user.id in conflicts:
        raise HTTPException(409, "处理人不得是申诉提交人、原登记人、原复核人或原作废人")
    row.status = "resolved"
    row.decision = decision
    row.resolution = resolution
    row.resolved_by = user.id
    row.resolved_by_name = user.name
    row.resolved_at = datetime.now()
    write_audit(db, user.employee, "处理治理事项", "governance_case", row.id, before={"status": "open"}, after={"status": row.status, "decision": decision}, reason=resolution, ip_address=client_ip(request))
    db.commit()
    return {"ok": True, "case": governance_case_payload(row)}


@router.get("/month-closes/{month}")
def month_close_status(
    month: str,
    attraction_id: int | None = None,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    month = parse_score_month(month)
    ensure_month_close_scope(db, user, attraction_id)
    return month_closure_payload(db, month, attraction_id, user)


@router.post("/month-closes/{month}/close")
def close_month(
    month: str,
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    month = parse_score_month(month)
    raw_attraction_id = payload.get("attraction_id")
    attraction_id = int(raw_attraction_id) if raw_attraction_id not in (None, "") else None
    attraction = ensure_month_close_scope(db, user, attraction_id)
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "关闭月结原因必填")
    effective = month_closure_scope(db, month, attraction_id)
    if effective:
        effective_attraction = db.get(Attraction, effective.attraction_id) if effective.attraction_id else None
        scope_name = effective_attraction.name if effective_attraction else "全部景点圈"
        raise HTTPException(409, f"{month} · {scope_name}已经关闭月结")
    checklist = month_close_checklist(db, month, attraction_id)
    blocking = [item for item in checklist if item["count"]]
    if blocking:
        raise HTTPException(
            409,
            detail={
                "code": "MONTH_CLOSE_CHECKLIST_INCOMPLETE",
                "message": "月结前仍有待办或待跟进事项，请先处理完成。",
                "items": blocking,
            },
        )
    scope_filter = MonthClosure.attraction_id == attraction_id if attraction_id is not None else MonthClosure.attraction_id.is_(None)
    row = db.query(MonthClosure).filter(MonthClosure.closure_month == month, scope_filter).first()
    before = {
        "status": row.status,
        "reopened_at": row.reopened_at.strftime("%Y-%m-%d %H:%M:%S") if row.reopened_at else "",
        "reopen_reason": row.reopen_reason or "",
    } if row else None
    now = datetime.now()
    values = {
        "status": "closed",
        "closed_by": user.id,
        "closed_by_name": user.name,
        "closed_at": now,
        "close_reason": reason,
        "reopened_by": None,
        "reopened_by_name": None,
        "reopened_at": None,
        "reopen_reason": None,
        "updated_at": now,
    }
    if row:
        claimed = db.execute(
            update(MonthClosure)
            .where(MonthClosure.id == row.id, MonthClosure.status == "open")
            .values(**values)
            .execution_options(synchronize_session=False)
        )
        if claimed.rowcount != 1:
            db.rollback()
            raise HTTPException(409, "该月份和景点圈刚刚已由其他管理员关闭月结")
        db.refresh(row)
    else:
        row = MonthClosure(closure_month=month, attraction_id=attraction_id, **values)
        db.add(row)
        try:
            db.flush()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(409, "该月份和景点圈刚刚已由其他管理员关闭月结") from exc
    scope_name = attraction.name if attraction else "全部景点圈"
    snapshot_count = capture_month_organization_snapshots(db, month, attraction_id)
    write_audit(
        db,
        user.employee,
        "关闭月结",
        "month_close",
        row.id,
        before=before,
        after={"month": month, "attraction_id": attraction_id, "scope": scope_name, "status": "closed", "organization_snapshots_created": snapshot_count, "checklist": "completed", "rule_effective_date": MONTH_CLOSE_EFFECTIVE_DATE},
        reason=reason,
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, **month_closure_payload(db, month, attraction_id, user)}


@router.post("/month-closes/{month}/reopen")
def reopen_month(
    month: str,
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    month = parse_score_month(month)
    raw_attraction_id = payload.get("attraction_id")
    attraction_id = int(raw_attraction_id) if raw_attraction_id not in (None, "") else None
    attraction = ensure_month_close_scope(db, user, attraction_id)
    if user.role.code not in MONTH_CLOSE_ROLE_CODES:
        raise HTTPException(403, "仅景点圈HR或最高管理员可以临时开放月结")
    if user.role.code == "HR_CIRCLE" and (attraction_id is None or attraction_id not in scoped_hr_attraction_ids(db, user)):
        raise HTTPException(403, "景点圈HR只能临时开放所属景点圈的月结")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "重新开启原因必填")
    scope_filter = MonthClosure.attraction_id == attraction_id if attraction_id is not None else MonthClosure.attraction_id.is_(None)
    row = db.query(MonthClosure).filter(
        MonthClosure.closure_month == month,
        scope_filter,
        MonthClosure.status == "closed",
    ).first()
    if not row:
        raise HTTPException(409, "该月份和景点圈当前没有可重开的月结")
    before = {
        "month": row.closure_month,
        "attraction_id": row.attraction_id,
        "status": row.status,
        "closed_by_name": row.closed_by_name or "",
        "closed_at": row.closed_at.strftime("%Y-%m-%d %H:%M:%S") if row.closed_at else "",
        "close_reason": row.close_reason or "",
    }
    row.status = "open"
    row.reopened_by = user.id
    row.reopened_by_name = user.name
    row.reopened_at = datetime.now()
    row.reopen_reason = reason
    row.updated_at = datetime.now()
    scope_name = attraction.name if attraction else "全部景点圈"
    correction_case = GovernanceCase(
        case_type="month_correction",
        record_type="month_close",
        record_id=row.id,
        attraction_id=attraction_id,
        score_month=month,
        subject_employee_id=None,
        submitted_by=user.id,
        submitted_by_name=user.name,
        reason=reason,
        status="resolved",
        decision="month_reopened",
        resolution="已按景点圈HR受控权限临时开放月结；后续更正仍通过既有受控写入与审计路径完成。",
        resolved_by=user.id,
        resolved_by_name=user.name,
        resolved_at=datetime.now(),
        due_at=datetime.now(),
    )
    db.add(correction_case)
    db.flush()
    write_audit(
        db,
        user.employee,
        "重新开启月结",
        "month_close",
        row.id,
        before=before,
        after={"month": month, "attraction_id": attraction_id, "scope": scope_name, "status": "open", "correction_case_id": correction_case.id},
        reason=reason,
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, **month_closure_payload(db, month, attraction_id, user)}


def months_between(start: date, end: date) -> list[str]:
    result = []
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        result.append(f"{year:04d}-{month:02d}")
        if month == 12:
            year, month = year + 1, 1
        else:
            month += 1
    return result


def ranking_employees(db: Session, attraction_id: int | None, on_date: str, role_codes: set[str]) -> tuple[list[Employee], dict[int, Role]]:
    circle_ids = [
        row.id
        for row in db.query(Attraction.id)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .order_by(Attraction.id)
        .all()
    ]
    query = db.query(Employee).filter(Employee.is_active.is_(True), Employee.attraction_id.in_(circle_ids or [-1]))
    if attraction_id is not None:
        query = query.filter(Employee.attraction_id == attraction_id)
    candidates = query.order_by(Employee.name, Employee.employee_no).all()
    role_map = roles_at(db, [employee.id for employee in candidates], on_date)
    selected = [employee for employee in candidates if role_map.get(employee.id) and role_map[employee.id].code in role_codes]
    return selected, role_map


def ranking_leader_names(db: Session, employee_ids: list[int], on_date: str) -> dict[int, str]:
    if not employee_ids:
        return {}
    memberships = (
        db.query(GroupMembership)
        .filter(
            GroupMembership.employee_id.in_(employee_ids),
            GroupMembership.status == "active",
            GroupMembership.starts_on <= on_date,
            or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= on_date),
        )
        .order_by(GroupMembership.employee_id, GroupMembership.starts_on.desc(), GroupMembership.id.desc())
        .all()
    )
    membership_by_employee: dict[int, GroupMembership] = {}
    for membership in memberships:
        membership_by_employee.setdefault(membership.employee_id, membership)
    group_ids = {membership.group_id for membership in membership_by_employee.values()}
    assignments = (
        db.query(GroupLeaderAssignment)
        .filter(
            GroupLeaderAssignment.group_id.in_(group_ids),
            GroupLeaderAssignment.status == "active",
            GroupLeaderAssignment.starts_on <= on_date,
            or_(GroupLeaderAssignment.ends_on.is_(None), GroupLeaderAssignment.ends_on >= on_date),
        )
        .order_by(GroupLeaderAssignment.group_id, GroupLeaderAssignment.starts_on.desc(), GroupLeaderAssignment.id.desc())
        .all()
        if group_ids
        else []
    )
    assignment_by_group: dict[int, GroupLeaderAssignment] = {}
    for assignment in assignments:
        assignment_by_group.setdefault(assignment.group_id, assignment)
    leader_ids = {assignment.leader_employee_id for assignment in assignment_by_group.values()}
    leader_names = {
        employee.id: employee.name
        for employee in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()
    } if leader_ids else {}
    return {
        employee_id: leader_names.get(assignment_by_group[membership.group_id].leader_employee_id, "未分配主管")
        if membership.group_id in assignment_by_group
        else "未分配主管"
        for employee_id, membership in membership_by_employee.items()
    }


def gsm_recognizer_ranking_payload(db: Session, start_value: str, end_value: str, subtype_id: int | None, keyword: str, page: int, page_size: int, attraction: Attraction | None) -> dict:
    subtype = db.get(RecognitionType, subtype_id) if subtype_id else None
    if subtype_id and (not subtype or not subtype.active):
        raise HTTPException(400, "请选择有效认可类型")
    query = db.query(RecognitionRecord).filter(
        RecognitionRecord.status == "confirmed",
        RecognitionRecord.recognition_date >= start_value,
        RecognitionRecord.recognition_date <= end_value,
    )
    if attraction:
        query = query.filter(RecognitionRecord.home_attraction_id == attraction.id)
    if subtype:
        query = query.filter(RecognitionRecord.recognition_type_id == subtype.id)
    aggregate: dict[int, dict] = {}
    for row in query.all():
        participants: dict[int, str | None] = {
            row.recognizer_employee_id: row.recognizer_role_code_snapshot,
            row.operator_employee_id: row.operator_role_code_snapshot,
        }
        for participant_id, snap_code in participants.items():
            # Older rows did not have role snapshots. Resolve them by the
            # recognition date, never by the employee's current role.
            code = snap_code or (role_at(db, participant_id, row.recognition_date).code if role_at(db, participant_id, row.recognition_date) else None)
            if code not in GSM_CODES:
                continue
            data = aggregate.setdefault(participant_id, {"count": 0, "score": 0.0, "recent_date": row.recognition_date, "role_code": code})
            data["count"] += 1
            data["score"] += float(row.credited_fraction or 0)
            data["recent_date"] = max(data["recent_date"], row.recognition_date)
    employee_rows = db.query(Employee).filter(Employee.id.in_(aggregate.keys()) if aggregate else Employee.id == -1).all()
    value = keyword.strip().lower()
    rows=[]
    for employee in employee_rows:
        if value and value not in employee.name.lower() and value not in employee.employee_no.lower():
            continue
        data=aggregate[employee.id]
        role = db.query(Role).filter(Role.code == data["role_code"]).first()
        rows.append({"employee_id":employee.id,"employee_no":employee.employee_no,"employee_name":employee.name,"role_name":role.name if role else data["role_code"],"leader_name":"","count":int(data["count"]),"score":round(float(data["score"]),2),"recent_date":data["recent_date"],"leave_days":0,"charged_days":0,"recognition_score":0,"deduction_score":0,"attendance_score":0,"total_score":0})
    rows.sort(key=lambda item:(-item["count"],-item["score"],item["employee_no"]))
    for index,item in enumerate(rows,1): item["rank"]=index
    offset=(page-1)*page_size
    return {"category":"gsm_leader","subtype_id":subtype_id,"subtype_name":subtype.name if subtype else "全部加分类别","sort_by":"count","start_date":start_value,"end_date":end_value,"attraction_id":attraction.id if attraction else None,"attraction_name":attraction.name if attraction else "全部景点圈","total":len(rows),"page":page,"page_size":page_size,"pages":max(1,ceil(len(rows)/page_size)),"rows":rows[offset:offset+page_size]}


def pr_ranking_payload(
    db: Session,
    user: V2User,
    start_value: str,
    end_value: str,
    category: str,
    subtype_id: int | None,
    sort_by: str,
    keyword: str,
    page: int,
    page_size: int,
    attraction_id: int | None = None,
) -> dict:
    if user.role.code not in GSM_CODES:
        raise HTTPException(403, "仅TA GSM和GSM可以查看PR排名数据")
    start = parse_iso_date(start_value, "开始日期")
    end = parse_iso_date(end_value, "结束日期")
    if start > end:
        raise HTTPException(400, "开始日期不能晚于结束日期")
    if end > date.today():
        raise HTTPException(400, "结束日期不能晚于今天")
    if (end - start).days > 366:
        raise HTTPException(400, "单次排名查询最多支持367天")
    if category not in {"overall", "recognition", "deduction", "absence", "leader", "gsm_leader"}:
        raise HTTPException(400, "排名类型无效")
    if sort_by not in {"score", "count"}:
        raise HTTPException(400, "排名依据无效")

    attraction = None
    if attraction_id is not None:
        attraction = (
            db.query(Attraction)
            .filter(
                Attraction.id == attraction_id,
                Attraction.active.is_(True),
                Attraction.employee_circle.is_(True),
            )
            .first()
        )
        if not attraction:
            raise HTTPException(400, "请选择有效的员工景点圈")
    if category == "gsm_leader":
        return gsm_recognizer_ranking_payload(db, start_value, end_value, subtype_id, keyword, page, page_size, attraction)
    keyword_value = keyword.strip().lower()
    end_iso = end.isoformat()
    role_codes = LEADER_CODES if category == "leader" else FRONTLINE_CODES
    employees, role_map = ranking_employees(db, attraction_id, end_iso, role_codes)
    if keyword_value:
        employees = [
            employee for employee in employees
            if keyword_value in employee.name.lower() or keyword_value in employee.employee_no.lower()
        ]
    employee_ids = [employee.id for employee in employees]
    leader_names = ranking_leader_names(db, employee_ids, end_iso) if category != "leader" else {}
    aggregates: dict[int, dict] = {employee_id: {} for employee_id in employee_ids}
    subtype_name = ""

    if category == "recognition":
        subtype = db.get(RecognitionType, subtype_id) if subtype_id else None
        if subtype_id and (not subtype or not subtype.active):
            raise HTTPException(400, "请选择有效的认可类型")
        subtype_name = subtype.name if subtype else "全部加分类别"
        rows = db.query(RecognitionRecord.employee_id, func.count(RecognitionRecord.id), text("SUM(CASE WHEN recognition_date < '2026-09-01' THEN fraction ELSE credited_fraction END)"), func.max(RecognitionRecord.recognition_date)).filter(
            RecognitionRecord.employee_id.in_(employee_ids) if employee_ids else RecognitionRecord.employee_id == -1,
            RecognitionRecord.status == "confirmed", RecognitionRecord.recognition_date >= start_value, RecognitionRecord.recognition_date <= end_value,
        )
        if subtype:
            rows = rows.filter(RecognitionRecord.recognition_type_id == subtype.id)
        rows = rows.group_by(RecognitionRecord.employee_id).all()
        for employee_id, count, score, recent_date in rows:
            aggregates[employee_id] = {"count": int(count or 0), "score": float(score or 0), "recent_date": recent_date or ""}
    elif category == "deduction":
        subtype = db.get(DeductionType, subtype_id) if subtype_id else None
        if subtype_id and (not subtype or not subtype.active):
            raise HTTPException(400, "请选择有效的扣分类型")
        subtype_name = subtype.name if subtype else "全部扣分类型"
        rows = db.query(DeductionRecord.employee_id, func.count(DeductionRecord.id), func.sum(DeductionRecord.points), func.max(DeductionRecord.occurred_on)).filter(
            DeductionRecord.employee_id.in_(employee_ids) if employee_ids else DeductionRecord.employee_id == -1,
            DeductionRecord.status == "active", DeductionRecord.occurred_on >= start_value, DeductionRecord.occurred_on <= end_value,
        )
        if subtype:
            rows = rows.filter(DeductionRecord.deduction_type_id == subtype.id)
        rows = rows.group_by(DeductionRecord.employee_id).all()
        for employee_id, count, score, recent_date in rows:
            aggregates[employee_id] = {"count": int(count or 0), "score": float(score or 0), "recent_date": recent_date or ""}
    elif category == "absence":
        if subtype_id not in {None, 1, 2}:
            raise HTTPException(400, "缺勤类型无效")
        subtype_name = "违规病假" if subtype_id == 2 else ("病假" if subtype_id == 1 else "全部缺勤类型")
        sick_rows = db.query(SickLeaveRecord).filter(
            SickLeaveRecord.employee_id.in_(employee_ids) if employee_ids else SickLeaveRecord.employee_id == -1,
            SickLeaveRecord.status == "active", SickLeaveRecord.leave_end_date >= start_value, SickLeaveRecord.leave_start_date <= end_value,
        )
        if subtype_id == 1:
            sick_rows = sick_rows.filter(SickLeaveRecord.is_violation.is_(False))
        elif subtype_id == 2:
            sick_rows = sick_rows.filter(SickLeaveRecord.is_violation.is_(True))
        sick_rows = sick_rows.all()
        selected_months: dict[int, set[str]] = {employee_id: set() for employee_id in employee_ids}
        for row in sick_rows:
            data = aggregates.setdefault(row.employee_id, {})
            data["count"] = int(data.get("count", 0)) + 1
            data["leave_days"] = float(data.get("leave_days", 0)) + float(row.leave_days)
            data["charged_days"] = float(data.get("charged_days", 0)) + float(row.charged_days)
            data["recent_date"] = max(str(data.get("recent_date") or ""), row.leave_start_date)
            selected_months.setdefault(row.employee_id, set()).add(row.attendance_month)
        months = months_between(start, end)
        for month in months:
            ensure_month_attendance(db, month, employee_ids)
        if months:
            db.commit()
        attendance_rows = (
            db.query(AttendanceMonthlyScore)
            .filter(
                AttendanceMonthlyScore.employee_id.in_(employee_ids) if employee_ids else AttendanceMonthlyScore.employee_id == -1,
                AttendanceMonthlyScore.attendance_month.in_(months),
            )
            .all()
        )
        for row in attendance_rows:
            if row.attendance_month in selected_months.get(row.employee_id, set()):
                aggregates[row.employee_id]["score"] = float(aggregates[row.employee_id].get("score", 0)) + float(row.sick_deduction or 0)
    elif category == "leader":
        subtype = db.get(RecognitionType, subtype_id) if subtype_id else None
        if subtype_id and (not subtype or not subtype.active):
            raise HTTPException(400, "请选择有效认可类型")
        subtype_name = subtype.name if subtype else "全部加分类别"
        query = db.query(RecognitionRecord).filter(
            *([RecognitionRecord.home_attraction_id == attraction_id] if attraction_id is not None else []),
            RecognitionRecord.status == "confirmed",
            RecognitionRecord.recognition_date >= start_value,
            RecognitionRecord.recognition_date <= end_value,
        )
        if subtype:
            query = query.filter(RecognitionRecord.recognition_type_id == subtype.id)
        for row in query.all():
            # One confirmed record is one independent scoring event. Attribute it to
            # participating current supervisors, while deduplicating only within it.
            for employee_id in {row.recognizer_employee_id, row.operator_employee_id}:
                if employee_id not in aggregates:
                    continue
                data = aggregates[employee_id]
                data["count"] = int(data.get("count", 0)) + 1
                data["score"] = float(data.get("score", 0)) + float(effective_recognition_credit(row))
                data["recent_date"] = max(str(data.get("recent_date") or ""), row.recognition_date)
    else:
        months = months_between(start, end)
        for month in months:
            ensure_month_attendance(db, month, employee_ids)
        if months:
            db.commit()
        recognition_totals = dict(
            db.query(RecognitionRecord.employee_id, text("SUM(CASE WHEN recognition_date < '2026-09-01' THEN fraction ELSE credited_fraction END)"))
            .filter(
                RecognitionRecord.employee_id.in_(employee_ids) if employee_ids else RecognitionRecord.employee_id == -1,
                RecognitionRecord.status == "confirmed",
                RecognitionRecord.recognition_date >= start_value,
                RecognitionRecord.recognition_date <= end_value,
            )
            .group_by(RecognitionRecord.employee_id)
            .all()
        )
        deduction_totals = dict(
            db.query(DeductionRecord.employee_id, func.sum(DeductionRecord.points))
            .filter(
                DeductionRecord.employee_id.in_(employee_ids) if employee_ids else DeductionRecord.employee_id == -1,
                DeductionRecord.status == "active",
                DeductionRecord.occurred_on >= start_value,
                DeductionRecord.occurred_on <= end_value,
            )
            .group_by(DeductionRecord.employee_id)
            .all()
        )
        attendance_totals = dict(
            db.query(AttendanceMonthlyScore.employee_id, func.sum(AttendanceMonthlyScore.final_score))
            .filter(
                AttendanceMonthlyScore.employee_id.in_(employee_ids) if employee_ids else AttendanceMonthlyScore.employee_id == -1,
                AttendanceMonthlyScore.attendance_month.in_(months),
                AttendanceMonthlyScore.eligible.is_(True),
            )
            .group_by(AttendanceMonthlyScore.employee_id)
            .all()
        )
        for employee_id in employee_ids:
            recognition_score = float(recognition_totals.get(employee_id) or 0)
            deduction_score = float(deduction_totals.get(employee_id) or 0)
            attendance_score = float(attendance_totals.get(employee_id) or 0)
            aggregates[employee_id] = {
                "recognition_score": recognition_score,
                "deduction_score": deduction_score,
                "attendance_score": attendance_score,
                "total_score": round(recognition_score + attendance_score - deduction_score, 2),
            }

    result_rows = []
    for employee in employees:
        data = aggregates.get(employee.id, {})
        role = role_map.get(employee.id)
        result_rows.append(
            {
                "employee_id": employee.id,
                "employee_no": employee.employee_no,
                "employee_name": employee.name,
                "role_name": role.name if role else "未配置",
                "leader_name": leader_names.get(employee.id, "") if category != "leader" else "",
                "count": int(data.get("count", 0)),
                "score": round(float(data.get("score", 0)), 2),
                "recent_date": str(data.get("recent_date") or ""),
                "leave_days": round(float(data.get("leave_days", 0)), 1),
                "charged_days": int(data.get("charged_days", 0)),
                "recognition_score": round(float(data.get("recognition_score", 0)), 2),
                "deduction_score": round(float(data.get("deduction_score", 0)), 2),
                "attendance_score": round(float(data.get("attendance_score", 0)), 2),
                "total_score": round(float(data.get("total_score", 0)), 2),
            }
        )

    if category == "overall":
        result_rows.sort(key=lambda row: (-row["total_score"], -row["recognition_score"], row["employee_no"]))
    elif category == "leader":
        result_rows.sort(key=lambda row: (-row["count"], -row["score"], row["employee_no"]))
    elif sort_by == "count":
        result_rows.sort(key=lambda row: (-row["count"], -row["score"], row["employee_no"]))
    else:
        result_rows.sort(key=lambda row: (-row["score"], -row["count"], row["employee_no"]))
    for index, row in enumerate(result_rows, start=1):
        row["rank"] = index
    total = len(result_rows)
    offset = (page - 1) * page_size
    return {
        "category": category,
        "subtype_id": subtype_id,
        "subtype_name": subtype_name,
        "sort_by": sort_by,
        "start_date": start_value,
        "end_date": end_value,
        "attraction_id": attraction_id,
        "attraction_name": attraction.name if attraction else "全部景点圈",
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": max(1, ceil(total / page_size)),
        "rows": result_rows[offset:offset + page_size],
    }


@router.get("/pr-rankings")
def pr_rankings(
    start_date: str,
    end_date: str,
    category: str = "overall",
    subtype_id: int | None = None,
    sort_by: str = "score",
    keyword: str = "",
    attraction_id: int | None = None,
    page: int = 1,
    page_size: int = 20,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    page = max(1, page)
    page_size = min(max(10, page_size), 100)
    cache_key = (user.id, start_date, end_date, category, subtype_id, sort_by, keyword.strip(), page, page_size, attraction_id)
    with _pr_ranking_cache_lock:
        now = monotonic()
        cached = _pr_ranking_response_cache.get(cache_key)
        if cached and cached[0] > now:
            content = cached[1]
        else:
            payload = pr_ranking_payload(db, user, start_date, end_date, category, subtype_id, sort_by, keyword, page, page_size, attraction_id)
            content = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
            _pr_ranking_response_cache[cache_key] = (now + PR_RANKING_CACHE_SECONDS, content)
            if len(_pr_ranking_response_cache) > 256:
                _pr_ranking_response_cache.clear()
                _pr_ranking_response_cache[cache_key] = (now + PR_RANKING_CACHE_SECONDS, content)
    return Response(content=content, media_type="application/json")


@router.get("/pr-rankings/export")
def export_pr_rankings(
    start_date: str,
    end_date: str,
    category: str = "overall",
    subtype_id: int | None = None,
    sort_by: str = "score",
    keyword: str = "",
    attraction_id: int | None = None,
    db: Session = Depends(get_db),
    user: V2User = Depends(current_user),
):
    if user.role.code != "GSM":
        raise HTTPException(403, "TAGSM仅支持查询PR排名，不能导出景点圈数据")
    data = pr_ranking_payload(db, user, start_date, end_date, category, subtype_id, sort_by, keyword, 1, 5000, attraction_id)
    wb = build_pr_rankings_workbook(db, data, category)
    watermark_workbook(wb, user.employee.employee_no)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    write_audit(db, user.employee, "导出PR排名", "pr_ranking_export", category, after={"start_date": start_date, "end_date": end_date, "subtype_id": subtype_id, "sort_by": sort_by, "keyword": keyword, "row_count": len(data["rows"])})
    db.commit()
    ascii_filename = f"pr_rankings_{category}_{start_date}_{end_date}.xlsx"
    display_filename = f"PR排名_{category}_{start_date}_{end_date}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": content_disposition(ascii_filename, display_filename)})


def statistics_hierarchy(db: Session, score_rows: list[dict], month_end: str, user: V2User) -> list[dict]:
    employee_ids = [int(score["employee_id"]) for score in score_rows]
    if not employee_ids:
        return []
    employees = {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(employee_ids)).all()}
    employee_roles = roles_at(db, employee_ids, month_end)

    memberships = (
        db.query(GroupMembership)
        .filter(
            GroupMembership.employee_id.in_(employee_ids),
            GroupMembership.status == "active",
            GroupMembership.starts_on <= month_end,
            or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= month_end),
        )
        .order_by(GroupMembership.employee_id, GroupMembership.starts_on.desc(), GroupMembership.id.desc())
        .all()
    )
    membership_by_employee: dict[int, GroupMembership] = {}
    for membership in memberships:
        membership_by_employee.setdefault(membership.employee_id, membership)
    group_ids = {membership.group_id for membership in membership_by_employee.values()}
    groups = {group.id: group for group in db.query(WorkGroup).filter(WorkGroup.id.in_(group_ids)).all()} if group_ids else {}
    leader_assignments = (
        db.query(GroupLeaderAssignment)
        .filter(
            GroupLeaderAssignment.group_id.in_(group_ids),
            GroupLeaderAssignment.status == "active",
            GroupLeaderAssignment.starts_on <= month_end,
            or_(GroupLeaderAssignment.ends_on.is_(None), GroupLeaderAssignment.ends_on >= month_end),
        )
        .order_by(GroupLeaderAssignment.group_id, GroupLeaderAssignment.starts_on.desc(), GroupLeaderAssignment.id.desc())
        .all()
        if group_ids
        else []
    )
    leader_assignment_by_group: dict[int, GroupLeaderAssignment] = {}
    for assignment in leader_assignments:
        leader_assignment_by_group.setdefault(assignment.group_id, assignment)
    leader_ids = {assignment.leader_employee_id for assignment in leader_assignment_by_group.values()}
    leaders = {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()} if leader_ids else {}
    leader_roles = roles_at(db, leader_ids, month_end)

    attraction_ids = {score.get("attraction_id") for score in score_rows if score.get("attraction_id") is not None}
    attraction_rows = db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).all() if attraction_ids else []
    attraction_names = {attraction.id: attraction.name for attraction in attraction_rows}
    managers_by_attraction: dict[int | None, dict[str, list[tuple[Employee, Role]]]] = {}
    if attraction_ids:
        scopes = (
            db.query(ManagementScope)
            .filter(
                ManagementScope.attraction_id.in_(attraction_ids),
                ManagementScope.starts_on <= month_end,
                or_(ManagementScope.ends_on.is_(None), ManagementScope.ends_on >= month_end),
            )
            .all()
        )
        manager_ids = {scope.employee_id for scope in scopes}
        manager_employees = {
            employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(manager_ids)).all()
        } if manager_ids else {}
        manager_roles = roles_at(db, manager_ids, month_end)
        candidates_by_attraction: dict[int, list[tuple[Employee, Role]]] = {}
        for scope in scopes:
            employee = manager_employees.get(scope.employee_id)
            role = manager_roles.get(scope.employee_id)
            if employee and role and role.code in GSM_CODES:
                candidates_by_attraction.setdefault(scope.attraction_id, []).append((employee, role))
        for attraction_id in attraction_ids:
            candidates = candidates_by_attraction.get(attraction_id, [])
            ordered = sorted(candidates, key=lambda item: (item[1].code != "GSM", item[0].employee_no))
            managers_by_attraction[attraction_id] = {
                "gsms": [item for item in ordered if item[1].code == "GSM"],
                "ta_gsms": [item for item in ordered if item[1].code == "TA_GSM"],
            }

    attractions: dict[int | None, dict] = {}
    for score in score_rows:
        employee = employees.get(score["employee_id"])
        if not employee:
            continue
        role = employee_roles.get(employee.id)
        # The team branch represents only scored frontline employees. GSM/TA GSM
        # are rendered once as circle-level management nodes, never as an
        # unassigned employee beneath a supervisor placeholder.
        if not role or role.code not in FRONTLINE_CODES:
            continue
        attraction_key = score.get("attraction_id")
        attraction_node = attractions.setdefault(
            attraction_key,
            {"name": attraction_names.get(attraction_key, "未设置景点圈"), "managers": {}},
        )
        manager_info = managers_by_attraction.get(attraction_key, {"gsms": [], "ta_gsms": []})
        gsms = manager_info["gsms"]
        ta_gsms = manager_info["ta_gsms"]
        manager_key = f"team-{attraction_key}" if len(gsms) > 1 else (gsms[0][0].id if gsms else None)
        manager_node = attraction_node["managers"].setdefault(
            manager_key,
            {
                "gsms": gsms,
                "ta_gsms": ta_gsms,
                "name": (f"主管组（由GSM共同承接：{'、'.join(employee.name for employee, _role in gsms)}）" if len(gsms) > 1 else (gsms[0][0].name if gsms else "未配置GSM")),
                "role_name": "" if len(gsms) != 1 else gsms[0][1].name,
                "leaders": {},
            },
        )
        membership = membership_by_employee.get(employee.id)
        group = groups.get(membership.group_id) if membership else None
        leader_assignment = leader_assignment_by_group.get(group.id) if group else None
        leader = leaders.get(leader_assignment.leader_employee_id) if leader_assignment else None
        leader_role = leader_roles.get(leader.id) if leader else None
        leader_key = f"employee-{leader.id}" if leader else (f"group-{group.id}" if group else "ungrouped")
        leader_node = manager_node["leaders"].setdefault(
            leader_key,
            {
                "name": leader.name if leader else "未配置主管",
                "role_name": leader_role.name if leader_role else "",
                "employees": [],
            },
        )
        leader_node["employees"].append(
            {
                "employee_id": employee.id,
                "employee_no": employee.employee_no,
                "employee_name": employee.name,
                "role_code": role.code if role else "",
                "role_name": role.name if role else "",
                "recognition_score": float(score["recognition_score"] or 0),
                "deduction_score": float(score["deduction_score"] or 0),
                "attendance_score": float(score["attendance_score"] or 0),
                "total_score": float(score["total_score"] or 0),
            }
        )

    result = []
    for attraction_key, attraction in sorted(attractions.items(), key=lambda item: item[1]["name"]):
        attraction_id = f"attraction-{attraction_key or 'none'}"
        result.append({"node_id": attraction_id, "parent_id": "", "level": 0, "node_type": "attraction", "name": attraction["name"]})
        for manager_key, manager in sorted(attraction["managers"].items(), key=lambda item: item[1]["name"]):
            manager_id = f"{attraction_id}-gsm-{manager_key or 'none'}"
            gsms = manager["gsms"]
            ta_gsms = manager["ta_gsms"]
            for gsm, gsm_role in gsms:
                result.append({"node_id": f"{manager_id}-employee-{gsm.id}", "parent_id": attraction_id, "level": 1, "node_type": "gsm", "name": gsm.name, "role_name": gsm_role.name})
            # TA GSM is a peer of GSM, not the parent of the supervisor branch.
            # Emit it before the common branch so table order also communicates that relationship.
            for ta_gsm, ta_gsm_role in ta_gsms:
                result.append({"node_id": f"{attraction_id}-ta-gsm-{ta_gsm.id}", "parent_id": attraction_id, "level": 1, "node_type": "gsm", "name": ta_gsm.name, "role_name": ta_gsm_role.name})
            branch_employees = [employee for leader in manager["leaders"].values() for employee in leader["employees"]]
            if len(gsms) > 1:
                result.append({
                    "node_id": manager_id,
                    "parent_id": attraction_id,
                    "level": 1,
                    "node_type": "gsm_team",
                    "name": manager["name"],
                    "role_name": "",
                    "supervisor_count": len(manager["leaders"]),
                    "member_count": len(branch_employees),
                    "cm_count": sum(employee["role_code"] == "CM" for employee in branch_employees),
                    "tr_count": sum(employee["role_code"] == "TR" for employee in branch_employees),
                })
            elif len(gsms) == 1:
                manager_id = f"{manager_id}-employee-{gsms[0][0].id}"
            else:
                result.append({"node_id": manager_id, "parent_id": attraction_id, "level": 1, "node_type": "gsm", "name": manager["name"], "role_name": manager["role_name"]})
            for leader_key, leader in sorted(manager["leaders"].items(), key=lambda item: item[1]["name"]):
                leader_id = f"{manager_id}-leader-{leader_key}"
                sorted_employees = sorted(leader["employees"], key=lambda item: (item["employee_name"], item["employee_no"]))
                leader_recognition_score = round(sum(employee["recognition_score"] for employee in leader["employees"]), 2)
                leader_deduction_score = round(sum(employee["deduction_score"] for employee in leader["employees"]), 2)
                leader_attendance_score = round(sum(employee["attendance_score"] for employee in leader["employees"]), 2)
                leader_total_score = round(sum(employee["total_score"] for employee in leader["employees"]), 2)
                leader_average_score = round(leader_total_score / len(leader["employees"]), 2) if leader["employees"] else 0
                result.append(
                    {
                        "node_id": leader_id,
                        "parent_id": manager_id,
                        "level": 2,
                        "node_type": "supervisor",
                        "name": leader["name"],
                        "role_name": leader["role_name"],
                        "member_count": len(leader["employees"]),
                        "employee_ids": [employee["employee_id"] for employee in sorted_employees],
                        "recognition_score": leader_recognition_score,
                        "deduction_score": leader_deduction_score,
                        "attendance_score": leader_attendance_score,
                        "total_score": leader_total_score,
                        "average_score": leader_average_score,
                    }
                )
                for employee in sorted_employees:
                    result.append(
                        {
                            "node_id": f"employee-{employee['employee_id']}",
                            "parent_id": leader_id,
                            "level": 3,
                            "node_type": "employee",
                            **employee,
                        }
                    )
    return result


def statistics_payload(
    db: Session,
    month: str,
    attraction_id: int | None = None,
    keyword: str | None = None,
    title: str | None = None,
    user: V2User | None = None,
    *,
    include_records: bool = False,
) -> dict:
    try:
        month_start = date.fromisoformat(f"{month}-01")
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "月份格式应为YYYY-MM") from exc
    month_end = month_start.replace(day=monthrange(month_start.year, month_start.month)[1]).isoformat()
    selected_title = (title or "").strip().upper()
    if selected_title and selected_title not in FRONTLINE_CODES:
        raise HTTPException(400, "Title只能选择CM或TR")
    allowed_attractions = scoped_hr_attraction_ids(db, user) if user else None
    if allowed_attractions is not None:
        if attraction_id is None:
            attraction_id = next(iter(allowed_attractions))
        elif attraction_id not in allowed_attractions:
            raise HTTPException(403, "景点圈HR只能查看所属景点圈数据")
    # Attendance materialization is retained for score-view compatibility, but the
    # savepoint is always rolled back so statistics and export remain read-only.
    statistics_savepoint = db.begin_nested()
    try:
        ensure_month_attendance(db, month)
        loa_employee_ids = {
            row[0]
            for row in db.query(EmployeeLOAPeriod.employee_id)
            .filter(
                EmployeeLOAPeriod.status != "cancelled",
                EmployeeLOAPeriod.starts_on <= month_end,
                or_(EmployeeLOAPeriod.ends_on.is_(None), EmployeeLOAPeriod.ends_on >= month_start.isoformat()),
            )
            .all()
        }
        if loa_employee_ids:
            for employee in db.query(Employee).filter(Employee.id.in_(loa_employee_ids)).all():
                recalculate_attendance(db, employee, month)
        db.flush()
        query = (
            "SELECT s.*, e.account_deleted_at AS account_deleted_at, COALESCE(os.attraction_id, e.attraction_id) AS attraction_id, "
            "COALESCE(os.attraction_name, a.name) AS attraction_name, "
            "CASE WHEN os.id IS NULL THEN 'current' ELSE 'month_close_snapshot' END AS organization_basis "
            "FROM v_employee_month_scores s "
            "JOIN employees e ON e.id=s.employee_id "
            "LEFT JOIN employee_month_organization_snapshots os ON os.employee_id=s.employee_id AND os.score_month=s.score_month "
            "LEFT JOIN attractions a ON a.id=e.attraction_id WHERE s.score_month=:month"
        )
        params = {"month": month}
        if attraction_id:
            query += " AND COALESCE(os.attraction_id, e.attraction_id)=:attraction_id"
            params["attraction_id"] = attraction_id
        if keyword:
            query += " AND (e.employee_no LIKE :keyword OR e.name LIKE :keyword)"
            params["keyword"] = f"%{keyword.strip()}%"
        query += " ORDER BY s.total_score DESC, e.employee_no"
        score_rows = [dict(row) for row in db.execute(text(query), params).mappings().all()]
    finally:
        if statistics_savepoint.is_active:
            statistics_savepoint.rollback()
    if selected_title and score_rows:
        month_end_roles = roles_at(db, [int(row["employee_id"]) for row in score_rows], month_end)
        score_rows = [
            row
            for row in score_rows
            if month_end_roles.get(int(row["employee_id"]))
            and month_end_roles[int(row["employee_id"])].code == selected_title
        ]
    loa_rows = []
    filtered_score_rows = []
    for row in score_rows:
        if full_month_loa(db, int(row["employee_id"]), month):
            loa_rows.append({**row, "employment_status": "LOA（长期病假）"})
        else:
            filtered_score_rows.append(row)
    score_rows = filtered_score_rows
    visible_employee_ids = {
        int(row["employee_id"])
        for row in [*score_rows, *loa_rows]
    }
    data_updated_at = ""
    if visible_employee_ids:
        update_values = [
            db.query(func.max(RecognitionRecord.submitted_at)).filter(RecognitionRecord.recognition_month == month, RecognitionRecord.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(RecognitionRecord.voided_at)).filter(RecognitionRecord.recognition_month == month, RecognitionRecord.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(DeductionRecord.submitted_at)).filter(DeductionRecord.deduction_month == month, DeductionRecord.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(DeductionRecord.voided_at)).filter(DeductionRecord.deduction_month == month, DeductionRecord.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(SickLeaveRecord.submitted_at)).filter(SickLeaveRecord.attendance_month == month, SickLeaveRecord.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(SickLeaveRecord.voided_at)).filter(SickLeaveRecord.attendance_month == month, SickLeaveRecord.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(AttendanceMonthlyScore.calculated_at)).filter(AttendanceMonthlyScore.attendance_month == month, AttendanceMonthlyScore.employee_id.in_(visible_employee_ids)).scalar(),
            db.query(func.max(Employee.updated_at)).filter(Employee.id.in_(visible_employee_ids)).scalar(),
        ]
        latest_update = max((value for value in update_values if value), default=None)
        if latest_update:
            data_updated_at = latest_update.strftime("%Y-%m-%d %H:%M")
    recognition_payloads: list[dict] = []
    deduction_payloads: list[dict] = []
    sick_leave_rows_payload: list[dict] = []
    if include_records:
        recognition_rows = db.query(RecognitionRecord).filter(RecognitionRecord.recognition_month == month).order_by(RecognitionRecord.submitted_at.desc()).all()
        deduction_rows = db.query(DeductionRecord).filter(DeductionRecord.deduction_month == month).order_by(DeductionRecord.submitted_at.desc()).all()
        sick_leave_rows = db.query(SickLeaveRecord).filter(SickLeaveRecord.attendance_month == month).order_by(SickLeaveRecord.submitted_at.desc()).all()
        all_record_rows = [*recognition_rows, *deduction_rows, *sick_leave_rows]
        record_employee_ids = {int(row.employee_id) for row in all_record_rows}
        employees = {
            employee.id: employee
            for employee in db.query(Employee).filter(Employee.id.in_(record_employee_ids)).all()
        } if record_employee_ids else {}
        historical_roles = roles_at(db, record_employee_ids, month_end) if record_employee_ids else {}
        keyword_value = (keyword or "").strip().casefold()

        def record_is_visible(row) -> bool:
            employee = employees.get(int(row.employee_id))
            if isinstance(row, RecognitionRecord):
                row_attraction_id = row.home_attraction_id
                employee_no = row.employee_no
                employee_name = row.employee_name
            elif isinstance(row, DeductionRecord):
                row_attraction_id = row.attraction_id_snapshot
                employee_no = row.employee_no
                employee_name = row.employee_name
            else:
                row_attraction_id = row.attraction_id_snapshot
                employee_no = row.employee_no_snapshot or (employee.employee_no if employee else "")
                employee_name = row.employee_name_snapshot or (employee.name if employee else "")
            row_attraction_id = row_attraction_id or (employee.attraction_id if employee else None)
            if attraction_id and row_attraction_id != attraction_id:
                return False
            if keyword_value and keyword_value not in f"{employee_no} {employee_name}".casefold():
                return False
            if selected_title:
                snapshot_title = str(getattr(row, "employee_role_snapshot", "") or "").strip().upper()
                historical_role = historical_roles.get(int(row.employee_id))
                row_title = snapshot_title if snapshot_title in FRONTLINE_CODES else (historical_role.code if historical_role else "")
                if row_title != selected_title:
                    return False
            return True

        recognition_payloads = [recognition_payload(row) for row in recognition_rows if record_is_visible(row)]
        deduction_payloads = [deduction_payload(row) for row in deduction_rows if record_is_visible(row)]
        sick_leave_rows_payload = sick_leave_payloads(db, [row for row in sick_leave_rows if record_is_visible(row)])
    payload = {
        "month": month,
        "title": selected_title,
        "filters": {
            "attraction_id": attraction_id or "",
            "title": selected_title,
            "keyword": (keyword or "").strip(),
        },
        "data_updated_at": data_updated_at,
        "organization_basis": "月结封存归属" if any(row.get("organization_basis") == "month_close_snapshot" for row in score_rows) else "当前组织归属（该月尚未月结）",
        "summary": {
            "employee_count": len(score_rows),
            "recognition_score": round(sum(float(row["recognition_score"] or 0) for row in score_rows), 2),
            "attendance_score": round(sum(float(row["attendance_score"] or 0) for row in score_rows), 2),
            "deduction_score": round(sum(float(row["deduction_score"] or 0) for row in score_rows), 2),
            "total_score": round(sum(float(row["total_score"] or 0) for row in score_rows), 2),
        },
        "hierarchy": statistics_hierarchy(db, score_rows, month_end, user) if user else [],
        "loa_rows": loa_rows,
    }
    if include_records:
        payload.update(
            {
                "scores": score_rows,
                "recognitions": recognition_payloads,
                "deductions": deduction_payloads,
                "sick_leaves": sick_leave_rows_payload,
            }
        )
    return payload


def statistics_details_payload(db: Session, month: str, employee_ids: list[int]) -> dict[str, dict]:
    if not employee_ids:
        return {}
    placeholders = []
    params: dict[str, object] = {"month": month}
    for index, employee_id in enumerate(employee_ids):
        key = f"employee_{index}"
        placeholders.append(f":{key}")
        params[key] = employee_id
    score_rows = {
        row["employee_id"]: dict(row)
        for row in db.execute(
            text(
                "SELECT employee_id, recognition_score, deduction_score, attendance_score, total_score "
                f"FROM v_employee_month_scores WHERE score_month=:month AND employee_id IN ({','.join(placeholders)})"
            ),
            params,
        ).mappings().all()
    }
    employees = {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(employee_ids)).all()}
    recognition_groups: dict[int, list[RecognitionRecord]] = {employee_id: [] for employee_id in employee_ids}
    deduction_groups: dict[int, list[DeductionRecord]] = {employee_id: [] for employee_id in employee_ids}
    sick_leave_groups: dict[int, list[SickLeaveRecord]] = {employee_id: [] for employee_id in employee_ids}
    for row in (
        db.query(RecognitionRecord)
        .filter(RecognitionRecord.employee_id.in_(employee_ids), RecognitionRecord.recognition_month == month)
        .order_by(RecognitionRecord.submitted_at.desc(), RecognitionRecord.id.desc())
        .all()
    ):
        recognition_groups[row.employee_id].append(row)
    for row in (
        db.query(DeductionRecord)
        .filter(DeductionRecord.employee_id.in_(employee_ids), DeductionRecord.deduction_month == month)
        .order_by(DeductionRecord.submitted_at.desc(), DeductionRecord.id.desc())
        .all()
    ):
        deduction_groups[row.employee_id].append(row)
    for row in (
        db.query(SickLeaveRecord)
        .filter(SickLeaveRecord.employee_id.in_(employee_ids), SickLeaveRecord.attendance_month == month)
        .order_by(SickLeaveRecord.submitted_at.desc(), SickLeaveRecord.id.desc())
        .all()
    ):
        sick_leave_groups[row.employee_id].append(row)
    attendance_rows = {
        row.employee_id: row
        for row in db.query(AttendanceMonthlyScore)
        .filter(AttendanceMonthlyScore.employee_id.in_(employee_ids), AttendanceMonthlyScore.attendance_month == month)
        .all()
    }
    empty_score = {"recognition_score": 0, "deduction_score": 0, "attendance_score": 0, "total_score": 0}
    result: dict[str, dict] = {}
    for employee_id in employee_ids:
        if employee_id not in employees:
            continue
        detail = member_score_detail_payload(
            db,
            employees[employee_id],
            score_rows.get(employee_id, empty_score),
            recognition_groups[employee_id],
            deduction_groups[employee_id],
            sick_leave_groups[employee_id],
            attendance_rows.get(employee_id),
        )
        # Keep the established detail lists while returning the identity fields
        # used by the dedicated statistics-detail page title.
        result[str(employee_id)] = {**detail["details"], **{key: detail[key] for key in ("employee_id", "employee_no", "employee_name", "role_name")}}
    return result


@router.get("/statistics")
def statistics(month: str, attraction_id: int | None = None, keyword: str | None = None, title: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("DATA_VIEW", "DATA_EXPORT"))):
    cache_key = (user.id, user.role.code, month, attraction_id, (keyword or "").strip(), (title or "").strip().upper())
    with _statistics_cache_lock:
        now = monotonic()
        cached = _statistics_response_cache.get(cache_key)
        if cached and cached[0] > now:
            content = cached[1]
        else:
            payload = statistics_payload(db, month, attraction_id, keyword, title, user)
            content = json.dumps(
                payload,
                ensure_ascii=False,
                separators=(",", ":"),
                default=lambda value: float(value) if isinstance(value, Decimal) else str(value),
            ).encode("utf-8")
            _statistics_response_cache[cache_key] = (now + STATISTICS_CACHE_SECONDS, content)
            if len(_statistics_response_cache) > 128:
                _statistics_response_cache.clear()
                _statistics_response_cache[cache_key] = (now + STATISTICS_CACHE_SECONDS, content)
    return Response(content=content, media_type="application/json")


@router.get("/statistics/details")
def statistics_details(
    month: str,
    employee_ids: str,
    effective_only: bool = False,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("DATA_VIEW", "DATA_EXPORT")),
):
    try:
        date.fromisoformat(f"{month}-01")
        selected_ids = list(dict.fromkeys(int(value) for value in employee_ids.split(",") if value.strip()))
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "月份或员工参数格式不正确") from exc
    if not selected_ids or len(selected_ids) > 200:
        raise HTTPException(400, "每次可查询1至200名员工明细")
    employees = db.query(Employee).filter(Employee.id.in_(selected_ids)).all()
    if len(employees) != len(selected_ids):
        raise HTTPException(404, "部分员工不存在")
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is not None and any(employee.attraction_id not in allowed_attractions for employee in employees):
        raise HTTPException(403, "景点圈HR只能查看所属景点圈数据")
    details_savepoint = db.begin_nested()
    try:
        ensure_month_attendance(db, month, selected_ids)
        db.flush()
        details = statistics_details_payload(db, month, selected_ids)
    finally:
        if details_savepoint.is_active:
            details_savepoint.rollback()
    if effective_only:
        for detail in details.values():
            detail["all_records"] = [record for record in detail.get("all_records", []) if record.get("included")]
    return {"month": month, "details": details}


@router.get("/statistics/my-exports")
def my_statistics_exports(
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("DATA_EXPORT")),
):
    rows = (
        db.query(AuditLog)
        .filter(AuditLog.operator_id == user.id, AuditLog.action == "导出统计数据")
        .order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        .limit(8)
        .all()
    )
    attraction_ids: set[int] = set()
    parsed_rows: list[tuple[AuditLog, dict]] = []
    for row in rows:
        try:
            details = json.loads(row.after_json or "{}")
        except (TypeError, ValueError):
            details = {}
        attraction_id = details.get("attraction_id")
        if isinstance(attraction_id, int):
            attraction_ids.add(attraction_id)
        parsed_rows.append((row, details))
    attraction_names = {
        attraction.id: attraction.name
        for attraction in db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).all()
    } if attraction_ids else {}
    return {
        "items": [
            {
                "id": row.id,
                "exported_at": row.created_at.strftime("%Y-%m-%d %H:%M"),
                "month": row.entity_id or "",
                "attraction_name": attraction_names.get(details.get("attraction_id"), "全部景点圈"),
                "title": details.get("title") or "CM/TR全部",
                "keyword": details.get("keyword") or "",
                "employee_count": int(details.get("employee_count") or 0),
            }
            for row, details in parsed_rows
        ]
    }


@router.get("/statistics/export")
def export_statistics(month: str, attraction_id: int | None = None, keyword: str | None = None, title: str | None = None, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("DATA_EXPORT"))):
    data = statistics_payload(db, month, attraction_id, keyword, title, user, include_records=True)
    wb, attraction_label, title_label = build_statistics_workbook(
        db,
        data,
        month=month,
        attraction_id=attraction_id,
        keyword=keyword,
        exporter_no=user.employee.employee_no,
        exporter_name=user.name,
        exporter_role_name=user.role.name,
        exporter_role_code=user.role.code,
    )
    watermark_workbook(wb, user.employee.employee_no)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    write_audit(db, user.employee, "导出统计数据", "statistics_export", month, after={"attraction_id": attraction_id, "title": data["title"], "keyword": keyword or "", "employee_count": len(data["scores"])})
    db.commit()
    ascii_filename = f"recognition_v2_{month.replace('-', '_')}_circle-{attraction_id or 'all'}_{data['title'] or 'all'}.xlsx"
    display_filename = f"认可数据_{month}_{attraction_label}_{title_label}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": content_disposition(ascii_filename, display_filename)})


@router.get("/hr/employees")
def hr_employees(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee)
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.employee_no).all()
    payloads = employee_payloads(db, employees)
    return [payloads[employee.id] for employee in employees]


@router.get("/hr/account-status")
def hr_account_status(request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    """Return approved account-status metadata without exposing passwords.

    This is deliberately separate from employee-management edit scope: a
    scoped HR can read CM/TR/TALEAD/Lead status across all three circles, but
    their existing edit endpoints remain limited to their own circle.
    """
    if user.role.code == SCOPED_HR_ROLE_CODE:
        permitted_roles = CIRCLE_HR_MANAGED_ROLE_CODES
        scope_label = "全部景点圈的 CM/TR、TALEAD、Lead"
    elif user.role.code == "SYSTEM_ADMIN":
        permitted_roles = REGULAR_ACCOUNT_ROLE_CODES
        scope_label = "全部常规账号"
    else:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以查看账号状态")

    employees = db.query(Employee).order_by(Employee.name, Employee.employee_no).all()
    role_map = roles_at(db, [employee.id for employee in employees])
    target_employees = [employee for employee in employees if role_map.get(employee.id) and role_map[employee.id].code in permitted_roles]
    target_ids = [employee.id for employee in target_employees]
    accounts = {
        account.employee_id: account
        for account in db.query(UserAccount).filter(UserAccount.employee_id.in_(target_ids)).all()
    } if target_ids else {}
    attraction_ids = {employee.attraction_id for employee in target_employees if employee.attraction_id}
    attractions = {
        row.id: row.name
        for row in db.query(Attraction).filter(Attraction.id.in_(attraction_ids)).all()
    } if attraction_ids else {}
    items = []
    for employee in target_employees:
        account = accounts.get(employee.id)
        role = role_map[employee.id]
        account_status, account_status_code = account_login_state(account, employee)
        password_status, password_status_code = account_password_status(account)
        items.append(
            {
                "employee_id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": attractions.get(employee.attraction_id, "未分配景点圈"),
                "login_account": account.login_account if account else "",
                "account_status": account_status,
                "account_status_code": account_status_code,
                "password_status": password_status,
                "password_status_code": password_status_code,
                "password_changed_at": account.password_changed_at.strftime("%Y-%m-%d %H:%M:%S") if account and account.password_changed_at else "",
                "login_status": "已登录" if account and account.last_login_at else "从未登录",
                "last_login_at": account.last_login_at.strftime("%Y-%m-%d %H:%M:%S") if account and account.last_login_at else "",
            }
        )
    write_audit(
        db,
        user.employee,
        "查看账号状态",
        "account_status",
        after={"scope": scope_label, "returned_count": len(items)},
        ip_address=client_ip(request),
    )
    db.commit()
    return {"scope_label": scope_label, "items": items}


def gsm_candidates_for_attraction(db: Session, attraction_id: int, on_date: str | None = None) -> list[tuple[Employee, Role]]:
    return gsm_candidates_for_attractions_bulk(db, [attraction_id], on_date).get(attraction_id, [])


def primary_gsm_for_attraction(db: Session, attraction_id: int, on_date: str | None = None) -> tuple[Employee | None, Role | None]:
    candidates = gsm_candidates_for_attraction(db, attraction_id, on_date)
    return candidates[0] if candidates else (None, None)


@router.get("/hr/organization")
def hr_organization(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    today_value = date.today().isoformat()
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee)
    attraction_query = db.query(Attraction)
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
        attraction_query = attraction_query.filter(Attraction.id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.name, Employee.employee_no).all()
    payloads = employee_payloads(db, employees, today_value)
    roles = roles_at(db, [employee.id for employee in employees], today_value)
    rows = []
    represented: set[int] = set()

    for attraction in attraction_query.order_by(Attraction.name).all():
        attraction_employees = [employee for employee in employees if employee.attraction_id == attraction.id and employee.is_active]
        if not attraction_employees:
            continue
        attraction_node = f"hr-attraction-{attraction.id}"
        rows.append({"node_id": attraction_node, "parent_id": "", "level": 0, "node_type": "attraction", "name": attraction.name})
        gsm_candidates = gsm_candidates_for_attraction(db, attraction.id, today_value)
        gsm_managers = [(employee, role) for employee, role in gsm_candidates if role.code == "GSM"]
        ta_gsm_managers = [(employee, role) for employee, role in gsm_candidates if role.code == "TA_GSM"]
        management_parent = attraction_node
        if len(gsm_managers) == 1:
            primary_gsm, _primary_gsm_role = gsm_managers[0]
            management_parent = f"hr-gsm-{attraction.id}-{primary_gsm.id}"
            rows.append({"node_id": management_parent, "parent_id": attraction_node, "level": 1, "node_type": "employee", "hierarchy_role": "gsm", "employee": payloads[primary_gsm.id]})
            represented.add(primary_gsm.id)
        elif len(gsm_managers) > 1:
            # Parallel GSMs jointly own every supervisor group in the circle.
            # A single team node prevents duplicating supervisors under each GSM.
            for gsm, _gsm_role in gsm_managers:
                gsm_node = f"hr-gsm-{attraction.id}-{gsm.id}"
                rows.append({"node_id": gsm_node, "parent_id": attraction_node, "level": 1, "node_type": "employee", "hierarchy_role": "gsm", "employee": payloads[gsm.id]})
                represented.add(gsm.id)
            management_parent = f"hr-gsm-team-{attraction.id}"
            manager_names = "、".join(gsm.name for gsm, _gsm_role in gsm_managers)
            rows.append({"node_id": management_parent, "parent_id": attraction_node, "level": 1, "node_type": "gsm_team", "name": f"主管组（由GSM共同承接：{manager_names}）"})
        else:
            management_parent = f"hr-gsm-team-{attraction.id}"
            rows.append({"node_id": management_parent, "parent_id": attraction_node, "level": 1, "node_type": "gsm_team", "name": "主管组（未配置GSM）"})

        leaders = [employee for employee in attraction_employees if roles[employee.id] and roles[employee.id].code in LEADER_CODES]
        assigned_frontline: set[int] = set()
        for leader in sorted(leaders, key=lambda employee: (employee.name, employee.employee_no)):
            leader_node = f"hr-leader-{leader.id}"
            rows.append({"node_id": leader_node, "parent_id": management_parent, "level": 2, "node_type": "employee", "hierarchy_role": "supervisor", "employee": payloads[leader.id]})
            represented.add(leader.id)
            leader_members = [
                member for member in attraction_employees
                if payloads[member.id]["leader_id"] == leader.id
                and roles.get(member.id)
                and roles[member.id].code in FRONTLINE_CODES
            ]
            for member in sorted(leader_members, key=lambda employee: (employee.name, employee.employee_no)):
                rows.append({"node_id": f"hr-employee-{member.id}", "parent_id": leader_node, "level": 3, "node_type": "employee", "hierarchy_role": "frontline", "employee": payloads[member.id]})
                assigned_frontline.add(member.id)
                represented.add(member.id)

        unassigned = [
            employee for employee in attraction_employees
            if roles[employee.id] and roles[employee.id].code in FRONTLINE_CODES and employee.id not in assigned_frontline
        ]
        if unassigned:
            placeholder = f"hr-unassigned-{attraction.id}"
            rows.append({"node_id": placeholder, "parent_id": management_parent, "level": 2, "node_type": "placeholder", "name": "未分配主管", "member_count": len(unassigned)})
            for employee in sorted(unassigned, key=lambda item: (item.name, item.employee_no)):
                rows.append({"node_id": f"hr-employee-{employee.id}", "parent_id": placeholder, "level": 3, "node_type": "employee", "hierarchy_role": "frontline", "employee": payloads[employee.id]})
                represented.add(employee.id)

        # TA GSM is a peer in the GSM layer and never owns a supervisor branch.
        for gsm, _gsm_role in ta_gsm_managers:
            gsm_node = f"hr-gsm-{attraction.id}-{gsm.id}"
            rows.append({"node_id": gsm_node, "parent_id": attraction_node, "level": 1, "node_type": "employee", "hierarchy_role": "gsm", "employee": payloads[gsm.id]})
            represented.add(gsm.id)

    other_employees = [employee for employee in employees if employee.id not in represented]
    if other_employees:
        other_node = "hr-other-employees"
        rows.append({"node_id": other_node, "parent_id": "", "level": 0, "node_type": "other", "name": "其他管理人员 / 未归属员工"})
        for employee in other_employees:
            rows.append({"node_id": f"hr-other-{employee.id}", "parent_id": other_node, "level": 1, "node_type": "employee", "hierarchy_role": "other", "employee": payloads[employee.id]})
    return {"rows": rows}


@router.get("/hr/import-template")
def employee_import_template(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    wb = build_employee_import_template()
    watermark_workbook(wb, user.employee.employee_no)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    write_audit(db, user.employee, "下载员工导入模板", "employee_import_template", None)
    db.commit()
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="employee_import_template_v2.xlsx"'})


@router.post("/hr/import-employees")
async def import_employees(request: Request, workbook: UploadFile = File(...), db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    if Path(workbook.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(400, "请上传xlsx文件")
    content = await workbook.read(5 * 1024 * 1024 + 1)
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "导入文件不能为空且不能超过5MB")
    try:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:
        raise HTTPException(400, "Excel文件无法读取") from exc
    expected = ["员工号", "姓名", "角色代码", "景点圈", "初始密码", "在职", "账号启用", "任职开始日", "任职结束日", "到期恢复角色代码"]
    header = [str(cell or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if header[:len(expected)] != expected:
        raise HTTPException(400, "表头与V2导入模板不一致")
    roles = {row.code: row for row in db.query(Role).all()}
    circles = {
        row.name: row
        for row in db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .all()
    }
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    allowed_circle_names = {
        row.name for row in db.query(Attraction).filter(Attraction.id.in_(allowed_attractions)).all()
    } if allowed_attractions is not None else None
    role_codes_allowed = set(roles) if "SYSTEM_ADMIN" in user.permissions else CIRCLE_HR_MANAGED_ROLE_CODES
    parsed, seen, errors = [], set(), []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values[:len(expected)]) + [None] * max(0, len(expected) - len(values))
        employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code = [str(v).strip() if v is not None else "" for v in values]
        if not any((employee_no, name, role_code, attraction_name)):
            continue
        if not employee_no or not name or role_code not in roles:
            errors.append(f"第{row_number}行：员工号、姓名必填且角色代码必须有效")
            continue
        if role_code not in role_codes_allowed:
            errors.append(f"第{row_number}行：景点圈HR只能创建LEAD及以下账号")
            continue
        if len(employee_no) != 7 or not employee_no.isdigit():
            errors.append(f"第{row_number}行：员工号必须为7位纯数字")
            continue
        if attraction_name and attraction_name not in circles:
            errors.append(f"第{row_number}行：景点圈只能填写热力追踪、矮人迷宫或小熊罐子")
            continue
        if allowed_circle_names is not None and attraction_name not in allowed_circle_names:
            errors.append(f"第{row_number}行：景点圈HR只能导入所属景点圈员工")
            continue
        if (
            employee_no in seen
            or db.query(Employee).filter(Employee.employee_no == employee_no).first()
            or db.query(UserAccount).filter(UserAccount.login_account == employee_no).first()
        ):
            errors.append(f"第{row_number}行：员工号{employee_no}重复")
            continue
        if ends_on and return_code not in roles:
            errors.append(f"第{row_number}行：临时角色必须填写有效的到期恢复角色")
            continue
        if ends_on and return_code not in role_codes_allowed:
            errors.append(f"第{row_number}行：景点圈HR只能将到期恢复角色设置为LEAD及以下")
            continue
        seen.add(employee_no)
        parsed.append((employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code))
    if errors:
        raise HTTPException(400, "；".join(errors[:20]))
    created = 0
    try:
        for employee_no, name, role_code, attraction_name, password, active, enabled, starts_on, ends_on, return_code in parsed:
            attraction = None
            if attraction_name:
                attraction = circles[attraction_name]
            employee = Employee(employee_no=employee_no, name=name, attraction_id=attraction.id if attraction else None, is_active=active not in {"否", "0", "false", "False"}, hired_on=starts_on or date.today().isoformat())
            db.add(employee)
            db.flush()
            db.add(EmployeeRoleAssignment(employee_id=employee.id, role_id=roles[role_code].id, starts_on=starts_on or date.today().isoformat(), ends_on=ends_on or None, assignment_type="temporary" if ends_on else "permanent", return_role_id=roles[return_code].id if ends_on else None, status="active", reason="HR Excel导入", created_by=user.id))
            # Keep Excel imports consistent with the documented employee-account
            # rule. An explicitly supplied initial password still takes priority.
            initial_password = password or default_initial_password(employee_no)
            db.add(
                UserAccount(
                    employee_id=employee.id,
                    login_account=employee_no,
                    password_hash=hash_password(initial_password),
                    enabled=enabled not in {"否", "0", "false", "False"},
                    must_change_password=True,
                    credential_initialized=True,
                )
            )
            created += 1
        write_audit(db, user.employee, "Excel导入员工", "employee_import", after={"created": created, "filename": workbook.filename}, ip_address=client_ip(request))
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True, "created": created}


@router.post("/hr/employees")
def create_employee(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee_no = str(payload.get("employee_no") or "").strip()
    name = str(payload.get("name") or "").strip()
    role = db.query(Role).filter(Role.code == str(payload.get("role_code") or "")).first()
    if not employee_no or not name or not role:
        raise HTTPException(400, "员工号、姓名和角色必填")
    if len(employee_no) != 7 or not employee_no.isdigit():
        raise HTTPException(400, "员工号必须为7位纯数字")
    ensure_hr_role_allowed(user, role.code)
    if db.query(Employee).filter(Employee.employee_no == employee_no).first() or db.query(UserAccount).filter(UserAccount.login_account == employee_no).first():
        raise HTTPException(400, "该员工号已存在，不能重复创建账号")
    attraction_id = employee_circle_id(db, payload.get("attraction_id"))
    ensure_scoped_hr_attraction(db, user, attraction_id)
    employee = Employee(employee_no=employee_no, name=name, attraction_id=attraction_id, is_active=True, hired_on=str(payload.get("hired_on") or date.today().isoformat()))
    db.add(employee)
    db.flush()
    ends_on = str(payload.get("ends_on") or "").strip() or None
    return_role = db.query(Role).filter(Role.code == str(payload.get("return_role_code") or "")).first() if ends_on else None
    if ends_on:
        if not return_role:
            raise HTTPException(400, "临时角色必须选择有效的到期恢复角色")
        ensure_hr_role_allowed(user, return_role.code, "到期恢复角色")
    db.add(
        EmployeeRoleAssignment(
            employee_id=employee.id,
            role_id=role.id,
            starts_on=str(payload.get("starts_on") or date.today().isoformat()),
            ends_on=ends_on,
            assignment_type="temporary" if ends_on else "permanent",
            return_role_id=return_role.id if return_role else None,
            status="active",
            reason="HR新建员工",
            created_by=user.id,
        )
    )
    db.flush()
    synchronize_gsm_management_scope(db, employee, role.code)
    requested_password = str(payload.get("password") or "").strip()
    # New accounts use the employee-number suffix by default. The user must
    # change it on the first successful login; HR may still explicitly provide
    # a different initial password when needed.
    temporary_password = requested_password or default_initial_password(employee_no)
    db.add(
        UserAccount(
            employee_id=employee.id,
            login_account=employee_no,
            password_hash=hash_password(temporary_password),
            enabled=True,
            must_change_password=True,
            credential_initialized=True,
        )
    )
    write_audit(db, user.employee, "新建员工", "employee", employee.id, after={"employee_no": employee_no, "role": role.name}, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {
        "ok": True,
        "employee": employee_payload(db, employee),
        "temporary_password": temporary_password,
        "must_change_password": True,
    }


@router.post("/accounts/reset-password")
def reset_employee_password(
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("PASSWORD_RESET")),
):
    employee_no = str(payload.get("employee_no") or "").strip()
    name = str(payload.get("name") or "").strip()
    if not employee_no or not name:
        raise HTTPException(400, "员工号和姓名必填")
    employee = db.query(Employee).filter(Employee.employee_no == employee_no).first()
    if not employee or employee.name != name:
        raise HTTPException(400, "员工号或姓名不匹配")
    target_role = role_at(db, employee.id)
    if not target_role:
        raise HTTPException(400, "该员工当前没有有效角色")
    ensure_scoped_hr_employee(db, user, employee)
    if target_role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR管理员、景点圈HR和系统管理员账号不能在此处重置")
    if user.role.code == SCOPED_HR_ROLE_CODE and target_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能重置LEAD及以下员工密码")
    if user.role.code == "GSM" and employee.attraction_id not in managed_attraction_ids(db, user.id):
        raise HTTPException(403, "只能重置本人管理景点圈内的员工密码")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if not account:
        raise HTTPException(400, "该员工尚未开通登录账号")
    temporary_password = new_temporary_password()
    account.password_hash = hash_password(temporary_password)
    account.failed_attempts = 0
    account.locked_until = None
    account.must_change_password = True
    account.credential_initialized = True
    account.password_changed_at = None
    db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "重置账号密码",
        "user_account",
        account.id,
        after={
            "employee_no": employee.employee_no,
            "employee_name": employee.name,
            "role": target_role.name,
            "account_enabled": account.enabled,
            "password_rule": "随机一次性临时密码",
            "must_change_password": True,
            "sessions_revoked": True,
        },
        ip_address=client_ip(request),
    )
    db.commit()
    return {
        "ok": True,
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "account_enabled": account.enabled,
        "password_rule": "随机一次性临时密码",
        "temporary_password": temporary_password,
        "must_change_password": True,
        "sessions_revoked": True,
    }


@router.get("/accounts/name-targets")
def account_name_targets(
    keyword: str = "",
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Return only accounts whose current display name this HR user may correct."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以修改账号中文姓名")
    normalized = str(keyword or "").strip()
    if not normalized:
        return {"items": [], "total": 0}
    escaped = normalized.replace("!", "!!").replace("%", "!%").replace("_", "!_")
    pattern = f"%{escaped}%"
    candidates = (
        db.query(Employee)
        .join(UserAccount, UserAccount.employee_id == Employee.id)
        .filter(
            or_(Employee.employee_no.like(pattern, escape="!"), Employee.name.like(pattern, escape="!")),
        )
        .order_by(Employee.name, Employee.employee_no)
        .limit(max(1, min(int(limit or 30), 50)))
        .all()
    )
    roles = roles_at(db, [employee.id for employee in candidates])
    allowed = []
    for employee in candidates:
        role = roles.get(employee.id)
        if not role or role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
            continue
        if user.role.code == SCOPED_HR_ROLE_CODE:
            if employee.attraction_id != user.employee.attraction_id or role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
                continue
        allowed.append(
            {
                "id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": employee.attraction.name if employee.attraction else "未分配景点圈",
            }
        )
    return {"items": allowed, "total": len(allowed)}


@router.get("/hr/employee-number-targets")
def employee_number_targets(
    keyword: str = "",
    limit: int = 30,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Search current or historic employee numbers without exposing HR accounts."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以变更员工号")
    normalized = str(keyword or "").strip()
    if not normalized:
        return {"items": [], "total": 0}
    escaped = normalized.replace("!", "!!").replace("%", "!%").replace("_", "!_")
    pattern = f"%{escaped}%"
    candidate_rows = (
        db.query(Employee, EmployeeNumberHistory.old_employee_no)
        .join(UserAccount, UserAccount.employee_id == Employee.id)
        .outerjoin(EmployeeNumberHistory, EmployeeNumberHistory.employee_id == Employee.id)
        .filter(
            or_(
                Employee.employee_no.like(pattern, escape="!"),
                Employee.name.like(pattern, escape="!"),
                EmployeeNumberHistory.old_employee_no.like(pattern, escape="!"),
            )
        )
        .order_by(Employee.name, Employee.employee_no, EmployeeNumberHistory.id.desc())
        .limit(max(1, min(int(limit or 30) * 3, 150)))
        .all()
    )
    items: list[dict] = []
    seen: set[int] = set()
    for employee, matched_old_number in candidate_rows:
        if employee.id in seen:
            continue
        role = role_at(db, employee.id)
        try:
            ensure_employee_number_change_target(db, user, employee, role)
        except HTTPException:
            continue
        seen.add(employee.id)
        items.append(
            {
                "id": employee.id,
                "employee_no": employee.employee_no,
                "name": employee.name,
                "role_code": role.code,
                "role_name": role.name,
                "attraction_name": employee.attraction.name if employee.attraction else "未分配景点圈",
                "matched_historical_no": matched_old_number or "",
            }
        )
        if len(items) >= max(1, min(int(limit or 30), 50)):
            break
    return {"items": items, "total": len(items)}


@router.post("/hr/employees/{employee_id}/employee-number")
def change_employee_number(
    employee_id: int,
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Migrate a login number while preserving the employee primary key and all records."""
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    role = ensure_employee_number_change_target(db, user, employee)
    new_employee_no = str(payload.get("new_employee_no") or "").strip()
    reason = str(payload.get("reason") or "").strip()
    reset_password = bool(payload.get("reset_password"))
    if len(new_employee_no) != 7 or not new_employee_no.isdigit():
        raise HTTPException(400, "新员工号必须为7位纯数字")
    if not reason:
        raise HTTPException(400, "请填写员工号变更原因")
    if len(reason) > 300:
        raise HTTPException(400, "员工号变更原因不能超过300个字符")
    old_employee_no = employee.employee_no
    if new_employee_no == old_employee_no:
        return {"ok": True, "unchanged": True, "employee_id": employee.id, "employee_no": employee.employee_no}
    collision = db.query(Employee).filter(Employee.employee_no == new_employee_no, Employee.id != employee.id).first()
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    login_collision = db.query(UserAccount).filter(UserAccount.login_account == new_employee_no).first()
    if collision or (login_collision and (not account or login_collision.id != account.id)):
        raise HTTPException(400, "新员工号已存在，不能合并或重复创建账号")
    if not account:
        raise HTTPException(400, "该员工尚未开通登录账号")
    before = {"employee_no": old_employee_no, "login_account": account.login_account, "role": role.name}
    employee.employee_no = new_employee_no
    employee.updated_at = datetime.now()
    account.login_account = new_employee_no
    account.failed_attempts = 0
    account.locked_until = None
    if reset_password:
        account.password_hash = hash_password(default_initial_password(new_employee_no))
        account.must_change_password = True
        account.credential_initialized = True
        account.password_changed_at = None
    db.add(
        EmployeeNumberHistory(
            employee_id=employee.id,
            old_employee_no=old_employee_no,
            new_employee_no=new_employee_no,
            effective_on=date.today().isoformat(),
            reason=reason,
            changed_by=user.id,
            changed_by_name=user.name,
        )
    )
    revoked_sessions = db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "变更员工号",
        "employee_number_change",
        employee.id,
        before=before,
        after={
            "employee_no": new_employee_no,
            "login_account": new_employee_no,
            "role": role.name,
            "password_reset": reset_password,
            "sessions_revoked": int(revoked_sessions),
        },
        reason=reason,
        ip_address=client_ip(request),
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(400, "新员工号已存在，不能合并或重复创建账号")
    invalidate_data_caches()
    return {
        "ok": True,
        "unchanged": False,
        "employee_id": employee.id,
        "old_employee_no": old_employee_no,
        "employee_no": new_employee_no,
        "sessions_revoked": int(revoked_sessions),
        "must_change_password": bool(account.must_change_password),
        "password_reset": reset_password,
    }


@router.post("/accounts/update-name")
def update_account_name(
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以修改账号中文姓名")
    employee_id = int(payload.get("employee_id") or 0)
    new_name = str(payload.get("name") or "").strip()
    if not employee_id or not new_name:
        raise HTTPException(400, "请选择账号并填写中文姓名")
    if len(new_name) > 100:
        raise HTTPException(400, "中文姓名不能超过100个字符")
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "账号不存在")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    role = role_at(db, employee.id)
    if not account or not role:
        raise HTTPException(400, "该员工尚未配置有效登录账号或角色")
    ensure_scoped_hr_employee(db, user, employee)
    if role.code in {"HR_ADMIN", "HR_CIRCLE", "SYSTEM_ADMIN"}:
        raise HTTPException(403, "HR和系统管理员账号不能在此处修改姓名")
    if user.role.code == SCOPED_HR_ROLE_CODE and role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能修改本圈CM、TR、TA主管或主管的姓名")
    old_name = employee.name
    if old_name == new_name:
        return {"ok": True, "employee_id": employee.id, "employee_no": employee.employee_no, "name": employee.name, "unchanged": True}
    employee.name = new_name
    employee.updated_at = datetime.now()
    write_audit(
        db,
        user.employee,
        "修改账号中文姓名",
        "employee",
        employee.id,
        before={"employee_no": employee.employee_no, "name": old_name, "role": role.name},
        after={"employee_no": employee.employee_no, "name": new_name, "role": role.name},
        reason="HR账号姓名纠错",
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "employee_id": employee.id, "employee_no": employee.employee_no, "name": employee.name, "unchanged": False}


@router.put("/hr/employees/{employee_id}")
def update_employee(employee_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    ensure_scoped_hr_employee(db, user, employee)
    existing_role = role_at(db, employee.id)
    if existing_role and existing_role.code not in CIRCLE_HR_MANAGED_ROLE_CODES and "SYSTEM_ADMIN" not in user.permissions:
        raise HTTPException(403, "景点圈HR只能编辑LEAD及以下员工")
    before = employee_payload(db, employee)
    score_sensitive_fields = {"is_active", "employment_status", "role_code", "attraction_id", "leader_id", "group_id"}
    if score_sensitive_fields & payload.keys():
        current_month = date.today().strftime("%Y-%m")
        ensure_month_open(db, current_month, employee.attraction_id, "修改当月人员计分状态")
        if "attraction_id" in payload and payload.get("attraction_id") not in (None, ""):
            requested_attraction_id = employee_circle_id(db, payload.get("attraction_id"))
            if requested_attraction_id != employee.attraction_id:
                ensure_month_open(db, current_month, requested_attraction_id, "调入员工")
        if str(payload.get("employment_status") or "").strip().lower() == "loa":
            loa_start = parse_iso_date(str(payload.get("loa_start_date") or date.today().isoformat()), "LOA开始日期")
            for loa_month in months_between(loa_start.replace(day=1), date.today().replace(day=1)):
                ensure_month_open(db, loa_month, employee.attraction_id, "设置LOA")
    attendance_state_changed = "is_active" in payload or "employment_status" in payload
    if "name" in payload:
        employee.name = str(payload["name"]).strip()
    if "attraction_id" in payload:
        employee.attraction_id = employee_circle_id(db, payload["attraction_id"])
        ensure_scoped_hr_attraction(db, user, employee.attraction_id)
    requested_employment_status = str(payload.get("employment_status") or "").strip().lower()
    if requested_employment_status:
        if requested_employment_status not in {"active", "loa", "terminated"}:
            raise HTTPException(400, "人员状态无效")
        current_loa = (
            db.query(EmployeeLOAPeriod)
            .filter(
                EmployeeLOAPeriod.employee_id == employee.id,
                EmployeeLOAPeriod.status == "active",
                EmployeeLOAPeriod.ends_on.is_(None),
            )
            .order_by(EmployeeLOAPeriod.starts_on.desc(), EmployeeLOAPeriod.id.desc())
            .first()
        )
        if requested_employment_status == "loa":
            if not existing_role or existing_role.code not in FRONTLINE_CODES:
                raise HTTPException(400, "仅可将CM/TR演职人员设置为LOA")
            starts_on = str(payload.get("loa_start_date") or date.today().isoformat())
            parse_iso_date(starts_on, "LOA开始日期")
            if not current_loa:
                employee.is_active = True
                employee.terminated_on = None
                db.add(
                    EmployeeLOAPeriod(
                        employee_id=employee.id,
                        starts_on=starts_on,
                        status="active",
                        note=str(payload.get("reason") or "HR设置LOA").strip() or None,
                        created_by=user.id,
                        created_by_name=user.name,
                    )
                )
                write_audit(
                    db,
                    user.employee,
                    "设置LOA（长期病假）",
                    "employee_loa",
                    employee.id,
                    after={"starts_on": starts_on, "status": "LOA（长期病假）"},
                    reason=str(payload.get("reason") or ""),
                    ip_address=client_ip(request),
                )
        else:
            if current_loa:
                end_value = date.today() - timedelta(days=1)
                if end_value.isoformat() < current_loa.starts_on:
                    current_loa.status = "cancelled"
                    current_loa.ends_on = current_loa.starts_on
                else:
                    current_loa.status = "ended"
                    current_loa.ends_on = end_value.isoformat()
                current_loa.ended_by = user.id
                current_loa.ended_by_name = user.name
                current_loa.ended_at = datetime.now()
                write_audit(
                    db,
                    user.employee,
                    "结束LOA（长期病假）",
                    "employee_loa",
                    employee.id,
                    after={"ends_on": current_loa.ends_on, "next_status": requested_employment_status},
                    reason=str(payload.get("reason") or ""),
                    ip_address=client_ip(request),
                )
            if requested_employment_status == "terminated" and groups_led_by(db, employee.id):
                raise HTTPException(400, "该员工仍在带组，请先整组移交")
            employee.is_active = requested_employment_status == "active"
            employee.terminated_on = None if employee.is_active else date.today().isoformat()
    elif "is_active" in payload:
        if not bool(payload["is_active"]) and groups_led_by(db, employee.id):
            raise HTTPException(400, "该员工仍在带组，请先整组移交")
        employee.is_active = bool(payload["is_active"])
        employee.terminated_on = None if employee.is_active else date.today().isoformat()
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if account and "account_enabled" in payload:
        next_enabled = bool(payload["account_enabled"])
        if next_enabled != account.enabled:
            account.enabled = next_enabled
            account.disabled_at = None if next_enabled else datetime.now()
    new_role_code = str(payload.get("role_code") or "")
    current_role = existing_role
    resulting_role = current_role
    if new_role_code and current_role and new_role_code != current_role.code:
        attendance_state_changed = True
        new_role = db.query(Role).filter(Role.code == new_role_code).first()
        if not new_role:
            raise HTTPException(400, "角色不存在")
        ensure_hr_role_allowed(user, new_role.code)
        led_groups = groups_led_by(db, employee.id)
        if new_role.code not in LEADER_CODES and led_groups:
            member_ids = {
                membership.employee_id
                for group in led_groups
                for membership in active_group_memberships(db, group.id)
            }
            if member_ids:
                raise HTTPException(400, "该员工仍有组员，请先整组移交")
            pending_review_count = db.query(RecognitionRecord).filter(
                RecognitionRecord.assigned_reviewer_id == employee.id,
                RecognitionRecord.status == "pending",
            ).count()
            if pending_review_count:
                raise HTTPException(400, f"该员工还有{pending_review_count}条待复核记录，请先完成审批后再更改角色")
            for group in led_groups:
                leader_assignment = active_group_leader(db, group.id)
                if leader_assignment:
                    leader_assignment.status = "ended"
                    leader_assignment.ends_on = date.today().isoformat()
                group.status = "closed"
                group.revision += 1
                db.query(SystemAlert).filter(
                    SystemAlert.group_id == group.id,
                    SystemAlert.status == "open",
                ).update(
                    {SystemAlert.status: "handled", SystemAlert.handled_by: user.id, SystemAlert.handled_at: datetime.now()},
                    synchronize_session=False,
                )
                write_audit(
                    db,
                    user.employee,
                    "关闭空工作组",
                    "work_group",
                    group.id,
                    before={"leader": employee.name, "status": "active", "member_count": 0},
                    after={"leader": "", "status": "closed", "member_count": 0},
                    reason=f"组长角色由{current_role.name}变更为{new_role.name}",
                    ip_address=client_ip(request),
                )
        current_assignment = (
            db.query(EmployeeRoleAssignment)
            .filter(EmployeeRoleAssignment.employee_id == employee.id, EmployeeRoleAssignment.status == "active")
            .order_by(EmployeeRoleAssignment.starts_on.desc())
            .first()
        )
        if current_assignment:
            current_assignment.ends_on = date.today().isoformat()
            current_assignment.status = "expired"
        ends_on = str(payload.get("role_ends_on") or "").strip() or None
        return_role = db.query(Role).filter(Role.code == str(payload.get("return_role_code") or "")).first() if ends_on else None
        if ends_on:
            if not return_role:
                raise HTTPException(400, "临时角色必须选择有效的到期恢复角色")
            ensure_hr_role_allowed(user, return_role.code, "到期恢复角色")
        db.add(
            EmployeeRoleAssignment(
                employee_id=employee.id,
                role_id=new_role.id,
                starts_on=date.today().isoformat(),
                ends_on=ends_on,
                assignment_type="temporary" if ends_on else "permanent",
                return_role_id=return_role.id if return_role else None,
                status="active",
                reason=str(payload.get("reason") or "HR变更角色"),
                created_by=user.id,
            )
        )
        resulting_role = new_role

    synchronize_gsm_management_scope(db, employee, resulting_role.code if resulting_role else None)

    group_change_requested = "leader_id" in payload or "group_id" in payload
    if group_change_requested:
        requested_leader_id = int(payload["leader_id"]) if payload.get("leader_id") else None
        requested_group_id = int(payload["group_id"]) if payload.get("group_id") else None
        current_membership = (
            db.query(GroupMembership)
            .filter(
                GroupMembership.employee_id == employee.id,
                GroupMembership.status == "active",
                GroupMembership.starts_on <= date.today().isoformat(),
                or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= date.today().isoformat()),
            )
            .order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc())
            .first()
        )
        if not resulting_role or resulting_role.code not in FRONTLINE_CODES:
            if requested_leader_id or requested_group_id:
                raise HTTPException(400, "只有CM/TR可以设置组长")
            if current_membership:
                current_membership.status = "ended"
                current_membership.ends_on = date.today().isoformat()
        else:
            if (requested_leader_id or requested_group_id) and not employee.is_active:
                raise HTTPException(400, "离职员工不能设置组长")
            target_group = db.get(WorkGroup, requested_group_id) if requested_group_id else None
            target_leader = None
            if target_group:
                target_assignment = active_group_leader(db, target_group.id)
                target_leader = target_assignment.leader if target_assignment else None
                if requested_leader_id and (not target_leader or target_leader.id != requested_leader_id):
                    raise HTTPException(409, "工作组组长已经变化，请刷新后重试")
            elif requested_leader_id:
                target_leader = db.get(Employee, requested_leader_id)
                candidate_groups = [group for group in groups_led_by(db, requested_leader_id) if group.attraction_id == employee.attraction_id]
                if len(candidate_groups) > 1:
                    raise HTTPException(409, "该组长有多个工作组，请选择具体工作组")
                if candidate_groups:
                    target_group = candidate_groups[0]
            if target_leader:
                target_role = role_at(db, target_leader.id)
                if not target_leader.is_active or not target_role or target_role.code not in LEADER_CODES:
                    raise HTTPException(400, "新组长必须是在职TA主管或主管")
                if target_leader.attraction_id != employee.attraction_id:
                    raise HTTPException(400, "新组长必须与员工属于同一景点圈")
                if not target_group:
                    target_group = WorkGroup(name=f"{target_leader.name}工作组", attraction_id=employee.attraction_id, status="active")
                    db.add(target_group)
                    db.flush()
                    db.add(GroupLeaderAssignment(group_id=target_group.id, leader_employee_id=target_leader.id, starts_on=date.today().isoformat(), status="active"))
            elif requested_group_id:
                raise HTTPException(400, "所选工作组当前没有有效组长")

            # Returning a former TA主管/主管 to CM/TR should restore the
            # latest viable historic group in the selected circle when that
            # choice is unambiguous.  We never guess between two historical
            # groups: HR must select the intended LEAD in that case.
            if not target_group and not requested_leader_id and new_role_code and current_role and current_role.code in LEADER_CODES:
                candidates: list[tuple[GroupMembership, WorkGroup, Employee]] = []
                history = db.query(GroupMembership).filter(
                    GroupMembership.employee_id == employee.id,
                    GroupMembership.status.in_(("ended", "active")),
                ).order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc()).all()
                seen_groups: set[int] = set()
                for historic in history:
                    if historic.group_id in seen_groups:
                        continue
                    seen_groups.add(historic.group_id)
                    group = db.get(WorkGroup, historic.group_id)
                    assignment = active_group_leader(db, group.id) if group and group.status == "active" and group.attraction_id == employee.attraction_id else None
                    leader = assignment.leader if assignment else None
                    leader_role = role_at(db, leader.id) if leader else None
                    if leader and leader.is_active and leader_role and leader_role.code in LEADER_CODES:
                        candidates.append((historic, group, leader))
                if len(candidates) == 1:
                    _, target_group, target_leader = candidates[0]
                elif len(candidates) > 1:
                    raise HTTPException(409, "该员工在目标景点圈有多个历史工作组，请选择明确组长后再保存")

            pending_query = db.query(RecognitionRecord).filter(RecognitionRecord.employee_id == employee.id, RecognitionRecord.status == "pending")
            current_group_id = current_membership.group_id if current_membership else None
            target_group_id = target_group.id if target_group else None
            if target_group_id != current_group_id:
                if not target_group and pending_query.count():
                    raise HTTPException(400, "该员工还有待复核记录，必须选择新组长")
                if current_membership:
                    current_membership.status = "ended"
                    current_membership.ends_on = date.today().isoformat()
                if target_group:
                    db.add(GroupMembership(group_id=target_group.id, employee_id=employee.id, starts_on=date.today().isoformat(), status="active", reason=str(payload.get("reason") or "HR调整组长")))
                    pending_query.update({RecognitionRecord.assigned_reviewer_id: target_leader.id}, synchronize_session=False)
                write_audit(
                    db,
                    user.employee,
                    "调整员工组长",
                    "employee_group",
                    employee.id,
                    before={"group_id": current_group_id, "leader_id": before.get("leader_id"), "leader_name": before.get("leader_name")},
                    after={"group_id": target_group_id, "leader_id": target_leader.id if target_leader else None, "leader_name": target_leader.name if target_leader else "未分配"},
                    reason=str(payload.get("reason") or "HR员工管理页面调整"),
                    ip_address=client_ip(request),
                )
    employee.updated_at = datetime.now()
    db.flush()
    if attendance_state_changed:
        recalculate_attendance(db, employee, date.today().strftime("%Y-%m"))
    write_audit(db, user.employee, "修改员工", "employee", employee.id, before=before, after=employee_payload(db, employee), reason=str(payload.get("reason") or ""), ip_address=client_ip(request))
    db.commit()
    return {"ok": True}


def ensure_login_account_archive_target(db: Session, user: V2User, employee: Employee) -> tuple[Role, UserAccount, dict]:
    """Authorize credential removal without ever deleting the employee archive."""
    if user.role.code not in {SCOPED_HR_ROLE_CODE, "SYSTEM_ADMIN"}:
        raise HTTPException(403, "仅景点圈HR和最高管理员可以删除登录账号")
    ensure_scoped_hr_employee(db, user, employee)
    role = role_at(db, employee.id)
    if not role or role.code not in REGULAR_ACCOUNT_ROLE_CODES:
        raise HTTPException(403, "HR和最高管理员账号不能在此删除")
    if user.role.code == SCOPED_HR_ROLE_CODE and role.code not in CIRCLE_HR_MANAGED_ROLE_CODES:
        raise HTTPException(403, "景点圈HR只能删除本圈CM、TR、TA主管或主管的登录账号")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    state = login_account_archive_state(employee, account, role)
    if not account:
        raise HTTPException(400, state["reason"] or "该员工没有可删除的登录账号")
    if not state["eligible"]:
        raise HTTPException(400, state["reason"])
    return role, account, state


@router.delete("/hr/employees/{employee_id}/account")
def archive_employee_login_account(
    employee_id: int,
    request: Request,
    payload: dict | None = None,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Remove only credentials and sessions after seven days; keep all data."""
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "员工不存在")
    role, account, _state = ensure_login_account_archive_target(db, user, employee)
    before = {
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "role": role.name,
        "login_account": account.login_account,
        "employee_active": employee.is_active,
        "account_enabled": account.enabled,
    }
    revoked_sessions = db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    db.delete(account)
    employee.account_deleted_at = datetime.now()
    employee.account_deleted_by_name = user.name
    employee.updated_at = datetime.now()
    write_audit(
        db,
        user.employee,
        "删除停用登录账号",
        "employee_login_archive",
        employee.id,
        before=before,
        after={
            "employee_no": employee.employee_no,
            "employee_name": employee.name,
            "role": role.name,
            "login_removed": True,
            "sessions_revoked": int(revoked_sessions),
            "employee_archive_retained": True,
            "business_history_retained": True,
        },
        reason=str((payload or {}).get("reason") or "离职/停用满7天后删除登录账号，保留员工与业务档案"),
        ip_address=client_ip(request),
    )
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "employee_id": employee.id, "sessions_revoked": int(revoked_sessions), "archive_retained": True}


@router.post("/hr/employees/batch-leaders")
def batch_assign_unclassified_leaders(
    payload: dict,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("HR_MANAGE")),
):
    """Assign multiple currently unclassified CM/TR employees in one atomic transaction."""
    raw_items = payload.get("items")
    if not isinstance(raw_items, list) or not raw_items:
        raise HTTPException(400, "请至少选择一名未分类组员的组长")
    if len(raw_items) > 200:
        raise HTTPException(400, "单次最多保存200名组员")
    reason = str(payload.get("reason") or "HR未分类组员批量分组").strip()[:200]
    today_value = date.today().isoformat()
    seen_employee_ids: set[int] = set()
    plans: list[dict] = []

    # Validate the full batch before creating or ending any relationship.
    for index, item in enumerate(raw_items, start=1):
        if not isinstance(item, dict):
            raise HTTPException(400, f"第{index}项格式无效")
        try:
            employee_id = int(item.get("employee_id") or 0)
            leader_id = int(item.get("leader_id") or 0)
            requested_group_id = int(item.get("group_id") or 0) or None
        except (TypeError, ValueError):
            raise HTTPException(400, f"第{index}项员工、组长或工作组编号无效")
        if not employee_id or not leader_id:
            raise HTTPException(400, f"第{index}项必须选择员工和组长")
        if employee_id in seen_employee_ids:
            raise HTTPException(400, f"第{index}项员工重复")
        seen_employee_ids.add(employee_id)

        employee = db.get(Employee, employee_id)
        employee_role = role_at(db, employee_id) if employee else None
        if not employee:
            raise HTTPException(404, f"第{index}项员工不存在")
        ensure_scoped_hr_employee(db, user, employee)
        if not employee.is_active or not employee_role or employee_role.code not in FRONTLINE_CODES:
            raise HTTPException(400, f"{employee.name}不是在职CM/TR，不能批量分组")
        if not employee.attraction_id:
            raise HTTPException(400, f"{employee.name}未配置景点圈")
        ensure_month_open(db, date.today().strftime("%Y-%m"), employee.attraction_id, "批量调整组长")

        current_membership = (
            db.query(GroupMembership)
            .filter(
                GroupMembership.employee_id == employee.id,
                GroupMembership.status == "active",
                GroupMembership.starts_on <= today_value,
                or_(GroupMembership.ends_on.is_(None), GroupMembership.ends_on >= today_value),
            )
            .order_by(GroupMembership.starts_on.desc(), GroupMembership.id.desc())
            .first()
        )
        current_group = db.get(WorkGroup, current_membership.group_id) if current_membership else None
        current_assignment = active_group_leader(db, current_group.id) if current_group else None
        if current_assignment:
            raise HTTPException(409, f"{employee.name}已分配组长，请刷新页面后重试")

        leader = db.get(Employee, leader_id)
        leader_role = role_at(db, leader_id) if leader else None
        if not leader or not leader.is_active or not leader_role or leader_role.code not in LEADER_CODES:
            raise HTTPException(400, f"{employee.name}选择的组长不是在职TA主管或主管")
        if leader.attraction_id != employee.attraction_id:
            raise HTTPException(400, f"{employee.name}与所选组长不属于同一景点圈")

        target_group = db.get(WorkGroup, requested_group_id) if requested_group_id else None
        if target_group:
            target_assignment = active_group_leader(db, target_group.id)
            if target_group.status == "closed" or target_group.attraction_id != employee.attraction_id:
                raise HTTPException(400, f"{employee.name}选择的工作组无效")
            if not target_assignment or target_assignment.leader_employee_id != leader.id:
                raise HTTPException(409, f"{employee.name}选择的工作组组长已经变化，请刷新后重试")
        else:
            existing_groups = [group for group in groups_led_by(db, leader.id) if group.attraction_id == employee.attraction_id]
            if len(existing_groups) > 1:
                raise HTTPException(409, f"{leader.name}有多个工作组，请刷新后选择具体工作组")
            target_group = existing_groups[0] if existing_groups else None

        plans.append(
            {
                "employee": employee,
                "leader": leader,
                "current_membership": current_membership,
                "current_group": current_group,
                "target_group": target_group,
            }
        )

    created_groups: dict[int, WorkGroup] = {}
    changes = []
    try:
        for plan in plans:
            employee = plan["employee"]
            leader = plan["leader"]
            target_group = plan["target_group"]
            if not target_group:
                target_group = created_groups.get(leader.id)
                if not target_group:
                    target_group = WorkGroup(
                        name=f"{leader.name}工作组",
                        attraction_id=employee.attraction_id,
                        status="active",
                    )
                    db.add(target_group)
                    db.flush()
                    db.add(
                        GroupLeaderAssignment(
                            group_id=target_group.id,
                            leader_employee_id=leader.id,
                            starts_on=today_value,
                            status="active",
                        )
                    )
                    created_groups[leader.id] = target_group

            current_membership = plan["current_membership"]
            if current_membership:
                current_membership.status = "ended"
                current_membership.ends_on = today_value
            db.add(
                GroupMembership(
                    group_id=target_group.id,
                    employee_id=employee.id,
                    starts_on=today_value,
                    status="active",
                    reason=reason,
                )
            )
            db.query(RecognitionRecord).filter(
                RecognitionRecord.employee_id == employee.id,
                RecognitionRecord.status == "pending",
            ).update({RecognitionRecord.assigned_reviewer_id: leader.id}, synchronize_session=False)
            change = {
                "employee_id": employee.id,
                "employee_no": employee.employee_no,
                "employee_name": employee.name,
                "group_id": target_group.id,
                "group_name": target_group.name,
                "leader_id": leader.id,
                "leader_name": leader.name,
            }
            changes.append(change)
            write_audit(
                db,
                user.employee,
                "批量调整未分类组员组长",
                "employee_group",
                employee.id,
                before={
                    "group_id": plan["current_group"].id if plan["current_group"] else None,
                    "leader_id": None,
                    "leader_name": "未分配",
                },
                after={"group_id": target_group.id, "leader_id": leader.id, "leader_name": leader.name},
                reason=reason,
                ip_address=client_ip(request),
            )
        write_audit(
            db,
            user.employee,
            "批量保存未分类组员",
            "employee_group_batch",
            None,
            after={"count": len(changes), "changes": changes},
            reason=reason,
            ip_address=client_ip(request),
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    invalidate_data_caches()
    return {"ok": True, "updated": len(changes), "changes": changes}


def circle_transfer_payload(
    db: Session,
    row: CircleTransferRequest,
    *,
    groups: dict[int, WorkGroup] | None = None,
    leaders: dict[int, Employee] | None = None,
) -> dict:
    group_map = groups or {}
    leader_map = leaders or {}
    source_group = group_map.get(row.source_group_id) if row.source_group_id else None
    source_leader = leader_map.get(row.source_leader_id) if row.source_leader_id else None
    target_group = group_map.get(row.target_group_id) if row.target_group_id else None
    target_leader = leader_map.get(row.target_leader_id) if row.target_leader_id else None
    group_display = group_display_metadata_bulk(
        db,
        [group.id for group in (source_group, target_group) if group],
    )
    try:
        migrated_counts = json.loads(row.migrated_record_counts_json or "{}")
    except (TypeError, ValueError):
        migrated_counts = {}
    return {
        "id": row.id,
        "employee_id": row.employee_id,
        "employee_no": row.employee_no,
        "employee_name": row.employee_name,
        "source_attraction_id": row.source_attraction_id,
        "source_attraction_name": row.source_attraction_name,
        "target_attraction_id": row.target_attraction_id,
        "target_attraction_name": row.target_attraction_name,
        "source_group_name": group_display.get(source_group.id, {}).get("name", source_group.name) if source_group else "未分组",
        "source_leader_name": source_leader.name if source_leader else "未分配",
        "target_group_id": row.target_group_id,
        "target_group_name": group_display.get(target_group.id, {}).get("name", target_group.name) if target_group else "",
        "target_leader_id": row.target_leader_id,
        "target_leader_name": target_leader.name if target_leader else "",
        "reason": row.reason,
        "status": row.status,
        "status_name": {"pending": "待目标HR确认", "completed": "已生效", "rejected": "已拒绝", "cancelled": "已撤回"}.get(row.status, row.status),
        "requested_by_name": row.requested_by_name,
        "requested_at": row.requested_at.strftime("%Y-%m-%d %H:%M:%S"),
        "reviewed_by_name": row.reviewed_by_name or "",
        "reviewed_at": row.reviewed_at.strftime("%Y-%m-%d %H:%M:%S") if row.reviewed_at else "",
        "review_note": row.review_note or "",
        "completed_at": row.completed_at.strftime("%Y-%m-%d %H:%M:%S") if row.completed_at else "",
        "migrated_record_counts": migrated_counts,
    }


@router.get("/hr/circle-transfers")
def circle_transfers(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    transfer_rows = db.query(CircleTransferRequest).order_by(CircleTransferRequest.requested_at.desc(), CircleTransferRequest.id.desc()).limit(500).all()
    if allowed_attractions is not None:
        transfer_rows = [
            row for row in transfer_rows
            if row.source_attraction_id in allowed_attractions or row.target_attraction_id in allowed_attractions
        ]
    circles = (
        db.query(Attraction)
        .filter(Attraction.active.is_(True), Attraction.employee_circle.is_(True))
        .order_by(Attraction.name)
        .all()
    )
    group_ids = {row.source_group_id for row in transfer_rows if row.source_group_id} | {
        row.target_group_id for row in transfer_rows if row.target_group_id
    }
    leader_ids = {row.source_leader_id for row in transfer_rows if row.source_leader_id} | {
        row.target_leader_id for row in transfer_rows if row.target_leader_id
    }
    groups = (
        {group.id: group for group in db.query(WorkGroup).filter(WorkGroup.id.in_(group_ids)).all()}
        if group_ids
        else {}
    )
    leaders = (
        {leader.id: leader for leader in db.query(Employee).filter(Employee.id.in_(leader_ids)).all()}
        if leader_ids
        else {}
    )
    return {
        "items": [circle_transfer_payload(db, row, groups=groups, leaders=leaders) for row in transfer_rows],
        "target_circles": [{"id": row.id, "name": row.name} for row in circles],
    }


@router.post("/hr/circle-transfers")
def create_circle_transfer(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    employee = db.get(Employee, int(payload.get("employee_id") or 0))
    employee_role = role_at(db, employee.id) if employee else None
    if not employee or not employee.is_active or not employee_role or employee_role.code not in FRONTLINE_CODES:
        raise HTTPException(400, "只能为在职CM/TR发起景点圈调动")
    ensure_scoped_hr_employee(db, user, employee)
    source = db.get(Attraction, employee.attraction_id) if employee.attraction_id else None
    target = db.get(Attraction, int(payload.get("target_attraction_id") or 0))
    if not source or not source.employee_circle or not target or not target.active or not target.employee_circle:
        raise HTTPException(400, "原景点圈或目标景点圈无效")
    if source.id == target.id:
        raise HTTPException(400, "目标景点圈不能与当前景点圈相同")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "调动原因必填")
    existing = db.query(CircleTransferRequest).filter(
        CircleTransferRequest.employee_id == employee.id,
        CircleTransferRequest.status == "pending",
    ).first()
    if existing:
        raise HTTPException(409, "该员工已有待确认的景点圈调动申请")
    source_group = current_group_for_employee(db, employee.id)
    source_leader = current_leader_for_employee(db, employee.id)
    row = CircleTransferRequest(
        employee_id=employee.id,
        employee_no=employee.employee_no,
        employee_name=employee.name,
        source_attraction_id=source.id,
        source_attraction_name=source.name,
        target_attraction_id=target.id,
        target_attraction_name=target.name,
        source_group_id=source_group.id if source_group else None,
        source_leader_id=source_leader.id if source_leader else None,
        reason=reason,
        status="pending",
        requested_by=user.id,
        requested_by_name=user.name,
    )
    db.add(row)
    db.flush()
    write_audit(
        db,
        user.employee,
        "发起跨景点圈调动",
        "circle_transfer",
        row.id,
        after=circle_transfer_payload(db, row),
        reason=reason,
        ip_address=client_ip(request),
    )
    db.commit()
    return {"ok": True, "transfer": circle_transfer_payload(db, row)}


def complete_circle_transfer(db: Session, row: CircleTransferRequest, target_leader: Employee, target_group: WorkGroup, reviewer: V2User, request: Request) -> dict:
    employee = db.get(Employee, row.employee_id)
    if not employee or employee.attraction_id != row.source_attraction_id:
        raise HTTPException(409, "员工景点圈已经变化，请撤回申请后重新发起")
    today_value = date.today().isoformat()
    current_month = today_value[:7]
    ensure_month_open(db, current_month, row.source_attraction_id, "迁出员工及当月数据")
    ensure_month_open(db, current_month, row.target_attraction_id, "迁入员工及当月数据")
    old_memberships = (
        db.query(GroupMembership)
        .filter(GroupMembership.employee_id == employee.id, GroupMembership.status == "active")
        .all()
    )
    for membership in old_memberships:
        membership.status = "ended"
        membership.ends_on = max(membership.starts_on, today_value)
        membership.reason = f"跨景点圈调动至{row.target_attraction_name}"
    db.add(
        GroupMembership(
            group_id=target_group.id,
            employee_id=employee.id,
            starts_on=today_value,
            status="active",
            reason=f"由{row.source_attraction_name}调入",
        )
    )
    employee.attraction_id = row.target_attraction_id
    employee.updated_at = datetime.now()
    pending_reviews = db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id == employee.id,
        RecognitionRecord.status == "pending",
    ).update({RecognitionRecord.assigned_reviewer_id: target_leader.id}, synchronize_session=False)
    pending_follow_ups = db.query(DeductionFollowUp).filter(
        DeductionFollowUp.employee_id == employee.id,
        DeductionFollowUp.status == "pending",
    ).update(
        {DeductionFollowUp.supervisor_id: target_leader.id, DeductionFollowUp.supervisor_name: target_leader.name},
        synchronize_session=False,
    )
    recognition_count = db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id == employee.id,
        RecognitionRecord.recognition_month == current_month,
    ).update(
        {RecognitionRecord.home_attraction_id: row.target_attraction_id, RecognitionRecord.home_attraction_name: row.target_attraction_name},
        synchronize_session=False,
    )
    deduction_count = db.query(DeductionRecord).filter(
        DeductionRecord.employee_id == employee.id,
        DeductionRecord.deduction_month == current_month,
    ).update(
        {
            DeductionRecord.attraction_id_snapshot: row.target_attraction_id,
            DeductionRecord.employee_group_id_snapshot: target_group.id,
        },
        synchronize_session=False,
    )
    sick_leave_count = db.query(SickLeaveRecord).filter(
        SickLeaveRecord.employee_id == employee.id,
        SickLeaveRecord.attendance_month == current_month,
    ).update({SickLeaveRecord.attraction_id_snapshot: row.target_attraction_id}, synchronize_session=False)
    recalculate_attendance(db, employee, current_month)
    migrated_counts = {
        "recognitions": recognition_count,
        "deductions": deduction_count,
        "sick_leaves": sick_leave_count,
        "pending_reviews": pending_reviews,
        "pending_follow_ups": pending_follow_ups,
    }
    row.target_group_id = target_group.id
    row.target_leader_id = target_leader.id
    row.status = "completed"
    row.reviewed_by = reviewer.id
    row.reviewed_by_name = reviewer.name
    row.reviewed_at = datetime.now()
    row.completed_at = datetime.now()
    row.migrated_record_counts_json = json.dumps(migrated_counts, ensure_ascii=False)
    write_audit(
        db,
        reviewer.employee,
        "确认跨景点圈调动并同步迁移数据",
        "circle_transfer",
        row.id,
        before={"employee_no": employee.employee_no, "attraction": row.source_attraction_name, "group_id": row.source_group_id},
        after={"attraction": row.target_attraction_name, "group_id": target_group.id, "leader": target_leader.name, "migrated": migrated_counts},
        reason=row.reason,
        ip_address=client_ip(request),
    )
    invalidate_data_caches()
    return migrated_counts


@router.post("/hr/circle-transfers/{transfer_id}/review")
def review_circle_transfer(transfer_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    row = db.get(CircleTransferRequest, transfer_id)
    if not row or row.status != "pending":
        raise HTTPException(404, "待确认调动申请不存在")
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is not None and row.target_attraction_id not in allowed_attractions:
        raise HTTPException(403, "只有目标景点圈HR可以确认该调动")
    action = str(payload.get("action") or "").strip().lower()
    note = str(payload.get("review_note") or "").strip()
    if action == "reject":
        if not note:
            raise HTTPException(400, "拒绝原因必填")
        row.status = "rejected"
        row.reviewed_by = user.id
        row.reviewed_by_name = user.name
        row.reviewed_at = datetime.now()
        row.review_note = note
        write_audit(db, user.employee, "拒绝跨景点圈调动", "circle_transfer", row.id, after=circle_transfer_payload(db, row), reason=note, ip_address=client_ip(request))
        db.commit()
        return {"ok": True, "transfer": circle_transfer_payload(db, row)}
    if action != "accept":
        raise HTTPException(400, "审批动作无效")
    target_leader = db.get(Employee, int(payload.get("target_leader_id") or 0))
    target_role = role_at(db, target_leader.id) if target_leader else None
    if not target_leader or not target_leader.is_active or target_leader.attraction_id != row.target_attraction_id or not target_role or target_role.code not in LEADER_CODES:
        raise HTTPException(400, "请选择目标景点圈内在职TA主管或主管")
    target_group_id = int(payload.get("target_group_id") or 0)
    target_group = db.get(WorkGroup, target_group_id) if target_group_id else None
    if target_group:
        if target_group.status == "closed" or target_group.attraction_id != row.target_attraction_id:
            raise HTTPException(400, "目标工作组无效")
        assignment = active_group_leader(db, target_group.id)
        if not assignment or assignment.leader_employee_id != target_leader.id:
            raise HTTPException(400, "目标工作组与所选LEAD不匹配")
    else:
        target_group = WorkGroup(name=f"{target_leader.name}工作组", attraction_id=row.target_attraction_id, status="active")
        db.add(target_group)
        db.flush()
        db.add(GroupLeaderAssignment(group_id=target_group.id, leader_employee_id=target_leader.id, starts_on=date.today().isoformat(), status="active"))
    row.review_note = note or None
    migrated_counts = complete_circle_transfer(db, row, target_leader, target_group, user, request)
    db.commit()
    return {"ok": True, "transfer": circle_transfer_payload(db, row), "migrated_record_counts": migrated_counts}


@router.post("/hr/circle-transfers/{transfer_id}/cancel")
def cancel_circle_transfer(transfer_id: int, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    row = db.get(CircleTransferRequest, transfer_id)
    if not row or row.status != "pending":
        raise HTTPException(404, "待确认调动申请不存在")
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    if allowed_attractions is not None and row.source_attraction_id not in allowed_attractions:
        raise HTTPException(403, "只有原景点圈HR可以撤回该申请")
    row.status = "cancelled"
    write_audit(db, user.employee, "撤回跨景点圈调动", "circle_transfer", row.id, after=circle_transfer_payload(db, row), reason=row.reason, ip_address=client_ip(request))
    db.commit()
    return {"ok": True, "transfer": circle_transfer_payload(db, row)}


@router.get("/hr/groups")
def hr_groups(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    group_query = db.query(WorkGroup).filter(WorkGroup.status != "closed")
    if allowed_attractions is not None:
        group_query = group_query.filter(WorkGroup.attraction_id.in_(allowed_attractions))
    groups = group_query.order_by(WorkGroup.name).all()
    group_ids = [group.id for group in groups]
    group_display = group_display_metadata_bulk(db, group_ids)
    leader_assignments = active_group_leaders_bulk(db, group_ids)
    memberships = active_group_memberships_bulk(db, group_ids)
    referenced_employees = {assignment.leader_employee_id for assignment in leader_assignments.values() if assignment}
    for member_rows in memberships.values():
        referenced_employees.update(row.employee_id for row in member_rows)
    employees = (
        {employee.id: employee for employee in db.query(Employee).filter(Employee.id.in_(referenced_employees)).all()}
        if referenced_employees
        else {}
    )
    attractions = {attraction.id: attraction.name for attraction in db.query(Attraction).all()}
    rows = []
    for group in groups:
        leader_assignment = leader_assignments.get(group.id)
        members = memberships.get(group.id, [])
        leader = employees.get(leader_assignment.leader_employee_id) if leader_assignment else None
        display = group_display.get(group.id, {})
        rows.append(
            {
                "id": group.id,
                "name": display.get("name", group.name),
                "stored_name": group.name,
                "previous_leader_name": display.get("previous_leader_name", ""),
                "previous_leader_until": display.get("previous_leader_until", ""),
                "attraction_id": group.attraction_id,
                "attraction_name": attractions.get(group.attraction_id, ""),
                "status": group.status,
                "revision": group.revision,
                "leader_id": leader_assignment.leader_employee_id if leader_assignment else None,
                "leader_name": leader.name if leader else "待接管",
                "member_count": len(members),
                "members": [
                    {
                        "id": employee.id,
                        "employee_no": employee.employee_no,
                        "name": employee.name,
                    }
                    for employee in (employees.get(row.employee_id) for row in members)
                    if employee
                ],
            }
        )
    return rows


@router.get("/hr/leader-options")
def leader_options(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    employee_query = db.query(Employee).filter(Employee.is_active.is_(True))
    if allowed_attractions is not None:
        employee_query = employee_query.filter(Employee.attraction_id.in_(allowed_attractions))
    employees = employee_query.order_by(Employee.name).all()
    roles = roles_at(db, [employee.id for employee in employees])
    leaders = [employee for employee in employees if (roles.get(employee.id) or None) and roles[employee.id].code in LEADER_CODES]
    led_groups = groups_led_by_bulk(db, [leader.id for leader in leaders])
    group_display = group_display_metadata_bulk(db, {group.id for rows in led_groups.values() for group in rows})
    gsm_by_attraction = gsm_candidates_for_attractions_bulk(db, {leader.attraction_id for leader in leaders if leader.attraction_id})
    attraction_names = {attraction.id: attraction.name for attraction in db.query(Attraction).all()}
    result = []
    for employee in leaders:
        role = roles[employee.id]
        groups = [group for group in led_groups.get(employee.id, []) if group.attraction_id == employee.attraction_id]
        if not groups:
            groups = [None]
        candidates = gsm_by_attraction.get(employee.attraction_id) if employee.attraction_id else None
        gsm, gsm_role = candidates[0] if candidates else (None, None)
        for group in groups:
            result.append(
                {
                    "id": employee.id,
                    "name": employee.name,
                    "employee_no": employee.employee_no,
                    "role_name": role.name,
                    "attraction_id": employee.attraction_id,
                    "attraction_name": attraction_names.get(employee.attraction_id, "") if employee.attraction_id else "",
                    "group_id": group.id if group else None,
                    "group_name": group_display.get(group.id, {}).get("name", group.name) if group else "新建工作组",
                    "gsm_id": gsm.id if gsm else None,
                    "gsm_name": gsm.name if gsm else "未配置GSM",
                    "gsm_role_name": gsm_role.name if gsm_role else "",
                }
            )
    return result


@router.post("/hr/groups/{group_id}/transfer")
def transfer_group(group_id: int, payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    group = db.get(WorkGroup, group_id)
    if not group:
        raise HTTPException(404, "工作组不存在")
    ensure_scoped_hr_attraction(db, user, group.attraction_id)
    if int(payload.get("revision") or 0) != group.revision:
        raise HTTPException(409, "工作组已被其他操作修改，请刷新后重试")
    new_leader = db.get(Employee, int(payload.get("new_leader_id") or 0))
    new_role = role_at(db, new_leader.id) if new_leader else None
    if not new_leader or not new_leader.is_active or not new_role or new_role.code not in LEADER_CODES:
        raise HTTPException(400, "新组长必须是在职TA主管或主管")
    if new_leader.attraction_id != group.attraction_id:
        raise HTTPException(400, "新组长必须与工作组属于同一景点圈")
    effective_date = str(payload.get("effective_date") or date.today().isoformat())
    if effective_date != date.today().isoformat():
        raise HTTPException(400, "当前版本整组移交仅支持当天生效")
    reason = str(payload.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "移交原因必填")
    old_assignment = active_group_leader(db, group.id)
    if old_assignment and old_assignment.leader_employee_id == new_leader.id:
        raise HTTPException(400, "新旧组长不能相同")
    members = active_group_memberships(db, group.id)
    pending_count = db.query(RecognitionRecord).filter(
        RecognitionRecord.employee_id.in_([member.employee_id for member in members]) if members else RecognitionRecord.id == -1,
        RecognitionRecord.status == "pending",
    ).count()
    transfer = GroupTransfer(
        group_id=group.id,
        old_leader_id=old_assignment.leader_employee_id if old_assignment else None,
        new_leader_id=new_leader.id,
        attraction_id=group.attraction_id,
        effective_date=effective_date,
        member_count=len(members),
        pending_review_count=pending_count,
        reason=reason,
        operator_id=user.id,
    )
    db.add(transfer)
    db.flush()
    if old_assignment:
        old_assignment.status = "ended"
        old_assignment.ends_on = effective_date
    db.add(GroupLeaderAssignment(group_id=group.id, leader_employee_id=new_leader.id, starts_on=effective_date, status="active", transfer_id=transfer.id))
    for member in members:
        db.add(GroupTransferMember(transfer_id=transfer.id, employee_id=member.employee_id, employee_no=member.employee.employee_no, employee_name=member.employee.name))
    member_ids = [member.employee_id for member in members]
    if member_ids:
        db.query(RecognitionRecord).filter(RecognitionRecord.employee_id.in_(member_ids), RecognitionRecord.status == "pending").update({RecognitionRecord.assigned_reviewer_id: new_leader.id}, synchronize_session=False)
    group.status = "active"
    group.revision += 1
    db.query(SystemAlert).filter(SystemAlert.group_id == group.id, SystemAlert.status == "open").update({SystemAlert.status: "handled", SystemAlert.handled_by: user.id, SystemAlert.handled_at: datetime.now()}, synchronize_session=False)
    write_audit(db, user.employee, "整组移交", "work_group", group.id, before={"leader": old_assignment.leader.name if old_assignment else "待接管"}, after={"leader": new_leader.name, "member_count": len(members)}, reason=reason, ip_address=client_ip(request))
    db.commit()
    invalidate_data_caches()
    return {"ok": True, "transfer_id": transfer.id}


@router.get("/hr/alerts")
def hr_alerts(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    process_role_expirations(db)
    allowed_attractions = scoped_hr_attraction_ids(db, user)
    rows = db.query(SystemAlert).order_by(SystemAlert.status.asc(), SystemAlert.created_at.desc()).limit(200).all()
    if allowed_attractions is not None:
        group_ids = {
            group.id
            for group in db.query(WorkGroup).filter(WorkGroup.attraction_id.in_(allowed_attractions)).all()
        }
        employee_ids = {
            employee.id
            for employee in db.query(Employee).filter(Employee.attraction_id.in_(allowed_attractions)).all()
        }
        rows = [row for row in rows if (row.group_id in group_ids if row.group_id else row.employee_id in employee_ids)]
    return [
        {"id": row.id, "type": row.alert_type, "message": row.message, "due_date": row.due_date or "", "status": row.status, "created_at": row.created_at.strftime("%Y-%m-%d %H:%M:%S")}
        for row in rows
    ]


@router.get("/hr/score-rules")
def hr_score_rules(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    result = []
    for role_code in RECOGNIZER_CODES:
        role = db.query(Role).filter(Role.code == role_code).one()
        rule = db.query(RecognitionScoreRule).filter_by(role_id=role.id, active=True).order_by(RecognitionScoreRule.effective_date.desc()).first()
        result.append({"role_id": role.id, "role_code": role.code, "role_name": role.name, "score": float(rule.score) if rule else 0, "effective_date": rule.effective_date if rule else ""})
    return sorted(result, key=lambda item: item["role_id"])


@router.post("/hr/score-rules")
def update_score_rule(payload: dict, request: Request, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("HR_MANAGE"))):
    role = db.query(Role).filter(Role.code == str(payload.get("role_code") or "")).first()
    if not role or role.code not in RECOGNIZER_CODES:
        raise HTTPException(400, "认可人角色无效")
    if user.role.code == SCOPED_HR_ROLE_CODE and role.code not in CIRCLE_HR_SCORE_RULE_ROLE_CODES:
        raise HTTPException(403, "景点圈HR仅可调整本圈TA主管/主管的分值；GSM及以上角色分值由最高管理员统一调整")
    try:
        score = Decimal(str(payload.get("score"))).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise HTTPException(400, "分值格式错误") from exc
    effective_date = str(payload.get("effective_date") or date.today().isoformat())
    parse_iso_date(effective_date, "生效日期")
    rule = db.query(RecognitionScoreRule).filter_by(role_id=role.id, effective_date=effective_date).first()
    before = {"score": float(rule.score), "active": rule.active} if rule else None
    if rule:
        rule.score = score
        rule.active = True
    else:
        rule = RecognitionScoreRule(role_id=role.id, score=score, effective_date=effective_date, active=True)
        db.add(rule)
    write_audit(db, user.employee, "修改认可角色分值", "recognition_score_rule", role.id, before=before, after={"role": role.name, "score": float(score), "effective_date": effective_date}, ip_address=client_ip(request))
    db.commit()
    return {"ok": True}


@router.get("/admin/logs")
def admin_logs(limit: int = 300, db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SYSTEM_ADMIN"))):
    return [
        {"id": row.id, "time": row.created_at.strftime("%Y-%m-%d %H:%M:%S"), "operator": row.operator_name or "系统", "action": row.action, "entity": f"{row.entity_type}:{row.entity_id or ''}", "reason": row.reason or ""}
        for row in db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(min(limit, 1000)).all()
    ]


@router.get("/admin/circle-hr-accounts")
def circle_hr_accounts(db: Session = Depends(get_db), user: V2User = Depends(require_permissions("SYSTEM_ADMIN"))):
    """Show the three scoped HR accounts without exposing stored passwords."""
    login_accounts = [item[0] for item in CIRCLE_HR_ACCOUNTS]
    employees = {
        employee.employee_no: employee
        for employee in db.query(Employee).filter(Employee.employee_no.in_(login_accounts)).all()
    }
    employee_ids = [employee.id for employee in employees.values()]
    accounts = {
        account.employee_id: account
        for account in db.query(UserAccount).filter(UserAccount.employee_id.in_(employee_ids)).all()
    } if employee_ids else {}
    items = []
    for login_account, name, circle_name in CIRCLE_HR_ACCOUNTS:
        employee = employees.get(login_account)
        account = accounts.get(employee.id) if employee else None
        items.append(
            {
                "employee_id": employee.id if employee else None,
                "login_account": login_account,
                "name": name,
                "attraction_name": circle_name,
                "account_enabled": bool(account and account.enabled),
                "must_change_password": bool(account.must_change_password) if account else True,
                "last_login_at": account.last_login_at.strftime("%Y-%m-%d %H:%M:%S") if account and account.last_login_at else "",
                "password_status": "待首次修改" if not account or account.must_change_password else "已设置（不可读取）",
            }
        )
    return {"items": items}


@router.post("/admin/circle-hr-accounts/{employee_id}/reset-password")
def reset_circle_hr_password(
    employee_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: V2User = Depends(require_permissions("SYSTEM_ADMIN")),
):
    employee = db.get(Employee, employee_id)
    role = role_at(db, employee.id) if employee else None
    if not employee or not role or role.code != "HR_CIRCLE":
        raise HTTPException(404, "景点圈HR账号不存在")
    account = db.query(UserAccount).filter(UserAccount.employee_id == employee.id).first()
    if not account:
        raise HTTPException(400, "该景点圈HR尚未开通登录账号")
    temporary_password = new_temporary_password()
    account.password_hash = hash_password(temporary_password)
    account.failed_attempts = 0
    account.locked_until = None
    account.must_change_password = True
    account.credential_initialized = True
    account.password_changed_at = None
    account.enabled = True
    db.query(UserSession).filter(UserSession.account_id == account.id).delete(synchronize_session=False)
    write_audit(
        db,
        user.employee,
        "重置景点圈HR密码",
        "user_account",
        account.id,
        after={
            "employee_no": employee.employee_no,
            "employee_name": employee.name,
            "attraction_name": employee.attraction.name if employee.attraction else "",
            "password_rule": "管理员生成的一次性临时密码",
            "must_change_password": True,
        },
        reason="最高管理员重置景点圈HR账号密码",
        ip_address=client_ip(request),
    )
    db.commit()
    return {
        "ok": True,
        "employee_no": employee.employee_no,
        "employee_name": employee.name,
        "temporary_password": temporary_password,
        "must_change_password": True,
        "message": "临时密码仅显示本次，请立即转交账号本人并要求登录后修改。",
    }
